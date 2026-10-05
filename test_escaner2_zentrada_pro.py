# -*- coding: utf-8 -*-
"""Banco del ESCANEO PRO DE ZENTRADA (encargo de la tarjeta de Zentrada, 05-oct-2026): la parte pura y los DOS programas
de verdad, con dobles. Sin red, sin secretos, sin produccion. 🔴 Repo PUBLICO: el Excel de Zentrada de aqui es
INVENTADO (la misma forma que el de Cowork, con precios redondos y EAN de mentira con su digito de control bueno).

PARTE A · LO PURO (escaner2_zentrada.py), en este proceso:
  A1  las guardas del Excel: el bueno pasa; cada guarda, rota a proposito, lo para y dice por que (hojas, cabeceras,
      «Leído» de mas de 2 dias o en el futuro, menos de 100 ofertas, llave repetida, precio 0 en una que se puede
      pedir, un sí/no que no lo es, un origen que no existe, una ficha mal);
  A2  los EAN: sin ceros a la izquierda; EAN-8 y UPC sin su cero caben en 13; el digito de control malo, fuera;
  A3  la foto: una fila por EAN con su oferta mas barata que se puede pedir (a igual precio, la de Mix), crudo
      (ofertas) = previas + foto, y lo de la compra (donde, el precio de Mix, el origen junto, el mejor puesto);
  A4  la hoja «Fichas»: las de EAN que se pueden pedir a la lista de ASIN; las nuestras sin oferta, a su hoja;
  A5  la puerta con un pack de «Fichas»: coste = unidades × precio, la misma ficha, y el detalle lo dice.

PARTE B · LOS PROGRAMAS (escaner2_zentrada_barrido.py y escaner2_zentrada_cruce.py), cada uno en su proceso, contra un
`supabase` en memoria:
  B1  [sin_llave] sin la llave de servicio no nace ningun cliente;
  B2  [bueno] el Excel en zentrada/<pasada>/ → barrido → 'esperando_csv' con su foto, eans.txt, asins.txt y
      barrido.json → los dos CSV de España (EAN y ASIN, con la CABECERA REAL del Visualizador) y uno de FR que se
      ignora → cruce → 'lista', cuadra, solo ES, y el Excel: las hojas del PRO, el UPC casado sin su cero, la oferta
      mas barata, las fichas de «Fichas» con su coste, las columnas de la compra al final, los COMPRAR por origen y
      solo del top 300, y la ultima hoja de nuestros agotados;
  B3  🔴 el registro (repo PUBLICO) no lleva ni un EAN, ASIN, precio ni nombre;
  B4  nada escribe fuera de escaner2_* ni del almacen zentrada/<pasada>/; productos acaba como empezo;
  B5  🔴 [viejo] un Excel leido hace mas de 2 dias: ROJO y NO se crea la pasada; [repetida] una pasada que ya existe:
      ROJO y no se toca.
"""
import base64
import csv
import io
import json
import os
import subprocess
import sys
import tempfile
import types
from datetime import datetime, timedelta, timezone

AQUI = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, AQUI)
fallos = []


def eq(nombre, obtenido, esperado):
    ok = obtenido == esperado
    if not ok:
        fallos.append(nombre)
    print(('OK ' if ok else 'XX ') + nombre + ('' if ok else '   got=%r exp=%r' % (obtenido, esperado)))


# ═══════════════════════════════════════════════════════════════════════════════
# EL DOBLE DE SUPABASE (el de test_escaner2_osma_pro.py, copiado: cada banco, el suyo)
# ═══════════════════════════════════════════════════════════════════════════════
class _Resp:
    def __init__(self, data, count=None):
        self.data, self.count = data, count


class _Consulta:
    def __init__(self, bd, tabla):
        self.bd, self.tabla = bd, tabla
        self.op, self.filtros, self.payload = 'select', [], None
        self.cuenta, self.orden, self.rango, self.tope = False, None, None, None

    def select(self, _cols='*', count=None):
        self.op, self.cuenta = 'select', (count == 'exact')
        return self

    def insert(self, filas):
        self.op, self.payload = 'insert', filas if isinstance(filas, list) else [filas]
        return self

    def update(self, cambios):
        self.op, self.payload = 'update', cambios
        return self

    def eq(self, k, v):
        self.filtros.append((k, v))
        return self

    def order(self, col, desc=False):
        self.orden = (col, desc)
        return self

    def range(self, a, b):
        self.rango = (a, b)
        return self

    def limit(self, n):
        self.tope = n
        return self

    def execute(self):
        self.bd.setdefault('ops', []).append([self.bd.get('programa'), self.op, self.tabla])
        filas = self.bd['tablas'].setdefault(self.tabla, [])
        if self.op == 'insert':
            for f in self.payload:
                filas.append(json.loads(json.dumps(f, default=str)))
            return _Resp(self.payload)
        sel = [f for f in filas if all(str(f.get(k)) == str(v) for k, v in self.filtros)]
        if self.op == 'update':
            for f in sel:
                f.update(json.loads(json.dumps(self.payload, default=str)))
            return _Resp(sel)
        if self.orden:
            sel = sorted(sel, key=lambda f: (f.get(self.orden[0]) is None, str(f.get(self.orden[0]))),
                         reverse=self.orden[1])
        n = len(sel)
        if self.rango:
            sel = sel[self.rango[0]:self.rango[1] + 1]
        if self.tope is not None:
            sel = sel[:self.tope]
        return _Resp(sel, n if self.cuenta else None)


class _Cubo:
    def __init__(self, bd, cubo):
        self.bd, self.cubo = bd, cubo

    def _objs(self):
        return self.bd['storage'].setdefault(self.cubo, {})

    def upload(self, ruta, datos, _opciones=None):
        self.bd.setdefault('subidas', []).append(ruta)
        self._objs()[ruta] = base64.b64encode(datos if isinstance(datos, bytes) else datos.encode()).decode()
        return {'Key': ruta}

    def download(self, ruta):
        if ruta not in self._objs():
            raise RuntimeError('404 Object not found: %s/%s' % (self.cubo, ruta))
        return base64.b64decode(self._objs()[ruta])

    def list(self, carpeta=None, opciones=None):
        op = dict({'limit': 100, 'offset': 0}, **(opciones or {}))
        pref = (carpeta or '').rstrip('/') + '/'
        nombres = sorted(r[len(pref):] for r in self._objs() if r.startswith(pref) and '/' not in r[len(pref):])
        trozo = nombres[op['offset']:op['offset'] + op['limit']]
        return [{'name': n, 'created_at': '2026-10-05T09:%02d:00Z' % i, 'metadata': {'size': 1}}
                for i, n in enumerate(trozo)]


class _Cliente:
    def __init__(self, bd):
        self.bd = bd
        self.storage = types.SimpleNamespace(from_=lambda cubo: _Cubo(bd, cubo))

    def table(self, nombre):
        return _Consulta(self.bd, nombre)


def hijo(ruta_estado, programa):
    with open(ruta_estado, encoding='utf-8') as fh:
        bd = json.load(fh)
    bd['programa'] = programa
    clientes = []

    def crear(url, llave):
        clientes.append(llave)
        return _Cliente(bd)

    sys.modules['supabase'] = types.ModuleType('supabase')
    sys.modules['supabase'].create_client = crear
    import runpy
    sys.argv = [programa]
    codigo = 0
    try:
        runpy.run_path(os.path.join(AQUI, programa), run_name='__main__')
    except SystemExit as ex:
        codigo = ex.code if isinstance(ex.code, int) else (0 if ex.code is None else 1)
    finally:
        with open(ruta_estado, 'w', encoding='utf-8') as fh:
            json.dump(bd, fh, default=str)
        print('CLIENTES_CREADOS=%d' % len(clientes))
    sys.exit(codigo)


if __name__ == '__main__' and len(sys.argv) >= 4 and sys.argv[1] == '--hijo':
    hijo(sys.argv[2], sys.argv[3])


def correr(ruta_estado, programa, env_extra):
    env = {k: v for k, v in os.environ.items() if k not in ('SUPABASE_SERVICE_KEY', 'PASADA', 'GITHUB_RUN_ID')}
    env.update(env_extra, PYTHONIOENCODING='utf-8', SUPABASE_URL='https://doble.invalid')
    p = subprocess.run([sys.executable, '-u', os.path.abspath(__file__), '--hijo', ruta_estado, programa],
                       capture_output=True, text=True, encoding='utf-8', errors='replace', env=env, cwd=AQUI)
    return p.returncode, (p.stdout or '') + (p.stderr or '')


import escaner2_motor as e2  # noqa: E402
import escaner2_zentrada as ez  # noqa: E402
import escaner2_heredado_pro as pro  # noqa: E402
from openpyxl import Workbook, load_workbook  # noqa: E402

M = e2.cargar_motor()


def ean13(cuerpo):
    return cuerpo + M._chk13(cuerpo)


def upc12(cuerpo11):
    return cuerpo11 + M._chk13('0' + cuerpo11)


# ═══════════════════════════════════════════════════════════════════════════════
# EL EXCEL DE ZENTRADA, INVENTADO (la forma del de Cowork del 05-oct-2026)
# ═══════════════════════════════════════════════════════════════════════════════
E_UNO = ean13('840000000001')           # dos ofertas: la mas barata fuera de Mix, otra en Mix
E_UPC = upc12('07000000002')            # Zentrada lo escribe SIN el cero; Keepa, con él
E_UPC_ZEN = E_UPC.lstrip('0')
E_COMP = ean13('840000000003')          # de la competencia, VALORAR
E_FICHA = ean13('840000000004')         # con dos fichas en «Fichas» (la suelta y un pack de 2)
E_PACK = ean13('840000000006')          # Keepa cae por EAN en el pack de 3 de «Fichas»
E_AGOT = ean13('840000000007')          # nuestro, su única oferta no se puede pedir
E_SIN = ean13('840000000008')           # nuestro, en «Fichas» y sin ninguna oferta
E_MALO = ean13('840000000010')[:-1] + str((int(ean13('840000000010')[-1]) + 1) % 10)
E_DUP = upc12('07000000011')            # dos ofertas al mismo precio, escritas con y sin el cero
RELLENO = [ean13('8410000%05d' % i) for i in range(95)]

ASIN_UNO, ASIN_UPC, ASIN_COMP = 'B0ZUNO0001', 'B0ZUPC0002', 'B0ZCOM0003'
ASIN_PACK2, ASIN_SUELTA, ASIN_PACK3 = 'B0ZPCK0004', 'B0ZSUE0005', 'B0ZPCK0006'
ASIN_AGOT, ASIN_SIN = 'B0ZAGO0007', 'B0ZSIN0008'
ASIN_AM24, ASIN_AM1 = 'B0ZAMZ0024', 'B0ZAMZ0001'


def leido_madrid(delta=timedelta(0)):
    from zoneinfo import ZoneInfo
    return (datetime.now(timezone.utc) - delta).astimezone(ZoneInfo('Europe/Madrid')).strftime('%Y-%m-%d %H:%M')


def oferta(ean, may, art, precio, mix=True, pedir=True, origen='top 300', puesto=None, alim=False, pelig=False,
           nombre=None, leido=None):
    sn = lambda b: 'sí' if b else 'no'  # noqa: E731
    return [ean, art, nombre or 'Producto %s' % art, may, sn(mix), sn(pedir), precio, 6, 1, 150.0, None, sn(alim),
            sn(pelig), puesto, origen, leido]


def ofertas_buenas(leido):
    filas = [oferta(E_UNO, 'Mayorista Alfa', 'A-1', 2.0, mix=False, origen='nuestro, top 300', puesto=40),
             oferta(E_UNO, 'Mayorista Beta', 'B-1', 2.5, origen='nuestro, top 300', puesto=12),
             oferta(E_UPC_ZEN, 'Mayorista Gamma', 'G-2', 3.0),
             oferta(E_COMP, 'Mayorista Alfa', 'A-3', 4.8, origen='competidor', alim=True, pelig=True),
             oferta(E_FICHA, 'Mayorista Delta', 'D-4', 1.0, origen='competidor, top 300'),
             oferta(E_PACK, 'Mayorista Delta', 'D-6', 1.5, origen='competidor'),
             oferta(E_AGOT, 'Mayorista Beta', 'B-7', 5.0, pedir=False, origen='nuestro'),
             oferta(E_MALO, 'Mayorista Beta', 'B-10', 1.0),
             oferta('', 'Mayorista Beta', 'B-0', 1.0),
             oferta(E_DUP, 'Mayorista Gamma', 'G-11', 6.0, mix=False),
             oferta(E_DUP.lstrip('0'), 'Mayorista Delta', 'D-11', 6.0)]
    # (AM) Dos de relleno son golosinas de 45 g: Keepa las casa con un multipack (ver los CSV de B2).
    filas += [oferta(e, 'Mayorista Relleno', 'R-%d' % i, 1.0, nombre=('Golosina R-%d 45g' % i) if i < 2 else None)
              for i, e in enumerate(RELLENO)]
    return [f[:-1] + [leido] for f in filas]


FICHAS = [[E_FICHA, ASIN_PACK2, 2, 'competidor', 'Producto D-4 pack de 2'],
          [E_FICHA, ASIN_SUELTA, 1, 'competidor', 'Producto D-4'],
          [E_PACK, ASIN_PACK3, 3, 'competidor', 'Producto D-6 pack de 3'],
          [E_AGOT, ASIN_AGOT, 1, 'nuestro', 'Producto B-7'],
          [E_SIN, ASIN_SIN, 1, 'nuestro', 'Producto sin oferta']]


def excel(ofertas=None, fichas=None, hojas=None, cab_ofertas=None, leido=None):
    wb = Workbook()
    wb.remove(wb.active)
    for h in hojas or ez.HOJAS:
        ws = wb.create_sheet(h)
        if h == 'Ofertas':
            ws.append(cab_ofertas or ez.CABECERA_OFERTAS)
            for f in (ofertas if ofertas is not None else ofertas_buenas(leido or leido_madrid())):
                ws.append(f)
        elif h == 'Fichas':
            ws.append(ez.CABECERA_FICHAS)
            for f in (fichas if fichas is not None else FICHAS):
                ws.append(f)
        else:
            ws.append(ez.CABECERA_RESUMEN)
            ws.append(['Ofertas en la hoja', 106])
    b = io.BytesIO()
    wb.save(b)
    return b.getvalue()


AHORA = datetime.now(timezone.utc)

# ═══════════════════════════════════════════════════════════════════════════════
# PARTE A · LO PURO
# ═══════════════════════════════════════════════════════════════════════════════
print('A1 · las guardas del Excel')
_x = ez.leer_excel(excel(), AHORA)
eq('A1 · el Excel bueno pasa: 106 ofertas, 5 fichas, la hoja Resumen leída', (len(_x['ofertas']), len(_x['fichas']),
                                                                             _x['resumen']),
   (106, 5, [('Ofertas en la hoja', 106)]))


def motivo(contenido):
    try:
        ez.leer_excel(contenido, AHORA)
    except ez.FalloZentrada as ex:
        return str(ex)
    return None


_b = ofertas_buenas(leido_madrid())
_casos = [
    ('sin la hoja «Fichas»', excel(hojas=['Ofertas', 'Resumen']), 'las hojas tienen que ser'),
    ('las hojas en otro orden', excel(hojas=['Fichas', 'Ofertas', 'Resumen']), 'en ese orden'),
    ('una cabecera cambiada', excel(cab_ofertas=ez.CABECERA_OFERTAS[:6] + ['Precio'] + ez.CABECERA_OFERTAS[7:]),
     'la cabecera de «Ofertas» no es la esperada'),
    ('«Leído» de hace 3 días', excel(leido=leido_madrid(timedelta(days=3))), 'hace más de 2 días'),
    ('«Leído» en el futuro', excel(leido=leido_madrid(-timedelta(hours=2))), 'en el futuro'),
    ('«Leído» sin la forma', excel(leido='05/10/2026 11:16'), 'no tiene la forma AAAA-MM-DD HH:MM'),
    ('99 ofertas', excel(ofertas=_b[:99]), 'trae 99 ofertas y hacen falta 100'),
    ('una llave repetida', excel(ofertas=_b + [_b[-1]]), 'la llave Mayorista + Artículo Zentrada ya está'),
    ('🔴 precio 0 en una que se puede pedir', excel(ofertas=[_b[0][:6] + [0] + _b[0][7:]] + _b[1:]),
     'se puede pedir y no tiene precio'),
    ('🔴 precio vacío en una que se puede pedir', excel(ofertas=[_b[0][:6] + [None] + _b[0][7:]] + _b[1:]),
     'se puede pedir y no tiene precio'),
    ('un sí/no que no lo es', excel(ofertas=[_b[0][:5] + ['quizá'] + _b[0][6:]] + _b[1:]),
     '«Se puede pedir» tiene que ser sí o no'),
    ('un origen que no existe', excel(ofertas=[_b[0][:14] + ['proveedor'] + _b[0][15:]] + _b[1:]), '«Origen» tiene que ser'),
    ('una ficha con 0 unidades', excel(fichas=[FICHAS[0][:2] + [0] + FICHAS[0][3:]]), 'un entero de 1 o más'),
    ('una ficha repetida', excel(fichas=FICHAS + [FICHAS[0]]), 'una ficha, una fila'),
    ('un fichero que no es Excel', b'no soy un xlsx', 'no se puede abrir como Excel'),
]
for _n, _c, _t in _casos:
    _m = motivo(_c)
    eq('A1 · %s → no pasa, y dice por qué' % _n, (_m is not None, _t in (_m or '')), (True, True))
eq('A1 · una oferta que NO se puede pedir puede ir sin precio',
   motivo(excel(ofertas=_b + [oferta(ean13('842000000001'), 'Mayorista Beta', 'B-99', None, pedir=False,
                                     leido=leido_madrid())])), None)
eq('A1 · «si» sin tilde y «SÍ» en mayúsculas valen', motivo(excel(ofertas=[_b[0][:4] + ['SÍ', 'si'] + _b[0][6:]]
                                                                   + _b[1:])), None)

print('A2 · los EAN')
eq('A2 · EAN-13, UPC sin su cero, EAN-8 y un EAN con su cero delante → 13 cifras, sin cambiar el número',
   [ez.core_de(x)[0] for x in (E_UNO, E_UPC_ZEN, '96385074', '0' + E_UPC)],
   [E_UNO, E_UPC.zfill(13), '96385074'.zfill(13), E_UPC.zfill(13)])
eq('A2 · el dígito de control malo, letras o 15 cifras → EAN con forma rara',
   [ez.core_de(x)[0] for x in (E_MALO, '84000A0000001', '123456789012345')], [None, None, None])
eq('A2 · sin ceros a la izquierda: el UPC con y sin el cero es el mismo', ez.ean_norm(E_UPC) == ez.ean_norm(E_UPC_ZEN),
   True)

print('A3 · la foto')
foto, apartados, cuentas, compra = ez.construir_foto(_x['ofertas'], M)
eq('A3 · 🔴 crudo 106 ofertas = previas 5 (1 no se puede pedir, 1 sin EAN, 1 EAN raro, 2 duplicadas) + foto 101',
   (cuentas['n_crudo'], cuentas['previas']['no_disponible'], cuentas['previas']['sin_gtin'],
    cuentas['previas']['ean_forma_rara'], cuentas['previas']['duplicado_proveedor'], cuentas['n_foto'],
    cuentas['cuadra_previo']), (106, 1, 1, 1, 2, 101, True))
_f = {ez.ean_norm(f['ean_core']): f for f in foto}
eq('A3 · E_UNO: la más barata (2,00, fuera de Mix), no la de Mix (2,50), tal cual: sin porte ni cuota',
   (_f[ez.ean_norm(E_UNO)]['precio_unidad'], _f[ez.ean_norm(E_UNO)]['producto_heo']), (2.0, 'Mayorista Alfa · A-1'))
eq('A3 · …y lo de la compra: fuera de Mix, el precio más barato en Mix, el origen junto y el mejor puesto',
   {k: compra[ez.ean_norm(E_UNO)][k] for k in ('mix', 'precio_mix', 'origen', 'puesto', 'alimentacion')},
   {'mix': False, 'precio_mix': 2.5, 'origen': ['nuestro', 'top 300'], 'puesto': 12, 'alimentacion': False})
eq('A3 · E_DUP: al mismo precio gana la de Mix, aunque una venga con el cero y otra sin él',
   (_f[ez.ean_norm(E_DUP)]['producto_heo'], compra[ez.ean_norm(E_DUP)]['precio_mix']), ('Mayorista Delta · D-11', None))
eq('A3 · el UPC sin su cero entra con su número de Zentrada en la lista, y 13 cifras en la foto',
   (_f[ez.ean_norm(E_UPC)]['ean_original'], _f[ez.ean_norm(E_UPC)]['ean_core'], E_UPC_ZEN in ez.lista_para_keepa(foto)),
   (E_UPC_ZEN, E_UPC.zfill(13), True))
eq('A3 · la lista: uno por EAN de la foto (101)', len(ez.lista_para_keepa(foto)), 101)
_e1 = ean13('100000000001')
_f1 = ez.construir_foto([dict(_x['ofertas'][2], ean=_e1)], M)[0][0]
eq('A3 · 🔴 un EAN de 13 cifras que empieza por 1 es un EAN: sin el rescate de GTIN de HEO, la lista lleva solo ese '
   '(con el Excel real del 05-oct-2026 colgaba otro código)', (_f1['codigos_keepa'], _f1['variantes']),
   ([_e1], [_e1]))
eq('A3 · las puertas previas listadas: el EAN raro y las dos ofertas más caras de su EAN',
   sorted((a['motivo'], a['producto_heo']) for a in apartados),
   [('duplicado_proveedor', 'Mayorista Beta · B-1'), ('duplicado_proveedor', 'Mayorista Gamma · G-11'),
    ('ean_forma_rara', 'Mayorista Beta · B-10')])
eq('A3 · la de la competencia: alimentación y mercancía peligrosa, solo para informar',
   (compra[ez.ean_norm(E_COMP)]['alimentacion'], compra[ez.ean_norm(E_COMP)]['peligrosa']), (True, True))

print('A4 · la hoja «Fichas»')
disp, agot = ez.mapa_de_fichas(_x['fichas'], foto)
eq('A4 · las de un EAN que se puede pedir, a la lista de ASIN; las demás, aparte',
   (ez.lista_asin(disp), [m['asin'] for m in agot]), ([ASIN_PACK2, ASIN_SUELTA, ASIN_PACK3], [ASIN_AGOT, ASIN_SIN]))
_ag = ez.agotados_nuestros(agot, ez.precios_sin_pedir(_x['ofertas']), {})
eq('A4 · nuestros agotados: el que no se puede pedir (con su precio) y el que no tiene oferta',
   [(r[1], r[5], r[6]) for r in _ag], [(ASIN_AGOT, 5.0, ez.TEXTO_SIN_PEDIR), (ASIN_SIN, None, ez.TEXTO_SIN_OFERTA)])
eq('A4 · una de la competencia sin oferta no va a esa hoja; una nuestra por productos sí',
   [len(ez.agotados_nuestros([dict(agot[0], origen=['competidor'])], {}, {})),
    len(ez.agotados_nuestros([dict(agot[0], origen=['competidor'])], {}, {ASIN_AGOT: {'ean': E_AGOT, 'nombre': 'x'}}))],
   [0, 1])
eq('A4 · el coste de una ficha de «Fichas»: unidades × precio, un redondeo', [ez.pa_mapa(1.0, 2), ez.pa_mapa(1.255, 2)],
   [2.0, 2.51])

print('A5 · la puerta con un pack de «Fichas»')
_params = {'umbral': 6, 'paises_filtro': ['ES'], 'paises_calculo': ['ES']}
_rec = {'asin': ASIN_PACK3, 'titulo': 'Producto D-6 pack de 3', 'rank': 9000, 'rank90': 11000, 'buybox': 12.95,
        'es_fba': True, 'nuevo': 12.95, 'fba': 3.10, 'compct': 15.01}
_fp = dict(_f[ez.ean_norm(E_PACK)], id='x')
_r1 = ez.decidir_zentrada(_fp, {'ES': [_rec]}, {'ES': {ASIN_PACK3: 30}}, _params, M, None, {}, {ASIN_PACK3: 3})
_r0 = ez.decidir_zentrada(_fp, {'ES': [_rec]}, {'ES': {ASIN_PACK3: 30}}, _params, M, None, {}, None)
eq('A5 · 🔴 con el pack de 3: coste 3 × 1,50 = 4,50 → VALORAR (sin él, 1,50 → COMPRAR), y el detalle lo dice',
   (_r1['factor'], _r1['pa'], _r1['paises']['ES']['pa'], _r1['puerta'], _r0['puerta'],
    'pack de 3 de «Fichas»: coste 3 × 1,500 € = 4,50 €' in _r1['detalle']), (3, 4.5, 4.5, 'e', 'f', True))

print('A6 · (AM) los packs de Amazon, con la regla de OSMA copiada')
_fg = dict(_f[ez.ean_norm(RELLENO[0])], id='g', nombre='Golosina R-0 45g')
_rg = dict(_rec, asin='B0ZAMZ0024', titulo='Golosinas 24 x 45 g')


def _sen(**k):
    return {'ES': {'B0ZAMZ0024': dict({'n_art': '', 'valor_ud': '', 'tipo_ud': '', 'paquete': '', 'tamano': '',
                                       'titulo': 'Golosinas 24 x 45 g'}, **k)}}


def _dec(sen, mapa=None):
    return ez.decidir_zentrada(_fg, {'ES': [_rg]}, {'ES': {'B0ZAMZ0024': 30}}, _params, M, None, {}, mapa, sen)


_a = _dec(_sen(tamano='45 g (Paquete de 24)'))
eq('A6 · 🔴 tamaño «45 g (Paquete de 24)» y título «24 x 45 g» (dos señales): coste 24 × 1,00 = 24,00, NO COMPRAR; sin el '
   'AM salía COMPRAR con el coste de una', (_a['pack'], _a['factor'], _a['pa'], _a['puerta'], _dec(None)['puerta']),
   ('amazon', 24, 24.0, 'd', 'f'))
_b = _dec(_sen(tamano='45 g (Paquete de 24)', titulo='Golosinas'))
eq('A6 · 🔴 con UNA sola señal: no se multiplica, y de COMPRAR baja a VALORAR con «posible pack en Amazon: revisar»',
   (_b['factor'], _b['pa'], _b['puerta'], _b['pack_amazon']['estado'], ez.TEXTO_POSIBLE_PACK in _b['detalle']),
   (1, 1.0, 'e', ez.PACK_DUDOSO, True))
_c = _dec(_sen(tamano='45 g (Paquete de 24)'), {'B0ZAMZ0024': 1})
_d = _dec(_sen(tamano='45 g (Paquete de 24)'), {'B0ZAMZ0024': 3})
eq('A6 · 🔴 con el ASIN en «Fichas», mandan sus unidades y Amazon ni se mira (1 → 1,00; 3 → 3,00)',
   [(_c['factor'], _c['pa'], _c['pack_amazon']), (_d['factor'], _d['pa'], _d['pack'], _d['pack_amazon'])],
   [(1, 1.0, None), (3, 3.0, 'fichas', None)])
_fh = dict(_fg, nombre='Golosina R-0 48g')
_h = ez.decidir_zentrada(_fh, {'ES': [_rg]}, {'ES': {'B0ZAMZ0024': 30}}, _params, M, None, {}, None,
                         _sen(tamano='45 g (Paquete de 24)', valor_ud='1080', tipo_ud='gramo', titulo='Golosinas'))
eq('A6 · 🔴 (2 bis, solo Zentrada; el Haribo real) «48g» en Zentrada y 1.080 g en Amazon: dos señales que dan 22,5, que no '
   'es entero → no se multiplica, pero nunca COMPRAR: VALORAR con «posible pack en Amazon: revisar»',
   (_h['factor'], _h['pa'], _h['puerta'], _h['pack_amazon']['estado'], 'tamaño «45 g (Paquete de 24)» → 22,5' in _h['detalle']),
   (1, 1.0, 'e', ez.PACK_DUDOSO, True))
_n = ez.factor_pack_amazon({'n_art': '', 'valor_ud': '50', 'tipo_ud': 'gramo', 'paquete': '', 'tamano': '', 'titulo': ''},
                           ez.cantidad_zentrada('Golosina 48g'))
eq('A6 · …y 50 g contra 48 g (1,04: el mismo artículo, con otra medida) sigue sin ser pack', _n['estado'], ez.PACK_NO)
eq('A6 · el nombre de Zentrada: «45g» es la medida; «24er» un recuento; sin nada, una unidad',
   [ez.cantidad_zentrada('Golosina R-0 45g')['medida'], ez.cantidad_zentrada('Gummibärchen 24er')['n'],
    ez.cantidad_zentrada('Producto R-9')['texto']], [(45.0, 'g'), 24, 'sin cantidad: 1 unidad'])


# ═══════════════════════════════════════════════════════════════════════════════
# PARTE B · LOS DOS PROGRAMAS, DE PUNTA A PUNTA
# ═══════════════════════════════════════════════════════════════════════════════
with io.open(os.path.join(AQUI, 'test_escaner2_osma_cabecera.json'), encoding='utf-8') as fh:
    CABECERA = json.load(fh)['cabecera']
COL_PAIS, COL_CAIDAS = e2.columnas_keepa()
COL_PADRE, COL_NVAR, COL_RANK = e2.columnas_ficha_compartida()
C = pro.CSV_COLS


def csv_real(pais, filas):
    """Un export del Visualizador con la CABECERA REAL (561 columnas, la del banco de OSMA) y solo lo que se lee."""
    buf = io.StringIO()
    w = csv.writer(buf)
    w.writerow(CABECERA)
    ix = {h: i for i, h in enumerate(CABECERA)}
    for asin, eans, caidas, precio, titulo, *extra in filas:
        fila = [''] * len(CABECERA)
        for h, v in ((C['asin'], asin), (COL_PAIS, pais), ('Título', titulo), (C['ean'], eans),
                     (C['rank'], '9000'), (C['rank90'], '11000'), (COL_CAIDAS, caidas), (C['buybox'], precio),
                     (C['es_fba'], 'yes'), (C['nuevo'], precio), (C['fba'], '3.10'), (C['compct'], '15.01 %'),
                     (C['nvar'], ''), (COL_PADRE, '')) + tuple((extra[0] if extra else {}).items()):
            fila[ix[h]] = v
        w.writerow(fila)
    return ('\ufeff' + buf.getvalue()).encode('utf-8')


PASADA = '00000000-0000-4000-8000-00000000c0de'
NOMBRE_EXCEL = 'Zentrada_2026-10-05.xlsx'
PRODUCTOS = [{'id': 1, 'ean': E_UNO, 'asin': ASIN_UNO, 'nombre': 'Producto Uno de Moloka', 'activo': True,
              'es_chase': False, 'iva_pct': 0.21, 'stock_moloka': 3, 'unidades_por_pack': None},
             {'id': 2, 'ean': E_UNO, 'asin': None, 'nombre': 'Producto Uno chase', 'activo': True, 'es_chase': True,
              'iva_pct': 0.21, 'stock_moloka': 1, 'unidades_por_pack': None}]
SEMBRADAS = ('productos', 'inventario_fba')


def estado_inicial(escena):
    leido = leido_madrid(timedelta(days=3)) if escena == 'viejo' else leido_madrid()
    pasadas = []
    if escena == 'repetida':
        pasadas = [{'id': PASADA, 'proveedor': 'ZENTRADA', 'estado': 'esperando_csv', 'run_id': 1}]
    return {'tablas': {
        'escaner2_parametros': [{'proveedor': 'HEO', 'umbral_caidas_30d': 6, 'paises_filtro': ['ES', 'IT', 'FR', 'DE'],
                                 'paises_calculo': ['ES', 'IT', 'FR', 'DE']},
                                {'proveedor': 'OSMA', 'umbral_caidas_30d': 6, 'paises_filtro': ['ES'],
                                 'paises_calculo': ['ES']},
                                # La fila de Zentrada que deja la migración de la v2: el corte de OSMA, solo España.
                                {'proveedor': 'ZENTRADA', 'umbral_caidas_30d': 6, 'paises_filtro': ['ES'],
                                 'paises_calculo': ['ES']}],
        'escaner2_pasada': pasadas,
        'productos': json.loads(json.dumps(PRODUCTOS)),
        'inventario_fba': [],
    }, 'storage': {'escaner2': {'zentrada/%s/%s' % (PASADA, NOMBRE_EXCEL):
                                base64.b64encode(excel(leido=leido)).decode()}}}


def caso(escena):
    tmp = tempfile.mkdtemp(prefix='e2zen_')
    ruta = os.path.join(tmp, 'estado.json')
    inicial = estado_inicial(escena)
    with open(ruta, 'w', encoding='utf-8') as fh:
        json.dump(inicial, fh)
    cod, log = correr(ruta, 'escaner2_zentrada_barrido.py', {'SUPABASE_SERVICE_KEY': 'svc-de-mentira', 'PASADA': PASADA,
                                                              'GITHUB_RUN_ID': '616161'})
    bd = json.load(open(ruta, encoding='utf-8'))
    return cod, log, bd, ruta, inicial


print('B1 · [sin_llave] sin la llave de servicio no se corre')
_tmp = tempfile.mkdtemp(prefix='e2zen_')
_ruta = os.path.join(_tmp, 'estado.json')
json.dump(estado_inicial('bueno'), open(_ruta, 'w', encoding='utf-8'))
for _prog in ('escaner2_zentrada_barrido.py', 'escaner2_zentrada_cruce.py'):
    _cod, _log = correr(_ruta, _prog, {'PASADA': PASADA})
    eq('B1 · %s: ROJO, la línea exacta y NINGÚN cliente' % _prog,
       (_cod, 'ESCANER2_NO_EJECUTADO: sin llave de servicio' in _log, 'CLIENTES_CREADOS=0' in _log), (1, True, True))

print('B2 · [bueno] el Excel → barrido → las dos listas → los CSV de ES (y uno de FR) → cruce → Excel')
cod, log, bd, ruta, inicial = caso('bueno')
eq('B2 · el barrido sale en VERDE', cod, 0)
pas = [p for p in bd['tablas']['escaner2_pasada'] if p.get('id') == PASADA][0]
eq('B2 · 🔴 la pasada: la del id de la v2, ZENTRADA, modo excel, esperando_csv, crudo 106 = previas 5 + foto 101',
   (pas['proveedor'], pas['modo'], pas['estado'], pas['run_id'], pas['n_crudo'], pas['n_foto'],
    sum(pas['p_' + p] for p in e2.PUERTAS_PREVIAS), pas['n_tandas'], pas['ruta_lista'], pas['n_eans_lista']),
   ('ZENTRADA', 'excel', 'esperando_csv', 616161, 106, 101, 5, 1, 'zentrada/%s/eans.txt' % PASADA, 101))
_alm = bd['storage']['escaner2']
_txt = lambda k: base64.b64decode(_alm[k]).decode('utf-8')  # noqa: E731
eq('B2 · en el almacén: el Excel de la v2, eans.txt, asins.txt y barrido.json',
   sorted(k.split('/')[-1] for k in _alm if k.startswith('zentrada/%s/' % PASADA)),
   sorted([NOMBRE_EXCEL, 'asins.txt', 'barrido.json', 'eans.txt']))
eq('B2 · asins.txt: los ASIN de «Fichas» de un EAN que se puede pedir', _txt('zentrada/%s/asins.txt' % PASADA).split('\n'),
   [ASIN_PACK2, ASIN_SUELTA, ASIN_PACK3])
eq('B2 · eans.txt: el UPC tal como lo escribe Zentrada (sin el cero)', E_UPC_ZEN in _txt('zentrada/%s/eans.txt' % PASADA)
   .split('\n'), True)

_carpeta = 'zentrada/%s/csv/' % PASADA
_filas_es = [(ASIN_UNO, E_UNO, '30', '12.95', 'Producto Uno'),
             (ASIN_UPC, E_UPC, '25', '12.95', 'Producto G-2'),
             (ASIN_COMP, E_COMP, '20', '12.95', 'Producto A-3'),
             (ASIN_SUELTA, E_FICHA, '30', '9.95', 'Producto D-4'),
             # 🔑 Amazon dice que el pack de 3 de «Fichas» lleva 6 (dos señales): mandan las unidades de «Fichas».
             (ASIN_PACK3, E_PACK, '30', '12.95', 'Producto D-6 pack de 3', {'Número de artículos': '6',
                                                                            'Paquete: Cantidad': '6'}),
             # (AM) La golosina de 45 g en un multipack de 24, dicho por TRES señales (tamaño, recuento y título)…
             (ASIN_AM24, RELLENO[0], '30', '12.95', 'Golosinas 24 x 45 g', {'Tamaño': '45 g (Paquete de 24)',
                                                                           'Número de artículos': '24'}),
             # …y la otra, con UNA sola (el tamaño): posible pack.
             (ASIN_AM1, RELLENO[1], '30', '12.95', 'Golosinas surtidas', {'Tamaño': '45 g (Paquete de 24)'})]
_filas_asin = [(ASIN_PACK2, '', '15', '16.95', 'Producto D-4 pack de 2'),
               (ASIN_SUELTA, E_FICHA, '30', '9.95', 'Producto D-4'),
               (ASIN_PACK3, E_PACK, '30', '12.95', 'Producto D-6 pack de 3')]
_alm[_carpeta + '20261005-100000-KeepaExport-2026-10-05-VisualizadorDeProductos.csv'] = \
    base64.b64encode(csv_real('es', _filas_es)).decode()
_alm[_carpeta + '20261005-100100-KeepaExport-2026-10-05-VisualizadorDeProductos (1).csv'] = \
    base64.b64encode(csv_real('es', _filas_asin)).decode()
_alm[_carpeta + '20261005-100200-KeepaExport-2026-10-05-VisualizadorDeProductos (2).csv'] = \
    base64.b64encode(csv_real('fr', _filas_es)).decode()
bd.pop('programa', None)
json.dump(bd, open(ruta, 'w', encoding='utf-8'), default=str)
cod2, log2 = correr(ruta, 'escaner2_zentrada_cruce.py', {'SUPABASE_SERVICE_KEY': 'svc-de-mentira', 'PASADA': PASADA,
                                                          'GITHUB_RUN_ID': '626262'})
bd = json.load(open(ruta, encoding='utf-8'))
T = bd['tablas']
_alm = bd['storage']['escaner2']
eq('B2 · el cruce sale en VERDE', cod2, 0)
cr = [c for c in T['escaner2_cruce'] if c.get('run_id') == 626262][0]
eq('B2 · 🔴 el cruce: lista, cuadra, ofertas 106 = previas 5 + 101 puertas, solo ES',
   (cr['estado'], cr['cuadra'], cr['n_crudo'], cr['n_previas'], cr['n_entradas'], sum(cr['n_' + x] for x in 'abcdef'),
    cr['paises_filtro'], cr['paises_calculo'], cr['paises_usados']),
   ('lista', True, 106, 5, 101, 101, ['ES'], ['ES'], ['ES']))
eq('B2 · los CSV: el de EAN y el de ASIN de ES cuentan, cada uno por su lista; el de FR se ignora y se avisa',
   (sorted((f['pais'], f['usado'], f.get('lista')) for f in cr['ficheros']), 'CSV de FR ignorado' in (cr['aviso'] or '')),
   ([('ES', True, 'asin'), ('ES', True, 'ean'), ('FR', False, 'ean')], True))
_res = {f['ean_original']: r for r in T['escaner2_resultado_ean'] for f in T['escaner2_foto'] if f['id'] == r['foto_id']}
_pais = {r['resultado_ean_id']: r for r in T['escaner2_resultado_pais'] if r['cruce_id'] == cr['id']}
eq('B2 · 🔴 el UPC que Zentrada escribe sin el cero casa con el de Keepa (con el cero)',
   (_res[E_UPC_ZEN]['asin'], _res[E_UPC_ZEN]['puerta']), (ASIN_UPC, 'f'))
eq('B2 · E_UNO: COMPRAR con la oferta más barata (2,00) en su cuenta de ES',
   (_res[E_UNO]['puerta'], _pais[_res[E_UNO]['id']]['pa']), ('f', 2.0))
eq('B2 · la de la competencia: VALORAR (4,80)', (_res[E_COMP]['puerta'], _pais[_res[E_COMP]['id']]['pa']), ('e', 4.8))
eq('B2 · 🔴 Keepa cae por EAN en el pack de 3 de «Fichas»: coste 4,50 y VALORAR (con 1,50 sería COMPRAR)',
   (_res[E_PACK]['asin'], _res[E_PACK]['puerta'], _pais[_res[E_PACK]['id']]['pa']), (ASIN_PACK3, 'e', 4.5))
eq('B2 · los de relleno y el duplicado no están en Keepa: puerta a',
   sorted({_res[e]['puerta'] for e in RELLENO[2:] + [E_DUP.lstrip('0')]}), ['a'])
_am24, _am1 = _res[RELLENO[0]], _res[RELLENO[1]]
eq('🔴 B2 · (AM) la golosina de 45 g en un multipack de 24 (tres señales): coste 24 × 1,00 = 24,00 y deja de ser COMPRAR '
   '(puerta d); el detalle dice el pack y sus señales',
   (_am24['asin'], _am24['puerta'], _pais[_am24['id']]['pa'],
    'pack de 24 en Amazon: coste 24 × 1,000 € = 24,00 €' in _am24['detalle'], 'señales: Zentrada 45 g' in _am24['detalle']),
   (ASIN_AM24, 'd', 24.0, True, True))
eq('🔴 B2 · (AM) la de UNA sola señal: no se multiplica (1,00) y nunca es COMPRAR: VALORAR, con «posible pack en Amazon: '
   'revisar»', (_am1['puerta'], _pais[_am1['id']]['pa'], _pais[_am1['id']]['decision'],
                'posible pack en Amazon: revisar' in _am1['detalle']), ('e', 1.0, 'VALORAR', True))
eq('B2 · (AM) …y el pack de 3 de «Fichas», aunque Amazon diga 6: mandan las de «Fichas» (4,50), no se multiplica dos veces',
   (_pais[_res[E_PACK]['id']]['pa'], 'en Amazon' in _res[E_PACK]['detalle']), (4.5, False))

_xl = [k for k in _alm if k.startswith('zentrada/%s/%s/Escaner2_ZENTRADA_' % (PASADA, cr['id']))]
eq('B2 · el Excel, en su carpeta y con su nombre (el que guarda el cruce)', (len(_xl), cr['ruta_excel'] == _xl[0]),
   (1, True))
_wb = load_workbook(io.BytesIO(base64.b64decode(_alm[_xl[0]])))
HOJAS_VIEJO = ['Análisis', 'Descartados', 'Ambiguos', 'Sin_rank', 'Precio por lote']
HOJAS_E2 = ['Resumen', 'Comparación', 'Varias fichas', 'Puertas', 'Puertas previas', ez.HOJA_AGOTADOS]
eq('B2 · las hojas del PRO, en su orden, y la ÚLTIMA la de nuestros agotados', _wb.sheetnames, HOJAS_VIEJO + HOJAS_E2)
_cab = [c.value for c in _wb['Análisis'][1]]
eq('B2 · «Análisis»: las columnas del PRO y, AL FINAL, «Ficha compartida», «Ficha del mapa» y las nueve de la compra',
   _cab, e2.columnas_analisis() + [e2.COLUMNA_FICHA_COMPARTIDA, ez.COLUMNA_MAPA] + ez.COLUMNAS_COMPRA)
_filas = list(_wb['Análisis'].iter_rows(min_row=2, values_only=True))
_ix = {n: _cab.index(n) for n in _cab}
_por = {(r[_ix['EAN']], r[_ix['ASIN']]): r for r in _filas}
eq('B2 · «Análisis» solo pinta ES, y cada ASIN una sola vez',
   (sorted({r[_ix['País']] for r in _filas}), len(_filas) == len({r[_ix['ASIN']] for r in _filas})), (['ES'], True))
_u = _por[(E_UNO, ASIN_UNO)]
eq('B2 · 🔴 E_UNO: PA 2,00, COMPRAR, comprado a «Mayorista Alfa · fuera de Mix», más barato en Mix 2,50, origen '
   '«nuestro, top 300», puesto 12, ni alimentación ni peligrosa',
   [_u[_ix[k]] for k in ('PA (€)', 'Decisión', 'Dónde se compra', 'Unidades por caja', 'Cajas mínimas',
                         'Compra mínima (€)', 'Precio más barato en Mix (€)', 'Origen', 'Puesto en más vendidos',
                         'Alimentación', 'Mercancía peligrosa')],
   [2.0, 'COMPRAR', 'Mayorista Alfa · fuera de Mix', 6, 1, 150.0, 2.5, 'nuestro, top 300', 12, 'no',
    'no'])
eq('B2 · la de la competencia: alimentación y peligrosa «sí», y sin precio de Mix (la elegida ya va en Mix)',
   [_por[(E_COMP, ASIN_COMP)][_ix[k]] for k in ('Decisión', 'Dónde se compra', 'Precio más barato en Mix (€)',
                                                 'Alimentación', 'Mercancía peligrosa', 'Origen')],
   ['VALORAR', 'Mayorista Alfa · va en Mix', None, 'sí', 'sí', 'competidor'])
eq('B2 · 🔴 las fichas de «Fichas» de E_FICHA: la suelta (la misma que la fila de siempre: UNA fila, la de «Fichas») y el '
   'pack de 2 (coste 2,00, que el CSV de EAN no traía: llega por el de ASIN)',
   [(_por[(E_FICHA, a)][_ix['PA (€)']], _por[(E_FICHA, a)][_ix[ez.COLUMNA_MAPA]], _por[(E_FICHA, a)][_ix['Origen']])
    for a in (ASIN_SUELTA, ASIN_PACK2)],
   [(1.0, 'ficha de «Fichas»', 'competidor, top 300'), (2.0, 'pack ×2 (de «Fichas»)', 'competidor, top 300')])
eq('B2 · el pack de 3: una fila, la de «Fichas», con su coste 4,50 y VALORAR',
   [(_por[(E_PACK, ASIN_PACK3)][_ix[k]]) for k in ('PA (€)', 'Decisión', ez.COLUMNA_MAPA)],
   [4.5, 'VALORAR', 'pack ×3 (de «Fichas»)'])
eq('B2 · el UPC sale con el número de Zentrada en la columna EAN', (E_UPC_ZEN, ASIN_UPC) in _por, True)

_resu = {r[0]: r[1] for r in _wb['Resumen'].iter_rows(values_only=True)}
# COMPRAR en «Análisis»: E_UNO (nuestro, top 300), E_UPC (top 300), la suelta y el pack de 2 de E_FICHA (competidor,
# top 300). VALORAR: la de la competencia, el pack de 3 y (AM) la golosina de una sola señal. La de tres señales, NO COMPRAR.
eq('🔴 B2 · el Resumen: COMPRAR 4 y VALORAR 3 en ES; por origen nuestro 1, competidor 2, top 300 4; y SOLO del top 300, 1',
   [_resu.get(k) for k in ('COMPRAR en ES (en «Análisis»)', 'VALORAR en ES (en «Análisis»)', 'COMPRAR · origen «nuestro»',
                           'COMPRAR · origen «competidor»', 'COMPRAR · origen «top 300»',
                           'COMPRAR que vienen SOLO del top 300 (ni nuestro ni competidor)')],
   [4, 3, 1, 2, 4, 1])
eq('B2 · (AM) el Resumen cuenta los packs de Amazon, como el de OSMA',
   _resu.get('Packs de Amazon detectados (la regla AM de OSMA)'), '1 multiplicados · 1 dudosos a VALORAR')
_g24 = _por[(RELLENO[0], ASIN_AM24)]
_g1 = _por[(RELLENO[1], ASIN_AM1)]
eq('B2 · (AM) en «Análisis»: la de 24 con PA 24,00, NO COMPRAR y el pack en el nombre y en «Coherencia caja»; la dudosa, '
   'VALORAR con el aviso en «Coherencia caja»',
   (_g24[_ix['PA (€)']], _g24[_ix['Decisión']], 'pack de 24 en Amazon' in _g24[_ix['Nombre']],
    _g24[_ix['Coherencia caja']], _g1[_ix['Decisión']], str(_g1[_ix['Coherencia caja']]).startswith(ez.TEXTO_POSIBLE_PACK)),
   (24.0, 'NO COMPRAR', True, 'pack de 24 en Amazon: coste 24 × 1,000 € = 24,00 €', 'VALORAR', True))
eq('B2 · …y el coste dicho en llano, el Excel de origen con su lectura, y el cuadre',
   (_resu.get('Coste de cada EAN'), str(_resu.get('Excel de Zentrada')).startswith(NOMBRE_EXCEL + ' · leído 20'),
    _resu.get('Cuadra'), _resu.get('Ofertas de Zentrada (el Excel)')),
   ('la oferta más barata que se puede pedir, tal cual: sin porte ni cuota', True, 'SÍ', 106))
_ag = list(_wb[ez.HOJA_AGOTADOS].iter_rows(values_only=True))
eq('🔴 B2 · la ÚLTIMA hoja: nuestros agotados (el que no se puede pedir y el que no tiene oferta)',
   [list(r) for r in _ag],
   [ez.COLUMNAS_AGOTADOS, [E_AGOT, ASIN_AGOT, 'Producto B-7', '—', 1, 5.0, ez.TEXTO_SIN_PEDIR],
    [E_SIN, ASIN_SIN, 'Producto sin oferta', '—', 1, None, ez.TEXTO_SIN_OFERTA]])
eq('B2 · «Puertas previas» con los nombres de Zentrada', sorted({r[4] for r in list(_wb['Puertas previas'].iter_rows(
    values_only=True))[1:]}), sorted(['EAN con forma rara', 'Otra oferta del mismo EAN, más cara']))
_pu = list(_wb['Puertas'].iter_rows(values_only=True))
eq('B2 · «Puertas»: las 101 de la foto y las 3 fichas de «Fichas» con su marca al final',
   (len(_pu) - 1, sorted(r[-1] for r in _pu[1:] if r[-1])),
   (104, sorted(['ficha de «Fichas» · ' + ASIN_SUELTA, 'pack ×2 (de «Fichas») · ' + ASIN_PACK2,
                 'pack ×3 (de «Fichas») · ' + ASIN_PACK3])))

print('B3 · 🔴 el registro (repo PÚBLICO) no lleva ni un EAN, ASIN, precio ni nombre')
_prohibido = [E_UNO, E_UPC, E_UPC_ZEN, E_COMP, E_FICHA, E_PACK, E_AGOT, E_SIN, E_DUP, ASIN_UNO, ASIN_UPC, ASIN_COMP,
              ASIN_PACK2, ASIN_SUELTA, ASIN_PACK3, ASIN_AGOT, ASIN_SIN, '12.95', '4.8', '2.5', 'Mayorista', 'Producto Uno',
              'Producto D-4', 'Producto B-7', 'Producto sin oferta', ASIN_AM24, ASIN_AM1, RELLENO[0], RELLENO[1],
              'Golosina']
eq('B3 · ni en el barrido ni en el cruce (si sale algo, la línea que lo lleva)',
   [(x, [ln for ln in (log + log2).splitlines() if x in ln][:2]) for x in _prohibido if x in log or x in log2], [])

print('B4 · solo se escribe en escaner2_* y en la carpeta de la pasada')
eq('B4 · las tablas escritas', sorted({t for _p, op, t in bd['ops'] if op != 'select'}),
   ['escaner2_apartado', 'escaner2_cruce', 'escaner2_foto', 'escaner2_pasada', 'escaner2_resultado_ean',
    'escaner2_resultado_pais'])
eq('B4 · y las sembradas acaban como empezaron',
   [t for t in SEMBRADAS if bd['tablas'][t] != json.loads(json.dumps(inicial['tablas'][t]))], [])
eq('B4 · en el almacén, solo zentrada/<pasada>/ (y el Excel de la v2 no se toca)',
   (sorted({k.split('/')[0] + '/' + k.split('/')[1] for k in _alm}),
    _alm['zentrada/%s/%s' % (PASADA, NOMBRE_EXCEL)] == inicial['storage']['escaner2']['zentrada/%s/%s'
                                                                                  % (PASADA, NOMBRE_EXCEL)],
    [s for s in bd.get('subidas', []) if not s.startswith('zentrada/%s/' % PASADA)]),
   (['zentrada/' + PASADA], True, []))

print('B5 · 🔴 sin guardas, sin pasada')
cod, log, bd, _r, _i0 = caso('viejo')
eq('B5 · [viejo] el Excel de hace 3 días: ROJO, lo dice, y NO se crea la pasada ni se escribe nada',
   (cod, 'no pasa sus guardas' in log, 'hace más de 2 días' in log, bd['tablas']['escaner2_pasada'],
    sorted({t for _p, op, t in bd.get('ops', []) if op != 'select'})), (1, True, True, [], []))
cod, log, bd, _r, _i0 = caso('repetida')
eq('B5 · [repetida] la pasada ya existe: ROJO y no se toca',
   (cod, 'ya existe' in log, bd['tablas']['escaner2_pasada'], sorted({t for _p, op, t in bd.get('ops', [])
                                                                      if op != 'select'})),
   (1, True, [{'id': PASADA, 'proveedor': 'ZENTRADA', 'estado': 'esperando_csv', 'run_id': 1}], []))

print()
if fallos:
    print('FALLAN %d: %s' % (len(fallos), fallos))
    sys.exit(1)
print('TODO OK')
