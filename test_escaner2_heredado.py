# -*- coding: utf-8 -*-
"""Banco del escaner 2 de HEO · INDEPENDIENTE DEL ESCANER VIEJO (encargo B7, 25-sep-2026).

Fernando: el Excel del escaner nuevo «se ha copiado y cuando apague el escaner viejo va a dejar de salir.
Hazlo bien y que se calcule bien el excel en la app nueva y que sea exacto a como me salia antes».

Desde el B7 el escaner 2 lleva COPIADAS LITERALMENTE las piezas del viejo que usa, en cinco ficheros
suyos (uno por fichero de origen, porque el viejo repite nombres entre ficheros: `norm`, `COLS`…):
    escaner2_heredado_nube.py       ← moloka_escaner_nube.py (formula, EAN/chase/caja, IVA, eleccion de
                                      ficha y la Celda 9 del Excel, hecha funcion: `excel_del_viejo`)
    escaner2_heredado_pro.py        ← moloka_escaner_pro.py (el lector del CSV del Visualizador)
    escaner2_heredado_director.py   ← director_heo_prep.py (el filtro del modo «marcas»)
    escaner2_heredado_descarga.py   ← descargar_heo.py (la API de HEO y la tanda del Visualizador)
    escaner2_heredado_escaparate.py ← procesador_keepa_escaparate.py (`TIPADAS`)
Los genera scripts/generar_escaner2_heredado.py desde un commit (`git show`), no se teclean.

SIN RED, SIN SECRETOS Y SIN BASE. QUE PRUEBA:
  (A) 🔴 LA GUARDA: ningun escaner2_*.py IMPORTA, ABRE o EJECUTA uno de los cinco ficheros del viejo
      (por estructura: `import`, y textos que son la ruta o el nombre del modulo; no los comentarios ni
      las explicaciones), ni ningun workflow escaner2-*.yml los corre. Y muerde: se prueba con textos
      que si lo hacen, y calla con los que solo lo mencionan.
  (B) PARIDAD DE PIEZAS, MIENTRAS EL VIEJO EXISTA: cada pieza heredada tiene el MISMO arbol (ast) y el
      MISMO texto (con sus comentarios) que la del viejo. Si el viejo cambia, NO se arrastra: sale en
      ROJO «el viejo ha cambiado en X: decide Fernando si se hereda». Si el viejo ya no esta, se SALTA
      y lo dice (no tumba el CI).
  (C) PARIDAD DEL EXCEL, MIENTRAS EL VIEJO EXISTA: el libro de la Celda 9 del viejo (su bloque, sacado
      de su fichero como hasta el B6) y el de `excel_del_viejo` son IDENTICOS celda a celda con los
      mismos datos (inventados: dos escenas, una llena y una vacia). La de los datos reales del cruce
      37eaa138 se hizo en local y va en el parte (el repo no lleva costes).
  (D) EL LIBRO DE REFERENCIA, QUE SOBREVIVE AL VIEJO: excel_referencia_escaner2.xlsx, generado con
      datos INVENTADOS (ni precios reales ni costes), se regenera y se compara celda a celda.
  (E) 🔴 LA COMPARACION MUERDE: con una formula de «Análisis» y un ancho de columna cambiados en una copia
      del fichero heredado, (C) y (D) se ponen rojas.

Para regenerar el libro de referencia (solo si Fernando decide cambiar el Excel):
    python test_escaner2_heredado.py --regenerar-referencia
"""
import ast
import copy
import hashlib
import io
import os
import re
import sys
import tempfile
from contextlib import redirect_stdout

import escaner2_huella_excel as HU
import escaner2_motor as e2

AQUI = os.path.dirname(os.path.abspath(__file__))
fallos = []


def eq(nombre, obtenido, esperado):
    ok = obtenido == esperado
    if not ok:
        fallos.append(nombre)
    print(('OK ' if ok else 'XX ') + nombre + ('' if ok else '   got=%r exp=%r' % (obtenido, esperado)))


# Heredado → origen en el viejo.
HEREDADOS = (('escaner2_heredado_nube.py', 'moloka_escaner_nube.py'),
             ('escaner2_heredado_pro.py', 'moloka_escaner_pro.py'),
             ('escaner2_heredado_director.py', 'director_heo_prep.py'),
             ('escaner2_heredado_descarga.py', 'descargar_heo.py'),
             ('escaner2_heredado_escaparate.py', 'procesador_keepa_escaparate.py'))
VIEJOS = tuple(v for _h, v in HEREDADOS)
MODULOS_VIEJOS = tuple(v[:-3] for v in VIEJOS)
# 🔑 Si Fernando decide NO heredar un cambio del viejo, se apunta aqui: (fichero del viejo, pieza) → md5 del
#    arbol (ast.dump) de la pieza del viejo que se ha decidido no heredar. Hoy, ninguna.
DIVERGENCIAS_ACEPTADAS = {}


def leer(ruta):
    with io.open(os.path.join(AQUI, ruta), encoding='utf-8') as fh:
        return fh.read()


# ═══════════════════════════════════════════════════════════════════════════════
print('(A) la guarda: el escáner 2 no lee, no importa y no ejecuta ninguno de los cinco ficheros del viejo')
_RE_RUTA = re.compile(r'(^|[/\\])(%s)(\.py)?$' % '|'.join(re.escape(m) for m in MODULOS_VIEJOS))
_RE_YML = re.compile(r'\b(%s)(\.py)?\b' % '|'.join(re.escape(m) for m in MODULOS_VIEJOS))


def lecturas_del_viejo(texto):
    """[(linea, que)] de un .py: `import`/`from` de un modulo del viejo, o un texto que ES la ruta o el
    nombre del modulo (lo que haria falta para abrirlo, importarlo o ejecutarlo). Los docstrings y los
    comentarios no cuentan: mencionar el viejo para explicar de donde sale algo no es leerlo."""
    arbol = ast.parse(texto)
    docs = {id(n.value) for n in ast.walk(arbol) if isinstance(n, ast.Expr) and isinstance(n.value, ast.Constant)
            and isinstance(n.value.value, str)}
    salida = []
    for n in ast.walk(arbol):
        if isinstance(n, ast.Import):
            salida += [(n.lineno, 'import ' + a.name) for a in n.names if a.name.split('.')[0] in MODULOS_VIEJOS]
        elif isinstance(n, ast.ImportFrom) and (n.module or '').split('.')[0] in MODULOS_VIEJOS:
            salida.append((n.lineno, 'from %s import …' % n.module))
        elif isinstance(n, ast.Constant) and isinstance(n.value, str) and id(n) not in docs:
            if _RE_RUTA.search(n.value.strip()):
                salida.append((n.lineno, 'texto %r' % n.value))
    return sorted(salida)


def lecturas_yml(texto):
    salida = []
    for i, linea in enumerate(texto.splitlines(), 1):
        codigo = linea.split(' #')[0] if not linea.lstrip().startswith('#') else ''
        if _RE_YML.search(codigo):
            salida.append((i, linea.strip()))
    return salida


_PY = sorted(f for f in os.listdir(AQUI) if f.startswith('escaner2_') and f.endswith('.py'))
_YML = sorted(os.path.join('.github', 'workflows', f) for f in os.listdir(os.path.join(AQUI, '.github', 'workflows'))
              if f.startswith('escaner2-') and f.endswith('.yml'))
print('    miro %d ficheros .py (%s) y %d workflows' % (len(_PY), ', '.join(_PY), len(_YML)))
eq('(A) hay qué mirar: los programas, el motor, la huella y los cinco heredados, y los dos workflows',
   ({'escaner2_motor.py', 'escaner2_heo_barrido.py', 'escaner2_heo_cruce.py'} | {h for h, _v in HEREDADOS}) <= set(_PY)
   and len(_YML) == 2, True)
_hallado = {f: lecturas_del_viejo(leer(f)) for f in _PY}
_hallado.update({f: lecturas_yml(leer(f)) for f in _YML})
eq('(A) 🔴 ninguno lee, importa o ejecuta el viejo', {f: v for f, v in _hallado.items() if v}, {})
eq('(A) …y las rutas del motor apuntan a los heredados, no al viejo',
   (e2.RUTA_MOTOR, e2.RUTA_DIRECTOR_HEO, e2.RUTA_DESCARGAR_HEO, e2.RUTA_ESCAPARATE),
   ('escaner2_heredado_nube.py', 'escaner2_heredado_director.py', 'escaner2_heredado_descarga.py',
    'escaner2_heredado_escaparate.py'))
for _texto, _debe in (("RUTA_MOTOR = 'moloka_escaner_nube.py'", True),
                      ("import moloka_escaner_pro as pro", True),
                      ("from descargar_heo import descargar_catalogo_heo", True),
                      ("exec(open('director_heo_prep.py').read())", True),
                      ("m = importlib.import_module('procesador_keepa_escaparate')", True),
                      ("ruta = os.path.join(AQUI, 'moloka_escaner_nube.py')", True),
                      ("ruta = f'{AQUI}/descargar_heo.py'", True),
                      ("# esto sale de moloka_escaner_nube.py", False),
                      ('"""Copiado de descargar_heo.py y director_heo_prep.py."""', False),
                      ("print('HEO dio 3 y descargar_heo devolvió 2')", False),
                      ("import escaner2_heredado_pro as pro", False)):
    eq('(A) 🔴 %s: %r' % ('lo caza' if _debe else 'calla con', _texto), bool(lecturas_del_viejo(_texto)), _debe)
eq('(A) 🔴 lo caza en un workflow: «run: python -u descargar_heo.py»',
   bool(lecturas_yml('      - run: python -u descargar_heo.py')), True)
eq('(A) …y calla con un comentario que lo menciona', bool(lecturas_yml('# con la funcion de descargar_heo.py')), False)


# ═══════════════════════════════════════════════════════════════════════════════
print('\n(B) paridad de piezas con el viejo (mientras exista)')


def _nombres(n):
    return tuple(e2._nombres_asignados(n))


def _unica(candidatas, que):
    arriba = [c for c in candidatas if getattr(c, '_arriba', False)]
    elegidas = arriba or candidatas
    if len(elegidas) != 1:
        return None, '%s: %d en el viejo' % (que, len(elegidas))
    return elegidas[0], None


def _en_viejo(arbol_v, n):
    """La pieza del viejo que corresponde a la heredada `n` (por nombre; si no esta a nivel superior, la
    UNICA de ese nombre en todo el fichero: `keyrank`, `TANDA`)."""
    for x in arbol_v.body:
        x._arriba = True
    if isinstance(n, ast.FunctionDef):
        return _unica([x for x in ast.walk(arbol_v) if isinstance(x, ast.FunctionDef) and x.name == n.name], 'def ' + n.name)
    if isinstance(n, ast.Assign):
        return _unica([x for x in ast.walk(arbol_v) if isinstance(x, ast.Assign) and _nombres(x) == _nombres(n)],
                      '%s = …' % ', '.join(_nombres(n)))
    iguales = [x for x in arbol_v.body if ast.dump(x) == ast.dump(n)]
    return (iguales[0], None) if iguales else (None, 'sentencia %r' % ast.unparse(n)[:60])


def _nombre_pieza(n):
    if isinstance(n, ast.FunctionDef):
        return n.name
    if isinstance(n, ast.Assign):
        return ', '.join(_nombres(n))
    return ast.unparse(n)[:50]


def _comentarios_encima(lineas, i0):
    i = i0 - 1
    while i >= 0 and lineas[i].lstrip().startswith('#'):
        i -= 1
    return [l.strip() for l in lineas[i + 1:i0] if not l.strip().startswith('# ──')]


def _texto(lineas, n, sangria=0):
    """El texto de la pieza (sus lineas, sin la sangria que tuviera) con los comentarios pegados encima."""
    cuerpo = [l[sangria:] if l.startswith(' ' * sangria) else l for l in lineas[n.lineno - 1:n.end_lineno]]
    return _comentarios_encima(lineas, n.lineno - 1) + [l.rstrip() for l in cuerpo]


def bloque_celda9(texto_v):
    """(lineas de texto, sentencias) de la Celda 9 del viejo: de los `import` de openpyxl que preceden a
    `COLS = …` a la sentencia anterior a `_sin_excel = …` (las anclas con que la leia el escaner 2 hasta el
    B6), con la cabecera de comentarios de encima."""
    cuerpo = ast.parse(texto_v).body
    idx = lambda nombre: [i for i, x in enumerate(cuerpo) if isinstance(x, ast.Assign) and nombre in _nombres(x)]
    ini, fin = idx('COLS'), idx('_sin_excel')
    assert len(ini) == 1 and len(fin) == 1, 'las anclas de la Celda 9 no son unicas en el viejo'
    ini, fin = ini[0], fin[0]
    while ini > 0 and isinstance(cuerpo[ini - 1], (ast.Import, ast.ImportFrom)):
        ini -= 1
    lineas = texto_v.split('\n')
    l_ini = cuerpo[ini].lineno - 1
    while l_ini > 0 and lineas[l_ini - 1].lstrip().startswith('#'):
        l_ini -= 1
    l_fin = cuerpo[fin - 1].end_lineno
    return [l.rstrip() for l in lineas[l_ini:l_fin]], cuerpo[ini:fin]


def paridad(heredado, viejo):
    """[(pieza, diferencia)] de un fichero heredado frente a su viejo; [] si todo es igual."""
    texto_h, texto_v = leer(heredado), leer(viejo)
    arbol_h, arbol_v = ast.parse(texto_h), ast.parse(texto_v)
    lin_h, lin_v = texto_h.split('\n'), texto_v.split('\n')
    difs, n_piezas = [], 0
    for n in arbol_h.body:
        if isinstance(n, ast.Expr) and isinstance(n.value, ast.Constant) and n is arbol_h.body[0]:
            continue                                                  # el docstring del heredado
        n_piezas += 1
        nombre = _nombre_pieza(n)
        if isinstance(n, ast.FunctionDef) and n.name == 'excel_del_viejo':
            # La Celda 9: el cuerpo de la funcion (sin el `return wb` final) es el bloque del viejo.
            lin_b, sent_b = bloque_celda9(texto_v)
            cuerpo = [l[4:] if l.startswith('    ') else l for l in lin_h[n.lineno:n.end_lineno - 1]]
            if [ast.dump(x) for x in n.body[:-1]] != [ast.dump(x) for x in sent_b]:
                difs.append((nombre, 'el árbol de la Celda 9 cambia'))
            elif [l.rstrip() for l in cuerpo] != lin_b:
                difs.append((nombre, 'el texto (o los comentarios) de la Celda 9 cambia'))
            if ast.unparse(n.body[-1]) != 'return wb' or tuple(a.arg for a in n.args.args) != e2.DATOS_CELDA9:
                difs.append((nombre, 'la línea def o el return no son los del B7'))
            continue
        if heredado == 'escaner2_heredado_pro.py' and isinstance(n, ast.Import) and [a.name for a in n.names] == ['csv']:
            # Documentado en su cabecera: de `import pandas as pd, csv` solo se trae `csv`.
            if not any(isinstance(x, ast.Import) and 'csv' in [a.name for a in x.names] for x in arbol_v.body):
                difs.append((nombre, 'el viejo ya no importa csv'))
            continue
        v, error = _en_viejo(arbol_v, n)
        if v is None:
            difs.append((nombre, error))
            continue
        if ast.dump(v) != ast.dump(n):
            clave = (viejo, nombre)
            huella_v = hashlib.md5(ast.dump(v).encode()).hexdigest()
            if DIVERGENCIAS_ACEPTADAS.get(clave) == huella_v:
                print('    (B) %s · %s: el viejo es distinto y Fernando decidió no heredarlo (aceptada)' % clave)
                continue
            difs.append((nombre, 'el árbol cambia'))
        elif _texto(lin_v, v, v.col_offset) != _texto(lin_h, n):
            difs.append((nombre, 'el texto o los comentarios cambian'))
    return n_piezas, difs


_RE_MD5 = re.compile(r'^# ── ORIGEN: .* · md5 ([0-9a-f]{32}) ──$')


def tocados_a_mano(heredado):
    """[pieza] del heredado cuyo texto ya no es el que dejo el generador: cada pieza lleva en su cabecera el
    md5 de su texto (sin las lineas `# ──`). No necesita el viejo: protege la copia tambien el dia que el
    viejo ya no este."""
    texto = leer(heredado)
    lineas = texto.split('\n')
    arbol = ast.parse(texto)
    salida = []
    for n in arbol.body:
        if n is arbol.body[0] and isinstance(n, ast.Expr):
            continue
        if heredado == 'escaner2_heredado_pro.py' and isinstance(n, ast.Import) and [a.name for a in n.names] == ['csv']:
            continue                                   # la linea `import csv`, documentada en la cabecera del fichero
        i = n.lineno - 2
        while i >= 0 and lineas[i].lstrip().startswith('#') and not _RE_MD5.match(lineas[i]):
            i -= 1
        m = _RE_MD5.match(lineas[i]) if i >= 0 else None
        if not m:
            salida.append(_nombre_pieza(n) + ' (sin cabecera de origen)')
            continue
        j = i + 1
        while lineas[j].startswith('# ──'):
            j += 1
        if hashlib.md5('\n'.join(l.rstrip() for l in lineas[j:n.end_lineno]).encode('utf-8')).hexdigest() != m.group(1):
            salida.append(_nombre_pieza(n))
    return salida


for _h, _v in HEREDADOS:
    _tocados = tocados_a_mano(_h)
    for _pieza in _tocados:
        print('XX (B) 🔴 EL HEREDADO SE HA TOCADO A MANO en %s · %s: no se toca a mano; si hay que cambiarlo, se '
              'regenera con scripts/generar_escaner2_heredado.py (y decide Fernando)' % (_h, _pieza))
    eq('(B) %s: ninguna pieza tocada a mano (cada una, con el md5 que dejó el generador)' % _h, _tocados, [])
    if not os.path.exists(os.path.join(AQUI, _v)):
        print('    (B) SALTADO: %s ya no está en el repo; la paridad con el viejo se salta (el escáner 2 no lo '
              'necesita: lleva su copia en %s)' % (_v, _h))
        continue
    _n, _difs = paridad(_h, _v)
    for _pieza, _que in _difs:
        if _pieza in _tocados:
            continue                                   # ya dicho arriba: el que ha cambiado es el heredado
        print('XX (B) 🔴 EL VIEJO HA CAMBIADO en %s · %s (%s): decide Fernando si se hereda. Si se hereda, se '
              'regenera %s con scripts/generar_escaner2_heredado.py; si no, se apunta en DIVERGENCIAS_ACEPTADAS.'
              % (_v, _pieza, _que, _h))
    eq('(B) %s: las %d piezas son las de %s (mismo árbol y mismo texto)' % (_h, _n, _v), _difs, [])

# La guarda de (B) muerde: el viejo, cambiado en una copia (una cifra de la formula y un comentario).
_orig_leer = leer


def _paridad_con(viejo_cambiado, heredado='escaner2_heredado_nube.py', viejo='moloka_escaner_nube.py'):
    global leer
    texto = _orig_leer(viejo)
    try:
        leer = lambda r: viejo_cambiado(texto) if r == viejo else _orig_leer(r)  # noqa: E731
        return paridad(heredado, viejo)[1]
    finally:
        leer = _orig_leer


if os.path.exists(os.path.join(AQUI, 'moloka_escaner_nube.py')):
    _TX = _orig_leer('moloka_escaner_nube.py')
    _a = "def decision_de(margen):\n    if margen is None: return 'Sin datos'\n    if margen*100 >= 10: return 'COMPRAR'"
    assert _TX.count(_a) == 1
    eq('(B) 🔴 si el viejo sube el umbral de COMPRAR del 10 al 12 %, lo dice (y el heredado NO cambia)',
       _paridad_con(lambda t: t.replace(_a, _a.replace('>= 10', '>= 12'))), [('decision_de', 'el árbol cambia')])
    eq('(B) 🔴 si en el viejo solo cambia un comentario de una pieza, también lo dice',
       _paridad_con(lambda t: t.replace("# 🔒 EL VALOR Y SU ORIGEN, EN LA MISMA FUNCION.", "# EL VALOR Y SU ORIGEN.")),
       [('ORIGEN_IVA_FICHA', 'el texto o los comentarios cambian')])
    eq('(B) 🔴 si el viejo cambia un ancho de la Celda 9, lo dice',
       _paridad_con(lambda t: t.replace("anchos = {'Nombre':50,", "anchos = {'Nombre':55,")),
       [('excel_del_viejo', 'el árbol de la Celda 9 cambia')])
    eq('(B) …y una línea nueva en el viejo, lejos de las piezas, no molesta (todo se busca por nombre)',
       _paridad_con(lambda t: '# linea nueva arriba del todo\n' + t), [])


def _tocados_con(cambio, heredado='escaner2_heredado_nube.py'):
    global leer
    texto = _orig_leer(heredado)
    try:
        leer = lambda r: cambio(texto) if r == heredado else _orig_leer(r)  # noqa: E731
        return tocados_a_mano(heredado)
    finally:
        leer = _orig_leer


_HT = _orig_leer('escaner2_heredado_nube.py')
eq('(B) 🔴 si alguien toca a mano el heredado (el umbral de COMPRAR), lo dice, con o sin viejo',
   _tocados_con(lambda t: t.replace("if margen*100 >= 10: return 'COMPRAR'", "if margen*100 >= 12: return 'COMPRAR'")),
   ['decision_de'])
eq('(B) 🔴 …y un ancho de la Celda 9 tocado a mano, también',
   _tocados_con(lambda t: t.replace("anchos = {'Nombre':50,", "anchos = {'Nombre':49,")), ['excel_del_viejo'])
eq('(B) …y una línea en blanco de más entre piezas no es tocar una pieza',
   _tocados_con(lambda t: t.replace('\n\n\n# ── ORIGEN', '\n\n\n\n# ── ORIGEN', 1)), [])


# ═══════════════════════════════════════════════════════════════════════════════
# LOS DATOS INVENTADOS DE (C), (D) y (E): ni precios reales ni costes (el repo es publico).
# Cubren todas las hojas y todas las ramas de la Celda 9: los cuatro paises con COMPRAR, VALORAR, NO
# COMPRAR y «Sin datos» (y uno sin pais), ISD de Francia, IVA de ficha y asumido, un producto sin PA (la
# fila no echa la cuenta), enlace de OcioStock, «Coincide» NO, «AMBIGUO», caja INCOHERENTE, precio por lote
# (con uno incoherente), las cuatro listas de Descartados, Ambiguos, Sin_rank y Chase_manual.
# ═══════════════════════════════════════════════════════════════════════════════
CATALOGO_PRUEBA = [{'ean': '8400000000017', 'iva_pct': 0.10, 'stock_moloka': 3, 'stock_fba': 5},
                   {'ean': '8400000000024', 'iva_pct': None, 'stock_moloka': 0, 'stock_fba': 2}]


def _pais(dec, precio, rank, margen, fee=3.1, ref=15.0, iva=0.21):
    return {'rank_act': rank, 'rank90': (rank * 2 if rank else rank), 'vendidos': 40, 'precio': precio, 'canal': 'BB-FBA',
            'n_of': 7, 'ref_pct': ref, 'fee': fee, 'iva': iva, 'decision': dec, 'margen': margen}


def datos_prueba(vacio=False):
    if vacio:
        return {'registros': [], 'problematicos': [], 'no_encontrados': [], 'chase_sueltos': [], '_dups': [],
                'ambiguos': [], 'sin_rank': [], 'chase_pendientes': [], 'cotejo_info': {}, 'PROVEEDOR': 'HEO'}
    registros = [
        {'nombre': 'Producto de prueba Alfa', 'ean': '8400000000017', 'asin': 'B0PRUEBA01', 'marca': 'Marca Uno',
         'core': '8400000000017', '_pa_efectivo': 4.0, 'ambiguo': False, 'titulo_amz': 'Titulo de prueba Alfa',
         'coincide': 'SI', 'coherencia_caja': None, 'url': '', 'volumen': None, '_margen_es': 0.2,
         '_paises_calc': {'ES': _pais('COMPRAR', 19.99, 1200, 0.2, iva=0.10), 'IT': _pais('VALORAR', 17.5, 3000, 0.05, iva=0.22),
                          'FR': _pais('NO COMPRAR', 12.0, 50000, -0.1, iva=0.20),
                          'DE': _pais('Sin datos', None, None, None, fee=None, ref=None, iva=0.19)}},
        {'nombre': 'Producto de prueba Beta (caja x6)', 'ean': '8400000000024 C6', 'asin': 'B0PRUEBA02', 'marca': 'Marca Dos',
         'core': '8400000000024', '_pa_efectivo': 2.5, 'ambiguo': True, 'titulo_amz': 'Titulo de prueba Beta',
         'coincide': 'NO', 'coherencia_caja': 'INCOHERENTE: la caja sale a 1.00 €/ud y la unidad suelta del MISMO EAN esta a 5.00 €',
         'url': 'https://example.invalid/ficha/beta', 'volumen': {'pa': 2.0, 'uds': 12}, '_margen_es': 0.08,
         '_paises_calc': {'ES': _pais('VALORAR', 9.99, 8000, 0.08), 'FR': _pais('COMPRAR', 11.5, 9000, 0.15, iva=0.20)}},
        {'nombre': 'Producto de prueba Gamma', 'ean': '8400000000031', 'asin': 'B0PRUEBA03', 'marca': 'Marca Tres',
         'core': '8400000000031', '_pa_efectivo': None, 'ambiguo': False, 'titulo_amz': '', 'coincide': None,
         'coherencia_caja': None, 'url': 'https://example.invalid/ficha/gamma', 'volumen': {'pa': 9.0, 'uds': 6},
         '_margen_es': None, '_paises_calc': {'ES': _pais('Sin datos', 14.0, 500, None)}},
        {'nombre': 'Producto de prueba Delta', 'ean': '8400000000048', 'asin': 'B0PRUEBA04', 'marca': 'Marca Uno',
         'core': '8400000000048', '_pa_efectivo': 6.0, 'ambiguo': False, 'titulo_amz': 'Titulo Delta', 'coincide': None,
         'coherencia_caja': None, 'url': '', 'volumen': {'pa': 7.0, 'uds': 24}, '_margen_es': -0.3,
         '_paises_calc': {'ES': _pais('NO COMPRAR', 8.0, 70000, -0.3), 'IT': _pais('NO COMPRAR', 8.5, 60000, -0.2, iva=0.22)}},
    ]
    return {
        'registros': registros,
        'problematicos': [{'EAN': '12345', 'Cabecera': 'Prueba forma rara', 'Motivo': 'EAN forma rara (len=5)'}],
        'no_encontrados': [{'EAN': '8400000000055', 'Cabecera': 'Prueba no aparece', 'Motivo': 'No aparece en ningún CSV'}],
        'chase_sueltos': [{'EAN': '8400000000062', 'Cabecera': 'Prueba Chase', 'Motivo': 'Chase SUELTO descartado (solo se compra en caja de 6)'}],
        '_dups': [{'EAN': '8400000000017', 'Cabecera': 'Prueba duplicada', 'Motivo': 'Duplicado del proveedor: me quedo con 4.0 (esta venía a 5.0)'}],
        'ambiguos': [{'EAN': '8400000000024 C6', 'asin_elegido': 'B0PRUEBA02'}, {'EAN': '8400000000079', 'asin_elegido': None}],
        'sin_rank': [{'ean_in': '8400000000086', 'asin': 'B0PRUEBA08', 'fila': {'nombre': 'Prueba sin rank'}, 'r_act': None, 'r_90': 88000}],
        'chase_pendientes': [{'nombre': 'Prueba caja con chase Surtido (6)', 'producto_heo': 'FK00000', 'ean_caja': '9990000000011',
                              'precio_caja': 30.0, 'estado': 'disponible', 'imagen': 'https://example.invalid/img.jpg',
                              'link_amazon': 'https://example.invalid/buscar'},
                             {'nombre': 'Prueba caja sin precio', 'producto_heo': 'FK00001', 'ean_caja': '', 'precio_caja': None,
                              'estado': 'agotado', 'imagen': '', 'link_amazon': ''}],
        'cotejo_info': {'8400000000017': {'veredicto': 'OK', 'detalle': '2/2 casan; entre esos, mejor rank (casa: alfa)'},
                        '8400000000024 C6': {'veredicto': '⚠ DUDOSO', 'detalle': 'ninguno pasa el filtro — REVISAR'},
                        '8400000000031': {'veredicto': 'n/d', 'detalle': 'n/d: sin texto que comparar'}},
        'PROVEEDOR': 'HEO'}


def _guardar(wb):
    b = io.BytesIO()
    wb.save(b)
    return b.getvalue()


M = e2.cargar_motor()
M.poner_catalogo_propio(CATALOGO_PRUEBA)


def libro_nuevo(datos, ruta=e2.RUTA_MOTOR):
    return _guardar(e2.escribir_celda9(copy.deepcopy(datos), M, ruta))


def libro_viejo(datos, viejo='moloka_escaner_nube.py'):
    """El libro de la Celda 9 DEL VIEJO: su bloque, sacado de su fichero y ejecutado con sus piezas (como
    lo hacia el escaner 2 hasta el B6). Solo mientras el viejo exista."""
    texto = leer(viejo)
    _lin, sentencias = bloque_celda9(texto)
    M_v = e2.cargar_motor(os.path.join(AQUI, viejo))
    M_v.poner_catalogo_propio(CATALOGO_PRUEBA)
    ns = e2.sacar_piezas(os.path.join(AQUI, viejo), ('pct_comision_celda', 'en_bd_txt'), (), base=M_v._ns)
    ns.update(copy.deepcopy(datos))
    with redirect_stdout(io.StringIO()):
        exec(compile(ast.Module(body=sentencias, type_ignores=[]), viejo, 'exec'), ns)
    return _guardar(ns['wb'])


REFERENCIA = os.path.join(AQUI, 'excel_referencia_escaner2.xlsx')
if '--regenerar-referencia' in sys.argv:
    with open(REFERENCIA, 'wb') as fh:
        fh.write(libro_nuevo(datos_prueba()))
    print('REGENERADO %s. Súbelo solo si Fernando ha decidido cambiar el Excel.' % REFERENCIA)
    sys.exit(0)

# ═══════════════════════════════════════════════════════════════════════════════
print('\n(C) paridad del Excel: la Celda 9 del viejo y la copia, celda a celda (mientras exista el viejo)')
HOJAS = ['Análisis', 'Descartados', 'Ambiguos', 'Sin_rank', 'Precio por lote', 'Chase_manual']
_NUEVO = libro_nuevo(datos_prueba())
if os.path.exists(os.path.join(AQUI, 'moloka_escaner_nube.py')):
    for _escena, _datos in (('llena', datos_prueba()), ('vacía', datos_prueba(vacio=True))):
        _n, _d = HU.celda_a_celda(libro_viejo(_datos), libro_nuevo(_datos))
        print('    escena %s: %d celdas comparadas, %d diferencias' % (_escena, _n, len(_d)))
        for x in _d[:10]:
            print('      ' + x)
        eq('(C) escena %s: el libro del viejo y el de la copia son idénticos celda a celda' % _escena, (_n > 0, _d), (True, []))
else:
    print('    (C) SALTADO: moloka_escaner_nube.py ya no está en el repo; queda (D), el libro de referencia')

# ═══════════════════════════════════════════════════════════════════════════════
print('\n(D) el libro de referencia del repo (datos inventados), regenerado y comparado celda a celda')
from openpyxl import load_workbook  # noqa: E402
_ref = open(REFERENCIA, 'rb').read()
_wr = load_workbook(io.BytesIO(_ref))
eq('(D) la referencia tiene las seis hojas del viejo, en su orden', _wr.sheetnames, HOJAS)
eq('(D) …y no está vacía: filas por hoja', [_wr[h].max_row for h in HOJAS], [17, 5, 3, 2, 6, 3])
_n, _d = HU.celda_a_celda(_ref, _NUEVO)
print('    %d celdas comparadas, %d diferencias' % (_n, len(_d)))
for x in _d[:10]:
    print('      ' + x)
eq('(D) 🔴 el libro que sale hoy es IDÉNTICO a la referencia celda a celda', (_n > 500, _d), (True, []))
_prec = [c.value for c in _wr['Análisis']['E'][1:] if isinstance(c.value, (int, float))]
eq('(D) la referencia no lleva costes reales: los PA son los inventados de este banco', sorted(set(_prec)), [2.5, 4.0, 6.0])

# ═══════════════════════════════════════════════════════════════════════════════
print('\n(E) la comparación muerde: una fórmula de «Análisis» y un ancho cambiados en una copia del heredado')
_H = leer(e2.RUTA_MOTOR)
_ROI = """f"={L['Beneficio (€)']}{r}/{L['PA (€)']}{r}" if hay_cuenta else None,"""
_ANCHO = "anchos = {'Nombre':50,"
assert _H.count(_ROI) == 1 and _H.count(_ANCHO) == 1
with tempfile.NamedTemporaryFile('w', suffix='.py', delete=False, encoding='utf-8') as _t:
    _t.write(_H.replace(_ROI, _ROI.replace('}{r}/{L', '}{r}*{L')).replace(_ANCHO, "anchos = {'Nombre':49,"))
try:
    _ROTO = libro_nuevo(datos_prueba(), ruta=_t.name)
finally:
    os.unlink(_t.name)
_n, _d = HU.celda_a_celda(_ref, _ROTO, max_difs=1000)
eq('(E) 🔴 (D) contra la copia rota: rojo, y dice el ancho y la fórmula del ROI',
   (bool(_d), any('anchos' in x and "'A'" in x for x in _d), any('Análisis · S' in x and '*' in x for x in _d)), (True, True, True))
if os.path.exists(os.path.join(AQUI, 'moloka_escaner_nube.py')):
    _n, _d = HU.celda_a_celda(libro_viejo(datos_prueba()), _ROTO, max_difs=1000)
    eq('(E) 🔴 (C) contra la copia rota: rojo también', bool(_d), True)

print()
if fallos:
    print('ROJO: %d comprobaciones fallan: %s' % (len(fallos), ', '.join(fallos)))
    sys.exit(1)
print('VERDE: el escáner 2 no lee el viejo, sus piezas heredadas son las del viejo y su Excel es el del viejo, celda a celda.')
