# -*- coding: utf-8 -*-
"""Banco del ESCANEO PRO DE DBLINE (encargo DB3, 09-oct-2026): la parte pura y los DOS programas de verdad, con dobles.
Sin red, sin secretos, sin produccion. 🔴 Repo PUBLICO (Seguridad 21): la foto de DBLine de aqui es INVENTADA (codigos
ZZ00xx, precios redondos, EAN de mentira con su digito de control bueno, marcas reales sin precio).

PARTE A · LO PURO (escaner2_dbline_pro.py), en este proceso:
  A1  el filtro de marca por modo: el del director (el nombre de la regla DENTRO del de la fila, como el viejo), todas
      (las sin marca tambien) y elegidas (coincidencia EXACTA, como HEO);
  A2  la foto: solo lo que tiene unidades hoy, sin chase suelto, «w/Chase» y TELEFONARE dentro, el EAN de 11 cifras
      (UPC con su cero) dentro y el vacio fuera, sin precio fuera, PA = el precio vigente (la promo del dia vale, la
      caducada desde la foto no), una fila por EAN con su codigo mas barato (a igualdad, el codigo menor); crudo =
      previas + foto, y si no, lo dice;
  A3  🔴 mutantes de lo puro: cada uno, A2 en rojo.

PARTE B · LOS PROGRAMAS (escaner2_dbline_barrido.py y escaner2_dbline_cruce.py), cada uno en su proceso, contra un
`supabase` en memoria:
  B1  [sin_llave] sin la llave de servicio no nace ningun cliente; [ofertas] el modo elegidas con ofertas, tampoco;
  B2  [marcas] foto → barrido → 'esperando_csv' con su foto, eans.txt y barrido.json → los CUATRO CSV (ES, IT, FR, DE,
      con la CABECERA REAL del Visualizador) → cruce → 'lista', cuadra en los cuatro, PA = el vigente, y el Excel: las
      hojas del PRO de HEO en su orden (sin «Chase_manual»), «Análisis» con «Ventas» y «Ficha compartida» y NADA MAS
      al final, y la biblioteca puede leer su ruta;
  B3  🔴 el registro (repo PUBLICO) no lleva ni un EAN, precio, codigo de DBLine ni nombre;
  B4  nada escribe fuera de escaner2_* ni del almacen dbline/<pasada>/; disp_*, reglas_director y productos acaban
      como empezaron;
  B5  [todas] y [elegidas]: su foto; [ajena] una foto que no es la de su ultima pasada aplicada: ROJO y la pasada
      'fallida' sin foto.
"""
import base64
import csv
import io
import json
import os
import subprocess
import sys
import tempfile
import types
from datetime import date

AQUI = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, AQUI)
fallos = []


def eq(nombre, obtenido, esperado):
    ok = obtenido == esperado
    if not ok:
        fallos.append(nombre)
    print(('OK ' if ok else 'XX ') + nombre + ('' if ok else '   got=%r exp=%r' % (obtenido, esperado)))


# ═══════════════════════════════════════════════════════════════════════════════
# EL DOBLE DE SUPABASE (el de test_escaner2_ociostock_pro.py, copiado: cada banco, el suyo)
# ═══════════════════════════════════════════════════════════════════════════════
class _Resp:
    def __init__(self, data, count=None):
        self.data, self.count = data, count


class _Consulta:
    def __init__(self, bd, tabla):
        self.bd, self.tabla = bd, tabla
        self.op, self.filtros, self.payload = 'select', [], None
        self.cuenta, self.orden, self.rango, self.tope = False, None, None, None

    def select(self, _cols='*', count=None):
        self.op, self.cuenta = 'select', (count == 'exact')
        return self

    def insert(self, filas):
        self.op, self.payload = 'insert', filas if isinstance(filas, list) else [filas]
        return self

    def update(self, cambios):
        self.op, self.payload = 'update', cambios
        return self

    def eq(self, k, v):
        self.filtros.append((k, v))
        return self

    def order(self, col, desc=False):
        self.orden = (col, desc)
        return self

    def range(self, a, b):
        self.rango = (a, b)
        return self

    def limit(self, n):
        self.tope = n
        return self

    def execute(self):
        self.bd.setdefault('ops', []).append([self.bd.get('programa'), self.op, self.tabla])
        filas = self.bd['tablas'].setdefault(self.tabla, [])
        if self.op == 'insert':
            for f in self.payload:
                filas.append(json.loads(json.dumps(f, default=str)))
            return _Resp(self.payload)
        sel = [f for f in filas if all(str(f.get(k)) == str(v) for k, v in self.filtros)]
        if self.op == 'update':
            for f in sel:
                f.update(json.loads(json.dumps(self.payload, default=str)))
            return _Resp(sel)
        if self.orden:
            sel = sorted(sel, key=lambda f: (f.get(self.orden[0]) is None, str(f.get(self.orden[0]))),
                         reverse=self.orden[1])
        n = len(sel)
        if self.rango:
            sel = sel[self.rango[0]:self.rango[1] + 1]
        if self.tope is not None:
            sel = sel[:self.tope]
        return _Resp(sel, n if self.cuenta else None)


class _Cubo:
    def __init__(self, bd, cubo):
        self.bd, self.cubo = bd, cubo

    def _objs(self):
        return self.bd['storage'].setdefault(self.cubo, {})

    def upload(self, ruta, datos, _opciones=None):
        self.bd.setdefault('subidas', []).append(ruta)
        self._objs()[ruta] = base64.b64encode(datos if isinstance(datos, bytes) else datos.encode()).decode()
        return {'Key': ruta}

    def download(self, ruta):
        if ruta not in self._objs():
            raise RuntimeError('404 Object not found: %s/%s' % (self.cubo, ruta))
        return base64.b64decode(self._objs()[ruta])

    def list(self, carpeta=None, opciones=None):
        op = dict({'limit': 100, 'offset': 0}, **(opciones or {}))
        pref = (carpeta or '').rstrip('/') + '/'
        nombres = sorted(r[len(pref):] for r in self._objs() if r.startswith(pref) and '/' not in r[len(pref):])
        trozo = nombres[op['offset']:op['offset'] + op['limit']]
        return [{'name': n, 'created_at': '2026-10-09T09:%02d:00Z' % i, 'metadata': {'size': 1}}
                for i, n in enumerate(trozo)]


class _Cliente:
    def __init__(self, bd):
        self.bd = bd
        self.storage = types.SimpleNamespace(from_=lambda cubo: _Cubo(bd, cubo))

    def table(self, nombre):
        return _Consulta(self.bd, nombre)


def hijo(ruta_estado, programa):
    with open(ruta_estado, encoding='utf-8') as fh:
        bd = json.load(fh)
    bd['programa'] = programa
    clientes = []

    def crear(url, llave):
        clientes.append(llave)
        return _Cliente(bd)

    sys.modules['supabase'] = types.ModuleType('supabase')
    sys.modules['supabase'].create_client = crear
    import runpy
    sys.argv = [programa]
    codigo = 0
    try:
        runpy.run_path(os.path.join(AQUI, programa), run_name='__main__')
    except SystemExit as ex:
        codigo = ex.code if isinstance(ex.code, int) else (0 if ex.code is None else 1)
    finally:
        with open(ruta_estado, 'w', encoding='utf-8') as fh:
            json.dump(bd, fh, default=str)
        print('CLIENTES_CREADOS=%d' % len(clientes))
    sys.exit(codigo)


if __name__ == '__main__' and len(sys.argv) >= 4 and sys.argv[1] == '--hijo':
    hijo(sys.argv[2], sys.argv[3])


def correr(ruta_estado, programa, env_extra):
    env = {k: v for k, v in os.environ.items()
           if k not in ('SUPABASE_SERVICE_KEY', 'PASADA', 'GITHUB_RUN_ID', 'MODO_BARRIDO', 'MARCAS_ELEGIDAS',
                        'OFERTAS_ELEGIDAS')}
    env.update(env_extra, PYTHONIOENCODING='utf-8', SUPABASE_URL='https://doble.invalid')
    p = subprocess.run([sys.executable, '-u', os.path.abspath(__file__), '--hijo', ruta_estado, programa],
                       capture_output=True, text=True, encoding='utf-8', errors='replace', env=env, cwd=AQUI)
    return p.returncode, (p.stdout or '') + (p.stderr or '')


import escaner2_motor as e2  # noqa: E402
import escaner2_dbline_pro as edb  # noqa: E402
import escaner2_heredado_pro as pro  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

M = e2.cargar_motor()


def ean13(cuerpo):
    return cuerpo + M._chk13(cuerpo)


# ═══════════════════════════════════════════════════════════════════════════════
# LA FOTO DE DBLINE, INVENTADA (la forma de disp_estado tras el encargo DB2)
# ═══════════════════════════════════════════════════════════════════════════════
E1 = ean13('840000000201')    # Funko, sin promo: PA = Prezzo (7,00)
E2 = ean13('840000000202')    # Funko con promo VIGENTE (fin = hoy en A; 2099 en B): PA = la promo (5,00)
E3 = ean13('840000000203')    # Funko con promo CADUCADA desde la foto: PA = Prezzo (9,00), no la promo (4,00)
E4 = ean13('840000000204')    # Pyramid
E5 = ean13('840000000205')    # Bandai: fuera de la lista del director
E6 = ean13('840000000206')    # Funko chase suelto (la guarda de DB2: regla)
E7 = ean13('840000000207')    # Funko sin unidades
E8 = ean13('840000000208')    # Funko PRENOTAZIONE, sin unidades (como las 975 reales)
U12 = ean13('001234567890')[1:]   # Ultra-Pro: el UPC-A de 12, con su cero delante y su digito de control bueno…
U11 = U12[1:]                     # …y como lo trae el fichero, de 11 (sin el cero); DB2 lo guarda con ean_core = U12
E11 = ean13('840000000211')   # Funko con Prezzo a 0: sin precio de compra
E12 = ean13('840000000212')   # dos codigos: 6,00 y 5,50 → gana el de 5,50
E13 = ean13('840000000213')   # dos codigos al mismo precio → el codigo menor
E14 = ean13('840000000214')   # Funko «w/Chase»: la figura normal, entra
E15 = ean13('840000000215')   # Funko TELEFONARE con unidades: entra
DISP = '00000000-0000-4000-8000-0000000db301'
OTRA = '00000000-0000-4000-8000-0000000db302'
HOY = date(2026, 10, 9)


def fila(cod, ean, marca, unidad, catalogo=None, oferta=False, fin=None, disp=True, pre=False, regla=None, core=None,
         chase=False, nombre=None, nota=None, pasada=DISP):
    return {'proveedor': 'DBLINE', 'producto_prov': cod, 'ean_original': ean,
            'ean_core': core if core is not None else (ean or None), 'marca': marca,
            'nombre': nombre or 'Figura inventada %s' % cod, 'categoria': 'Figure', 'precio_unidad': unidad,
            'precio_catalogo': unidad if catalogo is None else catalogo, 'en_oferta': oferta, 'fin_oferta': fin,
            'precio_escalon': None, 'uds_escalon': None, 'precio_pa': None, 'es_caja': False, 'uds_caja': None,
            'es_chase': chase, 'disponible': disp, 'disponibilidad': nota or ('PRENOTAZIONE' if pre else None),
            'preorder': pre, 'regla': regla, 'pasada_id': pasada, 'ausencias': 0}


def foto_disp(fin_vigente='2026-10-09', fin_caducada='2026-10-08', pasada_ajena=False):
    return [
        fila('ZZ0001', E1, 'Funko', 7.0),
        fila('ZZ0002', E2, 'Funko', 5.0, 8.0, oferta=True, fin=fin_vigente),
        fila('ZZ0003', E3, 'Funko', 4.0, 9.0, oferta=True, fin=fin_caducada),
        fila('ZZ0004', E4, 'Pyramid', 3.0),
        fila('ZZ0005', E5, 'Bandai', 20.0),
        fila('ZZ0006', E6, 'Funko', 7.0, regla='chase_suelto', chase=True, nombre='Figura inventada (Chase)'),
        fila('ZZ0007', E7, 'Funko', 7.0, disp=False),
        fila('ZZ0008', E8, 'Funko', 7.0, disp=False, pre=True),
        fila('ZZ0009', '', 'Funko', 7.0, regla='ean_forma_rara', core=''),
        fila('ZZ0010', U11, 'Ultra-Pro', 2.0, core=U12),
        fila('ZZ0011', E11, 'Funko', 0.0),
        fila('ZZ0013', E12, 'Funko', 6.0),
        fila('ZZ0012', E12, 'Funko', 5.5),
        fila('ZZ0015', E13, 'Funko', 6.0),
        fila('ZZ0014', E13, 'Funko', 6.0, pasada=OTRA if pasada_ajena else DISP),
        fila('ZZ0016', E14, 'Funko', 7.0, nombre='Figura inventada w/Chase'),
        fila('ZZ0017', E15, 'Funko', 7.0, nota='TELEFONARE'),
        fila('ZZ0018', '84000000011', 'Funko', 7.0, core='84000000011'),
    ]


N_SIN_EAN = 0                                   # DB2 sube todas las filas, tambien las sin EAN
N_CRUDO = len(foto_disp()) + N_SIN_EAN          # 18
DIRECTOR = ['Funko', 'Pyramid']

# ═══════════════════════════════════════════════════════════════════════════════
# PARTE A · LO PURO
# ═══════════════════════════════════════════════════════════════════════════════
print('A1 · el filtro de marca, por modo')
_q, _i = edb.filtro_marca('marcas', DIRECTOR)
eq('A1 · marcas: la regla del viejo (el nombre de la regla dentro del de la fila, sin mayúsculas)',
   [_q(m) for m in ('Funko', 'FUNKO', 'Pyramid', 'Bandai', None, 'Ultra-Pro')], [True, True, True, False, False, False])
eq('A1 · …y guarda las marcas de la regla', _i['marcas_reales'], DIRECTOR)
_q, _i = edb.filtro_marca('todas')
eq('A1 · todas: también las sin marca, y sin marcas que guardar', ([_q(m) for m in ('Bandai', None)], _i['marcas_reales']),
   ([True, True], None))
_q, _i = edb.filtro_marca('elegidas', elegidas=['Pyramid'])
eq('A1 · elegidas: coincidencia EXACTA (sin mayúsculas), como HEO', [_q(m) for m in ('PYRAMID', 'Pyramid X', ' pyramid ')],
   [True, False, True])
for _modo, _args in (('marcas', ([],)), ('elegidas', (None, [])), ('raro', ())):
    try:
        edb.filtro_marca(_modo, *_args)
        _r = 'pasa'
    except edb.FalloDBLine:
        _r = 'para'
    eq('A1 · %s sin marcas (o desconocido): para' % _modo, _r, 'para')


def comprobar_foto(mod):
    """A2 entero contra un modulo (el de verdad o un mutante). Devuelve lo que falla."""
    rotas = []

    def chk(nombre, obtenido, esperado):
        if obtenido != esperado:
            rotas.append((nombre, obtenido, esperado))
    quiere, _ = mod.filtro_marca('marcas', DIRECTOR)
    foto, apartados, cuentas = mod.construir_foto(foto_disp(), quiere, M, N_CRUDO, N_SIN_EAN, HOY, 'marcas')
    por = {f['ean_core']: f for f in foto}
    chk('foto', sorted(por), sorted([E1, E2, E3, E4, E12, E13, E14, E15]))
    chk('códigos', {k: por[k]['producto_heo'] for k in por},
        {E1: 'ZZ0001', E2: 'ZZ0002', E3: 'ZZ0003', E4: 'ZZ0004', E12: 'ZZ0012', E13: 'ZZ0014', E14: 'ZZ0016',
         E15: 'ZZ0017'})
    chk('PA: sin promo, promo del día, promo caducada', [(por[k]['precio_unidad'], por[k]['precio_catalogo'],
                                                          por[k]['en_oferta']) for k in (E1, E2, E3)],
        [(7.0, 7.0, False), (5.0, 8.0, True), (9.0, 9.0, False)])
    chk('sin cajas ni chase', {(f['es_caja'], f['uds_caja'], f['es_chase'], f['aviso_caja']) for f in foto},
        {(False, None, False, None)})
    previas = cuentas['previas']
    chk('previas', previas, {'chase_funko': 0, 'sin_gtin': 0, 'no_disponible': 2, 'marca_fuera': 2,
                             'estado_no_servible': 1, 'chase_suelto': 1, 'ean_forma_rara': 2, 'duplicado_proveedor': 2})
    chk('apartados', sorted((a['producto_heo'], a['motivo']) for a in apartados),
        sorted([('ZZ0005', 'marca_fuera'), ('ZZ0010', 'marca_fuera'), ('ZZ0006', 'chase_suelto'),
                ('ZZ0009', 'ean_forma_rara'), ('ZZ0018', 'ean_forma_rara'), ('ZZ0011', 'estado_no_servible'),
                ('ZZ0013', 'duplicado_proveedor'), ('ZZ0015', 'duplicado_proveedor')]))
    chk('cuadra', (cuentas['n_foto'], cuentas['n_previas'], cuentas['cuadra_previo']), (8, 10, True))
    chk('promos', (cuentas['n_en_oferta'], cuentas['n_promo_caducada']), (1, 1))
    _f, _a, c2 = mod.construir_foto(foto_disp(), quiere, M, N_CRUDO + 1, N_SIN_EAN, HOY, 'marcas')
    chk('fichero que no casa: no cuadra y lo dice', (c2['cuadra_previo'], 'el fichero tenía 19 filas' in
                                                     (c2['motivo_previo'] or '')), (False, True))
    return rotas


print('A2 · la foto de la pasada')
_rotas = comprobar_foto(edb)
for _n, _o, _e in _rotas:
    eq('A2 · ' + _n, _o, _e)
eq('A2 · las comprobaciones de la foto, en verde', len(_rotas), 0)
quiere_t, _ = edb.filtro_marca('todas')
_ft, _at, _ct = edb.construir_foto(foto_disp(), quiere_t, M, N_CRUDO, N_SIN_EAN, HOY, 'todas')
_pt = {f['ean_core']: f for f in _ft}
eq('A2 · todas: entran Bandai y Ultra-Pro (10), y cuadra', (len(_ft), _ct['cuadra_previo'], _ct['previas']['marca_fuera']),
   (10, True, 0))
eq('A2 · 🔴 el EAN de 11 cifras (Ultra-Pro) entra con su cero y su EAN tal como vino',
   (len(U11), U12 in _pt, _pt.get(U12, {}).get('ean_original')), (11, True, U11))
_qe, _ = edb.filtro_marca('elegidas', elegidas=['Pyramid'])
_fe, _ae, _ce = edb.construir_foto(foto_disp(), _qe, M, N_CRUDO, N_SIN_EAN, HOY, 'elegidas')
eq('A2 · elegidas: una, y la marca fuera dice «no elegida»', (
    [f['ean_core'] for f in _fe], sorted({a['detalle'].split("'")[0] for a in _ae if a['motivo'] == 'marca_fuera'})),
   ([E4], ['Marca ']))
eq('A2 · …y su puerta previa se llama «Marca no elegida»', edb.nombre_previa('marca_fuera', 'elegidas'), 'Marca no elegida')
eq('A2 · los nombres de las puertas previas de DBLine',
   [edb.nombre_previa(p, 'marcas') for p in ('no_disponible', 'estado_no_servible', 'ean_forma_rara')],
   ['Sin unidades hoy', 'Sin precio de compra', 'EAN vacío o con forma rara'])
eq('A2 · las rutas que admiten los checks de la v2',
   (edb.ruta_lista(DISP), edb.ruta_excel(DISP, OTRA, '20261009_1930')),
   ('dbline/%s/eans.txt' % DISP, 'dbline/%s/%s/Escaner2_DBLINE_20261009_1930.xlsx' % (DISP, OTRA)))
eq('A2 · precio_vigente: promo sin fecha de fin = la que dio la foto',
   edb.precio_vigente({'precio_unidad': 3.0, 'precio_catalogo': 4.0, 'en_oferta': True, 'fin_oferta': None}, HOY),
   (3.0, True, False))
# DB5: la promo vigente MAS CARA que Prezzo (7,99 frente a 7,83). La fila sale de la regla de la foto de
# disponibilidad (escaner2_dbline.precio_vigente) y el PRO compra a Prezzo, sin oferta.
import escaner2_dbline as DBF  # noqa: E402
_pu30, _of30 = DBF.precio_vigente(DBF.Decimal('7.83'), DBF.Decimal('7.99'), date(2026, 12, 31), HOY)
E30 = ean13('840000000230')
_f30, _a30, _c30 = edb.construir_foto([fila('ZZ0030', E30, 'Funko', float(_pu30), 7.83, oferta=_of30, fin=None)],
                                      lambda _m: True, M, 1, 0, HOY, 'todas')
eq('A2 · DB5: promo vigente más cara que Prezzo → PA = Prezzo, sin oferta (y nada caducado)',
   ([(f['precio_unidad'], f['precio_catalogo'], f['en_oferta']) for f in _f30], _c30['n_en_oferta'],
    _c30['n_promo_caducada'], _c30['cuadra_previo']), ([(7.83, 7.83, False)], 0, 0, True))

print('A3 · 🔴 mutantes de lo puro: cada uno pone A2 en rojo')
with io.open(os.path.join(AQUI, 'escaner2_dbline_pro.py'), encoding='utf-8') as _fh:
    FUENTE = _fh.read()
MUTANTES = [
    ('el PA es Prezzo aunque haya promo', "'precio_unidad': pa,", "'precio_unidad': _num(f.get('precio_catalogo')),"),
    ('la promo caducada sigue valiendo', 'if fin is not None and fin < hoy:', 'if False:'),
    ('la promo del día ya no vale', 'if fin is not None and fin < hoy:', 'if fin is not None and fin <= hoy:'),
    ('sin precio entra', 'if pa is None or pa <= 0:', 'if pa is None:'),
    ('gana el más caro', "if (pa, str(f.get('producto_prov'))) < (prev[1]", "if (pa, str(f.get('producto_prov'))) > (prev[1]"),
    ('el chase suelto entra', "if regla == 'chase_suelto' or f.get('es_chase'):", 'if False:'),
    ('la forma rara entra', 'if regla or not core.isdigit() or len(core) not in (12, 13):', 'if regla:'),
    ('lo que no tiene unidades entra', "if not f.get('disponible'):", 'if False:'),
    ('la promo no se marca', "'en_oferta': oferta,", "'en_oferta': None,"),
]
for _nombre, _antes, _despues in MUTANTES:
    if FUENTE.count(_antes) != 1:
        eq('A3 · el ancla del mutante «%s» es única' % _nombre, FUENTE.count(_antes), 1)
        continue
    _mod = types.ModuleType('mutante')
    exec(compile(FUENTE.replace(_antes, _despues), 'mutante', 'exec'), _mod.__dict__)
    eq('A3 · mutante «%s»: A2 en rojo' % _nombre, len(comprobar_foto(_mod)) > 0, True)

# ═══════════════════════════════════════════════════════════════════════════════
# PARTE B · LOS DOS PROGRAMAS, DE PUNTA A PUNTA
# ═══════════════════════════════════════════════════════════════════════════════
with io.open(os.path.join(AQUI, 'test_escaner2_osma_cabecera.json'), encoding='utf-8') as fh:
    CABECERA = json.load(fh)['cabecera']
COL_PAIS, COL_CAIDAS = e2.columnas_keepa()
COL_PADRE, COL_NVAR, COL_RANK = e2.columnas_ficha_compartida()
C = pro.CSV_COLS


def csv_real(pais, filas):
    """Un export del Visualizador con la CABECERA REAL (561 columnas, la del banco de OSMA) y solo lo que se lee."""
    buf = io.StringIO()
    w = csv.writer(buf)
    w.writerow(CABECERA)
    ix = {h: i for i, h in enumerate(CABECERA)}
    for asin, eans, caidas, precio, titulo in filas:
        fila_ = [''] * len(CABECERA)
        for h, v in ((C['asin'], asin), (COL_PAIS, pais), ('Título', titulo), (C['ean'], eans),
                     (C['rank'], '9000'), (C['rank90'], '11000'), (COL_CAIDAS, caidas), (C['buybox'], precio),
                     (C['es_fba'], 'yes'), (C['nuevo'], precio), (C['fba'], '3.10'), (C['compct'], '15.01 %'),
                     (C['nvar'], ''), (COL_PADRE, '')):
            fila_[ix[h]] = v
        w.writerow(fila_)
    return ('﻿' + buf.getvalue()).encode('utf-8')


ASIN1, ASIN2, ASIN3, ASIN4 = 'B0DBLN0001', 'B0DBLN0002', 'B0DBLN0003', 'B0DBLN0004'
PRODUCTOS = [{'id': 1, 'ean': E1, 'asin': ASIN1, 'nombre': 'Figura de Moloka', 'activo': True, 'es_chase': False,
              'iva_pct': 0.21, 'stock_moloka': 3}]
SEMBRADAS = ('productos', 'inventario_fba', 'disp_pasada', 'disp_estado', 'reglas_director')
# En los programas, el «hoy» es el de verdad: la promo vigente acaba en 2099 y la caducada en 2020.
FOTO_B = dict(fin_vigente='2099-12-31', fin_caducada='2020-01-01')


def estado_inicial(escena):
    return {'tablas': {
        'escaner2_parametros': [{'proveedor': 'HEO', 'umbral_caidas_30d': 6, 'paises_filtro': ['ES', 'IT', 'FR', 'DE'],
                                 'paises_calculo': ['ES', 'IT', 'FR', 'DE']},
                                # La fila de DBLine que deja la migracion de la v2: la de OcioStock con su nombre.
                                {'proveedor': 'DBLINE', 'umbral_caidas_30d': 6,
                                 'paises_filtro': ['ES', 'IT', 'FR', 'DE'], 'paises_calculo': ['ES', 'IT', 'FR', 'DE']}],
        'escaner2_pasada': [],
        'disp_pasada': [{'id': OTRA, 'proveedor': 'DBLINE', 'estado': 'aplicada', 'creada_en': '2026-10-08T07:00:00Z',
                         'n_leidas': 1, 'n_crudo': 1, 'n_sin_gtin': 0},
                        {'id': DISP, 'proveedor': 'DBLINE', 'estado': 'aplicada', 'creada_en': '2026-10-09T06:30:00Z',
                         'terminada_en': '2026-10-09T06:31:00Z', 'n_leidas': len(foto_disp()), 'n_crudo': N_CRUDO,
                         'n_sin_gtin': N_SIN_EAN, 'fichero_md5': 'f' * 32},
                        {'id': 'rechazada-al-dia', 'proveedor': 'DBLINE', 'estado': 'rechazada',
                         'creada_en': '2026-10-09T10:30:00Z'},
                        # Una foto de OcioStock mas nueva: el barrido de DBLine no la mira.
                        {'id': 'de-ociostock', 'proveedor': 'OCIOSTOCK', 'estado': 'aplicada',
                         'creada_en': '2026-10-09T11:00:00Z', 'n_leidas': 1, 'n_crudo': 1, 'n_sin_gtin': 0}],
        'disp_estado': foto_disp(pasada_ajena=(escena == 'ajena'), **FOTO_B)
        + [dict(fila('OC0001', E1, 'Funko', 1.0, pasada='de-ociostock'), proveedor='OCIOSTOCK')],
        'reglas_director': [{'proveedor': 'DBLINE', 'activo': True, 'marcas': DIRECTOR},
                            {'proveedor': 'OCIOSTOCK', 'activo': True, 'marcas': ['Otra']}],
        'productos': json.loads(json.dumps(PRODUCTOS)),
        'inventario_fba': [],
    }, 'storage': {'escaner2': {}}}


def caso(escena, env=None):
    tmp = tempfile.mkdtemp(prefix='e2db_')
    ruta = os.path.join(tmp, 'estado.json')
    inicial = estado_inicial(escena)
    with open(ruta, 'w', encoding='utf-8') as fh:
        json.dump(inicial, fh)
    cod, log = correr(ruta, 'escaner2_dbline_barrido.py', dict({'SUPABASE_SERVICE_KEY': 'svc-de-mentira',
                                                                'GITHUB_RUN_ID': '818181'}, **(env or {})))
    bd = json.load(open(ruta, encoding='utf-8'))
    return cod, log, bd, ruta, inicial


print('B1 · sin llave, o elegidas con ofertas: no se corre')
_tmp = tempfile.mkdtemp(prefix='e2db_')
_ruta = os.path.join(_tmp, 'estado.json')
json.dump(estado_inicial('marcas'), open(_ruta, 'w', encoding='utf-8'))
for _prog in ('escaner2_dbline_barrido.py', 'escaner2_dbline_cruce.py'):
    _cod, _log = correr(_ruta, _prog, {'PASADA': DISP})
    eq('B1 · %s: ROJO, la línea exacta y NINGÚN cliente' % _prog,
       (_cod, 'ESCANER2_NO_EJECUTADO: sin llave de servicio' in _log, 'CLIENTES_CREADOS=0' in _log), (1, True, True))
_cod, _log = correr(_ruta, 'escaner2_dbline_barrido.py', {'SUPABASE_SERVICE_KEY': 'x', 'MODO_BARRIDO': 'elegidas',
                                                          'MARCAS_ELEGIDAS': '["Funko"]', 'OFERTAS_ELEGIDAS': 'true'})
eq('B1 · elegidas con ofertas: ROJO, lo dice y NINGÚN cliente',
   (_cod, 'no tiene casilla de ofertas' in _log, 'CLIENTES_CREADOS=0' in _log), (1, True, True))
_cod, _log = correr(_ruta, 'escaner2_dbline_barrido.py', {'SUPABASE_SERVICE_KEY': 'x', 'MODO_BARRIDO': 'elegidas',
                                                          'MARCAS_ELEGIDAS': '["Funko", "a$b"]',
                                                          'OFERTAS_ELEGIDAS': 'false'})
eq('B1 · una marca con «$»: ROJO y NINGÚN cliente (la validación de HEO)',
   (_cod, 'selección de marcas no válida' in _log, 'CLIENTES_CREADOS=0' in _log), (1, True, True))

print('B2 · [marcas] la foto → barrido → la lista → los 4 CSV → cruce → Excel')
cod, log, bd, ruta, inicial = caso('marcas')
eq('B2 · el barrido sale en VERDE', cod, 0)
if cod != 0:
    print(log)
pas = [p for p in bd['tablas']['escaner2_pasada'] if p.get('run_id') == 818181][0]
PASADA = pas['id']
eq('B2 · 🔴 la pasada: DBLINE, modo marcas, esperando_csv, fichero 18 = previas 10 + foto 8, las marcas del director',
   (pas['proveedor'], pas['modo'], pas['estado'], pas['n_crudo'], pas['n_foto'],
    sum(pas['p_' + p] for p in e2.PUERTAS_PREVIAS), pas['n_tandas'], pas['ruta_lista'], pas['marcas'], pas['ofertas']),
   ('DBLINE', 'marcas', 'esperando_csv', 18, 8, 10, 1, 'dbline/%s/eans.txt' % PASADA, DIRECTOR, False))
_foto = [f for f in bd['tablas']['escaner2_foto'] if f['pasada_id'] == PASADA]
eq('B2 · la foto en la base: 8, y las listas de las puertas previas: 8',
   (len(_foto), len([a for a in bd['tablas']['escaner2_apartado'] if a['pasada_id'] == PASADA])), (8, 8))
_alm = bd['storage']['escaner2']
_txt = lambda k: base64.b64decode(_alm[k]).decode('utf-8')  # noqa: E731
eq('B2 · en el almacén: eans.txt y barrido.json', sorted(k.split('/')[-1] for k in _alm), ['barrido.json', 'eans.txt'])
eq('B2 · eans.txt: los ocho EAN', sorted(_txt('dbline/%s/eans.txt' % PASADA).split('\n')),
   sorted([E1, E2, E3, E4, E12, E13, E14, E15]))
_side = json.loads(_txt('dbline/%s/barrido.json' % PASADA))
eq('B2 · barrido.json: de qué foto sale y sus promos (una vigente, una caducada)',
   (_side['disp_pasada'], _side['n_en_oferta'], _side['n_promo_caducada']), (DISP, 1, 1))

_carpeta = 'dbline/%s/csv/' % PASADA
# E1, E2 y E3 se venden; E4 no se vende; los demás no están en Keepa.
_filas = [(ASIN1, E1, '30', '25.95', 'Figura uno'), (ASIN2, E2, '30', '25.95', 'Figura dos'),
          (ASIN3, E3, '30', '25.95', 'Figura tres'), (ASIN4, E4, '2', '25.95', 'Llavero cuatro')]
for _i, _p in enumerate(('es', 'it', 'fr', 'de')):
    _alm[_carpeta + '20261009-1000%02d-KeepaExport-2026-10-09-VisualizadorDeProductos (%d).csv' % (_i, _i)] = \
        base64.b64encode(csv_real(_p, _filas)).decode()
bd.pop('programa', None)
json.dump(bd, open(ruta, 'w', encoding='utf-8'), default=str)
cod2, log2 = correr(ruta, 'escaner2_dbline_cruce.py', {'SUPABASE_SERVICE_KEY': 'svc-de-mentira', 'PASADA': PASADA,
                                                        'GITHUB_RUN_ID': '828282'})
bd = json.load(open(ruta, encoding='utf-8'))
T = bd['tablas']
_alm = bd['storage']['escaner2']
eq('B2 · el cruce sale en VERDE', cod2, 0)
if cod2 != 0:
    print(log2)
cr = [c for c in T['escaner2_cruce'] if c.get('run_id') == 828282][0]
eq('B2 · 🔴 el cruce: lista, cuadra, fichero 18 = previas 10 + 8 puertas, en los CUATRO países',
   (cr['estado'], cr['cuadra'], cr['n_crudo'], cr['n_previas'], cr['n_entradas'], sum(cr['n_' + x] for x in 'abcdef'),
    cr['paises_filtro'], cr['paises_calculo'], cr['paises_usados'], cr['aviso']),
   ('lista', True, 18, 10, 8, 8, ['ES', 'IT', 'FR', 'DE'], ['ES', 'IT', 'FR', 'DE'], ['ES', 'IT', 'FR', 'DE'], None))
_res = {f['ean_core']: r for r in T['escaner2_resultado_ean'] for f in T['escaner2_foto'] if f['id'] == r['foto_id']}
_pais = {}
for _r in T['escaner2_resultado_pais']:
    _pais.setdefault(_r['resultado_ean_id'], {})[_r['pais']] = _r
eq('🔴 B2 · el PA de cada uno en los cuatro países: Prezzo 7, la promo vigente 5, la caducada con Prezzo 9',
   [(sorted(_pais[_res[k]['id']]), {r['pa'] for r in _pais[_res[k]['id']].values()}) for k in (E1, E2, E3)],
   [(['DE', 'ES', 'FR', 'IT'], {7.0}), (['DE', 'ES', 'FR', 'IT'], {5.0}), (['DE', 'ES', 'FR', 'IT'], {9.0})])
eq('B2 · E1 COMPRAR, E4 no se vende (puerta c) y E15 no está en Keepa (puerta a)',
   (_res[E1]['puerta'], _res[E4]['puerta'], _res[E15]['puerta']), ('f', 'c', 'a'))

_xl = [k for k in _alm if k.startswith('dbline/%s/%s/Escaner2_DBLINE_' % (PASADA, cr['id']))]
eq('B2 · el Excel, en su carpeta y con su nombre (el que guarda el cruce)', (len(_xl), cr['ruta_excel'] == _xl[0]),
   (1, True))
_wb = load_workbook(io.BytesIO(base64.b64decode(_alm[_xl[0]])))
eq('B2 · las hojas del PRO de HEO, en su orden (sin «Chase_manual», que es de HEO)', _wb.sheetnames,
   ['Análisis', 'Descartados', 'Ambiguos', 'Sin_rank', 'Precio por lote', 'Resumen', 'Comparación', 'Varias fichas',
    'Puertas', 'Puertas previas'])
_cab = [c.value for c in _wb['Análisis'][1]]
eq('🔴 B2 · «Análisis»: las del PRO de HEO (con «Ventas») y «Ficha compartida» al final, NADA MÁS',
   (_cab, 'Ventas' in _cab), (e2.columnas_analisis() + [e2.COLUMNA_FICHA_COMPARTIDA], True))
_ix = {n: _cab.index(n) for n in _cab}
_an = list(_wb['Análisis'].iter_rows(min_row=2, values_only=True))
_por = {}
for _r in _an:
    _por.setdefault(_r[_ix['EAN']], []).append(_r)
eq('B2 · «Análisis»: los que se venden, una fila por país (los cuatro)',
   sorted((k, sorted(r[_ix['País']] for r in v)) for k, v in _por.items()),
   sorted((k, ['DE', 'ES', 'FR', 'IT']) for k in (E1, E2, E3)))
eq('🔴 B2 · «Análisis»: el PA de cada uno (7 · 5 la promo · 9 la caducada)',
   [{r[_ix['PA (€)']] for r in _por[k]} for k in (E1, E2, E3)], [{7.0}, {5.0}, {9.0}])
_pu = list(_wb['Puertas'].iter_rows(values_only=True))
eq('B2 · «Puertas»: las columnas del PRO de HEO (con las de caja, vacías) y una fila por EAN de la foto',
   (list(_pu[0])[-5:], {r[-5:] for r in _pu[1:]}, len(_pu) - 1),
   (edb.COLUMNAS_CAJA, {(None, None, None, None, None)}, 8))
_resu = {r[0]: r[1] for r in _wb['Resumen'].iter_rows(values_only=True)}
eq('B2 · el Resumen: modo, marcas del director, cuadre, promos y las puertas previas con nombre de DBLine',
   (_resu.get('Modo del barrido'), _resu.get('Marcas del director'), _resu.get('Cuadra'),
    _resu.get('Puerta previa · Sin precio de compra'), _resu.get('Puerta previa · Sin unidades hoy'),
    _resu.get('Catálogo crudo de DBLine (el fichero)'), _resu.get('Promos caducadas desde la foto (con el precio normal)')),
   ('marcas de siempre', ', '.join(DIRECTOR), 'SÍ', 1, 2, 18, 1))

print('B3 · 🔴 el registro (repo PÚBLICO) no lleva ni un EAN, precio, código de DBLine ni nombre')
_prohibido = [E1, E2, E3, E4, E5, E6, E7, E8, U11, E11, E12, E13, E14, E15, ASIN1, ASIN2, ASIN3, ASIN4, '25.95',
              '7.0', '5.5', 'ZZ00', 'Figura inventada', 'Figura uno', 'Llavero']
eq('B3 · ni en el barrido ni en el cruce (si sale algo, la línea que lo lleva)',
   [(x, [ln for ln in (log + log2).splitlines() if x in ln][:2]) for x in _prohibido if x in log or x in log2], [])

print('B4 · solo se escribe en escaner2_* y en la carpeta de la pasada')
eq('B4 · las tablas escritas', sorted({t for _p, op, t in bd['ops'] if op != 'select'}),
   ['escaner2_apartado', 'escaner2_cruce', 'escaner2_foto', 'escaner2_pasada', 'escaner2_resultado_ean',
    'escaner2_resultado_pais'])
eq('B4 · y las sembradas acaban como empezaron',
   [t for t in SEMBRADAS if bd['tablas'][t] != json.loads(json.dumps(inicial['tablas'][t]))], [])
eq('B4 · en el almacén, solo dbline/<pasada>/', sorted({'/'.join(s.split('/')[:2]) for s in bd.get('subidas', [])}),
   ['dbline/' + PASADA])

print('B5 · los otros modos, y una foto que no es la suya')
cod, log, bd, _r, _i0 = caso('todas', {'MODO_BARRIDO': 'todas'})
_p = bd['tablas']['escaner2_pasada'][0]
eq('B5 · [todas] VERDE, 10 en la foto (Bandai y Ultra-Pro), sin marcas guardadas',
   (cod, _p['modo'], _p['n_foto'], _p['p_marca_fuera'], _p['marcas']), (0, 'todas', 10, 0, None))
cod, log, bd, _r, _i0 = caso('elegidas', {'MODO_BARRIDO': 'elegidas', 'MARCAS_ELEGIDAS': '["Pyramid"]',
                                          'OFERTAS_ELEGIDAS': 'false'})
_p = bd['tablas']['escaner2_pasada'][0]
eq('B5 · [elegidas] VERDE, una en la foto, con su lista y ofertas a false',
   (cod, _p['modo'], _p['n_foto'], _p['marcas'], _p['ofertas']), (0, 'elegidas', 1, ['Pyramid'], False))
cod, log, bd, _r, _i0 = caso('ajena')
_p = bd['tablas']['escaner2_pasada'][0]
eq('B5 · 🔴 [ajena] una fila de otra pasada de disponibilidad: ROJO, la pasada fallida con el motivo y SIN foto',
   (cod, _p['estado'], 'no es la de su última pasada aplicada' in (_p['motivo_fallo'] or ''),
    bd['tablas'].get('escaner2_foto', [])), (1, 'fallida', True, []))

print()
if fallos:
    print('FALLAN %d: %s' % (len(fallos), fallos))
    sys.exit(1)
print('TODO OK')
