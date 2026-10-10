# -*- coding: utf-8 -*-
"""Banco del ESCANEO PRO DE BEMS (encargo BE4, 10-oct-2026): la parte pura y los DOS programas de verdad, con dobles.
Sin red, sin secretos, sin produccion. 🔴 Repo PUBLICO (Seguridad 21): la foto de BEMS de aqui es INVENTADA
(referencias 9000xx, precios redondos, EAN de mentira con su digito de control bueno, marcas reales sin precio).

PARTE A · LO PURO (escaner2_bems_pro.py), en este proceso:
  A1  el filtro de marca por modo: la lista fija (el nombre de la lista DENTRO del de la fila: «Funko Tees» casa con
      «Funko»), todas (las sin fabricante tambien) y elegidas (coincidencia EXACTA, como HEO);
  A2  la foto: solo lo que tiene stock hoy; SIN regla de chase («w/Chase» y «… Chase» al final, dentro, como figura);
      la blind box a su precio de caja, sin dividir; el EAN vacio o raro fuera; sin precio (vacio o 0) fuera; PA = el
      PA del CSV; una fila por EAN con su referencia mas barata (a igualdad, la menor); crudo = previas + foto, y si
      no, lo dice; el aviso del dia del CSV;
  A3  🔴 mutantes de lo puro: cada uno, A2 en rojo.

PARTE B · LOS PROGRAMAS (escaner2_bems_barrido.py y escaner2_bems_cruce.py), cada uno en su proceso, contra un
`supabase` en memoria:
  B1  [sin_llave] sin la llave de servicio no nace ningun cliente; [ofertas] el modo elegidas con ofertas, tampoco;
  B2  [marcas] foto → barrido → 'esperando_csv' con su foto, eans.txt y barrido.json (el dia del CSV) → los CUATRO
      CSV (ES, IT, FR, DE, con la CABECERA REAL del Visualizador) → cruce → 'lista', cuadra en los cuatro, PA = el del
      CSV, y el Excel: las hojas del PRO de HEO en su orden (sin «Chase_manual»), «Análisis» con «Ventas» y «Ficha
      compartida» y NADA MAS al final, el dia del CSV y su aviso en el Resumen; el crudo son los ARTICULOS (no las
      lineas con las repetidas);
  B3  🔴 el registro (repo PUBLICO) no lleva ni un EAN, precio, referencia de BEMS ni nombre, y SI el dia del CSV;
  B4  nada escribe fuera de escaner2_* ni del almacen bems/<pasada>/; disp_*, escaner2_parametros, reglas_director y
      productos acaban como empezaron;
  B5  [todas] y [elegidas]: su foto; [ajena] una foto que no es la de su ultima pasada aplicada, [sin_foto] ninguna
      pasada aplicada y [descuadre] lineas ≠ articulos + repetidas: ROJO y la pasada 'fallida' sin foto.
"""
import base64
import csv
import io
import json
import os
import subprocess
import sys
import tempfile
import types
from datetime import date

AQUI = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, AQUI)
fallos = []


def eq(nombre, obtenido, esperado):
    ok = obtenido == esperado
    if not ok:
        fallos.append(nombre)
    print(('OK ' if ok else 'XX ') + nombre + ('' if ok else '   got=%r exp=%r' % (obtenido, esperado)))


# ═══════════════════════════════════════════════════════════════════════════════
# EL DOBLE DE SUPABASE (el de test_escaner2_dbline_pro.py, copiado: cada banco, el suyo)
# ═══════════════════════════════════════════════════════════════════════════════
class _Resp:
    def __init__(self, data, count=None):
        self.data, self.count = data, count


class _Consulta:
    def __init__(self, bd, tabla):
        self.bd, self.tabla = bd, tabla
        self.op, self.filtros, self.payload = 'select', [], None
        self.cuenta, self.orden, self.rango, self.tope = False, None, None, None

    def select(self, _cols='*', count=None):
        self.op, self.cuenta = 'select', (count == 'exact')
        return self

    def insert(self, filas):
        self.op, self.payload = 'insert', filas if isinstance(filas, list) else [filas]
        return self

    def update(self, cambios):
        self.op, self.payload = 'update', cambios
        return self

    def eq(self, k, v):
        self.filtros.append((k, v))
        return self

    def order(self, col, desc=False):
        self.orden = (col, desc)
        return self

    def range(self, a, b):
        self.rango = (a, b)
        return self

    def limit(self, n):
        self.tope = n
        return self

    def execute(self):
        self.bd.setdefault('ops', []).append([self.bd.get('programa'), self.op, self.tabla])
        filas = self.bd['tablas'].setdefault(self.tabla, [])
        if self.op == 'insert':
            for f in self.payload:
                filas.append(json.loads(json.dumps(f, default=str)))
            return _Resp(self.payload)
        sel = [f for f in filas if all(str(f.get(k)) == str(v) for k, v in self.filtros)]
        if self.op == 'update':
            for f in sel:
                f.update(json.loads(json.dumps(self.payload, default=str)))
            return _Resp(sel)
        if self.orden:
            sel = sorted(sel, key=lambda f: (f.get(self.orden[0]) is None, str(f.get(self.orden[0]))),
                         reverse=self.orden[1])
        n = len(sel)
        if self.rango:
            sel = sel[self.rango[0]:self.rango[1] + 1]
        if self.tope is not None:
            sel = sel[:self.tope]
        return _Resp(sel, n if self.cuenta else None)


class _Cubo:
    def __init__(self, bd, cubo):
        self.bd, self.cubo = bd, cubo

    def _objs(self):
        return self.bd['storage'].setdefault(self.cubo, {})

    def upload(self, ruta, datos, _opciones=None):
        self.bd.setdefault('subidas', []).append(ruta)
        self._objs()[ruta] = base64.b64encode(datos if isinstance(datos, bytes) else datos.encode()).decode()
        return {'Key': ruta}

    def download(self, ruta):
        if ruta not in self._objs():
            raise RuntimeError('404 Object not found: %s/%s' % (self.cubo, ruta))
        return base64.b64decode(self._objs()[ruta])

    def list(self, carpeta=None, opciones=None):
        op = dict({'limit': 100, 'offset': 0}, **(opciones or {}))
        pref = (carpeta or '').rstrip('/') + '/'
        nombres = sorted(r[len(pref):] for r in self._objs() if r.startswith(pref) and '/' not in r[len(pref):])
        trozo = nombres[op['offset']:op['offset'] + op['limit']]
        return [{'name': n, 'created_at': '2026-10-10T09:%02d:00Z' % i, 'metadata': {'size': 1}}
                for i, n in enumerate(trozo)]


class _Cliente:
    def __init__(self, bd):
        self.bd = bd
        self.storage = types.SimpleNamespace(from_=lambda cubo: _Cubo(bd, cubo))

    def table(self, nombre):
        return _Consulta(self.bd, nombre)


def hijo(ruta_estado, programa):
    with open(ruta_estado, encoding='utf-8') as fh:
        bd = json.load(fh)
    bd['programa'] = programa
    clientes = []

    def crear(url, llave):
        clientes.append(llave)
        return _Cliente(bd)

    sys.modules['supabase'] = types.ModuleType('supabase')
    sys.modules['supabase'].create_client = crear
    import runpy
    sys.argv = [programa]
    codigo = 0
    try:
        runpy.run_path(os.path.join(AQUI, programa), run_name='__main__')
    except SystemExit as ex:
        codigo = ex.code if isinstance(ex.code, int) else (0 if ex.code is None else 1)
    finally:
        with open(ruta_estado, 'w', encoding='utf-8') as fh:
            json.dump(bd, fh, default=str)
        print('CLIENTES_CREADOS=%d' % len(clientes))
    sys.exit(codigo)


if __name__ == '__main__' and len(sys.argv) >= 4 and sys.argv[1] == '--hijo':
    hijo(sys.argv[2], sys.argv[3])


def correr(ruta_estado, programa, env_extra):
    env = {k: v for k, v in os.environ.items()
           if k not in ('SUPABASE_SERVICE_KEY', 'PASADA', 'GITHUB_RUN_ID', 'MODO_BARRIDO', 'MARCAS_ELEGIDAS',
                        'OFERTAS_ELEGIDAS')}
    env.update(env_extra, PYTHONIOENCODING='utf-8', SUPABASE_URL='https://doble.invalid')
    p = subprocess.run([sys.executable, '-u', os.path.abspath(__file__), '--hijo', ruta_estado, programa],
                       capture_output=True, text=True, encoding='utf-8', errors='replace', env=env, cwd=AQUI)
    return p.returncode, (p.stdout or '') + (p.stderr or '')




import escaner2_motor as e2  # noqa: E402
import escaner2_bems_pro as ebe  # noqa: E402
import escaner2_heredado_pro as pro  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

M = e2.cargar_motor()


def ean13(cuerpo):
    return cuerpo + M._chk13(cuerpo)


# ═══════════════════════════════════════════════════════════════════════════════
# LA FOTO DE BEMS, INVENTADA (la forma de disp_estado tras el encargo BE2)
# ═══════════════════════════════════════════════════════════════════════════════
E1 = ean13('840000000301')    # Funko: PA 7,00
E2 = ean13('840000000302')    # Funko «w/Chase»: la figura normal, entra
E3 = ean13('840000000303')    # Funko con «Chase» al FINAL del titulo: en BEMS tambien entra (sin regla de chase)
E4 = ean13('840000000304')    # Pyramid
E5 = ean13('840000000305')    # Bandai Model Kit
E6 = ean13('840000000306')    # Funko Tees: en el modo de la lista fija casa con «Funko» (la regla de siempre)
E7 = ean13('840000000307')    # Paladone: fuera de la lista fija
E8 = ean13('840000000308')    # Funko sin stock
E10 = ean13('840000000310')   # sin fabricante: fuera de la lista fija; con «todas», dentro
E11 = ean13('840000000311')   # Funko sin PA (vacio)
E12 = ean13('840000000312')   # dos referencias: 6,00 y 5,50 → gana la de 5,50
E13 = ean13('840000000313')   # dos referencias al mismo precio → la referencia menor
E14 = ean13('840000000314')   # blind box de 6: un articulo a su precio de caja (64,00), sin dividir
E15 = ean13('840000000315')   # Funko con PA 0: sin precio
DISP = '00000000-0000-4000-8000-0000000be401'
OTRA = '00000000-0000-4000-8000-0000000be402'
FIJAS = ['Funko', 'Bandai Model Kit', 'Pyramid']


def fila(ref, ean, marca, pa, disp=True, regla=None, core=None, nombre=None, pasada=DISP):
    return {'proveedor': 'BEMS', 'producto_prov': ref, 'ean_original': ean or None,
            'ean_core': core if core is not None else (ean or None), 'marca': marca,
            'nombre': nombre or 'Figura inventada %s' % ref, 'categoria': 'Bobble Head POP', 'precio_unidad': pa,
            'precio_catalogo': pa, 'en_oferta': False, 'fin_oferta': None, 'precio_escalon': None,
            'uds_escalon': None, 'precio_pa': None, 'es_caja': False, 'uds_caja': None, 'es_chase': False,
            'disponible': disp, 'disponibilidad': '7' if disp else '0', 'preorder': False, 'regla': regla,
            'pasada_id': pasada, 'ausencias': 0}


def foto_disp(pasada_ajena=False):
    return [
        fila('900001', E1, 'Funko', 7.0),
        fila('900002', E2, 'Funko', 7.0, nombre='Figura inventada w/Chase'),
        fila('900003', E3, 'Funko', 7.0, nombre='Figura inventada Chase'),
        fila('900004', E4, 'Pyramid', 3.0),
        fila('900005', E5, 'Bandai Model Kit', 20.0),
        fila('900006', E6, 'Funko Tees', 9.0),
        fila('900007', E7, 'Paladone', 11.0),
        fila('900008', E8, 'Funko', 7.0, disp=False),
        fila('900009', '', 'Funko', 7.0, regla='ean_forma_rara', core=None),
        fila('900010', E10, None, 4.0),
        fila('900011', E11, 'Funko', None),
        fila('900013', E12, 'Funko', 6.0),
        fila('900012', E12, 'Funko', 5.5),
        fila('900015', E13, 'Funko', 6.0),
        fila('900014', E13, 'Funko', 6.0, pasada=OTRA if pasada_ajena else DISP),
        fila('900016', E14, 'Funko', 64.0, nombre='Figura inventada Blind Box (6pcs)'),
        fila('900017', E15, 'Funko', 0.0),
        fila('900018', '12345', 'Funko', 7.0, core='12345'),
    ]


N_SIN_EAN = 0                       # BE2 sube todas las filas, tambien las sin EAN
N_ARTICULOS = len(foto_disp())      # 18 (una por referencia)
N_REPETIDAS = 2                     # las lineas identicas repetidas del CSV, que la pasada ya se quedo una vez
N_LINEAS = N_ARTICULOS + N_REPETIDAS
EN_FOTO_MARCAS = [E1, E2, E3, E4, E5, E6, E12, E13, E14]

# ═══════════════════════════════════════════════════════════════════════════════
# PARTE A · LO PURO
# ═══════════════════════════════════════════════════════════════════════════════
print('A1 · el filtro de marca, por modo')
_q, _i = ebe.filtro_marca('marcas', FIJAS)
eq('A1 · marcas: la lista fija con la regla de siempre (dentro del nombre, sin mayúsculas)',
   [_q(m) for m in ('Funko', 'FUNKO', 'Funko Tees', 'Pyramid', 'Bandai Model Kit', 'Paladone', None, 'Bandai')],
   [True, True, True, True, True, False, False, False])
eq('A1 · …y guarda las marcas de la lista', _i['marcas_reales'], FIJAS)
_q, _i = ebe.filtro_marca('todas')
eq('A1 · todas: también las sin fabricante, y sin marcas que guardar', ([_q(m) for m in ('Paladone', None)],
                                                                         _i['marcas_reales']), ([True, True], None))
_q, _i = ebe.filtro_marca('elegidas', elegidas=['Funko'])
eq('A1 · elegidas: coincidencia EXACTA (sin mayúsculas), como HEO', [_q(m) for m in ('FUNKO', 'Funko Tees', ' funko ')],
   [True, False, True])
for _modo, _args in (('marcas', ([],)), ('elegidas', (None, [])), ('raro', ())):
    try:
        ebe.filtro_marca(_modo, *_args)
        _r = 'pasa'
    except ebe.FalloBEMS:
        _r = 'para'
    eq('A1 · %s sin marcas (o desconocido): para' % _modo, _r, 'para')


def comprobar_foto(mod):
    """A2 entero contra un modulo (el de verdad o un mutante). Devuelve lo que falla."""
    rotas = []

    def chk(nombre, obtenido, esperado):
        if obtenido != esperado:
            rotas.append((nombre, obtenido, esperado))
    quiere, _ = mod.filtro_marca('marcas', FIJAS)
    foto, apartados, cuentas = mod.construir_foto(foto_disp(), quiere, M, N_ARTICULOS, N_SIN_EAN, 'marcas')
    por = {f['ean_core']: f for f in foto}
    chk('foto', sorted(por), sorted(EN_FOTO_MARCAS))
    chk('referencias', {k: por[k]['producto_heo'] for k in por},
        {E1: '900001', E2: '900002', E3: '900003', E4: '900004', E5: '900005', E6: '900006', E12: '900012',
         E13: '900014', E14: '900016'})
    chk('PA = el del CSV; la blind box a su precio de caja, sin dividir',
        [por[k]['precio_unidad'] for k in (E1, E12, E14) if k in por], [7.0, 5.5, 64.0])
    chk('sin cajas, sin chase, sin ofertas', {(f['es_caja'], f['uds_caja'], f['es_chase'], f['en_oferta'],
                                               f['aviso_caja']) for f in foto}, {(False, None, False, False, None)})
    previas = cuentas['previas']
    chk('previas', previas, {'chase_funko': 0, 'sin_gtin': 0, 'no_disponible': 1, 'marca_fuera': 2,
                             'estado_no_servible': 2, 'chase_suelto': 0, 'ean_forma_rara': 2, 'duplicado_proveedor': 2})
    chk('apartados', sorted((a['producto_heo'], a['motivo']) for a in apartados),
        sorted([('900007', 'marca_fuera'), ('900010', 'marca_fuera'), ('900009', 'ean_forma_rara'),
                ('900018', 'ean_forma_rara'), ('900011', 'estado_no_servible'), ('900017', 'estado_no_servible'),
                ('900013', 'duplicado_proveedor'), ('900015', 'duplicado_proveedor')]))
    chk('cuadra', (cuentas['n_foto'], cuentas['n_previas'], cuentas['cuadra_previo']), (9, 9, True))
    _f, _a, c2 = mod.construir_foto(foto_disp(), quiere, M, N_ARTICULOS + 2, N_SIN_EAN, 'marcas')
    chk('un crudo que no casa (las líneas en vez de los artículos): no cuadra y lo dice',
        (c2['cuadra_previo'], 'el CSV tenía 20 artículos' in (c2['motivo_previo'] or '')), (False, True))
    return rotas


print('A2 · la foto de la pasada')
_rotas = comprobar_foto(ebe)
for _n, _o, _e in _rotas:
    eq('A2 · ' + _n, _o, _e)
eq('A2 · las comprobaciones de la foto, en verde', len(_rotas), 0)
quiere_t, _ = ebe.filtro_marca('todas')
_ft, _at, _ct = ebe.construir_foto(foto_disp(), quiere_t, M, N_ARTICULOS, N_SIN_EAN, 'todas')
eq('A2 · todas: entran Paladone y la sin fabricante (11), y cuadra',
   (len(_ft), _ct['cuadra_previo'], _ct['previas']['marca_fuera'], sorted(f['ean_core'] for f in _ft)),
   (11, True, 0, sorted(EN_FOTO_MARCAS + [E7, E10])))
_qe, _ = ebe.filtro_marca('elegidas', elegidas=['Funko'])
_fe, _ae, _ce = ebe.construir_foto(foto_disp(), _qe, M, N_ARTICULOS, N_SIN_EAN, 'elegidas')
eq('A2 · elegidas Funko: los Funko con chase en el título dentro, Funko Tees fuera, y la marca fuera dice «no elegida»',
   (sorted(f['ean_core'] for f in _fe), E6 in {f['ean_core'] for f in _fe},
    sorted({a['detalle'].split("'")[0] for a in _ae if a['motivo'] == 'marca_fuera'})),
   (sorted([E1, E2, E3, E12, E13, E14]), False, ['Marca ']))
eq('A2 · …y su puerta previa se llama «Marca no elegida»', ebe.nombre_previa('marca_fuera', 'elegidas'), 'Marca no elegida')
eq('A2 · los nombres de las puertas previas de BEMS',
   [ebe.nombre_previa(p, 'marcas') for p in ('no_disponible', 'estado_no_servible', 'ean_forma_rara', 'marca_fuera')],
   ['Sin stock hoy', 'Sin precio de compra', 'EAN vacío o con forma rara', 'Marca fuera de la lista fija de BEMS'])
eq('A2 · las rutas que admiten los checks de la v2',
   (ebe.ruta_lista(DISP), ebe.ruta_excel(DISP, OTRA, '20261010_1930')),
   ('bems/%s/eans.txt' % DISP, 'bems/%s/%s/Escaner2_BEMS_20261010_1930.xlsx' % (DISP, OTRA)))
eq('A2 · el día del CSV: de hoy, sin aviso; de ayer y de hace 5 días, lo dice con la fecha; sin fecha, también',
   (ebe.aviso_fecha_csv('2026-10-10T00:00:00', date(2026, 10, 10)),
    'CSV del 09-10-2026 (ayer), no de hoy' in (ebe.aviso_fecha_csv('2026-10-09', date(2026, 10, 10)) or ''),
    'CSV del 05-10-2026 (hace 5 días)' in (ebe.aviso_fecha_csv(date(2026, 10, 5), date(2026, 10, 10)) or ''),
    'no dice de qué día' in (ebe.aviso_fecha_csv(None, date(2026, 10, 10)) or '')),
   (None, True, True, True))

print('A3 · 🔴 mutantes de lo puro: cada uno pone A2 en rojo')
with io.open(os.path.join(AQUI, 'escaner2_bems_pro.py'), encoding='utf-8') as _fh:
    FUENTE = _fh.read()
MUTANTES = [
    ('sin precio (0) entra', 'if pa is None or pa <= 0:', 'if pa is None:'),
    ('gana la más cara', "if (pa, str(f.get('producto_prov'))) < (prev[1]", "if (pa, str(f.get('producto_prov'))) > (prev[1]"),
    ('la forma rara entra', 'if regla or not core.isdigit() or len(core) not in (12, 13):', 'if regla:'),
    ('lo que no tiene stock entra', "if not f.get('disponible'):", 'if False:'),
    ('la lista fija, exacta', "return (lambda m: any(b in str(m or '').lower() for b in bajas))",
     "return (lambda m: any(b == str(m or '').lower() for b in bajas))"),
    ('el chase con «Chase» al final, fuera', "        regla = f.get('regla')\n",
     "        regla = f.get('regla')\n        if str(f.get('nombre') or '').lower().endswith('chase'):\n            continue\n"),
    ('la caja se divide entre 6', "'precio_catalogo': _num(f.get('precio_catalogo')), 'precio_unidad': pa,",
     "'precio_catalogo': _num(f.get('precio_catalogo')), 'precio_unidad': pa / 6 if 'pcs' in str(f.get('nombre')) else pa,"),
]
for _nombre, _antes, _despues in MUTANTES:
    if FUENTE.count(_antes) != 1:
        eq('A3 · el ancla del mutante «%s» es única' % _nombre, FUENTE.count(_antes), 1)
        continue
    _mod = types.ModuleType('mutante')
    exec(compile(FUENTE.replace(_antes, _despues), 'mutante', 'exec'), _mod.__dict__)
    eq('A3 · mutante «%s»: A2 en rojo' % _nombre, len(comprobar_foto(_mod)) > 0, True)

# ═══════════════════════════════════════════════════════════════════════════════
# PARTE B · LOS DOS PROGRAMAS, DE PUNTA A PUNTA
# ═══════════════════════════════════════════════════════════════════════════════
with io.open(os.path.join(AQUI, 'test_escaner2_osma_cabecera.json'), encoding='utf-8') as fh:
    CABECERA = json.load(fh)['cabecera']
COL_PAIS, COL_CAIDAS = e2.columnas_keepa()
COL_PADRE, COL_NVAR, COL_RANK = e2.columnas_ficha_compartida()
C = pro.CSV_COLS


def csv_real(pais, filas):
    """Un export del Visualizador con la CABECERA REAL (561 columnas, la del banco de OSMA) y solo lo que se lee."""
    buf = io.StringIO()
    w = csv.writer(buf)
    w.writerow(CABECERA)
    ix = {h: i for i, h in enumerate(CABECERA)}
    for asin, eans, caidas, precio, titulo in filas:
        fila_ = [''] * len(CABECERA)
        for h, v in ((C['asin'], asin), (COL_PAIS, pais), ('Título', titulo), (C['ean'], eans),
                     (C['rank'], '9000'), (C['rank90'], '11000'), (COL_CAIDAS, caidas), (C['buybox'], precio),
                     (C['es_fba'], 'yes'), (C['nuevo'], precio), (C['fba'], '3.10'), (C['compct'], '15.01 %'),
                     (C['nvar'], ''), (COL_PADRE, '')):
            fila_[ix[h]] = v
        w.writerow(fila_)
    return ('﻿' + buf.getvalue()).encode('utf-8')


ASIN1, ASIN2, ASIN3, ASIN4 = 'B0BEMS0001', 'B0BEMS0002', 'B0BEMS0003', 'B0BEMS0004'
PRODUCTOS = [{'id': 1, 'ean': E1, 'asin': ASIN1, 'nombre': 'Figura de Moloka', 'activo': True, 'es_chase': False,
              'iva_pct': 0.21, 'stock_moloka': 3}]
SEMBRADAS = ('productos', 'inventario_fba', 'disp_pasada', 'disp_estado', 'reglas_director', 'escaner2_parametros')
# 🔑 El CSV de la foto es del 1-oct: en los programas el «hoy» es el de verdad, así que NO es de hoy y lo tienen que decir.
FECHA_CSV = '2026-10-01T00:00:00'


def estado_inicial(escena):
    aplicada = {'id': DISP, 'proveedor': 'BEMS', 'estado': 'aplicada', 'creada_en': '2026-10-01T14:46:39Z',
                'terminada_en': '2026-10-01T14:47:13Z', 'n_leidas': N_ARTICULOS,
                'n_crudo': N_LINEAS + (1 if escena == 'descuadre' else 0), 'n_duplicados': N_REPETIDAS,
                'n_sin_gtin': N_SIN_EAN, 'fichero_md5': 'f' * 32, 'fichero_fecha_max': FECHA_CSV}
    pasadas = [{'id': OTRA, 'proveedor': 'BEMS', 'estado': 'aplicada', 'creada_en': '2026-09-30T07:00:00Z',
                'n_leidas': 1, 'n_crudo': 1, 'n_duplicados': 0, 'n_sin_gtin': 0, 'fichero_fecha_max': '2026-09-30'},
               aplicada,
               # Un CSV con el mismo contenido despues: 'rechazada' «al día». No es la aplicada.
               {'id': 'rechazada-al-dia', 'proveedor': 'BEMS', 'estado': 'rechazada',
                'creada_en': '2026-10-01T14:49:37Z', 'n_crudo': N_LINEAS, 'n_duplicados': N_REPETIDAS,
                'fichero_fecha_max': FECHA_CSV},
               # Una foto de DBLine mas nueva: el barrido de BEMS no la mira.
               {'id': 'de-dbline', 'proveedor': 'DBLINE', 'estado': 'aplicada', 'creada_en': '2026-10-02T11:00:00Z',
                'n_leidas': 1, 'n_crudo': 1, 'n_duplicados': 0, 'n_sin_gtin': 0}]
    if escena == 'sin_foto':
        pasadas = [p for p in pasadas if not (p['proveedor'] == 'BEMS' and p['estado'] == 'aplicada')]
    return {'tablas': {
        'escaner2_parametros': [{'proveedor': 'HEO', 'umbral_caidas_30d': 6, 'paises_filtro': ['ES', 'IT', 'FR', 'DE'],
                                 'paises_calculo': ['ES', 'IT', 'FR', 'DE'], 'marcas_fijas': ['Funko', 'Ultimate Guard']},
                                # La fila de BEMS que deja la migracion de la v2.
                                {'proveedor': 'BEMS', 'umbral_caidas_30d': 6,
                                 'paises_filtro': ['ES', 'IT', 'FR', 'DE'], 'paises_calculo': ['ES', 'IT', 'FR', 'DE'],
                                 'marcas_fijas': FIJAS}],
        'escaner2_pasada': [],
        'disp_pasada': pasadas,
        'disp_estado': foto_disp(pasada_ajena=(escena == 'ajena'))
        + [dict(fila('DB0001', E1, 'Funko', 1.0, pasada='de-dbline'), proveedor='DBLINE')],
        # BEMS no tiene fila en reglas_director (medido el 10-oct-2026): el barrido no la lee.
        'reglas_director': [{'proveedor': 'DBLINE', 'activo': True, 'marcas': ['Funko', 'Pyramid']}],
        'productos': json.loads(json.dumps(PRODUCTOS)),
        'inventario_fba': [],
    }, 'storage': {'escaner2': {}}}


def caso(escena, env=None):
    tmp = tempfile.mkdtemp(prefix='e2be_')
    ruta = os.path.join(tmp, 'estado.json')
    inicial = estado_inicial(escena)
    with open(ruta, 'w', encoding='utf-8') as fh:
        json.dump(inicial, fh)
    cod, log = correr(ruta, 'escaner2_bems_barrido.py', dict({'SUPABASE_SERVICE_KEY': 'svc-de-mentira',
                                                              'GITHUB_RUN_ID': '818181'}, **(env or {})))
    bd = json.load(open(ruta, encoding='utf-8'))
    return cod, log, bd, ruta, inicial


print('B1 · sin llave, o elegidas con ofertas: no se corre')
_tmp = tempfile.mkdtemp(prefix='e2be_')
_ruta = os.path.join(_tmp, 'estado.json')
json.dump(estado_inicial('marcas'), open(_ruta, 'w', encoding='utf-8'))
for _prog in ('escaner2_bems_barrido.py', 'escaner2_bems_cruce.py'):
    _cod, _log = correr(_ruta, _prog, {'PASADA': DISP})
    eq('B1 · %s: ROJO, la línea exacta y NINGÚN cliente' % _prog,
       (_cod, 'ESCANER2_NO_EJECUTADO: sin llave de servicio' in _log, 'CLIENTES_CREADOS=0' in _log), (1, True, True))
_cod, _log = correr(_ruta, 'escaner2_bems_barrido.py', {'SUPABASE_SERVICE_KEY': 'x', 'MODO_BARRIDO': 'elegidas',
                                                        'MARCAS_ELEGIDAS': '["Funko"]', 'OFERTAS_ELEGIDAS': 'true'})
eq('B1 · elegidas con ofertas: ROJO, lo dice y NINGÚN cliente',
   (_cod, 'no tiene casilla de ofertas' in _log, 'CLIENTES_CREADOS=0' in _log), (1, True, True))
_cod, _log = correr(_ruta, 'escaner2_bems_barrido.py', {'SUPABASE_SERVICE_KEY': 'x', 'MODO_BARRIDO': 'elegidas',
                                                        'MARCAS_ELEGIDAS': '["Funko", "a$b"]',
                                                        'OFERTAS_ELEGIDAS': 'false'})
eq('B1 · una marca con «$»: ROJO y NINGÚN cliente (la validación de HEO)',
   (_cod, 'selección de marcas no válida' in _log, 'CLIENTES_CREADOS=0' in _log), (1, True, True))

print('B2 · [marcas] la foto → barrido → la lista → los 4 CSV → cruce → Excel')
cod, log, bd, ruta, inicial = caso('marcas')
eq('B2 · el barrido sale en VERDE', cod, 0)
if cod != 0:
    print(log)
pas = [p for p in bd['tablas']['escaner2_pasada'] if p.get('run_id') == 818181][0]
PASADA = pas['id']
eq('B2 · 🔴 la pasada: BEMS, modo marcas, esperando_csv, artículos 18 (no las 20 líneas) = previas 9 + foto 9, la lista fija',
   (pas['proveedor'], pas['modo'], pas['estado'], pas['n_crudo'], pas['n_foto'],
    sum(pas['p_' + p] for p in e2.PUERTAS_PREVIAS), pas['n_tandas'], pas['ruta_lista'], pas['marcas'], pas['ofertas']),
   ('BEMS', 'marcas', 'esperando_csv', 18, 9, 9, 1, 'bems/%s/eans.txt' % PASADA, FIJAS, False))
_foto = [f for f in bd['tablas']['escaner2_foto'] if f['pasada_id'] == PASADA]
eq('B2 · la foto en la base: 9, y las listas de las puertas previas: 8',
   (len(_foto), len([a for a in bd['tablas']['escaner2_apartado'] if a['pasada_id'] == PASADA])), (9, 8))
_alm = bd['storage']['escaner2']
_txt = lambda k: base64.b64decode(_alm[k]).decode('utf-8')  # noqa: E731
eq('B2 · en el almacén: eans.txt y barrido.json', sorted(k.split('/')[-1] for k in _alm), ['barrido.json', 'eans.txt'])
eq('B2 · eans.txt: los nueve EAN', sorted(_txt('bems/%s/eans.txt' % PASADA).split('\n')), sorted(EN_FOTO_MARCAS))
_side = json.loads(_txt('bems/%s/barrido.json' % PASADA))
eq('B2 · barrido.json: de qué foto sale, el día de su CSV y sus líneas',
   (_side['disp_pasada'], _side['fecha_csv'], _side['n_lineas_csv'], _side['n_duplicados_csv']),
   (DISP, '2026-10-01', N_LINEAS, N_REPETIDAS))
eq('B2 · 🔑 el registro del barrido dice de qué día es el CSV, y que no es de hoy',
   ('CSV del 01-10-2026' in log, 'no de hoy' in log), (True, True))

_carpeta = 'bems/%s/csv/' % PASADA
# E1, E2 y E14 se venden; E4 no se vende; los demás no están en Keepa.
_filas = [(ASIN1, E1, '30', '25.95', 'Figura uno'), (ASIN2, E2, '30', '25.95', 'Figura dos'),
          (ASIN3, E14, '30', '129.95', 'Caja tres'), (ASIN4, E4, '2', '25.95', 'Llavero cuatro')]
for _i, _p in enumerate(('es', 'it', 'fr', 'de')):
    _alm[_carpeta + '20261010-1000%02d-KeepaExport-2026-10-10-VisualizadorDeProductos (%d).csv' % (_i, _i)] = \
        base64.b64encode(csv_real(_p, _filas)).decode()
bd.pop('programa', None)
json.dump(bd, open(ruta, 'w', encoding='utf-8'), default=str)
cod2, log2 = correr(ruta, 'escaner2_bems_cruce.py', {'SUPABASE_SERVICE_KEY': 'svc-de-mentira', 'PASADA': PASADA,
                                                      'GITHUB_RUN_ID': '828282'})
bd = json.load(open(ruta, encoding='utf-8'))
T = bd['tablas']
_alm = bd['storage']['escaner2']
eq('B2 · el cruce sale en VERDE', cod2, 0)
if cod2 != 0:
    print(log2)
cr = [c for c in T['escaner2_cruce'] if c.get('run_id') == 828282][0]
eq('B2 · 🔴 el cruce: lista, cuadra, artículos 18 = previas 9 + 9 puertas, en los CUATRO países, y avisa del día del CSV',
   (cr['estado'], cr['cuadra'], cr['n_crudo'], cr['n_previas'], cr['n_entradas'], sum(cr['n_' + x] for x in 'abcdef'),
    cr['paises_filtro'], cr['paises_calculo'], cr['paises_usados'], 'CSV del 01-10-2026' in (cr['aviso'] or '')),
   ('lista', True, 18, 9, 9, 9, ['ES', 'IT', 'FR', 'DE'], ['ES', 'IT', 'FR', 'DE'], ['ES', 'IT', 'FR', 'DE'], True))
_res = {f['ean_core']: r for r in T['escaner2_resultado_ean'] for f in T['escaner2_foto'] if f['id'] == r['foto_id']}
_pais = {}
for _r in T['escaner2_resultado_pais']:
    _pais.setdefault(_r['resultado_ean_id'], {})[_r['pais']] = _r
eq('🔴 B2 · el PA de cada uno en los cuatro países: 7 · 7 · la caja entera, 64',
   [(sorted(_pais[_res[k]['id']]), {r['pa'] for r in _pais[_res[k]['id']].values()}) for k in (E1, E2, E14)],
   [(['DE', 'ES', 'FR', 'IT'], {7.0}), (['DE', 'ES', 'FR', 'IT'], {7.0}), (['DE', 'ES', 'FR', 'IT'], {64.0})])
eq('B2 · E1 COMPRAR, E4 no se vende (puerta c) y E5 no está en Keepa (puerta a)',
   (_res[E1]['puerta'], _res[E4]['puerta'], _res[E5]['puerta']), ('f', 'c', 'a'))

_xl = [k for k in _alm if k.startswith('bems/%s/%s/Escaner2_BEMS_' % (PASADA, cr['id']))]
eq('B2 · el Excel, en su carpeta y con su nombre (el que guarda el cruce)', (len(_xl), cr['ruta_excel'] == _xl[0]),
   (1, True))
_wb = load_workbook(io.BytesIO(base64.b64decode(_alm[_xl[0]])))
eq('B2 · las hojas del PRO de HEO, en su orden (sin «Chase_manual», que es de HEO)', _wb.sheetnames,
   ['Análisis', 'Descartados', 'Ambiguos', 'Sin_rank', 'Precio por lote', 'Resumen', 'Comparación', 'Varias fichas',
    'Puertas', 'Puertas previas'])
_cab = [c.value for c in _wb['Análisis'][1]]
eq('🔴 B2 · «Análisis»: las del PRO de HEO (con «Ventas») y «Ficha compartida» al final, NADA MÁS',
   (_cab, 'Ventas' in _cab), (e2.columnas_analisis() + [e2.COLUMNA_FICHA_COMPARTIDA], True))
_ix = {n: _cab.index(n) for n in _cab}
_an = list(_wb['Análisis'].iter_rows(min_row=2, values_only=True))
_por = {}
for _r in _an:
    _por.setdefault(_r[_ix['EAN']], []).append(_r)
eq('B2 · «Análisis»: los que se venden, una fila por país (los cuatro)',
   sorted((k, sorted(r[_ix['País']] for r in v)) for k, v in _por.items()),
   sorted((k, ['DE', 'ES', 'FR', 'IT']) for k in (E1, E2, E14)))
eq('🔴 B2 · «Análisis»: el PA de cada uno (7 · 7 · la caja, 64)',
   [{r[_ix['PA (€)']] for r in _por.get(k, [])} for k in (E1, E2, E14)], [{7.0}, {7.0}, {64.0}])
_pu = list(_wb['Puertas'].iter_rows(values_only=True))
eq('B2 · «Puertas»: las columnas del PRO de HEO (con las de caja, vacías) y una fila por EAN de la foto',
   (list(_pu[0])[-5:], {r[-5:] for r in _pu[1:]}, len(_pu) - 1),
   (ebe.COLUMNAS_CAJA, {(None, None, None, None, None)}, 9))
_resu = {}
for _r in _wb['Resumen'].iter_rows(values_only=True):
    _resu.setdefault(_r[0], _r[1])
_avisos = [r[1] for r in _wb['Resumen'].iter_rows(values_only=True) if r[0] == 'Aviso']
eq('B2 · el Resumen: modo, lista fija, cuadre, el día del CSV y su aviso, líneas y artículos, y las puertas previas de BEMS',
   (_resu.get('Modo del barrido'), _resu.get('Marcas de la lista fija'), _resu.get('Cuadra'),
    _resu.get('CSV de BEMS (el día del fichero)'), any('CSV del 01-10-2026' in (a or '') for a in _avisos),
    _resu.get('Líneas del CSV de BEMS (con las repetidas idénticas)'),
    _resu.get('Catálogo crudo de BEMS (artículos distintos del CSV)'),
    _resu.get('Puerta previa · Sin precio de compra'), _resu.get('Puerta previa · Sin stock hoy')),
   ('marcas de siempre', ', '.join(FIJAS), 'SÍ', '01-10-2026', True, N_LINEAS, 18, 2, 1))

print('B3 · 🔴 el registro (repo PÚBLICO) no lleva ni un EAN, precio, referencia de BEMS ni nombre')
_prohibido = [E1, E2, E3, E4, E5, E6, E7, E8, E10, E11, E12, E13, E14, E15, ASIN1, ASIN2, ASIN3, ASIN4, '25.95',
              '129.95', '64.0', '7.0', '5.5', '9000', 'Figura inventada', 'Figura uno', 'Caja tres', 'Llavero']
eq('B3 · ni en el barrido ni en el cruce (si sale algo, la línea que lo lleva)',
   [(x, [ln for ln in (log + log2).splitlines() if x in ln][:2]) for x in _prohibido if x in log or x in log2], [])
eq('B3 · …y el cruce SÍ dice de qué día es el CSV de la foto', 'CSV de BEMS de la foto: del 01-10-2026' in log2, True)

print('B4 · solo se escribe en escaner2_* y en la carpeta de la pasada')
eq('B4 · las tablas escritas', sorted({t for _p, op, t in bd['ops'] if op != 'select'}),
   ['escaner2_apartado', 'escaner2_cruce', 'escaner2_foto', 'escaner2_pasada', 'escaner2_resultado_ean',
    'escaner2_resultado_pais'])
eq('B4 · y las sembradas acaban como empezaron',
   [t for t in SEMBRADAS if bd['tablas'][t] != json.loads(json.dumps(inicial['tablas'][t]))], [])
eq('B4 · en el almacén, solo bems/<pasada>/', sorted({'/'.join(s.split('/')[:2]) for s in bd.get('subidas', [])}),
   ['bems/' + PASADA])

print('B5 · los otros modos, y una foto que no sirve')
cod, log, bd, _r, _i0 = caso('todas', {'MODO_BARRIDO': 'todas'})
_p = bd['tablas']['escaner2_pasada'][0]
eq('B5 · [todas] VERDE, 11 en la foto (Paladone y la sin fabricante), sin marcas guardadas',
   (cod, _p['modo'], _p['n_foto'], _p['p_marca_fuera'], _p['marcas']), (0, 'todas', 11, 0, None))
cod, log, bd, _r, _i0 = caso('elegidas', {'MODO_BARRIDO': 'elegidas', 'MARCAS_ELEGIDAS': '["Pyramid"]',
                                          'OFERTAS_ELEGIDAS': 'false'})
_p = bd['tablas']['escaner2_pasada'][0]
eq('B5 · [elegidas] VERDE, una en la foto, con su lista y ofertas a false',
   (cod, _p['modo'], _p['n_foto'], _p['marcas'], _p['ofertas']), (0, 'elegidas', 1, ['Pyramid'], False))
for _escena, _motivo in (('ajena', 'no es la de su última pasada aplicada'),
                         ('sin_foto', 'sube primero el CSV del día'),
                         ('descuadre', 'no cuadra: 21 líneas ≠ 18 artículos + 2 repetidas')):
    cod, log, bd, _r, _i0 = caso(_escena)
    _p = bd['tablas']['escaner2_pasada'][0]
    eq('B5 · 🔴 [%s] ROJO, la pasada fallida con el motivo y SIN foto' % _escena,
       (cod, _p['estado'], _motivo in (_p['motivo_fallo'] or ''), bd['tablas'].get('escaner2_foto', [])),
       (1, 'fallida', True, []))

print()
if fallos:
    print('FALLAN %d: %s' % (len(fallos), fallos))
    sys.exit(1)
print('TODO OK')
