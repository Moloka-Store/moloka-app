# -*- coding: utf-8 -*-
"""Banco de la PASADA DE DISPONIBILIDAD DE OSMA (encargo AE, tramo 2 del plano Y, 01-oct-2026).

SIN RED, SIN SECRETOS Y SIN BASE: se ejecuta el programa DE VERDAD (escaner2_osma_disponibilidad.py, con
runpy, como lo lanza el workflow) contra una web de OSMA y una base de mentira. 🔴 EL EXCEL ES INVENTADO:
el catalogo real de OSMA no se sube a este repo, que es publico. Codigos, EAN, nombres, marcas, precios y
stock son centinelas que no pueden salir en el registro.

QUE PRUEBA:
  (A) EL BUENO APLICA: entra una vez, localiza el enlace en la carpeta buena (habiendo otro
      osma_articles.xlsx en otra carpeta), sube las filas en lotes de 500 con la tabla del plano (llave = codigo;
      precio neto por unidad; disponible = stock > 0 o mas de 10.000) y declarado = llegados = filas; llama a
      disp_aplicar_pasada UNA vez, relee la pasada y SOLO DESPUES guarda la copia ENTERA en la ruta fija
      escaner2/osma/ultimo/osma_articles.xlsx (AE2). Si esa copia falla, la pasada sigue aplicada y el run sale
      en rojo diciendolo. Tambien con la carpeta en otra pagina.
  (B) NO APLICAN Y DICEN POR QUE (sin llamar a la funcion ni subir una fila): cabecera cambiada ('fallida',
      con el nombre), filas de menos ('rechazada_vaciado'), codigo repetido o vacio ('fallida'), fichero viejo
      o anterior a la ultima aplicada ('rechazada', «OSMA no ha regenerado el fichero»), lo bajado no es un
      xlsx, la descarga da 500. Y ninguno de ellos sube copia (un fichero rechazado no pisa la ultima buena).
  (C) SIN BAJAR NADA: entrada rechazada (UN solo intento), enlace no encontrado, enlace doble.
  (D) SIN SECRETOS: rojo en el primer paso, sin abrir pasada ni cliente de base ni de web.
  (E) 🔴 EL REGISTRO NO SUELTA DATOS, en TODOS los casos: ni el usuario, ni la clave, ni cookies, ni
      cabeceras, ni el HTML, ni el id o la URL del fichero, ni el cuerpo de un error de la web, ni el texto
      de un error de la base (que lleva una fila), ni codigos, EAN, nombres, marcas, precios o stock.
  (F) El rescate solo cierra la pasada de ESTE run, de OSMA y en 'leyendo'.
  (I) LA FECHA DEL FICHERO, CON LA HORA FIJADA (AE2): (a) relanzar el mismo dia con el fichero de hoy ya
      aplicado → VERDE (OSMA_AL_DIA), la pasada 'rechazada' «al día» (que no cuenta en ningun sitio: ni
      rechazada_vaciado ni fallida), sin filas, sin funcion y sin copia; el dia es el de MADRID (a las 00:10 de
      Madrid, la de las 23:50 es de ayer aunque en UTC sea el mismo dia); (b) laborable con el fichero de ayer
      ya aplicado → ROJO; (b) sabado → VERDE (OSMA_SIN_FICHERO_NUEVO).
  (G) Por estructura: el programa solo toca disp_pasada, disp_lectura y disp_parametros (lectura) y la
      funcion disp_aplicar_pasada; el workflow solo se lanza a mano, con grupo propio, los secretos en el
      primer paso y el rescate solo si estaban.
  (H) 🔴 LA MITAD ROJA: diez mutantes del programa (sin cabecera, sin fecha vieja, sin codigos repetidos,
      imprimiendo el error, con dos intentos de entrada; y los del AE2: «al día» en rojo, el dia en UTC, la copia
      antes de validar, la copia por pasada y la copia que falla en verde); el banco tiene que ponerse ROJO con
      cada uno, en la comprobacion que lo vigila.
"""
import ast
import contextlib
import io
import os
import runpy
import sys
import tempfile
import types
from datetime import datetime, timedelta, timezone

import openpyxl
import yaml

RAIZ = os.path.dirname(os.path.abspath(__file__))
PROGRAMA = os.path.join(RAIZ, 'escaner2_osma_disponibilidad.py')
WORKFLOW = os.path.join(RAIZ, '.github', 'workflows', 'escaner2-osma-disponibilidad.yml')
WEB = 'https://osma-werm.com'

fallos = []


def eq(nombre, obtenido, esperado):
    ok = obtenido == esperado
    if not ok:
        fallos.append(nombre)
    print(('OK ' if ok else 'XX ') + nombre + ('' if ok else '   got=%r exp=%r' % (obtenido, esperado)))


# ── Los centinelas: nada de esto puede salir en el registro ─────────────────────────────────
USUARIO = 'usuario.secreto.osma@example.test'
CLAVE = 'CLAVE-SECRETA-0sm4'
COOKIE = 'COOKIE-SECRETA-5e55'
CABECERA_WEB = 'X-Cabecera-Secreta-77'
HTML_SECRETO = 'HTML-SECRETO-DE-LA-WEB'
ID_FICHERO = '0f1e2d3c4b5a69788796a5b4c3d2e1f0'
ID_OTRA_CARPETA = '9a9a9a9a9a9a9a9a9a9a9a9a9a9a9a9a'
CUERPO_ERROR = 'CUERPO-ERROR-SECRETO'
FILA_BASE = 'FILA-DE-LA-BASE-SECRETA'
PRECIO = 3.777
STOCK = 987654
CENTINELAS = [USUARIO, CLAVE, COOKIE, CABECERA_WEB, HTML_SECRETO, ID_FICHERO, ID_OTRA_CARPETA, 'downloads/file',
              CUERPO_ERROR, FILA_BASE, 'CODSEC', '4999999', 'NOMBRE-SECRETO', 'MARCA-SECRETA', '3.777', '3,777',
              str(STOCK), 'Traceback']

COLUMNAS = ['BRAND', 'Description_short_ger_SDESC', 'Item-no_BAN', 'EAN_GTIN', 'MOQ', 'Price_net', 'Last price',
            'SALE', 'New', 'Stock_available_pcs', 'Stock_over10k_pcs', 'Ordered_pcs', 'Discontinued', 'MHD_EXP',
            'ArtGrp_ger', 'ArtGrp_eng',   # 🔑 columnas que no se comprueban (nombre no medido): pueden estar
            'Inner_Pack_pcs_IP', 'Outer_Case_pcs_OC', 'GTIN_Case', 'Layer_pcs', 'PAL_pcs', 'PAL_cm_HGT',
            'Weight_gram', 'VAT', 'Description_long_eng']


def ahora():
    return datetime.now(timezone.utc)


def fila_inventada(g):
    """Una fila de la hoja Data INVENTADA. g % 3: 0 con stock, 1 a cero, 2 sin stock (y la mitad de esas,
    con la marca de mas de 10.000). Cada 50, sin precio."""
    stock = {0: STOCK, 1: 0, 2: None}[g % 3]
    over = 'x' if g % 6 == 2 else None
    ean = 4999999000000 + g if g % 2 else str(4999999000000 + g)   # unas como numero y otras como texto
    return {'BRAND': 'MARCA-SECRETA', 'Description_short_ger_SDESC': 'NOMBRE-SECRETO %d' % g,
            'Item-no_BAN': 'CODSEC%05d' % g, 'EAN_GTIN': ean, 'MOQ': 6, 'Price_net': None if g % 50 == 7 else PRECIO,
            'Last price': PRECIO, 'SALE': 'x' if g % 10 == 0 else None, 'New': None,
            'Stock_available_pcs': stock, 'Stock_over10k_pcs': over, 'Ordered_pcs': None,
            'Discontinued': 'x' if g % 7 == 0 else None, 'MHD_EXP': None, 'ArtGrp_ger': 'Drogerie',
            'ArtGrp_eng': 'Drugstore', 'Inner_Pack_pcs_IP': 6, 'Outer_Case_pcs_OC': 24, 'GTIN_Case': None,
            'Layer_pcs': 120, 'PAL_pcs': 960, 'PAL_cm_HGT': 160, 'Weight_gram': 380, 'VAT': 19,
            'Description_long_eng': 'NOMBRE-SECRETO largo'}


def esperado(n):
    disp = sum(1 for g in range(n) if g % 3 == 0 or g % 6 == 2)
    return {'n': n, 'disponibles': disp, 'agotados': n - disp, 'sin_precio': sum(1 for g in range(n) if g % 50 == 7)}


_CACHE = {}


def xlsx(n=2600, cabecera=None, creado_hace=timedelta(hours=3), cambio=None, hoja='Data', creado=None):
    """Los bytes de un osma_articles.xlsx inventado. `cambio(filas)` toca las filas antes de escribir. `creado`
    (UTC) fija la fecha de creacion; si no, es la de hace `creado_hace`."""
    clave = (n, tuple(cabecera or ()), creado_hace, cambio.__name__ if cambio else None, hoja, creado)
    if clave not in _CACHE:
        cab = list(cabecera or COLUMNAS)
        filas = [fila_inventada(g) for g in range(n)]
        if cambio:
            cambio(filas)
        libro = openpyxl.Workbook(write_only=True)
        libro.properties.created = ((creado or (ahora() - creado_hace)).astimezone(timezone.utc)
                                    .replace(tzinfo=None, microsecond=0))
        h = libro.create_sheet(hoja)
        h.append(cab)
        for f in filas:
            # Por nombre; una columna renombrada lleva el dato de la que había en su sitio.
            h.append([f[c] if c in f else (f.get(COLUMNAS[i]) if i < len(COLUMNAS) else None) for i, c in enumerate(cab)])
        libro.create_sheet('Bilder').append(['NOMBRE-SECRETO foto'])
        b = io.BytesIO()
        libro.save(b)
        _CACHE[clave] = b.getvalue()
    return _CACHE[clave]


# ── La web de mentira ─────────────────────────────────────────────────────────────────────
def _li(id_, nombre):
    return f'<li class="file"><span class="name">{nombre}</span> <a href="/en/downloads/file/{id_}">Download</a></li>'


PAGINAS = {
    # La carpeta buena y otra con otro osma_articles.xlsx (que NO se puede coger).
    'buena': (f'<html><body><p>{HTML_SECRETO}</p><div class="albums">'
              f'<div class="album"><h3>Ordersatz_Updateliste</h3><ul>{_li(ID_FICHERO, "osma_articles.xlsx")}'
              f'{_li("aaaa1111aaaa1111aaaa1111aaaa1111", "Ordersatz_Update.xls")}</ul></div>'
              f'<div class="album"><h3>Archiv 2025</h3><ul>{_li(ID_OTRA_CARPETA, "osma_articles.xlsx")}</ul></div>'
              f'</div></body></html>'),
    # Las carpetas son otra página: aquí solo sus enlaces.
    'carpetas': (f'<html><body><p>{HTML_SECRETO}</p><a href="/en/downloads/album/c0ffee">Ordersatz_Updateliste</a>'
                 f'<a href="/en/downloads/album/beef">Archiv 2025</a></body></html>'),
    'carpeta_buena': (f'<html><body><h1>Ordersatz_Updateliste</h1><ul>{_li(ID_FICHERO, "osma_articles.xlsx")}'
                      f'{_li("aaaa1111aaaa1111aaaa1111aaaa1111", "Update_Gesamtsortiment_OSMA.xls")}</ul></body></html>'),
    # Sin ningún enlace (como sin sesión).
    'sin_enlace': f'<html><body><div class="account-download-center"><p>{HTML_SECRETO}</p></div></body></html>',
    # Dos osma_articles.xlsx en la misma carpeta.
    'doble': (f'<html><body><div class="album"><h3>Ordersatz_Updateliste</h3><ul>{_li(ID_FICHERO, "osma_articles.xlsx")}'
              f'{_li(ID_OTRA_CARPETA, "osma_articles.xlsx")}</ul></div></body></html>'),
}


class _Resp:
    def __init__(self, url, status=200, texto='', contenido=None):
        self.url, self.status_code = url, status
        self.text = texto
        self.content = contenido if contenido is not None else texto.encode('utf-8')
        self.headers = {'Set-Cookie': 'session=' + COOKIE, CABECERA_WEB: CABECERA_WEB}


class Web:
    """La web de OSMA de mentira: apunta cada llamada y responde segun el escenario."""

    def __init__(self, entrada_ok=True, pagina='buena', fichero=None, estado_fichero=200, error_red=False):
        self.entrada_ok, self.pagina, self.estado_fichero, self.error_red = entrada_ok, pagina, estado_fichero, error_red
        self.fichero = xlsx() if fichero is None else fichero
        self.llamadas = []

    def responder(self, metodo, url, datos):
        self.llamadas.append((metodo, url, datos))
        ruta = url.replace(WEB, '')
        if ruta == '/en/account/login' and metodo == 'GET':
            return _Resp(url, 200, f'<form>{HTML_SECRETO}</form>')
        if ruta == '/en/account/login' and metodo == 'POST':
            if self.entrada_ok:
                return _Resp(WEB + '/en/account', 200, f'<h1>{HTML_SECRETO}</h1>')
            return _Resp(WEB + '/en/account/login', 200, f'<div class="alert">{HTML_SECRETO}</div>')
        if ruta == '/en/downloads':
            return _Resp(url, 200, PAGINAS[self.pagina])
        if ruta == '/en/downloads/album/c0ffee':
            return _Resp(url, 200, PAGINAS['carpeta_buena'])
        if ruta.startswith('/en/downloads/file/'):
            if self.error_red:
                raise ConnectionError(f'HTTPSConnectionPool: Max retries exceeded with url: {url} ({CUERPO_ERROR})')
            if self.estado_fichero != 200:
                return _Resp(url, self.estado_fichero, f'<pre>{CUERPO_ERROR} {HTML_SECRETO}</pre>')
            return _Resp(url, 200, '', self.fichero)
        return _Resp(url, 404, f'{CUERPO_ERROR}')

    def de(self, metodo, trozo=''):
        return [x for x in self.llamadas if x[0] == metodo and trozo in x[1]]


def _modulo_requests(web, sesiones):
    m = types.ModuleType('requests')

    class Session:
        def __init__(self):
            self.headers = {}
            sesiones.append(self)

        def get(self, url, timeout=None, allow_redirects=True):
            return web.responder('GET', url, None)

        def post(self, url, data=None, timeout=None, allow_redirects=True):
            return web.responder('POST', url, data)

    m.Session = Session
    return m


# ── La base de mentira ────────────────────────────────────────────────────────────────────
class _Consulta:
    def __init__(self, base, tabla):
        self.base, self.tabla, self.accion, self.datos, self.filtros = base, tabla, None, None, {}

    def insert(self, d):
        self.accion, self.datos = 'insert', d
        return self

    def update(self, d):
        self.accion, self.datos = 'update', d
        return self

    def delete(self):
        self.accion = 'delete'
        return self

    def select(self, *_a, **_k):
        self.accion = 'select'
        return self

    def eq(self, col, val):
        self.filtros[col] = val
        return self

    def order(self, *_a, **_k):
        return self

    def limit(self, *_a):
        return self

    # 🆕 04-oct-2026 · lo que usa el repaso de las fechas de vuelta (paso 9). Aquí devuelve vacío: su banco es
    #    test_escaner2_osma_fecha_vuelta.py.
    def range(self, *_a):
        return self

    @property
    def not_(self):
        return self

    def is_(self, *_a):
        return self

    def execute(self):
        b = self.base
        b.ops.append((self.tabla, self.accion, dict(self.filtros)))
        r = []
        if self.tabla == 'disp_pasada':
            if self.accion == 'insert':
                pid = 'PASADA-OSMA-%d' % (len(b.pasadas) + 1)
                b.pasadas[pid] = dict(self.datos, id=pid)
                r = [{'id': pid}]
            elif self.accion == 'update':
                p = b.pasadas.get(self.filtros.get('id'))
                if p is not None and all(p.get(k) == v for k, v in self.filtros.items()):
                    p.update(self.datos)
            elif self.accion == 'select':
                if 'id' in self.filtros:
                    r = [dict(b.pasadas[self.filtros['id']])]
                elif self.filtros.get('estado') == 'aplicada':
                    r = [{'id': b.ultima_id, 'creada_en': b.ultima_aplicada.isoformat()}] if b.ultima_aplicada else []
                elif self.filtros.get('estado') == 'leyendo':
                    r = [{'id': i} for i, p in b.pasadas.items()
                         if p.get('estado') == 'leyendo' and p.get('run_id') == self.filtros.get('run_id')
                         and p.get('proveedor') == self.filtros.get('proveedor')]
        elif self.tabla == 'disp_parametros' and self.accion == 'select':
            r = [{'crudo_minimo': b.minimo}]
        elif self.tabla == 'disp_lectura':
            if self.accion == 'insert':
                if b.falla_lectura:
                    raise RuntimeError(f"APIError: new row violates check constraint · Failing row contains "
                                       f"({FILA_BASE}, CODSEC00001, 4999999000001, {PRECIO})")
                b.lectura.extend(self.datos)
                b.lotes.append(len(self.datos))
            elif self.accion == 'delete':
                b.lectura = [x for x in b.lectura if x['pasada_id'] != self.filtros.get('pasada_id')]
        return types.SimpleNamespace(data=r)


class Base:
    def __init__(self, minimo=2500, ultima_aplicada=None, falla_lectura=False, ultima_id='P-ANTES', falla_subida=False):
        self.minimo, self.ultima_aplicada, self.falla_lectura = minimo, ultima_aplicada, falla_lectura
        self.ultima_id, self.falla_subida = ultima_id, falla_subida
        self.ops, self.pasadas, self.lectura, self.lotes, self.subidas, self.rpcs = [], {}, [], [], [], []
        base = self

        class _Cubo:
            def __init__(self, nombre):
                self.nombre = nombre

            def upload(self, ruta, datos, opciones):
                base.ops.append(('storage', 'upload', {}))
                if base.falla_subida:
                    raise RuntimeError(f'StorageException 503 · {FILA_BASE} · /object/{ID_FICHERO}')
                base.subidas.append((self.nombre, ruta, datos, opciones))

        self.storage = types.SimpleNamespace(from_=_Cubo)

    def table(self, nombre):
        return _Consulta(self, nombre)

    def rpc(self, nombre, params):
        base = self

        class _Llamada:
            def execute(self):
                base.rpcs.append((nombre, params))
                base.ops.append(('rpc', nombre, {}))
                p = base.pasadas[params['p_pasada']]
                leidas = [x for x in base.lectura if x['pasada_id'] == p['id']]
                if p.get('n_leidas') == len(leidas) and p.get('n_crudo') == p.get('n_sin_gtin', 0) + len(leidas):
                    p.update(estado='aplicada', primera=True, n_en_catalogo=len(leidas),
                             n_disponibles_estado=sum(1 for x in leidas if x['disponible']), caida_aceptada=False)
                else:
                    p.update(estado='rechazada', motivo='no cuadra (de mentira)')
                base.lectura = [x for x in base.lectura if x['pasada_id'] != p['id']]
                return types.SimpleNamespace(data={'estado': p['estado']})
        return _Llamada()


def correr(programa, web=None, base=None, argv=(), secretos=True, run_id='4242', pasadas_previas=None):
    """El programa entero, como lo lanza el workflow (runpy, __main__), con `supabase` y `requests` de mentira.
    Devuelve (codigo, salida, base, web, sesiones, clientes)."""
    web = web or Web()
    base = base or Base()
    if pasadas_previas:
        base.pasadas.update(pasadas_previas)
    sesiones, clientes = [], []
    falso_supabase = types.ModuleType('supabase')

    def create_client(url, llave):
        clientes.append((url, llave))
        return base
    falso_supabase.create_client = create_client
    guardados = {k: sys.modules.get(k) for k in ('supabase', 'requests')}
    sys.modules['supabase'] = falso_supabase
    sys.modules['requests'] = _modulo_requests(web, sesiones)
    entorno = {'SUPABASE_URL': 'https://base.example.test', 'SUPABASE_SERVICE_KEY': 'LLAVE-DE-SERVICIO-SECRETA',
               'GITHUB_RUN_ID': run_id}
    if secretos:
        entorno.update(OSMA_USER=USUARIO, OSMA_PASS=CLAVE)
    viejo_env = {k: os.environ.get(k) for k in list(entorno) + ['OSMA_USER', 'OSMA_PASS']}
    for k in ('OSMA_USER', 'OSMA_PASS'):
        os.environ.pop(k, None)
    os.environ.update(entorno)
    viejo_argv = sys.argv
    sys.argv = [programa] + list(argv)
    salida = io.StringIO()
    codigo = None
    try:
        with contextlib.redirect_stdout(salida), contextlib.redirect_stderr(salida):
            try:
                runpy.run_path(programa, run_name='__main__')
                codigo = 0
            except SystemExit as e:
                codigo = e.code if isinstance(e.code, int) else (0 if e.code is None else 1)
    finally:
        sys.argv = viejo_argv
        for k, v in viejo_env.items():
            if v is None:
                os.environ.pop(k, None)
            else:
                os.environ[k] = v
        for k, v in guardados.items():
            if v is None:
                sys.modules.pop(k, None)
            else:
                sys.modules[k] = v
    return codigo, salida.getvalue(), base, web, sesiones, clientes


# ── Los casos, sobre un programa (el de verdad o un mutante) ─────────────────────────────────
def _repetido(filas):
    filas[10]['Item-no_BAN'] = filas[11]['Item-no_BAN']


def _vacio(filas):
    filas[20]['Item-no_BAN'] = None


def comprobar(programa, decir):
    """Corre todos los casos sobre `programa` y llama a decir(nombre, obtenido, esperado). Devuelve los
    nombres de los que fallan."""
    rotos = []

    def chk(nombre, obtenido, esperado_):
        if obtenido != esperado_:
            rotos.append(nombre)
        decir(nombre, obtenido, esperado_)

    salidas = {}

    def pasada_de(base):
        return next(iter(base.pasadas.values())) if base.pasadas else {}

    def sin_bajar(web):
        return web.de('GET', '/downloads/file/') == []

    def nada_en_la_base(base):
        return (base.rpcs, base.lotes)

    # (A) el bueno
    cod, out, base, web, ses, _ = correr(programa)
    salidas['bueno'] = out
    p = pasada_de(base)
    e = esperado(2600)
    chk('(A) el bueno: sale bien y la pasada queda aplicada', (cod, p.get('estado')), (0, 'aplicada'))
    chk('(A) …UNA entrada (un POST) con el usuario y la clave de los secretos',
        [(x[2] or {}).get('username') == USUARIO and (x[2] or {}).get('password') == CLAVE for x in web.de('POST')], [True])
    chk('(A) …baja el fichero de la carpeta buena, no el de la otra carpeta',
        [x[1] for x in web.de('GET', '/downloads/file/')], [WEB + '/en/downloads/file/' + ID_FICHERO])
    chk('(A) …la copia ENTERA, tras aplicar, en escaner2/osma/ultimo/osma_articles.xlsx',
        ([(s[0], s[1], s[2] == web.fichero, s[3].get('upsert')) for s in base.subidas],
         [o[0] for o in base.ops if o[0] in ('rpc', 'storage')], 'OSMA_COPIA_NO_GUARDADA' in out),
        ([('escaner2', 'osma/ultimo/osma_articles.xlsx', True, 'true')], ['rpc', 'storage'], False))
    cod_c, out_c, base_c, _w, _s, _c = correr(programa, base=Base(falla_subida=True))
    salidas['la copia falla'] = out_c
    chk('(A) si la copia falla con la pasada aplicada: ROJO, lo dice, y la pasada sigue aplicada',
        (cod_c, 'OSMA_COPIA_NO_GUARDADA' in out_c, pasada_de(base_c).get('estado'), base_c.subidas), (1, True, 'aplicada', []))
    chk('(A) …declarado = llegados = filas; sin GTIN y repetidos a 0',
        tuple(p.get(k) for k in ('n_declarado', 'n_crudo', 'n_declarado_precios', 'n_precios', 'n_declarado_disponibilidades',
                                 'n_disponibilidades', 'n_sin_gtin', 'n_duplicados')), (2600,) * 6 + (0, 0))
    chk('(A) …leídas, disponibles (stock > 0 o más de 10.000), agotados y sin precio',
        tuple(p.get(k) for k in ('n_leidas', 'n_disponibles', 'n_agotados', 'n_sin_dato_precio')),
        (e['n'], e['disponibles'], e['agotados'], e['sin_precio']))
    chk('(A) …en lotes de 500, y UNA llamada a disp_aplicar_pasada',
        (base.lotes, [r[0] for r in base.rpcs]), ([500, 500, 500, 500, 500, 100], ['disp_aplicar_pasada']))
    chk('(A) …sin fila en disp_lectura al acabar (es de paso)', base.lectura, [])
    chk('(A) …y DESPUÉS de aplicar y de la copia, el repaso de las fechas de vuelta (paso 9), que lo dice',
        ([o[0] for o in base.ops if o[0] in ('rpc', 'storage', 'disp_estado')][:3], 'FECHAS DE VUELTA' in out,
         'OSMA_FECHAS' in out),
        (['rpc', 'storage', 'disp_estado'], True, False))

    # La conversión, fila a fila, con la base que no aplica (para ver lo subido).
    base_ver = Base()
    base_ver.rpc = lambda n, prm: types.SimpleNamespace(execute=lambda: types.SimpleNamespace(data={}))
    cod_v, _o, base_ver, _w, _s, _c = correr(programa, base=base_ver)
    sub = {x['producto_prov']: x for x in base_ver.lectura}
    f0, f1, f2, f7 = sub.get('CODSEC00000', {}), sub.get('CODSEC00001', {}), sub.get('CODSEC00002', {}), sub.get('CODSEC00007', {})
    chk('(A) conversión: llave = código; precio neto por unidad en los dos precios; EAN de texto y de número',
        (f0.get('precio_catalogo'), f0.get('precio_unidad'), f0.get('ean_core'), f1.get('ean_core')),
        (PRECIO, PRECIO, '4999999000000', '4999999000001'))
    chk('(A) conversión: con stock disponible; a cero agotado; sin stock con «más de 10.000» disponible',
        (f0.get('disponible'), f1.get('disponible'), f2.get('disponible')), (True, False, True))
    chk('(A) conversión: ni caja ni chase; oferta y descatalogado de sus marcas; sin precio marcado',
        (f0.get('es_caja'), f0.get('uds_caja'), f0.get('es_chase'), f0.get('en_oferta'), f0.get('fin_de_vida'),
         f7.get('precio_unidad'), f7.get('sin_dato_precio')), (False, None, False, True, True, None, True))

    # (A) con la carpeta en otra página
    cod, out, base, web, ses, _ = correr(programa, web=Web(pagina='carpetas'))
    salidas['carpeta_aparte'] = out
    chk('(A) la carpeta en otra página: la abre, baja el fichero y aplica',
        (cod, pasada_de(base).get('estado'), [x[1] for x in web.de('GET', '/downloads/file/')]),
        (0, 'aplicada', [WEB + '/en/downloads/file/' + ID_FICHERO]))

    # (B) los que no aplican
    casos_b = [
        ('cabecera cambiada', Web(fichero=xlsx(cabecera=[('Price_net_EUR' if c == 'Price_net' else c) for c in COLUMNAS])),
         Base(), 'fallida', 'Price_net'),
        ('cabecera fuera de orden', Web(fichero=xlsx(cabecera=['Item-no_BAN', 'BRAND'] + [c for c in COLUMNAS
                                                                                          if c not in ('BRAND', 'Item-no_BAN')])),
         Base(), 'fallida', 'fuera de su orden'),
        ('sin hoja Data', Web(fichero=xlsx(hoja='Datos')), Base(), 'fallida', 'hoja Data'),
        ('filas de menos', Web(fichero=xlsx(n=2400)), Base(), 'rechazada_vaciado', 'por debajo del mínimo 2500'),
        ('código repetido', Web(fichero=xlsx(cambio=_repetido)), Base(), 'fallida', 'repetido'),
        ('código vacío', Web(fichero=xlsx(cambio=_vacio)), Base(), 'fallida', 'sin código de artículo'),
        ('fichero de hace 3 días', Web(fichero=xlsx(creado_hace=timedelta(days=3))), Base(), 'rechazada',
         'OSMA no ha regenerado el fichero'),
        ('lo bajado no es un xlsx', Web(fichero=b'<html>' + HTML_SECRETO.encode() + b'</html>'), Base(), 'fallida',
         'no es un xlsx'),
        ('la descarga da 500', Web(estado_fichero=500), Base(), 'fallida', 'respondió 500'),
        ('la red cae al bajar', Web(error_red=True), Base(), 'fallida', 'ConnectionError'),
        ('la base rechaza una fila', Web(), Base(falla_lectura=True), 'fallida', 'RuntimeError'),
    ]
    g = runpy.run_path(programa, run_name='osma_banco')
    finde = g['hora_de_madrid'](ahora()).weekday() >= 5
    for nombre, w, b, estado, texto in casos_b:
        cod, out, base, web, ses, _ = correr(programa, web=w, base=b)
        salidas[nombre] = out
        p = pasada_de(base)
        cod_esperado = 0 if (estado == 'rechazada' and finde) else 1
        chk('(B) %s: no aplica, sale %s y dice por qué' % (nombre, estado),
            (cod, p.get('estado'), texto in (p.get('motivo') or ''), texto in out), (cod_esperado, estado, True, True))
        chk('(B) %s: ni una fila en la base, ni llamada a la función, ni copia' % nombre,
            (base.rpcs, base.lectura, base.subidas), ([], [], []))
        chk('(B) %s: ni fechas de vuelta (solo tras una pasada aplicada)' % nombre,
            ([o for o in base.ops if o[0] == 'disp_estado'], 'FECHAS DE VUELTA' in out), ([], False))
    # Lo que no es un xlsx no se guarda en el almacén.
    cod, out, base, web, ses, _ = correr(programa, web=Web(fichero=b'<html>' + HTML_SECRETO.encode() + b'</html>'))
    chk('(B) lo que no es un xlsx no se guarda en el almacén', base.subidas, [])
    # El error de la base: su detalle (con la fila) va a la base, que es privada; al registro, no.
    cod, out, base, web, ses, _ = correr(programa, base=Base(falla_lectura=True))
    chk('(B) el error de la base: el detalle queda en disp_pasada.motivo (privado)', FILA_BASE in (pasada_de(base).get('motivo') or ''),
        True)

    # (C) sin bajar nada
    for nombre, w, texto in [('entrada rechazada', Web(entrada_ok=False), 'entrada rechazada'),
                             ('enlace no encontrado', Web(pagina='sin_enlace'), 'enlace no encontrado'),
                             ('enlace doble', Web(pagina='doble'), 'enlace')]:
        cod, out, base, web, ses, _ = correr(programa, web=w)
        salidas[nombre] = out
        p = pasada_de(base)
        chk('(C) %s: fallida, dice por qué, sin bajar nada ni guardar nada' % nombre,
            (cod, p.get('estado'), texto in (p.get('motivo') or ''), sin_bajar(web), base.subidas, base.rpcs),
            (1, 'fallida', True, True, [], []))
        if nombre == 'entrada rechazada':
            chk('(C) entrada rechazada: UN solo intento (un POST) y ni siquiera abre la página de descargas',
                (len(web.de('POST')), web.de('GET', '/en/downloads')), (1, []))

    # (D) sin secretos
    cod, out, base, web, ses, clientes = correr(programa, secretos=False)
    salidas['sin secretos'] = out
    chk('(D) sin OSMA_USER/OSMA_PASS: rojo, lo dice, y ni pasada, ni cliente de base, ni web',
        (cod, 'faltan OSMA_USER/OSMA_PASS' in out, base.pasadas, clientes, ses, web.llamadas), (1, True, {}, [], [], []))

    # (F) el rescate
    previas = {'P-ESTE': {'id': 'P-ESTE', 'proveedor': 'OSMA', 'estado': 'leyendo', 'run_id': 4242},
               'P-OTRO-RUN': {'id': 'P-OTRO-RUN', 'proveedor': 'OSMA', 'estado': 'leyendo', 'run_id': 1},
               'P-HEO': {'id': 'P-HEO', 'proveedor': 'HEO', 'estado': 'leyendo', 'run_id': 4242},
               'P-YA': {'id': 'P-YA', 'proveedor': 'OSMA', 'estado': 'aplicada', 'run_id': 4242}}
    cod, out, base, web, ses, _ = correr(programa, argv=['--rescate'], pasadas_previas=previas)
    salidas['rescate'] = out
    chk('(F) el rescate: solo la de ESTE run, de OSMA y leyendo, pasa a fallida; sin tocar la web',
        (cod, {k: v['estado'] for k, v in base.pasadas.items()}, web.llamadas),
        (0, {'P-ESTE': 'fallida', 'P-OTRO-RUN': 'leyendo', 'P-HEO': 'leyendo', 'P-YA': 'aplicada'}, []))

    # (I) LA FECHA DEL FICHERO, con la hora fijada (el módulo importado, sin __main__). Octubre de 2026: Madrid = UTC+2.
    def fija(nombre, ahora_utc, ultima_utc, creado_utc):
        w = Web(fichero=xlsx(creado=datetime.fromisoformat(creado_utc + '+00:00')))
        sys.modules['requests'] = _modulo_requests(w, [])
        b = Base(ultima_aplicada=datetime.fromisoformat(ultima_utc + '+00:00'), ultima_id='P-ANTES')
        salida = io.StringIO()
        try:
            with contextlib.redirect_stdout(salida):
                r = g['pasada'](b, sys.modules['requests'].Session(), USUARIO, CLAVE, '4242',
                                ahora=datetime.fromisoformat(ahora_utc + '+00:00'))
        finally:
            sys.modules.pop('requests', None)
        salidas[nombre] = salida.getvalue()
        return r, b, pasada_de(b), salida.getvalue()

    # (a) miércoles 07-oct, 09:00 de Madrid; la de las 07:17 de Madrid ya aplicada; el fichero, de las 03:33 UTC.
    r, b, p, out = fija('al día', '2026-10-07T07:00:00', '2026-10-07T05:17:00', '2026-10-07T03:33:00')
    chk('(a) relanzar el mismo día: VERDE, OSMA_AL_DIA con la pasada de esta mañana',
        (r, 'OSMA_AL_DIA: el fichero de hoy ya está aplicado (pasada P-ANTES); nada nuevo' in out), (0, True))
    chk('(a) …no cuenta como rechazo: «rechazada» por «al día» (ni rechazada_vaciado ni fallida), sin filas, sin función, sin copia',
        (p.get('estado'), (p.get('motivo') or '').startswith('al día'), b.rpcs, b.lectura, b.subidas),
        ('rechazada', True, [], [], []))
    # (b) el día es el de MADRID: miércoles 00:10 de Madrid (martes 22:10 UTC); la última, el martes a las 23:50 de
    #     Madrid (21:50 UTC: el MISMO día en UTC). Es de ayer → OSMA no ha renovado, y es laborable: ROJO.
    r, b, p, out = fija('medianoche de Madrid', '2026-10-06T22:10:00', '2026-10-06T21:50:00', '2026-10-06T03:33:00')
    chk('(b) a las 00:10 de Madrid, la de las 23:50 es de AYER (aunque en UTC sea hoy): no es «al día», ROJO',
        (r, p.get('estado'), 'OSMA_AL_DIA' in out, 'OSMA no ha regenerado' in out), (1, 'rechazada', False, True))
    # (b) laborable: miércoles 09:00 de Madrid; la última, la del martes 07:17; el fichero, del martes 03:33 UTC.
    r, b, p, out = fija('laborable sin renovar', '2026-10-07T07:00:00', '2026-10-06T05:17:00', '2026-10-06T03:33:00')
    chk('(b) laborable con el fichero de ayer ya aplicado: ROJO, «OSMA no ha regenerado», sin filas ni copia',
        (r, p.get('estado'), 'OSMA_NO_APLICADA' in out, 'OSMA no ha regenerado' in out, b.lectura, b.subidas),
        (1, 'rechazada', True, True, [], []))
    # (b) sábado 10-oct 09:00 de Madrid; la última, la del viernes; el fichero, del viernes.
    r, b, p, out = fija('sábado sin renovar', '2026-10-10T07:00:00', '2026-10-09T05:17:00', '2026-10-09T03:33:00')
    chk('(b) sábado con el fichero del viernes: VERDE (OSMA_SIN_FICHERO_NUEVO), sin filas ni copia',
        (r, p.get('estado'), 'OSMA_SIN_FICHERO_NUEVO' in out, b.lectura, b.subidas), (0, 'rechazada', True, [], []))
    # (b) y el fichero nuevo del día siguiente, con la de ayer aplicada: se aplica.
    r, b, p, out = fija('día siguiente renovado', '2026-10-08T07:00:00', '2026-10-07T05:17:00', '2026-10-08T03:33:00')
    chk('(b) …y el jueves, con el fichero del jueves y la del miércoles aplicada: se aplica y guarda copia',
        (r, p.get('estado'), [x[1] for x in b.subidas]), (0, 'aplicada', ['osma/ultimo/osma_articles.xlsx']))

    # (E) EL REGISTRO NO SUELTA DATOS, en todos los casos
    for nombre, out in salidas.items():
        chk('(E) %s: el registro no lleva ningún dato' % nombre, [c for c in CENTINELAS if c in out], [])
    return rotos


# ── Las piezas sueltas (puras) ──────────────────────────────────────────────────────────────
m = runpy.run_path(PROGRAMA, run_name='osma_piezas')
eq('ean_core: 13 cifras de texto', m['ean_core']('4002448039440'), ('4002448039440', '4002448039440', None))
eq('ean_core: número que perdió el 0 de delante (11 cifras) → 13', m['ean_core'](12345678905),
   ('12345678905', '0012345678905', None))
eq('ean_core: 14 con un 0 delante → 13', m['ean_core']('04002448039440')[1], '4002448039440')
eq('ean_core: vacío, sin regla; raro, regla', (m['ean_core'](None), m['ean_core']('ABC')[2]), ((None, None, None), 'ean_forma_rara'))
eq('marcado: x, 1, «yes» sí; vacío, 0, «no», «nein» no',
   [m['marcado'](v) for v in ('x', 1, 'yes', None, '', 0, 'no', 'Nein')], [True, True, True, False, False, False, False, False])
eq('numero: 3,777 y 3.777; vacío None', (m['numero']('3,777'), m['numero']('3.777'), m['numero'](' ')), (3.777, 3.777, None))
try:
    m['numero']('>10000')
    eq('numero: lo que no es un número no se inventa', 'pasó', 'ValueError')
except ValueError:
    eq('numero: lo que no es un número no se inventa', 'ValueError', 'ValueError')
eq('texto_codigo: 1538 y 1538.0 son «1538»', (m['texto_codigo'](1538), m['texto_codigo'](1538.0)), ('1538', '1538'))
eq('fecha de creación: la del docProps del Excel', m['fecha_de_creacion'](xlsx(creado_hace=timedelta(hours=5))).tzinfo is not None, True)


# ── (G) Por estructura ─────────────────────────────────────────────────────────────────────
arbol = ast.parse(io.open(PROGRAMA, encoding='utf-8').read())
tablas, rpcs = set(), set()
for n in ast.walk(arbol):
    if isinstance(n, ast.Call) and isinstance(n.func, ast.Attribute) and n.args and isinstance(n.args[0], ast.Constant):
        if n.func.attr == 'table':
            tablas.add(n.args[0].value)
        elif n.func.attr == 'rpc':
            rpcs.add(n.args[0].value)
eq('(G) el programa solo toca disp_pasada, disp_lectura y disp_parametros', sorted(tablas),
   ['disp_lectura', 'disp_parametros', 'disp_pasada'])
eq('(G) …y solo llama a disp_aplicar_pasada', sorted(rpcs), ['disp_aplicar_pasada'])
eq('(G) …sin traceback ni logging (nada que vuelque una petición)',
   sorted({a.name for n in ast.walk(arbol) if isinstance(n, (ast.Import, ast.ImportFrom))
           for a in n.names if a.name in ('traceback', 'logging')}), [])

wf = yaml.safe_load(io.open(WORKFLOW, encoding='utf-8'))
disparo = wf.get('on', wf.get(True))
pasos = wf['jobs']['leer']['steps']
eq('(G) el workflow solo se lanza a mano (sin schedule: el reloj es cron-job.org)', sorted(disparo), ['workflow_dispatch'])
eq('(G) …con grupo propio, sin cancelar el que corre', (wf['concurrency']['group'], wf['concurrency']['cancel-in-progress']),
   ('escaner2-osma', False))
eq('(G) …el PRIMER paso comprueba los secretos', (pasos[0].get('id'), 'OSMA_USER' in pasos[0]['env'],
                                                 'faltan OSMA_USER/OSMA_PASS' in pasos[0]['run']), ('secretos', True, True))
lanza = [p for p in pasos if 'escaner2_osma_disponibilidad.py' in (p.get('run') or '') and '--rescate' not in p['run']]
rescate = [p for p in pasos if '--rescate' in (p.get('run') or '')]
eq('(G) …el paso que lo lanza lleva OSMA_USER, OSMA_PASS y la llave de servicio',
   [sorted(k for k in p['env'] if k in ('OSMA_USER', 'OSMA_PASS', 'SUPABASE_SERVICE_KEY')) for p in lanza],
   [['OSMA_PASS', 'OSMA_USER', 'SUPABASE_SERVICE_KEY']])
eq('(G) …el rescate solo corre si el run falló Y los secretos estaban, y sin las credenciales de OSMA',
   [('steps.secretos.outcome' in p['if'], 'failure()' in p['if'], 'OSMA_PASS' in p['env']) for p in rescate],
   [(True, True, False)])
eq('(G) …ningún paso con modo verboso', [p.get('name') for p in pasos
                                        if 'ACTIONS_STEP_DEBUG' in str(p) or ' -v' in (p.get('run') or '')], [])


# ── El programa de verdad ──────────────────────────────────────────────────────────────────
print('\n── El programa de verdad ──')
rotos_reales = comprobar(PROGRAMA, eq)


# ── (H) LA MITAD ROJA: el banco se pone en rojo con cada mutante ───────────────────────────
print('\n── (H) La mitad roja ──')
fuente = io.open(PROGRAMA, encoding='utf-8').read()
MUTANTES = [
    ('sin_cabecera', '        if pos is None:\n', '        if False:\n',
     '(B) cabecera cambiada: no aplica, sale fallida y dice por qué'),
    ('sin_fecha_vieja', '    if ahora - creado > MAX_EDAD:\n', '    if False:\n',
     '(B) fichero de hace 3 días: no aplica, sale rechazada y dice por qué'),
    ('sin_codigos_repetidos', '    if repetidos:\n', '    if False:\n',
     '(B) código repetido: no aplica, sale fallida y dice por qué'),
    ('imprime_el_error', "        motivo_log = f'error inesperado ({type(ex).__name__}); el detalle queda en disp_pasada.motivo'\n",
     "        motivo_log = f'error inesperado ({type(ex).__name__}): {ex}'\n",
     '(E) la base rechaza una fila: el registro no lleva ningún dato'),
    ('dos_intentos', "        raise Rechazo('fallida', 'entrada rechazada: la web ha vuelto a la página de entrada; no se reintenta '\n",
     "        sesion.post(WEB + RUTA_ENTRADA, data={'username': usuario, 'password': clave}, timeout=ESPERA_S)\n"
     "        raise Rechazo('fallida', 'entrada rechazada: la web ha vuelto a la página de entrada; no se reintenta '\n",
     '(C) entrada rechazada: UN solo intento (un POST) y ni siquiera abre la página de descargas'),
    ('al_dia_en_rojo',
     '            print(f"OSMA_AL_DIA: el fichero de hoy ya está aplicado (pasada {ex.pasada_aplicada}); nada nuevo.", flush=True)\n'
     '            return 0\n',
     '            print(f"OSMA_AL_DIA: el fichero de hoy ya está aplicado (pasada {ex.pasada_aplicada}); nada nuevo.", flush=True)\n'
     '            return 1\n',
     '(a) relanzar el mismo día: VERDE, OSMA_AL_DIA con la pasada de esta mañana'),
    ('dia_en_utc', '        if hora_de_madrid(ultima_aplicada).date() == hora_de_madrid(ahora).date():\n',
     '        if ultima_aplicada.date() == ahora.date():\n',
     '(b) a las 00:10 de Madrid, la de las 23:50 es de AYER (aunque en UTC sea hoy): no es «al día», ROJO'),
    ('copia_antes_de_validar', '        print(f">>> Fichero bajado: {len(contenido)} bytes.", flush=True)\n',
     '        print(f">>> Fichero bajado: {len(contenido)} bytes.", flush=True)\n        guardar_copia(sb, contenido)\n',
     '(B) cabecera cambiada: ni una fila en la base, ni llamada a la función, ni copia'),
    ('copia_por_pasada', "RUTA_COPIA = 'osma/ultimo/' + FICHERO\n", "RUTA_COPIA = 'osma/' + FICHERO\n",
     '(A) …la copia ENTERA, tras aplicar, en escaner2/osma/ultimo/osma_articles.xlsx'),
    ('copia_falla_en_verde', '        codigo = 1  # la copia no se ha guardado: que se vea\n', '        codigo = 0\n',
     '(A) si la copia falla con la pasada aplicada: ROJO, lo dice, y la pasada sigue aplicada'),
]
with tempfile.TemporaryDirectory() as carpeta:
    for nombre, viejo, nuevo, vigila in MUTANTES:
        eq('(H) el ancla del mutante %s está UNA vez en el programa' % nombre, fuente.count(viejo), 1)
        ruta = os.path.join(carpeta, 'mutante_%s.py' % nombre)
        io.open(ruta, 'w', encoding='utf-8').write(fuente.replace(viejo, nuevo))
        rotos = comprobar(ruta, lambda *_a: None)
        eq('(H) con el mutante %s el banco se pone ROJO en «%s»' % (nombre, vigila), vigila in rotos, True)

print()
if fallos:
    print('ROJO: %d fallo(s): %s' % (len(fallos), '; '.join(fallos)))
    sys.exit(1)
print('VERDE: la pasada de OSMA entra una vez, baja el fichero de su carpeta, valida antes de subir nada, aplica con '
      'la función común y SOLO ENTONCES guarda la copia en la ruta fija; relanzar el mismo día sale en verde; su '
      'registro no suelta ningún dato; y cada mutante lo pone en rojo.')
