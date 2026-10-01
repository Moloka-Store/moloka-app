# -*- coding: utf-8 -*-
"""Banco del encargo AA (1-oct-2026): el umbral «7 o más» y la FICHA COMPARTIDA del escaner 2 de HEO.

SIN RED, SIN SECRETOS Y SIN BASE: el motor de verdad (escaner2_motor) con CSV de mentira escritos con las cabeceras
EXACTAS de los exports del Visualizador (las de `TIPADAS` del escaparate y `CSV_COLS` del Escaner Pro).

QUE PRUEBA:
  (A) EL UMBRAL SALE DEL PARAMETRO: con umbral_caidas_30d = 6, 7 caidas entran y 6 no; con 8, esas 7 no entran (no hay
      un 7 escrito en el codigo). Y 6,5 por figura NO pasa con 6: el corte es «7 o más», no «más de 6».
      Por estructura: en el motor, el cruce y novedades ninguna comparacion de caidas/ventas lleva un numero escrito, y
      los dos programas leen umbral_caidas_30d de escaner2_parametros.
  (B) LA DETECCION, por pais: dos hermanos con el mismo «ASIN Padre» y el mismo puesto → reparto; mismo padre y otro
      puesto → sin reparto; padre o puesto vacio → sin reparto; sin «Recuento de variaciones» → ÷ los hermanos de la
      lista, y el detalle lo dice. Un CSV sin la columna del padre: se usa igual, sin reparto, y se dice que falta.
  (C) EL REPARTO EN LAS PUERTAS: 30 ventas ÷ 3 variaciones = 10 → entra y lleva la marca (detalle y Excel); 8 ÷ 7 →
      puerta c con el motivo propio «c_ficha_compartida:» al principio del detalle (el motivo de la base sigue siendo
      c_pocas_caidas: su CHECK solo admite ese). Cada pais guarda las caidas de la FICHA, no las repartidas.
  (D) EL EXCEL: «Ficha compartida» es la ULTIMA columna de «Análisis», dentro de la tabla T_Analisis; la fila del pais
      compartido lleva la marca, las demas filas del EAN dicen en que pais lo es, y un EAN sin reparto la deja vacia.
"""
import ast
import csv
import io
import os
import shutil
import sys
import tempfile

import escaner2_motor as e2
import escaner2_heredado_pro as pro

AQUI = os.path.dirname(os.path.abspath(__file__))
fallos = []


def eq(nombre, obtenido, esperado):
    ok = obtenido == esperado
    if not ok:
        fallos.append(nombre)
    print(('OK ' if ok else 'XX ') + nombre + ('' if ok else '   got=%r exp=%r' % (obtenido, esperado)))


M = e2.cargar_motor()
M.poner_catalogo_propio([])
COL_PAIS, COL_CAIDAS = e2.columnas_keepa()
COLS_FICHA = e2.columnas_ficha_compartida()
COL_PADRE, COL_NVAR, COL_RANK = COLS_FICHA
C = pro.CSV_COLS
CABECERA = ['ASIN', COL_PAIS, 'Título', C['ean'], C['rank'], C['rank90'], COL_CAIDAS, C['buybox'], C['es_fba'],
            C['nuevo'], C['fba'], C['compct'], COL_PADRE, COL_NVAR]
_tmp = tempfile.mkdtemp(prefix='e2aa_')


def ean(n):
    cuerpo = '84355%07d' % n
    return cuerpo + M._chk13(cuerpo)


def fila(asin, n, caidas, padre='', rank='5000', nvar='', bb='22.00', pais='es'):
    return {'ASIN': asin, COL_PAIS: pais, 'Título': 'Funko Pop %s' % asin, C['ean']: ean(n), C['rank']: rank,
            C['rank90']: '6000', COL_CAIDAS: caidas, C['buybox']: bb, C['es_fba']: 'yes', C['nuevo']: bb,
            C['fba']: '3.50', C['compct']: '15.01 %', COL_PADRE: padre, COL_NVAR: nvar}


def escribir(nombre, filas, cabecera=CABECERA):
    ruta = os.path.join(_tmp, nombre)
    buf = io.StringIO()
    w = csv.DictWriter(buf, fieldnames=cabecera, extrasaction='ignore')
    w.writeheader()
    for f in filas:
        w.writerow(f)
    io.open(ruta, 'w', encoding='utf-8-sig', newline='').write(buf.getvalue())
    return ruta


def foto(n, precio=8.0):
    return {'id': 'F%d' % n, 'ean_original': ean(n), 'ean_core': ean(n), 'variantes': [M.norm(ean(n))],
            'nombre': 'Funko Pop %d' % n, 'marca': 'Funko', 'precio_unidad': precio, 'aviso_caja': None}


def cruzar(filas, umbral, n_foto, cabecera=CABECERA):
    """Un CSV de ES con esas filas → (resultados por numero de foto, compartidas, examen del CSV): el camino del
    cruce (examinar_csv → fichas_compartidas → decidir) sin red."""
    ruta = escribir('es_%d_%d.csv' % (umbral, len(filas)), filas, cabecera)
    ex = e2.examinar_csv(ruta, pro, COL_PAIS, COL_CAIDAS, COLS_FICHA)
    comp = e2.fichas_compartidas({ex['pais']: ex['fichas']})
    datos = pro.leer_csv_visualizador(ruta)
    params = {'umbral': umbral, 'paises_filtro': ['ES', 'IT', 'FR', 'DE'], 'paises_calculo': ['ES', 'IT', 'FR', 'DE']}
    res = {}
    for n in n_foto:
        f = foto(n)
        res[n] = e2.decidir(f, {'ES': e2.candidatos(f, datos)}, {'ES': ex['caidas']}, params, M, None, comp)
    return res, comp, ex


# ═══════════════════════════════════════════════════════════════════════════════
print('(A) el umbral sale del parámetro: «7 o más»')
_r6, _, _ = cruzar([fila('B0SIETE001', 1, '7'), fila('B0SEIS0001', 2, '6')], 6, (1, 2))
eq('(A) umbral 6: 7 caídas entran', _r6[1]['puerta'] in ('d', 'e', 'f'), True)
eq('(A) umbral 6: 6 caídas no entran, y el texto dice el umbral del parámetro',
   (_r6[2]['puerta'], _r6[2]['motivo'], _r6[2]['detalle']), ('c', 'c_pocas_caidas', '≤ 6 caídas en ES (IT, FR, DE: sin dato)'))
_r8, _, _ = cruzar([fila('B0SIETE001', 1, '7')], 8, (1,))
eq('(A) umbral 8 (el de antes): esas mismas 7 no entran → el 7 no está escrito en el código',
   (_r8[1]['puerta'], _r8[1]['detalle']), ('c', '≤ 8 caídas en ES (IT, FR, DE: sin dato)'))
eq('(A) «vende aquí» sigue el mismo corte', (_r6[1]['paises']['ES']['vende_aqui'], _r6[2]['paises']['ES']['vende_aqui']),
   (True, False))
eq('(A) el corte con las palabras de Fernando, sacado del parámetro', (e2.texto_corte(6), e2.texto_corte(8)),
   ('7 o más caídas en 30 días', '9 o más caídas en 30 días'))
eq('(A) 🔴 6,5 por figura con umbral 6 NO pasa (es «7 o más», no «más de 6»); 7 sí; sin dato, no',
   (e2.pasa_corte(6.5, 6), e2.pasa_corte(7, 6), e2.pasa_corte(6.99, 6), e2.pasa_corte(None, 6)), (False, True, False, False))
_r65, _, _ = cruzar([fila('B0MEDIO001', 1, '13', padre='B0PADRE065', nvar='2'),
                     fila('B0MEDIO002', 2, '13', padre='B0PADRE065', nvar='2')], 6, (1,))
eq('(A) 🔴 …y en las puertas: 13 caídas ÷ 2 figuras = 6,5 → puerta c', _r65[1]['puerta'], 'c')


def _comparaciones_con_numero(ruta):
    """Las comparaciones de caidas/ventas (por nombre) que llevan un NUMERO escrito: no deberia haber ninguna."""
    arbol = ast.parse(io.open(os.path.join(AQUI, ruta), encoding='utf-8').read())
    malas = []
    for n in ast.walk(arbol):
        if not isinstance(n, ast.Compare):
            continue
        partes = [n.left] + list(n.comparators)
        texto = ast.unparse(n)
        if any(k in texto for k in ('caidas', 'ventas', 'umbral')) and any(
                isinstance(p, ast.Constant) and isinstance(p.value, (int, float)) and not isinstance(p.value, bool)
                and p.value not in (0, 1) for p in partes):
            malas.append('%s:%d %s' % (ruta, n.lineno, texto))
    return malas


eq('(A) por estructura: en el motor, el cruce y novedades ninguna comparación de caídas lleva un número escrito',
   [m for r in ('escaner2_motor.py', 'escaner2_heo_cruce.py', 'escaner2_novedades.py') for m in _comparaciones_con_numero(r)], [])
def _constantes_en(ruta, nodo_tipo):
    """Las cadenas que aparecen como argumento de una llamada (`ast.Call`) o como indice (`ast.Subscript`): por
    estructura, asi un comentario o un texto de ayuda no cuentan."""
    arbol = ast.parse(io.open(os.path.join(AQUI, ruta), encoding='utf-8').read())
    salida = set()
    for n in ast.walk(arbol):
        if nodo_tipo is ast.Call and isinstance(n, ast.Call):
            salida |= {a.value for a in n.args if isinstance(a, ast.Constant) and isinstance(a.value, str)}
        if nodo_tipo is ast.Subscript and isinstance(n, ast.Subscript) and isinstance(n.slice, ast.Constant):
            salida.add(n.slice.value)
    return salida


for _prog in ('escaner2_heo_cruce.py', 'escaner2_novedades.py'):
    eq('(A) %s lee umbral_caidas_30d de escaner2_parametros (por estructura: una llamada y un índice)' % _prog,
       ('escaner2_parametros' in _constantes_en(_prog, ast.Call), 'umbral_caidas_30d' in _constantes_en(_prog, ast.Subscript)),
       (True, True))

# ═══════════════════════════════════════════════════════════════════════════════
print('\n(B) la detección de la ficha compartida')
eq('(B) las tres cabeceras salen de TIPADAS del escaparate (las de los exports reales)',
   COLS_FICHA, ('ASIN Padre', 'Recuento de variaciones', 'Clasificación de Ventas: Actual'))
_, _comp, _ex = cruzar([fila('B0HERMA001', 1, '30', padre='B0PADRE001', rank='2061', nvar='3'),
                        fila('B0HERMA002', 2, '30', padre='B0PADRE001', rank='2061', nvar='3'),
                        fila('B0OTROP001', 3, '30', padre='B0PADRE002', rank='100', nvar='4'),
                        fila('B0OTROP002', 4, '30', padre='B0PADRE002', rank='101', nvar='4'),
                        fila('B0SINPAD01', 5, '30', padre='', rank='777'),
                        fila('B0SINPAD02', 6, '30', padre='', rank='777'),
                        fila('B0SINPUE01', 7, '30', padre='B0PADRE003', rank=''),
                        fila('B0SINPUE02', 8, '30', padre='B0PADRE003', rank=''),
                        fila('B0SINREC01', 9, '30', padre='B0PADRE004', rank='900', nvar=''),
                        fila('B0SINREC02', 10, '30', padre='B0PADRE004', rank='900', nvar='0')], 6, ())
_es = _comp.get('ES', {})
eq('(B) mismo padre y mismo puesto → ficha compartida, los dos, con sus hermanos y el recuento de Amazon',
   [(_es[a]['hermanos'], _es[a]['divisor'], _es[a]['origen_divisor']) for a in ('B0HERMA001', 'B0HERMA002') if a in _es],
   [(['B0HERMA001', 'B0HERMA002'], 3.0, 'amazon')] * 2)
eq('(B) mismo padre y distinto puesto → sin reparto', [a for a in ('B0OTROP001', 'B0OTROP002') if a in _es], [])
eq('(B) padre vacío (aunque el puesto coincida) → sin reparto', [a for a in ('B0SINPAD01', 'B0SINPAD02') if a in _es], [])
eq('(B) puesto vacío (aunque el padre coincida) → sin reparto', [a for a in ('B0SINPUE01', 'B0SINPUE02') if a in _es], [])
eq('(B) recuento de variaciones vacío o 0 → ÷ los hermanos de la lista (2)',
   [(_es[a]['divisor'], _es[a]['origen_divisor'], _es[a]['nvar']) for a in ('B0SINREC01', 'B0SINREC02')],
   [(2, 'lista', None), (2, 'lista', None)])
eq('(B) solo esos cuatro son ficha compartida', sorted(_es), ['B0HERMA001', 'B0HERMA002', 'B0SINREC01', 'B0SINREC02'])
eq('(B) un ASIN solo en su padre no es compartido (hace falta OTRO producto de la lista)',
   e2.fichas_compartidas({'ES': {'B0SOLO0001': {'padre': 'B0P', 'rank': 5.0, 'nvar': 7.0}}}), {})
eq('(B) el padre y el puesto se cruzan POR PAÍS: el mismo par en ES y en DE no se junta',
   e2.fichas_compartidas({'ES': {'B0A': {'padre': 'B0P', 'rank': 5.0, 'nvar': 2.0}},
                          'DE': {'B0B': {'padre': 'B0P', 'rank': 5.0, 'nvar': 2.0}}}), {})
_sin_col = escribir('sin_padre.csv', [fila('B0X0000001', 1, '7')], [c for c in CABECERA if c != COL_PADRE])
_exs = e2.examinar_csv(_sin_col, pro, COL_PAIS, COL_CAIDAS, COLS_FICHA)
eq('(B) un CSV sin «ASIN Padre» se lee igual (las caídas están), sin fichas, y dice qué columna falta',
   (_exs['caidas'], _exs['fichas'], _exs['sin_columnas_ficha']), ({'B0X0000001': 7.0}, {}, ['ASIN Padre']))
eq('(B) sin pedir las columnas de la ficha (como antes del AA), examinar_csv devuelve lo de siempre',
   sorted(e2.examinar_csv(_sin_col, pro, COL_PAIS, COL_CAIDAS)), ['caidas', 'choques_caidas', 'filas', 'fuente_pais', 'pais'])

# ═══════════════════════════════════════════════════════════════════════════════
print('\n(C) el reparto en las puertas')
_FILAS_C = [fila('B0TRES0001', 1, '30', padre='B0PADRE030', rank='2061', nvar='3'),
            fila('B0TRES0002', 2, '30', padre='B0PADRE030', rank='2061', nvar='3'),
            fila('B0SIETE001', 3, '8', padre='B0PADRE008', rank='7622', nvar='7'),
            fila('B0SIETE002', 4, '8', padre='B0PADRE008', rank='7622', nvar='7'),
            fila('B0RECUE001', 5, '14', padre='B0PADRE014', rank='36542', nvar=''),
            fila('B0RECUE002', 6, '16', padre='B0PADRE014', rank='36542', nvar=''),
            fila('B0SOLA0001', 7, '30')]
_rc, _, _ = cruzar(_FILAS_C, 6, range(1, 8))
_t = _rc[1]
eq('(C) 30 ventas ÷ 3 variaciones = 10 por figura → entra (se vende)', (_t['puerta'] in ('d', 'e', 'f'), _t['compartida']['ES']['por_figura']),
   (True, 10.0))
eq('(C) …y lleva la marca al final del detalle',
   _t['detalle'].endswith(' · Ficha compartida en ES: 3 figuras · ventas de la ficha 30/mes · estimadas por figura 10/mes'), True)
eq('(C) …ES guarda las caídas de la FICHA (30), no las repartidas, y «vende aquí»',
   (_t['caidas']['ES'], _t['paises']['ES']['caidas_30d'], _t['paises']['ES']['vende_aqui']), (30.0, 30.0, True))
_s = _rc[3]
eq('(C) 8 ventas ÷ 7 variaciones → no pasa: puerta c, motivo de la base c_pocas_caidas',
   (_s['puerta'], _s['motivo']), ('c', 'c_pocas_caidas'))
eq('(C) …con el motivo propio delante del detalle, visible en «Puertas»', _s['detalle'],
   'c_ficha_compartida: ES 7 figuras, ventas de la ficha 8/mes → 1,1/mes por figura · por debajo del corte '
   '(7 o más caídas en 30 días) (IT, FR, DE: sin dato)')
eq('(C) …y sin el reparto (ficha sola) esas 8 caídas SÍ entrarían: lo que la deja fuera es la ficha compartida',
   cruzar([fila('B0SIETE001', 3, '8')], 6, (3,))[0][3]['puerta'] in ('d', 'e', 'f'), True)
eq('(C) sin recuento de Amazon: 14 y 16 ÷ 2 hermanos = 7 y 8 → entran los dos, y el detalle dice cómo se repartió',
   [(_rc[n]['puerta'] in ('d', 'e', 'f'), 'repartido entre los 2 hermanos de la lista' in _rc[n]['detalle']) for n in (5, 6)],
   [(True, True), (True, True)])
eq('(C) un producto sin hermanos: ni reparto ni marca', (_rc[7]['compartida'], 'Ficha compartida' in _rc[7]['detalle']),
   ({}, False))
_cq = e2.cuadre(['F%d' % n for n in range(1, 8)], [dict(_rc[n], foto_id='F%d' % n) for n in range(1, 8)])
eq('(C) el cuadre: la ficha compartida que no pasa cuenta como «pocas caídas» (c = pocas + sin dato en la base)',
   (_cq['cuadra'], _cq['conteo']['c'], _cq['n_c_pocas']), (True, 2, 2))

# ═══════════════════════════════════════════════════════════════════════════════
print('\n(D) el Excel: «Ficha compartida» al final de «Análisis»')
from openpyxl import load_workbook  # noqa: E402

_foto = [foto(n) for n in range(1, 8)]
_res = [dict(_rc[n], foto_id='F%d' % n) for n in range(1, 8)]
_buf = io.BytesIO()
e2.excel_como_el_viejo(_foto, _res, [], M).save(_buf)
_ws = load_workbook(io.BytesIO(_buf.getvalue()))['Análisis']
_cab = [c.value for c in _ws[1]]
eq('(D) es la ÚLTIMA columna, detrás de las del viejo (que no se mueven)',
   _cab, e2.columnas_analisis() + ['Ficha compartida'])
eq('(D) la tabla T_Analisis la cubre', [(t.displayName, t.ref.split(':')[1][:2]) for t in _ws.tables.values()], [('T_Analisis', 'AB')])
_i = {h: k for k, h in enumerate(_cab)}
_marcas = {(r[_i['EAN']], r[_i['País']]): r[_i['Ficha compartida']] for r in _ws.iter_rows(min_row=2, values_only=True)}
eq('(D) la fila de ES del EAN compartido lleva la marca de ES',
   _marcas[(ean(1), 'ES')], 'Ficha compartida: 3 figuras · ventas de la ficha 30/mes · estimadas por figura 10/mes')
eq('(D) las otras filas del mismo EAN dicen en qué país es compartida',
   _marcas[(ean(1), 'DE')], 'Ficha compartida en ES: 3 figuras · ventas de la ficha 30/mes · estimadas por figura 10/mes')
eq('(D) la del recuento que faltaba lo dice', 'repartido entre los 2 hermanos de la lista' in _marcas[(ean(5), 'ES')], True)
eq('(D) un EAN sin reparto: la celda, vacía', [_marcas[(ean(7), p)] for p in ('ES', 'IT', 'FR', 'DE')], [None] * 4)
eq('(D) la ficha compartida que no pasa no está en «Análisis» (va a la c)', (ean(3), 'ES') in _marcas, False)

shutil.rmtree(_tmp, ignore_errors=True)
print()
if fallos:
    print('ROJO: %d comprobaciones fallan: %s' % (len(fallos), ', '.join(fallos)))
    sys.exit(1)
print('VERDE: «se vende» = 7 o más caídas (del parámetro), y la ficha compartida se reparte, se marca y se ve.')
