# -*- coding: utf-8 -*-
"""ESCANER 2 · ESCANEO PRO DE OCIOSTOCK · LA PARTE PURA (encargo OC3, 07-oct-2026; pieza 3 del plano de OC1)

QUE ES. El calculo del Escaneo PRO de OcioStock, sin red y sin base: la foto de la pasada desde la FOTO DE
DISPONIBILIDAD (disp_estado, que deja escaner2_ociostock_disponibilidad.py, pieza 1), la lista para el Visualizador
y el Excel. La red vive en los dos programas que lo usan:
  · escaner2_ociostock_barrido.py (workflow escaner2-ociostock-barrido.yml): de la foto a la pasada y la lista.
  · escaner2_ociostock_cruce.py   (workflow escaner2-ociostock-cruce.yml):   los 4 CSV, las puertas y el Excel.

🔑 CADA ESCANER, A LA MEDIDA DE SU PROVEEDOR (Fernando, 02-oct-2026): leer y filtrar es de OcioStock y vive aqui. Lo
   que se REUTILIZA sin tocarlo son las REGLAS DE COMPRA del PRO de HEO (`escaner2_motor.decidir`, con el corte y los
   paises de escaner2_parametros: OCIOSTOCK = 6 y ES·IT·FR·DE, como HEO) y su Excel (la Celda 9 del viejo con
   «Ventas», la columna «Ficha compartida» y las hojas del escaner 2 en el MISMO orden).

LO PROPIO DE OCIOSTOCK (decidido por Fernando, 7-oct-2026):
  1. LA FOTO SALE DE disp_estado (las filas vistas en la ultima pasada APLICADA). Nada de la web.
  2. SOLO LO QUE SE PUEDE COMPRAR YA: disponible y sin preventa («Preventas evidentemente no entran»).
  3. LOS CHASE SUELTOS, FUERA («Los chase sueltos no los voy a comprar nunca»): `regla = 'chase_suelto'` de la foto,
     y cualquier chase que no sea caja. 🔴 Las cajas «5 + 1» llevan es_chase = true Y es_caja = true (como las de
     HEO): esas SI entran, valoradas con el EAN de la figura comun (el mismo EAN base).
  4. EL PRECIO, EL MAS BAJO («Calcula la rentabilidad al menor precio, me da igual comprar volumen»): PA =
     `precio_pa` de la foto (el escalon mas barato × 0,99 de la transferencia), POR FIGURA tambien en la caja (el
     precio de OcioStock de una caja ya es por figura: no se divide). El cruce no recalcula nada: `decidir` usa
     `precio_unidad` de la foto de la pasada, y ahi va el PA.
  5. UNA FILA POR EAN BASE: la de menor PA (a igualdad, la suelta, y luego el id menor); si gana una caja, se marca
     «caja de 6» (en «Coherencia caja» del Excel y en la columna «Caja» de «Puertas»). La otra, a la puerta previa
     `duplicado_proveedor`, listada.
  6. EL EXCEL, EL DEL PRO DE HEO, mismo formato y orden, con DOS columnas mas AL FINAL de «Análisis»: «Uds. escalón»
     (cuantas hay que pedir para ese precio) y «Precio unidad» (el de la ficha de OcioStock, sin escalon ni 0,99).
     Sin «Chase_manual» (es de HEO: el viejo de OcioStock tampoco la escribe) ni comparacion con el viejo (la
     compara la pieza 5, aparte).

🔴 REPO PUBLICO: aqui no hay ni un precio ni un EAN reales; los del banco son inventados (Seguridad 21).
"""
import io
import re

import escaner2_motor as e2

PROVEEDOR = 'OCIOSTOCK'
CARPETA = 'ociostock'
NOMBRE = 'OcioStock'
MODOS = ('marcas', 'todas', 'elegidas')
UDS_CAJA = 6
TEXTO_CAJA = 'caja de 6'

# Las columnas nuevas de «Análisis», AL FINAL (detras de «Ficha compartida»): hay quien lee la hoja por letra.
COLUMNA_UDS_ESCALON = 'Uds. escalón'
COLUMNA_PRECIO_UNIDAD = 'Precio unidad'
COLUMNAS_ESCALON = [COLUMNA_UDS_ESCALON, COLUMNA_PRECIO_UNIDAD]
ANCHOS_ESCALON = [12, 13]

# Los nombres de las puertas previas EN OCIOSTOCK (los mismos que NOMBRE_PREVIA_OCIOSTOCK de la v2). Las ocho de
# siempre (crudo = previas + foto); OcioStock usa seis y `chase_funko` y `estado_no_servible` salen a 0.
NOMBRE_PREVIA = dict(e2.NOMBRE_PUERTA_PREVIA,
                     sin_gtin='Sin EAN',
                     no_disponible='Sin stock o en preventa',
                     marca_fuera='Marca fuera de la lista del director',
                     chase_suelto='Chase suelto (no se compra nunca)',
                     ean_forma_rara='EAN con forma rara',
                     duplicado_proveedor='Mismo EAN que otra fila de OcioStock, más cara')


class FalloOcioStock(Exception):
    """Algo que impide seguir: la pasada o el cruce quedan 'fallida' con este motivo (privado, en la base)."""


def nombre_previa(p, modo):
    """El nombre de una puerta previa EN SU PASADA: en el modo «elegidas» la marca que se cae es «no elegida»."""
    if p == 'marca_fuera' and modo == 'elegidas':
        return e2.NOMBRE_MARCA_NO_ELEGIDA
    return NOMBRE_PREVIA[p]


# ═══════════════════════════════════════════════════════════════════════════════
# 1 · EL FILTRO DE MARCA, POR MODO
# ═══════════════════════════════════════════════════════════════════════════════
def filtro_marca(modo, marcas_director=None, elegidas=None):
    """(quiere, info) para el modo:
      · 'marcas'   · las de `reglas_director` (fila OCIOSTOCK, solo lectura) con la regla del viejo: el nombre de la
                     marca de la regla CONTENIDO en el de la fila, sin distinguir mayusculas (`m.lower() in marca`,
                     `_pasa_filtros` de moloka_escaner_nube.py);
      · 'todas'    · todas, sin leer reglas_director (las sin marca, tambien);
      · 'elegidas' · las del selector de la v2, COINCIDENCIA EXACTA (`escaner2_motor.clave_marca`), como HEO.
    `quiere(marca)` → bool. `info` lleva las marcas que se guardan en la pasada (None en 'todas')."""
    if modo == 'todas':
        return (lambda _m: True), {'marcas_reales': None}
    if modo == 'marcas':
        lista = [str(m).strip() for m in (marcas_director or []) if str(m or '').strip()]
        if not lista:
            raise FalloOcioStock('la fila OCIOSTOCK de reglas_director no trae marcas: no se sabe qué barrer')
        bajas = [m.lower() for m in lista]
        return (lambda m: any(b in str(m or '').lower() for b in bajas)), {'marcas_reales': lista}
    if modo == 'elegidas':
        claves = frozenset(e2.clave_marca(m) for m in (elegidas or []))
        if not claves:
            raise FalloOcioStock('modo elegidas sin ninguna marca')
        return (lambda m: e2.clave_marca(m) in claves), {'marcas_reales': list(elegidas)}
    raise FalloOcioStock('modo de barrido desconocido: %r' % (modo,))


# ═══════════════════════════════════════════════════════════════════════════════
# 2 · LA FOTO DE LA PASADA, DESDE LA FOTO DE DISPONIBILIDAD
# ═══════════════════════════════════════════════════════════════════════════════
def _clave_id(v):
    s = str(v or '')
    return (0, int(s), s) if s.isdigit() else (1, 0, s)


def _num(x):
    return None if x is None else float(x)


def _apartado(f, motivo, detalle):
    return {'producto_heo': str(f.get('producto_prov') or ''), 'ean_original': str(f.get('ean_original') or ''),
            'nombre': f.get('nombre') or '', 'marca': (f.get('marca') or '').strip(),
            'precio_catalogo': _num(f.get('precio_unidad')), 'motivo': motivo, 'detalle': detalle, 'en_oferta': None}


def _gana(a, b):
    """¿`a` va antes que `b` para quedarse con el EAN? Menor PA (sin PA, la ultima); a igualdad, la suelta antes que
    la caja (se pide de una en una), y despues el id de OcioStock menor (para que salga siempre la misma)."""
    def clave(f):
        pa = f.get('precio_pa')
        return (pa is None, float(pa) if pa is not None else 0.0, bool(f.get('es_caja')), _clave_id(f.get('producto_prov')))
    return clave(a) < clave(b)


def construir_foto(filas, quiere_marca, M, n_crudo, n_sin_gtin, modo='marcas'):
    """De la foto de disponibilidad de OcioStock a la foto de la pasada, sin perder a nadie: cada fila sale por UNA
    puerta previa o entra en la foto, y crudo (el fichero) = previas + foto. En este orden:
      1. no se puede comprar ya (no disponible, o preventa) → `no_disponible` (se cuenta);
      2. marca que no pasa el filtro del modo → `marca_fuera` (se lista);
      3. chase suelto (`regla` de la foto, o chase que no es caja) → `chase_suelto` (se lista);
      4. EAN base que no son 12 o 13 cifras, u otra `regla` de la foto → `ean_forma_rara` (se lista);
      5. dos filas con el mismo EAN base → la de menor PA (`_gana`); la otra → `duplicado_proveedor` (se lista).
    `filas`: disp_estado de OCIOSTOCK vistas en la ultima pasada aplicada; `n_crudo` y `n_sin_gtin`: los de esa
    pasada (las filas sin EAN no estan en disp_estado: solo se cuentan). Devuelve (foto, apartados, cuentas, escalon),
    con `escalon` = {producto_prov: {uds_escalon, precio_unidad, precio_escalon, precio_pa}} de las filas de la foto."""
    previas = {p: 0 for p in e2.PUERTAS_PREVIAS}
    previas['sin_gtin'] = n_sin_gtin
    apartados, sirven = [], []
    no_elegida = modo == 'elegidas'
    for f in sorted(filas or [], key=lambda x: _clave_id(x.get('producto_prov'))):
        if not f.get('disponible') or f.get('preorder'):
            previas['no_disponible'] += 1
            continue
        if not quiere_marca(f.get('marca')):
            apartados.append(_apartado(f, 'marca_fuera', ('Marca %r no elegida' if no_elegida
                                                          else 'Marca %r fuera de la lista del director')
                                       % ((f.get('marca') or '').strip(),)))
            continue
        regla = f.get('regla')
        if regla == 'chase_suelto' or (f.get('es_chase') and not f.get('es_caja')):
            apartados.append(_apartado(f, 'chase_suelto', 'Chase suelto: no se compra nunca (solo la caja de 6)'))
            continue
        core = str(f.get('ean_core') or '').strip()
        if regla or not core.isdigit() or len(core) not in (12, 13):
            apartados.append(_apartado(f, 'ean_forma_rara', 'EAN con forma rara (%s%s)'
                                       % (f.get('ean_original') or '', ', regla %s' % regla if regla else '')))
            continue
        sirven.append(f)

    uno = {}
    for f in sirven:
        k = M.norm(f['ean_core'])
        prev = uno.get(k)
        if prev is None:
            uno[k] = f
            continue
        gana, pierde = (f, prev) if _gana(f, prev) else (prev, f)
        uno[k] = gana
        apartados.append(_apartado(pierde, 'duplicado_proveedor', 'Mismo EAN que el %s de OcioStock%s: me quedo con '
                                   'el más barato (%s frente a %s)' % (
                                       gana.get('producto_prov'), ' (caja de 6)' if gana.get('es_caja') else '',
                                       gana.get('precio_pa'), pierde.get('precio_pa'))))

    foto, escalon = [], {}
    for f in sorted(uno.values(), key=lambda x: _clave_id(x.get('producto_prov'))):
        core = str(f['ean_core']).strip()
        caja = bool(f.get('es_caja'))
        variantes = sorted({M.norm(v) for v in M.variantes_ean(core)} - {''})
        foto.append({
            'producto_heo': str(f.get('producto_prov') or ''), 'ean_original': str(f.get('ean_original') or core),
            'ean_core': core, 'variantes': variantes, 'codigos_keepa': e2.codigos_para_keepa(core, M),
            'nombre': f.get('nombre') or '', 'marca': (f.get('marca') or '').strip(),
            'categoria': f.get('categoria') or '',
            # 🔑 precio_unidad = el PA (escalon × 0,99, por figura): con el decide `escaner2_motor.decidir`.
            #    precio_catalogo = el precio de la ficha de OcioStock, sin escalon ni 0,99 (columna «Precio unidad»).
            'precio_catalogo': _num(f.get('precio_unidad')), 'precio_unidad': _num(f.get('precio_pa')),
            'es_caja': caja, 'uds_caja': (int(f.get('uds_caja') or UDS_CAJA) if caja else None),
            'es_chase': bool(f.get('es_chase')), 'en_oferta': None, 'campana': None, 'disponibilidad': None,
            'fin_de_vida': None, 'preorder': False, 'imagen': None,
            'aviso_caja': TEXTO_CAJA if caja else None, 'origen_ean': None, 'aviso_ean': None,
        })
        escalon[str(f.get('producto_prov') or '')] = {
            'uds_escalon': f.get('uds_escalon'), 'precio_unidad': _num(f.get('precio_unidad')),
            'precio_escalon': _num(f.get('precio_escalon')), 'precio_pa': _num(f.get('precio_pa'))}

    for m in e2.MOTIVOS_APARTADO:
        previas[m] = sum(1 for a in apartados if a['motivo'] == m)
    cuentas = {'n_crudo': n_crudo, 'n_filas': len(filas or []), 'n_foto': len(foto), 'previas': previas}
    cuentas.update(cuadre_previo(cuentas))
    return foto, apartados, cuentas, escalon


def cuadre_previo(cuentas):
    """🔴 crudo = previas + foto, y el crudo es el del fichero: filas de la foto de disponibilidad + las sin EAN.
    Un recuento que falta (None) NO es un cero."""
    previas = cuentas['previas']
    faltan = [p for p in e2.PUERTAS_PREVIAS if previas.get(p) is None]
    if cuentas.get('n_crudo') is None:
        faltan.insert(0, 'crudo')
    if faltan:
        return {'n_previas': None, 'cuadra_previo': False, 'motivo_previo': 'sin recuento de: ' + ', '.join(faltan)}
    n_previas = sum(previas[p] for p in e2.PUERTAS_PREVIAS)
    if cuentas['n_filas'] + previas['sin_gtin'] != cuentas['n_crudo']:
        return {'n_previas': n_previas, 'cuadra_previo': False,
                'motivo_previo': 'el fichero tenía %d filas y la foto de disponibilidad trae %d + %d sin EAN'
                                 % (cuentas['n_crudo'], cuentas['n_filas'], previas['sin_gtin'])}
    if cuentas['n_crudo'] != n_previas + cuentas['n_foto']:
        return {'n_previas': n_previas, 'cuadra_previo': False,
                'motivo_previo': 'catálogo crudo %d ≠ puertas previas %d + foto %d'
                                 % (cuentas['n_crudo'], n_previas, cuentas['n_foto'])}
    return {'n_previas': n_previas, 'cuadra_previo': True, 'motivo_previo': None}


def corpus_cotejo(foto, apartados):
    """Los nombres con que la regla del viejo mide sus palabras distintivas al elegir entre varias fichas: los de la
    foto y los de lo disponible apartado por MARCA (como el B5-bis de HEO: que la eleccion no dependa de las marcas
    elegidas). `escaner2_motor.corpus_cotejo` lee el perfil y los EAN de HEO; este no mira EAN."""
    nombres = [f.get('nombre') or '' for f in foto]
    nombres += [a.get('nombre') or '' for a in apartados if a.get('motivo') == 'marca_fuera']
    return nombres


def ruta_lista(pasada):
    return '%s/%s/eans.txt' % (CARPETA, pasada)


def ruta_excel(pasada, cruce, sello):
    """`ociostock/<pasada>/<cruce>/Escaner2_OCIOSTOCK_<AAAAMMDD_HHMM>.xlsx`: la forma que admite
    escaner2_cruce.ruta_excel desde la migracion de la v2 de este encargo."""
    return '%s/%s/%s/Escaner2_OCIOSTOCK_%s.xlsx' % (CARPETA, pasada, cruce, sello)


def nombre_modo(modo):
    """Como se llama el modo de la pasada en el Excel (el mismo texto que el cruce de HEO)."""
    return {'todas': 'todas las marcas', 'elegidas': 'marcas elegidas'}.get(modo or 'marcas', 'marcas de siempre')


# ═══════════════════════════════════════════════════════════════════════════════
# 3 · EL EXCEL: EL DEL PRO DE HEO, CON «Uds. escalón» Y «Precio unidad» AL FINAL DE «Análisis»
# ═══════════════════════════════════════════════════════════════════════════════
COLUMNAS_CAJA = ['Caja', 'Precio caja (€)', 'EAN de la figura', 'Origen del EAN', 'Aviso del EAN']


def _columnas_caja(f):
    """Las cinco columnas de caja de «Puertas» del PRO de HEO. En OcioStock la caja se compra a precio POR FIGURA:
    el precio de la caja son las seis figuras a ese precio (el de la ficha, sin escalon ni 0,99)."""
    if not f.get('es_caja'):
        return [None, None, None, None, None]
    pc = f.get('precio_catalogo')
    uds = f.get('uds_caja') or UDS_CAJA
    return [e2.rotulo_caja(f) or None, None if pc is None else round(pc * uds, 2), None, None, None]


def poner_columnas_escalon(wb, foto, escalon):
    """«Uds. escalón» y «Precio unidad» detras de la ultima columna de «Análisis», fila a fila por su EAN, y la tabla
    alargada (como `poner_columna_ficha_compartida`). Devuelve cuantas filas llevan dato."""
    from openpyxl.styles import Font
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.table import TableColumn
    ws = wb['Análisis']
    cab = [c.value for c in ws[1]]
    i_ean = cab.index('EAN')
    primera = len(cab) + 1
    for j, (nombre, ancho) in enumerate(zip(COLUMNAS_ESCALON, ANCHOS_ESCALON)):
        ws.cell(row=1, column=primera + j, value=nombre).font = Font(bold=True)
        ws.column_dimensions[get_column_letter(primera + j)].width = ancho
    por_ean = {}
    for f in foto:
        por_ean[str(f['ean_original'])] = (escalon or {}).get(str(f.get('producto_heo') or '')) or {}
    n = 0
    for fila in range(2, ws.max_row + 1):
        dato = por_ean.get(str(ws.cell(row=fila, column=i_ean + 1).value))
        if not dato:
            continue
        ws.cell(row=fila, column=primera, value=dato.get('uds_escalon'))
        c = ws.cell(row=fila, column=primera + 1, value=dato.get('precio_unidad'))
        c.number_format = '0.00'
        n += 1
    ultima = primera + len(COLUMNAS_ESCALON) - 1
    for tabla in ws.tables.values():
        ini, fin = tabla.ref.split(':')
        fila_fin = re.sub(r'^[A-Z]+', '', fin)
        tabla.ref = '%s:%s%s' % (ini, get_column_letter(ultima), fila_fin)
        if tabla.tableColumns:
            for nombre in COLUMNAS_ESCALON:
                tabla.tableColumns.append(TableColumn(id=len(tabla.tableColumns) + 1, name=nombre))
        if tabla.autoFilter is not None:
            tabla.autoFilter.ref = tabla.ref
    return n


def escribir_excel(foto, resultados, apartados, M, info):
    """El Excel del cruce de OcioStock: el del PRO de HEO (`escaner2_heo_cruce.escribir_excel`), mismo formato y orden.
    DELANTE las hojas del viejo (la Celda 9: Análisis con «Ventas», Descartados, Ambiguos, Sin_rank, Precio por lote),
    con «Ficha compartida» y las DOS columnas de OcioStock al final de «Análisis»; DETRAS Resumen, Comparación (vacía:
    la comparacion con el viejo es la pieza 5), Varias fichas, Puertas y Puertas previas. Solo lo calculado: aqui no
    se decide nada. → bytes."""
    from openpyxl.styles import Font
    por_foto = {f['id']: f for f in foto}
    datos = e2.datos_como_el_viejo(foto, resultados, apartados, M, eleccion=info.get('eleccion'))
    # El viejo de OcioStock tampoco escribe «Chase_manual» (solo con PROVEEDOR == 'HEO').
    datos['PROVEEDOR'] = PROVEEDOR
    wb = e2.escribir_celda9(datos, M)
    e2.poner_columna_ficha_compartida(wb, foto, resultados)
    poner_columnas_escalon(wb, foto, info.get('escalon'))

    def hoja(nombre, cabecera, filas, anchos=None):
        ws = wb.create_sheet(nombre)
        ws.append(cabecera)
        for c in ws[1]:
            c.font = Font(bold=True)
        for fila in filas:
            ws.append(fila)
        for letra, ancho in (anchos or {}).items():
            ws.column_dimensions[letra].width = ancho
        ws.freeze_panes = 'A2'
        return ws

    p = info['params']
    modo = info.get('modo')
    filas_res = [['Pasada', info['pasada']], ['Cruce', info['cruce']], ['Modo del barrido', nombre_modo(modo)]]
    if modo == 'elegidas':
        filas_res += [['Marcas elegidas', ', '.join(info.get('marcas') or []) or '(ninguna)']]
    elif modo == 'marcas':
        filas_res += [['Marcas del director', ', '.join(info.get('marcas') or []) or '(ninguna)']]
    filas_res += [['Foto de OcioStock (pasada de disponibilidad)', info.get('foto') or '—'],
                  ['Precio de compra (PA)', 'el más barato: el escalón más bajo × 0,99 (transferencia), por figura; '
                                            '«Uds. escalón» dice cuántas pedir para ese precio'],
                  ['Umbral de caídas (30 días)', '%s (> %d)' % (e2.texto_corte(p['umbral']), p['umbral'])],
                  ['Fichas compartidas (productos de la lista, por país)',
                   ' · '.join('%s %d' % (k, n) for k, n in (info.get('n_compartidas') or {}).items()) or '—'],
                  ['Países del filtro de ventas', ', '.join(p['paises_filtro'])],
                  ['Países que se calculan (si traen CSV)', ', '.join(p['paises_calculo'])],
                  ['Países con CSV', ', '.join(info['usados'])],
                  ['Catálogo crudo de OcioStock (el fichero)', info['n_crudo']]]
    filas_res += [['Puerta previa · %s' % nombre_previa(x, modo), info['previas'][x]] for x in e2.PUERTAS_PREVIAS]
    filas_res += [['Entradas (filas de la foto: un EAN, su fila más barata)', info['n_entradas']]]
    filas_res += [['Puerta %s · %s' % (x, e2.NOMBRE_PUERTA[x]), info['n_bd'][x]] for x in e2.PUERTAS]
    filas_res += [['Puertas previas + suma de puertas', sum(info['previas'].values()) + sum(info['n_bd'].values())],
                  ['Cuadra', 'SÍ' if info['cuadra'] else 'NO'],
                  ['Comparación con el escáner viejo', 'no se hace aquí: es la pieza 5 (aparte)']]
    filas_res += [['CSV', '%s · %s · %s filas%s%s' % (
        f['nombre'], f.get('pais') or '¿?', f.get('filas') or 0,
        '' if f.get('usado') else ' · NO USADO', (' · ' + f['error']) if f.get('error') else '')]
        for f in info['ficheros']]
    filas_res += [['Aviso', a] for a in info['avisos']]
    hoja('Resumen', ['Qué', 'Valor'], filas_res, {'A': 44, 'B': 90})

    hoja('Comparación', ['EAN', 'Nombre', 'Categoría', 'Viejo', 'País viejo', 'Fecha viejo (UTC)', 'Excel viejo',
                         'Puerta nuevo', 'Nuevo', 'País nuevo', 'Fecha nuevo (UTC)', 'Diferencia de criterio',
                         'Explicación'], [], {'A': 15, 'B': 50, 'C': 16, 'G': 36, 'M': 110})

    varias = []
    for res in resultados:
        if res['puerta'] == 'b':
            f = por_foto[res['foto_id']]
            for fi in res['fichas']:
                varias.append([f['ean_original'], f['nombre'], fi['pais'], fi['asin'], fi['titulo'], fi['rank'],
                               fi['caidas_30d'], fi['precio_venta']])
    hoja('Varias fichas', ['EAN', 'Nombre OcioStock', 'País', 'ASIN', 'Título Amazon', 'Puesto', 'Caídas 30 d',
                           'Precio venta (€)'], varias, {'A': 15, 'B': 50, 'E': 60})

    hoja('Puertas', ['EAN', 'Nombre', 'Marca', 'Precio compra (€)', 'Puerta', 'Motivo', 'Detalle'] + COLUMNAS_CAJA,
         [[por_foto[r['foto_id']]['ean_original'], por_foto[r['foto_id']]['nombre'], por_foto[r['foto_id']]['marca'],
           por_foto[r['foto_id']]['precio_unidad'], '%s · %s' % (r['puerta'], e2.NOMBRE_PUERTA[r['puerta']]),
           r['motivo'], r['detalle']] + _columnas_caja(por_foto[r['foto_id']]) for r in resultados],
         {'A': 15, 'B': 55, 'C': 16, 'E': 24, 'F': 16, 'G': 70, 'H': 26})

    hoja('Puertas previas', ['EAN tal como vino', 'Nombre', 'Marca', 'Precio catálogo (€)', 'Puerta previa', 'Detalle'],
         [[a['ean_original'], a['nombre'], a['marca'], a['precio_catalogo'],
           nombre_previa(a['motivo'], modo) if a['motivo'] in NOMBRE_PREVIA else a['motivo'], a['detalle']]
          for a in apartados],
         {'A': 18, 'B': 55, 'C': 16, 'E': 32, 'F': 80})
    salida = io.BytesIO()
    wb.save(salida)
    return salida.getvalue()
