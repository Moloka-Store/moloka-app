# -*- coding: utf-8 -*-
"""ESCANER 2 · EL MOTOR PURO DEL BARRIDO DE HEO, EN SOMBRA (encargo B, trozo 1, 24-sep-2026)

QUE ES. La parte de CALCULO del escaner nuevo de HEO, sin red y sin base: construir la
foto del catalogo, leer los CSV del Visualizador de Keepa, decidir por que PUERTA sale
cada EAN, comprobar el CUADRE y compararse con el escaner viejo. La red (API de HEO,
Supabase, Storage) vive en los dos programas que lo usan:
  · escaner2_heo_barrido.py  (workflow escaner2-heo-barrido.yml): baja HEO y deja la foto.
  · escaner2_heo_cruce.py    (workflow escaner2-heo-cruce.yml):   cruza los CSV y decide.

🔴 NO SE COPIA NI UNA LINEA DEL ESCANER VIEJO: SE LEE DE SU FICHERO, POR ESTRUCTURA.
   `moloka_escaner_nube.py` no se puede importar (crea el cliente de Keepa y sale a la red
   al cargar) y el encargo prohibe modificarlo. Asi que `calc_rentabilidad`, `decision_de`,
   las reglas de EAN/chase/caja (`partir_ean`, `clasificar_chase`, `variantes_ean` con el
   rescate de GTIN), el IVA con su origen, `ISD_PAIS`, `UNIDADES_CASE_TCG` y el perfil HEO
   se SACAN del fichero real con `ast` -- buscando cada `def` y cada asignacion de NIVEL
   SUPERIOR por su nombre, no con un grep -- y se EJECUTAN. Es el mismo metodo con el que
   `test_escaner_memoria.py` prueba las guardas del escaner. Consecuencias:
     · la formula que usa el escaner nuevo ES la del viejo, letra por letra, hoy y manana:
       si alguien la cambia alli, el nuevo la hereda sin tocar este fichero;
     · si alguien la BORRA o la RENOMBRA, esto no adivina: revienta con el nombre que falta.
   Lo mismo con el filtro del director (`_quiere` y sus marcas, de director_heo_prep.py) y
   con la tanda del Visualizador (descargar_heo.py). Las columnas del CSV salen del Escaner
   Pro (`CSV_COLS`, `leer_csv_visualizador`) y de `procesador_keepa_escaparate.py`
   (`TIPADAS`, medidas contra los exports reales del 17 y 18-ago-2026).

🔒 LO QUE NO SE PUEDE SACAR POR NOMBRE, porque en el viejo va escrito EN LINEA dentro de
   un bucle y no en una funcion, se replica aqui con el sitio de donde sale al lado:
   el dedup del proveedor y el guardarrail caja-vs-suelta (Celda 4 y siguientes), la
   division del PA por la caja (Celda 8), el diccionario de IVA por pais (Celda 8) y la
   eleccion del precio de venta del CSV (escanear_pro del Escaner Pro). El banco
   (test_escaner2_motor.py) coteja cada una contra su original ejecutandolo.
"""
import ast
import csv
import io
import math
import os
import re
from datetime import datetime, timezone

RUTA_MOTOR = 'moloka_escaner_nube.py'
RUTA_DIRECTOR_HEO = 'director_heo_prep.py'
RUTA_DESCARGAR_HEO = 'descargar_heo.py'
RUTA_ESCAPARATE = 'procesador_keepa_escaparate.py'

PROVEEDOR = 'HEO'

# Las seis puertas del encargo. Cada EAN de la foto sale por UNA y solo una.
PUERTAS = ('a', 'b', 'c', 'd', 'e', 'f')
NOMBRE_PUERTA = {
    'a': 'No está en Amazon',
    'b': 'Dos o más fichas',
    'c': 'No se vende',
    'd': 'Se vende sin margen',
    'e': 'VALORAR',
    'f': 'COMPRAR',
}
# Los motivos, por puerta. La c tiene los DOS del encargo; el resto afina el porque.
MOTIVOS = {
    'a_no_aparece': 'a', 'a_sin_asin': 'a',
    'b_varias_fichas': 'b',
    'c_pocas_caidas': 'c', 'c_sin_dato': 'c',
    'd_sin_margen': 'd', 'd_sin_datos': 'd',
    'e_valorar': 'e',
    'f_comprar': 'f',
}
# Lo que se aparta ANTES de la foto y se guarda EAN a EAN, con su motivo (escaner2_apartado).
# `marca_fuera` se lista desde el encargo B2 (25-sep-2026): antes solo se contaba.
MOTIVOS_APARTADO = ('chase_funko', 'marca_fuera', 'estado_no_servible', 'chase_suelto', 'ean_forma_rara',
                    'duplicado_proveedor')
# 🔑 LOS MODOS DE BARRIDO, elegidos en la pantalla en cada pasada:
#    · 'marcas'   · el filtro del director, con las marcas de `reglas_director` (como siempre);
#    · 'todas'    · (B2) solo «disponible»: sin filtro de marca y SIN LEER `reglas_director`;
#    · 'elegidas' · (B4) las marcas que Fernando marca en el selector de la v2 y, si quiere, las
#                   ofertas de cualquier marca. Tampoco lee `reglas_director`.
MODOS = ('marcas', 'todas', 'elegidas')
# El tope del Visualizador de Keepa: 10.000 códigos por lista. La tanda del motor no lo pasa.
TOPE_VISUALIZADOR = 10000
# 🔴 LAS PUERTAS PREVIAS (Fernando, 24-sep-2026): el cuadre empieza en el catalogo CRUDO de HEO,
#    no en la foto. Todo producto que devuelve HEO sale por UNA de estas o entra en la foto:
#        crudo = puertas previas + foto        y, al cruzar,        crudo = previas + a..f
#    En el orden en que se aplican. Ninguna regla de descarte cambia: solo se cuentan y se ven.
#    `sin_gtin` y `no_disponible` solo se CUENTAN (sin GTIN no hay EAN que listar; no disponible
#    es el catalogo entero de HEO menos lo que se puede pedir); las otras seis (MOTIVOS_APARTADO)
#    ademas se LISTAN en escaner2_apartado.
#    🔑 Desde el B2, `chase_funko` es SOLO la caja con chase de la que no se puede sacar el EAN de
#       la figura comun: las demas entran en la foto (disponibles), cuentan como no disponibles
#       (agotadas) o, si el nombre no dice unidades, siguen el flujo normal como figura suelta.
PUERTAS_PREVIAS = ('chase_funko', 'sin_gtin', 'no_disponible', 'marca_fuera', 'estado_no_servible',
                   'chase_suelto', 'ean_forma_rara', 'duplicado_proveedor')
NOMBRE_PUERTA_PREVIA = {
    'chase_funko': 'Caja con chase sin EAN de la figura',
    'sin_gtin': 'Sin GTIN',
    'no_disponible': 'No disponible',
    'marca_fuera': 'Marca fuera de la lista',
    'estado_no_servible': 'Estado no servible',
    'chase_suelto': 'Chase suelto',
    'ean_forma_rara': 'Código de barras con forma rara',
    'duplicado_proveedor': 'Duplicado del proveedor',
}
# (B4) En el modo 'elegidas' la marca que se cae no está «fuera de la lista del director»: es una
#      marca que Fernando no ha elegido. La puerta es la MISMA (`marca_fuera`), cambia el rótulo.
NOMBRE_MARCA_NO_ELEGIDA = 'Marca no elegida'


def nombre_previa(p, modo):
    """El nombre de una puerta previa EN SU PASADA (el mismo que pinta la v2 con `nombrePrevia`)."""
    if p == 'marca_fuera' and modo == 'elegidas':
        return NOMBRE_MARCA_NO_ELEGIDA
    return NOMBRE_PUERTA_PREVIA[p]


class PiezaNoEncontrada(Exception):
    """Una pieza que el escaner nuevo lee del viejo ya no esta (o esta dos veces)."""


class CsvIlegible(Exception):
    """Un CSV del Visualizador que no se puede usar: se dice cual y por que, nunca se traga."""


# ═══════════════════════════════════════════════════════════════════════════════
# 1 · SACAR PIEZAS DE UN FICHERO, POR ESTRUCTURA
# ═══════════════════════════════════════════════════════════════════════════════
def _nombres_asignados(nodo):
    """Los nombres a los que asigna un `ast.Assign` de nivel superior (`A = …` o `A, B = …`)."""
    salida = []
    for t in nodo.targets:
        if isinstance(t, ast.Name):
            salida.append(t.id)
        elif isinstance(t, (ast.Tuple, ast.List)):
            salida.extend(e.id for e in t.elts if isinstance(e, ast.Name))
    return salida


def sacar_piezas(ruta, defs, nombres, base=None):
    """Ejecuta, en un espacio de nombres propio, los `def` y las asignaciones de NIVEL SUPERIOR
    de `ruta` que se piden por nombre, EN EL ORDEN DEL FICHERO (una constante puede depender
    de otra: `ORIGEN_IVA_ASUMIDO` usa `IVA_DEFAULT_ES`). Devuelve ese espacio de nombres.

    🔴 Falla CERRADO: si falta una pieza, o si esta asignada DOS veces a nivel superior (no se
       sabria cual es la buena), lanza PiezaNoEncontrada con los nombres. Nunca devuelve medio
       motor."""
    with io.open(ruta, encoding='utf-8') as fh:
        arbol = ast.parse(fh.read(), ruta)
    por_def, por_nombre = {}, {}
    for nodo in arbol.body:
        if isinstance(nodo, ast.FunctionDef) and nodo.name in defs:
            por_def.setdefault(nodo.name, []).append(nodo)
        elif isinstance(nodo, ast.Assign):
            for n in _nombres_asignados(nodo):
                if n in nombres:
                    por_nombre.setdefault(n, []).append(nodo)
    faltan = [d for d in defs if d not in por_def] + [n for n in nombres if n not in por_nombre]
    if faltan:
        raise PiezaNoEncontrada('%s: ya no tiene (a nivel superior) %s' % (ruta, ', '.join(faltan)))
    dobles = ([d for d, v in por_def.items() if len(v) > 1]
              + [n for n, v in por_nombre.items() if len(v) > 1])
    if dobles:
        raise PiezaNoEncontrada('%s: %s aparece(n) mas de una vez a nivel superior' % (ruta, ', '.join(dobles)))
    nodos = {}
    for lista in list(por_def.values()) + list(por_nombre.values()):
        nodos[id(lista[0])] = lista[0]
    cuerpo = sorted(nodos.values(), key=lambda n: n.lineno)
    ns = dict(base or {})
    exec(compile(ast.fix_missing_locations(ast.Module(body=cuerpo, type_ignores=[])), ruta, 'exec'), ns)
    return ns


# Lo que el escaner nuevo usa del viejo. Si alguna desaparece, el arranque lo dice por su nombre.
DEFS_MOTOR = ('norm', '_num', 'partir_ean', 'clasificar', 'core_ean', 'clasificar_chase',
              '_chk13', '_ean_ok', '_gtin14_ok', 'rescatar_gtin', 'variantes_ean', 'aviso_caja_incoherente',
              'calc_rentabilidad', 'decision_de',
              '_sup', 'iva_es_con_origen', 'iva_es_de', 'origen_iva_fila', 'es_propio')
NOMBRES_MOTOR = ('PAISES', 'IVA_DEFAULT_ES', 'IVA_IT', 'IVA_FR', 'IVA_DE', 'ALMACEN',
                 'COM_DIGITALES', 'UNIDADES_CASE_TCG', 'ISD_PAIS', 'SIN_ISD_HISTORICO',
                 '_RE_SUFIJO', '_RE_CAJA6', '_RE_CHASE_NOM', '_RE_CON_CHASE',
                 'UMBRAL_CAJA_VS_SUELTA', 'ORIGEN_IVA_FICHA', 'ORIGEN_IVA_ASUMIDO', 'PERFILES')


class Motor:
    """Las piezas del escaner viejo, ya ejecutadas. Se leen como atributos (`M.calc_rentabilidad`).

    🔑 Las funciones sacadas llevan como globales ESTE espacio de nombres, asi que el catalogo
       propio (`sup`, el que usa `iva_es_con_origen`) se pone aqui con `poner_catalogo_propio`
       y ellas lo ven igual que en el escaner viejo."""

    def __init__(self, ns):
        self._ns = ns

    def __getattr__(self, k):
        try:
            return self._ns[k]
        except KeyError:
            raise AttributeError(k)

    def poner_catalogo_propio(self, filas_productos):
        """Lo mismo que la Celda 5 del escaner viejo: `sup[norm(ean)] = fila`, solo filas con EAN."""
        sup = {}
        for p in filas_productos or []:
            if p.get('ean'):
                sup[self._ns['norm'](p['ean'])] = p
        self._ns['sup'] = sup
        return len(sup)

    def iva_por_pais(self, core):
        """El IVA de cada pais para ESTE producto, como la Celda 8 del escaner viejo:
        ES de la ficha (o el general, asumido); IT, FR y DE, el general del pais."""
        return {'ES': self.iva_es_de(core), 'IT': self.IVA_IT, 'FR': self.IVA_FR, 'DE': self.IVA_DE}


def cargar_motor(ruta=RUTA_MOTOR):
    return Motor(sacar_piezas(ruta, DEFS_MOTOR, NOMBRES_MOTOR, base={'re': re, 'sup': {}}))


def cargar_filtro_director(regla, ruta=RUTA_DIRECTOR_HEO):
    """El filtro de director_heo_prep.py, el suyo: `_quiere` con las marcas de `reglas_director`.
    Devuelve (quiere, info) con las marcas reales, si pide ofertas y el rank maximo."""
    ns = sacar_piezas(ruta, ('_quiere',), ('marcas', 'quiere_ofertas', 'marcas_reales', 'rank_max'),
                      base={'regla': regla})
    return ns['_quiere'], {'marcas_reales': list(ns['marcas_reales']),
                           'quiere_ofertas': bool(ns['quiere_ofertas']),
                           'rank_max': ns['rank_max']}


def _quiere_todas(f):
    """El filtro del modo «todas las marcas»: SOLO lo disponible, la primera condicion del
    `_quiere` del director. Ni marcas ni ofertas: no se lee `reglas_director`."""
    return f.get('estado') == 'disponible'


def filtro_todas():
    """(quiere, info) del modo 'todas', con la misma forma que `cargar_filtro_director`: sin
    marcas ni ofertas (None: no se han leido, no es que la lista este vacia)."""
    return _quiere_todas, {'marcas_reales': None, 'quiere_ofertas': None, 'rank_max': None}


# ── (B4) EL MODO «MARCAS ELEGIDAS» ─────────────────────────────────────────────────────────
# La lista llega de la v2 como input del workflow (JSON, por env: NUNCA interpolada dentro de un
# `run:`), y aqui se vuelve a validar ANTES de abrir ningun cliente: si no vale, no se corre.
class SeleccionInvalida(ValueError):
    """La seleccion de marcas que llega del workflow no se puede usar: la pasada no arranca."""


# Topes de la seleccion. HEO tiene 265 marcas con algo disponible (25-sep-2026) y la mas larga
# mide 37 caracteres: los topes dejan holgura y cortan cualquier otra cosa.
MAX_MARCAS = 500
MAX_LARGO_MARCA = 80
MAX_TEXTO_MARCAS = 20000
# 🔴 Lo que no puede llevar un nombre de marca: comillas dobles, acento grave, barra invertida,
#    `$ < > ; | { }` y cualquier caracter de control o salto de linea. El apostrofo SI pasa:
#    HEO tiene una marca que lo lleva («Loop' », medido el 25-sep-2026) y cae en «Otras».
_PROHIBIDOS_MARCA = set('"`\\$<>;|{}')


def _caracter_prohibido(c):
    import unicodedata
    return c in _PROHIBIDOS_MARCA or unicodedata.category(c) in ('Cc', 'Cf', 'Zl', 'Zp')


def validar_seleccion(texto_marcas, texto_ofertas):
    """(marcas, ofertas) de los dos inputs del modo 'elegidas', o SeleccionInvalida con el porque.
    `texto_marcas`: una lista JSON de textos; `texto_ofertas`: 'true' o 'false' (el booleano de
    `workflow_dispatch`). Las marcas salen sin espacios a los lados y sin repetir (sin distinguir
    mayusculas), en el orden en que llegan."""
    if texto_ofertas not in ('true', 'false'):
        raise SeleccionInvalida('la casilla de ofertas no es true ni false: %r' % (texto_ofertas,))
    ofertas = texto_ofertas == 'true'
    if not isinstance(texto_marcas, str) or not texto_marcas.strip():
        raise SeleccionInvalida('no llega la lista de marcas')
    if len(texto_marcas) > MAX_TEXTO_MARCAS:
        raise SeleccionInvalida('la lista de marcas pasa de %d caracteres' % MAX_TEXTO_MARCAS)
    import json
    try:
        lista = json.loads(texto_marcas)
    except ValueError as ex:
        raise SeleccionInvalida('la lista de marcas no es JSON: %s' % ex)
    if not isinstance(lista, list):
        raise SeleccionInvalida('la lista de marcas no es una lista')
    if len(lista) > MAX_MARCAS:
        raise SeleccionInvalida('la lista trae %d marcas; el tope es %d' % (len(lista), MAX_MARCAS))
    marcas, vistas = [], set()
    for i, m in enumerate(lista):
        if not isinstance(m, str):
            raise SeleccionInvalida('la marca %d no es un texto: %r' % (i + 1, m))
        malos = sorted({c for c in m if _caracter_prohibido(c)})
        if malos:
            raise SeleccionInvalida('la marca %d lleva caracteres que no puede llevar: %s'
                                    % (i + 1, ' '.join(repr(c) for c in malos)))
        limpia = m.strip()
        if len(limpia) > MAX_LARGO_MARCA:
            raise SeleccionInvalida('la marca %d pasa de %d caracteres' % (i + 1, MAX_LARGO_MARCA))
        if clave_marca(limpia) not in vistas:
            vistas.add(clave_marca(limpia))
            marcas.append(limpia)
    if not marcas and not ofertas:
        raise SeleccionInvalida('selección vacía: ni una marca ni las ofertas')
    return marcas, ofertas


def clave_marca(m):
    """La marca con la que se compara: sin espacios a los lados y sin distinguir mayusculas."""
    return str(m or '').strip().lower()


def filtro_elegidas(marcas, ofertas):
    """(quiere, info) del modo 'elegidas', con la misma forma que `cargar_filtro_director`.

    🔴 COINCIDENCIA EXACTA con el nombre de marca de HEO (sin espacios a los lados, sin
       distinguir mayusculas), NO por trozo como el `_quiere` del director (`mr.lower() in m`):
       con 265 marcas, «CID» metería cualquier marca que lleve esas tres letras. Lo disponible,
       y de una marca elegida o (con la casilla) en oferta; el resto de lo disponible es marca no
       elegida (la puerta `marca_fuera`)."""
    elegidas = frozenset(clave_marca(m) for m in marcas)
    ofertas = bool(ofertas)

    def quiere(f):
        if f.get('estado') != 'disponible':
            return False
        if clave_marca(f.get('marca')) in elegidas:
            return True
        return ofertas and f.get('en_oferta') == 'SI'
    return quiere, {'marcas_reales': list(marcas), 'quiere_ofertas': ofertas, 'rank_max': None}


def tanda_visualizador(ruta=RUTA_DESCARGAR_HEO):
    """Cuantos EAN caben en una tanda del Visualizador: la `TANDA` de descargar_heo.py (modo
    completo), evaluada igual que alli (acepta HEO_TANDA del entorno). Una sola asignacion o
    nada: si hubiera dos, no se sabria cual manda. 🔴 Y nunca mas que el tope del Visualizador
    (10.000 por lista): si alguien la sube por encima, no arranca, en vez de dejar tandas que
    Keepa no acepta enteras."""
    with io.open(ruta, encoding='utf-8') as fh:
        arbol = ast.parse(fh.read(), ruta)
    asignaciones = [n for n in ast.walk(arbol)
                    if isinstance(n, ast.Assign) and 'TANDA' in _nombres_asignados(n)]
    if len(asignaciones) != 1:
        raise PiezaNoEncontrada('%s: se esperaba UNA asignacion a TANDA y hay %d' % (ruta, len(asignaciones)))
    valor = eval(compile(ast.Expression(asignaciones[0].value), ruta, 'eval'), {'int': int, 'os': os})
    if not isinstance(valor, int) or valor <= 0:
        raise PiezaNoEncontrada('%s: TANDA no es un entero positivo (%r)' % (ruta, valor))
    if valor > TOPE_VISUALIZADOR:
        raise PiezaNoEncontrada('%s: TANDA %d pasa del tope del Visualizador (%d por lista)'
                                % (ruta, valor, TOPE_VISUALIZADOR))
    return valor


def columnas_keepa(ruta=RUTA_ESCAPARATE):
    """Las dos cabeceras que el lector del Escaner Pro NO lee y aqui hacen falta, sacadas de
    `TIPADAS` del procesador del escaparate (copiadas LITERALMENTE de los exports reales del
    Visualizador): el pais del fichero y las caidas del puesto en 30 dias."""
    ns = sacar_piezas(ruta, (), ('TIPADAS',))
    por_columna = {col: cab for cab, col, _t in ns['TIPADAS']}
    faltan = [c for c in ('dominio', 'rank_drops_30d') if c not in por_columna]
    if faltan:
        raise PiezaNoEncontrada('%s: TIPADAS ya no declara %s' % (ruta, ', '.join(faltan)))
    return por_columna['dominio'], por_columna['rank_drops_30d']


# ═══════════════════════════════════════════════════════════════════════════════
# 2 · LA FOTO DEL CATALOGO (paso 2 del encargo)
# ═══════════════════════════════════════════════════════════════════════════════
def _apartado(fila, motivo, detalle, ean=None, precio=None, en_oferta=None):
    # (B4) `en_oferta` solo se rellena en la marca fuera: es lo que el selector de marcas de la v2
    #      cuenta como «N en oferta» de las marcas que no entraron en la foto. En el resto, None.
    return {'producto_heo': fila.get('productNumber') or fila.get('producto_heo'),
            'ean_original': str(ean if ean is not None else (fila.get('ean') or '')).strip(),
            'nombre': fila.get('nombre') or '', 'marca': fila.get('marca') or '',
            'precio_catalogo': precio, 'motivo': motivo, 'detalle': detalle, 'en_oferta': en_oferta}


# ── Las CAJAS CON CHASE de HEO (encargo B2, 25-sep-2026) ─────────────────────────────────
# `descargar_heo._es_funko_chase` (regex por NOMBRE) desvia a una lista aparte todo Funko que
# «suena» a chase, con el codigo de la CAJA y no el de la figura. El escaner viejo los manda a la
# puente `escaner_chase_asin` para un ASIN a mano que nunca se ha puesto (0 de 120 el 25-sep), asi
# que ninguna caja chase de HEO se ha valorado nunca. Aqui se valoran con la regla de siempre de
# la caja con chase: precio de la caja ÷ unidades, sobre la figura COMUN (su EAN, y el ASIN que
# Keepa le da a ese EAN). 🔴 El chase nunca recibe ASIN propio, y nada de esto escribe en
# `productos` ni en la puente.
#
# Unidades: SOLO del nombre. «Surtido (6)», «Surtido 6», «Asst. (6)» (medido en las 83 del
# 25-sep-2026: «Surtido (6)», «Surtido (3)» y un «Asst.  (6)» con dos espacios). Un «Surtido 9 cm»
# no son 9 unidades. Si el nombre no lo dice, NO es caja (el regex tambien pesca un apellido,
# «Ja'Marr Chase», o un pack «w/CH» a precio de unidad): va al flujo normal como figura suelta.
_RE_UDS_CAJA = re.compile(r'\b(?:surtido|asst\.?)\s*(?:\(\s*(\d+)\s*\)|(\d+)\b(?!\s*(?:cm|mm)\b))', re.I)
# El numero de HEO de un Funko (FK87245, FK72611-01, FK95150-1): sus cinco cifras son las del UPC
# de Funko, 889698 + esas cinco + control (medido en las 38 cajas con codigo GS1 del 25-sep).
_RE_NUMERO_FUNKO = re.compile(r'^FK(\d{5})(?!\d)', re.I)
PREFIJO_FUNKO = '889698'
# De donde sale el EAN de la figura comun de una caja (escaner2_foto.origen_ean).
ORIGEN_EAN = {
    'gs1_caja': 'código GS1 de la caja',
    'ean_caja': 'EAN que trae la caja',
    'numero_heo': 'deducido del número de HEO',
}


def unidades_caja_chase(nombre):
    """Las unidades de la caja segun el nombre, o None si el nombre no las dice."""
    m = _RE_UDS_CAJA.search(str(nombre or ''))
    if not m:
        return None
    n = int(m.group(1) or m.group(2))
    return n if n > 0 else None


def _upc_ok(s):
    return s.isdigit() and len(s) == 12


def _a_upc(ean13):
    """Un EAN-13 que empieza por 0 es un UPC-A de 12: se guarda en su forma de 12 (la de Funko)."""
    return ean13[1:] if len(ean13) == 13 and ean13.startswith('0') else ean13


def ean_de_la_figura(ean_caja, producto_heo, M):
    """El EAN de la figura COMUN de una caja con chase de HEO → (ean, origen, aviso) o
    (None, None, motivo) si no se puede sacar. Por este orden (encargo B2):
      1. codigo GS1 «01» + GTIN-14 (+ «21»…) con su control valido → el interior del GTIN-14 con
         el control recalculado (`rescatar_gtin` del escaner viejo, leida de su fichero);
      2. si el codigo ya es un EAN-13 o un UPC-12 valido → ese;
      3. si no → 889698 + las cinco cifras del numero FK de HEO + control, y lo dice.
    Si 1 o 2 no coinciden con 3, se avisa en la fila (no se descarta)."""
    codigo = re.sub(r'\s+', '', str(ean_caja or ''))
    m = _RE_NUMERO_FUNKO.match(str(producto_heo or '').strip())
    deducido = None
    if m:
        cuerpo = PREFIJO_FUNKO + m.group(1)
        deducido = cuerpo + M._chk13('0' + cuerpo)
    visto, origen = None, None
    if codigo.isdigit() and codigo.startswith('01') and len(codigo) >= 16 and M._gtin14_ok(codigo[2:16]):
        visto, origen = _a_upc(M.rescatar_gtin(codigo[2:16])), 'gs1_caja'
    elif codigo.isdigit() and len(codigo) == 13 and M._ean_ok(codigo):
        visto, origen = _a_upc(codigo), 'ean_caja'
    elif _upc_ok(codigo) and M._ean_ok('0' + codigo):
        visto, origen = codigo, 'ean_caja'
    if visto is None:
        if deducido is None:
            return None, None, ('el código de la caja (%r) no es GS1 ni EAN, y el número de HEO (%r) no es '
                                'FK + cinco cifras' % (codigo, producto_heo))
        return deducido, 'numero_heo', None
    aviso = None
    if deducido is not None and M.norm(visto) != M.norm(deducido):
        aviso = ('el %s da %s y el número de HEO %s da %s'
                 % (ORIGEN_EAN[origen], visto, producto_heo, deducido))
    return visto, origen, aviso


# ── (B6) EL CODIGO DE 14 CIFRAS CON UN 0 DELANTE ──────────────────────────────────────────
# 🔑 DESVIO DELIBERADO DEL VIEJO (Fernando, 25-sep-2026). El viejo tira como «EAN forma rara» todo
#    codigo de 14 cifras (Celda 4, antes del rescate). Si el primero es un 0, es el MISMO EAN-13 con un
#    cero de relleno delante (GTIN-14 con indicador 0): el digito de control del GTIN no cambia con
#    ceros a la izquierda, asi que el EAN-13 que queda al quitarlo es el del producto. En ed95d086 eran
#    51 (Hasbro, Mattel, Phat Mojo, ThreeZero, Wizards) y ninguno estaba en la foto con sus 13 cifras.
#    Se busca en Keepa y se cruza con las 13; `ean_original` se guarda tal cual, con sus 14.
#    Los de 14 cifras que empiezan por 1-9 (codigo de CAJA: otra unidad) y los de 8 NO se tocan.
def core_de_heo(ean_in, M):
    """El EAN con el que se busca y se cruza una fila de HEO: `core_ean` del viejo y, si son 14 cifras
    que empiezan por 0, esas 14 sin el 0 (13 cifras). Todo lo demas, tal cual lo deja el viejo."""
    core = M.core_ean(ean_in)
    if core.isdigit() and len(core) == 14 and core.startswith('0'):
        return core[1:]
    return core


def rotulo_caja(f):
    """Lo que dice la fila de su caja, en pantalla y en el Excel: «caja con chase · N uds» (la
    caja con chase de HEO y la 5+1 del viejo, que es lo mismo), «caja x N» si es caja sin chase, o
    '' si es una unidad suelta."""
    if not f.get('es_caja'):
        return ''
    if f.get('es_chase'):
        return 'caja con chase · %s uds' % f.get('uds_caja')
    return 'caja x%s' % f.get('uds_caja')


def _fila_de_chase(c):
    """Una entrada de la lista `chase` de descargar_heo con la forma de una fila del catalogo,
    para que pase por el MISMO filtro y las MISMAS reglas que las demas. descargar_heo no calcula
    la oferta de estas (no se toca): sin dato, no es oferta."""
    return {'productNumber': c.get('producto_heo'), 'ean': c.get('ean_caja') or '',
            'nombre': c.get('nombre') or '', 'marca': c.get('marca') or '', 'categoria': '',
            'precio': '' if c.get('precio_caja') is None else c.get('precio_caja'),
            'precio_base': '', 'en_oferta': '', 'campana': '', 'estado': c.get('estado') or '',
            'disponibilidad': '', 'imagen': c.get('imagen') or '', 'fin_de_vida': '', 'preorder': ''}


def construir_foto(filas_heo, chase_heo, quiere, M, n_crudo=None, n_sin_gtin=None, n_declarado=None,
                   modo='marcas'):
    """Del catalogo CRUDO de HEO a la foto de la pasada, sin perder a nadie por el camino.

    `filas_heo, chase_heo` es lo que devuelve `descargar_catalogo_heo(con_chase=True)`;
    `n_crudo` (productos que se bajaron de HEO), `n_sin_gtin` (los que descargar_heo tira por no
    tener GTIN) y `n_declarado` (los que HEO DICE que tiene, `totalElements` de la API) salen de
    su log, porque esa funcion no los devuelve y NO se toca.

    Devuelve (foto, apartados, cuentas). Cada producto crudo sale por UNA puerta previa
    (PUERTAS_PREVIAS) o entra en la foto, en el MISMO orden que el viejo para el perfil HEO:
      0. descargar_heo desvia los Funko chase (codigo de caja) y tira los que no traen GTIN. De
         los desviados (encargo B2): los que el nombre dice CAJA de N van al paso 1 como caja con
         chase; los que no, vuelven al flujo normal como figura suelta con su codigo;
      1. el filtro (`quiere`: el del director o, en el modo «todas», solo disponible): lo que no
         esta disponible, y lo disponible cuya marca no esta en la regla (ni es oferta, si la
         regla pide ofertas), que se LISTA con su oferta. (B4) En el modo 'elegidas' la regla son
         las marcas elegidas, y el detalle dice «no elegida» (`modo`);
      1 bis. la caja con chase que pasa el filtro: el EAN de su figura comun
         (`ean_de_la_figura`) o, si no hay forma de sacarlo, puerta previa `chase_funko`;
      2. Celda 4: estado servible, chase SUELTO fuera, EAN de forma rara fuera (los GTIN-14 caen
         aqui: el viejo los rechaza ANTES del rescate). (B6) Salvo los de 14 cifras con un 0 delante,
         que entran con sus 13 (`core_de_heo`); si esas 13 ya estan, el dedup del paso 3 los junta;
      3. dedup del proveedor: una fila por (EAN, caja), la mas barata;
      4. guardarrail caja-vs-suelta (marca, no borra) y el precio POR UNIDAD (Celda 8)."""
    perfil = M.PERFILES[PROVEEDOR]
    previas = {p: 0 for p in PUERTAS_PREVIAS}
    chase_todo = list(chase_heo or [])
    sueltas_chase, cajas = [], []
    for c in chase_todo:
        uds = unidades_caja_chase(c.get('nombre'))
        if uds is None:
            sueltas_chase.append(_fila_de_chase(c))
        else:
            cajas.append((c, _fila_de_chase(c), uds))
    apartados = []

    def filtro(f):
        """🔑 `quiere` es quien DECIDE; esto solo pone nombre al porque, con su misma primera
        condicion: lo que no esta disponible no pasa; lo disponible que no pasa es por marca."""
        if quiere(f):
            return True
        if f.get('estado') != 'disponible':
            previas['no_disponible'] += 1
        else:
            detalle = 'Marca %r no elegida' if modo == 'elegidas' else 'Marca %r fuera de la lista del director'
            apartados.append(_apartado(f, 'marca_fuera', detalle % (f.get('marca') or ''),
                                       precio=M._num(f.get('precio', '')), en_oferta=f.get('en_oferta') == 'SI'))
        return False

    sel = [f for f in list(filas_heo) + sueltas_chase if filtro(f)]

    filas = []
    for c, f, uds in cajas:
        if not filtro(f):
            continue
        ean, origen, aviso = ean_de_la_figura(c.get('ean_caja'), c.get('producto_heo'), M)
        if ean is None:
            apartados.append(_apartado(c, 'chase_funko', 'Caja con chase: ' + aviso,
                                       ean=c.get('ean_caja') or '', precio=M._num(f['precio'])))
            continue
        filas.append({
            'producto_heo': c.get('producto_heo'), 'ean_in': f['ean'], 'core': ean,
            'nombre': f['nombre'], 'marca': f['marca'], 'categoria': '', 'pa': M._num(f['precio']),
            'es_chase': True, 'es_caja': True, 'uds_caja': uds,
            'en_oferta': False, 'campana': '', 'disponibilidad': '', 'fin_de_vida': False,
            'preorder': False, 'imagen': f['imagen'], 'origen_ean': origen, 'aviso_ean': aviso,
        })

    for f in sel:
        # Celda 4 del viejo, perfil HEO: estado permitido, sin columna de stock (se asume 1).
        if perfil.get('estados_ok'):
            if str(f.get(perfil['col_estado'], '')).strip() not in perfil['estados_ok']:
                apartados.append(_apartado(f, 'estado_no_servible',
                                           'Estado %r fuera de %s' % (f.get(perfil['col_estado']),
                                                                      perfil['estados_ok'])))
                continue
        ean_in = str(f.get(perfil['col_ean']) or '').strip()
        nombre = f.get(perfil['col_nombre']) or ''
        es_case, es_caja6, descartar = M.clasificar_chase(nombre, ean_in)
        if descartar:
            apartados.append(_apartado(f, 'chase_suelto',
                                       'Chase SUELTO descartado (solo se compra en caja de 6)'))
            continue
        core = core_de_heo(ean_in, M)
        if (not core.isdigit()) or len(core) not in (12, 13):
            apartados.append(_apartado(f, 'ean_forma_rara', 'EAN forma rara (len=%d)' % len(core)
                                       + (': GTIN-14, el escáner viejo lo rechaza antes del rescate'
                                          if core.isdigit() and len(core) == 14 else '')))
            continue
        filas.append({
            'producto_heo': f.get('productNumber'), 'ean_in': ean_in, 'core': core,
            'nombre': nombre, 'marca': f.get(perfil['col_marca']) or '',
            'categoria': f.get('categoria') or '',
            'pa': M._num(f.get(perfil['col_pa'], '')),
            'es_chase': bool(es_case), 'es_caja': bool(es_caja6),
            'uds_caja': (M.partir_ean(ean_in)[2] or M.UNIDADES_CASE_TCG),
            'en_oferta': f.get('en_oferta') == 'SI', 'campana': f.get('campana') or '',
            'disponibilidad': f.get('disponibilidad') or '', 'fin_de_vida': f.get('fin_de_vida') == 'SI',
            'preorder': f.get('preorder') == 'SI', 'imagen': f.get('imagen') or '',
            'origen_ean': None, 'aviso_ean': None,
        })

    # DEDUP del proveedor (moloka_escaner_nube.py, "DEDUP del proveedor"): una fila por
    # (norm(core), es_chase), la MAS BARATA; la otra se lista, no se tira en silencio.
    uni = {}
    for f in filas:
        k = (M.norm(f['core']), bool(f['es_chase']))
        prev = uni.get(k)
        if prev is None:
            uni[k] = f
            continue
        if f['pa'] is not None and (prev['pa'] is None or f['pa'] < prev['pa']):
            barato, caro = f, prev
        else:
            barato, caro = prev, f
        uni[k] = barato
        apartados.append(_apartado({'productNumber': caro['producto_heo'], 'nombre': caro['nombre'],
                                    'marca': caro['marca']}, 'duplicado_proveedor',
                                   'Duplicado del proveedor: me quedo con %s (esta venía a %s)'
                                   % (barato['pa'], caro['pa']), ean=caro['ean_in'], precio=caro['pa']))
    filas = list(uni.values())

    # GUARDARRAIL caja-vs-suelta (mismo bloque del viejo): marca, no borra.
    sueltas = {}
    for f in filas:
        if not f['es_caja'] and f['pa']:
            k = M.norm(f['core'])
            if k not in sueltas or f['pa'] < sueltas[k]:
                sueltas[k] = f['pa']
    for f in filas:
        f['aviso_caja'] = None
        if f['es_caja'] and f['pa']:
            su = sueltas.get(M.norm(f['core']))
            if su:
                pa_ud = (f['pa'] / (f.get('uds_caja') or M.UNIDADES_CASE_TCG)
                         if perfil.get('precio_caja6') == 'caja' else f['pa'])
                f['aviso_caja'] = M.aviso_caja_incoherente(pa_ud, su)

    foto = []
    for f in filas:
        # Celda 8: solo se divide donde el proveedor da el precio de la CAJA COMPLETA.
        pa = f['pa']
        if f['es_caja'] and pa and perfil.get('precio_caja6') == 'caja':
            pa = pa / (f.get('uds_caja') or M.UNIDADES_CASE_TCG)
        variantes = sorted({M.norm(v) for v in M.variantes_ean(f['core'])} - {''})
        foto.append({
            'producto_heo': f['producto_heo'], 'ean_original': f['ean_in'], 'ean_core': f['core'],
            'variantes': variantes, 'codigos_keepa': codigos_para_keepa(f['core'], M),
            'nombre': f['nombre'], 'marca': f['marca'], 'categoria': f['categoria'],
            'precio_catalogo': f['pa'], 'precio_unidad': pa,
            'es_caja': f['es_caja'], 'uds_caja': (f['uds_caja'] if f['es_caja'] else None),
            'es_chase': f['es_chase'], 'en_oferta': f['en_oferta'], 'campana': f['campana'],
            'disponibilidad': f['disponibilidad'], 'fin_de_vida': f['fin_de_vida'],
            'preorder': f['preorder'], 'imagen': f['imagen'], 'aviso_caja': f['aviso_caja'],
            'origen_ean': f['origen_ean'], 'aviso_ean': f['aviso_ean'],
        })

    for m in MOTIVOS_APARTADO:
        previas[m] = sum(1 for a in apartados if a['motivo'] == m)
    previas['sin_gtin'] = n_sin_gtin
    cuentas = {'n_crudo': n_crudo, 'n_declarado': n_declarado, 'n_foto': len(foto), 'previas': previas,
               'n_devueltos': len(filas_heo) + len(chase_todo)}
    cuentas.update(cuadre_previo(cuentas))
    return foto, apartados, cuentas


def cuadre_previo(cuentas):
    """🔴 El cuadre de ANTES de las puertas, desde el catalogo CRUDO: crudo = previas + foto.
    Un recuento que falta (None) NO es un cero: sin el no se puede afirmar que cuadra.
    Tambien se exige que lo que descargar_heo DEVOLVIO mas lo que tiro sin GTIN sea el crudo:
    es la prueba de que los dos numeros sacados de su log son los buenos."""
    previas = cuentas['previas']
    faltan = [p for p in PUERTAS_PREVIAS if previas.get(p) is None]
    if cuentas.get('n_crudo') is None:
        faltan.insert(0, 'crudo')
    if cuentas.get('n_declarado') is None:
        faltan.insert(0, 'total que declara HEO')
    if faltan:
        return {'n_previas': None, 'cuadra_previo': False,
                'motivo_previo': 'sin recuento de: ' + ', '.join(faltan)}
    n_previas = sum(previas[p] for p in PUERTAS_PREVIAS)
    # 🔴 El catalogo «tal como llega de HEO» es el que HEO DICE que tiene, no lo que se bajo:
    #    `_paginar` de descargar_heo se para EN SILENCIO si una pagina falla cinco veces, y sin
    #    esto el cuadre diria CUADRA sobre un catalogo recortado.
    if cuentas['n_crudo'] != cuentas['n_declarado']:
        return {'n_previas': n_previas, 'cuadra_previo': False,
                'motivo_previo': 'HEO dice que tiene %d productos y se bajaron %d: la descarga se cortó'
                                 % (cuentas['n_declarado'], cuentas['n_crudo'])}
    if cuentas['n_devueltos'] + previas['sin_gtin'] != cuentas['n_crudo']:
        return {'n_previas': n_previas, 'cuadra_previo': False,
                'motivo_previo': 'HEO dio %d productos y descargar_heo devolvió %d + %d sin GTIN'
                                 % (cuentas['n_crudo'], cuentas['n_devueltos'], previas['sin_gtin'])}
    if cuentas['n_crudo'] != n_previas + cuentas['n_foto']:
        return {'n_previas': n_previas, 'cuadra_previo': False,
                'motivo_previo': 'catálogo crudo %d ≠ puertas previas %d + foto %d'
                                 % (cuentas['n_crudo'], n_previas, cuentas['n_foto'])}
    return {'n_previas': n_previas, 'cuadra_previo': True, 'motivo_previo': None}


def codigos_para_keepa(core, M):
    """Lo que se pega en el Visualizador por esta fila: el EAN tal cual y, si el rescate de
    GTIN del escaner viejo (`variantes_ean`) reconstruye OTRO EAN distinto, tambien ese. Las
    variantes que solo cambian ceros a la izquierda no se repiten: Keepa y el lector del Pro
    comparan sin esos ceros."""
    codigos = [core]
    for v in sorted(M.variantes_ean(core)):
        if M.norm(v) and M.norm(v) != M.norm(core) and all(M.norm(v) != M.norm(c) for c in codigos):
            codigos.append(v)
    return codigos


def lista_para_keepa(foto):
    """Los codigos de toda la foto, sin repetir y en el orden de la foto: uno por linea."""
    return list(dict.fromkeys(c for f in foto for c in f['codigos_keepa']))


def partir_en_tandas(codigos, tanda):
    return [codigos[i:i + tanda] for i in range(0, len(codigos), tanda)] or [[]]


# ═══════════════════════════════════════════════════════════════════════════════
# 3 · LOS CSV DEL VISUALIZADOR (paso 3 y 4)
# ═══════════════════════════════════════════════════════════════════════════════
_RE_AMAZON = re.compile(r'amazon\.(es|de|fr|it)(?![a-z.])', re.I)


def _limpio(v):
    return ('' if v is None else str(v)).replace('﻿', '').replace('\xa0', ' ').strip()


def _leer_crudo(ruta):
    with open(ruta, encoding='utf-8-sig', newline='') as fh:
        filas = [r for r in csv.reader(fh) if any((c or '').strip() for c in r)]
    if not filas:
        raise CsvIlegible('el fichero está vacío')
    return [_limpio(h) for h in filas[0]], filas[1:]


def columnas_requeridas(pro, col_caidas):
    """Sin estas, el CSV no sirve para decidir nada: se rechaza con su nombre."""
    c = pro.CSV_COLS
    return [c['ean'], c['asin'], c['buybox'], c['nuevo'], c['fba'], c['compct'], col_caidas]


def examinar_csv(ruta, pro, col_pais, col_caidas):
    """Lee UN export del Visualizador y devuelve {'pais', 'filas', 'fuente_pais', 'caidas'}.

    🔑 EL PAIS LO DICE EL DATO, no se pregunta: la columna `Localización` (la que usa la
       Guarda 5 bis del procesador del escaparate) y, si no estuviera, el dominio de las URL
       de Amazon del propio fichero. Tiene que ser UNO en todo el fichero.
    `caidas` = {asin: caidas_30d o None}, leidas con el mismo `_num_csv` del Escaner Pro."""
    cab, filas = _leer_crudo(ruta)
    if not filas:
        raise CsvIlegible('sin filas de datos (solo cabecera)')
    idx = {}
    for i, h in enumerate(cab):
        idx.setdefault(h, i)
    faltan = [h for h in columnas_requeridas(pro, col_caidas) if h not in idx]
    if faltan:
        raise CsvIlegible('le faltan columnas del Visualizador: ' + ', '.join(repr(h) for h in faltan))

    def celda(fila, h):
        i = idx.get(h)
        return _limpio(fila[i]) if (i is not None and i < len(fila)) else ''

    if col_pais in idx:
        vistos = set()
        for n, fila in enumerate(filas, 2):
            v = celda(fila, col_pais).lower()
            if not v:
                raise CsvIlegible("fila %d: '%s' vacía; es la única fuente de país del fichero" % (n, col_pais))
            vistos.add(v)
        fuente = col_pais
    else:
        cols_url = [h for h in cab if 'url' in h.lower()]
        vistos = set()
        for fila in filas:
            doms = {m.group(1).lower() for h in cols_url for m in _RE_AMAZON.finditer(celda(fila, h))}
            vistos |= doms
        fuente = 'URL de Amazon'
        if not vistos:
            raise CsvIlegible("no trae '%s' ni URL de Amazon: no se puede saber de qué país es" % col_pais)
    if len(vistos) != 1:
        raise CsvIlegible('mezcla países (%s); cada export del Visualizador es de UN dominio'
                          % ', '.join(sorted(vistos)))
    pais = vistos.pop().upper()

    caidas, choques = {}, 0
    col_asin = pro.CSV_COLS['asin']
    for fila in filas:
        asin = celda(fila, col_asin)
        if not asin:
            continue
        v = pro._num_csv(celda(fila, col_caidas))
        if asin in caidas and caidas[asin] != v:
            choques += 1          # el mismo ASIN dos veces con caidas distintas: se queda la primera
            continue
        caidas.setdefault(asin, v)
    return {'pais': pais, 'filas': len(filas), 'fuente_pais': fuente, 'caidas': caidas,
            'choques_caidas': choques}


def candidatos(fila_foto, datos_pais):
    """Las fichas del CSV de UN pais para esta fila de la foto: todas las que casan con alguna de
    sus variantes (sin ceros a la izquierda, como el lector del Pro), sin repetir ASIN. Se
    conservan tambien las filas SIN ASIN: 'aparece sin ASIN' no es lo mismo que 'no aparece'."""
    salida, vistos = [], set()
    for v in fila_foto['variantes']:
        for rec in datos_pais.get(v, []):
            clave = rec.get('asin') or ('sin-asin', id(rec))
            if clave in vistos:
                continue
            vistos.add(clave)
            salida.append(rec)
    return salida


def precio_de_venta(rec):
    """El precio de venta de una ficha del CSV, COMO EL ESCANER PRO (escanear_pro): la Caja de
    Compra si hay; si no, «Nuevo: Actual»; si no, ninguno. Devuelve (precio, canal)."""
    if rec.get('buybox') and rec['buybox'] > 0:
        return rec['buybox'], ('BB-FBA' if rec.get('es_fba') else 'BB-FBM')
    if rec.get('nuevo') and rec['nuevo'] > 0:
        return rec['nuevo'], 'SIN BB'
    return None, 'sin precio'


def calcular_pais(pais, rec, pa, core, M):
    """La rentabilidad de UNA ficha en UN pais, con `calc_rentabilidad` del viejo y el recargo
    ISD de ese pais (lo que hace hoy la Celda 8). Misma condicion para calcular: precio, PA,
    comision y tarifa FBA; si falta una, 'Sin datos' y ni un numero inventado."""
    precio, canal = precio_de_venta(rec)
    iva = M.iva_por_pais(core)[pais]
    isd = M.ISD_PAIS[pais]
    out = {'pais': pais, 'asin': rec.get('asin'), 'titulo': rec.get('titulo') or '',
           'rank': rec.get('rank'), 'rank_90d': rec.get('rank90'),
           'precio_venta': precio, 'canal': canal, 'ref_pct': rec.get('compct'),
           'fee_fba': rec.get('fba'), 'iva': iva, 'iva_origen': M.origen_iva_fila(pais, iva, core),
           'almacen': M.ALMACEN, 'com_digitales': M.COM_DIGITALES,
           'isd_pct': isd['pct'], 'isd_incluye_fba': bool(isd['incluye_fba']), 'pa': pa}
    if precio and pa and rec.get('compct') is not None and rec.get('fba') is not None:
        r = M.calc_rentabilidad(precio, pa, rec['compct'], rec['fba'], iva,
                                almacen=M.ALMACEN, com_digitales=M.COM_DIGITALES, isd=isd)
        out.update(com_amazon=r['com_amazon'], beneficio=r['beneficio'], roi=r['roi'],
                   margen=r['margen'], decision=M.decision_de(r['margen']))
    else:
        out.update(com_amazon=None, beneficio=None, roi=None, margen=None, decision='Sin datos')
    return out


def _pct(x):
    return ('%.1f %%' % (x * 100)).replace('.', ',') if x is not None else '—'


def decidir(fila_foto, cands_por_pais, caidas_por_pais, params, M, eleccion=None):
    """La PUERTA de una fila de la foto. Devuelve {'puerta', 'motivo', 'detalle', 'asin',
    'fichas', 'paises': {pais: calculo}, 'caidas': {pais: n o None}, 'mejor': {...} o None}.

    `cands_por_pais`: {pais: [fichas del CSV]}, los paises que se CALCULAN de los que se ha
    subido CSV. `params` (escaner2_parametros, Fernando 24-sep-2026), DOS listas distintas:
      · 'paises_filtro'  (ES, DE): donde se mira si SE VENDE (> `umbral` caidas en 30 dias en
        alguno de ellos);
      · 'paises_calculo' (ES, IT, FR, DE): donde se CALCULA la rentabilidad, si traen CSV. El
        mejor pais sale de TODOS los calculados, venda o no alli: cada pais lleva `vende_aqui`
        y la pantalla marca «no vende aquí», pero se ve (decide Fernando).

    (B4 → B5) `eleccion` (`cargar_eleccion_viejo`): con dos o mas fichas, la del viejo elige UNA
    entre las de ES y el EAN sigue como si solo tuviera esa; su porque va en `detalle` y cada ficha
    lleva `elegida`. (B6) El titulo se coteja en todos los paises con CSV, no solo en ES.

    Orden de las puertas (una y solo una):
      a · ninguna ficha con ASIN en ningun CSV;
      b · dos o mas ASIN distintos para el mismo EAN y la regla del viejo no puede elegir (B5: sin
          `eleccion`, o sin ninguna ficha en ES, que es donde elige el viejo); se listan;
      c · no se vende: ningun pais del FILTRO con MAS de `umbral` caidas en 30 dias;
      d/e/f · se vende: COMPRAR si algun pais da COMPRAR, VALORAR si alguno da VALORAR, y si
              ninguno, sin margen. El mejor pais es el de mas margen dentro de esa decision."""
    todas = [(p, r) for p in cands_por_pais for r in cands_por_pais[p]]
    asins = list(dict.fromkeys(r['asin'] for _p, r in todas if r.get('asin')))
    base = {'asin': None, 'fichas': None, 'paises': {}, 'caidas': {}, 'mejor': None}

    if not asins:
        if not todas:
            return dict(base, puerta='a', motivo='a_no_aparece', detalle='No aparece en ningún CSV')
        return dict(base, puerta='a', motivo='a_sin_asin', detalle='Aparece en el CSV pero sin ASIN')

    if len(asins) >= 2:
        fichas = []
        for p, r in todas:
            if not r.get('asin'):
                continue
            precio, canal = precio_de_venta(r)
            fichas.append({'pais': p, 'asin': r['asin'], 'titulo': r.get('titulo') or '',
                           'rank': r.get('rank'), 'rank_90d': r.get('rank90'),
                           'caidas_30d': caidas_por_pais.get(p, {}).get(r['asin']),
                           'precio_venta': precio, 'canal': canal})
        detalle = '%d fichas para el mismo EAN: %s' % (len(asins), ', '.join(asins))
        # (B6) Las candidatas, las de ES; el titulo se coteja tambien en los demas paises con CSV.
        otros = {p: v for p, v in cands_por_pais.items() if p != 'ES'}
        elegida = (eleccion.elegir(fila_foto.get('nombre') or '', cands_por_pais.get('ES') or [], otros)
                   if eleccion else None)
        if elegida is None:
            if eleccion is not None:
                detalle += ' · ninguna en ES: el viejo elige entre las de ES, y aquí no hay ninguna'
            return dict(base, puerta='b', motivo='b_varias_fichas', detalle=detalle, fichas=fichas)
        # (B5) La del viejo ha elegido: el EAN sigue con ESA ficha, como si fuera la unica.
        for fi in fichas:
            fi['elegida'] = fi['asin'] == elegida['asin']
        salida = _decidir_con_asin(fila_foto, cands_por_pais, caidas_por_pais, params, M, elegida['asin'], base)
        descartadas = [a for a in asins if a != elegida['asin']]
        return dict(salida, fichas=fichas, eleccion=elegida,
                    detalle='Ficha %s elegida con la regla del viejo y el título en los cuatro países (%s: %s; descartadas %s) · %s'
                            % (elegida['asin'], elegida['veredicto'], elegida['detalle'],
                               ', '.join(descartadas), salida['detalle']))
    return _decidir_con_asin(fila_foto, cands_por_pais, caidas_por_pais, params, M, asins[0], base)


def _decidir_con_asin(fila_foto, cands_por_pais, caidas_por_pais, params, M, asin, base):
    """Las puertas c, d, e y f para UNA ficha (la unica, o la que eligio la regla del viejo)."""
    umbral = params['umbral']
    filtro = list(params['paises_filtro'])
    # Los que se calculan, en el orden de los parametros; solo los que traen CSV.
    paises = [p for p in params['paises_calculo'] if p in cands_por_pais]
    calculos, caidas = {}, {}
    for p in paises:
        rec = next((r for r in cands_por_pais.get(p, []) if r.get('asin') == asin), None)
        caidas[p] = caidas_por_pais.get(p, {}).get(asin) if rec is not None else None
        if rec is not None:
            calculos[p] = calcular_pais(p, rec, fila_foto['precio_unidad'], fila_foto['ean_core'], M)
            calculos[p]['caidas_30d'] = caidas[p]
            # 🔑 La marca «no vende aquí»: se decide AQUI y se guarda; la pantalla no la recalcula.
            calculos[p]['vende_aqui'] = caidas[p] is not None and caidas[p] > umbral
    salida = dict(base, asin=asin, paises=calculos, caidas=caidas)

    se_vende = any(caidas.get(p) is not None and caidas[p] > umbral for p in filtro)
    if not se_vende:
        con_dato = [p for p in filtro if caidas.get(p) is not None]
        sin_dato = [p for p in filtro if caidas.get(p) is None]
        if not con_dato:
            return dict(salida, puerta='c', motivo='c_sin_dato',
                        detalle='Sin dato de caídas en ' + ' ni en '.join(filtro))
        det = '≤ %d caídas en %s' % (umbral, ' y en '.join(con_dato))
        if sin_dato:
            det += ' (%s: sin dato)' % ', '.join(sin_dato)
        return dict(salida, puerta='c', motivo='c_pocas_caidas', detalle=det)

    def mejor_de(decision):
        cand = [p for p in paises if p in calculos and calculos[p]['decision'] == decision]
        if not cand:
            return None
        # El de mas margen; a igualdad, el primero de los paises configurados.
        return max(cand, key=lambda p: (calculos[p]['margen'], -paises.index(p)))

    for decision, puerta, motivo in (('COMPRAR', 'f', 'f_comprar'), ('VALORAR', 'e', 'e_valorar')):
        p = mejor_de(decision)
        if p:
            c = calculos[p]
            mejor = {'pais': p, 'margen': c['margen'], 'beneficio': c['beneficio'],
                     'precio_venta': c['precio_venta']}
            return dict(salida, puerta=puerta, motivo=motivo, mejor=mejor,
                        detalle='%s en %s (margen %s)' % (decision, p, _pct(c['margen'])))
    con_margen = [p for p in paises if p in calculos and calculos[p]['margen'] is not None]
    if con_margen:
        return dict(salida, puerta='d', motivo='d_sin_margen',
                    detalle='Margen por debajo de VALORAR: ' + ' · '.join(
                        '%s %s' % (p, _pct(calculos[p]['margen'])) for p in con_margen))
    return dict(salida, puerta='d', motivo='d_sin_datos',
                detalle='Se vende pero el CSV no trae precio, comisión o tarifa FBA para calcular')


# ═══════════════════════════════════════════════════════════════════════════════
# 4 · EL CUADRE (paso 5): entradas = suma de puertas, y cada EAN por UNA sola
# ═══════════════════════════════════════════════════════════════════════════════
def cuadre(ids_foto, resultados):
    """Comprueba que cada fila de la foto sale por UNA y solo una puerta.

    `ids_foto`: los ids de las filas de la foto (las ENTRADAS). `resultados`: dicts con 'foto_id'
    y 'puerta'. Devuelve {'cuadra', 'n_entradas', 'conteo', 'suma', 'n_c_pocas', 'n_c_sin_dato',
    'faltan', 'sobran', 'repetidos', 'puerta_rara'}. `cuadra` solo es True si NO falta ninguna
    entrada, NO sobra ninguna salida, NINGUNA sale dos veces, todas las puertas son de las seis y
    la suma de las puertas es igual a las entradas."""
    ids = list(ids_foto)
    conteo = {p: 0 for p in PUERTAS}
    vistos, repetidos, rara = set(), [], []
    for r in resultados:
        if r['foto_id'] in vistos:
            repetidos.append(r['foto_id'])
        vistos.add(r['foto_id'])
        if r.get('puerta') in conteo:
            conteo[r['puerta']] += 1
        else:
            rara.append(r.get('puerta'))
    faltan = sorted(set(ids) - vistos)
    sobran = sorted(vistos - set(ids))
    suma = sum(conteo.values())
    return {'cuadra': (not faltan and not sobran and not repetidos and not rara
                       and suma == len(ids) and len(set(ids)) == len(ids)),
            'n_entradas': len(ids), 'conteo': conteo, 'suma': suma,
            'n_c_pocas': sum(1 for r in resultados if r.get('motivo') == 'c_pocas_caidas'),
            'n_c_sin_dato': sum(1 for r in resultados if r.get('motivo') == 'c_sin_dato'),
            'faltan': faltan, 'sobran': sobran, 'repetidos': repetidos, 'puerta_rara': rara}


# ═══════════════════════════════════════════════════════════════════════════════
# 5 · LA COMPARACION CON EL ESCANER VIEJO (paso 6)
# ═══════════════════════════════════════════════════════════════════════════════
_RE_NOMBRE_VIEJO = re.compile(r'Escaneo_HEO_[^/]*_(\d{8})_(\d{4})\.xlsx$')
POSITIVAS = ('COMPRAR', 'VALORAR')


def fecha_del_excel_viejo(ruta):
    """`resultados/Escaneo_HEO_TODAS_20260923_2006.xlsx` → 2026-09-23 20:06 UTC. El sello lo pone
    el escaner viejo con `datetime.now()` en el runner de GitHub, que va en UTC."""
    m = _RE_NOMBRE_VIEJO.search(ruta or '')
    if not m:
        return None
    return datetime.strptime(m.group(1) + m.group(2), '%Y%m%d%H%M').replace(tzinfo=timezone.utc)


def leer_excel_viejo(contenido, M):
    """La hoja 'Análisis' de un Excel del escaner viejo → {ean_norm: {'ean', 'nombre',
    'decisiones': {pais: decision}}}. Por NOMBRE de columna, no por letra (el viejo avisa de que
    hay columnas que se han ido anadiendo al final)."""
    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO(contenido), read_only=True, data_only=False)
    if 'Análisis' not in wb.sheetnames:
        raise ValueError("el Excel no tiene la hoja 'Análisis' (tiene %s)" % wb.sheetnames)
    filas = wb['Análisis'].iter_rows(values_only=True)
    cab = [str(c or '').strip() for c in next(filas)]
    faltan = [c for c in ('EAN', 'País', 'Decisión') if c not in cab]
    if faltan:
        raise ValueError('a la hoja Análisis le faltan las columnas %s' % faltan)
    i_ean, i_pais, i_dec = cab.index('EAN'), cab.index('País'), cab.index('Decisión')
    i_nom = cab.index('Nombre') if 'Nombre' in cab else None
    salida = {}
    for fila in filas:
        if fila is None or i_ean >= len(fila) or fila[i_ean] in (None, ''):
            continue
        ean = str(fila[i_ean]).strip()
        k = M.norm(M.core_ean(ean))
        d = salida.setdefault(k, {'ean': ean, 'nombre': (fila[i_nom] if i_nom is not None else '') or '',
                                  'decisiones': {}})
        pais = str(fila[i_pais] or '').strip().upper()
        if pais:
            d['decisiones'][pais] = str(fila[i_dec] or '').strip()
    return salida


def fusionar_viejos(excels):
    """[(meta, datos)] con meta = {'fichero', 'fecha', 'modo'} → {ean_norm: {..., 'fichero',
    'fecha', 'modo'}}. Se recorren de MAS VIEJO a MAS NUEVO y el mas nuevo manda: es la ultima
    opinion del escaner viejo sobre cada EAN."""
    salida = {}
    for meta, datos in sorted(excels, key=lambda md: md[0]['fecha']):
        for k, d in datos.items():
            salida[k] = dict(d, fichero=meta['fichero'], fecha=meta['fecha'], modo=meta.get('modo'))
    return salida


def elegir_excels_viejos(filas_resultados):
    """De las filas de `escaner_resultados` de HEO con fichero (las mas nuevas primero) → las que
    hacen falta: el ULTIMO escaneo completo ('todo') y todas las pasadas de novedades
    ('nuevos') posteriores. Si no hay ninguno completo, todas las que haya (y se dice)."""
    filas = [f for f in filas_resultados if f.get('fichero')]
    for i, f in enumerate(filas):
        if (f.get('modo') or '').lower() == 'todo':
            return filas[:i + 1], True
    return filas, False


def _positiva(decisiones):
    """De {pais: decision} → (mejor, [paises con esa decision]) o (None, [])."""
    for d in POSITIVAS:
        ps = sorted(p for p, v in decisiones.items() if v == d)
        if ps:
            return d, ps
    return None, []


_ORDEN_PUERTA = {'f': 6, 'e': 5, 'd': 4, 'c': 3, 'b': 2, 'a': 1}


def comparar(viejo, nuevo, contexto, M):
    """Cruza la ultima opinion del viejo con la del nuevo, EAN a EAN.

    `viejo`: salida de fusionar_viejos. `nuevo`: {ean_norm: {'puerta', 'motivo', 'detalle',
    'mejor_pais', 'decision', 'ean', 'nombre', 'core', 'rank_es', 'rank90_es'}} (si un EAN tiene
    dos filas en la foto --unidad y caja-- manda la de puerta mas alta). `contexto`: {'umbral',
    'paises' (los calculados), 'paises_filtro', 'rank_max', 'apartados': {ean_norm: detalle},
    'fecha_nuevo'}.

    Solo entran los EAN que son COMPRAR o VALORAR en ALGUNO de los dos lados. Se rotula como
    DIFERENCIA DE CRITERIO lo que se explica por lo que el encargo nombra:
      · el viejo filtra por puesto ≤ rank_max en ES (actual o media de 90 dias) y el nuevo
        por caidas de 30 dias en ES o DE;
      · el viejo calcula ES/IT/FR/DE y el nuevo solo los paises configurados;
      · (B2) el viejo no valora cajas con chase de HEO: toda fila cuyo lado nuevo es una;
      · (B2) en el modo «todas», la marca que el viejo no mira: `contexto['marca_fuera_viejo']`,
        una funcion (fila nueva → texto o None) con la lista del viejo; None si no aplica.
    Todo lo demas queda como SIN EXPLICAR, con una nota de lo que dice cada lado."""
    # `paises_filtro`: donde el nuevo mira si se vende; `paises`: los que el nuevo CALCULO (con CSV).
    umbral, paises, rank_max = contexto['umbral'], list(contexto['paises']), contexto['rank_max']
    filtro = list(contexto.get('paises_filtro') or paises)
    filas = []
    for k in sorted(set(viejo) | set(nuevo)):
        v, n = viejo.get(k), nuevo.get(k)
        dec_v, paises_v = _positiva(v['decisiones']) if v else (None, [])
        dec_n = n['decision'] if (n and n['puerta'] in ('e', 'f')) else None
        if not dec_v and not dec_n:
            continue
        if dec_v and dec_n:
            categoria = 'ambos'
        elif dec_v:
            categoria = 'solo_viejo'
        else:
            categoria = 'solo_nuevo'

        criterio, notas = [], []
        # 🔑 El viejo manda las cajas con chase de HEO a la puente de ASIN a mano y no las valora
        #    nunca: cualquier diferencia con una de ellas es de criterio, no un misterio.
        if categoria != 'ambos' and n is not None and n.get('caja_chase_heo'):
            criterio.append('el viejo no valora cajas con chase de HEO')
        if categoria == 'solo_viejo':
            if n is None:
                # No es criterio: es que HOY no esta en el catalogo filtrado (o se aparto antes).
                notas.append(contexto['apartados'].get(k) or 'no está en el catálogo filtrado de esta pasada')
            elif n['puerta'] in ('a', 'b'):
                # El nuevo ni siquiera llego a calcularlo: eso hay que mirarlo, no explicarlo.
                notas.append('nuevo: puerta %s · %s' % (n['puerta'], n['detalle']))
            else:
                # c o d: el nuevo lo evaluo y dijo que no. Solo aqui caben los dos criterios.
                if paises_v and not [p for p in paises_v if p in paises]:
                    criterio.append('el viejo solo lo da en %s y el nuevo calcula %s'
                                    % (', '.join(paises_v), ' y '.join(paises)))
                # 🔑 Solo «pocas caidas» es criterio: «sin dato de caidas» es un HUECO de dato
                #    (el CSV no las trae), y un hueco se mira, no se explica.
                if n['motivo'] == 'c_pocas_caidas':
                    criterio.append('el nuevo pide más de %d caídas en 30 días en %s (%s); el viejo '
                                    'filtra por puesto ≤ %s en ES' % (umbral, ' o '.join(filtro), n['detalle'],
                                                                     '{:,}'.format(rank_max).replace(',', '.')))
                if not criterio:
                    notas.append('nuevo: puerta %s · %s' % (n['puerta'], n['detalle']))
        elif categoria == 'solo_nuevo':
            fuera = contexto.get('marca_fuera_viejo')
            fuera = fuera(n) if (fuera and v is None) else None
            if fuera:
                criterio.append(fuera)
            if v is not None:
                notas.append('el viejo lo evaluó: ' + ', '.join('%s %s' % (p, d) for p, d in sorted(v['decisiones'].items())))
            else:
                r, r90 = n.get('rank_es'), n.get('rank90_es')
                pasa = any(x and x > 0 and x <= rank_max for x in (r, r90))
                if not pasa and not n.get('propio'):
                    criterio.append('el viejo solo mira puesto ≤ %s en ES (actual o media de 90 días) y '
                                    'este tiene %s / %s' % ('{:,}'.format(rank_max).replace(',', '.'),
                                                            _num_o_raya(r), _num_o_raya(r90)))
                else:
                    notas.append('no está en los Excel del viejo')
        filas.append({
            'ean_norm': k,
            'ean': (n or {}).get('ean') or (v or {}).get('ean'),
            'nombre': (n or {}).get('nombre') or (v or {}).get('nombre') or '',
            'categoria': categoria,
            'decision_viejo': dec_v, 'pais_viejo': ', '.join(paises_v) or None,
            'fecha_viejo': v['fecha'] if v else None, 'excel_viejo': v['fichero'] if v else None,
            'puerta_nuevo': n['puerta'] if n else None, 'motivo_nuevo': n['motivo'] if n else None,
            'decision_nuevo': dec_n, 'pais_nuevo': n.get('mejor_pais') if n else None,
            'fecha_nuevo': contexto.get('fecha_nuevo'),
            'diferencia_criterio': bool(criterio) and categoria != 'ambos',
            'explicacion': (('Diferencia de criterio: ' + '; '.join(criterio)) if criterio
                            else ('Sin explicar: ' + '; '.join(notas)) if notas and categoria != 'ambos'
                            else None),
        })
    return filas


def marca_fuera_del_viejo(marcas, ofertas, ruta=RUTA_DIRECTOR_HEO):
    """(B2) Para comparar una pasada en modo «todas»: una funcion fila nueva → texto si su marca
    NO la miraria el viejo (su `_quiere`, el del director, con `marcas` y `ofertas`), o None.
    La lista NO sale de `reglas_director` (el modo «todas» no la lee): quien llama la toma de la
    ultima pasada de marcas de siempre, que la guardo al barrer (escaner2_pasada.marcas)."""
    quiere, info = cargar_filtro_director({'marcas': list(marcas or []) + (['OFERTAS'] if ofertas else [])}, ruta)
    lista = ', '.join(info['marcas_reales']) + (' y ofertas' if info['quiere_ofertas'] else '')

    def fuera(n):
        f = {'estado': 'disponible', 'marca': n.get('marca') or '', 'en_oferta': 'SI' if n.get('en_oferta') else ''}
        return None if quiere(f) else 'marca fuera de la lista del viejo (%s)' % lista
    return fuera


def _num_o_raya(x):
    return '—' if x is None or (isinstance(x, float) and math.isnan(x)) or x <= 0 else '{:,}'.format(int(x)).replace(',', '.')


def resumen_comparacion(filas):
    solo = [f for f in filas if f['categoria'] != 'ambos']
    return {'n_cmp_ambos': sum(1 for f in filas if f['categoria'] == 'ambos'),
            'n_cmp_solo_viejo': sum(1 for f in filas if f['categoria'] == 'solo_viejo'),
            'n_cmp_solo_nuevo': sum(1 for f in filas if f['categoria'] == 'solo_nuevo'),
            'n_cmp_criterio': sum(1 for f in solo if f['diferencia_criterio']),
            'n_cmp_sin_explicar': sum(1 for f in solo if not f['diferencia_criterio'])}


def nuevo_por_ean(foto, resultados_por_foto, M):
    """{ean_norm: lo que el nuevo dice de ese EAN}, para `comparar`. Si un EAN vive en dos filas
    de la foto (unidad y caja), manda la de puerta mas alta."""
    salida = {}
    for f in foto:
        r = resultados_por_foto[f['id']]
        k = M.norm(f['ean_core'])
        es = (r.get('paises') or {}).get('ES') or {}
        cand = {'puerta': r['puerta'], 'motivo': r['motivo'], 'detalle': r['detalle'],
                'mejor_pais': (r.get('mejor') or {}).get('pais'),
                'decision': {'f': 'COMPRAR', 'e': 'VALORAR'}.get(r['puerta']),
                'ean': f['ean_original'], 'nombre': f['nombre'], 'core': f['ean_core'],
                'rank_es': es.get('rank'), 'rank90_es': es.get('rank_90d'),
                'propio': M.es_propio(f['ean_core']),
                'caja_chase_heo': bool(f.get('origen_ean')), 'marca': f.get('marca') or '',
                'en_oferta': bool(f.get('en_oferta'))}
        prev = salida.get(k)
        if prev is None or _ORDEN_PUERTA[cand['puerta']] > _ORDEN_PUERTA[prev['puerta']]:
            salida[k] = cand
    return salida


# ═══════════════════════════════════════════════════════════════════════════════
# 6 · (B4) EL EXCEL CON EL FORMATO DEL VIEJO
# ═══════════════════════════════════════════════════════════════════════════════
# Fernando (25-sep-2026): «yo necesito exactamente el mismo formato de excel del escaner antiguo».
# 🔴 NO SE COPIA: SE EJECUTA EL SUYO. La «Celda 9» de moloka_escaner_nube.py (la que escribe
#    Análisis, Descartados, Ambiguos, Sin_rank, Precio por lote y Chase_manual) va EN LINEA, no en una
#    funcion, asi que se saca del fichero por ESTRUCTURA: las sentencias de nivel superior desde los
#    `import` de openpyxl que preceden a `COLS = …` hasta la anterior a `_sin_excel = …`, dos anclas
#    unicas o no arranca. Se ejecutan con los datos del escaner 2 vestidos como los suyos
#    (`registros`, `problematicos`, …): mismas hojas, columnas, formulas vivas, anchos y semaforo, y
#    si alguien cambia el viejo, el nuevo lo hereda sin tocar este fichero. El viejo NO se modifica.
# 🔑 Si el bloque empieza a usar un nombre que aqui no se le da, NO se adivina: revienta con el nombre.
PUERTAS_ANALISIS = ('d', 'e', 'f')
_ANCLA_INICIO_EXCEL, _ANCLA_FIN_EXCEL = 'COLS', '_sin_excel'


def _nombres_libres(nodos):
    """Los nombres que el bloque LEE y no define el mismo (ni son del lenguaje)."""
    import builtins
    leidos, definidos = set(), set()
    for n in nodos:
        for x in ast.walk(n):
            if isinstance(x, ast.Name):
                (leidos if isinstance(x.ctx, ast.Load) else definidos).add(x.id)
            elif isinstance(x, (ast.FunctionDef, ast.ClassDef)):
                definidos.add(x.name)
            elif isinstance(x, ast.arg):
                definidos.add(x.arg)
            elif isinstance(x, (ast.Import, ast.ImportFrom)):
                definidos.update((a.asname or a.name).split('.')[0] for a in x.names)
    return leidos - definidos - set(dir(builtins))


def sacar_bloque_excel(ruta=RUTA_MOTOR):
    """(codigo, nombres_que_necesita) de la Celda 9 del viejo. Falla CERRADO si una ancla falta,
    esta dos veces o estan al reves."""
    with io.open(ruta, encoding='utf-8') as fh:
        arbol = ast.parse(fh.read(), ruta)
    cuerpo = arbol.body

    def ancla(nombre):
        idx = [i for i, n in enumerate(cuerpo) if isinstance(n, ast.Assign) and nombre in _nombres_asignados(n)]
        if len(idx) != 1:
            raise PiezaNoEncontrada('%s: se esperaba UNA asignacion de nivel superior a %s y hay %d'
                                    % (ruta, nombre, len(idx)))
        return idx[0]
    ini, fin = ancla(_ANCLA_INICIO_EXCEL), ancla(_ANCLA_FIN_EXCEL)
    while ini > 0 and isinstance(cuerpo[ini - 1], (ast.Import, ast.ImportFrom)):
        ini -= 1
    if not ini < fin:
        raise PiezaNoEncontrada('%s: %s no va antes de %s' % (ruta, _ANCLA_INICIO_EXCEL, _ANCLA_FIN_EXCEL))
    nodos = cuerpo[ini:fin]
    return compile(ast.Module(body=nodos, type_ignores=[]), ruta, 'exec'), _nombres_libres(nodos)


def _pais_viejo(c):
    """Un pais calculado por el cruce, con las llaves con que la Celda 9 del viejo lo lee.
    `vendidos` y `n_of` no los guarda el cruce: van vacios, no inventados."""
    return {'rank_act': c.get('rank'), 'rank90': c.get('rank_90d'), 'vendidos': None, 'precio': c.get('precio_venta'),
            'canal': c.get('canal'), 'n_of': None, 'ref_pct': c.get('ref_pct'), 'fee': c.get('fee_fba'),
            'iva': c.get('iva'), 'decision': c.get('decision'), 'margen': c.get('margen')}


def datos_como_el_viejo(foto, resultados, apartados, M, eleccion=None):
    """Los datos del cruce con la forma de los del viejo, hoja a hoja (la correspondencia):
      · Análisis        ← los que SE VENDEN (puertas d, e y f), del de mas margen en ES al de menos,
                          como el viejo (que dejaba fuera por puesto lo que no se vende);
      · Descartados     ← EAN de forma rara y estado no servible (sus `problematicos`) + puerta a
                          (sus `no_encontrados`) + chase suelto + duplicado del proveedor, en su orden;
      · Ambiguos        ← (B5) cada EAN con dos o mas fichas en ES, con las filas que escribe el
                          viejo: `asin_elegido` = la que va ganando POR PUESTO al registrar cada
                          ficha (no la del cotejo); y los que siguen en la puerta b (sin ficha en ES),
                          con `asin_elegido` vacio;
      · Sin_rank        ← puerta c sin dato de caidas;
      · Precio por lote ← nada (es de OcioStock: en HEO el viejo tambien la deja solo con su cabecera);
      · Chase_manual    ← las cajas con chase sin EAN de la figura (puerta previa `chase_funko`).
    La marca fuera (o no elegida) NO va a Descartados: el viejo nunca la vio (el director la
    filtraba antes); esta en «Puertas previas»."""
    por_foto = {f['id']: f for f in foto}
    registros = []
    for r in resultados:
        if r['puerta'] not in PUERTAS_ANALISIS:
            continue
        f = por_foto[r['foto_id']]
        calc = r.get('paises') or {}
        titulo = next((calc[p].get('titulo') for p in M.PAISES if p in calc and calc[p].get('titulo')), '')
        registros.append({
            'nombre': f.get('nombre') or '', 'ean': f['ean_original'], 'asin': r['asin'], 'marca': f.get('marca') or '',
            'core': f['ean_core'], '_pa_efectivo': f.get('precio_unidad'), 'ambiguo': False, 'titulo_amz': titulo,
            # «Coincide» y «Cotejo» los saca el viejo con Keepa; el cruce no: vacios (None), no inventados.
            'coincide': None, 'coherencia_caja': f.get('aviso_caja'), 'url': '', 'volumen': None,
            '_paises_calc': {p: _pais_viejo(c) for p, c in calc.items()},
            '_margen_es': (calc.get('ES') or {}).get('margen')})
    registros.sort(key=lambda x: x['_margen_es'] if x['_margen_es'] is not None else -10 ** 9, reverse=True)

    def fila(ean, nombre, motivo):
        return {'EAN': ean, 'Cabecera': nombre or '', 'Motivo': motivo}

    def de(*motivos):
        return [fila(a['ean_original'], a['nombre'], a['detalle']) for a in apartados if a['motivo'] in motivos]
    ambiguos, sin_rank, no_encontrados, cotejo = [], [], [], {}
    for r in resultados:
        f = por_foto[r['foto_id']]
        # (B5) Las fichas de ES en el orden en que llegaron, con su puesto de 90 dias.
        es = [{'asin': x['asin'], 'rank90': x.get('rank_90d')} for x in (r.get('fichas') or []) if x.get('pais') == 'ES']
        if eleccion is not None and len({x['asin'] for x in es}) >= 2:
            ambiguos += [{'EAN': f['ean_original'], 'asin_elegido': a} for a in eleccion.ganador_por_puesto(es)]
        elif r['puerta'] == 'b':
            ambiguos.append({'EAN': f['ean_original'], 'asin_elegido': None})
        if (r.get('eleccion') or {}).get('n', 0) >= 2:
            cotejo[f['ean_original']] = {'veredicto': r['eleccion']['veredicto'], 'detalle': r['eleccion']['detalle']}
        if r['puerta'] == 'a':
            no_encontrados.append(fila(f['ean_original'], f.get('nombre'), r['detalle']))
        elif r['motivo'] == 'c_sin_dato':
            es = (r.get('paises') or {}).get('ES') or {}
            sin_rank.append({'ean_in': f['ean_original'], 'asin': r['asin'], 'fila': {'nombre': f.get('nombre') or ''},
                             'r_act': es.get('rank'), 'r_90': es.get('rank_90d')})
    # Las cajas con chase sin EAN de la figura pasaron el filtro, que exige «disponible»: de ahi su estado.
    chase = [{'nombre': a['nombre'] or '', 'producto_heo': a.get('producto_heo') or '', 'ean_caja': a['ean_original'],
              'precio_caja': a['precio_catalogo'], 'estado': 'disponible', 'imagen': '', 'link_amazon': ''}
             for a in apartados if a['motivo'] == 'chase_funko']
    return {'registros': registros, 'problematicos': de('ean_forma_rara', 'estado_no_servible'),
            'no_encontrados': no_encontrados, 'chase_sueltos': de('chase_suelto'), '_dups': de('duplicado_proveedor'),
            'ambiguos': ambiguos, 'sin_rank': sin_rank, 'chase_pendientes': chase,
            # Cotejo: el del viejo donde se ha elegido entre varias fichas (B5); en el resto la celda
            # va vacia (el viejo coteja tambien las de una sola ficha, y aqui eso no se hace).
            'cotejo_info': {x['ean']: cotejo.get(x['ean'], {'veredicto': None, 'detalle': None}) for x in registros},
            'PROVEEDOR': PROVEEDOR}


def excel_como_el_viejo(foto, resultados, apartados, M, ruta=RUTA_MOTOR, eleccion=None):
    """El libro de openpyxl con las SEIS hojas del viejo, escritas por SU codigo. El catalogo propio
    (`M.poner_catalogo_propio`) tiene que estar puesto: «En mi BD» sale de el, con `en_bd_txt` del viejo."""
    from contextlib import redirect_stdout
    codigo, necesita = sacar_bloque_excel(ruta)
    ns = sacar_piezas(ruta, ('pct_comision_celda', 'en_bd_txt'), (), base=M._ns)
    ns.update(datos_como_el_viejo(foto, resultados, apartados, M, eleccion))
    faltan = sorted(n for n in necesita if n not in ns)
    if faltan:
        raise PiezaNoEncontrada('%s: la Celda 9 del viejo usa %s y el escaner 2 no se lo da' % (ruta, ', '.join(faltan)))
    with redirect_stdout(io.StringIO()):     # sus `print` de la hoja no ensucian el log del cruce
        exec(codigo, ns)
    return ns['wb']


# ═══════════════════════════════════════════════════════════════════════════════
# 7 · (B5) ELEGIR LA FICHA CUANDO UN EAN TIENE VARIAS, CON LA REGLA DEL VIEJO
# ═══════════════════════════════════════════════════════════════════════════════
# moloka_escaner_nube.py, Celda 6 (Fase 1) y su COTEJO: el viejo pregunta a Keepa SOLO en ES; todas
# las fichas que devuelve para un EAN son candidatas (`cands_por_ean`), y `elegir_candidato` elige:
#   1. coteja el TITULO de cada ficha con el nombre del proveedor (`cotejar`: palabras distintivas
#      segun lo raras que son en el CATALOGO DEL PROPIO ESCANEO, `construir_idf`);
#   2. si alguna casa, de las que casan, la de menor puesto medio de 90 dias en ES (`keyrank`: sin
#      puesto de 90 dias, va la ultima);
#   3. si ninguna casa, la MAS PARECIDA (y a igualdad, la de mejor puesto), marcada «⚠ DUDOSO»; si
#      ninguna se parece en nada, la de mejor puesto, tambien «⚠ DUDOSO»;
#   4. si no hay texto con que cotejar, la de mejor puesto («n/d»). NUNCA se rinde: siempre elige.
# 🔴 NO SE COPIA: se sacan del fichero del viejo, por estructura, `_tok_cot`, `construir_idf`, `_idf`,
#    `_distintivo`, `cotejar`, `elegir_candidato`, `UMBRAL_COTEJO` y `keyrank` (que va anidada dentro
#    del `if filas:`, por eso se busca por nombre en todo el arbol y tiene que ser UNA). El cotejo
#    corre como en el viejo con HEO: activo (ningun workflow le pasa COTEJO_MODO=off y HEO trae nombre).
# 🔒 El ASIN elegido vive SOLO en el resultado del cruce: ni `productos` ni ninguna tabla de identidad.
DEFS_ELECCION = ('_tok_cot', 'construir_idf', '_idf', '_distintivo', 'cotejar', 'elegir_candidato')
NOMBRES_ELECCION = ('UMBRAL_COTEJO',)


def _def_anidada(ruta, nombre):
    """El `def` de ese nombre, este donde este dentro del fichero; tiene que haber UNO."""
    with io.open(ruta, encoding='utf-8') as fh:
        arbol = ast.parse(fh.read(), ruta)
    defs = [n for n in ast.walk(arbol) if isinstance(n, ast.FunctionDef) and n.name == nombre]
    if len(defs) != 1:
        raise PiezaNoEncontrada('%s: se esperaba UN def %s y hay %d' % (ruta, nombre, len(defs)))
    return defs[0]


class EleccionViejo:
    """La regla del viejo, lista para usar en un cruce. `nombres`: los del catalogo del escaneo (la
    foto de la pasada), sobre los que se calcula que palabras distinguen, como `construir_idf(filas)`."""

    def __init__(self, nombres, ruta=RUTA_MOTOR):
        from collections import Counter
        self._ns = sacar_piezas(ruta, DEFS_ELECCION, NOMBRES_ELECCION,
                                base={'re': re, 'math': math, 'unicodedata': __import__('unicodedata'),
                                      'Counter': Counter, '_DF': Counter(), '_NDOC': 1, 'COTEJO_ACTIVO': True})
        exec(compile(ast.Module(body=[_def_anidada(ruta, 'keyrank')], type_ignores=[]), ruta, 'exec'), self._ns)
        self._ns['construir_idf'](list(nombres))
        # (B6) Otra copia de `elegir_candidato` del viejo, cuyo `cotejar` es `_cotejar_paises`.
        self._ns6 = sacar_piezas(ruta, ('elegir_candidato',), (), base={'cotejar': self._cotejar_paises})

    @property
    def n_nombres(self):
        return self._ns['_NDOC']

    def elegir(self, nombre, registros_es, otros=None):
        """`registros_es`: las filas del CSV de ES de este EAN; `otros`: {pais: filas del CSV} de los demas
        paises (B6). → {'asin', 'veredicto', 'detalle', 'n', 'paises_cotejo', 'asin_viejo'} o None si en
        ES no hay ninguna ficha con ASIN (el viejo solo mira ES).

        🔑 (B6) DESVIO DELIBERADO DEL VIEJO (Fernando, 25-sep-2026), SOLO EN EL COTEJO. Las candidatas
           siguen siendo las de ES y el orden, `keyrank` (puesto medio de 90 dias en ES). Pero una
           candidata CASA si `cotejar` del viejo casa con su titulo en CUALQUIERA de los paises donde
           aparece (PAISES_COTEJO). El viejo solo miraba el titulo de ES, y un titulo de ES mal traducido
           le hacia elegir otro producto (Deck Case 100+ Black, 4260250075074: ES «Êltimo Guardia…», sin
           «deck» ni «case»; FR «Ultimate Guard Deck Case 100+ - Black»). Si ninguna casa en ningun pais,
           se sigue EXACTAMENTE como el viejo (su `elegir_candidato`: el mas parecido en ES, «⚠ DUDOSO»).
           `asin_viejo` es la que elegiria el viejo, y si es otra, el detalle lo dice."""
        cands, vistos = [], set()
        for r in registros_es:
            a = r.get('asin')
            if a and a not in vistos:
                vistos.add(a)
                cands.append({'asin': a, 'title': r.get('titulo') or '', 'r_90': r.get('rank90')})
        if not cands:
            return None
        if len(cands) == 1:
            return {'asin': cands[0]['asin'], 'veredicto': 'única en ES', 'n': 1, 'paises_cotejo': [],
                    'asin_viejo': cands[0]['asin'], 'detalle': 'solo hay una ficha en ES, que es donde mira el viejo'}
        kr = self._ns['keyrank']
        viejo, _veredicto_v, _detalle_v = self._ns['elegir_candidato'](nombre, cands, kr)
        # 🔑 La eleccion la hace el MISMO `elegir_candidato` del viejo (sacado de su fichero otra vez, en
        #    `self._ns6`), y lo unico que cambia es a que `cotejar` llama: `_cotejar_paises`, que pasa el
        #    `cotejar` del viejo por el titulo de cada pais. Ni una linea de su regla se copia aqui.
        titulos = titulos_por_pais(registros_es, otros)
        cands6 = [dict(c, title=_Titulos(titulos.get(c['asin']) or [('ES', c['title'])])) for c in cands]
        elegido, veredicto, detalle = self._ns6['elegir_candidato'](nombre, cands6, kr)
        paises = self._paises_que_casan(nombre, elegido['title'])
        if not paises:
            detalle += ' · ninguna casa en ningún país (%s): como el viejo' % '/'.join(PAISES_COTEJO)
        if elegido['asin'] != viejo['asin']:
            detalle += ' · distinta del viejo: el viejo elegiría %s' % viejo['asin']
        return {'asin': elegido['asin'], 'veredicto': veredicto, 'detalle': detalle, 'n': len(cands),
                'paises_cotejo': paises, 'asin_viejo': viejo['asin']}

    def _paises_que_casan(self, nombre, titulos):
        """Los paises (en el orden de PAISES_COTEJO) en los que `cotejar` del viejo casa con el titulo."""
        salida = []
        for pais, titulo in titulos:
            if pais not in salida and self._ns['cotejar'](nombre, titulo)[0] is True:
                salida.append(pais)
        return salida

    def _cotejar_paises(self, nombre_prov, titulos):
        """(B6) El `cotejar` que ve `elegir_candidato` en `self._ns6`: casa si el del viejo casa con el
        titulo de ALGUN pais, y lo dice («casó en IT, FR, DE · casa: case, deck»). Si no casa en ninguno,
        devuelve TAL CUAL lo que dice el del viejo con el titulo de ES (su «n/d», su «no casa» y su
        parecido), para que desde ahi la eleccion sea exactamente la del viejo."""
        paises = self._paises_que_casan(nombre_prov, titulos)
        if paises:
            _casa, score, motivo = self._ns['cotejar'](nombre_prov, dict(titulos)[paises[0]])
            return True, score, 'casó en %s · %s' % (', '.join(paises), motivo)
        titulo_es = dict(titulos).get('ES', '')
        return self._ns['cotejar'](nombre_prov, titulo_es)

    def ganador_por_puesto(self, registros_es):
        """Las filas de la hoja «Ambiguos» del viejo para este EAN: al registrar cada ficha nueva, la
        que va ganando POR PUESTO (`keyrank`, con «<» estricto: a igualdad se queda la de antes). Es lo
        que el viejo escribe en `asin_elegido`, que NO es la elegida por el cotejo."""
        kr = self._ns['keyrank']
        filas, prev = [], None
        vistos = set()
        for r in registros_es:
            a = r.get('asin')
            if not a or a in vistos:
                continue
            vistos.add(a)
            c = {'asin': a, 'r_90': r.get('rank90')}
            if prev is not None:
                filas.append(c['asin'] if kr(c) < kr(prev) else prev['asin'])
                if kr(c) < kr(prev):
                    prev = c
            else:
                prev = c
        return filas


# (B6) Los paises cuyo titulo se coteja, en este orden (ES primero: es el del viejo).
PAISES_COTEJO = ('ES', 'IT', 'FR', 'DE')


class _Titulos(tuple):
    """Los titulos de una ficha por pais, ((pais, titulo), …): lo que `elegir_candidato` del viejo lleva en
    `title` y le pasa a `cotejar` (en `EleccionViejo._ns6`, `_cotejar_paises`)."""


def titulos_por_pais(registros_es, otros=None):
    """{asin: [(pais, titulo)]}: los titulos de cada ficha en cada pais de PAISES_COTEJO donde aparece,
    en ese orden y sin repetir. `otros`: {pais: filas del CSV}; los paises de fuera de la lista no cuentan."""
    por_pais = dict(otros or {})
    por_pais['ES'] = registros_es
    salida = {}
    for pais in PAISES_COTEJO:
        for r in por_pais.get(pais) or []:
            a, t = r.get('asin'), (r.get('titulo') or '')
            if a and t and (pais, t) not in salida.setdefault(a, []):
                salida[a].append((pais, t))
    return salida


def cargar_eleccion_viejo(nombres, ruta=RUTA_MOTOR):
    return EleccionViejo(nombres, ruta)


# ── (B5-bis) EL CORPUS DEL COTEJO NO DEPENDE DE LAS MARCAS ELEGIDAS ─────────────────────────
# Las palabras «distintivas» se miden sobre un catalogo. Si fuera solo la FOTO, con marcas elegidas el
# catalogo encoge y cambia que palabra distingue (en d053378c, «deck» y «case» dejan de serlo): la
# eleccion dependeria de lo que Fernando marco. El corpus es el que tendria la pasada en modo «todas»:
# la foto + lo apartado por MARCA (marca fuera / no elegida) que habria pasado las demas puertas previas.
#   · estado no servible: la marca fuera se lista SOLO si esta «disponible», que es lo unico que HEO
#     admite (PERFILES['HEO']['estados_ok']): la pasa siempre;
#   · chase suelto y EAN de forma rara: las mismas funciones del viejo (`clasificar_chase`, `core_ean`) y,
#     desde el B6, el mismo `core_de_heo` que la foto (14 cifras con un 0 delante → sus 13);
#   · caja con chase (lista `chase` de descargar_heo, con unidades en el nombre): el mismo camino que en
#     `construir_foto` (`ean_de_la_figura`); si no sale EAN de la figura, iria a su puerta previa;
#   · duplicado del proveedor: una por (EAN, chase), la MAS BARATA, contando tambien con la foto.
def corpus_cotejo(foto, apartados, M):
    """(nombres, n_de_fuera): los nombres del corpus del cotejo y cuantos vienen de fuera de la foto."""
    perfil = M.PERFILES[PROVEEDOR]
    mejor = {}                                       # clave → (precio, nombre, de_fuera)

    def poner(clave, precio, nombre, de_fuera):
        prev = mejor.get(clave)
        if prev is None or (precio is not None and (prev[0] is None or precio < prev[0])):
            mejor[clave] = (precio, nombre, de_fuera)
    for f in foto:
        poner((M.norm(f['ean_core']), bool(f.get('es_chase'))), f.get('precio_catalogo'), f.get('nombre') or '', False)
    for a in apartados:
        if a.get('motivo') != 'marca_fuera':
            continue
        if perfil.get('estados_ok') and 'disponible' not in perfil['estados_ok']:
            continue
        nombre, ean = a.get('nombre') or '', str(a.get('ean_original') or '').strip()
        es_case, _es_caja6, descartar = M.clasificar_chase(nombre, ean)
        core = core_de_heo(ean, M)
        if (not core.isdigit()) or len(core) not in (12, 13):
            # ¿Una caja con chase de la lista `chase`? Entonces el EAN es el de su figura.
            if unidades_caja_chase(nombre) is None:
                continue
            figura, _origen, _aviso = ean_de_la_figura(ean, a.get('producto_heo'), M)
            if figura is None:
                continue
            poner((M.norm(figura), True), a.get('precio_catalogo'), nombre, True)
            continue
        if descartar:
            continue
        poner((M.norm(core), bool(es_case)), a.get('precio_catalogo'), nombre, True)
    nombres = [v[1] for v in mejor.values()]
    return nombres, sum(1 for v in mejor.values() if v[2])
