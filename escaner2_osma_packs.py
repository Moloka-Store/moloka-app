#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""ESCANER 2 · LOS PACKS DE OSMA QUE EXISTEN EN AMAZON Y NO VENDEMOS (encargo AN, plano Y, tramo 7, 03-oct-2026).

Fernando (01-oct-2026): «imagina que no vendieramos el lenor pack de dos, el escaner no se enteraria que existe esa
posibilidad… las kukident 99… en pack de dos que nosotros no tenemos, como se si eso me interesa?». Por EAN un pack de
Amazon no sale nunca (lleva su propio EAN). Se busca POR FAMILIA, en tres pasos:

  1. ENCOLAR (`python escaner2_osma_packs.py encolar`, ULTIMO paso del workflow escaner2-osma-cruce.yml, detras del
     cruce y sin tocar nada de lo que produce ni su Excel): una fila en `osma_packs_cola` por familia a buscar, las de
     las puertas d, e y f del cruce y TODAS nuestras fichas de OSMA de la foto (por codigo o por EAN). Por fila: el
     codigo, los EAN, el ASIN suelto de referencia (y cuantas unidades de OSMA lleva), su marca y su titulo EN AMAZON
     ESPAÑA (del CSV de Keepa, de amz_ficha o de keepa_escaparate; nunca el nombre aleman de OSMA), el PA de una unidad
     con porte y el Price_net. HEO no encola nada (este programa solo abre cruces de OSMA).
  2. BUSCAR: el cartero de la v2 (sp-api/cartero-osma-packs.mjs), en su vuelta horaria. Deja en `osma_packs_candidato`
     lo que contesto Amazon: el suelto, los «mismo», los «parecido» y los nuestros, con precio y tarifa de Amazon ES.
  3. VALORAR (`python escaner2_osma_packs.py valorar`, workflow escaner2-osma-packs.yml; AN-2: lo lanza cron-job.org a
     y 40 SIN cruce —el ultimo con su cola entera buscada y sin Excel de packs; si no hay, nada— y «Valorar packs» de la
     pestaña OSMA como atajo manual): el N de cada pack, su coste, la rentabilidad con la formula del viejo, el Excel
     «Escaner2_OSMA_Packs_<AAAAMMDD_HHMM>.xlsx» en la biblioteca y el aviso por Telegram.

🔑 EL N DEL PACK (la regla del AM, `escaner2_osma.factor_pack_amazon`: se multiplica SOLO con dos señales distintas que
   dicen el mismo N; una sola o contradictorias → «posible pack», nunca COMPRAR). Aqui la vara de medir es el SUELTO de
   referencia (sus lecturas de Amazon, guardadas por el cartero), no el nombre de OSMA: el «2» de «paquete de 2» del pack
   de Kukident se compara con el paquete del suelto (1), no con las 99 tabletas del nombre de OSMA. Señales:
     · paquete   item_package_quantity del pack ÷ el del suelto (1 si el suelto no lo dice);
     · recuento  unit_count del pack ÷ el del suelto (mismo tipo), o number_of_items ÷ number_of_items: UNA señal;
     · titulo    «pack de N», «paquete de N», «lote de N», «N x 99», «N unidades»… del titulo del pack, menos los
                 numeros que ya estan en el titulo del suelto (el «99» de las tabletas, el «34» de las toallitas).
   Veredicto, por este orden: (1) si el recuento da 1, el pack lleva lo mismo que el suelto: no es pack; (2) si ninguna
   lectura da 2 o mas: no es pack; (3) todas las lecturas dan el mismo N ≥ 2 desde dos señales distintas: pack de N;
   (4) si no: posible pack, con el N MAYOR que dice alguna señal (el coste mas prudente) y nunca COMPRAR.
   N de OSMA = N × las unidades de OSMA que lleva el suelto (`factor_ref` de la cola; si el suelto es a su vez un posible
   pack, el pack es posible pack). COSTE = N de OSMA × el PA de una unidad con porte YA REDONDEADO (Protefix 3 × 2,13).
🔑 LA RENTABILIDAD: la MISMA de siempre, `escaner2_motor.calcular_pais` → `calc_rentabilidad` y `decision_de` del viejo
   (COMPRAR ≥ 10 %, VALORAR ≥ 1 %), con el precio y la tarifa de Amazon como los usan las novedades (`rec_de_fila`:
   comision en % = ReferralFee ÷ precio; FBAFees tal cual; el ÷ 1,21 donde toque lo hace calc_rentabilidad, como hoy).
🔑 LAS VENTAS: caidas de 30 dias de Keepa de ESE ASIN en España, si estan guardadas desde hace 15 dias o menos: las de
   las novedades (`nov_keepa`) o las del CSV del Visualizador de ESTA pasada (los EAN de los packs encontrados entran
   solos en la lista del proximo «Barrer OSMA»: `eans_de_packs`). Con el corte del PRO (7 o mas: `escaner2_parametros`).
   (AN-2) Los packs que siguen sin ventas se piden a Keepa POR ASIN (stats=90, sin historial, ~1 token cada uno, con la
   via y la llave de las novedades, respetando su reserva de tokens) y se guardan en osma_packs_keepa.
   Sin ventas: el Excel enseña el puesto de ventas de Amazon y la decision queda en VALORAR como mucho, con el motivo
   «sin ventas de Keepa: entra en la lista del proximo PRO». Con ventas por debajo del corte: no se vende (fuera de
   «Análisis», contado en el Resumen).
🔴 REPO PUBLICO: LOS REGISTROS LOS VE CUALQUIERA. Solo estados y recuentos; nunca precios, EAN, ASIN ni nombres.
🔒 SOLO ESCRIBE `osma_packs_cola` (encolar: UNA insercion, entera o nada), `osma_packs_excel`, `osma_packs_keepa` y su Excel en
   `escaner2/osma/<pasada>/<id>/` (valorar).
   Ni productos, ni identidad: un pack encontrado es una PROPUESTA en un Excel, nunca un ASIN pegado a una ficha.
"""
import csv
import gzip
import io
import json
import os
import re
import sys
import tempfile
import uuid
from datetime import datetime, timedelta, timezone
from decimal import Decimal, ROUND_HALF_UP

import escaner2_motor as e2
import escaner2_osma as eo

BUCKET = 'escaner2'
PAIS = 'ES'
PUERTAS_COLA = ('d', 'e', 'f')
DIAS_VENTAS = 15
TOLERANCIA = eo.TOLERANCIA_PACK
TEXTO_SIN_VENTAS = 'sin ventas de Keepa: entra en la lista del próximo PRO'
TEXTO_PARECIDO = 'parecido: revisar'
# Las columnas que se AÑADEN al final de «Análisis» (detras de las del Análisis de OSMA, en este orden).
COLUMNAS_EXTRA = ['Unidades del pack', 'Señales del pack', 'ASIN suelto de referencia', 'Puesto de ventas de Amazon']
ANCHOS_EXTRA = [12, 70, 16, 18]


class FalloPacks(Exception):
    """Algo que impide seguir: se dice con su motivo (en la base o en el log, sin datos)."""


def _dec(x):
    return None if x is None else Decimal(str(x))


def _r2(x):
    return None if x is None else float(_dec(x).quantize(Decimal('0.01'), rounding=ROUND_HALF_UP))


def ean_norm(v):
    return eo.ean_norm(v)


# ═══════════════════════════════════════════════════════════════════════════════
# 1 · LA COLA: las familias de un cruce
# ═══════════════════════════════════════════════════════════════════════════════
def senales_csv_de(ruta):
    """{asin: {'titulo', 'marca'}} y {asin: señal del AM} de UN CSV del Visualizador (el mas nuevo manda). Puro."""
    titulos = {}
    with open(ruta, encoding='utf-8-sig', newline='') as fh:
        filas = csv.reader(fh)
        cab = next(filas, [])
        ix = {h: i for i, h in enumerate(cab)}
        col_tit = next((c for c in eo.COLUMNAS_TITULO if c in ix), None)
        for fila in filas:
            def celda(c):
                return fila[ix[c]].strip() if (c in ix and ix[c] < len(fila)) else ''
            asin = celda('ASIN')
            if asin and asin not in titulos:
                titulos[asin] = {'titulo': celda(col_tit) if col_tit else '', 'marca': celda('Marca')}
    return titulos, eo.senales_pack_csv(ruta)


def familias_para_cola(cruce, foto, resultados, enlaces, productos, csv_titulos, senales, amz_ficha, keepa):
    """Las filas de `osma_packs_cola` de UN cruce de OSMA → (filas, cuentas). Puro.
      · `foto`: escaner2_foto de la pasada; `resultados`: escaner2_resultado_ean del cruce (foto_id, puerta, asin);
      · `enlaces`: los de barrido.json (codigo → producto_id, ean_ficha, asins);
      · `productos`: id, asin, ean, nombre, activo, es_chase, unidades_por_pack;
      · `csv_titulos` / `senales`: del CSV de España de la pasada (`senales_csv_de`);
      · `amz_ficha` / `keepa`: {asin: {'titulo', 'marca'}} de España (amz_ficha, keepa_escaparate).
    Una familia por articulo de la foto que (a) cae en las puertas d, e o f, o (b) es NUESTRO (codigo enlazado, o EAN
    de una ficha viva). El suelto de referencia: el de MENOS unidades de nuestras fichas de la familia; si no es nuestra,
    el ASIN que eligio el cruce (con las unidades de OSMA que lleva, por la regla del AM; NULL si es posible pack)."""
    por_foto = {r['foto_id']: r for r in resultados}
    vivas = [p for p in productos if eo._ficha_viva(p)]
    por_ean = {}
    for p in vivas:
        k = ean_norm(p.get('ean'))
        if k:
            por_ean.setdefault(k, []).append(p)
    factores = eo.factores_por_ficha(productos)
    factor_asin, _ = eo.factor_por_asin(productos)
    nombre_de = {p['asin']: p.get('nombre') for p in vivas}
    filas, cuentas = [], {'puertas': 0, 'nuestras': 0, 'sin_marca': 0, 'sin_titulo': 0, 'sin_precio': 0, 'sin_asin': 0}
    for f in foto:
        r = por_foto.get(f['id']) or {}
        puerta = r.get('puerta') if r.get('puerta') in PUERTAS_COLA else None
        enlace = enlaces.get(str(f.get('producto_heo')))
        ean_ficha = (enlace or {}).get('ean_ficha')
        fichas = [x for x in vivas if x['asin'] in (enlace or {}).get('asins', [])] if enlace else \
            por_ean.get(ean_norm(f.get('ean_core')), [])
        nuestra = bool(fichas)
        if not puerta and not nuestra:
            continue
        if nuestra:
            suelta = sorted(fichas, key=lambda x: (factores.get(x['id'], 1), x['asin']))[0]
            asin_ref, factor_ref = suelta['asin'], factores.get(suelta['id'], 1)
        elif r.get('asin'):
            asin_ref = r['asin']
            if asin_ref in factor_asin:
                factor_ref = factor_asin[asin_ref]
            else:
                v = eo.factor_pack_amazon(senales.get(asin_ref), eo.cantidad_osma(f.get('nombre')))
                factor_ref = v['factor'] if v['estado'] == eo.PACK_SI else (None if v['estado'] == eo.PACK_DUDOSO else 1)
        else:
            cuentas['sin_asin'] += 1
            continue
        if f.get('precio_unidad') is None or f.get('precio_catalogo') is None:
            cuentas['sin_precio'] += 1
            continue
        fuente_tit = next(((d[asin_ref]['titulo'], o) for d, o in ((csv_titulos, 'csv'), (amz_ficha, 'amz_ficha'), (keepa, 'keepa_escaparate'))
                           if (d.get(asin_ref) or {}).get('titulo')), None)
        if fuente_tit is None and nombre_de.get(asin_ref):
            fuente_tit = (nombre_de[asin_ref], 'productos')
        fuente_marca = next(((d[asin_ref]['marca'], o) for d, o in ((csv_titulos, 'csv'), (amz_ficha, 'amz_ficha'), (keepa, 'keepa_escaparate'))
                             if (d.get(asin_ref) or {}).get('marca')), None)
        if fuente_tit is None:
            cuentas['sin_titulo'] += 1
            continue
        if fuente_marca is None:
            cuentas['sin_marca'] += 1
            continue
        eans = [x for x in dict.fromkeys(re.sub(r'\D', '', str(e or '')) for e in (f.get('ean_original'), ean_ficha)) if x]
        filas.append({
            'cruce_id': cruce, 'codigo_osma': str(f['producto_heo']), 'nombre_osma': f.get('nombre') or str(f['producto_heo']),
            'eans': eans[:4], 'asin_ref': asin_ref, 'factor_ref': factor_ref,
            'marca': fuente_marca[0], 'marca_origen': fuente_marca[1], 'titulo_ref': fuente_tit[0], 'titulo_origen': fuente_tit[1],
            'pa_unidad': f['precio_unidad'], 'price_net': f['precio_catalogo'], 'puerta': puerta, 'nuestra': nuestra,
            'fin_de_vida': bool(f.get('fin_de_vida')),
        })
        # Dos cuentas que se solapan (el Lenor es de la puerta f Y nuestro).
        cuentas['puertas'] += bool(puerta)
        cuentas['nuestras'] += nuestra
    return filas, cuentas


# ═══════════════════════════════════════════════════════════════════════════════
# 2 · EL N DEL PACK, contra el suelto (la regla del AM)
# ═══════════════════════════════════════════════════════════════════════════════
_RES_PACK_TITULO = [re.compile(p, re.I) for p in (
    r'\bpack\s*(?:de\s*)?(\d{1,3})\b', r'\b(\d{1,3})\s*-?\s*pack\b', r'\blote\s+de\s+(\d{1,3})\b',
    r'\bpaquete\s+de\s+(\d{1,3})\b', r'\bcaja\s+de\s+(\d{1,3})\b', r'\b(\d{1,2})\s*x\s*\d+', r'\b(\d{1,3})\s+(?:piezas|unidades)\b',
    r'\bconfezione\s+da\s+(\d{1,3})\b', r'\blot\s+de\s+(\d{1,3})\b')]


def _numeros(texto):
    return set(re.findall(r'(?<![\d.,])(\d+)(?![\d.,])', str(texto or '')))


def _entero(v):
    return eo._entero_pack(v)


def lecturas_n(pack, suelto):
    """Las lecturas de N (cuantos sueltos lleva el pack) → [{'grupo', 'texto', 'valor'}]. Puro."""
    s = suelto or {}
    salida = []
    if pack.get('paquete'):
        base = s.get('paquete') or 1
        salida.append({'grupo': 'paquete', 'texto': 'paquete %s (suelto %s)' % (eo._cifra_pack(pack['paquete']), eo._cifra_pack(base)),
                       'valor': pack['paquete'] / base})
    tipo_p, tipo_s = str(pack.get('tipo_ud') or '').lower(), str(s.get('tipo_ud') or '').lower()
    if pack.get('valor_ud') and s.get('valor_ud') and tipo_p == tipo_s:
        salida.append({'grupo': 'recuento', 'texto': 'contenido %s %s (suelto %s)' % (eo._cifra_pack(pack['valor_ud']), pack.get('tipo_ud') or '',
                                                                                     eo._cifra_pack(s['valor_ud'])),
                       'valor': pack['valor_ud'] / s['valor_ud']})
    elif pack.get('n_art') and s.get('n_art'):
        salida.append({'grupo': 'recuento', 'texto': 'n.º de artículos %s (suelto %s)' % (eo._cifra_pack(pack['n_art']), eo._cifra_pack(s['n_art'])),
                       'valor': pack['n_art'] / s['n_art']})
    del_suelto = _numeros(s.get('titulo'))
    vistos = set()
    for rx in _RES_PACK_TITULO:
        for m in rx.finditer(str(pack.get('titulo') or '')):
            n = m.group(1)
            if int(n) >= 2 and n not in del_suelto and n not in vistos:
                vistos.add(n)
                salida.append({'grupo': 'título', 'texto': 'título «%s»' % m.group(0).strip(), 'valor': float(n)})
    return salida


def n_del_pack(pack, suelto):
    """El veredicto → {'estado': PACK_SI | PACK_DUDOSO | PACK_NO, 'n', 'senales': texto}. Puro."""
    lect = lecturas_n(pack, suelto)
    texto = ' · '.join('%s → %s' % (x['texto'], eo._cifra_pack(x['valor'])) for x in lect) or 'sin señales'
    if not suelto:
        texto = 'sin las lecturas del suelto · ' + texto
    rec = [x for x in lect if x['grupo'] == 'recuento']
    if rec and all(_entero(x['valor']) == 1 for x in rec):
        return {'estado': eo.PACK_NO, 'n': 1, 'senales': texto}
    packs = [x for x in lect if (_entero(x['valor']) or 0) >= 2]
    if not packs:
        return {'estado': eo.PACK_NO, 'n': 1, 'senales': texto}
    ns = {_entero(x['valor']) for x in packs}
    contra = [x for x in lect if _entero(x['valor']) not in ns]
    if len(ns) == 1 and not contra and len({x['grupo'] for x in packs}) >= 2:
        return {'estado': eo.PACK_SI, 'n': ns.pop(), 'senales': texto}
    return {'estado': eo.PACK_DUDOSO, 'n': max(ns), 'senales': texto}


# ═══════════════════════════════════════════════════════════════════════════════
# 3 · LA VALORACION
# ═══════════════════════════════════════════════════════════════════════════════
def rec_de_candidato(c):
    """El candidato con la forma de la fila del CSV del Visualizador que lee `calcular_pais`, como `rec_de_fila` de las
    novedades: el precio como caja (BB-FBA / BB-FBM) o como «nuevo» (SIN BB); la comision en % del precio. Puro."""
    precio, canal = c.get('precio'), c.get('canal')
    ref = c.get('ref_eur')
    ref_pct = (float(ref) / float(precio) * 100) if (c.get('resultado') == 'dato' and precio and ref is not None) else None
    fee = c.get('fee_fba') if c.get('resultado') == 'dato' else None
    return {'asin': c['asin'], 'titulo': c.get('titulo') or '', 'rank': c.get('rank'), 'rank90': None,
            'buybox': float(precio) if (precio and canal in ('BB-FBA', 'BB-FBM')) else None, 'es_fba': canal == 'BB-FBA',
            'nuevo': float(precio) if (precio and canal == 'SIN BB') else None, 'compct': ref_pct,
            'fba': None if fee is None else float(fee)}


def ventas_de(asin, nov_keepa, csv_caidas, ahora, packs_keepa=None):
    """(caidas de 30 dias, de donde) de Keepa en España guardadas hace DIAS_VENTAS o menos, o (None, None). Puro.
    `nov_keepa`: filas (pais ES) con `fichas` y `consultada_en`; `csv_caidas`: {asin: caidas} con 'fecha' (la subida);
    (AN-2) `packs_keepa`: filas de osma_packs_keepa (asin, caidas_30d, consultada_en), las pedidas por ASIN. La mas nueva manda."""
    limite = ahora - timedelta(days=DIAS_VENTAS)
    mejor = None
    for k in nov_keepa or []:
        cuando = _fecha(k.get('consultada_en'))
        if cuando is None or cuando < limite:
            continue
        for fi in k.get('fichas') or []:
            if fi.get('asin') == asin and fi.get('caidas_30d') is not None and (mejor is None or cuando > mejor[0]):
                mejor = (cuando, fi['caidas_30d'], 'Keepa (novedades, %s)' % cuando.strftime('%d/%m'))
    for k in packs_keepa or []:
        cuando = _fecha(k.get('consultada_en'))
        if k.get('asin') == asin and k.get('caidas_30d') is not None and cuando is not None and cuando >= limite                 and (mejor is None or cuando > mejor[0]):
            mejor = (cuando, k['caidas_30d'], 'Keepa (por ASIN, %s)' % cuando.strftime('%d/%m'))
    if csv_caidas and asin in (csv_caidas.get('caidas') or {}) and csv_caidas['caidas'][asin] is not None:
        cuando = _fecha(csv_caidas.get('fecha'))
        if cuando is not None and cuando >= limite and (mejor is None or cuando > mejor[0]):
            mejor = (cuando, csv_caidas['caidas'][asin], 'Keepa (CSV del PRO, %s)' % cuando.strftime('%d/%m'))
    return (mejor[1], mejor[2]) if mejor else (None, None)


def _fecha(v):
    if v is None:
        return None
    if isinstance(v, datetime):
        return v if v.tzinfo else v.replace(tzinfo=timezone.utc)
    try:
        d = datetime.fromisoformat(str(v).replace('Z', '+00:00'))
        return d if d.tzinfo else d.replace(tzinfo=timezone.utc)
    except ValueError:
        return None


def valorar(cola, candidatos, M, umbral, nov_keepa, csv_caidas, ahora, packs_keepa=None):
    """La valoracion de los packs de un cruce → (packs, cuentas). Puro (salvo `M`, el motor del viejo ya cargado).
    Cada pack: {'fam', 'c' (el candidato), 'n', 'n_osma', 'estado', 'senales', 'pa', 'calc', 'tope', 'motivos',
    'caidas', 'ventas_de', 'decision'}. Solo «mismo» y «parecido» que sean pack (o posible pack); nunca nuestros."""
    por_cola = {}
    for c in candidatos:
        por_cola.setdefault(c['cola_id'], []).append(c)
    packs = []
    cuentas = {'familias': len(cola), 'buscadas': 0, 'fallidas': 0, 'pendientes': 0, 'candidatos': 0, 'no_pack': 0,
               'no_se_vende': [], 'nuestros': [], 'sin_precio': 0, 'sin_ventas': 0, 'repetidos': 0}
    vistos = set()
    for fam in sorted(cola, key=lambda x: eo._clave_codigo(x['codigo_osma'])):
        cuentas[{'buscada': 'buscadas', 'fallida': 'fallidas'}.get(fam['estado'], 'pendientes')] += 1
        cs = por_cola.get(fam['id'], [])
        suelto = next((c for c in cs if c['clase'] == 'suelto'), None)
        cuentas['nuestros'] += [(fam['codigo_osma'], c['asin']) for c in cs if c['clase'] == 'nuestro']
        for c in cs:
            if c['clase'] not in ('mismo', 'parecido'):
                continue
            cuentas['candidatos'] += 1
            # El mismo pack encontrado desde dos familias: se valora UNA vez, con la primera (la cola va por codigo).
            if c['asin'] in vistos:
                cuentas['repetidos'] += 1
                continue
            vistos.add(c['asin'])
            v = n_del_pack(c, suelto)
            if v['estado'] == eo.PACK_NO:
                cuentas['no_pack'] += 1
                continue
            estado = v['estado']
            if fam.get('factor_ref') is None and estado == eo.PACK_SI:
                estado = eo.PACK_DUDOSO
            n_osma = v['n'] * (fam.get('factor_ref') or 1)
            pa = _r2(_dec(fam['pa_unidad']) * n_osma)
            caidas, origen = ventas_de(c['asin'], nov_keepa, csv_caidas, ahora, packs_keepa)
            if caidas is not None and not e2.pasa_corte(caidas, umbral):
                cuentas['no_se_vende'].append((fam['codigo_osma'], c['asin'], caidas))
                continue
            core = (fam.get('eans') or [''])[-1]
            calc = e2.calcular_pais(PAIS, rec_de_candidato(c), pa, core, M)
            calc['caidas_30d'] = caidas
            motivos = []
            if estado == eo.PACK_DUDOSO:
                motivos.append('%s (%s)' % (eo.TEXTO_POSIBLE_PACK, v['senales']))
            if c['clase'] == 'parecido':
                motivos.append(TEXTO_PARECIDO)
            if caidas is None:
                motivos.append(TEXTO_SIN_VENTAS)
                cuentas['sin_ventas'] += 1
            if calc['decision'] == 'Sin datos':
                cuentas['sin_precio'] += 1
            tope = bool(motivos)
            decision = 'VALORAR' if (tope and calc['decision'] == 'COMPRAR') else calc['decision']
            packs.append({'fam': fam, 'c': c, 'n': v['n'], 'n_osma': n_osma, 'estado': estado, 'senales': v['senales'], 'pa': pa,
                          'calc': dict(calc, decision=decision), 'tope': tope, 'motivos': motivos, 'caidas': caidas,
                          'ventas_de': origen, 'decision': decision})
    return packs, cuentas


# ═══════════════════════════════════════════════════════════════════════════════
# 4 · EL EXCEL: «Análisis» con las columnas del de OSMA (solo ES) y «Resumen» con el pedido y el porte vivos
# ═══════════════════════════════════════════════════════════════════════════════
def _puerta(decision):
    return {'COMPRAR': 'f', 'VALORAR': 'e'}.get(decision, 'd')


def foto_y_resultados(packs):
    """Los packs con la forma de la foto y los resultados del cruce, para la Celda 9 del viejo. Puro."""
    foto, resultados = [], []
    for i, p in enumerate(packs):
        fam, c = p['fam'], p['c']
        clave = 'pack-%d' % i
        aviso = eo.texto_pack(p['n_osma'], fam['pa_unidad'], p['pa'], 'amazon')
        if p['motivos']:
            aviso += ' · ' + ' · '.join(p['motivos'])
        foto.append({'id': clave, 'ean_original': (fam.get('eans') or [''])[0], 'ean_core': (fam.get('eans') or [''])[-1],
                     'nombre': '%s · pack de %d en Amazon (%d × %s €)' % (fam['nombre_osma'], p['n_osma'], p['n_osma'], eo._coma(float(fam['pa_unidad']))),
                     'marca': fam.get('marca') or '', 'precio_unidad': p['pa'], 'aviso_caja': aviso,
                     'fin_de_vida': bool(fam.get('fin_de_vida')), '_asin': c['asin'],
                     '_sin_porte': float(_dec(fam['price_net']) * p['n_osma']), '_factor': p['n_osma']})
        resultados.append({'foto_id': clave, 'puerta': _puerta(p['decision']), 'motivo': 'pack', 'detalle': aviso, 'asin': c['asin'],
                           'paises': {PAIS: p['calc']}, 'fichas': None,
                           'pack_amazon': {'estado': eo.PACK_DUDOSO if p['tope'] else eo.PACK_SI}})
    return foto, resultados


def poner_pedido_vivo_packs(wb, foto_excel, resultados):
    """`escaner2_osma.poner_pedido_vivo`, con la fila buscada por su ASIN (aqui dos packs de una familia comparten EAN):
    «Precio OSMA sin porte (€)» al final, «PA (€)» y «Decisión» como formulas (`formula_pa` y `formula_decision` del AM;
    con tope —posible pack, parecido o sin ventas— la Decisión no llega a COMPRAR)."""
    from openpyxl.styles import Font
    from openpyxl.utils import get_column_letter
    ws = wb['Análisis']
    cab = [c.value for c in ws[1]]
    i = {n: cab.index(n) + 1 for n in ('ASIN', 'PA (€)', 'Beneficio (€)', 'Margen', 'Decisión')}
    col = len(cab) + 1
    letra = get_column_letter(col)
    ws.cell(row=1, column=col, value=eo.COLUMNA_SIN_PORTE).font = Font(bold=True)
    ws.column_dimensions[letra].width = eo.ANCHO_SIN_PORTE
    por_asin = {f['_asin']: f for f in foto_excel}
    tope = {r['asin']: (r.get('pack_amazon') or {}).get('estado') == eo.PACK_DUDOSO for r in resultados}
    for fila in range(2, ws.max_row + 1):
        f = por_asin[str(ws.cell(row=fila, column=i['ASIN']).value)]
        c = ws.cell(row=fila, column=col, value=f['_sin_porte'])
        c.number_format = '0.000'
        if ws.cell(row=fila, column=i['PA (€)']).value is not None:
            ws.cell(row=fila, column=i['PA (€)']).value = eo.formula_pa(letra, fila, f['_factor'])
        if str(ws.cell(row=fila, column=i['Beneficio (€)']).value or '').startswith('='):
            ws.cell(row=fila, column=i['Decisión']).value = eo.formula_decision(get_column_letter(i['Margen']), fila, tope[f['_asin']])
    _alargar_tabla(ws, col, eo.COLUMNA_SIN_PORTE)


def _alargar_tabla(ws, col, nombre):
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.table import TableColumn
    for tabla in ws.tables.values():
        ini, fin = tabla.ref.split(':')
        tabla.ref = '%s:%s%s' % (ini, get_column_letter(col), re.sub(r'^[A-Z]+', '', fin))
        if tabla.tableColumns:
            tabla.tableColumns.append(TableColumn(id=len(tabla.tableColumns) + 1, name=nombre))
        if tabla.autoFilter is not None:
            tabla.autoFilter.ref = tabla.ref


def poner_columnas_extra(wb, packs):
    """Las cuatro columnas propias, al final de «Análisis», por ASIN."""
    from openpyxl.styles import Font
    from openpyxl.utils import get_column_letter
    ws = wb['Análisis']
    cab = [c.value for c in ws[1]]
    i_asin = cab.index('ASIN') + 1
    por_asin = {p['c']['asin']: p for p in packs}
    for nombre, ancho in zip(COLUMNAS_EXTRA, ANCHOS_EXTRA):
        col = ws.max_column + 1
        ws.cell(row=1, column=col, value=nombre).font = Font(bold=True)
        ws.column_dimensions[get_column_letter(col)].width = ancho
        for fila in range(2, ws.max_row + 1):
            p = por_asin[str(ws.cell(row=fila, column=i_asin).value)]
            valor = {'Unidades del pack': p['n_osma'],
                     'Señales del pack': '%s%s%s' % ('pack de %d' % p['n'] if p['estado'] == eo.PACK_SI else eo.TEXTO_POSIBLE_PACK,
                                                     ' · ' + p['senales'], (' · ' + p['c']['porque']) if p['c'].get('porque') else ''),
                     'ASIN suelto de referencia': p['fam']['asin_ref'],
                     'Puesto de ventas de Amazon': ('%d · %s' % (p['c']['rank'], p['c'].get('rank_categoria') or '')).rstrip(' ·') if p['c'].get('rank') else None}[nombre]
            ws.cell(row=fila, column=col, value=valor)
        _alargar_tabla(ws, col, nombre)


def escribir_excel(packs, cuentas, M, info):
    """El Excel de packs → bytes. Hoja «Análisis» (las columnas del Análisis de OSMA, solo ES, con la Celda 9 del viejo
    y las mismas columnas añadidas, mas las cuatro propias) y hoja «Resumen» (el pedido y el porte en dos casillas
    amarillas vivas, como el AM). Las demas hojas de la Celda 9 (vacias aqui) no se dejan. Sin linea final de control."""
    from openpyxl.styles import Font, PatternFill
    foto, resultados = foto_y_resultados(packs)
    datos = e2.datos_como_el_viejo(foto, resultados, [], M)
    datos['PROVEEDOR'] = eo.PROVEEDOR
    wb = e2.escribir_celda9(datos, M, paises=[PAIS])
    e2.poner_columna_ficha_compartida(wb, foto, resultados)
    eo.poner_columna_no_habra_mas(wb, foto)
    poner_pedido_vivo_packs(wb, foto, resultados)
    poner_columnas_extra(wb, packs)
    for nombre in list(wb.sheetnames):
        if nombre != 'Análisis':
            del wb[nombre]
    an = wb['Análisis']
    cab = [c.value for c in an[1]]
    from openpyxl.utils import get_column_letter
    l_pais, l_dec = (get_column_letter(cab.index(n) + 1) for n in ('País', 'Decisión'))
    po = info.get('porte') or {}
    vivas = {d: "=COUNTIFS('Análisis'!$%s:$%s,\"%s\",'Análisis'!$%s:$%s,\"%s\")" % (l_pais, l_pais, PAIS, l_dec, l_dec, d)
             for d in ('COMPRAR', 'VALORAR')}
    n_dec = {d: sum(1 for p in packs if p['decision'] == d) for d in ('COMPRAR', 'VALORAR')}
    filas = [
        ['Pedido previsto (€)', po.get('pedido_previsto')],
        ['Porte del envío (€)', po.get('gastos_envio')],
        ['Porte en el precio', '=IF(N(%s)>0,N(%s)/%s,0)' % (eo.NOMBRE_PEDIDO, eo.NOMBRE_PORTE, eo.NOMBRE_PEDIDO)],
        ['Nota', 'Al cambiar estas dos celdas cambian «PA (€)», el beneficio y la «Decisión» de «Análisis», y estos recuentos.'],
        ['Recuento en ES', 'con el pedido de arriba (fórmula)', 'al valorar'],
        ['COMPRAR', vivas['COMPRAR'], n_dec['COMPRAR']],
        ['VALORAR', vivas['VALORAR'], n_dec['VALORAR']],
        ['Qué es', 'Packs de Amazon España que existen y NO vendemos, de las familias del cruce PRO de OSMA (puertas d, e y f '
                   'y nuestras fichas). Solo una PROPUESTA: ninguno se pega a una ficha.'],
        ['Cruce PRO de OSMA', info['cruce']],
        ['Pasada', info['pasada']],
        ['Familias en la cola', cuentas['familias']],
        ['Buscadas en Amazon · sin respuesta · aún esperando', '%d · %d · %d' % (cuentas['buscadas'], cuentas['fallidas'], cuentas['pendientes'])],
        ['Candidatos de la misma marca (mismo o parecido)', cuentas['candidatos']],
        ['Packs en «Análisis»', len(packs)],
        ['Candidatos que no son pack (llevan lo mismo que el suelto, o ninguna señal)', cuentas['no_pack']],
        ['No se venden (Keepa, %s)' % e2.texto_corte(info['umbral']),
         ' | '.join('%s · %s (%s caídas)' % (cod, a, e2._cifra(n)) for cod, a, n in cuentas['no_se_vende']) or 'ninguno'],
        ['Sin ventas de Keepa (VALORAR como mucho)', cuentas['sin_ventas']],
        ['Sin precio o sin tarifa de Amazon (Sin datos)', cuentas['sin_precio']],
        ['Packs NUESTROS encontrados (excluidos)', ' | '.join('%s · %s' % x for x in cuentas['nuestros']) or 'ninguno'],
        ['Regla del pack', 'se multiplica solo con dos señales distintas que dicen el mismo N (paquete, recuento y título, '
                           'contra el suelto de referencia); con una o contradictorias, posible pack (el N mayor) y nunca COMPRAR'],
        ['Coste', 'unidades de OSMA del pack × el PA de una unidad con porte ya redondeado'],
        ['Porte en el precio (el de la pasada)', ('%s € ÷ pedido previsto %s €' % (eo._coma(po.get('gastos_envio')), eo._coma(po.get('pedido_previsto'))))
         if po.get('gastos_envio') is not None and po.get('pedido_previsto') else 'sin porte'],
        ['Ventas', 'caídas de 30 días de Keepa en España guardadas hace %d días o menos (novedades o CSV del PRO); '
                   'los EAN de los packs entran en la lista del próximo «Barrer OSMA»' % DIAS_VENTAS],
    ]
    filas += [['Aviso', a] for a in info.get('avisos') or []]
    wr = wb.create_sheet('Resumen')
    wr.append(['Qué', 'Valor'])
    for c in wr[1]:
        c.font = Font(bold=True)
    for f in filas:
        wr.append(f)
    for letra, ancho in {'A': 52, 'B': 100, 'C': 14}.items():
        wr.column_dimensions[letra].width = ancho
    wr.freeze_panes = 'A2'
    amarillo = PatternFill(start_color='FFFF00', end_color='FFFF00', fill_type='solid')
    for celda in (eo.CELDA_PEDIDO, eo.CELDA_PORTE):
        wr[celda.replace('$', '')].fill = amarillo
        wr[celda.replace('$', '')].number_format = '#,##0.00'
        wr['A' + celda.split('$')[-1]].font = Font(bold=True)
    wr['B4'].number_format = '0.00%'
    eo.poner_celdas_pedido(wb, wr)
    salida = io.BytesIO()
    wb.save(salida)
    return salida.getvalue()


def ruta_excel(pasada, excel_id, sello):
    """`osma/<pasada>/<id de su fila de osma_packs_excel>/Escaner2_OSMA_Packs_<AAAAMMDD_HHMM>.xlsx`: la forma que exige
    `osma_packs_excel.ruta_excel` y la papelera de la biblioteca (la de un cruce, con el id de ESTA fila)."""
    return '%s/%s/%s/Escaner2_OSMA_Packs_%s.xlsx' % (eo.CARPETA, pasada, excel_id, sello)


def mensaje_telegram(packs, cuentas):
    """El aviso al dejar el Excel. Sale siempre (aunque no haya COMPRAR). Puro."""
    import html
    comprar = [p for p in packs if p['decision'] == 'COMPRAR']
    valorar_ = [p for p in packs if p['decision'] == 'VALORAR']
    lineas = ['%s <b>Packs de OSMA en Amazon</b>: %d que no vendemos · %d para COMPRAR · %d para VALORAR'
              % ('🟢' if comprar else '⚪', len(packs), len(comprar), len(valorar_))]
    for p in sorted(comprar, key=lambda x: -(x['calc'].get('margen') or 0))[:10]:
        lineas.append('· %s · %s · margen %s' % (html.escape(p['fam']['nombre_osma'][:50]), p['c']['asin'], e2._pct(p['calc'].get('margen'))))
    if cuentas['fallidas'] or cuentas['pendientes']:
        lineas.append('⚠️ %d familia(s) sin respuesta de Amazon y %d aún esperando' % (cuentas['fallidas'], cuentas['pendientes']))
    lineas.append('El Excel está en la biblioteca de escaneos.')
    return '\n'.join(lineas)


# ═══════════════════════════════════════════════════════════════════════════════
# 5 · LA LISTA DEL PROXIMO «BARRER OSMA»: los EAN de los packs encontrados
# ═══════════════════════════════════════════════════════════════════════════════
def eans_de_packs(candidatos, lista, M):
    """Los EAN de los packs encontrados («mismo» y «parecido») que la lista del Visualizador no trae ya → (nuevos,
    sin_ean). Van DETRAS de la lista, sin repetir y como los compara Keepa (sin ceros delante). Un pack que solo lista
    el EAN de la unidad ya viene con ella. Puro."""
    vistos = {M.norm(c) for c in lista}
    nuevos, sin_ean = [], 0
    for c in candidatos:
        if c.get('clase') not in ('mismo', 'parecido'):
            continue
        propios = [e for e in (c.get('eans') or []) if M.norm(e) and M.norm(e) not in vistos]
        if not (c.get('eans') or []):
            sin_ean += 1
        for e in propios:
            vistos.add(M.norm(e))
            nuevos.append(str(e))
    return nuevos, sin_ean


def recortar_al_tope(lista, eans_packs, tope):
    """(AN-2) Los EAN de packs que CABEN detras de la lista sin pasar del tope del Visualizador → (los que caben, cuantos
    se recortan). Se recortan los de PACK, nunca los de OSMA (la lista ya hecha no se toca). Puro."""
    cabe = max(0, tope - len(lista))
    return list(eans_packs)[:cabe], max(0, len(eans_packs) - cabe)


# ═══════════════════════════════════════════════════════════════════════════════
# 6 · LOS DOS PROGRAMAS (con la base inyectada: el test pasa un doble)
# ═══════════════════════════════════════════════════════════════════════════════
TROZO_IDS = 100


def _todas(sb, tabla, columnas, orden, filtros=(), en=None):
    """Todas las filas, paginando de 1.000 en 1.000. (AN-2) Con `en` = (columna, ids), los ids van en TROZOS de
    TROZO_IDS (una URL con cientos de uuid no cabe) y cada trozo se pagina entero."""
    if en:
        salida = []
        for i in range(0, len(en[1]), TROZO_IDS):
            salida += _paginar(sb, tabla, columnas, orden, filtros, (en[0], en[1][i:i + TROZO_IDS]))
        return salida
    return _paginar(sb, tabla, columnas, orden, filtros, None)


def _paginar(sb, tabla, columnas, orden, filtros, en):
    salida, desde = [], 0
    while True:
        q = sb.table(tabla).select(columnas)
        for op, col, v in filtros:
            q = getattr(q, op)(col, v)
        if en:
            q = q.in_(en[0], en[1])
        pagina = q.order(orden).range(desde, desde + 999).execute().data or []
        salida.extend(pagina)
        if len(pagina) < 1000:
            return salida
        desde += 1000


def _csv_de_la_pasada(sb, pasada, M):
    """El CSV de España de la pasada (el mas nuevo): (ruta local, fecha de subida) o (None, None)."""
    from foto_comun import descargar_buzon, listar_buzon
    carpeta = '%s/%s/csv' % (eo.CARPETA, pasada)
    objetos = sorted([o for o in listar_buzon(sb, BUCKET, carpeta) if (o.get('name') or '').lower().endswith(('.csv', '.csv.gz'))],
                     key=lambda x: x['name'], reverse=True)
    if not objetos:
        return None, None
    o = objetos[0]
    datos = descargar_buzon(sb, BUCKET, '%s/%s' % (carpeta, o['name']))
    if datos[:2] == b'\x1f\x8b':
        datos = gzip.decompress(datos)
    ruta = os.path.join(tempfile.mkdtemp(prefix='osma_packs_'), 'es.csv')
    with open(ruta, 'wb') as fh:
        fh.write(datos)
    return ruta, o.get('created_at') or o.get('updated_at')


# (AN-2) Lo que exige `osma_packs_cola` (sus CHECK), mirado ANTES de insertar: una fila que no lo cumple se salta y se
# cuenta, en vez de tumbar la cola entera en la base.
_RE_CODIGO = re.compile(r'^[A-Za-z0-9._-]{1,40}$')
_RE_ASIN = re.compile(r'^[A-Z0-9]{10}$')
_RE_EAN = re.compile(r'^[0-9]{1,14}$')


def motivo_fila_invalida(f):
    """None si la fila cumple los CHECK de osma_packs_cola; si no, el motivo (una palabra, para contar). Puro."""
    def positivo(x, tope):
        try:
            v = Decimal(str(x))
        except Exception:
            return False
        return v > 0 and v < tope
    if not _RE_CODIGO.match(str(f.get('codigo_osma') or '')):
        return 'codigo'
    if not str(f.get('nombre_osma') or '').strip():
        return 'nombre'
    eans = f.get('eans') or []
    if not (1 <= len(eans) <= 4) or not all(_RE_EAN.match(str(e)) for e in eans):
        return 'eans'
    if not _RE_ASIN.match(str(f.get('asin_ref') or '')):
        return 'asin_ref'
    if f.get('factor_ref') is not None and not (isinstance(f['factor_ref'], int) and f['factor_ref'] >= 1):
        return 'factor_ref'
    if not str(f.get('marca') or '').strip() or f.get('marca_origen') not in ('csv', 'amz_ficha', 'keepa_escaparate'):
        return 'marca'
    if not str(f.get('titulo_ref') or '').strip() or f.get('titulo_origen') not in ('csv', 'amz_ficha', 'keepa_escaparate', 'productos'):
        return 'titulo'
    if not positivo(f.get('pa_unidad'), Decimal('1e8')):
        return 'pa_unidad'
    if not positivo(f.get('price_net'), Decimal('1e7')):
        return 'price_net'
    if f.get('puerta') not in PUERTAS_COLA + (None,) or not (f.get('puerta') or f.get('nuestra')):
        return 'puerta'
    return None


def encolar(sb, pasada, run_id, imprimir=print):
    """Paso 1: la cola de UN cruce de OSMA (el de este run). Devuelve cuantas familias se han apuntado ahora.
    (AN-2) La cola de un cruce se apunta ENTERA O NADA: UNA sola insercion (una sentencia, una transaccion en la base),
    con las filas que cumplen los CHECK de la tabla (las que no, se saltan y se cuentan). Si el cruce ya tiene una cola
    incompleta (de antes de esto), se COMPLETA con las que le faltan."""
    cruces = sb.table('escaner2_cruce').select('id,estado,pasada_id,run_id').eq('pasada_id', pasada).eq('estado', 'lista') \
        .order('creado_en', desc=True).execute().data or []
    mio = [c for c in cruces if run_id and str(c.get('run_id')) == str(run_id)] or ([] if run_id else cruces[:1])
    if not mio:
        raise FalloPacks('no hay un cruce «lista» de esta pasada en este run')
    cruce = mio[0]['id']
    pas = (sb.table('escaner2_pasada').select('proveedor').eq('id', pasada).limit(1).execute().data or [{}])[0]
    if pas.get('proveedor') != eo.PROVEEDOR:
        raise FalloPacks('la pasada no es de OSMA: HEO no encola nada')
    ya = {x['codigo_osma'] for x in _todas(sb, 'osma_packs_cola', 'codigo_osma', 'codigo_osma', [('eq', 'cruce_id', cruce)])}
    M = e2.cargar_motor()
    barrido = json.loads(_descargar(sb, '%s/%s/barrido.json' % (eo.CARPETA, pasada)).decode('utf-8'))
    foto = _todas(sb, 'escaner2_foto', 'id,producto_heo,ean_original,ean_core,nombre,marca,precio_unidad,precio_catalogo,fin_de_vida',
                  'id', [('eq', 'pasada_id', pasada)])
    resultados = _todas(sb, 'escaner2_resultado_ean', 'foto_id,puerta,asin', 'id', [('eq', 'cruce_id', cruce)])
    productos = _todas(sb, 'productos', 'id,asin,ean,nombre,activo,es_chase,unidades_por_pack', 'id')
    ruta, _subido = _csv_de_la_pasada(sb, pasada, M)
    csv_titulos, senales = senales_csv_de(ruta) if ruta else ({}, {})
    asins = sorted({r['asin'] for r in resultados if r.get('asin')} | {p['asin'] for p in productos if p.get('asin')})
    amz, keepa = {}, {}
    for i in range(0, len(asins), TROZO_IDS):
        lote = asins[i:i + TROZO_IDS]
        for x in _paginar(sb, 'amz_ficha', 'asin,titulo,marca', 'asin', [('eq', 'pais', PAIS)], ('asin', lote)):
            amz[x['asin']] = {'titulo': x.get('titulo'), 'marca': x.get('marca')}
        for x in _paginar(sb, 'keepa_escaparate', 'asin,titulo,marca', 'asin', [('eq', 'dominio', 'es')], ('asin', lote)):
            keepa[x['asin']] = {'titulo': x.get('titulo'), 'marca': x.get('marca')}
    filas, cuentas = familias_para_cola(cruce, foto, resultados, barrido.get('enlaces') or {}, productos, csv_titulos, senales, amz, keepa)
    invalidas = {}
    validas = []
    for f in filas:
        m = motivo_fila_invalida(f)
        if m:
            invalidas[m] = invalidas.get(m, 0) + 1
        else:
            validas.append(f)
    nuevas = [f for f in validas if f['codigo_osma'] not in ya]
    if nuevas:
        # 🔑 UNA sola llamada: PostgREST la convierte en UNA sentencia INSERT (todas o ninguna).
        sb.table('osma_packs_cola').insert(nuevas).execute()
    en_bd = sb.table('osma_packs_cola').select('id', count='exact').eq('cruce_id', cruce).limit(1).execute().count
    imprimir('>>> PACKS: familias de este cruce %d (de las puertas d/e/f %d · nuestras %d; una puede ser las dos) · ya estaban %d · '
             'apuntadas ahora %d · en la base %s · saltadas por no cumplir la tabla %d %s · fuera: sin marca en Amazon %d, '
             'sin título %d, sin precio %d, sin ASIN %d'
             % (len(filas), cuentas['puertas'], cuentas['nuestras'], len(ya), len(nuevas), en_bd, sum(invalidas.values()),
                ' '.join('%s=%d' % kv for kv in sorted(invalidas.items())), cuentas['sin_marca'], cuentas['sin_titulo'],
                cuentas['sin_precio'], cuentas['sin_asin']), flush=True)
    if en_bd != len(ya) + len(nuevas):
        raise FalloPacks('la cola no quedó entera en la base')
    return len(nuevas)


def _descargar(sb, ruta):
    from foto_comun import descargar_buzon
    return descargar_buzon(sb, BUCKET, ruta)


# ═══════════════════════════════════════════════════════════════════════════════
# 7 · (AN-2) QUÉ CRUCE SE VALORA SIN BOTON, Y LAS VENTAS DE KEEPA POR ASIN
# ═══════════════════════════════════════════════════════════════════════════════
CRUCES_A_MIRAR = 10


def cruce_a_valorar(sb):
    """(AN-2) El reloj de y 40 lanza la valoracion SIN cruce: el ULTIMO cruce de OSMA «lista» con cola, ninguna familia
    pendiente y SIN NINGUN Excel de packs (tampoco uno quitado: si Fernando o Elena lo quitaron, el reloj no lo vuelve a
    hacer; el boton si puede). None si no hay ninguno."""
    pasadas = sb.table('escaner2_pasada').select('id').eq('proveedor', eo.PROVEEDOR).order('creado_en', desc=True) \
        .limit(CRUCES_A_MIRAR).execute().data or []
    if not pasadas:
        return None
    cruces = sb.table('escaner2_cruce').select('id,creado_en').eq('estado', 'lista').in_('pasada_id', [p['id'] for p in pasadas]) \
        .order('creado_en', desc=True).limit(CRUCES_A_MIRAR).execute().data or []
    for c in cruces:
        cola = _todas(sb, 'osma_packs_cola', 'estado', 'id', [('eq', 'cruce_id', c['id'])])
        if not cola or any(f['estado'] == 'pendiente' for f in cola):
            continue
        if sb.table('osma_packs_excel').select('id').eq('cruce_id', c['id']).limit(1).execute().data:
            continue
        return c['id']
    return None


def filas_keepa(productos, cruce, ahora):
    """Los productos de Keepa (por ASIN, España) → las filas de `osma_packs_keepa`, con la MISMA lectura que las novedades
    (`escaner2_novedades.ficha_de_keepa`: caidas de 30 dias, puesto actual y de 90 dias). Puro."""
    from escaner2_novedades import ficha_de_keepa
    filas = []
    for p in productos or []:
        fi = ficha_de_keepa(p)
        if not fi.get('asin') or not _RE_ASIN.match(str(fi['asin'])):
            continue
        filas.append({'cruce_id': cruce, 'asin': fi['asin'], 'pais': PAIS, 'caidas_30d': fi['caidas_30d'], 'rank': fi['rank'],
                      'rank_90d': fi['rank_90d'], 'titulo': (fi.get('titulo') or '')[:500] or None, 'consultada_en': ahora.isoformat()})
    return filas


def ventas_por_asin(sb, cruce, asins, ahora, env=None, keepa=None, imprimir=print):
    """(AN-2) Las ventas de Keepa de los packs que no tienen otras guardadas: por ASIN, en España, con la via de las
    novedades (stats=90, sin historial: ~1 token por ASIN) y la llave KEEPA_API_KEY de la v1; se GUARDAN en
    osma_packs_keepa (pelicula). Respeta la reserva de tokens de las novedades (`nov_parametros.keepa_reserva` de HEO, la
    misma llave): si no cabe, no pide. → (filas guardadas, tokens gastados, aviso o None). Nunca lanza."""
    if not asins:
        return [], 0, None
    env = os.environ if env is None else env
    k = keepa
    try:
        if k is None:
            from escaner2_novedades import Keepa
            k = Keepa(env.get('KEEPA_API_KEY'))
        par =(sb.table('nov_parametros').select('keepa_reserva,keepa_tope_peticion').eq('proveedor', 'HEO').limit(1)
               .execute().data or [{}])[0]
        reserva = int(par.get('keepa_reserva') or 0)
        tope = max(1, min(100, int(par.get('keepa_tope_peticion') or 100)))
        saldo = k.leer_saldo()
        if saldo - len(asins) < reserva:
            return [], k.tokens, ('Keepa: no se piden las ventas de %d pack(s): con %d tokens se bajaría de la reserva de %d'
                                  % (len(asins), saldo, reserva))
        productos = []
        for i in range(0, len(asins), tope):
            productos += k.productos(PAIS, asins[i:i + tope], por='asin')
        filas = filas_keepa(productos, cruce, ahora)
        if filas:
            sb.table('osma_packs_keepa').insert(filas).execute()
        return filas, k.tokens, None
    except Exception as ex:
        return [], getattr(k, 'tokens', 0) or 0, \
            'Keepa no respondió (%s): esos packs siguen sin ventas (VALORAR como mucho)' % type(ex).__name__


def valorar_cruce(sb, cruce, run_id, ahora=None, imprimir=print, enviar=None, env=None, keepa=None):
    """Paso 3: la valoracion de UN cruce, su Excel en la biblioteca y el Telegram. Devuelve la fila de osma_packs_excel.
    (AN-2) Sin `cruce` (el reloj de y 40), el de `cruce_a_valorar`; si no hay ninguno, no hace nada (sin Telegram)."""
    ahora = ahora or datetime.now(timezone.utc)
    if not cruce:
        cruce = cruce_a_valorar(sb)
        if not cruce:
            imprimir('>>> PACKS: ningún cruce de OSMA con su cola entera buscada y sin Excel de packs: nada que valorar.', flush=True)
            return None
    c = (sb.table('escaner2_cruce').select('id,estado,pasada_id').eq('id', cruce).limit(1).execute().data or [None])[0]
    if not c or c.get('estado') != 'lista':
        raise FalloPacks('el cruce no existe o no está «lista»')
    pasada = c['pasada_id']
    pas = (sb.table('escaner2_pasada').select('proveedor').eq('id', pasada).limit(1).execute().data or [{}])[0]
    if pas.get('proveedor') != eo.PROVEEDOR:
        raise FalloPacks('el cruce no es de OSMA')
    ya = sb.table('osma_packs_excel').select('id').eq('cruce_id', cruce).is_('quitado_biblioteca_en', 'null').execute().data or []
    if ya:
        imprimir('>>> PACKS: este cruce ya tiene su Excel en la biblioteca; no se hace otro.', flush=True)
        return None
    cola = _todas(sb, 'osma_packs_cola', 'id,codigo_osma,nombre_osma,eans,asin_ref,factor_ref,marca,pa_unidad,price_net,puerta,'
                  'nuestra,fin_de_vida,estado', 'id', [('eq', 'cruce_id', cruce)])
    if not cola:
        raise FalloPacks('este cruce no tiene cola de packs')
    if any(f['estado'] == 'pendiente' for f in cola):
        raise FalloPacks('aún hay familias esperando la búsqueda en Amazon')
    candidatos = _todas(sb, 'osma_packs_candidato', 'id,cola_id,asin,clase,porque,titulo,marca,eans,paquete,n_art,valor_ud,tipo_ud,'
                        'tamano,rank,rank_categoria,resultado,precio,canal,ref_eur,fee_fba,error', 'id',
                        en=('cola_id', [f['id'] for f in cola]))
    par = (sb.table('escaner2_parametros').select('umbral_caidas_30d').eq('proveedor', eo.PROVEEDOR).limit(1).execute().data or [None])[0]
    if not par:
        raise FalloPacks('no hay fila OSMA en escaner2_parametros (el corte de ventas)')
    umbral = int(par['umbral_caidas_30d'])
    productos = _todas(sb, 'productos', 'id,ean,asin,iva_pct,stock_moloka,es_chase,activo,unidades_por_pack', 'id')
    M = e2.cargar_motor()
    M.poner_catalogo_propio([p for p in productos if p.get('activo')])
    import en_mi_bd
    M.poner_foto_fba(en_mi_bd.leer_foto_fba(sb, imprimir=lambda linea: imprimir(linea, flush=True)))
    desde = (ahora - timedelta(days=DIAS_VENTAS)).isoformat()
    nov_keepa = _todas(sb, 'nov_keepa', 'pais,fichas,consultada_en', 'consultada_en', [('eq', 'pais', PAIS), ('gte', 'consultada_en', desde)])
    avisos, csv_caidas = [], None
    try:
        ruta, subido = _csv_de_la_pasada(sb, pasada, M)
        if ruta:
            col_pais, col_caidas = e2.columnas_keepa()
            import escaner2_heredado_pro as pro
            ex = e2.examinar_csv(ruta, pro, col_pais, col_caidas)
            if ex['pais'] == PAIS:
                csv_caidas = {'caidas': ex['caidas'], 'fecha': subido}
    except Exception as ex:
        avisos.append('No se pudo leer el CSV de la pasada para las ventas (%s): solo las de las novedades' % type(ex).__name__)
    candidatos_asin = sorted({x['asin'] for x in candidatos if x['clase'] in ('mismo', 'parecido')})
    por_asin = _todas(sb, 'osma_packs_keepa', 'asin,caidas_30d,consultada_en', 'consultada_en',
                      [('eq', 'pais', PAIS), ('gte', 'consultada_en', desde)], en=('asin', candidatos_asin))
    packs, cuentas = valorar(cola, candidatos, M, umbral, nov_keepa, csv_caidas, ahora, por_asin)
    # (AN-2) Los packs que siguen SIN ventas: se piden a Keepa por ASIN (y se guardan), y se valora otra vez.
    faltan = sorted({p['c']['asin'] for p in packs if p['caidas'] is None})
    nuevas, tokens, aviso_keepa = ventas_por_asin(sb, cruce, faltan, ahora, env=env, keepa=keepa, imprimir=imprimir)
    if aviso_keepa:
        avisos.append(aviso_keepa)
    if faltan:
        avisos.append('Ventas de Keepa por ASIN: %d pack(s) sin ventas guardadas · traídas %d (con caídas %d) · tokens gastados %d'
                      % (len(faltan), len(nuevas), sum(1 for f in nuevas if f['caidas_30d'] is not None), tokens))
    if nuevas:
        packs, cuentas = valorar(cola, candidatos, M, umbral, nov_keepa, csv_caidas, ahora, por_asin + nuevas)
    imprimir('>>> PACKS: ventas de Keepa por ASIN · sin ventas antes %d · traídas %d · tokens gastados %d%s'
             % (len(faltan), len(nuevas), tokens, ' · AVISO: Keepa no se pudo usar' if aviso_keepa else ''), flush=True)
    barrido = json.loads(_descargar(sb, '%s/%s/barrido.json' % (eo.CARPETA, pasada)).decode('utf-8'))
    from zoneinfo import ZoneInfo
    sello = ahora.astimezone(ZoneInfo('Europe/Madrid')).strftime('%Y%m%d_%H%M')
    excel_id = str(uuid.uuid4())
    ruta = ruta_excel(pasada, excel_id, sello)
    contenido = escribir_excel(packs, cuentas, M, {'cruce': cruce, 'pasada': pasada, 'porte': barrido.get('porte'),
                                                   'umbral': umbral, 'avisos': avisos})
    sb.storage.from_(BUCKET).upload(ruta, contenido, {
        'content-type': 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet', 'upsert': 'false'})
    fila = {'id': excel_id, 'cruce_id': cruce, 'ruta_excel': ruta, 'n_familias': cuentas['familias'], 'n_packs': len(packs),
            'n_comprar': sum(1 for p in packs if p['decision'] == 'COMPRAR'),
            'n_valorar': sum(1 for p in packs if p['decision'] == 'VALORAR'),
            'run_id': int(run_id) if str(run_id or '').isdigit() else None}
    sb.table('osma_packs_excel').insert(fila).execute()
    imprimir('>>> PACKS: familias %d · candidatos %d · packs %d (COMPRAR %d · VALORAR %d) · no son pack %d · no se venden %d · '
             'sin ventas %d · nuestros encontrados %d. Excel en la biblioteca.'
             % (cuentas['familias'], cuentas['candidatos'], len(packs), fila['n_comprar'], fila['n_valorar'], cuentas['no_pack'],
                len(cuentas['no_se_vende']), cuentas['sin_ventas'], len(cuentas['nuestros'])), flush=True)
    try:
        from escaner2_novedades import enviar_telegram
        (enviar or enviar_telegram)(mensaje_telegram(packs, cuentas), env, None, imprimir)
    except Exception as ex:
        imprimir('AVISO Telegram (no se envió, la corrida sigue): %s' % type(ex).__name__, flush=True)
    return fila


def main(argv=None):
    argv = sys.argv[1:] if argv is None else argv
    if len(argv) != 1 or argv[0] not in ('encolar', 'valorar'):
        print('ESCANER2_NO_EJECUTADO: uso: escaner2_osma_packs.py encolar|valorar')
        sys.exit(1)
    llave = os.environ.get('SUPABASE_SERVICE_KEY')
    if not llave:
        print('ESCANER2_NO_EJECUTADO: sin llave de servicio')
        sys.exit(1)
    from supabase import create_client
    sb = create_client(os.environ['SUPABASE_URL'], llave)
    run_id = os.environ.get('GITHUB_RUN_ID')
    uuid_ok = r'[0-9a-f]{8}-[0-9a-f]{4}-[0-9a-f]{4}-[0-9a-f]{4}-[0-9a-f]{12}'
    try:
        if argv[0] == 'encolar':
            pasada = (os.environ.get('PASADA') or '').strip().lower()
            if not re.fullmatch(uuid_ok, pasada):
                raise FalloPacks('la pasada no es un id válido')
            encolar(sb, pasada, run_id)
        else:
            cruce = (os.environ.get('CRUCE') or '').strip().lower()
            # (AN-2) Sin cruce: el reloj de y 40 (busca el que toca; si no hay ninguno, sale en verde sin hacer nada).
            if cruce and not re.fullmatch(uuid_ok, cruce):
                raise FalloPacks('el cruce no es un id válido')
            valorar_cruce(sb, cruce, run_id)
    except FalloPacks as ex:
        print('ESCANER2_NO_EJECUTADO: %s' % ex)
        sys.exit(1)
    except Exception as ex:
        # 🔴 REPO PUBLICO: solo el tipo, nunca el texto (podria llevar EAN, ASIN o precios).
        print('ESCANER2_NO_EJECUTADO: %s (%s)' % (argv[0], type(ex).__name__))
        sys.exit(1)


if __name__ == '__main__':
    main()
