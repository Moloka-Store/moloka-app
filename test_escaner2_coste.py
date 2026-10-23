# -*- coding: utf-8 -*-
"""Banco del COSTE DE PREPARACION del escaner 2 (encargo RF4, 23-oct-2026).

Fernando, 23-oct-2026: «tiene que ser consistente, que en todos los sitios donde se vea la rentabilidad de un producto
se calcule con su coste correcto». Desde el RF4 el escaner 2 no resta ALMACEN = 0,15 a todo: resta el coste de cada
ficha de la fuente unica del RF2 (escaner2_motor.leer_coste_preparacion + coste_almacen). Aqui:
  (A) la lectura de la pasada: v_fijo_preparacion_asin por tandas, eur_kg de mv_coste_preparacion y el global VIGENTE
      de parametro_coste ese mes;
  (B) los tres caminos, en su orden: (a) el ASIN en la fuente (envios, peso, excepcion) → ese; (b) no esta, pero hay
      peso del paquete → eur_kg × peso / 1000 a 4 decimales, con el redondeo de Postgres (la mitad, hacia fuera) y la
      mediana de los paises; (c) nada → el global;
  (C) 🔴 la fuente AUSENTE (RF2 sin aplicar) o ilegible: la pasada sigue, todo con el global (o el respaldo), y avisa;
  (D) el COMPRAR/VALORAR de un caso conocido cambia SOLO por el almacen: la misma cuenta del viejo con otro numero;
  (E) el Excel: «Almacén (€)» de cada fila con su coste (la formula del beneficio resta esa casilla) y «Origen almacén»;
  (F) el peso del CSV del Visualizador (el desvio RF4 del lector heredado);
  (G) 🔴 nadie vuelve a usar el ALMACEN del viejo a pelo en el escaner 2 (y la guarda muerde);
  (H) el log publico no lleva ni el eur_kg ni el global ni un coste.
🔒 Repo PUBLICO: todos los ASIN, pesos y costes de aqui son INVENTADOS.
"""
import ast
import io
import os
import sys
import tempfile
from datetime import date

AQUI = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, AQUI)
os.chdir(AQUI)

import escaner2_motor as e2  # noqa: E402
import escaner2_heredado_pro as pro  # noqa: E402

M = e2.cargar_motor()
fallos = []


def eq(nombre, obtenido, esperado):
    ok = obtenido == esperado
    if not ok:
        fallos.append(nombre)
    print(('OK ' if ok else 'XX ') + nombre + ('' if ok else '   got=%r exp=%r' % (obtenido, esperado)))


# ── El doble de Supabase: tablas en memoria; una tabla en `rotas` contesta con error, como PostgREST ──────────
class _Resp:
    def __init__(self, data):
        self.data = data


class _Q:
    def __init__(self, sb, tabla):
        self.sb, self.tabla, self.filtros, self.rango, self.tope, self.orden = sb, tabla, [], None, None, None

    def select(self, cols):
        self.cols = cols.split(',')
        return self

    def eq(self, k, v):
        self.filtros.append((k, v))
        return self

    def order(self, col):
        self.orden = col
        return self

    def range(self, a, b):
        self.rango = (a, b)
        return self

    def limit(self, n):
        self.tope = n
        return self

    def execute(self):
        self.sb.lecturas.append(self.tabla)
        if self.tabla in self.sb.rotas:
            raise RuntimeError("{'code': 'PGRST205', 'message': \"Could not find the table 'public.%s'\"}" % self.tabla)
        filas = [f for f in self.sb.tablas.get(self.tabla, []) if all(f.get(k) == v for k, v in self.filtros)]
        if self.orden:
            filas = sorted(filas, key=lambda f: f[self.orden])
        if self.rango:
            filas = filas[self.rango[0]:self.rango[1] + 1]
        if self.tope is not None:
            filas = filas[:self.tope]
        return _Resp([{c: f.get(c) for c in self.cols} for f in filas])


class SbFalso:
    def __init__(self, tablas, rotas=()):
        self.tablas, self.rotas, self.lecturas = tablas, set(rotas), []

    def table(self, nombre):
        return _Q(self, nombre)


# Inventado: tres ASIN con coste, uno de cada origen; eur_kg 0,7141; el global vigente cambia el 1-oct-2026.
VFA = [{'asin': 'B0RF4ENV01', 'valor': 0.96, 'origen': 'envios'},
       {'asin': 'B0RF4PES01', 'valor': 0.0357, 'origen': 'peso'},
       {'asin': 'B0RF4EXC01', 'valor': 0.50, 'origen': 'excepcion'}]
MV = [{'producto_id': 'p%d' % i, 'eur_kg': 0.7141, 'refrescada_el': '2026-10-23T05:10:00+00:00'} for i in range(4)]
PC = [{'clave': 'coste_almacen_ud', 'valor': 0.12, 'desde': '2026-01-01', 'hasta': '2026-10-01'},
      {'clave': 'coste_almacen_ud', 'valor': 0.15, 'desde': '2026-10-01', 'hasta': None},
      {'clave': 'otra_cosa', 'valor': 9.99, 'desde': '2026-01-01', 'hasta': None}]
HOY = date(2026, 10, 23)


def sb_entero(**kw):
    return SbFalso({'v_fijo_preparacion_asin': list(VFA), 'mv_coste_preparacion': list(MV), 'parametro_coste': list(PC)},
                   **kw)


def leer(sb, hoy=HOY):
    avisos, log = [], []
    f = e2.leer_coste_preparacion(sb, M.ALMACEN_GLOBAL_RESPALDO, hoy=hoy, avisos=avisos, imprimir=log.append)
    return f, avisos, log


# ═══════════════════════════════════════════════════════════════════════════════
print('(A) la lectura de la pasada')
F, AV, LOG = leer(sb_entero())
eq('(A) fuente leída: estado ok, los tres ASIN con su valor y origen, eur_kg y el global vigente en octubre',
   (F['estado'], F['por_asin'], F['eur_kg'], F['global'], F['global_de'], AV),
   ('ok', {'B0RF4ENV01': (0.96, 'envios'), 'B0RF4PES01': (0.0357, 'peso'), 'B0RF4EXC01': (0.5, 'excepcion')},
    0.7141, 0.15, 'parametro_coste', []))
eq('(A) el global es el VIGENTE ese mes: en septiembre, el de entonces (0,12)', leer(sb_entero(), date(2026, 9, 15))[0]['global'], 0.12)
_sb = sb_entero()
leer(_sb)
eq('(A) UNA lectura por pasada: el global, la fuente y la lista (una tanda), en este orden', _sb.lecturas,
   ['parametro_coste', 'mv_coste_preparacion', 'v_fijo_preparacion_asin'])
_tanda = e2.TANDA_COSTE
try:
    e2.TANDA_COSTE = 2
    _sb = SbFalso({'v_fijo_preparacion_asin': [{'asin': 'B0RF4T%04d' % i, 'valor': 0.01 * i, 'origen': 'envios'}
                                               for i in range(1, 6)],
                   'mv_coste_preparacion': list(MV), 'parametro_coste': list(PC)})
    _f, _a, _l = leer(_sb)
finally:
    e2.TANDA_COSTE = _tanda
eq('(A) 🔴 la lista se lee ENTERA por tandas (5 filas en tandas de 2: tres lecturas de la vista, ninguna fila perdida)',
   (len(_f['por_asin']), _sb.lecturas.count('v_fijo_preparacion_asin')), (5, 3))
_f, _a, _l = leer(SbFalso({'v_fijo_preparacion_asin': VFA + [{'asin': 'B0RF4RARO1', 'valor': 0.2, 'origen': 'inventado'}],
                           'mv_coste_preparacion': list(MV), 'parametro_coste': list(PC)}))
eq('(A) un origen que el escáner no conoce se GRITA (aviso), no se traga', (_f['estado'], len(_a), 'inventado' in _a[0]),
   ('ok', 1, True))
_f, _a, _l = leer(SbFalso({'v_fijo_preparacion_asin': list(VFA), 'mv_coste_preparacion': [dict(MV[0], eur_kg=None)],
                           'parametro_coste': list(PC)}))
eq('(A) fuente sin eur_kg (ningún envío repartido por peso): se lee igual, sin estimación por peso, y el log lo dice',
   (_f['estado'], _f['eur_kg'], any('NO (sin estimación por peso)' in x for x in _l)), ('ok', None, True))

# ═══════════════════════════════════════════════════════════════════════════════
print('\n(B) los tres caminos, en su orden')
eq('(B) (a) en la fuente por envíos → su valor', e2.coste_almacen(F, 'B0RF4ENV01', None, 0.15), (0.96, 'envios'))
eq('(B) (a) en la fuente por peso (amz_ficha) → su valor', e2.coste_almacen(F, 'B0RF4PES01', None, 0.15), (0.0357, 'peso'))
eq('(B) (a) con excepción vigente → la excepción', e2.coste_almacen(F, 'B0RF4EXC01', None, 0.15), (0.5, 'excepcion'))
eq('(B) 🔴 (a) manda sobre (b): en la fuente Y con peso en el CSV → el de la fuente', e2.coste_almacen(F, 'B0RF4ENV01', 500, 0.15),
   (0.96, 'envios'))
# 0,7141 × 500 / 1000 = 0,35705 exacto: Postgres (la mitad hacia fuera) da 0,3571; el round de Python, 0,357.
eq('(B) (b) fuera de la fuente, con peso → eur_kg × peso / 1000 a 4 decimales, como round(numeric) de Postgres',
   e2.coste_almacen(F, 'B0RF4NUEVO', 500, 0.15), (0.3571, 'peso'))
eq('(B) (b) …y con decimales de gramo (redondeo del peso a 1 decimal, como la mediana del RF2)',
   e2.coste_almacen(F, 'B0RF4NUEVO', 250.04, 0.15), (float(e2._redondear(0.7141 * 250.0 / 1000, 4)), 'peso'))
eq('(B) (c) fuera de la fuente y sin peso → el global vigente', e2.coste_almacen(F, 'B0RF4NUEVO', None, 0.15), (0.15, 'global'))
eq('(B) (c) peso 0 no es peso → el global', e2.coste_almacen(F, 'B0RF4NUEVO', 0, 0.15), (0.15, 'global'))
eq('(B) sin ASIN no está en la fuente: con peso, (b); sin peso, el global (el orden del encargo, literal)',
   (e2.coste_almacen(F, None, 500, 0.15), e2.coste_almacen(F, None, None, 0.15)), ((0.3571, 'peso'), (0.15, 'global')))
_sin_kg = dict(F, eur_kg=None)
eq('(B) (c) con peso pero sin eur_kg en la fuente → el global', e2.coste_almacen(_sin_kg, 'B0RF4NUEVO', 500, 0.15), (0.15, 'global'))
eq('(B) peso de la ficha: la MEDIANA de sus países (480, 500, 520 → 500; sin peso o 0, fuera)',
   e2.peso_de_fichas([{'peso': 480.0}, {'peso': 520.0}, {'peso': None}, {'peso': 0.0}, {'peso': 500.0}, {}]), 500.0)
eq('(B) …con dos países, la media de los dos (percentile_cont): 490 y 511 → 500,5',
   e2.peso_de_fichas([{'peso': 490.0}, {'peso': 511.0}]), 500.5)
eq('(B) …y sin ninguno, None', e2.peso_de_fichas([{'peso': None}, {}]), None)
M.poner_coste_preparacion(F)
eq('(B) el motor con la lectura puesta pregunta lo mismo', M.coste_almacen('B0RF4ENV01'), (0.96, 'envios'))
M.poner_coste_preparacion(None)
eq('(B) sin lectura puesta (un programa que no la lee): el respaldo, «global», como antes del RF4',
   M.coste_almacen('B0RF4ENV01', 500), (M.ALMACEN, 'global'))

# ═══════════════════════════════════════════════════════════════════════════════
print('\n(C) 🔴 la fuente ausente o ilegible: la pasada sigue con el global, y lo dice')
_f, _a, _l = leer(sb_entero(rotas=('mv_coste_preparacion',)))
eq('(C) RF2 sin aplicar (mv_coste_preparacion no existe): estado sin_fuente, NADA por ASIN, el global de parametro_coste',
   (_f['estado'], _f['por_asin'], _f['global'], _f['global_de']), ('sin_fuente', {}, 0.15, 'parametro_coste'))
eq('(C) …todo con el global, aunque el ASIN estuviera en la lista y hubiera peso',
   [e2.coste_almacen(_f, a, 500, 0.15) for a in ('B0RF4ENV01', 'B0RF4NUEVO')], [(0.15, 'global'), (0.15, 'global')])
eq('(C) …y lo dice: un aviso para el Excel y una línea «!!!» en el log',
   (len(_a), 'SIN la fuente única' in _a[0], 'PGRST205' in _a[0], sum(x.startswith('!!! COSTE_PREPARACION') for x in _l)),
   (1, True, True, 1))
_f, _a, _l = leer(sb_entero(rotas=('v_fijo_preparacion_asin',)))
eq('(C) la lista ilegible (la fuente sí): tampoco se usa a medias → sin_fuente', (_f['estado'], _f['por_asin']), ('sin_fuente', {}))
_f, _a, _l = leer(SbFalso({'v_fijo_preparacion_asin': list(VFA), 'mv_coste_preparacion': [], 'parametro_coste': list(PC)}))
eq('(C) la fuente vacía (0 filas: no se ha podido leer, no «nadie tiene coste») → sin_fuente', _f['estado'], 'sin_fuente')
_f, _a, _l = leer(SbFalso({}, rotas=('mv_coste_preparacion', 'parametro_coste')))
eq('(C) ni fuente ni parametro_coste: el respaldo del escáner (el ALMACEN del viejo, 0,15), con dos avisos',
   (_f['global'], _f['global_de'], len(_a), e2.coste_almacen(_f, 'B0RF4ENV01', 500, M.ALMACEN_GLOBAL_RESPALDO)),
   (0.15, 'respaldo', 2, (0.15, 'global')))
_f, _a, _l = leer(SbFalso({'v_fijo_preparacion_asin': list(VFA), 'mv_coste_preparacion': list(MV),
                           'parametro_coste': [dict(PC[0])]}))
eq('(C) parametro_coste sin global vigente este mes: el respaldo, avisado, y la fuente se usa igual',
   (_f['estado'], _f['global_de'], len(_a)), ('ok', 'respaldo', 1))
try:
    _r = e2.leer_coste_preparacion(None, 0.15, hoy=HOY, avisos=[], imprimir=lambda x: None)['estado']
except Exception as ex:   # pragma: no cover
    _r = 'LANZA %s' % type(ex).__name__
eq('(C) 🔴 NUNCA lanza (ni sin cliente)', _r, 'sin_fuente')

# ═══════════════════════════════════════════════════════════════════════════════
print('\n(D) el COMPRAR/VALORAR de un caso conocido cambia SOLO por el almacén')
# Inventado: precio 20, PA 7,50, comisión 15 %, FBA 3,50, España (IVA asumido 21 %).
REC = {'asin': 'B0RF4ENV01', 'buybox': 20.0, 'es_fba': True, 'compct': 15.0, 'fba': 3.5, 'titulo': 'x', 'rank': 1,
       'rank90': 1, 'peso': 500.0}
CORE = '8400000000017'
M.poner_coste_preparacion(None)
_antes = e2.calcular_pais('ES', REC, 7.5, CORE, M)
M.poner_coste_preparacion(F)
_ahora = e2.calcular_pais('ES', REC, 7.5, CORE, M)
_nuevo = e2.calcular_pais('ES', dict(REC, asin='B0RF4NUEVO'), 7.5, CORE, M)
M.poner_coste_preparacion(None)
eq('(D) antes del RF4 (sin fuente): 0,15 «global» y COMPRAR', (_antes['almacen'], _antes['almacen_origen'], _antes['decision']),
   (0.15, 'global', 'COMPRAR'))
eq('(D) con su coste por envíos (0,96): VALORAR', (_ahora['almacen'], _ahora['almacen_origen'], _ahora['decision']),
   (0.96, 'envios', 'VALORAR'))
_iva, _isd = M.iva_por_pais(CORE)['ES'], M.ISD_PAIS['ES']
for _x, _alm in ((_antes, 0.15), (_ahora, 0.96), (_nuevo, 0.3571)):
    _r = M.calc_rentabilidad(20.0, 7.5, 15.0, 3.5, _iva, almacen=_alm, com_digitales=M.COM_DIGITALES, isd=_isd)
    eq('(D) 🔒 la cuenta con %s es la del viejo con almacen=%s (beneficio, margen, decisión)' % (_x['almacen_origen'], _alm),
       (_x['beneficio'], _x['margen'], _x['decision']), (_r['beneficio'], _r['margen'], M.decision_de(_r['margen'])))
_igual = [k for k in _antes if k not in ('almacen', 'almacen_origen', 'beneficio', 'roi', 'margen', 'decision')
          and _antes[k] != _ahora[k]]
eq('(D) …y nada más cambia en la fila (precio, comisión, IVA, ISD, PA, com_amazon…)', _igual, [])
eq('(D) la diferencia de beneficio es exactamente la del almacén (0,96 − 0,15 = 0,81)',
   round(_antes['beneficio'] - _ahora['beneficio'], 10), 0.81)
eq('(D) fuera de la fuente, con 500 g: 0,3571 «peso», y sigue en COMPRAR', (_nuevo['almacen'], _nuevo['almacen_origen'],
                                                                           _nuevo['decision']), (0.3571, 'peso', 'COMPRAR'))
_sin = e2.calcular_pais('ES', dict(REC, fba=None), 7.5, CORE, M)
eq('(D) sin tarifa: «Sin datos», pero el almacén y su origen se guardan igual', (_sin['decision'], _sin['almacen_origen']),
   ('Sin datos', 'global'))

# El mismo caso por las puertas (lo que hacen los cruces): la mediana del peso de sus países entra en la cuenta.
PARAMS = {'umbral': 6, 'paises_filtro': ['ES'], 'paises_calculo': ['ES', 'IT', 'FR', 'DE']}
FOTO = {'id': 'f1', 'nombre': 'Funko Pop Prueba', 'ean_core': CORE, 'ean_original': CORE, 'precio_unidad': 7.5,
        'marca': 'Funko', 'aviso_caja': None}
_r_nuevo = dict(REC, asin='B0RF4NUEVO')
CANDS = {'ES': [dict(_r_nuevo, peso=480.0)], 'IT': [dict(_r_nuevo, peso=520.0)], 'DE': [dict(_r_nuevo, peso=None)]}
CAIDAS = {p: {'B0RF4NUEVO': 30} for p in CANDS}
M.poner_coste_preparacion(F)
RES = e2.decidir(FOTO, CANDS, CAIDAS, PARAMS, M)
eq('(D) por las puertas: el peso es la mediana de ES (480) e IT (520) = 500, y TODOS los países pagan lo mismo',
   sorted({(p, c['almacen'], c['almacen_origen']) for p, c in RES['paises'].items()}),
   [('DE', 0.3571, 'peso'), ('ES', 0.3571, 'peso'), ('IT', 0.3571, 'peso')])

# ═══════════════════════════════════════════════════════════════════════════════
print('\n(E) el Excel: «Almacén (€)» de cada fila con su coste y «Origen almacén»')
from openpyxl import load_workbook  # noqa: E402
RES2 = e2.decidir(dict(FOTO, id='f2', ean_core='8400000000024', ean_original='8400000000024'),
                  {'ES': [dict(REC)]}, {'ES': {'B0RF4ENV01': 30}}, PARAMS, M)
RES['foto_id'], RES2['foto_id'] = 'f1', 'f2'
FOTOS = [FOTO, dict(FOTO, id='f2', ean_core='8400000000024', ean_original='8400000000024')]
M.poner_catalogo_propio([])
_wb = e2.excel_como_el_viejo(FOTOS, [RES, RES2], [], M)
_buf = io.BytesIO()
_wb.save(_buf)
_ws = load_workbook(io.BytesIO(_buf.getvalue()))['Análisis']
_cab = [c.value for c in _ws[1]]
_ia = {h: i for i, h in enumerate(_cab)}
_filas = [[c.value for c in f] for f in _ws.iter_rows(min_row=2)]
eq('(E) «Origen almacén» detrás de «Origen IVA» (la última del viejo), antes de «Ficha compartida»',
   _cab[_cab.index('Origen IVA') + 1:], [e2.COLUMNA_ORIGEN_ALMACEN, e2.COLUMNA_FICHA_COMPARTIDA])
eq('(E) cada fila (los cuatro países de cada ficha, también los sin cálculo) con SU coste y su origen',
   sorted({(f[_ia['ASIN']], f[_ia['Almacén (€)']], f[_ia[e2.COLUMNA_ORIGEN_ALMACEN]]) for f in _filas}),
   [('B0RF4ENV01', 0.96, 'envíos'), ('B0RF4NUEVO', 0.3571, 'peso')])
eq('(E) …ocho filas: dos fichas por los cuatro países', len(_filas), 8)
_es = [f for f in _filas if f[_ia['ASIN']] == 'B0RF4ENV01' and f[_ia['País']] == 'ES'][0]
_L = {h: chr(ord('A') + i) for i, h in enumerate(_cab) if i < 26}
eq('(E) 🔒 la fórmula del beneficio es la del viejo y resta la casilla «Almacén (€)» de su fila',
   ('-%s' % _L['Almacén (€)']) in str(_es[_ia['Beneficio (€)']]), True)
from openpyxl.utils import get_column_letter  # noqa: E402
eq('(E) la tabla T_Analisis cubre hasta la última columna', [t.ref.split(':')[1].rstrip('0123456789') for t in _ws.tables.values()],
   [get_column_letter(len(_cab))])
_n = e2.poner_almacen_por_fila(load_workbook(io.BytesIO(_buf.getvalue())), [{'asin': 'B0RF4ENV01', '_almacen': None}])
eq('(E) un registro sin su coste no inventa nada: su fila se queda como la escribió el viejo', _n, {})
M.poner_coste_preparacion(None)

# ═══════════════════════════════════════════════════════════════════════════════
print('\n(F) el peso del CSV del Visualizador (desvío RF4 del lector heredado)')
_d = tempfile.mkdtemp(prefix='e2coste_')


def _csv(nombre, cab, fila):
    ruta = os.path.join(_d, nombre)
    with io.open(ruta, 'w', encoding='utf-8', newline='') as fh:
        fh.write(','.join('"%s"' % c for c in cab) + '\n' + ','.join('"%s"' % c for c in fila) + '\n')
    return ruta


_C = pro.CSV_COLS
_cab = [_C['ean'], _C['asin'], _C['buybox'], _C['fba'], _C['compct'], 'Paquete: Peso (g)']
_rec = pro.leer_csv_visualizador(_csv('con.csv', _cab, ['8400000000017', 'B0RF4NUEVO', '20', '3.5', '15', '470']))
eq('(F) «Paquete: Peso (g)» → peso en gramos', _rec['8400000000017'][0]['peso'], 470.0)
_rec = pro.leer_csv_visualizador(_csv('sin.csv', _cab[:-1], ['8400000000017', 'B0RF4NUEVO', '20', '3.5', '15']))
eq('(F) un CSV sin esa columna: peso None (no 0)', _rec['8400000000017'][0]['peso'], None)
eq('(F) el desvío dice su nombre en el lector heredado', sum('DESVÍO RF4' in l for l in io.open(
    'escaner2_heredado_pro.py', encoding='utf-8')), 2)

# ═══════════════════════════════════════════════════════════════════════════════
print('\n(G) 🔴 nadie usa el ALMACEN del viejo a pelo en el escaner 2')


def usos_de_almacen(textos):
    """[(fichero, linea)] de cada `<algo>.ALMACEN` (atributo) en el codigo del escaner 2. Por estructura, con ast."""
    salida = []
    for nombre, texto in textos.items():
        for n in ast.walk(ast.parse(texto)):
            if isinstance(n, ast.Attribute) and n.attr == 'ALMACEN':
                salida.append((nombre, n.lineno))
    return salida


_prog = {f: io.open(f, encoding='utf-8').read() for f in sorted(os.listdir('.'))
         if f.startswith('escaner2_') and f.endswith('.py') and not f.startswith('escaner2_heredado')}
eq('(G) los programas del escáner 2 que se miran (más de 40, el motor entre ellos)',
   (len(_prog) > 40, 'escaner2_motor.py' in _prog, 'escaner2_factura.py' in _prog), (True, True, True))
eq('(G) 🔴 ningún `M.ALMACEN` en el escáner 2: el almacén de una cuenta sale de M.coste_almacen', usos_de_almacen(_prog), [])
_roto = dict(_prog, **{'escaner2_factura.py': _prog['escaner2_factura.py'] + '\nalm = M.ALMACEN\n'})
eq('(G) …y la guarda muerde: un `M.ALMACEN` metido en una copia, cazado con su fichero',
   [f for f, _l in usos_de_almacen(_roto)], ['escaner2_factura.py'])
eq('(G) el respaldo es el ALMACEN del viejo, con su nombre', (M.ALMACEN_GLOBAL_RESPALDO, M.ALMACEN), (0.15, 0.15))

# ═══════════════════════════════════════════════════════════════════════════════
print('\n(H) 🔒 el log es PUBLICO: ni el eur_kg, ni el global, ni un coste')
_todo = ' '.join(LOG + leer(sb_entero(rotas=('mv_coste_preparacion',)))[2])
eq('(H) el log de la lectura dice recuentos y orígenes', any('3 ASIN (envios 1, excepcion 1, peso 1)' in x for x in LOG), True)
eq('(H) 🔴 …y ninguna cifra de coste', [c for c in ('0.7141', '0,7141', '0.15', '0,15', '0.96', '0,96', '0.12') if c in _todo], [])

print()
if fallos:
    print('FALLOS (%d): %s' % (len(fallos), ' | '.join(fallos)))
    sys.exit(1)
print('TODO OK')
