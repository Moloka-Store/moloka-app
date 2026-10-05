# -*- coding: utf-8 -*-
"""ESCANER 2 · ESCANEO PRO DE ZENTRADA · LA PARTE PURA (encargo de la tarjeta de Zentrada, 05-oct-2026)

QUE ES. El calculo del Escaneo PRO de Zentrada, sin red y sin base: leer el Excel del dia (con sus guardas), la foto,
las dos listas para el Visualizador, el cruce de cada EAN con sus fichas de Keepa y el Excel. La red vive en los dos
programas que lo usan:
  · escaner2_zentrada_barrido.py (workflow escaner2-zentrada-barrido.yml): del Excel a la foto y las listas.
  · escaner2_zentrada_cruce.py   (workflow escaner2-zentrada-cruce.yml):   los CSV del Visualizador y el Excel.

🔑 EL CATALOGO NO LO BAJA LA APP (Fernando, 05-oct-2026: «eso ya hemos probado en el pasado y no funciona, lo
   bloquean, lo unico que nos funciona es cuando te metes tu»). Lo lee Cowork con el navegador de la app y la sesion
   de Moloka y deja un Excel del dia (hojas «Ofertas», «Fichas» y «Resumen»); Fernando lo suelta en la tarjeta de
   Zentrada de la v2, que lo guarda en `escaner2/zentrada/<pasada>/` y lanza el barrido. NINGUN programa llama a
   zentrada.com.
🔑 CADA ESCANER, A LA MEDIDA DE SU PROVEEDOR (Fernando, 29-sep-2026): nada de tronco comun con adaptadores. Leer el
   Excel, filtrar y cruzar es codigo de Zentrada y vive aqui (lo que se parece a OSMA se ha copiado y adaptado, no se
   importa de escaner2_osma.py: un cambio en OSMA no mueve Zentrada). Lo que se REUTILIZA del escaner 2 de HEO, sin
   tocarlo, son las REGLAS DE COMPRA y el formato del Excel, como hace OSMA:
     · `escaner2_motor.decidir` (las seis puertas, el corte de ventas y los paises de escaner2_parametros: para
       Zentrada, solo ES, como OSMA: «para Zentrada vamos a hacer como Osma y miramos solo para españa»), que por
       debajo usa `calc_rentabilidad` y `decision_de` del viejo (escaner2_heredado_nube.py). La formula y los umbrales
       COMPRAR/VALORAR no se tocan;
     · `escaner2_motor.examinar_csv`, `fichas_compartidas` y el lector del CSV del Visualizador (escaner2_heredado_pro.py);
     · la Celda 9 del viejo (`escaner2_motor.datos_como_el_viejo` + `escribir_celda9`) y la columna «Ficha compartida».

LO PROPIO DE ZENTRADA:
  1. LA ENTRADA ES UN EXCEL, CON GUARDAS (`leer_excel`): hojas y cabeceras exactas y en orden; «Leído» de hace 2 dias
     como mucho; 100 ofertas o mas; sin llaves (Mayorista + Artículo Zentrada) repetidas; y ninguna oferta que se pueda
     pedir con precio 0 o vacio (un precio 0 es «sin sesión», nunca «gratis»). Si falla una, no hay pasada. Las mismas
     guardas, con las mismas reglas, las pasa la v2 al soltar el fichero (lib/escaner2/zentrada.ts), para decirlo en
     pantalla antes de gastar una corrida.
  2. UNA OFERTA POR FILA; LA FOTO, UNA FILA POR EAN. El catalogo «crudo» son las ofertas: cada una sale por una puerta
     previa o es la que entra en la foto (crudo = previas + foto). Coste de un EAN = la oferta MAS BARATA que se pueda
     pedir, tal cual (Fernando: «Gastos de envío no cobran»; la cuota: «no quiero que la metas en los costes de los
     productos, tu obviala»). Mix o no, cuenta (Fernando: «Pedido Mix es mas facil llegar al minimo pero yo no
     descartaría lo que no sea Mix»): solo a igual precio gana la que va en Mix.
  3. LOS CODIGOS DE BARRAS SE COMPARAN SIN CEROS A LA IZQUIERDA: en Zentrada y en Keepa el mismo UPC sale con y sin el
     cero de delante (medido el 05-oct-2026: 19 codigos solo salian en Zentrada sin el cero).
  4. LA HOJA «Fichas» HACE DE MAPA FIJO (como osma_mapa_ficha): EAN → ASIN → unidades de Zentrada por ficha. Sus ASIN
     (los de un EAN que se puede pedir) van en una lista aparte, `asins.txt` (modo «ASIN» del Visualizador); cada ficha
     se valora por su ASIN con coste = unidades × precio, y si el ASIN de la fila de siempre es de «Fichas» de su EAN,
     mandan sus unidades. Las de origen «nuestro» sin ninguna oferta que se pueda pedir, a la ultima hoja.
  5. AL FINAL DE «Análisis», lo de la compra: donde se compra (mayorista y si va en Mix), unidades por caja, cajas
     minimas, compra minima, el precio mas barato en Mix (solo si la mas barata no va en Mix y hay otra que si), el
     origen, el puesto en mas vendidos, alimentacion y mercancia peligrosa (estas dos solo informan).
  6. EN «Resumen», cuantos COMPRAR hay por origen y, aparte, cuantos vienen SOLO del top 300 (ni nuestro ni
     competidor): con ese numero decide Fernando si se baja a 1.000 (Fernando: «vamos a empezar con los 300 y lo que
     vende la competencia»).

🔴 EL REPO ES PUBLICO: aqui no hay ni un precio real (las pruebas van con precios inventados y redondos).
"""
import io
import re
from datetime import datetime, timedelta, timezone
from decimal import Decimal, ROUND_HALF_UP

import escaner2_motor as e2

PROVEEDOR = 'ZENTRADA'
CARPETA = 'zentrada'
NOMBRE = 'Zentrada'
# El modo de la pasada (check de escaner2_pasada.modo, migracion de la v2 de este encargo): la foto sale de un Excel.
MODO = 'excel'
TOPE_LISTA = e2.TOPE_VISUALIZADOR

# ═══════════════════════════════════════════════════════════════════════════════
# 1 · EL EXCEL DE ZENTRADA Y SUS GUARDAS
# ═══════════════════════════════════════════════════════════════════════════════
# 🔒 LAS MISMAS que lib/escaner2/zentrada.ts de la v2 (HOJAS, CABECERAS, MIN_OFERTAS, HORAS_LEIDO, ORIGENES, el formato
#    de «Leído» y los sí/no). Si se cambia una, se cambian las dos.
HOJAS = ['Ofertas', 'Fichas', 'Resumen']
CABECERA_OFERTAS = ['EAN', 'Artículo Zentrada', 'Nombre en Zentrada', 'Mayorista', 'Va en Mix', 'Se puede pedir',
                    'Precio por unidad (€)', 'Unidades por caja', 'Cajas mínimas', 'Compra mínima (€)',
                    'Precios por cantidad', 'Alimentación', 'Mercancía peligrosa', 'Puesto en más vendidos', 'Origen',
                    'Leído']
CABECERA_FICHAS = ['EAN', 'ASIN', 'Unidades de Zentrada por ficha', 'Origen', 'Nombre de la ficha']
CABECERA_RESUMEN = ['Dato', 'Valor']
CABECERAS = {'Ofertas': CABECERA_OFERTAS, 'Fichas': CABECERA_FICHAS, 'Resumen': CABECERA_RESUMEN}
MIN_OFERTAS = 100
# «“Leído” de hace 2 días como mucho»: 48 horas. Y nunca en el futuro (con 10 minutos de holgura de reloj).
HORAS_LEIDO = 48
HOLGURA_FUTURO_MIN = 10
ORIGENES = ('nuestro', 'competidor', 'top 300')
_RE_LEIDO = re.compile(r'^(\d{4})-(\d{2})-(\d{2}) (\d{2}):(\d{2})$')
_RE_ASIN = re.compile(r'^[A-Z0-9]{10}$')
# Cuantos motivos se dicen como mucho (el resto, «y N más»).
MAX_MOTIVOS = 8


class FalloZentrada(Exception):
    """Algo que impide seguir: el Excel no pasa sus guardas, o la pasada o el cruce quedan 'fallida' con este motivo."""


def _texto(v):
    """Una celda como texto, sin blancos en los extremos; vacia → ''. Un numero entero de Excel, sin «.0»."""
    if v is None:
        return ''
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    return str(v).strip()


def _si_no(v):
    """«sí»/«no» (también «si» sin tilde, en mayúsculas o minúsculas) → True/False; otra cosa → None."""
    t = _texto(v).lower()
    return True if t in ('sí', 'si') else (False if t == 'no' else None)


def _numero(v):
    """Un numero de la celda: el de Excel tal cual, o un texto con cifras (coma o punto decimal); vacio → None;
    otra cosa → ValueError."""
    if v is None or (isinstance(v, str) and not v.strip()):
        return None
    if isinstance(v, bool):
        raise ValueError(v)
    if isinstance(v, (int, float)):
        return v
    t = str(v).strip().replace(',', '.')
    if not re.fullmatch(r'-?\d+(\.\d+)?', t):
        raise ValueError(v)
    return float(t) if '.' in t else int(t)


def _entero(v):
    n = _numero(v)
    if n is None:
        return None
    if float(n) != int(n):
        raise ValueError(v)
    return int(n)


def origenes_de(v):
    """«competidor, top 300» → ['competidor', 'top 300'] (en el orden de ORIGENES); None si trae algo que no es un
    origen o no trae ninguno."""
    trozos = [t.strip().lower() for t in _texto(v).split(',')]
    trozos = [t for t in trozos if t]
    if not trozos or any(t not in ORIGENES for t in trozos):
        return None
    return [o for o in ORIGENES if o in trozos]


def leido_utc(v):
    """«AAAA-MM-DD HH:MM» en hora de Madrid → datetime en UTC; si no tiene esa forma, None."""
    from zoneinfo import ZoneInfo
    m = _RE_LEIDO.match(_texto(v))
    if not m:
        return None
    try:
        local = datetime(*(int(x) for x in m.groups()), tzinfo=ZoneInfo('Europe/Madrid'))
    except ValueError:
        return None
    return local.astimezone(timezone.utc)


def ean_norm(v):
    """Solo cifras y sin ceros delante (`moloka_ean_norm` de la base, `eanNorm` de la v2); vacio → None."""
    s = re.sub(r'[^0-9]', '', _texto(v)).lstrip('0')
    return s or None


def _gtin_ok(cifras):
    """El digito de control de un GTIN (EAN-8, UPC-12, EAN-13, GTIN-14): rellenar con ceros delante no lo cambia."""
    c = cifras.zfill(14)
    d = [int(x) for x in c[:13]][::-1]
    return (10 - sum(v * (3 if i % 2 == 0 else 1) for i, v in enumerate(d)) % 10) % 10 == int(c[13])


def core_de(ean):
    """El EAN de la foto (`escaner2_foto.ean_core`, 12 o 13 cifras): sin ceros delante y rellenado a 13. → (core,
    motivo): con motivo None, vale; si no, el porque de la puerta «EAN con forma rara». Un EAN-8 o un UPC sin su cero
    delante caben en 13 con ceros (Keepa y el lector del Pro comparan sin ellos); un GTIN-14 de caja, no."""
    t = _texto(ean)
    if not t.isdigit():
        return None, 'EAN %s con letras o signos' % t
    s = t.lstrip('0')
    if len(s) < 8 or len(s) > 13:
        return None, 'EAN %s con %d cifras' % (t, len(t))
    if not _gtin_ok(s):
        return None, 'EAN %s con el dígito de control mal' % t
    return s.zfill(13), None


def _cola(motivos):
    if len(motivos) <= MAX_MOTIVOS:
        return motivos
    return motivos[:MAX_MOTIVOS] + ['y %d más' % (len(motivos) - MAX_MOTIVOS)]


def leer_excel(contenido, ahora):
    """El Excel del dia, leido y con sus GUARDAS. `contenido`: los bytes del .xlsx; `ahora`: datetime con zona.
    → {'ofertas': [...], 'fichas': [...], 'resumen': [(dato, valor)], 'leido': datetime UTC (el mas viejo)}.
    Si no pasa una guarda, FalloZentrada con TODOS los motivos que encuentra (como mucho MAX_MOTIVOS)."""
    from openpyxl import load_workbook
    try:
        wb = load_workbook(io.BytesIO(contenido), read_only=True, data_only=True)
    except Exception as ex:
        raise FalloZentrada('el fichero no se puede abrir como Excel (%s)' % type(ex).__name__)
    if wb.sheetnames != HOJAS:
        raise FalloZentrada('las hojas tienen que ser %s, en ese orden, y son %s'
                            % (', '.join('«%s»' % h for h in HOJAS), ', '.join('«%s»' % h for h in wb.sheetnames)))
    filas = {}
    motivos = []
    for h in HOJAS:
        todas = [list(f) for f in wb[h].iter_rows(values_only=True)]
        cab = [_texto(x) for x in (todas[0] if todas else [])]
        while cab and cab[-1] == '':
            cab.pop()
        if cab != CABECERAS[h]:
            motivos.append('la cabecera de «%s» no es la esperada (%s)' % (h, ', '.join(CABECERAS[h])))
        filas[h] = [f for f in todas[1:] if any(_texto(x) for x in f)]
    if motivos:
        raise FalloZentrada(' · '.join(motivos))

    ofertas, llaves, leidos = [], {}, []
    for n, f in enumerate(filas['Ofertas'], 2):
        f = (f + [None] * len(CABECERA_OFERTAS))[:len(CABECERA_OFERTAS)]
        v = dict(zip(CABECERA_OFERTAS, f))
        malos = []
        o = {'fila': n, 'ean': _texto(v['EAN']), 'articulo': _texto(v['Artículo Zentrada']),
             'nombre': _texto(v['Nombre en Zentrada']), 'mayorista': _texto(v['Mayorista']),
             'precios_cantidad': _texto(v['Precios por cantidad'])}
        if not o['mayorista'] or not o['articulo']:
            malos.append('sin «Mayorista» o sin «Artículo Zentrada»')
        for k, col in (('mix', 'Va en Mix'), ('pedible', 'Se puede pedir'), ('alimentacion', 'Alimentación'),
                       ('peligrosa', 'Mercancía peligrosa')):
            o[k] = _si_no(v[col])
            if o[k] is None:
                malos.append('«%s» tiene que ser sí o no' % col)
        for k, col, fn in (('precio', 'Precio por unidad (€)', _numero), ('uds_caja', 'Unidades por caja', _entero),
                           ('cajas_min', 'Cajas mínimas', _entero), ('compra_min', 'Compra mínima (€)', _numero),
                           ('puesto', 'Puesto en más vendidos', _entero)):
            try:
                o[k] = fn(v[col])
            except ValueError:
                o[k] = None
                malos.append('«%s» no es un número' % col)
        o['origen'] = origenes_de(v['Origen'])
        if o['origen'] is None:
            malos.append('«Origen» tiene que ser %s (separados por coma)' % ', '.join(ORIGENES))
        leido = leido_utc(v['Leído'])
        if leido is None:
            malos.append('«Leído» no tiene la forma AAAA-MM-DD HH:MM')
        else:
            leidos.append(leido)
        if o['pedible'] and (o['precio'] is None or o['precio'] <= 0):
            malos.append('se puede pedir y no tiene precio (un precio 0 es «sin sesión», nunca «gratis»)')
        llave = (o['mayorista'], o['articulo'])
        if o['mayorista'] and o['articulo']:
            if llave in llaves:
                malos.append('la llave Mayorista + Artículo Zentrada ya está en la fila %d' % llaves[llave])
            else:
                llaves[llave] = n
        if malos:
            motivos.append('Ofertas, fila %d: %s' % (n, '; '.join(malos)))
        ofertas.append(o)

    fichas, asins = [], {}
    for n, f in enumerate(filas['Fichas'], 2):
        f = (f + [None] * len(CABECERA_FICHAS))[:len(CABECERA_FICHAS)]
        v = dict(zip(CABECERA_FICHAS, f))
        malos = []
        fi = {'fila': n, 'ean': _texto(v['EAN']), 'asin': _texto(v['ASIN']).upper(),
              'nombre': _texto(v['Nombre de la ficha'])}
        if not ean_norm(fi['ean']) or not fi['ean'].isdigit():
            malos.append('«EAN» vacío o con algo que no son cifras')
        if not _RE_ASIN.match(fi['asin']):
            malos.append('«ASIN» no es un ASIN')
        elif fi['asin'] in asins:
            malos.append('el ASIN %s ya está en la fila %d (una ficha, una fila)' % (fi['asin'], asins[fi['asin']]))
        else:
            asins[fi['asin']] = n
        try:
            fi['unidades'] = _entero(v['Unidades de Zentrada por ficha'])
        except ValueError:
            fi['unidades'] = None
        if fi['unidades'] is None or fi['unidades'] < 1:
            malos.append('«Unidades de Zentrada por ficha» tiene que ser un entero de 1 o más')
        fi['origen'] = origenes_de(v['Origen'])
        if fi['origen'] is None:
            malos.append('«Origen» tiene que ser %s (separados por coma)' % ', '.join(ORIGENES))
        if malos:
            motivos.append('Fichas, fila %d: %s' % (n, '; '.join(malos)))
        fichas.append(fi)

    if len(ofertas) < MIN_OFERTAS:
        motivos.insert(0, 'trae %d ofertas y hacen falta %d como poco (¿se cortó la lectura?)'
                       % (len(ofertas), MIN_OFERTAS))
    leido = min(leidos) if leidos else None
    if leidos:
        if ahora - min(leidos) > timedelta(hours=HORAS_LEIDO):
            motivos.insert(0, 'se leyó en Zentrada hace más de 2 días (%s): vuelve a pedirlo a Cowork'
                           % _madrid(min(leidos)))
        if max(leidos) - ahora > timedelta(minutes=HOLGURA_FUTURO_MIN):
            motivos.insert(0, '«Leído» está en el futuro (%s)' % _madrid(max(leidos)))
    if motivos:
        raise FalloZentrada(' · '.join(_cola(motivos)))
    resumen = [(_texto(f[0]) if f else '', f[1] if len(f) > 1 else None) for f in filas['Resumen']]
    return {'ofertas': ofertas, 'fichas': fichas, 'resumen': resumen, 'leido': leido}


def _madrid(d):
    from zoneinfo import ZoneInfo
    return d.astimezone(ZoneInfo('Europe/Madrid')).strftime('%Y-%m-%d %H:%M')


# ═══════════════════════════════════════════════════════════════════════════════
# 2 · LA FOTO: DE LAS OFERTAS A UNA FILA POR EAN
# ═══════════════════════════════════════════════════════════════════════════════
NOMBRE_PREVIA = dict(e2.NOMBRE_PUERTA_PREVIA,
                     no_disponible='Oferta que no se puede pedir',
                     sin_gtin='Sin EAN',
                     ean_forma_rara='EAN con forma rara',
                     duplicado_proveedor='Otra oferta del mismo EAN, más cara')


def _clave_oferta(o):
    """El orden para elegir la oferta de un EAN: la mas barata; a igual precio, la que va en Mix (es mas facil
    llegar al minimo); despues la de menor compra minima; y por ultimo, por mayorista y articulo (para que salga
    siempre la misma)."""
    cm = o.get('compra_min')
    return (Decimal(str(o['precio'])), not o.get('mix'), (cm is None, Decimal(str(cm)) if cm is not None else 0),
            o.get('mayorista') or '', o.get('articulo') or '')


def _apartado(o, motivo, detalle):
    return {'producto_heo': '%s · %s' % (o.get('mayorista') or '', o.get('articulo') or ''),
            'ean_original': o.get('ean') or '', 'nombre': o.get('nombre') or '', 'marca': '',
            'precio_catalogo': o.get('precio'), 'motivo': motivo, 'detalle': detalle, 'en_oferta': None}


def construir_foto(ofertas, M):
    """De las ofertas del Excel a la foto de la pasada, sin perder ninguna: cada oferta sale por UNA puerta previa
    o es la que entra en la foto, y crudo (ofertas) = previas + foto. En este orden:
      1. no se puede pedir → `no_disponible` (se cuenta);
      2. sin EAN → `sin_gtin` (se cuenta);
      3. EAN con forma rara (no son cifras, longitud imposible, digito de control mal) → `ean_forma_rara` (se lista);
      4. varias ofertas del mismo EAN (comparado sin ceros a la izquierda) → la de `_clave_oferta`; las demas →
         `duplicado_proveedor` (se lista: «me quedo con la de …»).
    → (foto, apartados, cuentas, compra): `compra` = {ean_norm: lo de la compra de ese EAN} (`info_compra`)."""
    previas = {p: 0 for p in e2.PUERTAS_PREVIAS}
    apartados, sirven = [], []
    for o in ofertas or []:
        if not o.get('pedible'):
            previas['no_disponible'] += 1
            continue
        if not (o.get('ean') or '').strip():
            previas['sin_gtin'] += 1
            continue
        core, porque = core_de(o['ean'])
        if core is None:
            apartados.append(_apartado(o, 'ean_forma_rara', porque))
            continue
        sirven.append((o, core))

    grupos = {}
    for o, core in sirven:
        grupos.setdefault(ean_norm(core), []).append((o, core))
    todas_por_ean = {}
    for o in ofertas or []:
        k = ean_norm(o.get('ean'))
        if k:
            todas_por_ean.setdefault(k, []).append(o)

    foto, compra = [], {}
    for k in sorted(grupos, key=lambda x: (len(x), x)):
        grupo = sorted(grupos[k], key=lambda x: _clave_oferta(x[0]))
        o, core = grupo[0]
        for otra, _c in grupo[1:]:
            apartados.append(_apartado(otra, 'duplicado_proveedor',
                                       'Otra oferta del mismo EAN: me quedo con la de %s (%s), la más barata que se '
                                       'puede pedir' % (o['mayorista'], o['articulo'])))
        foto.append({
            'producto_heo': '%s · %s' % (o['mayorista'], o['articulo']), 'ean_original': o['ean'], 'ean_core': core,
            # Sin el rescate de GTIN de HEO (era para sus codigos de caja truncados): en Zentrada un EAN de 13 cifras
            # que empieza por 1 es un EAN, y el rescate le colgaba OTRO codigo (medido con el Excel del 05-oct-2026).
            'variantes': [ean_norm(core)], 'codigos_keepa': [o['ean']],
            'nombre': o.get('nombre') or '', 'marca': '', 'categoria': '',
            'precio_catalogo': o['precio'], 'precio_unidad': o['precio'],
            'es_caja': False, 'uds_caja': None, 'es_chase': False, 'en_oferta': False, 'campana': '',
            'disponibilidad': '', 'fin_de_vida': False, 'preorder': False, 'imagen': '', 'aviso_caja': None,
            'origen_ean': None, 'aviso_ean': None,
        })
        compra[k] = info_compra([x[0] for x in grupo], todas_por_ean.get(k, []))
    for m in e2.MOTIVOS_APARTADO:
        previas[m] = sum(1 for a in apartados if a['motivo'] == m)
    n_previas = sum(previas.values())
    n = len(ofertas or [])
    cuentas = {'n_crudo': n, 'n_foto': len(foto), 'previas': previas, 'n_previas': n_previas,
               'cuadra_previo': n == n_previas + len(foto)}
    return foto, apartados, cuentas, compra


def info_compra(pedibles, todas):
    """Lo de la compra de UN EAN, para las columnas del final de «Análisis». `pedibles`: sus ofertas que se pueden
    pedir y con EAN valido, la elegida DELANTE; `todas`: todas sus ofertas (tambien las que no se pueden pedir).
      · donde, mix, uds_caja, cajas_min, compra_min, precio: los de la oferta elegida (la mas barata);
      · precio_mix: el precio de la mas barata que va en Mix, SOLO si la elegida no va en Mix y hay otra que si;
      · origen: el de todas sus ofertas, juntos; puesto: el mejor (el menor) de todas;
      · alimentacion, peligrosa: «sí» si lo dice alguna de sus ofertas (solo informan: ni filtran ni avisan)."""
    o = pedibles[0]
    mix = [x for x in pedibles if x.get('mix')]
    puestos = [x['puesto'] for x in todas if x.get('puesto') is not None]
    return {'mayorista': o['mayorista'], 'articulo': o['articulo'], 'mix': bool(o.get('mix')),
            'uds_caja': o.get('uds_caja'), 'cajas_min': o.get('cajas_min'), 'compra_min': o.get('compra_min'),
            'precio': o['precio'],
            'precio_mix': (min(x['precio'] for x in mix) if (mix and not o.get('mix')) else None),
            'origen': [g for g in ORIGENES if any(g in (x.get('origen') or []) for x in todas)],
            'puesto': min(puestos) if puestos else None,
            'alimentacion': any(x.get('alimentacion') for x in todas),
            'peligrosa': any(x.get('peligrosa') for x in todas)}


def lista_para_keepa(foto):
    """LA lista de EAN para el Visualizador de España: los de la foto (los EAN con alguna oferta que se puede pedir),
    uno por linea, sin repetir y tal como los escribe Zentrada. Sin el rescate de GTIN de HEO (ver `construir_foto`)."""
    return e2.lista_para_keepa(foto)


# ═══════════════════════════════════════════════════════════════════════════════
# 3 · LA HOJA «Fichas»: EL MAPA FIJO EAN → ASIN → UNIDADES
# ═══════════════════════════════════════════════════════════════════════════════
COLUMNA_MAPA = 'Ficha del mapa'
ANCHO_MAPA = 24
HOJA_AGOTADOS = 'Nuestros agotados en Zentrada'
TEXTO_SIN_PEDIR = 'no se puede pedir en Zentrada'
TEXTO_SIN_OFERTA = 'sin oferta en Zentrada'
COLUMNAS_AGOTADOS = ['EAN', 'ASIN', 'Nombre de la ficha', 'Nuestra ficha', 'Unidades por ficha',
                     'Precio Zentrada (€)', 'Estado']
# El CSV de ASIN se reconoce por su contenido: casi todos sus ASIN estan en la lista de «Fichas».
PARTE_CSV_MAPA = 0.8


def mapa_de_fichas(fichas, foto):
    """Las filas de «Fichas» → (disponibles, agotados): las de un EAN de la foto (con alguna oferta que se puede
    pedir) y las demas, en el orden de la hoja. Cada una: {'ean': EAN de la foto o el de la ficha, 'clave': sin ceros,
    'asin', 'unidades', 'origen', 'nombre'}. Puro."""
    en_foto = {ean_norm(f['ean_core']): f for f in foto}
    disponibles, agotados = [], []
    for fi in fichas or []:
        k = ean_norm(fi['ean'])
        f = en_foto.get(k)
        fila = {'ean': f['ean_original'] if f else fi['ean'], 'clave': k, 'asin': fi['asin'],
                'unidades': int(fi['unidades']), 'origen': list(fi.get('origen') or []), 'nombre': fi.get('nombre') or ''}
        (disponibles if f else agotados).append(fila)
    return disponibles, agotados


def lista_asin(disponibles):
    """La lista de ASIN para el Visualizador (modo «ASIN»): los de «Fichas» de un EAN que se puede pedir, sin repetir."""
    return list(dict.fromkeys(m['asin'] for m in disponibles or []))


def mapa_por_ean(disponibles):
    """{ean sin ceros: {asin: unidades}}: lo que manda sobre la fila de siempre de ese EAN."""
    salida = {}
    for m in disponibles or []:
        salida.setdefault(m['clave'], {})[m['asin']] = m['unidades']
    return salida


def pa_mapa(precio, unidades):
    """El coste de una ficha de «Fichas»: unidades × precio de Zentrada, tal cual (sin porte ni cuota), redondeado a 2
    decimales una vez, al final."""
    if precio is None:
        return None
    return float((Decimal(str(precio)) * int(unidades)).quantize(Decimal('0.01'), rounding=ROUND_HALF_UP))


def es_csv_del_mapa(asins_csv, asins_mapa):
    """¿Es este CSV el de la lista de ASIN? Si casi todos sus ASIN (`PARTE_CSV_MAPA`) son de «Fichas»."""
    asins_csv, asins_mapa = set(asins_csv or ()), set(asins_mapa or ())
    if not asins_csv or not asins_mapa:
        return False
    return len(asins_csv & asins_mapa) >= PARTE_CSV_MAPA * len(asins_csv)


def leer_csv_mapa(rutas, pro):
    """El CSV de la lista de ASIN, leido con el MISMO lector del Escaner Pro (`leer_csv_visualizador`, sin tocarlo)
    pero por ASIN: en una copia temporal, la columna de los EAN lleva el ASIN de la fila (asi una ficha sin EAN en
    Keepa, la de un pack, no se pierde). → {asin: registro}; el mas nuevo manda (`rutas` del mas nuevo al mas viejo)."""
    import csv
    import tempfile
    col_ean, col_asin = pro.CSV_COLS['ean'], pro.CSV_COLS['asin']
    copias = []
    for ruta in ([rutas] if isinstance(rutas, str) else rutas):
        with open(ruta, encoding='utf-8-sig', newline='') as fh:
            filas = list(csv.reader(fh))
        if not filas or col_asin not in filas[0]:
            continue
        cab = filas[0]
        i_asin = cab.index(col_asin)
        if col_ean not in cab:
            cab = cab + [col_ean]
        i_ean = cab.index(col_ean)
        fd, copia = tempfile.mkstemp(suffix='.csv')
        with open(fd, 'w', encoding='utf-8', newline='') as fh:
            w = csv.writer(fh)
            w.writerow(cab)
            for fila in filas[1:]:
                fila = fila + [''] * (len(cab) - len(fila))
                fila[i_ean] = fila[i_asin].strip()
                w.writerow(fila)
        copias.append(copia)
    salida = {}
    for recs in (pro.leer_csv_visualizador(copias) if copias else {}).values():
        for r in recs:
            if r.get('asin') and r['asin'] not in salida:
                salida[r['asin']] = r
    return salida


def indice_por_asin(datos_por_ean):
    """{asin: registro} de lo que devuelve `leer_csv_visualizador` (indexado por EAN): el primero de cada ASIN."""
    salida = {}
    for recs in (datos_por_ean or {}).values():
        for r in recs:
            if r.get('asin') and r['asin'] not in salida:
                salida[r['asin']] = r
    return salida


def texto_mapa(unidades):
    """La marca de la columna «Ficha del mapa»."""
    return 'pack ×%d (de «Fichas»)' % unidades if unidades > 1 else 'ficha de «Fichas»'


def _coma(x):
    return '—' if x is None else ('%.2f' % x).replace('.', ',')


def _coma3(x):
    return '—' if x is None else ('%.3f' % x).replace('.', ',')


def valorar_mapa(disponibles, foto, por_asin_mapa, por_asin_csv, caidas_por_pais, params, usados, M, compartidas):
    """Cada fila de «Fichas» de un EAN que se puede pedir, valorada con las reglas de compra del PRO
    (`escaner2_motor.decidir`, sin tocarlas) y SU coste (`pa_mapa`) → [resultado de `decidir` + lo de la ficha]. Puro.
    La ficha, por su ASIN: la del CSV de ASIN (`por_asin_mapa`: {pais: {asin: registro}}) y, si no la trae, la del de
    EAN (`por_asin_csv`, igual); nunca por un EAN."""
    por_clave = {ean_norm(f['ean_core']): f for f in foto}
    salida = []
    for m in disponibles or []:
        f = por_clave[m['clave']]
        pa = pa_mapa(f['precio_unidad'], m['unidades'])
        cands, fuente = {}, None
        for p in usados:
            rec = ((por_asin_mapa or {}).get(p) or {}).get(m['asin'])
            if rec is not None:
                fuente = fuente or 'asin'
            else:
                rec = ((por_asin_csv or {}).get(p) or {}).get(m['asin'])
                if rec is not None:
                    fuente = fuente or 'ean'
            cands[p] = [rec] if rec is not None else []
        r = e2.decidir({'nombre': f.get('nombre') or '', 'ean_core': f['ean_core'], 'precio_unidad': pa}, cands,
                       caidas_por_pais, params, M, None, compartidas)
        coste = '%d × %s € = %s €' % (m['unidades'], _coma3(f['precio_unidad']), _coma(pa))
        detalle = 'Ficha de «Fichas» (%d ud de Zentrada en la ficha; coste %s)%s · %s' % (
            m['unidades'], coste, '' if fuente else ' · ningún CSV la trae', r['detalle'])
        salida.append(dict(r, detalle=detalle, clave=m['clave'], asin_mapa=m['asin'], unidades=m['unidades'],
                           origen_ficha=m['origen'], pa=pa, precio_catalogo=f['precio_unidad'],
                           nombre_zentrada=f.get('nombre') or '', ean_original=f['ean_original'],
                           ean_core=f['ean_core'], fuente=fuente, coste=coste, marca_mapa=texto_mapa(m['unidades'])))
    return salida


def fichas_nuestras_de(productos):
    """{asin: {'ean', 'nombre'}} de nuestras fichas vivas (activas, no chase, con ASIN; el chase nunca lleva ASIN)."""
    salida = {}
    for p in productos or []:
        if p.get('activo') and not p.get('es_chase') and p.get('asin') and p['asin'] not in salida:
            salida[p['asin']] = {'ean': p.get('ean'), 'nombre': p.get('nombre') or ''}
    return salida


def agotados_nuestros(agotados, precio_sin_pedir, fichas_nuestras):
    """La hoja «Nuestros agotados en Zentrada»: las fichas de «Fichas» de origen «nuestro» sin ninguna oferta que se
    pueda pedir (y, como en OSMA, tambien las que no lo dicen pero son una ficha viva nuestra en `productos`). Sin
    Keepa ni margen: que no se esconda. `precio_sin_pedir`: {ean sin ceros: el precio mas barato de sus ofertas que no
    se pueden pedir}; sin ninguna oferta, el estado lo dice. Puro."""
    filas = []
    for m in agotados or []:
        ficha = (fichas_nuestras or {}).get(m['asin'])
        if 'nuestro' not in (m.get('origen') or []) and ficha is None:
            continue
        precio = (precio_sin_pedir or {}).get(m['clave'])
        filas.append([m['ean'], m['asin'], m.get('nombre') or '', (ficha or {}).get('nombre') or '—', m['unidades'],
                      precio, TEXTO_SIN_PEDIR if m['clave'] in (precio_sin_pedir or {}) else TEXTO_SIN_OFERTA])
    return filas


def precios_sin_pedir(ofertas):
    """{ean sin ceros: el precio mas barato} de las ofertas que NO se pueden pedir (para la hoja de agotados); un EAN
    con ofertas sin precio sale con None."""
    salida = {}
    for o in ofertas or []:
        k = ean_norm(o.get('ean'))
        if not k or o.get('pedible'):
            continue
        p = o.get('precio')
        if k not in salida or (p is not None and (salida[k] is None or p < salida[k])):
            salida[k] = p
    return salida


# ═══════════════════════════════════════════════════════════════════════════════
# 3 bis · LOS PACKS DE AMAZON: LA REGLA AM DE OSMA, COPIADA (devolución de Cowork, 05-oct-2026)
# ═══════════════════════════════════════════════════════════════════════════════
# 🔑 COPIADA de escaner2_osma.py (secciones «4 bis», encargo AM del 02-oct-2026), NO importada: cada escáner a la medida
#    de su proveedor (Fernando, 29-sep-2026). Misma regla y mismo orden de veredictos; solo cambian los nombres
#    (`cantidad_zentrada` lee el «Nombre en Zentrada» de la oferta que entra en la foto; `_numero_pack`, para no chocar
#    con el `_numero` de las celdas del Excel). Por qué hace falta aquí (Cowork, 05-oct-2026, con los CSV reales): el
#    Haribo ruleta Mega 48 g (EAN 4001686372586, sin fila en «Fichas») casaba con un ASIN de 24 paquetes y salía
#    COMPRAR con el coste de UNO. Lo de abajo es el texto de OSMA con «OSMA» cambiado por «Zentrada».
# Zentrada vende la unidad suelta, y el cruce por EAN cae a veces en un ASIN que es un MULTIPACK con el EAN de la unidad
# (en OSMA, las Melody Pops de 15 g en el «Pack 48», su primera pasada real 655b3e06; aquí, el Haribo de arriba): con el
# coste de UNA unidad salen COMPRAR falsos. Aqui se lee, del CSV del Visualizador, cuantas unidades de Zentrada lleva el
# ASIN, y se multiplica el coste SOLO cuando lo sostienen dos señales independientes que dicen lo mismo.
#
# LA REGLA. Cada señal es una LECTURA de N = unidades de Zentrada que lleva el ASIN, comparando lo que dice Amazon con la
# cantidad del nombre de Zentrada (`cantidad_zentrada`: el recuento «24er», «8 Stück», «56x10», y la medida «45g», «50ml»,
# «100ml + 250ml»; sin recuento en el nombre, el articulo de Zentrada es UNA unidad):
#   · contenido   «Detalles de la unidad»: valor y tipo. En g/ml, ÷ la medida de Zentrada; en unidades, ÷ su recuento.
#   · tamaño      «Tamaño» («45 g (Paquete de 24)» = 1080 g; «24 unità (Confezione da 1)» = 24 ud), igual.
#   · recuento    «Número de artículos» (solo si es 2 o más: Amazon pone 1 casi siempre), ÷ el recuento de Zentrada.
#   · paquete     «Paquete: Cantidad» (solo si es 2 o más), igual.
#   · título      «Pack 48», «paquete de 4», «lote de 12», «3 x 50 ml», «10 piezas», «24 unidades» (2 o más), igual.
#   El contenido en unidades y el recuento son, los dos, el recuento que escribe el vendedor: cuentan como UNA señal.
# Y el veredicto, por este orden:
#   1. Si las lecturas de CONTENIDO y TAMAÑO que hay dan todas 1 contra una cantidad LEIDA en el nombre de Zentrada, el
#      ASIN lleva justo lo que vende Zentrada: no es pack, diga lo que diga el recuento (el set de Adidas: 100 + 250 ml =
#      350 ml, aunque Amazon diga «6 artículos»).
#   2. Si ninguna lectura da 2 o más: no es pack.
#   3. Se MULTIPLICA por N si todas las lecturas dan N (un entero ≥ 2), vienen de DOS señales distintas como poco, y
#      el nombre de Zentrada se ha leido sin ambigüedad.
#   4. Si no (una sola señal, o señales que se contradicen, o el nombre de Zentrada con dos cantidades distintas): NO se
#      multiplica y la fila nunca es COMPRAR: baja a VALORAR con «posible pack en Amazon: revisar (…las señales…)».
# 🔒 Si el ASIN está en «Fichas» de ese EAN, mandan sus unidades y esto no se mira: no se multiplica dos veces.
COLUMNAS_PACK = {'n_art': 'Número de artículos', 'valor_ud': 'Detalles de la unidad: Valor de la unidad',
                 'tipo_ud': 'Detalles de la unidad: Tipo de unidad', 'paquete': 'Paquete: Cantidad',
                 'tamano': 'Tamaño'}
# La columna del titulo, como la busca el lector del CSV del viejo (escaner2_heredado_pro.leer_csv_visualizador).
COLUMNAS_TITULO = ('Título', 'Titulo', 'Title', 'Título principal', 'Titulo principal')
PACK_SI, PACK_DUDOSO, PACK_NO = 'pack', 'dudoso', 'no'
# Lo que se admite como «el mismo numero» (58,33 ml × 6 = 349,98 ml frente a 350 ml).
TOLERANCIA_PACK = 0.02
TEXTO_POSIBLE_PACK = 'posible pack en Amazon: revisar'

# Unidades, a g, ml o 'ud' (recuento). Lo que no esta aqui (cm, metro, onzas…) no es una cantidad que se compare.
_UNIDAD = {'ml': ('ml', 1), 'mililitro': ('ml', 1), 'millilitro': ('ml', 1), 'milliliter': ('ml', 1),
           'milliliters': ('ml', 1), 'mililitros': ('ml', 1), 'cl': ('ml', 10), 'l': ('ml', 1000),
           'ltr': ('ml', 1000), 'liter': ('ml', 1000), 'litro': ('ml', 1000), 'litros': ('ml', 1000),
           'g': ('g', 1), 'gr': ('g', 1), 'gramo': ('g', 1), 'gramos': ('g', 1), 'gram': ('g', 1), 'grams': ('g', 1),
           'kg': ('g', 1000), 'kilogramo': ('g', 1000), 'kilogramos': ('g', 1000),
           'unidad': ('ud', 1), 'unidades': ('ud', 1), 'unità': ('ud', 1), 'unité': ('ud', 1), 'unités': ('ud', 1),
           'stück': ('ud', 1), 'stückzahl': ('ud', 1), 'count': ('ud', 1), 'pieza': ('ud', 1), 'piezas': ('ud', 1)}
_MEDIDA = r'(ml|cl|ltr|liter|l|kg|gr|g)'
_FIN_PALABRA = r'(?![a-zäöüßà-ÿ])'
_RE_N_X_MEDIDA = re.compile(r'(?<![\d.,])(\d{1,3})\s*[xX]\s*(\d+(?:[.,]\d+)?)\s*' + _MEDIDA + _FIN_PALABRA, re.I)
_RE_MEDIDA_ZEN = re.compile(r'(?<![\d.,])(\d+(?:[.,]\d+)?)\s*' + _MEDIDA + r'\.?' + _FIN_PALABRA, re.I)
_RE_ER = re.compile(r'(?<![\d.,])(\d+)\s*er' + _FIN_PALABRA, re.I)
_RE_STUECK = re.compile(r'(?<![\d.,])(\d+)\s*-?\s*(?:stück|stck|stk|st|teilig|teil|tlg)\.?' + _FIN_PALABRA, re.I)
# «56x10» (Tempo: 56 paquetes de 10) es un recuento; «190x68» (una vela) o «20x15cm» son medidas: solo cifras de una
# o dos y sin unidad ni otra «x» detras.
_RE_NXM = re.compile(r'(?<![\d.,])(\d{1,2})\s*[xX]\s*(\d{1,2})(?![\d.,])(?!\s*(?:[xX]|cm|mm|m' + _FIN_PALABRA + r'|'
                     + _MEDIDA + _FIN_PALABRA + r'))', re.I)
# «im 18er Tray», «12 St. im Aufsteller»: el expositor, no lo que se vende.
_RE_EXPOSITOR = re.compile(r'^\W*(?:im\s+)?(?:tray|display|aufsteller|karton)' + _FIN_PALABRA, re.I)
_RE_TAMANO = re.compile(r'^\s*(?:(\d{1,3})\s*x\s*)?(\d+(?:[.,]\d+)?)\s*([a-zà-ÿ]+)\b(?:\s*\((?:paquete de|pack of|'
                        r'confezione da|lot de|packung mit)\s*(\d+)\)|\s*\((\d+)er[\s-]*pack\))?', re.I)
_RES_TITULO = [re.compile(p, re.I) for p in (
    r'\bpack\s*(?:de\s*)?(\d+)\b', r'\b(\d+)\s*-?\s*pack\b', r'\blote\s+de\s+(\d+)\b', r'\bpaquete\s+de\s+(\d+)\b',
    r'\bcaja\s+de\s+(\d+)\b', r'\b(\d+)\s*x\s*\d+(?:[.,]\d+)?\s*' + _MEDIDA + r'\b', r'\b(\d+)\s+(?:piezas|unidades)\b')]


def _numero_pack(s):
    try:
        return float(str(s).strip().replace(',', '.'))
    except (TypeError, ValueError):
        return None


def senales_pack_csv(rutas):
    """{asin: {'n_art', 'valor_ud', 'tipo_ud', 'paquete', 'tamano', 'titulo'}} del CSV del Visualizador, tal cual
    (texto). `rutas`: una o varias, del mas nuevo al mas viejo, y EL MAS NUEVO MANDA (como el lector del viejo). Una
    columna que el CSV no trae sale vacia."""
    import csv
    salida = {}
    for ruta in ([rutas] if isinstance(rutas, str) else rutas):
        with open(ruta, encoding='utf-8-sig', newline='') as fh:
            filas = csv.reader(fh)
            cab = next(filas, [])
            ix = {h: i for i, h in enumerate(cab)}
            if 'ASIN' not in ix:
                continue
            col_tit = next((c for c in COLUMNAS_TITULO if c in ix), None)

            def celda(fila, col):
                return fila[ix[col]].strip() if (col in ix and ix[col] < len(fila)) else ''
            for fila in filas:
                asin = celda(fila, 'ASIN')
                if asin and asin not in salida:
                    salida[asin] = dict({k: celda(fila, c) for k, c in COLUMNAS_PACK.items()},
                                        titulo=celda(fila, col_tit) if col_tit else '')
    return salida


def cantidad_zentrada(nombre):
    """Lo que dice el nombre de Zentrada de su cantidad → {'n': recuento o None, 'medida': (total, 'g'|'ml') o None,
    'ambigua': bool, 'texto': lo leido}. El recuento: «24er», «8 Stück», «132 St.», «77Stk», «32 teil.», «4tlg»,
    «56x10», o el N de «6x50g»; el de un expositor («im 18er Tray») no cuenta. La medida, en g o ml: «45g», «1,35 Liter», «6x50g» = 300 g,
    y las de un set con «+» se suman («100ml + 250ml» = 350 ml). Dos recuentos distintos, o dos medidas distintas sin
    «+», no se adivinan: 'ambigua'."""
    s = str(nombre or '')
    recuentos, medidas, tapado = [], [], []
    for m in _RE_N_X_MEDIDA.finditer(s):
        dim, k = _UNIDAD[m.group(3).lower()]
        recuentos.append(int(m.group(1)))
        medidas.append((m.start(), m.end(), int(m.group(1)) * _numero_pack(m.group(2)) * k, dim))
        tapado.append((m.start(), m.end()))
    for m in _RE_MEDIDA_ZEN.finditer(s):
        if any(a <= m.start() < b for a, b in tapado):
            continue
        dim, k = _UNIDAD[m.group(2).lower()]
        medidas.append((m.start(), m.end(), _numero_pack(m.group(1)) * k, dim))
    for rx in (_RE_ER, _RE_STUECK):
        for m in rx.finditer(s):
            if not _RE_EXPOSITOR.match(s[m.end():]):
                recuentos.append(int(m.group(1)))
    for m in _RE_NXM.finditer(s):
        recuentos.append(int(m.group(1)) * int(m.group(2)))
    recuentos = [n for n in recuentos if n > 0]
    ambigua = len(set(recuentos)) > 1
    n = recuentos[0] if (recuentos and not ambigua) else None
    medida = None
    if medidas:
        medidas.sort()
        dims = {x[3] for x in medidas}
        unidas = all('+' in s[medidas[i][1]:medidas[i + 1][0]] for i in range(len(medidas) - 1))
        if len(dims) == 1 and len(medidas) > 1 and unidas:
            medida = (sum(x[2] for x in medidas), medidas[0][3])
        elif len(dims) == 1 and len({round(x[2], 6) for x in medidas}) == 1:
            medida = (medidas[0][2], medidas[0][3])
        else:
            ambigua = True
    partes = ([('%d ud' % n)] if n else []) + ([_cifra_pack(medida[0]) + ' ' + medida[1]] if medida else [])
    return {'n': n, 'medida': medida, 'ambigua': ambigua,
            'texto': ' · '.join(partes) or ('cantidad ambigua' if ambigua else 'sin cantidad: 1 unidad')}


def _cifra_pack(x):
    return ('%d' % round(x)) if abs(x - round(x)) < 1e-9 else ('%.2f' % x).rstrip('0').rstrip('.').replace('.', ',')


def _entero_pack(v):
    """La lectura como entero si lo es (con la tolerancia), y si no None."""
    if v is None or v <= 0:
        return None
    n = round(v)
    return n if (n >= 1 and abs(v - n) <= TOLERANCIA_PACK * n) else None


def lecturas_pack(senal, cant):
    """Las lecturas de N (unidades de Zentrada que lleva el ASIN) de cada señal del CSV → [{'fuente', 'grupo', 'texto',
    'valor', 'trivial'}]. `trivial`: comparada contra el «1 unidad» que se supone a un nombre de Zentrada sin recuento
    (no contra una cantidad leida)."""
    if not senal:
        return []
    recuento = cant.get('n') or 1
    sin_recuento = cant.get('n') is None
    medida = cant.get('medida')
    salida = []

    def contra(total, dim, fuente, texto, grupo=None):
        if dim == 'ud':
            salida.append({'fuente': fuente, 'grupo': grupo or fuente, 'texto': texto, 'valor': total / recuento,
                           'trivial': sin_recuento})
        elif medida and medida[1] == dim and medida[0] > 0:
            salida.append({'fuente': fuente, 'grupo': grupo or fuente, 'texto': texto, 'valor': total / medida[0],
                           'trivial': False})

    v, tipo = _numero_pack(senal.get('valor_ud')), _UNIDAD.get(str(senal.get('tipo_ud') or '').strip().lower())
    if v and v > 0 and tipo:
        contra(v * tipo[1], tipo[0], 'contenido', 'contenido %s %s' % (senal['valor_ud'], senal['tipo_ud']),
               grupo='recuento' if tipo[0] == 'ud' else 'contenido')
    m = _RE_TAMANO.match(str(senal.get('tamano') or ''))
    if m and m.group(3).lower() in _UNIDAD:
        dim, k = _UNIDAD[m.group(3).lower()]
        por = int(m.group(1) or 1) * int(m.group(4) or m.group(5) or 1)
        total = _numero_pack(m.group(2)) * k * por
        texto = 'tamaño «%s»' % senal['tamano']
        if dim == 'ud' or (medida and medida[1] == dim):
            contra(total, dim, 'tamaño', texto)
        elif por >= 2:
            # «20 g (Paquete de 3)» sin medida en el nombre de Zentrada: cuenta el «Paquete de».
            salida.append({'fuente': 'tamaño', 'grupo': 'tamaño', 'texto': texto, 'valor': por / recuento,
                           'trivial': sin_recuento})
    for clave, fuente, grupo, rotulo in (('n_art', 'recuento', 'recuento', 'n.º de artículos'),
                                         ('paquete', 'paquete', 'paquete', 'paquete: cantidad')):
        x = _numero_pack(senal.get(clave))
        if x is not None and x >= 2 and x == int(x):
            salida.append({'fuente': fuente, 'grupo': grupo, 'texto': '%s %d' % (rotulo, x), 'valor': x / recuento,
                           'trivial': sin_recuento})
    vistos = set()
    for rx in _RES_TITULO:
        for mt in rx.finditer(str(senal.get('titulo') or '')):
            x = int(mt.group(1))
            if x >= 2 and x not in vistos:
                vistos.add(x)
                salida.append({'fuente': 'título', 'grupo': 'título', 'texto': 'título «%s»' % mt.group(0).strip(),
                               'valor': x / recuento, 'trivial': sin_recuento})
    return salida


def factor_pack_amazon(senal, cant):
    """El veredicto de la regla de arriba → {'estado': PACK_SI | PACK_DUDOSO | PACK_NO, 'factor', 'senales': texto}."""
    lecturas = lecturas_pack(senal, cant)
    texto = 'Zentrada %s; Amazon: %s' % (cant['texto'], ' · '.join(
        '%s → %s' % (x['texto'], _cifra_pack(x['valor'])) for x in lecturas) or 'sin señales')
    no = {'estado': PACK_NO, 'factor': 1, 'senales': texto}
    totales = [x for x in lecturas if x['fuente'] in ('contenido', 'tamaño') and not x['trivial']]
    if totales and all(_entero_pack(x['valor']) == 1 for x in totales):
        return no
    packs = [x for x in lecturas if (_entero_pack(x['valor']) or 0) >= 2]
    if not packs:
        return no
    ns = {_entero_pack(x['valor']) for x in packs}
    contra = [x for x in lecturas if _entero_pack(x['valor']) not in ns]
    if len(ns) == 1 and not contra and not cant['ambigua'] and len({x['grupo'] for x in packs}) >= 2:
        return {'estado': PACK_SI, 'factor': ns.pop(), 'senales': texto}
    return {'estado': PACK_DUDOSO, 'factor': 1, 'senales': texto}


# ═══════════════════════════════════════════════════════════════════════════════
# 4 · EL CRUCE DE UNA FILA DE LA FOTO
# ═══════════════════════════════════════════════════════════════════════════════
def decidir_zentrada(fila, cands_por_pais, caidas_por_pais, params, M, eleccion, compartidas, mapa_ean=None,
                     senales_pack=None):
    """La puerta de UNA fila de la foto con las reglas de compra del PRO (`escaner2_motor.decidir`, sin tocarlas). Lo
    propio, en este orden (como `decidir_osma`):
      · si el ASIN que sale esta en «Fichas» de este EAN (`mapa_ean`: {asin: unidades}), MANDAN SUS UNIDADES (y las
        señales de Amazon ni se miran: no se multiplica dos veces);
      · si no, y la fila se vende (puertas d, e y f), el PACK DE AMAZON (`factor_pack_amazon`, con las señales de
        `senales_pack`: {pais: {asin: señal}} de `senales_pack_csv`, contra la cantidad del «Nombre en Zentrada»): con
        pack, N × precio; con «posible pack», la fila nunca es COMPRAR (baja a VALORAR);
      · con un factor > 1, coste = factor × precio (un redondeo) y se decide OTRA VEZ con ese coste (la eleccion de
        ficha no depende del precio: sale la misma, y se comprueba).
    → el resultado de `decidir` con 'factor', 'pa', 'pack' (None, 'fichas' o 'amazon') y 'pack_amazon' (el veredicto,
    o None si no se ha mirado)."""
    r = e2.decidir(fila, cands_por_pais, caidas_por_pais, params, M, eleccion, compartidas)
    asin = r.get('asin')
    de_fichas = (mapa_ean or {}).get(asin) if asin else None
    factor, pack, pack_amz = 1, None, None
    if de_fichas is not None:
        factor, pack = de_fichas, ('fichas' if de_fichas > 1 else None)
    elif senales_pack is not None and asin and r['puerta'] in e2.PUERTAS_ANALISIS:
        senal = next((senales_pack[p][asin] for p in params['paises_calculo'] if asin in (senales_pack.get(p) or {})),
                     None)
        pack_amz = factor_pack_amazon(senal, cantidad_zentrada(fila.get('nombre')))
        if pack_amz['estado'] == PACK_SI:
            factor, pack = pack_amz['factor'], 'amazon'
    pa = fila['precio_unidad']
    if factor > 1 and pa is not None:
        pa = pa_mapa(pa, factor)
        r2 = e2.decidir(dict(fila, precio_unidad=pa), cands_por_pais, caidas_por_pais, params, M, eleccion, compartidas)
        if r2.get('asin') != asin:
            raise FalloZentrada('la ficha elegida cambia con el coste del pack (%s → %s): no debería' % (asin, r2.get('asin')))
        r = r2
    if pack_amz and pack_amz['estado'] == PACK_DUDOSO and r['puerta'] == 'f':
        r = _sin_comprar(r, params)
    notas = []
    if factor > 1:
        notas.append(texto_pack(factor, fila['precio_unidad'], pa, pack))
    if pack == 'amazon':
        notas.append('señales: %s' % pack_amz['senales'])
    if pack_amz and pack_amz['estado'] == PACK_DUDOSO:
        notas.append('%s (%s)' % (TEXTO_POSIBLE_PACK, pack_amz['senales']))
    if notas:
        r = dict(r, detalle='%s · %s' % (r['detalle'], ' · '.join(notas)))
    return dict(r, factor=factor, pa=pa, pack=pack, pack_amazon=pack_amz)


def texto_pack(factor, precio, pa, pack):
    """«pack de 3 de «Fichas»: coste 3 × 1,500 € = 4,50 €» o «pack de 24 en Amazon: coste 24 × 0,370 € = 8,88 €»."""
    return 'pack de %d %s: coste %d × %s € = %s €' % (factor, 'en Amazon' if pack == 'amazon' else 'de «Fichas»', factor,
                                                      _coma3(precio), _coma(pa))


def _sin_comprar(r, params):
    """(AM, copiada de escaner2_osma.py) Un «posible pack en Amazon» nunca es COMPRAR: cada pais que decia COMPRAR pasa
    a VALORAR y la puerta se saca OTRA VEZ con `escaner2_motor._puerta_def` (la e, con su mejor pais). En el detalle se
    cambia SOLO el trozo que escribe `_puerta_def` (lo de delante, la eleccion de ficha, y lo de detras, la ficha
    compartida, se quedan)."""
    paises = [p for p in params['paises_calculo'] if p in (r.get('paises') or {})]
    antes = e2._puerta_def(dict(r), paises, r['paises'])['detalle']
    calculos = {p: dict(c, decision='VALORAR') if c.get('decision') == 'COMPRAR' else c for p, c in r['paises'].items()}
    nuevo = e2._puerta_def(dict(r, paises=calculos), paises, calculos)
    if r['detalle'].count(antes) != 1:
        raise FalloZentrada('el detalle de la puerta no se puede rehacer (posible pack): %r no está una vez' % antes)
    return dict(nuevo, detalle=r['detalle'].replace(antes, nuevo['detalle']))


# ═══════════════════════════════════════════════════════════════════════════════
# 5 · EL EXCEL: EL DEL ESCANER 2, CON LO DE LA COMPRA AL FINAL
# ═══════════════════════════════════════════════════════════════════════════════
COLUMNAS_COMPRA = ['Dónde se compra', 'Unidades por caja', 'Cajas mínimas', 'Compra mínima (€)',
                   'Precio más barato en Mix (€)', 'Origen', 'Puesto en más vendidos', 'Alimentación',
                   'Mercancía peligrosa']
ANCHOS_COMPRA = [34, 10, 9, 11, 13, 26, 11, 11, 11]
COLUMNAS_PUERTAS = ['EAN', 'Nombre', 'Marca', 'Precio compra (€)', 'Puerta', 'Motivo', 'Detalle', 'Caja',
                    'Precio caja (€)', 'EAN de la figura', 'Origen del EAN', 'Aviso del EAN']


def texto_donde(c):
    """«Mayorista X · va en Mix» o «Mayorista X · fuera de Mix»."""
    return '%s · %s' % (c.get('mayorista') or '—', 'va en Mix' if c.get('mix') else 'fuera de Mix')


def texto_origen(origenes):
    return ', '.join(o for o in ORIGENES if o in (origenes or [])) or None


def solo_top(origenes):
    """¿Viene SOLO del top 300 (ni nuestro ni competidor)?"""
    o = set(origenes or [])
    return 'top 300' in o and not (o & {'nuestro', 'competidor'})


def foto_para_excel(foto, resultados):
    """La foto con la que se escribe la Celda 9: si la ficha que sale es un pack (de «Fichas» o, AM, de Amazon), «PA (€)»
    es su coste (unidades × precio) y el nombre y «Coherencia caja» lo dicen; un «posible pack en Amazon», en «Coherencia
    caja» con sus señales. Y el ASIN de la fila (con el EAN, la llave de cada
    fila de «Análisis»: una ficha de «Fichas» puede llevar el mismo EAN que la fila de siempre)."""
    por_foto = {r['foto_id']: r for r in resultados}
    salida = []
    for f in foto:
        r = por_foto.get(f['id']) or {}
        g = dict(f)
        factor = r.get('factor', 1)
        if factor > 1:
            g['precio_unidad'] = r['pa']
            g['nombre'] = '%s · pack de %d %s (%d × %s €)' % (
                f.get('nombre') or '', factor, 'en Amazon' if r.get('pack') == 'amazon' else 'de «Fichas»', factor,
                _coma3(f['precio_unidad']))
            g['aviso_caja'] = texto_pack(factor, f['precio_unidad'], r['pa'], r.get('pack'))
        elif (r.get('pack_amazon') or {}).get('estado') == PACK_DUDOSO:
            g['aviso_caja'] = '%s (%s)' % (TEXTO_POSIBLE_PACK, r['pack_amazon']['senales'])
        g['asin'] = r.get('asin')
        salida.append(g)
    return salida


def filas_mapa_excel(mapa):
    """Las fichas de «Fichas» que se venden (puertas d, e y f), como filas de la foto y resultados para las hojas del
    viejo → (foto, resultados). Cada una con su id propio ('mapa:<ean>:<asin>') y el coste ya calculado."""
    foto, resultados = [], []
    for r in mapa or []:
        if r['puerta'] not in e2.PUERTAS_ANALISIS:
            continue
        fid = 'mapa:%s:%s' % (r['clave'], r['asin_mapa'])
        foto.append({
            'id': fid, 'producto_heo': '', 'ean_original': r['ean_original'], 'ean_core': r['ean_core'],
            'nombre': '%s · %s' % (r['nombre_zentrada'], r['marca_mapa']), 'marca': '',
            'precio_catalogo': r['precio_catalogo'], 'precio_unidad': r['pa'],
            'aviso_caja': 'coste de «Fichas»: %s' % r['coste'] if r['unidades'] > 1 else None,
            'asin': r['asin_mapa'], '_mapa': r['marca_mapa'], '_origen_ficha': r['origen_ficha']})
        resultados.append(dict(r, foto_id=fid, id=fid, fichas=None, compartida={}, eleccion=None, factor=1))
    return foto, resultados


def _alargar_tabla(ws, col, nombres):
    """La tabla de «Análisis» (T_Analisis del viejo) se alarga para cubrir las columnas nuevas del final."""
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.table import TableColumn
    for tabla in ws.tables.values():
        ini, fin = tabla.ref.split(':')
        tabla.ref = '%s:%s%s' % (ini, get_column_letter(col), re.sub(r'^[A-Z]+', '', fin))
        if tabla.tableColumns:
            for nombre in nombres:
                tabla.tableColumns.append(TableColumn(id=len(tabla.tableColumns) + 1, name=nombre))
        if tabla.autoFilter is not None:
            tabla.autoFilter.ref = tabla.ref


def poner_columna_mapa(wb, foto_excel):
    """«Ficha del mapa» detras de la ultima columna de «Análisis», fila a fila por (EAN, ASIN), con la marca de las
    fichas de «Fichas». Devuelve cuantas filas."""
    from openpyxl.styles import Font
    from openpyxl.utils import get_column_letter
    ws = wb['Análisis']
    cab = [c.value for c in ws[1]]
    i_ean, i_asin = cab.index('EAN'), cab.index('ASIN')
    col = len(cab) + 1
    ws.cell(row=1, column=col, value=COLUMNA_MAPA).font = Font(bold=True)
    ws.column_dimensions[get_column_letter(col)].width = ANCHO_MAPA
    marcas = {(str(f['ean_original']), f.get('asin')): f['_mapa'] for f in foto_excel if f.get('_mapa')}
    n = 0
    for fila in range(2, ws.max_row + 1):
        m = marcas.get((str(ws.cell(row=fila, column=i_ean + 1).value), ws.cell(row=fila, column=i_asin + 1).value))
        if m:
            ws.cell(row=fila, column=col, value=m)
            n += 1
    _alargar_tabla(ws, col, [COLUMNA_MAPA])
    return n


def poner_columnas_compra(wb, foto_excel, compra):
    """Las nueve columnas de la compra (`COLUMNAS_COMPRA`) AL FINAL de «Análisis» (columnas nuevas → al final,
    SIEMPRE: hay quien lee la hoja por letra), fila a fila por (EAN, ASIN): lo de `info_compra` del EAN de la fila;
    en una ficha de «Fichas», el origen suma el de la ficha. → filas con dato."""
    from openpyxl.styles import Font
    from openpyxl.utils import get_column_letter
    ws = wb['Análisis']
    cab = [c.value for c in ws[1]]
    i_ean, i_asin = cab.index('EAN'), cab.index('ASIN')
    col0 = len(cab) + 1
    for j, (nombre, ancho) in enumerate(zip(COLUMNAS_COMPRA, ANCHOS_COMPRA)):
        ws.cell(row=1, column=col0 + j, value=nombre).font = Font(bold=True)
        ws.column_dimensions[get_column_letter(col0 + j)].width = ancho
    por_clave = {(str(f['ean_original']), f.get('asin')): f for f in foto_excel}
    n = 0
    for fila in range(2, ws.max_row + 1):
        ean = str(ws.cell(row=fila, column=i_ean + 1).value)
        f = por_clave.get((ean, ws.cell(row=fila, column=i_asin + 1).value))
        c = (compra or {}).get(ean_norm(ean))
        if not c:
            continue
        origen = list(c['origen']) + list((f or {}).get('_origen_ficha') or [])
        valores = [texto_donde(c), c.get('uds_caja'), c.get('cajas_min'), c.get('compra_min'), c.get('precio_mix'),
                   texto_origen(origen), c.get('puesto'), 'sí' if c.get('alimentacion') else 'no',
                   'sí' if c.get('peligrosa') else 'no']
        for j, v in enumerate(valores):
            celda = ws.cell(row=fila, column=col0 + j, value=v)
            if COLUMNAS_COMPRA[j] in ('Compra mínima (€)', 'Precio más barato en Mix (€)') and v is not None:
                celda.number_format = '0.00'
        n += 1
    _alargar_tabla(ws, col0 + len(COLUMNAS_COMPRA) - 1, COLUMNAS_COMPRA)
    return n


def origen_de_fila(r, foto_por_id, compra):
    """El origen de una fila de «Análisis»: el de las ofertas de su EAN, y en una ficha de «Fichas», tambien el suyo."""
    f = foto_por_id.get(r['foto_id']) or {}
    c = (compra or {}).get(ean_norm(f.get('ean_core') or r.get('ean_core'))) or {}
    return set(c.get('origen') or []) | set(r.get('origen_ficha') or [])


def comprar_por_origen(en_hoja, foto_por_id, compra, pais='ES'):
    """Los COMPRAR de «Análisis» (en `pais`) por origen, y aparte los que vienen SOLO del top 300 → (dict, solo_top)."""
    cuenta = {o: 0 for o in ORIGENES}
    solo = 0
    for r in en_hoja:
        if ((r.get('paises') or {}).get(pais) or {}).get('decision') != 'COMPRAR':
            continue
        org = origen_de_fila(r, foto_por_id, compra)
        for o in org:
            cuenta[o] += 1
        solo += solo_top(org)
    return cuenta, solo


def texto_packs_amazon(en_hoja):
    """La línea del Resumen, como la de OSMA: «N multiplicados · M dudosos a VALORAR · K dudosos sin margen»."""
    pk = [r.get('pack_amazon') or {} for r in en_hoja]
    n_mult = sum(1 for r in en_hoja if r.get('pack') == 'amazon')
    n_val = sum(1 for r, x in zip(en_hoja, pk) if x.get('estado') == PACK_DUDOSO and r['puerta'] == 'e')
    n_no = sum(1 for r, x in zip(en_hoja, pk) if x.get('estado') == PACK_DUDOSO and r['puerta'] != 'e')
    return '%d multiplicados · %d dudosos a VALORAR%s' % (n_mult, n_val, (' · %d dudosos sin margen (NO COMPRAR)' % n_no)
                                                          if n_no else '')


def escribir_excel(foto, resultados, apartados, M, info):
    """El Excel del cruce de Zentrada: DELANTE las hojas del viejo (la Celda 9, con «Análisis» solo de ES), DETRAS las
    del escaner 2 con los mismos nombres que el de OSMA (Resumen, Comparación, Varias fichas, Puertas y Puertas previas)
    y la ULTIMA, «Nuestros agotados en Zentrada». En «Análisis», detras de «Ficha compartida»: «Ficha del mapa» y las
    nueve columnas de la compra. Solo lo calculado: aqui no se decide nada. → bytes.
    `info['mapa']`: las fichas de «Fichas» valoradas (`valorar_mapa`); `info['compra']`: {ean sin ceros: info_compra};
    `info['agotados']`: las filas de la ultima hoja (`agotados_nuestros`)."""
    from openpyxl.styles import Font
    por_foto = {f['id']: f for f in foto}
    mapa = info.get('mapa') or []
    asins_mapa = {r['asin_mapa'] for r in mapa}
    foto_mapa, res_mapa = filas_mapa_excel(mapa)
    foto_excel = foto_para_excel(foto, resultados) + foto_mapa
    paises = [p for p in M.PAISES if p in (info['params'].get('paises_calculo') or [])]
    datos = e2.datos_como_el_viejo(foto_excel, resultados, apartados, M, info.get('eleccion'))
    if res_mapa or asins_mapa:
        extra = e2.datos_como_el_viejo(foto_excel, res_mapa, [], M, None)['registros']
        regs = [x for x in datos['registros'] if x['asin'] not in asins_mapa] + extra
        regs.sort(key=lambda x: x['_margen_es'] if x['_margen_es'] is not None else -10 ** 9, reverse=True)
        datos['registros'] = regs
        for x in extra:
            datos['cotejo_info'].setdefault(x['ean'], {'veredicto': None, 'detalle': None})
    datos['PROVEEDOR'] = NOMBRE
    wb = e2.escribir_celda9(datos, M, paises=paises if len(paises) == 1 else None)
    e2.poner_columna_ficha_compartida(wb, foto_excel, resultados)
    poner_columna_mapa(wb, foto_excel)
    poner_columnas_compra(wb, foto_excel, info.get('compra'))

    def hoja(nombre, cabecera, filas, anchos=None):
        ws = wb.create_sheet(nombre)
        ws.append(cabecera)
        for c in ws[1]:
            c.font = Font(bold=True)
        for fila in filas:
            ws.append(fila)
        for letra, ancho in (anchos or {}).items():
            ws.column_dimensions[letra].width = ancho
        ws.freeze_panes = 'A2'
        return ws

    p = info['params']
    en_hoja = [r for r in resultados if r['puerta'] in e2.PUERTAS_ANALISIS and r.get('asin') not in asins_mapa] + res_mapa
    sustituidas = [r for r in resultados if r['puerta'] in e2.PUERTAS_ANALISIS and r.get('asin') in asins_mapa]
    foto_por_id = dict(por_foto, **{f['id']: f for f in foto_mapa})
    dec = {d: sum(1 for r in en_hoja if ((r.get('paises') or {}).get('ES') or {}).get('decision') == d)
           for d in ('COMPRAR', 'VALORAR')}
    por_origen, solo = comprar_por_origen(en_hoja, foto_por_id, info.get('compra'))
    filas_res = [['Pasada', info['pasada']], ['Cruce', info['cruce']],
                 ['Proveedor', 'Zentrada · el Excel que deja Cowork (ofertas leídas con la sesión de Moloka)'],
                 ['Excel de Zentrada', '%s · leído %s' % (info.get('excel') or '—', info.get('leido') or '—')],
                 ['Coste de cada EAN', 'la oferta más barata que se puede pedir, tal cual: sin porte ni cuota'],
                 ['COMPRAR en ES (en «Análisis»)', dec['COMPRAR']], ['VALORAR en ES (en «Análisis»)', dec['VALORAR']]]
    filas_res += [['COMPRAR · origen «%s»' % o, por_origen[o]] for o in ORIGENES]
    filas_res += [['COMPRAR que vienen SOLO del top 300 (ni nuestro ni competidor)', solo],
                  ['Umbral de caídas (30 días)', '%s (> %d)' % (e2.texto_corte(p['umbral']), p['umbral'])],
                  ['Fichas compartidas (productos de la lista, por país)',
                   ' · '.join('%s %d' % (k, n) for k, n in (info.get('n_compartidas') or {}).items()) or '—'],
                  ['Países del filtro de ventas', ', '.join(p['paises_filtro'])],
                  ['Países que se calculan (si traen CSV)', ', '.join(p['paises_calculo'])],
                  ['Países con CSV', ', '.join(info['usados'])],
                  ['Fichas de «Fichas» (EAN → ASIN → unidades), de EAN que se pueden pedir',
                   '%d valoradas (%d por el CSV de ASIN, %d por el de EAN, %d sin dato) · %d en «Análisis» · %d filas '
                   'normales sustituidas por la de «Fichas»' % (
                       len(mapa), sum(1 for r in mapa if r.get('fuente') == 'asin'),
                       sum(1 for r in mapa if r.get('fuente') == 'ean'), sum(1 for r in mapa if not r.get('fuente')),
                       len(res_mapa), len(sustituidas))],
                  ['Packs de «Fichas» en la fila de siempre (unidades × precio)', info.get('n_packs', 0)],
                  ['Packs de Amazon detectados (la regla AM de OSMA)', texto_packs_amazon(en_hoja)],
                  [HOJA_AGOTADOS, '%d (hoja al final del libro)' % len(info.get('agotados') or [])],
                  ['Ofertas de Zentrada (el Excel)', info['n_crudo']]]
    filas_res += [['Puerta previa · %s' % NOMBRE_PREVIA[x], info['previas'][x]] for x in e2.PUERTAS_PREVIAS]
    filas_res += [['Entradas (filas de la foto: un EAN, su oferta más barata)', info['n_entradas']]]
    filas_res += [['Puerta %s · %s' % (x, e2.NOMBRE_PUERTA[x]), info['n_bd'][x]] for x in e2.PUERTAS]
    filas_res += [['Puertas previas + suma de puertas', sum(info['previas'].values()) + sum(info['n_bd'].values())],
                  ['Cuadra', 'SÍ' if info['cuadra'] else 'NO'],
                  ['Comparación con el escáner viejo', 'no se hace: el viejo no tiene escaneos de Zentrada']]
    filas_res += [['CSV', '%s · %s · %s filas%s%s' % (
        f['nombre'], f.get('pais') or '¿?', f.get('filas') or 0,
        '' if f.get('usado') else ' · NO USADO', (' · ' + f['error']) if f.get('error') else '')]
        for f in info['ficheros']]
    filas_res += [['Aviso', a] for a in info['avisos']]
    hoja('Resumen', ['Qué', 'Valor'], filas_res, {'A': 58, 'B': 90})

    hoja('Comparación', ['EAN', 'Nombre', 'Categoría', 'Viejo', 'País viejo', 'Fecha viejo (UTC)', 'Excel viejo',
                         'Puerta nuevo', 'Nuevo', 'País nuevo', 'Fecha nuevo (UTC)', 'Diferencia de criterio',
                         'Explicación'], [], {'A': 15, 'B': 50, 'C': 16, 'G': 36, 'M': 110})

    varias = []
    for res in resultados:
        if res['puerta'] == 'b':
            f = por_foto[res['foto_id']]
            for fi in res['fichas']:
                varias.append([f['ean_original'], f['nombre'], fi['pais'], fi['asin'], fi['titulo'], fi['rank'],
                               fi['caidas_30d'], fi['precio_venta']])
    hoja('Varias fichas', ['EAN', 'Nombre Zentrada', 'País', 'ASIN', 'Título Amazon', 'Puesto', 'Caídas 30 d',
                           'Precio venta (€)'], varias, {'A': 15, 'B': 50, 'E': 60})

    hoja('Puertas', COLUMNAS_PUERTAS + [COLUMNA_MAPA],
         [[por_foto[r['foto_id']]['ean_original'], por_foto[r['foto_id']]['nombre'], por_foto[r['foto_id']]['marca'],
           r.get('pa', por_foto[r['foto_id']]['precio_unidad']), '%s · %s' % (r['puerta'], e2.NOMBRE_PUERTA[r['puerta']]),
           r['motivo'], r['detalle'], None, None, None, None, None, None] for r in resultados]
         + [[r['ean_original'], r['nombre_zentrada'], None, r['pa'],
             '%s · %s' % (r['puerta'], e2.NOMBRE_PUERTA[r['puerta']]), r['motivo'], r['detalle'], None, None, None, None,
             None, '%s · %s' % (r['marca_mapa'], r['asin_mapa'])] for r in mapa],
         {'A': 15, 'B': 55, 'C': 16, 'E': 24, 'F': 16, 'G': 70, 'H': 26, 'M': 36})

    hoja('Puertas previas', ['EAN tal como vino', 'Nombre', 'Mayorista · artículo', 'Precio por unidad (€)',
                             'Puerta previa', 'Detalle'],
         [[a['ean_original'], a['nombre'], a.get('producto_heo'), a['precio_catalogo'],
           NOMBRE_PREVIA.get(a['motivo'], a['motivo']), a['detalle']] for a in apartados],
         {'A': 18, 'B': 55, 'C': 30, 'E': 32, 'F': 80})

    hoja(HOJA_AGOTADOS, COLUMNAS_AGOTADOS, info.get('agotados') or [], {'A': 15, 'B': 14, 'C': 50, 'D': 50, 'G': 28})
    salida = io.BytesIO()
    wb.save(salida)
    return salida.getvalue()


def ruta_excel(pasada, cruce, sello):
    """`zentrada/<pasada>/<cruce>/Escaner2_ZENTRADA_<AAAAMMDD_HHMM>.xlsx`: la forma que admite escaner2_cruce.ruta_excel
    desde la migracion de la v2 de este encargo."""
    return '%s/%s/%s/Escaner2_ZENTRADA_%s.xlsx' % (CARPETA, pasada, cruce, sello)
