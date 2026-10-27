# -*- coding: utf-8 -*-
"""Banco del encargo EX (27-oct-2026): «Almacén» y «FBA» en números y «También lo tiene» detrás de «Decisión».

SIN RED, SIN SECRETOS Y SIN BASE: el motor de verdad, su Celda 9 y fichas, fotos y precios INVENTADOS (repo publico).

QUE PRUEBA:
  (A) LOS NUMEROS (en_mi_bd.numeros_en_mi_bd y cabecera_fba): 17 disponibles + 3 en camino → 20; fc_transfer suma;
      foto vieja, sin foto, ASIN fuera de la foto, sin ASIN, chase, ficha sin stock_moloka y sin ficha → vacio (None),
      nunca 0; un 0 medido sigue siendo 0; la cabecera «FBA (foto 26-oct)», «FBA (foto vieja del 23-oct)», «FBA (sin
      foto)»; y `texto_en_mi_bd` (escaner viejo y Pro) NO cambia.
  (B) LA HOJA: de «Nombre» a «Decisión» (A–Q) celda a celda igual (valor, formula, estilo, enlace); detras, «Almacén»,
      «FBA (…)» y «También lo tiene»; despues TODAS las demas en su orden de antes, sin «En mi BD»; cada columna con
      sus mismos valores por nombre; las formulas que apuntaban a columnas movidas («Com. Amazon (€)» → «ISD s/ Fee
      Log.»), el formato condicional, la tabla T_Analisis, los anchos y los nombres definidos, en su columna nueva.
  (C) LA FICHA DE CADA FILA es la que la Celda 9 usaba para «En mi BD» (el `core` del registro, no el EAN de la fila):
      «Almacén» y «FBA» casan con el texto de «En mi BD» de antes en todas las filas.
  (D) LA NEGRITA de «También lo tiene»: solo los tramos de precio MENOR que el precio propio (a centimos), en texto
      enriquecido; igual no; «sin consultar», vacia o sin precio propio, texto plano; el precio de cada tramo se lee
      bien en todas las formas de escaner2_tambien.texto_tambien (caja, oferta, dia); y el precio propio que se le
      pasa (OSMA sin porte) manda sobre el «PA (€)».
  (E) LO QUE NO SE ADIVINA: una formula que apuntase a «En mi BD» → ColumnaBorrada; un libro que no sale de la Celda 9
      → error, no columnas vacias.
  (F) POR ESTRUCTURA: los ocho Excel del escaner 2 (siete Escaneo PRO/cruce y novedades) llaman al paso final DESPUES
      de «También lo tiene» y justo antes de guardar; el escaner viejo, el Pro viejo y los heredados no.
"""
import ast
import copy
import io
import os
import re
import sys
from datetime import date

from openpyxl import load_workbook
from openpyxl.cell.rich_text import CellRichText, TextBlock
from openpyxl.workbook.defined_name import DefinedName

AQUI = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, AQUI)
import en_mi_bd as E  # noqa: E402
import escaner2_motor as e2  # noqa: E402

fallos = []


def eq(nombre, obtenido, esperado):
    ok = obtenido == esperado
    if not ok:
        fallos.append(nombre)
    print(('OK ' if ok else 'XX ') + nombre + ('' if ok else '   got=%r exp=%r' % (obtenido, esperado)))


HOY = date(2026, 10, 27)


def fila_foto(asin, disp, fct=0, env=0, rec=0, fecha='2026-10-26', sku=None):
    return {'sku': sku or 'SKU-' + asin, 'asin': asin, 'available': disp, 'fc_transfer': fct, 'inbound_shipped': env,
            'inbound_receiving': rec, 'fecha_foto': fecha}


FILAS = [fila_foto('B0EX000001', 17, env=3),                      # 17 disponibles + 3 en camino → 20
         fila_foto('B0EX000002', 4, fct=3, rec=2),                # fc_transfer suma, y recibiendo tambien: 9
         fila_foto('B0EX000003', 10, sku='SKU-VIDA1'), fila_foto('B0EX000003', 5, sku='SKU-VIDA2'),   # dos vidas: 15
         fila_foto('B0EX000004', 0)]                              # un 0 MEDIDO
FOTO = E.foto_de_filas(FILAS, '2026-10-26', hoy=HOY)
VIEJA = E.foto_de_filas(FILAS, '2026-10-23', hoy=HOY)
SIN = E.foto_de_filas([], None, hoy=HOY, error='caida')


def ficha(asin, alm=5, **k):
    return dict({'ean': '8400000000017', 'asin': asin, 'stock_moloka': alm, 'es_chase': False}, **k)


# ═══════════════════════════════════════════════════════════════════════════════
# (A) LOS NUMEROS
# ═══════════════════════════════════════════════════════════════════════════════
eq('(A) 🔴 17 disponibles + 3 en camino → FBA 20; almacén de la ficha',
   E.numeros_en_mi_bd(ficha('B0EX000001', 2), FOTO), {'almacen': 2, 'fba': 20})
eq('(A) fc_transfer y recibiendo suman: 4 + 3 + 2 = 9', E.numeros_en_mi_bd(ficha('B0EX000002'), FOTO)['fba'], 9)
eq('(A) dos vidas de SKU del mismo ASIN: 10 + 5 = 15', E.numeros_en_mi_bd(ficha('B0EX000003'), FOTO)['fba'], 15)
eq('(A) un 0 medido es 0, no vacío', E.numeros_en_mi_bd(ficha('B0EX000004', 0), FOTO), {'almacen': 0, 'fba': 0})
eq('(A) 🔴 foto vieja → FBA vacío (nunca 0); el almacén sigue', E.numeros_en_mi_bd(ficha('B0EX000001', 2), VIEJA),
   {'almacen': 2, 'fba': None})
eq('(A) sin foto → FBA vacío', E.numeros_en_mi_bd(ficha('B0EX000001', 2), SIN), {'almacen': 2, 'fba': None})
eq('(A) foto sin poner (None) → FBA vacío', E.numeros_en_mi_bd(ficha('B0EX000001', 2), None)['fba'], None)
eq('(A) ASIN fuera de la foto → FBA vacío', E.numeros_en_mi_bd(ficha('B0NOESTA00'), FOTO)['fba'], None)
eq('(A) sin ASIN (o en blanco) → FBA vacío', (E.numeros_en_mi_bd(ficha(None), FOTO)['fba'],
                                             E.numeros_en_mi_bd(ficha('  '), FOTO)['fba']), (None, None))
eq('(A) 🔴 chase (aunque traiga ASIN) → FBA vacío', E.numeros_en_mi_bd(ficha('B0EX000001', es_chase=True), FOTO)['fba'],
   None)
eq('(A) ficha sin stock_moloka (nulo o no pedido) → Almacén vacío, nunca 0',
   (E.numeros_en_mi_bd(ficha('B0EX000001', None), FOTO), E.numeros_en_mi_bd({'asin': 'B0EX000001'}, FOTO)),
   ({'almacen': None, 'fba': 20}, {'almacen': None, 'fba': 20}))
eq('(A) sin ficha → los dos vacíos', (E.numeros_en_mi_bd(None, FOTO), E.numeros_en_mi_bd({}, FOTO)),
   ({'almacen': None, 'fba': None}, {'almacen': None, 'fba': None}))
eq('(A) cabecera con la fecha de la foto, una vez', E.cabecera_fba(FOTO), 'FBA (foto 26-oct)')
eq('(A) cabecera de la foto vieja', E.cabecera_fba(VIEJA), 'FBA (foto vieja del 23-oct)')
eq('(A) cabecera sin foto (ilegible, sin poner)', (E.cabecera_fba(SIN), E.cabecera_fba(None)),
   ('FBA (sin foto)', 'FBA (sin foto)'))
eq('(A) el día sin cero delante, como «También lo tiene»', E.dia_mes(date(2026, 11, 5)), '5-nov')
eq('(A) 🔒 texto_en_mi_bd (escáner viejo y Pro) NO cambia',
   (E.texto_en_mi_bd(ficha('B0EX000001', 2), FOTO), E.texto_en_mi_bd(ficha('B0EX000001', 2), VIEJA)),
   ('OK Alm:2 FBA:17 +3 en camino (foto 2026-10-26)', 'OK Alm:2 FBA: foto vieja del 23-10'))

# ═══════════════════════════════════════════════════════════════════════════════
# (B) y (C) LA HOJA: un libro de la Celda 9 de verdad, con las columnas que añade el escaner 2 detrás
# ═══════════════════════════════════════════════════════════════════════════════
M = e2.cargar_motor()
FICHAS = [{'ean': '8400000000017', 'asin': 'B0EX000001', 'stock_moloka': 2, 'es_chase': False},
          {'ean': '8400000000024', 'asin': 'B0EX000002', 'stock_moloka': None, 'es_chase': False},
          {'ean': '8400000000031', 'asin': None, 'stock_moloka': 7, 'es_chase': True},
          {'ean': '8400000000048', 'asin': 'B0NOESTA00', 'stock_moloka': 0, 'es_chase': False}]
M.poner_catalogo_propio(FICHAS)
M.poner_foto_fba(FOTO)


def reg(n, ean, asin, pa, margen, core=None, decision='COMPRAR', almacen=None):
    r = {'nombre': 'Producto %d' % n, 'ean': ean, 'asin': asin, 'marca': 'Marca', 'core': core or ean, '_pa_efectivo': pa,
         'ambiguo': False, 'titulo_amz': 'Título %d' % n, 'coincide': 'SI', 'coherencia_caja': None,
         'url': 'https://example.invalid/%d' % n, 'volumen': None, '_margen_es': margen,
         '_paises_calc': {'ES': {'precio': 20.0 + n, 'canal': 'FBA', 'ref_pct': 0.15, 'fee': 3.1, 'iva': 0.21,
                                 'decision': decision, 'margen': margen, 'caidas_30d': 9 + n},
                          'FR': {'precio': 21.0 + n, 'canal': 'FBA', 'ref_pct': 0.15, 'fee': 3.3, 'iva': 0.2,
                                 'decision': 'VALORAR', 'margen': 0.05, 'caidas_30d': 0}}}
    if almacen:
        r['_almacen'] = almacen
    return r


REGISTROS = [reg(1, '8400000000017', 'B0EX000001', 6.0, 0.3, almacen=(0.21, 'envios')),
             reg(2, '8400000000024', 'B0EX000002', 7.0, 0.2, almacen=(0.15, 'global')),
             # 🔑 OSMA/Zentrada: la fila dice el EAN del proveedor y la ficha es la NUESTRA (otro EAN, el `core`)
             reg(3, '8400000000994', 'B0EX000001', 5.5, 0.15, core='8400000000017', almacen=(0.21, 'envios')),
             reg(4, '8400000000031', 'B0EX000009', 8.0, 0.1, decision='VALORAR'),        # chase
             reg(5, '8400000000048', 'B0NOESTA00', 9.0, 0.05, decision='NO COMPRAR'),    # ASIN fuera de la foto
             reg(6, '8400000000055', 'B0EX000006', 4.0, 0.01, decision='NO COMPRAR')]    # no es nuestro
DATOS = {'registros': REGISTROS, 'problematicos': [], 'no_encontrados': [], 'chase_sueltos': [], '_dups': [],
         'ambiguos': [], 'sin_rank': [], 'chase_pendientes': [], 'cotejo_info': {}, 'PROVEEDOR': 'HEO'}
TL = {'8400000000017': 'TCG 5,95 € oferta · OcioStock 6,00 € · BEMS 7,40 € (10-oct)',
      '8400000000024': 'HEO 4,52 € caja de 6 · BEMS 6,48 €',
      '8400000000994': 'DBLine 5,49 €'}


def libro(datos=DATOS, tambien=True):
    wb = e2.escribir_celda9(copy.deepcopy(datos), M)
    e2.poner_columna_ficha_compartida(wb, [], [])
    e2.poner_columna_amazon(wb, lambda asins: {})
    if tambien:
        e2.poner_columna_tambien(wb, lambda eans: {e: TL.get(e) for e in eans})
    return wb


def guardado(wb, rich=False):
    b = io.BytesIO()
    wb.save(b)
    return load_workbook(io.BytesIO(b.getvalue()), rich_text=rich)


ANTES = guardado(libro())
wb = libro()
# Un nombre definido que apunta a «Análisis» (columna «Origen IVA», que se corre) y otro a otra hoja (no se toca).
wb.defined_names['EX_Prueba'] = DefinedName('EX_Prueba', attr_text="'Análisis'!$AA$2:$AA$9")
wb.defined_names['EX_Otra'] = DefinedName('EX_Otra', attr_text="'Descartados'!$B$2")
cuenta = e2.poner_almacen_fba_tambien(wb, M)
DESP = guardado(wb)
a, d = ANTES['Análisis'], DESP['Análisis']
ca, cd = [c.value for c in a[1]], [c.value for c in d[1]]
i_dec = ca.index('Decisión')
eq('(B) 🔴 «Decisión» sigue en la Q', (ca.index('Decisión'), cd.index('Decisión')), (16, 16))
eq('(B) 🔴 detrás de «Decisión»: «Almacén», «FBA (foto 26-oct)» y «También lo tiene»', cd[17:20],
   ['Almacén', 'FBA (foto 26-oct)', 'También lo tiene'])
eq('(B) 🔴 y luego TODAS las demás, en su orden de antes, sin «En mi BD»', cd[20:],
   [c for c in ca[17:] if c not in ('En mi BD', 'También lo tiene')])
eq('(B) «En mi BD» ya no está', 'En mi BD' in cd, False)


def _rgb(color):
    v = getattr(color, 'rgb', None)
    return v if isinstance(v, str) else None


def por_nombre(v, cab):
    """Una formula con cada referencia escrita por el NOMBRE de su columna («=[Precio venta (€)]2*…»): la misma formula
    antes y despues aunque su columna haya cambiado de letra."""
    if not (isinstance(v, str) and v.startswith('=')):
        return v
    from openpyxl.utils import column_index_from_string
    return re.sub(r'(?<![A-Za-z])([A-Z]{1,3})(\d+)', lambda m: '[%s]%s' % (cab[column_index_from_string(m.group(1)) - 1],
                                                                        m.group(2)), v)


def huella_celda(c, cab):
    return (por_nombre(c.value, cab), c.number_format, c.font.b, _rgb(c.font.color), _rgb(c.fill.fgColor),
            c.hyperlink.target if c.hyperlink is not None else None)


eq('(B) 🔴 A–Q celda a celda iguales (valor, fórmula por nombre de columna, formato, letra, relleno, enlace)',
   [[huella_celda(d.cell(r, c), cd) for c in range(1, 18)] for r in range(1, d.max_row + 1)],
   [[huella_celda(a.cell(r, c), ca) for c in range(1, 18)] for r in range(1, a.max_row + 1)])
_movidas = {}
for j, col in enumerate(ca, 1):
    if col in ('En mi BD',) or j <= 17:
        continue
    k = cd.index(col) + 1
    _movidas[col] = ([a.cell(r, j).value for r in range(2, a.max_row + 1)] == [d.cell(r, k).value for r in range(2, d.max_row + 1)])
eq('(B) cada columna movida lleva sus mismos valores (por nombre)', [c for c, ok in _movidas.items() if not ok], [])
_c = {n: cd.index(n) + 1 for n in cd}
from openpyxl.utils import get_column_letter as _L  # noqa: E402
eq('(B) 🔴 «Com. Amazon (€)» sigue sumando «ISD s/ Fee Log. (€)» en su columna NUEVA',
   d.cell(2, _c['Com. Amazon (€)']).value, '=%s2*%s2+%s2' % (_L(_c['Precio venta (€)']), _L(_c['% Comisión']),
                                                              _L(_c['ISD s/ Fee Log. (€)'])))
eq('(B) …y antes apuntaba a otra letra (la prueba distingue)', a.cell(2, ca.index('Com. Amazon (€)') + 1).value ==
   d.cell(2, _c['Com. Amazon (€)']).value, False)
eq('(B) el beneficio, que solo usa A–Q, igual', d.cell(2, _c['Beneficio (€)']).value,
   a.cell(2, ca.index('Beneficio (€)') + 1).value)
_t = d.tables['T_Analisis']
eq('(B) la tabla T_Analisis cubre todas las columnas, con sus nombres en su orden',
   (_t.ref, [x.name for x in _t.tableColumns]), ('A1:%s%d' % (_L(len(cd)), d.max_row), [str(x) for x in cd]))
_cf = {str(cf.sqref): [r.formula for r in reglas] for cf, reglas in d.conditional_formatting._cf_rules.items()}
_cfa = {str(cf.sqref): [r.formula for r in reglas] for cf, reglas in a.conditional_formatting._cf_rules.items()}


def _cf_de(col, hoja_cab, cf):
    letra = _L(hoja_cab.index(col) + 1)
    return next(v for k, v in cf.items() if k.startswith(letra + '2:' + letra))


eq('(B) el semáforo de «Decisión», el de «Margen»: los mismos', (_cf_de('Decisión', cd, _cf), _cf_de('Margen', cd, _cf)),
   (_cf_de('Decisión', ca, _cfa), _cf_de('Margen', ca, _cfa)))
eq('(B) 🔴 «Coincide», «Cotejo» y «Coherencia caja» se corren y su regla apunta a SU columna nueva',
   [_cf_de(c, cd, _cf)[0][0] for c in ('Coincide', 'Cotejo', 'Coherencia caja')],
   ['ISNUMBER(SEARCH("NO",%s2))' % _L(_c['Coincide']), 'ISNUMBER(SEARCH("DUDOSO",%s2))' % _L(_c['Cotejo']),
    'ISNUMBER(SEARCH("INCOHERENTE",%s2))' % _L(_c['Coherencia caja'])])
eq('(B) la banda gris de cada producto va de «Nombre» a «Origen IVA» (con las tres nuevas dentro)',
   [k for k, v in _cf.items() if v == [['ISODD(INT((ROW()-2)/4))']]], ['A2:%s%d' % (_L(_c['Origen IVA']), d.max_row)])
eq('(B) anchos: «Almacén» 10, «FBA» 24, «También lo tiene» 60, y el de «Amazon» viaja con ella',
   [d.column_dimensions[_L(_c[n])].width for n in ('Almacén', 'FBA (foto 26-oct)', 'También lo tiene', 'Amazon')],
   [10, 24, 60, 40])
eq('(B) cabeceras nuevas en negrita, como las demás', [d.cell(1, _c[n]).font.b for n in ('Almacén', 'FBA (foto 26-oct)')],
   [True, True])
eq('(B) el nombre definido que apuntaba a «Origen IVA» la sigue; el de otra hoja, intacto',
   (DESP.defined_names['EX_Prueba'].attr_text, DESP.defined_names['EX_Otra'].attr_text),
   ("'Análisis'!$%s$2:$%s$9" % (_L(_c['Origen IVA']), _L(_c['Origen IVA'])), "'Descartados'!$B$2"))
eq('(B) el enlace del ASIN y el de «OcioStock» siguen en su celda', (d.cell(2, 3).hyperlink.target,
   d.cell(2, _c['OcioStock']).hyperlink.target), (a.cell(2, 3).hyperlink.target, 'https://example.invalid/1'))
eq('(B) las otras hojas, celda a celda iguales',
   [[[c.value for c in r] for r in ANTES[h].iter_rows()] for h in ANTES.sheetnames[1:]],
   [[[c.value for c in r] for r in DESP[h].iter_rows()] for h in DESP.sheetnames[1:]])

# (C) La ficha de cada fila: la de «En mi BD».
_esperado, _obtenido = [], []
_bd = ca.index('En mi BD') + 1
for r in range(2, a.max_row + 1):
    _esperado.append((a.cell(r, 2).value, a.cell(r, 6).value, a.cell(r, _bd).value))
    _obtenido.append((d.cell(r, 2).value, d.cell(r, 6).value, d.cell(r, _c['Almacén']).value,
                      d.cell(r, _c['FBA (foto 26-oct)']).value))
_por_ean = {(o[0], o[1]): (o[2], o[3]) for o in _obtenido}
eq('(C) 🔴 «En mi BD» de antes → los números de ahora, fila a fila (ES)',
   [(e[2], _por_ean[(e[0], e[1])]) for e in _esperado if e[1] == 'ES'],
   [('OK Alm:2 FBA:17 +3 en camino (foto 2026-10-26)', (2, 20)),
    ('OK Alm: — FBA:7 +2 en camino (foto 2026-10-26)', (None, 9)),   # 4 + 3 de fc_transfer, + 2 recibiendo
    ('OK Alm:2 FBA:17 +3 en camino (foto 2026-10-26)', (2, 20)),     # EAN del proveedor, ficha nuestra por el core
    ('OK Alm:7 FBA: —', (7, None)),                                    # chase: almacén sí, FBA nunca
    ('OK Alm:0 FBA: — (foto 2026-10-26)', (0, None)),                  # ASIN fuera de la foto
    (None, (None, None))])                                             # no es nuestro: las dos vacías
eq('(C) cada país de un producto lleva los mismos números', len({(o[0], o[2], o[3]) for o in _obtenido}),
   len({o[0] for o in _obtenido}))
eq('(C) recuento que devuelve el paso (cuatro países por producto)', cuenta, {'almacen': 4 * 4, 'fba': 4 * 3,
                                                                             'negrita': 4 * 3})
eq('(C) las celdas con número llevan formato de unidades', d.cell(2, _c['Almacén']).number_format, '0')

_wv = libro()
_m2 = e2.cargar_motor()
_m2.poner_catalogo_propio(FICHAS)
_m2.poner_foto_fba(VIEJA)
e2.poner_almacen_fba_tambien(_wv, _m2)
_hv = _wv['Análisis']
_cv = [c.value for c in _hv[1]]
eq('(C) 🔴 foto vieja: cabecera «FBA (foto vieja del 23-oct)» y TODAS sus celdas vacías; el almacén sigue',
   (_cv[18], {_hv.cell(r, 19).value for r in range(2, _hv.max_row + 1)}, _hv.cell(2, 18).value),
   ('FBA (foto vieja del 23-oct)', {None}, 2))
_m3 = e2.cargar_motor()
_m3.poner_catalogo_propio(FICHAS)
_ws = libro()
e2.poner_almacen_fba_tambien(_ws, _m3)
eq('(C) sin foto puesta: «FBA (sin foto)» y vacías', ([c.value for c in _ws['Análisis'][1]][18],
   {_ws['Análisis'].cell(r, 19).value for r in range(2, _ws['Análisis'].max_row + 1)}), ('FBA (sin foto)', {None}))
_wn = libro(tambien=False)
e2.poner_almacen_fba_tambien(_wn, M)
_cn = [c.value for c in _wn['Análisis'][1]]
eq('(C) sin lector de «También lo tiene» (no hay columna): solo «Almacén» y «FBA» detrás de «Decisión»',
   (_cn[17:19], 'También lo tiene' in _cn, _cn[19]), (['Almacén', 'FBA (foto 26-oct)'], False, 'EAN ambiguo'))
_vacio = dict(DATOS, registros=[])
_we = e2.escribir_celda9(copy.deepcopy(_vacio), M)
e2.poner_columna_tambien(_we, lambda eans: {})
e2.poner_almacen_fba_tambien(_we, M)
eq('(C) un Excel sin filas (sin tabla) también sale, con la cabecera nueva',
   [c.value for c in guardado(_we)['Análisis'][1]][16:20], ['Decisión', 'Almacén', 'FBA (foto 26-oct)', 'También lo tiene'])

# ═══════════════════════════════════════════════════════════════════════════════
# (D) LA NEGRITA
# ═══════════════════════════════════════════════════════════════════════════════
DR = guardado(wb, rich=True)['Análisis']
_ct = _c['También lo tiene']


def tramos(v):
    if isinstance(v, CellRichText):
        return [('N' if isinstance(t, TextBlock) and t.font.b else '-', t.text if isinstance(t, TextBlock) else t)
                for t in v]
    return v


eq('(D) 🔴 PA 6,00: TCG 5,95 en negrita; OcioStock 6,00 (igual) y BEMS 7,40 no; el texto, el de siempre',
   (tramos(DR.cell(2, _ct).value), str(DR.cell(2, _ct).value)),
   ([('N', 'TCG 5,95 € oferta'), ('-', ' · OcioStock 6,00 € · BEMS 7,40 € (10-oct)')], TL['8400000000017']))
eq('(D) PA 7,00: HEO 4,52 (caja de 6) y BEMS 6,48, los dos en negrita, el separador sin negrita',
   tramos(DR.cell(2 + 4, _ct).value), [('N', 'HEO 4,52 € caja de 6'), ('-', ' · '), ('N', 'BEMS 6,48 €')])
eq('(D) PA 5,50 y DBLine 5,49: un céntimo por debajo, negrita', tramos(DR.cell(2 + 8, _ct).value),
   [('N', 'DBLine 5,49 €')])
eq('(D) a céntimos: 6,004 de PA y 6,00 de otro → no; 6,006 → sí (se ve 6,01)',
   (e2.tambien_en_negrita('X 6,00 €', 6.004), tramos(e2.tambien_en_negrita('X 6,00 €', 6.006))),
   ('X 6,00 €', [('N', 'X 6,00 €')]))
eq('(D) «sin consultar», vacía, sin precio propio o a 0: texto plano',
   (e2.tambien_en_negrita('sin consultar', 9.0), e2.tambien_en_negrita(None, 9.0), e2.tambien_en_negrita('', 9.0),
    e2.tambien_en_negrita('TCG 1,00 €', None), e2.tambien_en_negrita('TCG 1,00 €', 0)),
   ('sin consultar', None, '', 'TCG 1,00 €', 'TCG 1,00 €'))
import escaner2_tambien as etl  # noqa: E402
_of = [{'proveedor': 'TCG', 'precio': 6.95, 'oferta': True, 'dia': None, 'caja': None},
       {'proveedor': 'HEO', 'precio': 1234.5, 'oferta': False, 'dia': date(2026, 10, 10), 'caja': 6},
       {'proveedor': 'ZENTRADA', 'precio': 0.07, 'oferta': False, 'dia': HOY, 'caja': 12},
       {'proveedor': 'OSMA', 'precio': 12.0, 'oferta': True, 'dia': date(2026, 9, 30), 'caja': None}]
_txt = etl.texto_tambien(_of, HOY)
eq('(D) 🔴 el precio de cada tramo se lee bien en todas las formas de texto_tambien (caja, oferta, día, miles)',
   [str(e2.precio_de_parte_tambien(p)) for p in _txt.split(e2.SEPARADOR_TAMBIEN)], ['0.07', '6.95', '12.00', '1234.50'])
eq('(D) …y el separador es el suyo', _txt.count(e2.SEPARADOR_TAMBIEN), 3)
_wp = libro()
_k = ('8400000000017', 'B0EX000001')
e2.poner_almacen_fba_tambien(_wp, M, precio_propio={_k: 5.80})
eq('(D) 🔴 el precio propio que se pasa (OSMA: sin porte) manda sobre el «PA (€)»: con 5,80, TCG 5,95 ya no',
   guardado(_wp, rich=True)['Análisis'].cell(2, _ct).value, TL['8400000000017'])
eq('(D) …y una fila sin precio propio en el mapa, sin negrita (no se adivina)',
   guardado(_wp, rich=True)['Análisis'].cell(6, _ct).value, TL['8400000000024'])

# ═══════════════════════════════════════════════════════════════════════════════
# (E) LO QUE NO SE ADIVINA
# ═══════════════════════════════════════════════════════════════════════════════
_wx = libro()
_hx = _wx['Análisis']
_hx.cell(2, len([c.value for c in _hx[1]]) + 1, value='=R2&"x"')     # una formula que lee «En mi BD» (R)
try:
    e2.poner_almacen_fba_tambien(_wx, M)
    _r = 'no lanzó'
except e2.ColumnaBorrada as ex:
    _r = 'ColumnaBorrada' if 'En mi BD' in str(ex) else str(ex)
eq('(E) 🔴 una fórmula que apunta a «En mi BD» → ColumnaBorrada (no se adivina)', _r, 'ColumnaBorrada')
from openpyxl import Workbook  # noqa: E402
_wz = Workbook()
_wz.active.title = 'Análisis'
_wz.active.append(['EAN', 'ASIN', 'Decisión', 'En mi BD'])
try:
    e2.poner_almacen_fba_tambien(_wz, M)
    _r = 'no lanzó'
except ValueError as ex:
    _r = 'escribir_celda9' in str(ex)
eq('(E) un libro que no sale de la Celda 9 → error, no columnas vacías', _r, True)

# ═══════════════════════════════════════════════════════════════════════════════
# (F) POR ESTRUCTURA: quién llama al paso final y dónde
# ═══════════════════════════════════════════════════════════════════════════════
def _llamadas(nodo):
    salida = []
    for n in ast.walk(nodo):
        if isinstance(n, ast.Call):
            f = n.func
            nombre = f.attr if isinstance(f, ast.Attribute) else (f.id if isinstance(f, ast.Name) else None)
            salida.append((n.lineno, n.col_offset, nombre))
    return sorted(salida)


ESCRITORES = {'escaner2_heo_cruce.py': 'escribir_excel', 'escaner2_bems_pro.py': 'escribir_excel',
              'escaner2_dbline_pro.py': 'escribir_excel', 'escaner2_ociostock_pro.py': 'escribir_excel',
              'escaner2_tcg_pro.py': 'escribir_excel', 'escaner2_osma.py': 'escribir_excel',
              'escaner2_zentrada.py': 'escribir_excel', 'escaner2_novedades.py': 'excel_de_novedades'}
_bien = {}
for fichero, funcion in ESCRITORES.items():
    with io.open(os.path.join(AQUI, fichero), encoding='utf-8') as fh:
        arbol = ast.parse(fh.read())
    fn = [n for n in arbol.body if isinstance(n, ast.FunctionDef) and n.name == funcion]
    seq = [x[2] for x in _llamadas(fn[0])] if len(fn) == 1 else []
    # La secuencia: «También lo tiene» → el paso final → save, con nada de la hoja entre el paso y el save.
    try:
        i_t, i_p, i_s = seq.index('poner_columna_tambien'), seq.index('poner_almacen_fba_tambien'), seq.index('save')
        _bien[fichero] = (seq.count('poner_almacen_fba_tambien'), i_t < i_p < i_s,
                          [x for x in seq[i_p + 1:i_s] if x not in ('BytesIO', 'precio_propio_excel')])
    except ValueError:
        _bien[fichero] = None
eq('(F) 🔴 los ocho Excel llaman al paso final UNA vez, tras «También lo tiene» y justo antes de guardar',
   _bien, {f: (1, True, []) for f in ESCRITORES})
_fuera = {}
for fichero in ('moloka_escaner_nube.py', 'moloka_escaner_pro.py', 'escaner2_heredado_nube.py',
                'escaner2_heredado_pro.py', 'escaner2_desvios.py'):
    with io.open(os.path.join(AQUI, fichero), encoding='utf-8') as fh:
        _fuera[fichero] = 'poner_almacen_fba_tambien' in fh.read()
eq('(F) el escáner viejo, el Pro viejo, los heredados y los desvíos no lo llaman (siguen con «En mi BD»)',
   _fuera, {f: False for f in _fuera})

print()
if fallos:
    print('ROJO: %d fallos' % len(fallos))
    for f in fallos:
        print('  · ' + f)
    sys.exit(1)
print('VERDE: «Almacén» y «FBA» en números y «También lo tiene» detrás de «Decisión», con A–Q intactas, las fórmulas, '
      'la tabla y el formato en su sitio, y la negrita solo en el proveedor más barato.')
