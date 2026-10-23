# -*- coding: utf-8 -*-
"""ESCANER 2 · ESCANEO PRO DE TCG · LA PARTE PURA (encargo T3, 14-oct-2026; copiada de la de OcioStock, encargo OC3)

QUE ES. El calculo del Escaneo PRO de TCG, sin red y sin base: la foto de la pasada desde la FOTO DE DISPONIBILIDAD de
TCG (disp_estado, que deja escaner2_tcg_disponibilidad.py, encargo T2: el Excel de TCG con el precio del tramo mas
barato de su web ya aplicado), la lista para el Visualizador y el Excel. La red vive en los dos programas que lo usan:
  · escaner2_tcg_barrido.py (workflow escaner2-tcg-barrido.yml): de la foto a la pasada y la lista.
  · escaner2_tcg_cruce.py   (workflow escaner2-tcg-cruce.yml):   los 4 CSV, las puertas y el Excel.

🔑 CADA ESCANER, A LA MEDIDA DE SU PROVEEDOR (Fernando, 29-sep-2026: «Quiero que cada escaner sea personalizado a cada
   proveedor»): leer y filtrar es de TCG y vive aqui. Lo que se REUTILIZA sin tocarlo son las REGLAS DE COMPRA del PRO
   de HEO (`escaner2_motor.decidir`, con el corte y los paises de escaner2_parametros: TCG = 6 y ES·IT·FR·DE) y su
   Excel (la Celda 9 del viejo con «Ventas», la columna «Ficha compartida» y las hojas del escaner 2 en el MISMO orden).

LO PROPIO DE TCG (decidido por Fernando; partes T2 y encargo T3):
  1. LA FOTO SALE DE disp_estado (las filas vistas en la ultima pasada APLICADA de TCG). Nada de la web ni del Excel.
  2. SOLO LO QUE SE PUEDE COMPRAR YA («Quiero ver solo lo que esta disponible», 12-oct): `disponible` y sin preventa.
     Backorder («En reposición») y PreOrder (preventa) ya llegan con `disponible` falso; la preventa se mira igual.
  3. CHASE SUELTO, FUERA si lo hubiera (un chase que no es caja). El Excel de TCG no los trae: la «C» es la CAJA DE 6
     con precio por figura (es_caja, y es_chase si dice «5 + 1»): esa SI entra, valorada con el EAN de la figura.
  4. EL PRECIO, EL DEL TRAMO MAS BARATO («siempre hay que calcular con el tramo mas barato», 8-oct): PA =
     `precio_escalon` de la foto (ya lleva el cociente del tramo, y en la caja es por figura), SIN NINGUN DESCUENTO
     AÑADIDO (el × 0,99 de la transferencia es solo de OcioStock). `uds_escalon` dice cuantas pedir para ese precio
     (en la caja, FIGURAS: D7 de T2). Una fila con el aviso «sin tramo leído» entra con su precio suelto (es el que ya
     trae `precio_escalon`) y el aviso se ve en el Excel.
  5. SIN PRECIO DE COMPRA (vacio o 0), FUERA: un PA de 0 daria un COMPRAR falso. Va a `estado_no_servible`, listada.
  6. EAN CON FORMA RARA (la regla 'ean_forma_rara' de la foto: 14 cifras, el GS1 largo, colas raras, codigos cortos),
     FUERA y listada, SALVO el GTIN-14 que empieza por 0: es el MISMO codigo que el EAN-13 de detras (mismo digito de
     control), y entra con esas 13 cifras (decidido por Code, a revisar: parte T3).
  7. UNA FILA POR EAN BASE: la de menor PA (a igualdad, la suelta antes que la caja, y luego la llave menor); si gana
     una caja, «caja de 6». La otra, a `duplicado_proveedor`, listada.
  8. EL CRUDO SON LOS ARTICULOS DEL EXCEL (una llave = un EAN crudo), no sus lineas: la pasada de disponibilidad ya se
     quedo con UNA de cada EAN repetido (n_duplicados). crudo = n_leidas + n_sin_gtin; el Excel dice las dos cifras.
  9. MODOS: 'marcas' = las de la fila TCG de `reglas_director` (SOLO LECTURA: `marcas` y `marcas_es`) con la regla del
     viejo (el nombre de la regla DENTRO del de la fila, sin mayusculas), pero Magic y Yu-Gi-Oh (`marcas_es`) en
     CUALQUIER idioma y SIN el corte de rank del viejo (Fernando, 13-oct: «El escaner pro es de todas las marcas y yo
     elijo»); 'todas' (tambien las sin marca); 'elegidas' (las del selector, coincidencia EXACTA, como HEO).
 10. EL EXCEL, EL DEL PRO DE HEO, mismo formato y orden, con TRES columnas mas AL FINAL de «Análisis»: «Uds. escalón»,
     «Precio unidad» (el del Excel de TCG, por figura, sin tramo) y «Aviso del tramo» («sin tramo leído» o «tramo leído
     hace más de 7 días»). Sin «Chase_manual» (es de HEO) ni comparacion con el viejo.

🔴 REPO PUBLICO: aqui no hay ni un precio ni un EAN reales; los del banco son inventados (Seguridad 21).
"""
import io
import re

import escaner2_motor as e2

PROVEEDOR = 'TCG'
CARPETA = 'tcg'
NOMBRE = 'TCG'
MODOS = ('marcas', 'todas', 'elegidas')
UDS_CAJA = 6
TEXTO_CAJA = 'caja de 6'

# Las columnas nuevas de «Análisis», AL FINAL (detras de «Ficha compartida»): hay quien lee la hoja por letra.
COLUMNA_UDS_ESCALON = 'Uds. escalón'
COLUMNA_PRECIO_UNIDAD = 'Precio unidad'
COLUMNA_AVISO_TRAMO = 'Aviso del tramo'
COLUMNAS_ESCALON = [COLUMNA_UDS_ESCALON, COLUMNA_PRECIO_UNIDAD, COLUMNA_AVISO_TRAMO]
ANCHOS_ESCALON = [12, 13, 30]
# Los avisos de la foto de TCG (escaner2_tcg.AVISO_SIN_TRAMO y AVISO_TRAMO_VIEJO): se copian, no se importan, para que
# este modulo no dependa del lector del Excel.
AVISO_SIN_TRAMO = 'sin tramo leído'

# Los nombres de las puertas previas EN TCG (los mismos que NOMBRE_PREVIA_TCG de la v2). Las ocho de siempre
# (crudo = previas + foto); `chase_funko` sale siempre a 0.
NOMBRE_PREVIA = dict(e2.NOMBRE_PUERTA_PREVIA,
                     sin_gtin='Sin EAN',
                     no_disponible='Sin stock, en reposición o en preventa',
                     marca_fuera='Marca fuera de las marcas del director',
                     estado_no_servible='Sin precio de compra',
                     chase_suelto='Chase suelto (no se compra)',
                     ean_forma_rara='EAN con forma rara',
                     duplicado_proveedor='Mismo EAN que otra fila de TCG, más cara')

_RE_GTIN14_CERO = re.compile(r'0(\d{13})')


class FalloTCG(Exception):
    """Algo que impide seguir: la pasada o el cruce quedan 'fallida' con este motivo (privado, en la base)."""


def nombre_previa(p, modo):
    """El nombre de una puerta previa EN SU PASADA: en el modo «elegidas» la marca que se cae es «no elegida»."""
    if p == 'marca_fuera' and modo == 'elegidas':
        return e2.NOMBRE_MARCA_NO_ELEGIDA
    return NOMBRE_PREVIA[p]


# ═══════════════════════════════════════════════════════════════════════════════
# 1 · EL FILTRO DE MARCA, POR MODO
# ═══════════════════════════════════════════════════════════════════════════════
def marcas_del_director(regla):
    """Las marcas del modo 'marcas' desde la fila TCG de reglas_director: `marcas` y despues `marcas_es` (Magic y
    Yu-Gi-Oh), estas SIN su idioma. Nada mas de la regla: ni `rank_maximo` ni `incluir_estados`."""
    regla = regla or {}
    salida = []
    for m in list(regla.get('marcas') or []) + list(regla.get('marcas_es') or []):
        m = str(m or '').strip()
        if m and m not in salida:
            salida.append(m)
    return salida


def filtro_marca(modo, marcas_director=None, elegidas=None):
    """(quiere, info) para el modo:
      · 'marcas'   · las del director (`marcas_del_director`) con la regla del viejo: el nombre de la marca de la regla
                     CONTENIDO en el de la fila, sin distinguir mayusculas (`_pasa_filtros` de moloka_escaner_nube.py),
                     y sin mirar el idioma;
      · 'todas'    · todas (las sin marca, tambien);
      · 'elegidas' · las del selector de la v2, COINCIDENCIA EXACTA (`escaner2_motor.clave_marca`), como HEO.
    `quiere(marca)` → bool. `info` lleva las marcas que se guardan en la pasada (None en 'todas')."""
    if modo == 'todas':
        return (lambda _m: True), {'marcas_reales': None}
    if modo == 'marcas':
        lista = [str(m).strip() for m in (marcas_director or []) if str(m or '').strip()]
        if not lista:
            raise FalloTCG('la fila TCG de reglas_director no trae marcas: no se sabe qué barrer')
        bajas = [m.lower() for m in lista]
        return (lambda m: any(b in str(m or '').lower() for b in bajas)), {'marcas_reales': lista}
    if modo == 'elegidas':
        claves = frozenset(e2.clave_marca(m) for m in (elegidas or []))
        if not claves:
            raise FalloTCG('modo elegidas sin ninguna marca')
        return (lambda m: e2.clave_marca(m) in claves), {'marcas_reales': list(elegidas)}
    raise FalloTCG('modo de barrido desconocido: %r' % (modo,))


# ═══════════════════════════════════════════════════════════════════════════════
# 2 · LA FOTO DE LA PASADA, DESDE LA FOTO DE DISPONIBILIDAD
# ═══════════════════════════════════════════════════════════════════════════════
def _num(x):
    return None if x is None else float(x)


def _apartado(f, motivo, detalle):
    return {'producto_heo': str(f.get('producto_prov') or ''), 'ean_original': str(f.get('ean_original') or ''),
            'nombre': f.get('nombre') or '', 'marca': (f.get('marca') or '').strip(),
            'precio_catalogo': _num(f.get('precio_unidad')), 'motivo': motivo, 'detalle': detalle, 'en_oferta': None}


def ean_de_la_foto(f):
    """(ean_core, rescatado) con que la fila entra al cruce, o (None, False) si su EAN no sirve (regla 6). El
    ean_core de la foto (12 o 13 cifras) tal cual; con la regla 'ean_forma_rara', solo el GTIN-14 con un 0 delante, que
    es el EAN-13 de detras: mismas cifras y mismo digito de control."""
    regla = f.get('regla')
    core = str(f.get('ean_core') or '').strip()
    if not regla:
        return (core, False) if core.isdigit() and len(core) in (12, 13) else (None, False)
    if regla == 'ean_forma_rara':
        m = _RE_GTIN14_CERO.fullmatch(str(f.get('ean_original') or '').strip())
        if m:
            return m.group(1), True
    return None, False


def _gana(a, b):
    """¿`a` va antes que `b` para quedarse con el EAN? Menor PA (precio_escalon; sin PA, la ultima); a igualdad, la
    suelta antes que la caja (se pide de una en una), y despues la llave menor (para que salga siempre la misma)."""
    def clave(f):
        pa = f.get('precio_escalon')
        return (pa is None, float(pa) if pa is not None else 0.0, bool(f.get('es_caja')), str(f.get('producto_prov') or ''))
    return clave(a) < clave(b)


def construir_foto(filas, quiere_marca, M, n_crudo, n_sin_gtin, modo='marcas'):
    """De la foto de disponibilidad de TCG a la foto de la pasada, sin perder a nadie: cada fila sale por UNA puerta
    previa o entra en la foto, y crudo (los articulos del Excel) = previas + foto. En este orden:
      1. no se puede comprar ya (no disponible, o preventa) → `no_disponible` (se cuenta);
      2. marca que no pasa el filtro del modo → `marca_fuera` (se lista);
      3. chase suelto (chase que no es caja) → `chase_suelto` (se lista);
      4. EAN que no sirve (`ean_de_la_foto`) → `ean_forma_rara` (se lista);
      5. sin precio de compra (precio_escalon vacio o 0) → `estado_no_servible` (se lista);
      6. dos filas con el mismo EAN base → la de menor PA (`_gana`); la otra → `duplicado_proveedor` (se lista).
    `filas`: disp_estado de TCG vistas en la ultima pasada aplicada; `n_crudo` y `n_sin_gtin`: los articulos de esa
    pasada (n_leidas + n_sin_gtin) y los sin EAN (no estan en disp_estado: solo se cuentan). Devuelve (foto,
    apartados, cuentas, escalon), con `escalon` = {producto_prov: {uds_escalon, precio_unidad, precio_escalon, aviso}}
    de las filas de la foto, y en `cuentas` cuantas entran sin tramo leido y cuantas con el EAN de 14 rescatado."""
    previas = {p: 0 for p in e2.PUERTAS_PREVIAS}
    previas['sin_gtin'] = n_sin_gtin
    apartados, sirven = [], []
    no_elegida = modo == 'elegidas'
    for f in sorted(filas or [], key=lambda x: str(x.get('producto_prov') or '')):
        if not f.get('disponible') or f.get('preorder'):
            previas['no_disponible'] += 1
            continue
        if not quiere_marca(f.get('marca')):
            apartados.append(_apartado(f, 'marca_fuera', ('Marca %r no elegida' if no_elegida
                                                          else 'Marca %r fuera de las marcas del director')
                                       % ((f.get('marca') or '').strip(),)))
            continue
        if f.get('es_chase') and not f.get('es_caja'):
            apartados.append(_apartado(f, 'chase_suelto', 'Chase suelto: no se compra (solo la caja de 6)'))
            continue
        core, rescatado = ean_de_la_foto(f)
        if core is None:
            apartados.append(_apartado(f, 'ean_forma_rara', 'EAN con forma rara (%s%s)'
                                       % (str(f.get('ean_original') or '').strip() or 'vacío',
                                          ', regla %s' % f['regla'] if f.get('regla') else '')))
            continue
        pa = _num(f.get('precio_escalon'))
        if pa is None or pa <= 0:
            apartados.append(_apartado(f, 'estado_no_servible', 'Sin precio de compra en la foto de TCG'))
            continue
        sirven.append(dict(f, _core=core, _rescatado=rescatado))

    uno = {}
    for f in sirven:
        k = M.norm(f['_core'])
        prev = uno.get(k)
        if prev is None:
            uno[k] = f
            continue
        gana, pierde = (f, prev) if _gana(f, prev) else (prev, f)
        uno[k] = gana
        apartados.append(_apartado(pierde, 'duplicado_proveedor', 'Mismo EAN que el %s de TCG%s: me quedo con el más '
                                   'barato (%s frente a %s)' % (
                                       gana.get('producto_prov'), ' (caja de 6)' if gana.get('es_caja') else '',
                                       gana.get('precio_escalon'), pierde.get('precio_escalon'))))

    foto, escalon = [], {}
    for f in sorted(uno.values(), key=lambda x: str(x.get('producto_prov') or '')):
        core = f['_core']
        caja = bool(f.get('es_caja'))
        variantes = sorted({M.norm(v) for v in M.variantes_ean(core)} - {''})
        foto.append({
            'producto_heo': str(f.get('producto_prov') or ''), 'ean_original': str(f.get('ean_original') or core),
            'ean_core': core, 'variantes': variantes, 'codigos_keepa': e2.codigos_para_keepa(core, M),
            'nombre': f.get('nombre') or '', 'marca': (f.get('marca') or '').strip(),
            'categoria': f.get('categoria') or '',
            # 🔑 precio_unidad = el precio del TRAMO MAS BARATO (precio_escalon, por figura): con el decide
            #    `escaner2_motor.decidir`, sin descuento añadido. precio_catalogo = el del Excel de TCG, por figura y
            #    sin tramo (columna «Precio unidad»).
            'precio_catalogo': _num(f.get('precio_unidad')), 'precio_unidad': _num(f.get('precio_escalon')),
            'es_caja': caja, 'uds_caja': (int(f.get('uds_caja') or UDS_CAJA) if caja else None),
            'es_chase': bool(f.get('es_chase')), 'en_oferta': bool(f.get('en_oferta')), 'campana': None,
            'disponibilidad': None, 'fin_de_vida': None, 'preorder': False, 'imagen': None,
            'aviso_caja': TEXTO_CAJA if caja else None, 'origen_ean': None, 'aviso_ean': None,
        })
        escalon[str(f.get('producto_prov') or '')] = {
            'uds_escalon': f.get('uds_escalon'), 'precio_unidad': _num(f.get('precio_unidad')),
            'precio_escalon': _num(f.get('precio_escalon')), 'aviso': f.get('aviso') or None}

    for m in e2.MOTIVOS_APARTADO:
        previas[m] = sum(1 for a in apartados if a['motivo'] == m)
    cuentas = {'n_crudo': n_crudo, 'n_filas': len(filas or []), 'n_foto': len(foto), 'previas': previas,
               'n_sin_tramo': sum(1 for f in uno.values() if f.get('aviso') == AVISO_SIN_TRAMO),
               'n_ean14_rescatado': sum(1 for f in uno.values() if f.get('_rescatado'))}
    cuentas.update(cuadre_previo(cuentas))
    return foto, apartados, cuentas, escalon


def cuadre_previo(cuentas):
    """🔴 crudo = previas + foto, y el crudo son los articulos del Excel: filas de la foto de disponibilidad + las sin
    EAN. Un recuento que falta (None) NO es un cero."""
    previas = cuentas['previas']
    faltan = [p for p in e2.PUERTAS_PREVIAS if previas.get(p) is None]
    if cuentas.get('n_crudo') is None:
        faltan.insert(0, 'crudo')
    if faltan:
        return {'n_previas': None, 'cuadra_previo': False, 'motivo_previo': 'sin recuento de: ' + ', '.join(faltan)}
    n_previas = sum(previas[p] for p in e2.PUERTAS_PREVIAS)
    if cuentas['n_filas'] + previas['sin_gtin'] != cuentas['n_crudo']:
        return {'n_previas': n_previas, 'cuadra_previo': False,
                'motivo_previo': 'el Excel tenía %d artículos y la foto de disponibilidad trae %d + %d sin EAN'
                                 % (cuentas['n_crudo'], cuentas['n_filas'], previas['sin_gtin'])}
    if cuentas['n_crudo'] != n_previas + cuentas['n_foto']:
        return {'n_previas': n_previas, 'cuadra_previo': False,
                'motivo_previo': 'catálogo crudo %d ≠ puertas previas %d + foto %d'
                                 % (cuentas['n_crudo'], n_previas, cuentas['n_foto'])}
    return {'n_previas': n_previas, 'cuadra_previo': True, 'motivo_previo': None}


def corpus_cotejo(foto, apartados):
    """Los nombres con que la regla del viejo mide sus palabras distintivas al elegir entre varias fichas: los de la
    foto y los de lo disponible apartado por MARCA (que la eleccion no dependa de las marcas elegidas)."""
    nombres = [f.get('nombre') or '' for f in foto]
    nombres += [a.get('nombre') or '' for a in apartados if a.get('motivo') == 'marca_fuera']
    return nombres


def ruta_lista(pasada):
    return '%s/%s/eans.txt' % (CARPETA, pasada)


def ruta_excel(pasada, cruce, sello):
    """`tcg/<pasada>/<cruce>/Escaner2_TCG_<AAAAMMDD_HHMM>.xlsx`: la forma que admite escaner2_cruce.ruta_excel desde la
    migracion de la v2 de este encargo."""
    return '%s/%s/%s/Escaner2_TCG_%s.xlsx' % (CARPETA, pasada, cruce, sello)


def nombre_modo(modo):
    """Como se llama el modo de la pasada en el Excel (el mismo texto que el cruce de HEO)."""
    return {'todas': 'todas las marcas', 'elegidas': 'marcas elegidas'}.get(modo or 'marcas', 'marcas de siempre')


# ═══════════════════════════════════════════════════════════════════════════════
# 3 · EL EXCEL: EL DEL PRO DE HEO, CON «Uds. escalón», «Precio unidad» Y «Aviso del tramo» AL FINAL DE «Análisis»
# ═══════════════════════════════════════════════════════════════════════════════
COLUMNAS_CAJA = ['Caja', 'Precio caja (€)', 'EAN de la figura', 'Origen del EAN', 'Aviso del EAN']


def _columnas_caja(f):
    """Las cinco columnas de caja de «Puertas» del PRO de HEO. En TCG la caja se compra a precio POR FIGURA: el precio
    de la caja son las seis figuras a ese precio (el del Excel, sin tramo)."""
    if not f.get('es_caja'):
        return [None, None, None, None, None]
    pc = f.get('precio_catalogo')
    uds = f.get('uds_caja') or UDS_CAJA
    return [e2.rotulo_caja(f) or None, None if pc is None else round(pc * uds, 2), None, None, None]


def poner_columnas_escalon(wb, foto, escalon):
    """«Uds. escalón», «Precio unidad» y «Aviso del tramo» detras de la ultima columna de «Análisis», fila a fila por
    su EAN, y la tabla alargada (como `poner_columna_ficha_compartida`). Devuelve cuantas filas llevan dato."""
    from openpyxl.styles import Font
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.table import TableColumn
    ws = wb['Análisis']
    cab = [c.value for c in ws[1]]
    i_ean = cab.index('EAN')
    primera = len(cab) + 1
    for j, (nombre, ancho) in enumerate(zip(COLUMNAS_ESCALON, ANCHOS_ESCALON)):
        ws.cell(row=1, column=primera + j, value=nombre).font = Font(bold=True)
        ws.column_dimensions[get_column_letter(primera + j)].width = ancho
    por_ean = {}
    for f in foto:
        por_ean[str(f['ean_original'])] = (escalon or {}).get(str(f.get('producto_heo') or '')) or {}
    n = 0
    for fila in range(2, ws.max_row + 1):
        dato = por_ean.get(str(ws.cell(row=fila, column=i_ean + 1).value))
        if not dato:
            continue
        ws.cell(row=fila, column=primera, value=dato.get('uds_escalon'))
        c = ws.cell(row=fila, column=primera + 1, value=dato.get('precio_unidad'))
        c.number_format = '0.00'
        ws.cell(row=fila, column=primera + 2, value=dato.get('aviso'))
        n += 1
    ultima = primera + len(COLUMNAS_ESCALON) - 1
    for tabla in ws.tables.values():
        ini, fin = tabla.ref.split(':')
        fila_fin = re.sub(r'^[A-Z]+', '', fin)
        tabla.ref = '%s:%s%s' % (ini, get_column_letter(ultima), fila_fin)
        if tabla.tableColumns:
            for nombre in COLUMNAS_ESCALON:
                tabla.tableColumns.append(TableColumn(id=len(tabla.tableColumns) + 1, name=nombre))
        if tabla.autoFilter is not None:
            tabla.autoFilter.ref = tabla.ref
    return n


def escribir_excel(foto, resultados, apartados, M, info):
    """El Excel del cruce de TCG: el del PRO de HEO (`escaner2_heo_cruce.escribir_excel`), mismo formato y orden.
    DELANTE las hojas del viejo (la Celda 9: Análisis con «Ventas», Descartados, Ambiguos, Sin_rank, Precio por lote),
    con «Ficha compartida» y las TRES columnas de TCG al final de «Análisis»; DETRAS Resumen, Comparación (vacía),
    Varias fichas, Puertas y Puertas previas. Solo lo calculado: aqui no se decide nada. → bytes."""
    from openpyxl.styles import Font
    por_foto = {f['id']: f for f in foto}
    datos = e2.datos_como_el_viejo(foto, resultados, apartados, M, eleccion=info.get('eleccion'))
    # Sin «Chase_manual» (solo con PROVEEDOR == 'HEO').
    datos['PROVEEDOR'] = PROVEEDOR
    wb = e2.escribir_celda9(datos, M)
    e2.poner_columna_ficha_compartida(wb, foto, resultados)
    poner_columnas_escalon(wb, foto, info.get('escalon'))

    def hoja(nombre, cabecera, filas, anchos=None):
        ws = wb.create_sheet(nombre)
        ws.append(cabecera)
        for c in ws[1]:
            c.font = Font(bold=True)
        for fila in filas:
            ws.append(fila)
        for letra, ancho in (anchos or {}).items():
            ws.column_dimensions[letra].width = ancho
        ws.freeze_panes = 'A2'
        return ws

    p = info['params']
    modo = info.get('modo')
    filas_res = [['Pasada', info['pasada']], ['Cruce', info['cruce']], ['Modo del barrido', nombre_modo(modo)]]
    if modo == 'elegidas':
        filas_res += [['Marcas elegidas', ', '.join(info.get('marcas') or []) or '(ninguna)']]
    elif modo == 'marcas':
        filas_res += [['Marcas del director (Magic y Yu-Gi-Oh en cualquier idioma, sin corte de rank)',
                       ', '.join(info.get('marcas') or []) or '(ninguna)']]
    filas_res += [['Foto de TCG (pasada de disponibilidad)', info.get('foto') or '—'],
                  ['Precio de compra (PA)', 'el del tramo más barato de la web de TCG (el de la foto), sin descuento '
                                            'añadido; por figura en la caja de 6; «Uds. escalón» dice cuántas pedir '
                                            'para ese precio (en la caja, figuras). Sin tramo leído, el suelto del '
                                            'Excel, y «Aviso del tramo» lo dice'],
                  ['Filas con su precio suelto («sin tramo leído»)', info.get('n_sin_tramo', '—')],
                  ['Umbral de caídas (30 días)', '%s (> %d)' % (e2.texto_corte(p['umbral']), p['umbral'])],
                  ['Fichas compartidas (productos de la lista, por país)',
                   ' · '.join('%s %d' % (k, n) for k, n in (info.get('n_compartidas') or {}).items()) or '—'],
                  ['Países del filtro de ventas', ', '.join(p['paises_filtro'])],
                  ['Países que se calculan (si traen CSV)', ', '.join(p['paises_calculo'])],
                  ['Países con CSV', ', '.join(info['usados'])],
                  ['Líneas del Excel de TCG (con los EAN repetidos)', info.get('n_lineas_excel', '—')],
                  ['Catálogo crudo de TCG (artículos distintos del Excel)', info['n_crudo']]]
    filas_res += [['Puerta previa · %s' % nombre_previa(x, modo), info['previas'][x]] for x in e2.PUERTAS_PREVIAS]
    filas_res += [['Entradas (filas de la foto: un EAN, su fila más barata)', info['n_entradas']]]
    filas_res += [['Puerta %s · %s' % (x, e2.NOMBRE_PUERTA[x]), info['n_bd'][x]] for x in e2.PUERTAS]
    filas_res += [['Puertas previas + suma de puertas', sum(info['previas'].values()) + sum(info['n_bd'].values())],
                  ['Cuadra', 'SÍ' if info['cuadra'] else 'NO'],
                  ['Comparación con el escáner viejo', 'no se hace aquí']]
    filas_res += [['CSV', '%s · %s · %s filas%s%s' % (
        f['nombre'], f.get('pais') or '¿?', f.get('filas') or 0,
        '' if f.get('usado') else ' · NO USADO', (' · ' + f['error']) if f.get('error') else '')]
        for f in info['ficheros']]
    filas_res += [['Aviso', a] for a in info['avisos']]
    hoja('Resumen', ['Qué', 'Valor'], filas_res, {'A': 44, 'B': 90})

    hoja('Comparación', ['EAN', 'Nombre', 'Categoría', 'Viejo', 'País viejo', 'Fecha viejo (UTC)', 'Excel viejo',
                         'Puerta nuevo', 'Nuevo', 'País nuevo', 'Fecha nuevo (UTC)', 'Diferencia de criterio',
                         'Explicación'], [], {'A': 15, 'B': 50, 'C': 16, 'G': 36, 'M': 110})

    varias = []
    for res in resultados:
        if res['puerta'] == 'b':
            f = por_foto[res['foto_id']]
            for fi in res['fichas']:
                varias.append([f['ean_original'], f['nombre'], fi['pais'], fi['asin'], fi['titulo'], fi['rank'],
                               fi['caidas_30d'], fi['precio_venta']])
    hoja('Varias fichas', ['EAN', 'Nombre TCG', 'País', 'ASIN', 'Título Amazon', 'Puesto', 'Caídas 30 d',
                           'Precio venta (€)'], varias, {'A': 15, 'B': 50, 'E': 60})

    hoja('Puertas', ['EAN', 'Nombre', 'Marca', 'Precio compra (€)', 'Puerta', 'Motivo', 'Detalle'] + COLUMNAS_CAJA,
         [[por_foto[r['foto_id']]['ean_original'], por_foto[r['foto_id']]['nombre'], por_foto[r['foto_id']]['marca'],
           por_foto[r['foto_id']]['precio_unidad'], '%s · %s' % (r['puerta'], e2.NOMBRE_PUERTA[r['puerta']]),
           r['motivo'], r['detalle']] + _columnas_caja(por_foto[r['foto_id']]) for r in resultados],
         {'A': 15, 'B': 55, 'C': 16, 'E': 24, 'F': 16, 'G': 70, 'H': 26})

    hoja('Puertas previas', ['EAN tal como vino', 'Nombre', 'Marca', 'Precio catálogo (€)', 'Puerta previa', 'Detalle'],
         [[a['ean_original'], a['nombre'], a['marca'], a['precio_catalogo'],
           nombre_previa(a['motivo'], modo) if a['motivo'] in NOMBRE_PREVIA else a['motivo'], a['detalle']]
          for a in apartados],
         {'A': 18, 'B': 55, 'C': 16, 'E': 32, 'F': 80})
    salida = io.BytesIO()
    # (RS, 23-oct-2026) «Amazon», la ÚLTIMA columna de «Análisis» (detrás de las de cada proveedor: nadie se corre)
    e2.poner_columna_amazon(wb, info.get('restricciones'))
    wb.save(salida)
    return salida.getvalue()
