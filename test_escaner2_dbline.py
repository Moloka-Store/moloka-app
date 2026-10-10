# -*- coding: utf-8 -*-
"""Banco de la FOTO DE DISPONIBILIDAD DE DBLINE (encargo DB2-A, 09-oct-2026): escaner2_dbline.py (las reglas).

SIN RED, SIN SECRETOS Y SIN BASE. 🔴 LOS DATOS SON INVENTADOS: el Excel se fabrica aqui con la MISMA FORMA que el real
(fila 1 «Catalogo generale <dd-mm-aaaa>», cabecera de 29 columnas en la fila 3, la llave en una formula =HYPERLINK con
el codigo en base64 dentro del enlace), pero ni un precio, ni un EAN, ni un codigo son de DBLine: el repo es publico.

QUE PRUEBA: la llave por formula (y que no se lee con data_only); codigo vacio, repetido o que no casa con el enlace →
fallida; cabeceras en italiano y en ingles (y mezcladas o de menos → fallida); fila 1; EAN con espacios, de 11 cifras
(UPC-A con su control bien y mal) y raro; promo vigente, del dia y caducada; TELEFONARE con unidades; PRENOTAZIONE sin
unidades; «w/Chase» normal y chase suelto; huella igual con la hora de las fechas distinta; minimos del cuadre; los
recuentos por marca; y que los errores no sueltan valores.
DB2-B (devolucion de Cowork): la huella cambia con el EAN y con la nota; convertir() no deja salir avisos de openpyxl
(con su CONTROL: sin la regla, el aviso sale); y los cambios de una pasada contados por marca.
DB2-B, la fila 1 de la descarga por servidor (seccion 11): vale con la fecha de HOY en italiano o en ingles, con - / .,
en los dos ordenes (dd-mm y mm-dd), partida en dos celdas o como fecha de Excel; ayer, mañana, otro año o un dia
imposible → fallida. Si falla la fila 1 o la cabecera, al texto solo su FORMA y donde esta «Publisher»/«EAN»; su
texto real, recortado a 80, solo en para_la_base(); la fila 2, en ningun sitio.
"""
import base64
import contextlib
import io
import json
import os
import sys
from datetime import date, datetime

import openpyxl

RAIZ = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, RAIZ)
import escaner2_dbline as DB  # noqa: E402

fallos = []
HOY = date(2026, 1, 10)


def eq(nombre, obtenido, esperado):
    ok = obtenido == esperado
    if not ok:
        fallos.append(nombre)
    print(('OK ' if ok else 'XX ') + nombre + ('' if ok else '   got=%r exp=%r' % (obtenido, esperado)))


def formula(codigo, cod_enlace=None):
    """La celda «Codice/Link» como la trae DBLine, con un codigo inventado."""
    cod_enlace = codigo if cod_enlace is None else cod_enlace
    param = base64.b64encode(json.dumps({'COD_PRODOTTO': cod_enlace}).encode()).decode().rstrip('=')
    return '=HYPERLINK("https://tienda.invalid/?id=SCHEDA_PRODOTTO&param=%s","%s")' % (param, codigo)


def fila(codigo='ZZ0001', **k):
    """Una fila de 29 columnas con valores inventados; lo que no se diga, neutro."""
    d = {c: None for c in DB.COLUMNAS_IT}
    d.update({'Publisher': 'OTRA', 'Codice/Link': formula(codigo), 'Descrizione': 'Cosa inventada',
              'EAN': '4999999000016       ', 'Disponibili': 5, 'Prezzo (€)': 10.0, 'Prezzo promo (€)': 0,
              'Data uscita': datetime(2025, 5, 1, 12, 30)})
    d.update(k)
    return [d[c] for c in DB.COLUMNAS_IT]


def excel(filas, cabecera=DB.COLUMNAS_IT, titulo='Catalogo generale 10-01-2026', fila2=None):
    """titulo: el valor de A1, o una lista (una celda por valor, desde A1). fila2: el valor de A2 (como el codigo de
    cliente y el nombre de Moloka del fichero de verdad), si se da."""
    libro = openpyxl.Workbook()
    h = libro.active
    for j, t in enumerate(titulo if isinstance(titulo, (list, tuple)) else [titulo], 1):
        h.cell(row=1, column=j, value=t)
    if fila2 is not None:
        h.cell(row=2, column=1, value=fila2)
    for j, c in enumerate(cabecera, 1):
        h.cell(row=3, column=j, value=c)
    for i, f in enumerate(filas, 4):
        for j, val in enumerate(f, 1):
            if val is not None:
                h.cell(row=i, column=j, value=val)
    b = io.BytesIO()
    libro.save(b)
    return b.getvalue()


def leer(filas, **k):
    return DB.convertir(excel(filas, **{x: y for x, y in k.items() if x in ('cabecera', 'titulo', 'fila2')}), HOY,
                        min_filas=k.get('min_filas', 1), min_disponibles=k.get('min_disponibles', 0))


def invalida(filas, **k):
    try:
        leer(filas, **k)
    except DB.LecturaInvalida as e:
        return str(e)
    return None


def excepcion(filas, **k):
    """La LecturaInvalida entera (para mirar lo que va a la base), o None."""
    try:
        leer(filas, **k)
    except DB.LecturaInvalida as e:
        return e
    return None


def por_codigo(salida):
    return {f['producto_prov']: f for f in salida}


# ── (1) La llave por formula ──────────────────────────────────────────────────────────────
salida, c = leer([fila('ZZ0001'), fila('ZZZZ0002')])
eq('llave: el codigo sale del texto de la formula', sorted(f['producto_prov'] for f in salida), ['ZZ0001', 'ZZZZ0002'])
eq('llave: pieza suelta, formula buena', DB.codigo_de_formula(formula('AB1234')), 'AB1234')
eq('llave: el codigo no casa con COD_PRODOTTO → None', DB.codigo_de_formula(formula('AB1234', 'AB9999')), None)
eq('llave: texto plano, sin formula → None', DB.codigo_de_formula('AB1234'), None)
eq('llave: param que no es base64/json → None',
   DB.codigo_de_formula('=HYPERLINK("https://tienda.invalid/?param=@@@","AB1234")'), None)
libro = openpyxl.load_workbook(io.BytesIO(excel([fila('ZZ0001')])), data_only=True)
eq('llave: con data_only la celda sale vacia (por eso se lee sin el)', libro.active.cell(row=4, column=8).value, None)

err = invalida([fila('ZZ0001'), fila('ZZ0001')])
eq('llave: codigo repetido → fallida', err is not None and '1 código(s) repetido(s)' in err, True)
err = invalida([fila('ZZ0001'), fila('', **{'Codice/Link': formula('')})])
eq('llave: codigo vacio → fallida', err is not None and '1 fila(s) sin código' in err, True)
err = invalida([fila('ZZ0001', **{'Codice/Link': formula('ZZ0001', 'ZZ0009')})])
eq('llave: codigo que no casa con el enlace → fallida', err is not None and 'no casa' in err, True)
eq('llave: el error no suelta el codigo', 'ZZ0001' in (err or ''), False)

# ── (2) Cabecera y fila 1 ─────────────────────────────────────────────────────────────────
_s, c_it = leer([fila()])
_s, c_en = leer([fila()], cabecera=DB.COLUMNAS_EN)
eq('cabecera: italiano', c_it['idioma_cabecera'], 'it')
eq('cabecera: ingles', c_en['idioma_cabecera'], 'en')
eq('cabecera: 29 columnas en cada idioma', (len(DB.COLUMNAS_IT), len(DB.COLUMNAS_EN)), (29, 29))
eq('cabecera: 13 se llaman igual en los dos', sum(a == b for a, b in zip(DB.COLUMNAS_IT, DB.COLUMNAS_EN)), 13)
mezcla = list(DB.COLUMNAS_IT)
mezcla[3] = DB.COLUMNAS_EN[3]
eq('cabecera: idiomas mezclados → fallida', 'mezcla' in (invalida([fila()], cabecera=tuple(mezcla)) or ''), True)
err = invalida([fila()], cabecera=DB.COLUMNAS_IT[:28])
eq('cabecera: 28 columnas → fallida con la cuenta', (err or '').startswith('cabecera: 28 columnas, y son 29 · su forma: «'),
   True)
movida = list(DB.COLUMNAS_IT)
movida[13], movida[17] = movida[17], movida[13]
eq('cabecera: columnas cambiadas de sitio → fallida', 'columna 14' in (invalida([fila()], cabecera=tuple(movida)) or ''), True)
eq('fila 1: fecha del catalogo', c_it['fecha_catalogo'], '2026-01-10')
eq('fila 1: otra cosa → fallida', (invalida([fila()], titulo='Listino') or '').startswith(
    'fila 1: no trae la fecha de hoy (10-01-2026) · su forma: «aaaaaaa»'), True)
eq('fila 1: dia imposible → fallida', (invalida([fila()], titulo='Catalogo generale 31-02-2026') or '').startswith(
    'fila 1: no trae la fecha de hoy'), True)

# ── (3) EAN ───────────────────────────────────────────────────────────────────────────────
eq('ean: espacios de relleno fuera (13)', DB.ean_de_cruce('4999999000016       '), ('4999999000016', None))
eq('ean: 12 cifras tal cual', DB.ean_de_cruce('012345678905        '), ('012345678905', None))
eq('ean: 11 cifras con UPC-A bien → 0 delante', DB.ean_de_cruce('12345678905         '), ('012345678905', None))
eq('ean: 11 cifras con el control mal → forma rara', DB.ean_de_cruce('12345678904'), (None, 'ean_forma_rara'))
eq('ean: vacio → forma rara', DB.ean_de_cruce('                    '), (None, 'ean_forma_rara'))
eq('ean: None → forma rara', DB.ean_de_cruce(None), (None, 'ean_forma_rara'))
eq('ean: con letras → forma rara', DB.ean_de_cruce('4999999000016 C6'), (None, 'ean_forma_rara'))
eq('ean: 14 cifras → forma rara', DB.ean_de_cruce('14999999000013'), (None, 'ean_forma_rara'))
eq('ean: numero de la celda (int)', DB.ean_de_cruce(4999999000016), ('4999999000016', None))
salida, c = leer([fila('ZZ0001', EAN='   '), fila('ZZ0002', EAN='12345678905   '), fila('ZZ0003')])
p = por_codigo(salida)
eq('ean: la fila rara NO se tira, se cuenta', (len(salida), c['n_ean_forma_rara'], p['ZZ0001']['regla']),
   (3, 1, 'ean_forma_rara'))
eq('ean: el de 11 se cuenta aparte', (c['n_ean_11_upc'], p['ZZ0002']['ean_core']), (1, '012345678905'))
eq('ean: ean_original guardado sin el relleno', p['ZZ0003']['ean_original'], '4999999000016')

# ── (4) Disponible, notas y preventa ──────────────────────────────────────────────────────
salida, c = leer([fila('D1', Disponibili=3, Note='TELEFONARE'), fila('D2', Disponibili=0, Note='PRENOTAZIONE'),
                  fila('D3', Disponibili=2, Note='NEW'), fila('D4', Disponibili=0, Note='ESAURITO'),
                  fila('D5', Disponibili=None)])
p = por_codigo(salida)
eq('disponible: TELEFONARE con unidades entra', (p['D1']['disponible'], p['D1']['disponibilidad']), (True, 'TELEFONARE'))
eq('disponible: PRENOTAZIONE sin unidades, preorder y no disponible', (p['D2']['disponible'], p['D2']['preorder']),
   (False, True))
eq('disponible: NEW con unidades entra', p['D3']['disponible'], True)
eq('disponible: 0 unidades fuera', p['D4']['disponible'], False)
eq('disponible: vacio = sin dato', (p['D5']['disponible'], p['D5']['sin_dato_disponibilidad']), (False, True))
eq('disponible: recuentos', (c['n_disponibles'], c['n_agotados'], c['n_preventa'], c['n_telefonare_disponibles'],
                             c['n_new_disponibles']), (2, 3, 1, 1, 1))
eq('disponible: texto en Disponibili → fallida', 'Disponibili' in (invalida([fila(Disponibili='muchas')]) or ''), True)

# ── (5) Precio ────────────────────────────────────────────────────────────────────────────
salida, c = leer([
    fila('P1', **{'Prezzo (€)': 10.0, 'Prezzo promo (€)': 8.0, 'Scadenza promo': datetime(2026, 1, 20, 9, 15)}),
    fila('P2', **{'Prezzo (€)': 10.0, 'Prezzo promo (€)': 8.0, 'Scadenza promo': datetime(2026, 1, 9, 23, 59)}),
    fila('P3', **{'Prezzo (€)': 10.0, 'Prezzo promo (€)': 8.0, 'Scadenza promo': datetime(2026, 1, 10, 0, 1)}),
    fila('P4', **{'Prezzo (€)': 10.0}),
    fila('P5', **{'Prezzo (€)': '12,50', 'Prezzo promo (€)': 0, 'Scadenza promo': '20/01/2026 17:40'}),
])
p = por_codigo(salida)
eq('precio: promo vigente → precio_unidad = promo',
   (p['P1']['precio_unidad'], p['P1']['precio_catalogo'], p['P1']['en_oferta'], p['P1']['fin_oferta']),
   (8.0, 10.0, True, '2026-01-20'))
eq('precio: promo caducada → Prezzo', (p['P2']['precio_unidad'], p['P2']['en_oferta'], p['P2']['fin_oferta']),
   (10.0, False, '2026-01-09'))
eq('precio: la promo que acaba hoy vale (solo el dia)', (p['P3']['precio_unidad'], p['P3']['en_oferta']), (8.0, True))
eq('precio: sin promo → Prezzo', (p['P4']['precio_unidad'], p['P4']['en_oferta'], p['P4']['fin_oferta']),
   (10.0, False, None))
eq('precio: texto con coma y fecha en texto con hora', (p['P5']['precio_unidad'], p['P5']['fin_oferta']),
   (12.5, '2026-01-20'))
eq('precio: sin escalones, sin PA, sin cajas',
   {(f['precio_escalon'], f['uds_escalon'], f['precio_pa'], f['es_caja'], f['uds_caja']) for f in salida},
   {(None, None, None, False, None)})
eq('precio: recuentos de oferta', (c['n_en_oferta'], c['n_promo_caducada']), (2, 1))
eq('precio: Prezzo que no es numero → fallida', 'Prezzo' in (invalida([fila(**{'Prezzo (€)': 'diez'})]) or ''), True)
eq('precio: promo sin fin que se entienda → fallida',
   'Scadenza promo' in (invalida([fila(**{'Prezzo promo (€)': 8.0, 'Scadenza promo': 'pronto'})]) or ''), True)

# ── (5b) DB5: el precio de compra es el MINIMO de Prezzo y la promo vigente ────────────────
_FIN = datetime(2026, 12, 31, 10, 0)
salida, c = leer([
    fila('M1', **{'Prezzo (€)': 7.83, 'Prezzo promo (€)': 7.0, 'Scadenza promo': _FIN}),
    fila('M2', **{'Prezzo (€)': 7.83, 'Prezzo promo (€)': 7.83, 'Scadenza promo': _FIN}),
    fila('M3', **{'Prezzo (€)': 7.83, 'Prezzo promo (€)': 7.99, 'Scadenza promo': _FIN}),
    fila('M4', **{'Prezzo (€)': 7.83, 'Prezzo promo (€)': 8.49, 'Scadenza promo': datetime(2026, 1, 10, 0, 1)}),
    fila('M5', **{'Prezzo (€)': 7.83, 'Prezzo promo (€)': 7.99, 'Scadenza promo': datetime(2026, 1, 9, 23, 59)}),
    fila('M6', **{'Prezzo (€)': 0, 'Prezzo promo (€)': 5.0, 'Scadenza promo': _FIN}),
    fila('M7', **{'Prezzo (€)': None, 'Prezzo promo (€)': 5.0, 'Scadenza promo': _FIN}),
])
p = por_codigo(salida)


def _precio(k):
    return (p[k]['precio_unidad'], p[k]['precio_catalogo'], p[k]['en_oferta'], p[k]['fin_oferta'])


eq('DB5: promo vigente MAS BARATA que Prezzo → la promo, en oferta', _precio('M1'), (7.0, 7.83, True, '2026-12-31'))
eq('DB5: promo vigente IGUAL a Prezzo → Prezzo, sin oferta y sin fin', _precio('M2'), (7.83, 7.83, False, None))
eq('DB5: promo vigente MAS CARA que Prezzo → Prezzo, sin oferta y sin fin', _precio('M3'), (7.83, 7.83, False, None))
eq('DB5: promo mas cara que acaba hoy → Prezzo, sin oferta y sin fin', _precio('M4'), (7.83, 7.83, False, None))
eq('DB5: promo mas cara y CADUCADA → como antes (Prezzo, su fin, sin oferta)', _precio('M5'),
   (7.83, 7.83, False, '2026-01-09'))
eq('DB5: Prezzo a 0 → como antes (la promo vigente vale)', _precio('M6'), (5.0, 0.0, True, '2026-12-31'))
eq('DB5: Prezzo vacio → como antes (la promo vigente vale, sin dato de precio)',
   _precio('M7') + (p['M7']['sin_dato_precio'],), (5.0, None, True, '2026-12-31', True))
eq('DB5: recuentos (en oferta, mas cara, caducada, sin fin)',
   (c['n_en_oferta'], c['n_promo_mas_cara'], c['n_promo_caducada'], c['n_promo_sin_fin']), (3, 3, 1, 0))
eq('DB5: precio_vigente directo (mas barata, igual, mas cara, Prezzo 0, Prezzo vacio)',
   [DB.precio_vigente(DB.Decimal(a), DB.Decimal(b), date(2026, 1, 10), HOY) for a, b in
    (('7.83', '7'), ('7.83', '7.83'), ('7.83', '7.99'), ('0', '5'))]
   + [DB.precio_vigente(None, DB.Decimal('5'), date(2026, 1, 10), HOY)],
   [(DB.Decimal('7'), True), (DB.Decimal('7.83'), False), (DB.Decimal('7.83'), False), (DB.Decimal('5'), True),
    (DB.Decimal('5'), True)])
# Una promo mas cara deja la fila (y la huella) igual que la del mismo producto sin promo.
_s_cara, _c_cara = leer([fila('M3', **{'Prezzo (€)': 7.83, 'Prezzo promo (€)': 7.99, 'Scadenza promo': _FIN})])
_s_sin, _c_sin = leer([fila('M3', **{'Prezzo (€)': 7.83})])
eq('DB5: promo mas cara = la fila de un producto sin promo', (_s_cara, _c_cara['huella_contenido']),
   (_s_sin, _c_sin['huella_contenido']))
eq('DB5: …y se cuenta solo en la de la promo', (_c_cara['n_promo_mas_cara'], _c_sin['n_promo_mas_cara']), (1, 0))

# ── (6) Chase ─────────────────────────────────────────────────────────────────────────────
salida, c = leer([fila('C1', Publisher='FUNKO', Descrizione='POP Heroe Inventado w/Chase'),
                  fila('C2', Publisher='FUNKO', Descrizione='POP Heroe Inventado Chase'),
                  fila('C3', Publisher='FUNKO', Descrizione='POP Heroe Inventado (Glow Chase)'),
                  fila('C4', Publisher='FUNKO', Descrizione='POP Heroe Inventado with Chase'),
                  fila('C5', Publisher='OTRA', Descrizione='Perro Inventado Chase')])
p = por_codigo(salida)
eq('chase: «w/Chase» = figura normal', (p['C1']['es_chase'], p['C1']['regla']), (False, None))
eq('chase: FUNKO «… Chase» al final → chase_suelto', (p['C2']['es_chase'], p['C2']['regla']), (True, 'chase_suelto'))
eq('chase: FUNKO «(… Chase)» → chase_suelto', p['C3']['regla'], 'chase_suelto')
eq('chase: «with Chase» = figura normal', p['C4']['es_chase'], False)
eq('chase: otra marca, nunca', p['C5']['es_chase'], False)
eq('chase: recuento', c['n_chase_suelto'], 2)

# ── (7) Fechas y huella ───────────────────────────────────────────────────────────────────
def lote(hora):
    return [fila('H1', **{'Prezzo promo (€)': 8.0, 'Scadenza promo': datetime(2026, 1, 20, *hora),
                          'Data uscita': datetime(2025, 3, 3, *hora)}),
            fila('H2')]


s1, c1 = leer(lote((9, 15)))
s2, c2 = leer(list(reversed(lote((17, 40)))))
eq('fechas: solo el dia', (s1[0]['fin_oferta'], s1[0]['fecha_salida']), ('2026-01-20', '2025-03-03'))
eq('huella: igual con la hora de las fechas distinta (y otro orden de filas)',
   c1['huella_contenido'] == c2['huella_contenido'], True)
eq('huella: el md5 del fichero si cambia', c1['md5_fichero'] != c2['md5_fichero'], True)
eq('huella: bytes del fichero', c1['bytes_fichero'] > 0, True)
s3, c3 = leer([fila('H1', **{'Prezzo promo (€)': 8.0, 'Scadenza promo': datetime(2026, 1, 21)}), fila('H2')])
eq('huella: cambia si cambia el fin de la oferta', c1['huella_contenido'] != c3['huella_contenido'], True)
s4, c4 = leer([fila('H1', **{'Prezzo promo (€)': 8.0, 'Scadenza promo': datetime(2026, 1, 20)}),
               fila('H2', Disponibili=0)])
eq('huella: cambia si cambia el disponible', c1['huella_contenido'] != c4['huella_contenido'], True)
s5, c5 = leer([fila('H1', **{'Prezzo promo (€)': 8.0, 'Scadenza promo': datetime(2026, 1, 20)}),
               fila('H2', Disponibili=99, Descrizione='Otro nombre')])
eq('huella: NO cambia si cambian unidades (sigue disponible) o nombre',
   c1['huella_contenido'] == c5['huella_contenido'], True)
# (DB2-B) El EAN y la nota, dentro: un cambio de EAN o de reserva no es «al día».
s6, c6 = leer([fila('H1', **{'Prezzo promo (€)': 8.0, 'Scadenza promo': datetime(2026, 1, 20)}),
               fila('H2', Note='PRENOTAZIONE')])
eq('huella (DB2-B): cambia si cambia la nota (PRENOTAZIONE)', c1['huella_contenido'] != c6['huella_contenido'], True)
s7, c7 = leer([fila('H1', **{'Prezzo promo (€)': 8.0, 'Scadenza promo': datetime(2026, 1, 20)}),
               fila('H2', EAN='4999999000023')])
eq('huella (DB2-B): cambia si cambia el EAN', c1['huella_contenido'] != c7['huella_contenido'], True)
s8, c8 = leer([fila('H1', **{'Prezzo promo (€)': 8.0, 'Scadenza promo': datetime(2026, 1, 20)}),
               fila('H2', EAN='  4999999000016')])
eq('huella (DB2-B): NO cambia si el EAN es el mismo con otros espacios (mira ean_core)',
   c1['huella_contenido'] == c8['huella_contenido'], True)
eq('filas: no salen las columnas internas', any(k.startswith('_') for f in s1 for k in f), False)

# ── (8) Cuadre y recuentos por marca ──────────────────────────────────────────────────────
filas = [fila('M%d' % i, Publisher=m, Disponibili=d) for i, (m, d) in enumerate(
    [('FUNKO', 1), ('FUNKO', 0), ('Pyramid International', 2), ('OTRA', 3), ('OTRA', 0)])]
_s, c = leer(filas)
eq('marcas: total, Funko y Pyramid (filas, disponibles)', c['por_marca'],
   {'total': {'filas': 5, 'disponibles': 3}, 'funko': {'filas': 2, 'disponibles': 1},
    'pyramid': {'filas': 1, 'disponibles': 1}})
eq('cuadre: los minimos de la casa', (DB.MIN_FILAS, DB.MIN_DISPONIBLES), (16500, 11000))
eq('cuadre: menos filas del minimo → fallida', invalida(filas, min_filas=6), 'cuadre: 5 filas, y el mínimo es 6')
eq('cuadre: menos disponibles del minimo → fallida', invalida(filas, min_disponibles=4),
   'cuadre: 3 disponibles, y el mínimo es 4')
try:
    DB.convertir(excel(filas), HOY)
    err = None
except DB.LecturaInvalida as e:
    err = str(e)
eq('cuadre: con los minimos por defecto, 5 filas no pasan', err, 'cuadre: 5 filas, y el mínimo es 16500')
try:
    DB.convertir(b'esto no es un excel', HOY)
    err = None
except DB.LecturaInvalida as e:
    err = str(e)
eq('fichero: no es un .xlsx → fallida', err, 'el fichero no se abre como .xlsx')

# ── (9) Devolucion de Cowork del 9-oct-2026, un caso por punto ─────────────────────────────
# 1. EAN que llega como float entero
eq('cowork 1: EAN float entero = el int', DB.ean_de_cruce(4999999000016.0), ('4999999000016', None))
eq('cowork 1: EAN float con decimales → forma rara', DB.ean_de_cruce(4999999000016.5), (None, 'ean_forma_rara'))
# 2. n_ean_forma_rara cuenta por el EAN, no por la regla
salida, c = leer([fila('E1', Publisher='FUNKO', Descrizione='POP Heroe Inventado Chase', EAN='  ')])
eq('cowork 2: chase suelto con EAN raro cuenta en los dos', (c['n_chase_suelto'], c['n_ean_forma_rara'],
                                                           salida[0]['regla'], salida[0]['ean_core']),
   (1, 1, 'chase_suelto', None))
# 3. promo > 0 sin fin: con Prezzo y contada
salida, c = leer([fila('S1', **{'Prezzo (€)': 10.0, 'Prezzo promo (€)': 8.0}), fila('S2')])
eq('cowork 3: promo sin fin → Prezzo, sin oferta', (por_codigo(salida)['S1']['precio_unidad'],
                                                    por_codigo(salida)['S1']['en_oferta']), (10.0, False))
eq('cowork 3: n_promo_sin_fin', c['n_promo_sin_fin'], 1)
# 4. hoy como datetime
salida, c = DB.convertir(excel([fila('T1', **{'Prezzo promo (€)': 8.0, 'Scadenza promo': datetime(2026, 1, 10, 0, 1)})]),
                         datetime(2026, 1, 10, 18, 30), min_filas=1, min_disponibles=0)
eq('cowork 4: hoy como datetime, sin TypeError y por el dia', (salida[0]['precio_unidad'], salida[0]['en_oferta']),
   (8.0, True))
# 5. NaN e Infinity no son numeros
for nombre, val in (('NaN', float('nan')), ('Infinity', float('inf')), ('-Infinity', float('-inf')),
                    ('Decimal NaN', DB.Decimal('NaN'))):
    try:
        DB.numero(val)
        r = 'pasa'
    except ValueError:
        r = 'ValueError'
    eq('cowork 5: numero(%s) → ValueError' % nombre, r, 'ValueError')
# 6. una dimension mal declarada no recorta columnas
import re as _re
import zipfile as _zip


def dimension_falsa(contenido, ref='A1:C3'):
    ent, sal = _zip.ZipFile(io.BytesIO(contenido)), io.BytesIO()
    with _zip.ZipFile(sal, 'w', _zip.ZIP_DEFLATED) as z:
        for it in ent.infolist():
            datos = ent.read(it.filename)
            if it.filename == 'xl/worksheets/sheet1.xml':
                datos, n = _re.subn(rb'<dimension ref="[^"]*"\s*/>', b'<dimension ref="%s"/>' % ref.encode(), datos)
                assert n == 1
            z.writestr(it, datos)
    return sal.getvalue()


falsa = dimension_falsa(excel([fila('R1'), fila('R2')]))
try:
    salida, c = DB.convertir(falsa, HOY, min_filas=1, min_disponibles=0)
    r = (len(salida), c['idioma_cabecera'], salida[0]['precio_unidad'])
except DB.LecturaInvalida as e:
    r = str(e)
eq('cowork 6: dimension declarada A1:C3 → se leen las 29 columnas y las filas', r, (2, 'it', 10.0))

# ── (10) Devolucion de Cowork al DB2-B ─────────────────────────────────────────────────────
# d) Ningun aviso de openpyxl sale de convertir(): uno de mentira (con un codigo y un EAN inventados dentro) en cada
#    apertura del libro. CONTROL: la lectura SIN la regla 10 (_convertir) si lo deja salir; si no, la prueba no prueba.
import warnings as _w

_load_real = DB.openpyxl.load_workbook


def _load_con_aviso(*a, **k):
    _w.warn('AVISO-DE-OPENPYXL ZZW001 4999999000016', UserWarning)
    return _load_real(*a, **k)


_libro_w = excel([fila('ZZW001')])
DB.openpyxl.load_workbook = _load_con_aviso
try:
    with _w.catch_warnings(record=True) as _vistos:
        _w.simplefilter('always')
        DB._convertir(_libro_w, HOY, 1, 0)
    _control = len(_vistos)
    with _w.catch_warnings(record=True) as _vistos:
        _w.simplefilter('always')
        DB.convertir(_libro_w, HOY, min_filas=1, min_disponibles=0)
    _con_regla = len(_vistos)
    _err = io.StringIO()
    with contextlib.redirect_stderr(_err), _w.catch_warnings():
        _w.simplefilter('default')
        DB.convertir(_libro_w, HOY, min_filas=1, min_disponibles=0)
finally:
    DB.openpyxl.load_workbook = _load_real
eq('avisos (DB2-B): CONTROL, sin la regla 10 el aviso de mentira sale', _control >= 1, True)
eq('avisos (DB2-B): convertir() no deja salir ninguno', _con_regla, 0)
eq('avisos (DB2-B): nada al registro (stderr)', _err.getvalue(), '')

# c) Los cambios de UNA pasada, por marca (las filas de disp_cambio: tipo y marca).
_cambios = [{'tipo': 'entra_catalogo', 'marca': 'FUNKO'}, {'tipo': 'pasa_disponible', 'marca': 'Funko Pop'},
            {'tipo': 'cambia_precio', 'marca': 'Pyramid International'}, {'tipo': 'sale_catalogo', 'marca': None},
            {'tipo': 'agotado_sin_dato', 'marca': 'FUNKO'}, {'tipo': 'vuelve_catalogo', 'marca': 'OTRA'},
            {'tipo': 'pasa_agotado', 'marca': 'PYRAMID'}, {'tipo': 'cambia_precio', 'marca': 'FUNKO'}]
_r = DB.recuentos_cambios(_cambios)
_cero = {'entran': 0, 'vuelven': 0, 'salen': 0, 'a_disponible': 0, 'a_agotado': 0, 'cambio_precio': 0}
eq('cambios (DB2-B): Funko', _r['funko'], dict(_cero, entran=1, a_disponible=1, cambio_precio=1))
eq('cambios (DB2-B): Pyramid', _r['pyramid'], dict(_cero, cambio_precio=1, a_agotado=1))
eq('cambios (DB2-B): total (agotado_sin_dato no se mide)', _r['total'],
   {'entran': 1, 'vuelven': 1, 'salen': 1, 'a_disponible': 1, 'a_agotado': 1, 'cambio_precio': 2})
eq('cambios (DB2-B): sin cambios, todo a 0', DB.recuentos_cambios([]),
   {'total': _cero, 'funko': _cero, 'pyramid': _cero})

# ── (11) DB2-B: la fila 1 de la descarga por servidor ─────────────────────────────────────
# Titulos INVENTADOS (el de servidor no lo ha visto nadie): vale cualquiera que traiga la fecha de HOY (10-01-2026).
for _t in ('Catalogo generale 10-01-2026', 'General catalogue 10-01-2026', 'General catalog 10/01/2026',
           'Catalogue général 10.01.2026', 'Price list 2026-01-10', 'Catalogo generale 10-1-2026',
           'General catalogue 01-10-2026', 'General catalogue 01/10/2026', ['General catalogue', '10-01-2026'],
           [datetime(2026, 1, 10)]):
    try:
        _r = leer([fila()], titulo=_t)[1]['fecha_catalogo']
    except DB.LecturaInvalida as e:
        _r = str(e)
    eq('fila 1 (DB2-B): vale %r' % (_t,), _r, '2026-01-10')
eq('fila 1 (DB2-B): las dos lecturas de 01-10-2026 (1-oct y 10-ene)', DB.fechas_de_la_fila('x 01-10-2026'),
   {date(2026, 10, 1), date(2026, 1, 10)})
for _t, _que in (('Catalogo generale 09-01-2026', 'ayer'), ('General catalogue 11-01-2026', 'mañana (y 1-nov)'),
                 ('General catalogue 10-01-2025', 'otro año'), ('Catalogo generale 31-02-2026', 'dia imposible'),
                 ('General catalogue', 'sin fecha'), ('', 'vacía'), ('General catalogue 10-01-26', 'año de 2 cifras'),
                 ('Ref 110-01-2026', 'con una cifra pegada delante')):
    eq('fila 1 (DB2-B): %s → fallida' % _que, (invalida([fila()], titulo=_t) or '').startswith(
        'fila 1: no trae la fecha de hoy (10-01-2026)'), True)

# Lo que va al texto (registro) y lo que va a la base.
_e = excepcion([fila()], titulo='General catalogue SECRETO 09-01-2026', fila2='CLIENTE-FILA2 12345 Moloka Store')
_pub, _base = str(_e), _e.para_la_base()
eq('fila 1 (DB2-B): al texto, la forma', 'su forma: «aaaaaaa aaaaaaaaa aaaaaaa 99-99-9999»' in _pub, True)
eq('fila 1 (DB2-B): al texto, nada del titulo', [x for x in ('General', 'catalogue', 'SECRETO') if x in _pub], [])
eq('fila 1 (DB2-B): al texto, dónde está la cabecera',
   '«Publisher» y «EAN» en la fila 3; 29 columnas en la fila 3' in _pub, True)
eq('fila 1 (DB2-B): a la base, además su texto', _base == _pub + ' · su texto: «General catalogue SECRETO 09-01-2026»',
   True)
eq('fila 1 (DB2-B): la fila 2, en ningún sitio', [x for x in ('CLIENTE', '12345', 'Moloka') if x in _pub + _base], [])
_largo = 'Titolo ' + 'x' * 100 + ' 09-01-2026'
_e = excepcion([fila()], titulo=_largo)
eq('fila 1 (DB2-B): el texto a la base, recortado a 80',
   ('«' + _largo[:80] + '»' in _e.para_la_base(), _largo[:81] in _e.para_la_base()), (True, False))
eq('fila 1 (DB2-B): la forma, recortada a 80', len(str(_e).split('su forma: «')[1].split('»')[0]), 80)
_e = excepcion([fila()], titulo='Catalogo generale 09-01-2026', cabecera=tuple('Col%d' % i for i in range(29)))
eq('fila 1 (DB2-B): sin «Publisher» ni «EAN», lo dice',
   'sin «Publisher» y «EAN» en las 6 primeras filas; 29 columnas en la fila 3' in str(_e), True)
# La cabecera: igual.
_cab = list(DB.COLUMNAS_IT)
_cab[13] = 'ColumnaInventada'
_e = excepcion([fila()], cabecera=tuple(_cab), fila2='CLIENTE-FILA2')
eq('cabecera (DB2-B): fallida en la columna 14', str(_e).startswith('cabecera: la columna 14 no es «Disponibili»'), True)
eq('cabecera (DB2-B): al texto la forma, no el nombre real',
   ('su forma: «aaa 9 aaa 9 aaa 9' in str(_e), 'ColumnaInventada' in str(_e)), (True, False))
eq('cabecera (DB2-B): a la base, su texto', 'su texto: «Cat 1 Cat 2 Cat 3' in _e.para_la_base(), True)
eq('cabecera (DB2-B): la fila 2, en ningún sitio', 'CLIENTE' in _e.para_la_base(), False)
eq('fila 1 (DB2-B): lo que no falla no lleva privado', DB.LecturaInvalida('x').para_la_base(), 'x')

print()
if fallos:
    print('ROJO: %d comprobación(es) fallida(s):' % len(fallos))
    for f in fallos:
        print('  - ' + f)
    sys.exit(1)
print('VERDE: todas las comprobaciones pasan.')
