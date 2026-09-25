# -*- coding: utf-8 -*-
"""Banco de punta a punta del escaner 2 de HEO: los DOS programas de verdad, con dobles.

`escaner2_heo_barrido.py` y `escaner2_heo_cruce.py` son scripts (se corren con `runpy` y
hacen `sys.exit`), asi que cada caso va en su PROCESO. `supabase` (base y Storage) y la API de
HEO (`descargar_heo`) se sustituyen por dobles EN MEMORIA que guardan su estado en un fichero
entre proceso y proceso: el barrido deja la foto y el cruce la lee. Sin red, sin secretos y
sin tocar produccion. El escaner viejo NO se toca: sus piezas se leen de su fichero real.

CASOS:
  1. [sin_llave]  sin SUPABASE_SERVICE_KEY → ROJO, la linea exacta y NINGUN cliente creado.
  2. [bueno]      barrido → la pasada queda 'esperando_csv' con su foto, sus apartados y la
                  lista de EAN en el bucket cerrado; subo un CSV de ES y otro de DE; cruce →
                  'lista', CUADRA contra la base, compara con el viejo y deja el Excel.
  3. [frenado]    la base ACEPTA el insert de las puertas y se come las de la puerta b (lo que
                  haria una politica que frena: 200 y cero filas) → el cruce CUENTA en la base,
                  ve que no cuadra y queda 'fallida' en ROJO. Es el cuadre del encargo en vivo.
  6. [todas]      (B2) barrido «todas las marcas»: sin leer reglas_director, sin filtro de marca,
                  cuadra; y la marca que el viejo no mira sale como diferencia de criterio.
  7. [cajas]      (B2) cajas con chase REALES de HEO: la disponible entra con el EAN de su figura
                  y a precio de caja ÷ 6; la agotada y la suelta del apellido, como toca; cuadra.
  8. [el viejo]   (B2) en todos los casos, NI UNA escritura en reglas_director, escaner_chase_asin,
                  escaner_memoria, escaner_resultados, escaner_detalle ni productos: el doble
                  apunta cada operacion, y esas tablas acaban byte a byte como empezaron. (B4) Y
                  el buzon del viejo (`informes`, carpeta «resultados») acaba como empezo.
  9. [elegidas]   (B4) barrido «marcas elegidas» con solo Funko: la foto solo trae Funko, el resto
                  va a «marca no elegida» y cuadra crudo = previas + foto; ni se lee reglas_director;
                  el cruce deja el Excel con las seis hojas del viejo delante, con SU formato.
  (B4) En los casos 2 y 9, el formato de las seis hojas del viejo se COMPARA con la huella del Excel
  viejo de verdad (huella_excel_viejo_heo.json): orden de hojas, cabeceras, anchos, formulas,
  enlaces, tablas y formato condicional. Si cambia algo, en el nuevo o en el viejo, sale ROJO.
 10. [mala]       (B4) una seleccion malformada (no JSON, no lista, comillas, saltos de linea,
                  vacia) → ROJO, la linea exacta y NINGUN cliente creado: no se ejecuta nada.
"""
import base64
import csv
import io
import json
import os
import re
import subprocess
import sys
import tempfile
import types

AQUI = os.path.dirname(os.path.abspath(__file__))
fallos = []


def eq(nombre, obtenido, esperado):
    ok = obtenido == esperado
    if not ok:
        fallos.append(nombre)
    print(('OK ' if ok else 'XX ') + nombre + ('' if ok else '   got=%r exp=%r' % (obtenido, esperado)))


# ═══════════════════════════════════════════════════════════════════════════════
# EL DOBLE DE SUPABASE: tablas en memoria y un Storage que pagina como storage3 2.31.0
# ═══════════════════════════════════════════════════════════════════════════════
class _Resp:
    def __init__(self, data, count=None):
        self.data, self.count = data, count


class _Consulta:
    def __init__(self, bd, tabla):
        self.bd, self.tabla = bd, tabla
        self.op, self.filtros, self.payload = 'select', [], None
        self.cuenta, self.orden, self.rango, self.tope = False, None, None, None

    def select(self, _cols='*', count=None):
        self.op, self.cuenta = 'select', (count == 'exact')
        return self

    def upsert(self, filas, **_kw):
        self.op, self.payload = 'upsert', filas if isinstance(filas, list) else [filas]
        return self

    def delete(self):
        self.op = 'delete'
        return self

    def insert(self, filas):
        self.op, self.payload = 'insert', filas if isinstance(filas, list) else [filas]
        return self

    def update(self, cambios):
        self.op, self.payload = 'update', cambios
        return self

    def eq(self, k, v):
        self.filtros.append((k, v))
        return self

    def order(self, col, desc=False):
        self.orden = (col, desc)
        return self

    def range(self, a, b):
        self.rango = (a, b)
        return self

    def limit(self, n):
        self.tope = n
        return self

    def execute(self):
        # (B2) Cada operacion queda apuntada: quien la hizo, cual y sobre que tabla.
        self.bd.setdefault('ops', []).append([self.bd.get('programa'), self.op, self.tabla])
        filas = self.bd['tablas'].setdefault(self.tabla, [])
        if self.op in ('upsert', 'delete'):
            return _Resp([])          # apuntada y sin efecto: el banco exige que no ocurra
        if self.op == 'insert':
            for f in self.payload:
                if self.bd.get('frenar') and self.tabla == self.bd['frenar'][0] and f.get('puerta') == self.bd['frenar'][1]:
                    continue       # 200 y cero filas: lo que hace una politica que frena
                filas.append(json.loads(json.dumps(f, default=str)))
            return _Resp(self.payload)
        sel = [f for f in filas if all(str(f.get(k)) == str(v) for k, v in self.filtros)]
        if self.op == 'update':
            for f in sel:
                f.update(json.loads(json.dumps(self.payload, default=str)))
            return _Resp(sel)
        if self.orden:
            sel = sorted(sel, key=lambda f: (f.get(self.orden[0]) is None, str(f.get(self.orden[0]))),
                         reverse=self.orden[1])
        n = len(sel)
        if self.rango:
            sel = sel[self.rango[0]:self.rango[1] + 1]
        if self.tope is not None:
            sel = sel[:self.tope]
        return _Resp(sel, n if self.cuenta else None)


class _Cubo:
    def __init__(self, bd, cubo):
        self.bd, self.cubo = bd, cubo

    def _objs(self):
        return self.bd['storage'].setdefault(self.cubo, {})

    def upload(self, ruta, datos, _opciones=None):
        self._objs()[ruta] = base64.b64encode(datos if isinstance(datos, bytes) else datos.encode()).decode()
        return {'Key': ruta}

    def download(self, ruta):
        if ruta not in self._objs():
            raise RuntimeError('404 Object not found: %s/%s' % (self.cubo, ruta))
        return base64.b64decode(self._objs()[ruta])

    def list(self, carpeta=None, opciones=None):
        # Como storage3 2.31.0: DEFAULT_SEARCH_OPTIONS = limit 100, offset 0, por nombre asc.
        op = dict({'limit': 100, 'offset': 0}, **(opciones or {}))
        pref = (carpeta or '').rstrip('/') + '/'
        nombres = sorted(r[len(pref):] for r in self._objs() if r.startswith(pref) and '/' not in r[len(pref):])
        trozo = nombres[op['offset']:op['offset'] + op['limit']]
        return [{'name': n, 'created_at': '2026-09-24T12:%02d:00Z' % i, 'metadata': {'size': 1}}
                for i, n in enumerate(trozo)]


class _Cliente:
    def __init__(self, bd):
        self.bd = bd
        self.storage = types.SimpleNamespace(from_=lambda cubo: _Cubo(bd, cubo))

    def table(self, nombre):
        return _Consulta(self.bd, nombre)


# ═══════════════════════════════════════════════════════════════════════════════
# LA ESCENA: el catalogo de HEO, los dos CSV y el escaner viejo
# ═══════════════════════════════════════════════════════════════════════════════
def _escena():
    sys.path.insert(0, AQUI)
    import escaner2_motor as e2
    import moloka_escaner_pro as pro
    M = e2.cargar_motor(os.path.join(AQUI, e2.RUTA_MOTOR))
    return e2, pro, M


def _ean(M, n):
    cuerpo = '84355%07d' % n
    return cuerpo + M._chk13(cuerpo)


def filas_heo(M):
    def h(n, nombre, precio, marca='FUNKO'):
        return {'productNumber': 'HEO%04d' % n, 'ean': _ean(M, n), 'nombre': nombre, 'marca': marca,
                'categoria': 'Figuras', 'precio': precio, 'precio_base': precio, 'en_oferta': '',
                'campana': '', 'estado': 'disponible', 'disponibilidad': 'GREEN', 'imagen': '',
                'fin_de_vida': '', 'preorder': ''}
    return [h(1, 'Funko Pop Alfa', 8.00), h(2, 'Funko Pop Beta', 10.00), h(3, 'Funko Pop Gamma', 8.00),
            h(4, 'Funko Pop Delta', 8.00), h(5, 'Funko Pop Epsilon', 8.00),
            h(6, 'Funko Pop Zeta Chase', 30.00), h(7, 'Hasbro fuera', 5.00, marca='Hasbro')]


def csv_visualizador(e2, pro, M, pais, filas):
    col_pais, col_caidas = e2.columnas_keepa(os.path.join(AQUI, e2.RUTA_ESCAPARATE))
    C = pro.CSV_COLS
    cab = ['ASIN', col_pais, 'Título', C['ean'], C['rank'], C['rank90'], col_caidas, C['buybox'], C['es_fba'],
           C['nuevo'], C['fba'], C['compct'], C['nof'], C['vendidos'], C['vendidos2'], C['nvar']]
    buf = io.StringIO()
    w = csv.writer(buf)
    w.writerow(cab)
    for asin, n, caidas, bb in filas:
        # `n` entero: el EAN de la escena; texto: un EAN real (el de la figura de una caja, B2).
        w.writerow([asin, pais, 'Funko %s' % n, _ean(M, n) if isinstance(n, int) else n, '12000', '15000', caidas, bb,
                    'yes', bb, '3.50', '15.01 %', '10', '50', '', '0'])
    return ('﻿' + buf.getvalue()).encode('utf-8')


def excel_viejo(e2, M, filas):
    from openpyxl import Workbook
    cols = e2.sacar_piezas(os.path.join(AQUI, e2.RUTA_MOTOR), (), ('COLS',))['COLS']
    wb = Workbook()
    ws = wb.active
    ws.title = 'Análisis'
    ws.append(cols)
    for n, pais, dec in filas:
        fila = [None] * len(cols)
        fila[cols.index('EAN')], fila[cols.index('País')], fila[cols.index('Decisión')] = _ean(M, n), pais, dec
        ws.append(fila)
    b = io.BytesIO()
    wb.save(b)
    return b.getvalue()


# ═══════════════════════════════════════════════════════════════════════════════
# EL HIJO: monta los dobles y corre el programa de verdad
# ═══════════════════════════════════════════════════════════════════════════════
def chase_de(escena):
    """La lista `chase` de descargar_heo en cada escena. 'base': una caja con chase de la que no
    se puede sacar el EAN de la figura (codigo 999… y numero que no es FK) → su puerta previa.
    'cajas' (B2): cuatro filas REALES de la pasada 8ac9a9de (ver test_escaner2_cajas_chase.py;
    precio redondo donde el encargo no lo cita)."""
    if escena == 'cajas':
        return [
            {'producto_heo': 'FK87245', 'nombre': '*heo Exclusive Edition* One Piece POP!&Buddy Animation Vinyl '
             'Figuren Rob Lucci with Hattori w/Chase 10 cm Surtido (6)', 'ean_caja': '01108896988724512110000',
             'marca': 'Funko', 'precio_caja': 60.0, 'estado': 'disponible', 'imagen': '', 'link_amazon': ''},
            {'producto_heo': 'FK93057', 'nombre': "NFL Figura POP! Vinyl : Bengals- Ja'Marr Chase (clear visor) 9 cm",
             'ean_caja': '889698930574', 'marca': 'Funko', 'precio_caja': 8.62, 'estado': 'agotado', 'imagen': '',
             'link_amazon': ''},
            {'producto_heo': 'FK91391', 'nombre': 'X-Men Pack de 4 Figuras Bitty POP! Vinyl Jean Grey w/CH 2,5 cm',
             'ean_caja': '889698913911', 'marca': 'Funko', 'precio_caja': 6.59, 'estado': 'disponible', 'imagen': '',
             'link_amazon': ''},
            {'producto_heo': 'FK13318', 'nombre': 'Stranger Things POP! TV Vinyl Figuren Eleven With Eggos 9 cm Surtido (6)',
             'ean_caja': '01108896981331872110000', 'marca': 'Funko', 'precio_caja': 60.0, 'estado': 'agotado',
             'imagen': '', 'link_amazon': ''},
        ]
    return [{'producto_heo': 'HEO9001', 'nombre': 'Funko Pop Omega w/CH Surtido (6)', 'ean_caja': '9990000000011',
             'marca': 'FUNKO', 'precio_caja': 70.0, 'estado': 'disponible', 'imagen': '', 'link_amazon': ''}]


def hijo(ruta_estado, programa):
    with open(ruta_estado, encoding='utf-8') as fh:
        bd = json.load(fh)
    bd['programa'] = programa
    clientes = []

    def crear(url, llave):
        clientes.append(llave)
        return _Cliente(bd)

    sys.modules['supabase'] = types.ModuleType('supabase')
    sys.modules['supabase'].create_client = crear
    e2, _pro, M = _escena()

    def descargar_catalogo_heo(max_paginas=None, con_chase=False):
        # Las dos lineas del log de la funcion de verdad (descargar_heo.py) de las que el barrido
        # saca el catalogo CRUDO y los tirados sin GTIN: 7 con EAN + la lista chase + 3 sin GTIN.
        # …y la de `_paginar`: lo que HEO DICE que tiene (totalElements). En la de verdad sale antes.
        chase = chase_de(os.environ.get('E2_ESCENA', 'base'))
        crudo = 7 + len(chase) + 3
        print('  catalog/products: %d items | 1 paginas | pageSize 500' % int(os.environ.get('E2_DECLARADO', crudo)))
        print('>>> Cruzando: %d productos | 0 precios | 0 disponibilidades' % int(os.environ.get('E2_CRUDO', crudo)))
        print('>>> Catalogo cruzado: 7 filas con EAN (descartadas 3 sin GTIN)')
        return (filas_heo(M), chase) if con_chase else filas_heo(M)

    sys.modules['descargar_heo'] = types.ModuleType('descargar_heo')
    sys.modules['descargar_heo'].descargar_catalogo_heo = descargar_catalogo_heo
    import runpy
    sys.argv = [programa] + os.environ.get('E2_ARGS', '').split()
    codigo = 0
    try:
        runpy.run_path(os.path.join(AQUI, programa), run_name='__main__')
    except SystemExit as ex:
        codigo = ex.code if isinstance(ex.code, int) else (0 if ex.code is None else 1)
    finally:
        with open(ruta_estado, 'w', encoding='utf-8') as fh:
            json.dump(bd, fh, default=str)
        print('CLIENTES_CREADOS=%d' % len(clientes))
    sys.exit(codigo)


def correr(ruta_estado, programa, env_extra):
    env = {k: v for k, v in os.environ.items()
           if k not in ('SUPABASE_SERVICE_KEY', 'HEO_USER', 'HEO_PASS', 'PASADA', 'MODO_BARRIDO', 'E2_ESCENA',
                        'MARCAS_ELEGIDAS', 'OFERTAS_ELEGIDAS')}
    env.update(env_extra, PYTHONIOENCODING='utf-8', SUPABASE_URL='https://doble.invalid')
    p = subprocess.run([sys.executable, '-u', os.path.abspath(__file__), '--hijo', ruta_estado, programa],
                       capture_output=True, text=True, encoding='utf-8', errors='replace', env=env, cwd=AQUI)
    return p.returncode, (p.stdout or '') + (p.stderr or '')


# (B2) Las tablas del escaner viejo, sembradas: al final de cada caso tienen que seguir IGUAL.
TABLAS_VIEJO = ('reglas_director', 'escaner_chase_asin', 'escaner_memoria', 'escaner_resultados', 'escaner_detalle',
                'productos')


def estado_inicial(e2, M, frenar=None):
    viejo = excel_viejo(e2, M, [(1, 'ES', 'COMPRAR'), (1, 'DE', 'COMPRAR'), (5, 'ES', 'COMPRAR'), (3, 'IT', 'VALORAR')])
    return {
        'frenar': frenar,
        'tablas': {
            # (B2) La ultima pasada de marcas de siempre, de la que el modo «todas» saca la lista
            #      del viejo para comparar (sin leer reglas_director).
            'escaner2_pasada': [{'id': '00000000-0000-4000-8000-000000000001', 'proveedor': 'HEO', 'modo': 'marcas',
                                 'estado': 'esperando_csv', 'marcas': ['Funko', 'Ultimate Guard'], 'ofertas': True,
                                 'creada_en': '2026-09-25T05:48:00+00:00', 'run_id': 1}],
            'escaner_chase_asin': [{'producto_heo': 'FK87245', 'nombre': 'Rob Lucci', 'ean_caja': '01108896988724512110000',
                                    'asin': None, 'estado': 'disponible'}],
            'escaner_memoria': [{'id': 1, 'proveedor': 'HEO', 'ean': '889698872454'}],
            'escaner_detalle': [{'id': 1, 'ean': '889698872454', 'decision': 'COMPRAR'}],
            'reglas_director': [{'proveedor': 'HEO', 'activo': True, 'marcas': ['Funko', 'OFERTAS'], 'rank_maximo': 30000}],
            'escaner2_parametros': [{'proveedor': 'HEO', 'umbral_caidas_30d': 8, 'paises_filtro': ['ES', 'DE'],
                                     'paises_calculo': ['ES', 'IT', 'FR', 'DE']}],
            'productos': [{'id': 'p1', 'ean': _ean(M, 1), 'asin': 'B0ALFA0001', 'iva_pct': 0.10, 'activo': True,
                           'stock_moloka': 0, 'stock_fba': 3}],
            'escaner_resultados': [{'id': 1, 'proveedor': 'HEO', 'modo': 'todo', 'rank_maximo': 30000,
                                    'fecha': '2026-09-24T00:16:10+00:00',
                                    'fichero': 'resultados/Escaneo_HEO_TODAS_20260923_2006.xlsx'}],
        },
        'storage': {'informes': {'resultados/Escaneo_HEO_TODAS_20260923_2006.xlsx': base64.b64encode(viejo).decode()}},
    }


SIEMBRA = {}


def caso(nombre, frenar=None, crudo=None, modo=None, escena=None, extra=None):
    e2, pro, M = _escena()
    tmp = tempfile.mkdtemp(prefix='e2pp_')
    ruta = os.path.join(tmp, 'estado.json')
    inicial = estado_inicial(e2, M, frenar)
    SIEMBRA[nombre] = {t: json.loads(json.dumps(inicial['tablas'][t])) for t in TABLAS_VIEJO}
    with open(ruta, 'w', encoding='utf-8') as fh:
        json.dump(inicial, fh)
    llave = {'SUPABASE_SERVICE_KEY': 'svc-de-mentira', 'HEO_USER': 'u', 'HEO_PASS': 'p', 'GITHUB_RUN_ID': '424242'}
    if crudo is not None:
        llave['E2_CRUDO'] = str(crudo)
    if modo is not None:
        llave['MODO_BARRIDO'] = modo
    if escena is not None:
        llave['E2_ESCENA'] = escena
    llave.update(extra or {})
    cod, log = correr(ruta, 'escaner2_heo_barrido.py', llave)
    bd = json.load(open(ruta, encoding='utf-8'))
    pasada = [p for p in bd['tablas']['escaner2_pasada'] if p.get('run_id') == 424242][0]
    # Fernando exporta del Visualizador y sube los dos CSV (el nombre NO dice el pais).
    carpeta = 'heo/%s/csv/' % pasada['id']
    bd['storage'].setdefault('escaner2', {})
    for nombre_fichero, pais, filas in (
            ('KeepaExport-2026-09-24-VisualizadorDeProductos.csv', 'es',
             [('B0ALFA0001', 1, '30', '22.00'), ('B0BETA0001', 2, '12', '22.00'), ('B0GAMA0001', 3, '4', '22.00'),
              ('B0DELT0001', 4, '20', '22.00'), ('B0DELT0002', 4, '20', '21.00'),
              # (B2) el Hasbro (EAN 7: solo entra en el modo «todas») y la figura comun de FK87245.
              ('B0HASB0001', 7, '30', '22.00'), ('B0CAJA0001', '889698872454', '30', '30.00')]),
            ('KeepaExport-2026-09-24-VisualizadorDeProductos (1).csv', 'de',
             [('B0ALFA0001', 1, '12', '21.00'), ('B0GAMA0001', 3, '2', '21.00'),
              # (B5) el EPSILON: dos fichas, y SOLO en DE. El viejo elige entre las de ES y aquí no hay
              #      ninguna → sigue en la puerta b. (El DELTA, con dos en ES, ahora se elige.)
              ('B0EPSI0001', 5, '12', '21.00'), ('B0EPSI0002', 5, '12', '21.00')])):
        bd['storage']['escaner2'][carpeta + '20260924-1200-' + nombre_fichero] = base64.b64encode(
            csv_visualizador(e2, pro, M, pais, filas)).decode()
    if nombre == 'bueno':
        # 🔑 Un ES MAS VIEJO con otro precio y 0 caidas para el ALFA: el mas reciente tiene que
        #    mandar (antes mandaba el primero por nombre, o sea el viejo).
        bd['storage']['escaner2'][carpeta + '20260924-1100-KeepaExport-antiguo.csv'] = base64.b64encode(
            csv_visualizador(e2, pro, M, 'es', [('B0ALFA0001', 1, '0', '99.00')])).decode()
        # 🔑 Y un fichero que no se puede leer: se dice y se sigue (no bloquea la pasada para siempre).
        bd['storage']['escaner2'][carpeta + '20260924-1300-roto.csv'] = base64.b64encode(b'esto no es un csv').decode()
        # 🔑 Un tercer pais (IT): se CALCULA y se ve, pero no cuenta para «se vende» (filtro ES y DE).
        bd['storage']['escaner2'][carpeta + '20260924-1205-KeepaExport-italia.csv'] = base64.b64encode(
            csv_visualizador(e2, pro, M, 'it', [('B0ALFA0001', 1, '3', '25.00')])).decode()
    bd.pop('programa', None)
    json.dump(bd, open(ruta, 'w', encoding='utf-8'), default=str)
    cod2, log2 = correr(ruta, 'escaner2_heo_cruce.py', {'SUPABASE_SERVICE_KEY': 'svc-de-mentira', 'PASADA': pasada['id'],
                                                        'GITHUB_RUN_ID': '434343'})
    return (cod, log, cod2, log2, json.load(open(ruta, encoding='utf-8')), pasada['id'], M, e2)


if __name__ == '__main__' and len(sys.argv) >= 4 and sys.argv[1] == '--hijo':
    hijo(sys.argv[2], sys.argv[3])

# ═══════════════════════════════════════════════════════════════════════════════
print('1 · [sin_llave] sin la llave de servicio no se corre')
_e2, _pro, _M = _escena()
_tmp = tempfile.mkdtemp(prefix='e2pp_')
_ruta = os.path.join(_tmp, 'estado.json')
json.dump(estado_inicial(_e2, _M), open(_ruta, 'w', encoding='utf-8'))
for _prog in ('escaner2_heo_barrido.py', 'escaner2_heo_cruce.py'):
    _cod, _log = correr(_ruta, _prog, {'HEO_USER': 'u', 'HEO_PASS': 'p', 'PASADA': '00000000-0000-0000-0000-000000000000'})
    eq('1 · %s sin llave → ROJO' % _prog, _cod, 1)
    eq('1 · …con la linea exacta', 'ESCANER2_NO_EJECUTADO: sin llave de servicio' in _log, True)
    eq('1 · …y sin crear NINGUN cliente', 'CLIENTES_CREADOS=0' in _log, True)

print('\n2 · [bueno] barrido → CSV de ES y DE → cruce')
cod, log, cod2, log2, bd, pasada, M, e2 = caso('bueno')
T = bd['tablas']
p = [x for x in T['escaner2_pasada'] if x['id'] == pasada][0]
eq('2 · el barrido sale en VERDE', cod, 0)
eq('2 · la pasada queda esperando los CSV, con el run apuntado', (p['estado'], p['run_id']), ('esperando_csv', 424242))
eq('2 · foto: los 5 Funko (el chase suelto y Hasbro, fuera)', (p['n_foto'], len(T['escaner2_foto'])), (5, 5))
eq('2 · apartados: la caja con chase sin EAN posible, el chase suelto y (B2) la marca fuera',
   sorted(a['motivo'] for a in T['escaner2_apartado']), ['chase_funko', 'chase_suelto', 'marca_fuera'])
eq('2 · (B2) la pasada guarda su modo, y las marcas que leyó', (p['modo'], p['marcas'], p['ofertas']),
   ('marcas', ['Funko'], True))
eq('2 · los descartados sin GTIN se rescatan del log de descargar_heo', p['p_sin_gtin'], 3)
eq('2 · 🔴 el catálogo CRUDO sale del log, y cuadra: crudo = previas + foto',
   (p['n_crudo'], sum(p['p_' + x] for x in e2.PUERTAS_PREVIAS) + p['n_foto']), (11, 11))
_lista = base64.b64decode(bd['storage']['escaner2']['heo/%s/eans.txt' % pasada]).decode().split('\n')
eq('2 · la lista para el Visualizador: un EAN por linea, en el bucket escaner2', (len(_lista), p['n_eans_lista'], p['n_tandas']),
   (5, 5, 1))
eq('2 · ni una tabla nueva fuera de escaner2_', sorted(t for t in T if not t.startswith('escaner2_')), sorted(TABLAS_VIEJO))
c = T['escaner2_cruce'][0]
eq('2 · el cruce sale en VERDE', cod2, 0)
eq('2 · …y queda LISTA, cuadrado', (c['estado'], c['cuadra'], c['motivo_fallo']), ('lista', True, None))
eq('2 · el pais de cada CSV sale del dato', sorted((f['pais'] or '', f['usado']) for f in c['ficheros']),
   [('', False), ('DE', True), ('ES', True), ('ES', True), ('IT', True)])
eq('2 · 🔴 el CSV roto se dice (en ficheros) y NO tumba el cruce',
   [(f['nombre'].endswith('roto.csv'), bool(f['error'])) for f in c['ficheros'] if f['error']], [(True, True)])
eq('2 · (B5) puertas: entradas 5 = a0 b1 c1 d0 e1 f2 (el DELTA, con dos fichas en ES, se elige; el EPSILON, solo en DE, sigue en b)',
   (c['n_entradas'], c['n_a'], c['n_b'], c['n_c'], c['n_d'], c['n_e'], c['n_f']), (5, 0, 1, 1, 0, 1, 2))
_delta = [r for r in T['escaner2_resultado_ean'] if r['asin'] and r['asin'].startswith('B0DELT')][0]
eq('2 · 🔴 (B5) el DELTA sigue con UNA ficha, la del viejo, y el porqué y las descartadas quedan en el resultado',
   (_delta['puerta'], _delta['asin'], [(fi['asin'], fi['elegida']) for fi in _delta['fichas'] if fi['pais'] == 'ES'],
    _delta['detalle'].startswith('Ficha B0DELT0001 elegida como el viejo (⚠ DUDOSO'), 'descartadas B0DELT0002' in _delta['detalle']),
   ('f', 'B0DELT0001', [('B0DELT0001', True), ('B0DELT0002', False)], True, True))
_epsi = [r for r in T['escaner2_resultado_ean'] if r['puerta'] == 'b'][0]
eq('2 · (B5) el EPSILON sigue en b, y dice por qué no se pudo elegir',
   (_epsi['asin'], 'ninguna en ES' in _epsi['detalle']), (None, True))
eq('2 · …una fila por EAN, ni mas ni menos', len(T['escaner2_resultado_ean']), 5)
_f1 = [r for r in T['escaner2_resultado_ean'] if r['puerta'] == 'f']
eq('2 · COMPRAR con su mejor pais (ES: IVA de ficha 10 %; y el DELTA, ya elegido)', sorted((r['asin'], r['mejor_pais']) for r in _f1),
   [('B0ALFA0001', 'ES'), ('B0DELT0001', 'ES')])
_es1 = [r for r in T['escaner2_resultado_pais'] if r['asin'] == 'B0ALFA0001' and r['pais'] == 'ES'][0]
_alfa_p = {r['pais']: r for r in T['escaner2_resultado_pais'] if r['asin'] == 'B0ALFA0001'}
eq('2 · 🔴 el ALFA se calcula en ES, IT y DE (el IT también), cada uno con su marca «vende aquí»',
   {p: r['vende_aqui'] for p, r in _alfa_p.items()}, {'ES': True, 'IT': False, 'DE': True})
eq('2 · …y el cruce guarda las DOS listas de países', (c['paises_filtro'], c['paises_calculo']),
   (['ES', 'DE'], ['ES', 'IT', 'FR', 'DE']))
eq('2 · 🔴 manda el CSV MAS RECIENTE de ES: precio 22 y 30 caídas, no 99 y 0 del antiguo',
   (_es1['precio_venta'], _es1['caidas_30d']), (22.0, 30))
eq('2 · el IVA de ES sale de la ficha y la fila lo dice', (_es1['iva'], _es1['iva_origen']), (0.1, 'ficha'))
eq('2 · la cuenta es la del viejo (margen guardado, decision COMPRAR)', (_es1['decision'], round(_es1['margen'], 4)),
   ('COMPRAR', round(M.calc_rentabilidad(22.0, 8.0, 15.01, 3.5, 0.10, almacen=M.ALMACEN, com_digitales=M.COM_DIGITALES,
                                          isd=M.ISD_PAIS['ES'])['margen'], 4)))
_cmp = {r['ean']: r for r in T['escaner2_comparacion']}
eq('2 · comparacion: el EAN 1, COMPRAR en los dos', _cmp[_ean(M, 1)]['categoria'], 'ambos')
eq('2 · comparacion: el EAN 3 (IT en el viejo, pocas caidas en el nuevo) → diferencia de criterio',
   (_cmp[_ean(M, 3)]['categoria'], _cmp[_ean(M, 3)]['diferencia_criterio']), ('solo_viejo', True))
eq('2 · …y la fecha del viejo es la de su Excel (UTC)', str(_cmp[_ean(M, 1)]['fecha_viejo'])[:16], '2026-09-23T20:06')
_xl = [k for k in bd['storage']['escaner2'] if k.startswith('heo/%s/%s/Escaner2_HEO_' % (pasada, c['id']))]
eq('2 · el Excel del cruce queda en el bucket escaner2 y apuntado en el cruce', (len(_xl), c['ruta_excel'] == (_xl or [None])[0]),
   (1, True))
from openpyxl import load_workbook  # noqa: E402
_wb = load_workbook(io.BytesIO(base64.b64decode(bd['storage']['escaner2'][_xl[0]])), read_only=True)
# (B4) «exactamente el mismo formato de excel del escaner antiguo»: sus seis hojas DELANTE, en su
#      orden y escritas por SU codigo; las del escaner 2, DETRAS.
HOJAS_VIEJO = ['Análisis', 'Descartados', 'Ambiguos', 'Sin_rank', 'Precio por lote', 'Chase_manual']
HOJAS_E2 = ['Resumen', 'Comparación', 'Varias fichas', 'Puertas', 'Puertas previas']
eq('2 · 🔴 (B4) las seis hojas del viejo delante, en su orden, y las del escáner 2 detrás',
   _wb.sheetnames, HOJAS_VIEJO + HOJAS_E2)
_COLS = e2.sacar_piezas(os.path.join(AQUI, e2.RUTA_MOTOR), (), ('COLS',))['COLS']
_an = list(_wb['Análisis'].iter_rows(values_only=True))
eq('2 · 🔴 (B4) «Análisis» lleva las cabeceras de COLS del viejo, en su orden', list(_an[0]), list(_COLS))
_ia = {h: k for k, h in enumerate(_an[0])}
# Lo esperado sale de la BASE del doble, no de la hoja: los productos de las puertas d, e y f, del
# de más margen en ES al de menos, y cada uno en los cuatro países del viejo.
_foto_ean = {f['id']: f['ean_original'] for f in T['escaner2_foto']}
_mes = {r['foto_id']: r['margen'] for r in T['escaner2_resultado_pais'] if r['pais'] == 'ES'}
_def = sorted((r['foto_id'] for r in T['escaner2_resultado_ean'] if r['puerta'] in 'def'),
              key=lambda fid: _mes.get(fid) if _mes.get(fid) is not None else -10 ** 9, reverse=True)
eq('2 · (B4) una fila por país (ES, IT, FR, DE) de cada producto que se vende (puertas d, e, f: 3 aquí, con el DELTA elegido), por margen de ES',
   ([(r[_ia['EAN']], r[_ia['País']]) for r in _an[1:]], len(_def)),
   ([(_foto_ean[fid], pais) for fid in _def for pais in ('ES', 'IT', 'FR', 'DE')], 3))
_fil_es = [r for r in _an[1:] if r[_ia['EAN']] == _ean(M, 1) and r[_ia['País']] == 'ES'][0]
eq('2 · (B4) el ALFA en ES: la decisión es la GUARDADA y el beneficio, la fórmula VIVA del viejo con su IVA de ficha (10 %)',
   (_fil_es[_ia['Decisión']], _fil_es[_ia['Beneficio (€)']], _fil_es[_ia['Margen']]),
   (_es1['decision'], '=(J2/1.1)-E2-N2-O2-P2', '=R2/J2'))
eq('2 · (B4) «En mi BD» sale del catálogo propio con la función del viejo (el ALFA es nuestro: 3 en FBA)',
   _fil_es[_ia['En mi BD']], 'OK Alm:0 FBA:3')
_fil_fr = [r for r in _an[1:] if r[_ia['EAN']] == _ean(M, 1) and r[_ia['País']] == 'FR'][0]
eq('2 · (B4) sin CSV de FR, su fila dice «Sin datos» y no inventa números',
   (_fil_fr[_ia['Decisión']], _fil_fr[_ia['Margen']], _fil_fr[_ia['Precio venta (€)']]), ('Sin datos', None, None))
_VACIAS = ('Vendidos/mes', 'Nº ofertas', 'Coincide', 'Promo activa', 'OcioStock')
eq('2 · (B4) las columnas que el escáner 2 no tiene van VACÍAS, con su cabecera',
   {c: [r[_ia[c]] for r in _an[1:] if r[_ia[c]] not in (None, '')] for c in _VACIAS}, {c: [] for c in _VACIAS})
eq('2 · (B5) «Cotejo» lleva el veredicto del viejo SOLO donde se eligió entre varias fichas (el DELTA)',
   sorted({(r[_ia['EAN']], r[_ia['Cotejo']]) for r in _an[1:] if r[_ia['Cotejo']]}), [(_ean(M, 4), '⚠ DUDOSO')])
_puerta = {_foto_ean[r['foto_id']]: r['puerta'] for r in T['escaner2_resultado_ean']}
eq('2 · (B5) Ambiguos: el DELTA con el ganador POR PUESTO (lo que escribe el viejo) y el EPSILON, que sigue en b, vacío',
   # (sin orden: las filas siguen el de la foto, que aquí sale de un id aleatorio)
   sorted((tuple(r) for r in list(_wb['Ambiguos'].iter_rows(values_only=True))[1:]), key=str),
   sorted([(_ean(M, 4), 'B0DELT0001'), (_ean(M, 5), None)], key=str))
eq('2 · (B4) Descartados ← puerta a + apartados (no la marca fuera)',
   ([],
    sorted(r[0] for r in list(_wb['Descartados'].iter_rows(values_only=True))[1:])),
   ([],
    sorted([e for e, pu in _puerta.items() if pu == 'a']
           + [a['ean_original'] for a in T['escaner2_apartado'] if a['motivo'] in ('chase_suelto', 'ean_forma_rara',
                                                                                  'duplicado_proveedor', 'estado_no_servible')])))
eq('2 · (B4) Chase_manual ← la caja con chase sin EAN de la figura, con su precio de caja',
   [r[:5] for r in list(_wb['Chase_manual'].iter_rows(values_only=True))[1:]],
   [('Funko Pop Omega w/CH Surtido (6)', 'HEO9001', '9990000000011', 70.0, 11.67)])
import escaner2_huella_excel as HU  # noqa: E402
_REF = json.load(open(os.path.join(AQUI, 'huella_excel_viejo_heo.json'), encoding='utf-8'))
_h2 = HU.huella(base64.b64decode(bd['storage']['escaner2'][_xl[0]]))
eq('2 · 🔴 (B4) el FORMATO de las seis hojas es el del Excel viejo de verdad (huella_excel_viejo_heo.json)',
   HU.diferencias(_REF, _h2, solo_hojas=HOJAS_VIEJO, vacias=('Sin_rank',)), [])
eq('2 · (B4) …con Sin_rank vacía (ningún «sin dato» en esta escena), escrita como la escribe el viejo',
   list(next(_wb['Sin_rank'].iter_rows(values_only=True))), ['(vacio)'])
eq('2 · el log deja el CUADRE, cuadre o no', 'CUADRE [HEO]: crudo=11 | previas=6' in log2
   and '| entradas=5 | suma de puertas=5' in log2, True)
eq('2 · 🔴 el cruce guarda el crudo y las previas: crudo 11 = previas 6 + puertas 5',
   (c['n_crudo'], c['n_previas'], sum(c['n_' + x] for x in 'abcdef')), (11, 6, 5))

BDS = {'bueno': bd}

print('\n3 · [frenado] la base se come las filas de una puerta → el cuadre lo caza')
cod, log, cod2, log2, bd, pasada, M, e2 = caso('frenado', frenar=('escaner2_resultado_ean', 'b'))
c = bd['tablas']['escaner2_cruce'][0]
eq('3 · 🔴 el cruce sale en ROJO', cod2, 1)
eq('3 · 🔴 …y queda FALLIDA, sin cuadrar, con el motivo', (c['estado'], c['cuadra'], 'NO CUADRA' in (c['motivo_fallo'] or '')),
   ('fallida', False, True))
eq('3 · 🔴 …contando en la BASE: 5 entradas y 4 en las puertas', (c['n_entradas'], sum(c['n_' + x] for x in 'abcdef')), (5, 4))
eq('3 · el log lo dice', 'NO CUADRA' in log2, True)
BDS['frenado'] = bd

print('\n4 · [crudo_de_mas] HEO dice un producto más de los que salen → la pasada FALLA')
cod, log, cod2, log2, bd, pasada, M, e2 = caso('crudo_de_mas', crudo=12)
BDS['crudo_de_mas'] = bd
p = [x for x in bd['tablas']['escaner2_pasada'] if x['id'] == pasada][0]
eq('4 · 🔴 el barrido sale en ROJO', cod, 1)
eq('4 · 🔴 …y la pasada queda FALLIDA con el motivo', (p['estado'], 'NO CUADRA antes de la foto' in (p['motivo_fallo'] or '')),
   ('fallida', True))
eq('4 · 🔴 …sin lista para Keepa: no se puede cruzar', (p.get('ruta_lista'), cod2), (None, 1))
eq('4 · 🔴 …y con los recuentos GUARDADOS aunque no cuadre (sin ellos no se sabe dónde se descuadró)',
   (p.get('n_crudo'), p.get('p_sin_gtin'), p.get('p_chase_funko')), (12, 3, 1))

print('\n6 · [todas] (B2) barrido «todas las marcas» → cruce')
cod, log, cod2, log2, bd, pasada, M, e2 = caso('todas', modo='todas')
BDS['todas'] = bd
T = bd['tablas']
p = [x for x in T['escaner2_pasada'] if x['id'] == pasada][0]
eq('6 · el barrido sale en VERDE y la pasada dice su modo', (cod, p['estado'], p['modo']), (0, 'esperando_csv', 'todas'))
eq('6 · 🔴 sin regla leída: ni marcas, ni ofertas, ni «regla activa»', (p['marcas'], p['ofertas'], p['regla_activa']),
   (None, None, None))
eq('6 · 🔴 el barrido NO toca reglas_director, ni para leer',
   [o for o in bd['ops'] if o[2] == 'reglas_director'], [])
eq('6 · 🔴 el Hasbro entra (no hay filtro de marca): foto 6 y marca fuera 0',
   (p['n_foto'], p['p_marca_fuera'], sorted(f['marca'] for f in T['escaner2_foto']).count('Hasbro')), (6, 0, 1))
eq('6 · 🔴 CUADRA: crudo 11 = previas 5 + foto 6',
   (p['n_crudo'], sum(p['p_' + x] for x in e2.PUERTAS_PREVIAS), p['n_foto']), (11, 5, 6))
c = T['escaner2_cruce'][0]
eq('6 · el cruce sale en VERDE, LISTA y cuadrado', (cod2, c['estado'], c['cuadra']), (0, 'lista', True))
_cmp = {r['ean']: r for r in T['escaner2_comparacion']}
eq('6 · 🔴 el Hasbro COMPRAR que el viejo no mira → «diferencia de criterio: marca fuera de la lista del viejo»',
   (_cmp[_ean(M, 7)]['categoria'], _cmp[_ean(M, 7)]['diferencia_criterio'], _cmp[_ean(M, 7)]['explicacion']),
   ('solo_nuevo', True, 'Diferencia de criterio: marca fuera de la lista del viejo (Funko, Ultimate Guard y ofertas)'))
_xl = [k for k in bd['storage']['escaner2'] if k.startswith('heo/%s/%s/Escaner2_HEO_' % (pasada, c['id']))]
_wb = load_workbook(io.BytesIO(base64.b64decode(bd['storage']['escaner2'][_xl[0]])), read_only=True)
_res = {r[0]: r[1] for r in _wb['Resumen'].iter_rows(values_only=True)}
eq('6 · el Excel dice el modo y de dónde sale la lista del viejo',
   (_res.get('Modo del barrido'), 'Funko, Ultimate Guard y ofertas' in (_res.get('Lista del viejo (para comparar)') or '')),
   ('todas las marcas', True))

print('\n7 · [cajas] (B2) cajas con chase REALES de HEO → cruce')
cod, log, cod2, log2, bd, pasada, M, e2 = caso('cajas', escena='cajas')
BDS['cajas'] = bd
T = bd['tablas']
p = [x for x in T['escaner2_pasada'] if x['id'] == pasada][0]
eq('7 · el barrido sale en VERDE', (cod, p['estado']), (0, 'esperando_csv'))
_caja = [f for f in T['escaner2_foto'] if f['producto_heo'] == 'FK87245']
eq('7 · 🔴 FK87245 (disponible) entra con el EAN de su FIGURA, caja con chase de 6 a 60 ÷ 6',
   [(f['ean_core'], f['es_caja'], f['es_chase'], f['uds_caja'], f['precio_catalogo'], f['precio_unidad'], f['origen_ean'])
    for f in _caja], [('889698872454', True, True, 6, 60.0, 10.0, 'gs1_caja')])
eq('7 · FK91391 (pack «w/CH», sin unidades) entra como figura suelta con su EAN',
   [(f['ean_core'], f['es_caja'], f['origen_ean']) for f in T['escaner2_foto'] if f['producto_heo'] == 'FK91391'],
   [('889698913911', False, None)])
eq('7 · 🔴 FK13318 (caja agotada) y FK93057 (el apellido, agotado) → no disponibles; ninguna a «sin EAN»',
   (p['p_no_disponible'], p['p_chase_funko'], {f['producto_heo'] for f in T['escaner2_foto']} & {'FK13318', 'FK93057'}),
   (2, 0, set()))
eq('7 · 🔴 CUADRA: crudo 14 = previas 7 + foto 7',
   (p['n_crudo'], sum(p['p_' + x] for x in e2.PUERTAS_PREVIAS), p['n_foto']), (14, 7, 7))
_lista = base64.b64decode(bd['storage']['escaner2']['heo/%s/eans.txt' % pasada]).decode().split('\n')
eq('7 · a Keepa va el EAN de la figura, nunca el código de la caja',
   ('889698872454' in _lista, any(x.startswith('0110889') for x in _lista)), (True, False))
c = T['escaner2_cruce'][0]
eq('7 · el cruce sale en VERDE, LISTA y cuadrado', (cod2, c['estado'], c['cuadra']), (0, 'lista', True))
_rc = [r for r in T['escaner2_resultado_ean'] if r['foto_id'] == _caja[0]['id']][0]
_pc = [r for r in T['escaner2_resultado_pais'] if r['foto_id'] == _caja[0]['id'] and r['pais'] == 'ES'][0]
eq('7 · 🔴 se valora con el ASIN de la figura común y a 10 €/ud', (_rc['puerta'], _rc['asin'], _pc['pa']),
   ('f', 'B0CAJA0001', 10.0))
_cmp = {r['ean']: r for r in T['escaner2_comparacion']}
eq('7 · 🔴 en la comparación: «diferencia de criterio: el viejo no valora cajas con chase de HEO»',
   (_cmp[_caja[0]['ean_original']]['diferencia_criterio'], _cmp[_caja[0]['ean_original']]['explicacion']),
   (True, 'Diferencia de criterio: el viejo no valora cajas con chase de HEO'))
_xl = [k for k in bd['storage']['escaner2'] if k.startswith('heo/%s/%s/Escaner2_HEO_' % (pasada, c['id']))]
_wb = load_workbook(io.BytesIO(base64.b64decode(bd['storage']['escaner2'][_xl[0]])), read_only=True)
_filas = list(_wb['Puertas'].iter_rows(values_only=True))
_i = {h: k for k, h in enumerate(_filas[0])}
_fc = [f for f in _filas[1:] if f[0] == _caja[0]['ean_original']][0]
eq('7 · 🔴 el Excel (hoja «Puertas») dice «caja con chase · 6 uds», el precio de la CAJA y el EAN de la figura',
   (_fc[_i['Caja']], _fc[_i['Precio caja (€)']], _fc[_i['EAN de la figura']], _fc[_i['Origen del EAN']]),
   ('caja con chase · 6 uds', 60.0, '889698872454', 'código GS1 de la caja'))
_pp = list(_wb['Puertas previas'].iter_rows(values_only=True))
eq('7 · (B2) la hoja «Puertas previas» lleva la marca fuera, con su EAN, nombre, marca y precio',
   [(r[0], r[1], r[2], r[3]) for r in _pp[1:] if r[4] == 'Marca fuera de la lista'],
   [(_ean(M, 7), 'Hasbro fuera', 'Hasbro', 5.0)])

print('\n9 · [elegidas] (B4) barrido «marcas elegidas», solo Funko → cruce')
cod, log, cod2, log2, bd, pasada, M, e2 = caso('elegidas', modo='elegidas',
                                               extra={'MARCAS_ELEGIDAS': '["Funko"]', 'OFERTAS_ELEGIDAS': 'false'})
BDS['elegidas'] = bd
T = bd['tablas']
p = [x for x in T['escaner2_pasada'] if x['id'] == pasada][0]
eq('9 · el barrido sale en VERDE y la pasada guarda su modo, lo elegido y la casilla',
   (cod, p['estado'], p['modo'], p['marcas'], p['ofertas'], p['regla_activa']),
   (0, 'esperando_csv', 'elegidas', ['Funko'], False, None))
eq('9 · 🔴 la foto SOLO trae Funko (HEO lo escribe «FUNKO»: sin distinguir mayúsculas)',
   (p['n_foto'], sorted({f['marca'] for f in T['escaner2_foto']})), (5, ['FUNKO']))
_mf = [a for a in T['escaner2_apartado'] if a['motivo'] == 'marca_fuera']
eq('9 · 🔴 el Hasbro va a la puerta de la marca fuera, «no elegida», con su oferta apuntada',
   [(a['marca'], a['detalle'], a['en_oferta']) for a in _mf], [('Hasbro', "Marca 'Hasbro' no elegida", False)])
eq('9 · 🔴 CUADRA: crudo 11 = previas 6 + foto 5',
   (p['n_crudo'], sum(p['p_' + x] for x in e2.PUERTAS_PREVIAS), p['n_foto']), (11, 6, 5))
eq('9 · 🔴 el modo «elegidas» NO toca reglas_director, ni para leer', [o for o in bd['ops'] if o[2] == 'reglas_director'], [])
c = T['escaner2_cruce'][0]
eq('9 · el cruce sale en VERDE, LISTA y cuadrado', (cod2, c['estado'], c['cuadra']), (0, 'lista', True))
_linea = [x for x in log2.splitlines() if 'ELECCION DE FICHA' in x]
eq('9 · (B5-bis) el log del cruce dice el corpus del cotejo y cuántos vienen de fuera de la foto (la foto 5 + el Hasbro no elegido)',
   [('corpus del cotejo 6 nombres' in x, '1 de fuera de la foto' in x) for x in _linea], [(True, True)])
_xl = [k for k in bd['storage']['escaner2'] if k.startswith('heo/%s/%s/Escaner2_HEO_' % (pasada, c['id']))]
_wb = load_workbook(io.BytesIO(base64.b64decode(bd['storage']['escaner2'][_xl[0]])), read_only=True)
eq('9 · 🔴 el Excel lleva las seis hojas del viejo delante, con las cabeceras de COLS en «Análisis»',
   (_wb.sheetnames, list(next(_wb['Análisis'].iter_rows(values_only=True)))), (HOJAS_VIEJO + HOJAS_E2, list(_COLS)))
eq('9 · 🔴 …y con el FORMATO del Excel viejo de verdad',
   HU.diferencias(_REF, HU.huella(base64.b64decode(bd['storage']['escaner2'][_xl[0]])), solo_hojas=HOJAS_VIEJO, vacias=('Sin_rank',)), [])
_res = {r[0]: r[1] for r in _wb['Resumen'].iter_rows(values_only=True)}
eq('9 · el Resumen dice el modo, las marcas elegidas y la casilla, y «Marca no elegida» en su puerta previa',
   (_res.get('Modo del barrido'), _res.get('Marcas elegidas'), _res.get('Ofertas de cualquier marca'),
    _res.get('Puerta previa · Marca no elegida')), ('marcas elegidas', 'Funko', 'no', 1))
_pp = list(_wb['Puertas previas'].iter_rows(values_only=True))
eq('9 · la hoja «Puertas previas» llama al Hasbro «Marca no elegida»',
   [(r[2], r[4]) for r in _pp[1:] if r[2] == 'Hasbro'], [('Hasbro', 'Marca no elegida')])

print('\n10 · [mala] (B4) una selección que no vale → no se ejecuta nada')
_tmp10 = tempfile.mkdtemp(prefix='e2pp_')
_ruta10 = os.path.join(_tmp10, 'estado.json')
_ini10 = estado_inicial(_e2, _M)
json.dump(_ini10, open(_ruta10, 'w', encoding='utf-8'))
for _nombre, _marcas, _ofertas in (
        ('no es JSON', 'Funko, CID', 'false'),
        ('JSON pero no una lista', '{"marca": "Funko"}', 'false'),
        ('una lista con un número', '["Funko", 7]', 'false'),
        ('comillas dobles dentro de una marca', '["Funko", "Mal\\"a"]', 'false'),
        ('un salto de línea escapado dentro de una marca', '["Fun\\nko"]', 'false'),
        ('un salto de línea de verdad en el input', '["Funko"]\n$(rm -rf /)', 'false'),
        ('un acento grave', '["Fun`ko`"]', 'false'),
        ('selección vacía', '[]', 'false'),
        ('la casilla de ofertas que no es true ni false', '["Funko"]', 'si'),
        ('sin lista', '', 'true')):
    _cod, _log = correr(_ruta10, 'escaner2_heo_barrido.py',
                        {'SUPABASE_SERVICE_KEY': 'svc-de-mentira', 'HEO_USER': 'u', 'HEO_PASS': 'p', 'GITHUB_RUN_ID': '1',
                         'MODO_BARRIDO': 'elegidas', 'MARCAS_ELEGIDAS': _marcas, 'OFERTAS_ELEGIDAS': _ofertas})
    eq('10 · %s → ROJO, «selección de marcas no válida» y NINGÚN cliente creado' % _nombre,
       (_cod, 'ESCANER2_NO_EJECUTADO: selección de marcas no válida' in _log, 'CLIENTES_CREADOS=0' in _log), (1, True, True))
eq('10 · 🔴 …y la base acaba exactamente como empezó (ni una pasada abierta)',
   json.load(open(_ruta10, encoding='utf-8'))['tablas'], _ini10['tablas'])

print('\n8 · [el viejo] (B2) ni una escritura en sus tablas, en ningún caso')
for _n, _bd in BDS.items():
    eq('8 · %s: ni insert, ni update, ni upsert, ni delete en %s' % (_n, ', '.join(TABLAS_VIEJO)),
       [o for o in _bd['ops'] if o[2] in TABLAS_VIEJO and o[1] != 'select'], [])
    eq('8 · %s: …y esas tablas acaban exactamente como empezaron' % _n,
       {t: _bd['tablas'].get(t) for t in TABLAS_VIEJO}, SIEMBRA[_n])
    eq('8 · %s: ni se leen reglas del chase ni la puente (escaner_chase_asin, escaner_memoria, escaner_detalle)' % _n,
       [o for o in _bd['ops'] if o[2] in ('escaner_chase_asin', 'escaner_memoria', 'escaner_detalle')], [])
    # (B4) Y el buzon del viejo: los Excel del escaner 2 NO van a `informes/resultados/` (sus Excel los
    #      leen otros programas por letra de columna); ahi solo sigue el Excel viejo sembrado.
    eq('8 · %s: 🔴 el buzón del viejo (informes) acaba con lo sembrado y nada más' % _n,
       sorted(_bd['storage'].get('informes', {})), ['resultados/Escaneo_HEO_TODAS_20260923_2006.xlsx'])

print('\n5 · [rescate] el run muere a medias → SU fila queda fallida, las demás no se tocan')
_ruta = os.path.join(_tmp, 'rescate.json')
json.dump({'tablas': {
    'escaner2_pasada': [{'id': 'p-colgada', 'run_id': 777, 'estado': 'descargando'},
                        {'id': 'p-otro-run', 'run_id': 888, 'estado': 'descargando'},
                        {'id': 'p-cerrada', 'run_id': 777, 'estado': 'esperando_csv'}],
    'escaner2_cruce': [{'id': 'c-colgado', 'run_id': 777, 'estado': 'cruzando'},
                       {'id': 'c-lista', 'run_id': 777, 'estado': 'lista'}]}, 'storage': {}},
          open(_ruta, 'w', encoding='utf-8'))
for _que in ('pasada', 'cruce'):
    _cod, _log = correr(_ruta, 'escaner2_rescate.py', {'SUPABASE_SERVICE_KEY': 'svc-de-mentira', 'GITHUB_RUN_ID': '777',
                                                       'E2_ARGS': _que})
    eq('5 · el rescate de %s sale en verde' % _que, _cod, 0)
_bd = json.load(open(_ruta, encoding='utf-8'))['tablas']
eq('5 · 🔴 la pasada colgada de ESTE run queda fallida, con el motivo',
   [(f['id'], f['estado'], 'run 777' in (f.get('motivo_fallo') or '')) for f in _bd['escaner2_pasada']],
   [('p-colgada', 'fallida', True), ('p-otro-run', 'descargando', False), ('p-cerrada', 'esperando_csv', False)])
eq('5 · 🔴 el cruce colgado también; el que ya estaba lista, no',
   [(f['id'], f['estado']) for f in _bd['escaner2_cruce']], [('c-colgado', 'fallida'), ('c-lista', 'lista')])
_cod, _log = correr(_ruta, 'escaner2_rescate.py', {'GITHUB_RUN_ID': '777', 'E2_ARGS': 'pasada'})
eq('5 · sin llave, el rescate tampoco corre', (_cod, 'CLIENTES_CREADOS=0' in _log), (1, True))

print()
if fallos:
    print('ROJO: %d comprobaciones fallan: %s' % (len(fallos), ', '.join(fallos)))
    sys.exit(1)
print('VERDE: barrido y cruce de punta a punta; sin llave no se corre, y un cruce que no cuadra sale en ROJO.')
