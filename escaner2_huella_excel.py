# -*- coding: utf-8 -*-
"""ESCANER 2 · LA HUELLA DEL FORMATO DE UN EXCEL DEL ESCANER (encargo B4, 25-sep-2026).

Fernando: «yo necesito exactamente el mismo formato de excel del escaner antiguo». Para poder
PROBARLO sin publicar un Excel real (llevan el precio de coste y este repo es publico), se compara
la HUELLA del formato: sin un solo dato, solo como esta hecho el fichero.

  · las hojas, con su nombre y en su orden;
  · de cada hoja: la cabecera, los anchos de columna, la fila congelada, las tablas (nombre, estilo
    y columnas), el formato condicional (a que columnas se aplica, su formula, sus colores y si
    corta) y, por COLUMNA, el formato de numero, si lleva enlace, el color de la letra y si es una
    FORMULA (con la formula escrita en R1C1 relativo, para que no dependa de la fila).

`huella(contenido)` → dict JSON. `diferencias(a, b)` → lista de lo que cambia, legible.
La huella del Excel viejo de referencia vive en `huella_excel_viejo_heo.json`, sacada del
fichero real `informes/resultados/Escaneo_HEO_TODAS_20260923_2006.xlsx` (23-sep-2026).
"""
import io
import re

_RE_REF = re.compile(r'(\$?)([A-Z]{1,3})(\$?)(\d+)')


def _col(letra):
    n = 0
    for ch in letra:
        n = n * 26 + ord(ch) - 64
    return n


def _relativa(formula, fila, columna):
    """`=J5*M5+AD5` escrita en la fila 5, columna N → `=C[-4]R[0]*…`: la misma para cualquier fila."""
    def sust(m):
        c = ('C%s' % m.group(2)) if m.group(1) else 'C[%d]' % (_col(m.group(2)) - columna)
        r = ('R%s' % m.group(4)) if m.group(3) else 'R[%d]' % (int(m.group(4)) - fila)
        return c + r
    # Los numeros sueltos (el 1,21 del IVA de ESA fila, el 0,03 del ISD) son DATO, no formato: «k».
    return re.sub(r'(?<![\[\-\d.])\d+(?:\.\d+)?', 'k', _RE_REF.sub(sust, formula))


def _color(c):
    try:
        v = c.rgb if c is not None else None
        return v if isinstance(v, str) else None
    except Exception:
        return None


def _columnas_de_rango(rango, cab):
    salida = []
    for trozo in str(rango).split():
        letras = re.findall(r'([A-Z]{1,3})\d+', trozo)
        if letras:
            a, b = _col(letras[0]), _col(letras[-1])
            salida.append([cab[a - 1] if a - 1 < len(cab) else letras[0], cab[b - 1] if b - 1 < len(cab) else letras[-1]])
    return salida


def huella(contenido):
    from openpyxl import load_workbook
    from openpyxl.utils import get_column_letter
    wb = load_workbook(io.BytesIO(contenido))
    hojas = []
    for ws in wb.worksheets:
        cab = [c.value for c in ws[1]] if ws.max_row >= 1 else []
        cab = [str(v) if v is not None else None for v in cab]
        anchos = {}
        for i, h in enumerate(cab, 1):
            d = ws.column_dimensions.get(get_column_letter(i))
            if d is not None and d.width is not None and d.customWidth:
                anchos[h or get_column_letter(i)] = round(float(d.width), 2)
        columnas = {}
        # Por columna, lo que dicen las filas de datos (se miran hasta 60; basta con que una lo tenga).
        for i, h in enumerate(cab, 1):
            formatos, enlaces, letras, formulas = set(), False, set(), set()
            for r in range(2, min(ws.max_row, 61) + 1):
                c = ws.cell(row=r, column=i)
                if c.value is None:
                    continue
                formatos.add(c.number_format)
                enlaces = enlaces or bool(c.hyperlink)
                if c.font is not None and _color(c.font.color):
                    letras.add(_color(c.font.color))
                if isinstance(c.value, str) and c.value.startswith('='):
                    formulas.add(_relativa(c.value, r, i))
            columnas[h or get_column_letter(i)] = {'formato': sorted(formatos), 'enlace': enlaces,
                                                    'color_letra': sorted(letras), 'formula': sorted(formulas)}
        cf = []
        for rango, reglas in ws.conditional_formatting._cf_rules.items():
            for regla in reglas:
                dx = regla.dxf
                cf.append({'columnas': _columnas_de_rango(rango.sqref, cab), 'tipo': regla.type,
                           'operador': regla.operator,
                           'formula': [re.sub(r'\d+', 'n', f) for f in (regla.formula or [])],
                           'corta': bool(regla.stopIfTrue),
                           'relleno': _color(dx.fill.fgColor) if dx is not None and dx.fill is not None else None,
                           'letra': _color(dx.font.color) if dx is not None and dx.font is not None else None})
        tablas = [{'nombre': t.displayName, 'estilo': t.tableStyleInfo.name if t.tableStyleInfo else None,
                   'rayas_filas': bool(t.tableStyleInfo.showRowStripes) if t.tableStyleInfo else None,
                   'columnas': _columnas_de_rango(t.ref, cab)} for t in ws.tables.values()]
        hojas.append({'hoja': ws.title, 'cabecera': cab, 'negrita_cabecera': all(ws.cell(1, i).font.b for i in range(1, len(cab) + 1)),
                      'congelada': ws.freeze_panes, 'anchos': anchos, 'columnas': columnas,
                      'formato_condicional': sorted(cf, key=lambda x: (str(x['columnas']), str(x['formula']), str(x['operador']))),
                      'tablas': tablas})
    return {'hojas': hojas}


def diferencias(ref, otro, solo_hojas=None, ignorar_columnas=(), vacias=()):
    """Lo que cambia del Excel `otro` respecto a la referencia, hoja a hoja. `solo_hojas`: comparar
    esas (en su orden, al principio del libro); `ignorar_columnas`: {hoja: [columnas]} cuyas
    propiedades de DATO (formato, enlace, color, formula) no se comparan porque la referencia no
    trae filas con dato ahi. `vacias`: hojas que en `otro` no llevan ni una fila; el viejo las escribe
    con la cabecera «(vacio)» (su `hoja()`), y eso es lo que se exige en ellas."""
    difs = []
    hr = [h for h in ref['hojas'] if solo_hojas is None or h['hoja'] in solo_hojas]
    ho = otro['hojas'][:len(hr)]
    nombres_r, nombres_o = [h['hoja'] for h in hr], [h['hoja'] for h in ho]
    if nombres_r != nombres_o:
        return ['hojas: %s ≠ %s' % (nombres_r, nombres_o)]
    for a, b in zip(hr, ho):
        n = a['hoja']
        if n in vacias:
            if b['cabecera'] != ['(vacio)']:
                difs.append('%s · vacía: el viejo pone «(vacio)» y aquí %r' % (n, b['cabecera']))
            continue
        for k in ('cabecera', 'negrita_cabecera', 'congelada', 'anchos', 'formato_condicional', 'tablas'):
            if a[k] != b[k]:
                difs.append('%s · %s: %r ≠ %r' % (n, k, a[k], b[k]))
        ign = set((ignorar_columnas or {}).get(n, ()))
        for col, pa in a['columnas'].items():
            pb = b['columnas'].get(col)
            if col in ign or pb is None:
                continue
            # Lo de DATO solo se compara donde las DOS hojas tienen dato en esa columna: una columna
            # vacia no dice nada de su formato.
            if pa['formato'] and pb['formato']:
                for k in ('formato', 'enlace', 'color_letra'):
                    if pa[k] != pb[k]:
                        difs.append('%s · %s · %s: %r ≠ %r' % (n, col, k, pa[k], pb[k]))
            # Y ni una formula que el viejo no escriba.
            if not set(pb['formula']) <= set(pa['formula']):
                difs.append('%s · %s · formula: %r no es de las del viejo %r' % (n, col, pb['formula'], pa['formula']))
    return difs
