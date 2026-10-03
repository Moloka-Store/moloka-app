# -*- coding: utf-8 -*-
"""TEST · LOS PACKS DE OSMA QUE EXISTEN EN AMAZON Y NO VENDEMOS (encargo AN, plano Y, tramo 7, 03-oct-2026).

Sin red y sin base de verdad: las piezas puras de escaner2_osma_packs.py y los dos programas (encolar y valorar) contra
una base de mentira EN EL MISMO PROCESO.

  A  LA COLA (`familias_para_cola`): nuestras fichas por codigo y por EAN (el suelto, la de menos unidades), las puertas
     d/e/f con el ASIN del cruce y sus unidades de OSMA por la regla del AM, la marca y el titulo de Amazon ES (nunca el
     nombre de OSMA), y lo que se queda fuera (puerta c que no es nuestra, sin marca en Amazon).
  B  EL N DEL PACK contra el suelto: Kukident Aktiv Plus 99 → B08629V7YQ pack de 2 por TRES señales; una sola señal →
     posible pack; el que lleva lo mismo que el suelto → no es pack; señales que se contradicen → posible pack (el N mayor).
  C  LA VALORACION: coste 2 × 3,52 = 7,04 € (con 2.500 € de pedido y 168,19 € de porte); la cuenta es la de
     `escaner2_motor.calcular_pais`; sin ventas de Keepa NUNCA es COMPRAR; un parecido tampoco; bajo el corte, no se vende;
     el nuestro (Lenor B07HCJQ45L) queda fuera y se dice.
  D  EL EXCEL: solo «Análisis» y «Resumen»; «Análisis» con LAS MISMAS columnas que el Análisis de OSMA (solo ES) y las
     cuatro propias detras; PA y Decisión vivas (el posible pack sin la rama de COMPRAR); las dos casillas amarillas.
  E  LA LISTA DEL PROXIMO BARRIDO: los EAN de los packs que no estan ya.
  F  LOS PROGRAMAS contra la base de mentira: encolar apunta la cola de SU cruce (y no dos veces); valorar deja el Excel con
     la ruta de SU fila (la forma que exige la papelera), la fila en osma_packs_excel y el Telegram; no valora con
     familias esperando ni dos veces el mismo cruce.
  G  QUE LOS TESTS MUERDEN: con «una señal basta» y sin el tope de las ventas, se ponen en rojo.
"""
import io
import json
import os
import sys
import types
from datetime import datetime, timezone

AQUI = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, AQUI)

import escaner2_motor as e2  # noqa: E402
import escaner2_osma as eo  # noqa: E402
import escaner2_osma_packs as op  # noqa: E402

FALLOS = []
N_OK = [0]


def eq(nombre, obtenido, esperado):
    if obtenido == esperado:
        N_OK[0] += 1
        print('OK  %s' % nombre)
    else:
        FALLOS.append(nombre)
        print('XX  %s\n    obtenido: %r\n    esperado: %r' % (nombre, obtenido, esperado))


AHORA = datetime(2026, 10, 3, 14, 0, tzinfo=timezone.utc)
PRODUCTOS = [
    {'id': 'p-kuki', 'asin': 'B001PASC5E', 'ean': '4002448039440', 'nombre': 'Kukident Active Plus 99', 'activo': True, 'es_chase': False,
     'unidades_por_pack': None, 'iva_pct': 0.21, 'stock_moloka': 0},
    {'id': 'p-lenor', 'asin': 'B014DGG0OQ', 'ean': '8001090747723', 'nombre': 'Lenor se está secando. abril fresco, paquete de 34',
     'activo': True, 'es_chase': False, 'unidades_por_pack': None, 'iva_pct': 0.21, 'stock_moloka': 0},
    {'id': 'p-lenor2', 'asin': 'B07HCJQ45L', 'ean': '8001090747723', 'nombre': 'Pack 2. Lenor se está secando. abril fresco, paquete de 34',
     'activo': True, 'es_chase': False, 'unidades_por_pack': None, 'iva_pct': 0.21, 'stock_moloka': 0},
]
TIT_KUKI = 'Kukident Aktiv Plus Zahnersatz-Reinigungstabletten, 99 pzas Tabletas'

# ── A · LA COLA ─────────────────────────────────────────────────────────────────────────
FOTO = [
    {'id': 'f1', 'producto_heo': '4213', 'ean_original': '4002448039440', 'ean_core': '4002448039440', 'nombre': 'Kukident Aktiv Plus 99er',
     'marca': 'Kukident', 'precio_unidad': 3.52, 'precio_catalogo': 3.299, 'fin_de_vida': False},
    {'id': 'f2', 'producto_heo': '18459', 'ean_original': '8001090747723', 'ean_core': '8001090747723', 'nombre': 'Lenor Trocknertücher 34er Aprilfrisch',
     'marca': 'Lenor', 'precio_unidad': 1.71, 'precio_catalogo': 1.60, 'fin_de_vida': True},
    {'id': 'f3', 'producto_heo': '1929', 'ean_original': '5054563000001', 'ean_core': '5054563000001', 'nombre': 'Corega Tabs 108er',
     'marca': 'Corega', 'precio_unidad': 4.48, 'precio_catalogo': 4.199, 'fin_de_vida': False},
    {'id': 'f4', 'producto_heo': '7777', 'ean_original': '4000000000007', 'ean_core': '4000000000007', 'nombre': 'Algo que no se vende',
     'marca': 'X', 'precio_unidad': 1.0, 'precio_catalogo': 0.9, 'fin_de_vida': False},
    {'id': 'f5', 'producto_heo': '8888', 'ean_original': '4000000000008', 'ean_core': '4000000000008', 'nombre': 'Sin marca en Amazon',
     'marca': 'Y', 'precio_unidad': 1.0, 'precio_catalogo': 0.9, 'fin_de_vida': False},
]
RESULTADOS = [
    {'foto_id': 'f1', 'puerta': 'c', 'asin': 'B001PASC5E'},
    {'foto_id': 'f2', 'puerta': 'f', 'asin': 'B0794VHRVZ'},
    {'foto_id': 'f3', 'puerta': 'e', 'asin': 'B0NUEVA001'},
    {'foto_id': 'f4', 'puerta': 'c', 'asin': 'B0OTRO0004'},
    {'foto_id': 'f5', 'puerta': 'd', 'asin': 'B0SINMARCA'},
]
ENLACES = {'4213': {'producto_id': 'p-kuki', 'ean_ficha': '4002448039440', 'asins': ['B001PASC5E']}}
CSV_TIT = {'B001PASC5E': {'titulo': TIT_KUKI, 'marca': 'Kukident'},
           'B0NUEVA001': {'titulo': 'Corega Tabs Intensiv limpiador, 108 tabletas', 'marca': 'Corega'},
           'B0SINMARCA': {'titulo': 'Algo', 'marca': ''}}
AMZ = {'B014DGG0OQ': {'titulo': 'Lenor Aprilfrisch Toallitas de secado, 34 toallitas', 'marca': 'LENOR'}}
filas, cuentas = op.familias_para_cola('cruce-1', FOTO, RESULTADOS, ENLACES, PRODUCTOS, CSV_TIT, {}, AMZ, {})
por = {f['codigo_osma']: f for f in filas}
eq('A1 · tres familias: Kukident (nuestra por código, aunque su puerta sea c), Lenor (nuestra por EAN, puerta f) y Corega (puerta e)',
   sorted(por), ['18459', '1929', '4213'])
eq('A2 · Kukident: suelto B001PASC5E ×1, marca y título de Amazon ES (del CSV), PA con porte y Price_net, nuestra sin puerta',
   [por['4213'][k] for k in ('asin_ref', 'factor_ref', 'marca', 'marca_origen', 'titulo_ref', 'titulo_origen', 'pa_unidad', 'price_net', 'puerta', 'nuestra', 'eans')],
   ['B001PASC5E', 1, 'Kukident', 'csv', TIT_KUKI, 'csv', 3.52, 3.299, None, True, ['4002448039440']])
eq('A3 · Lenor: el suelto es NUESTRA ficha de 1 (B014DGG0OQ), no el pack de 2 ni el ASIN del cruce; título de amz_ficha; no habrá más',
   [por['18459'][k] for k in ('asin_ref', 'factor_ref', 'marca', 'marca_origen', 'titulo_origen', 'puerta', 'nuestra', 'fin_de_vida')],
   ['B014DGG0OQ', 1, 'LENOR', 'amz_ficha', 'amz_ficha', 'f', True, True])
eq('A4 · nunca el nombre alemán de OSMA como título', all(f['titulo_ref'] != f['nombre_osma'] for f in filas), True)
eq('A5 · fuera: la puerta c que no es nuestra (no se cuenta) y la que no tiene marca en Amazon; de puerta 2 y nuestras 2 (el Lenor, las dos)', [cuentas['sin_marca'], cuentas['puertas'], cuentas['nuestras']], [1, 2, 2])
# Un ASIN del cruce que es un pack de Amazon (la regla del AM): sus unidades de OSMA van en la cola; posible pack → NULL.
_sen = {'B0NUEVA001': {'n_art': '', 'valor_ud': '324', 'tipo_ud': 'unidad', 'paquete': '', 'tamano': '108 unidad (Paquete de 3)', 'titulo': ''}}
f_amz, _ = op.familias_para_cola('cruce-1', FOTO[2:3], RESULTADOS[2:3], {}, PRODUCTOS, CSV_TIT, _sen, {}, {})
_sen_d = {'B0NUEVA001': dict(_sen['B0NUEVA001'], tamano='')}
f_dud, _ = op.familias_para_cola('cruce-1', FOTO[2:3], RESULTADOS[2:3], {}, PRODUCTOS, CSV_TIT, _sen_d, {}, {})
eq('A6 · la Corega del cruce es un pack de 3 en Amazon (dos señales): factor_ref 3; con una sola señal, NULL (posible pack)',
   [f_amz[0]['factor_ref'], f_dud[0]['factor_ref']], [3, None])

# ── B · EL N DEL PACK ───────────────────────────────────────────────────────────────────
SUELTO = {'asin': 'B001PASC5E', 'clase': 'suelto', 'paquete': 1, 'valor_ud': 99, 'tipo_ud': 'Count', 'n_art': None, 'titulo': TIT_KUKI}
PACK = {'asin': 'B08629V7YQ', 'clase': 'mismo', 'porque': 'lista el EAN de la unidad (4002448039440)', 'paquete': 2, 'valor_ud': 198, 'tipo_ud': 'Count',
        'n_art': None, 'titulo': 'Kukident Active Plus Tabletas limpiadoras, 2 x 99 unidades', 'eans': ['4002448039440', '4002448126515'],
        'rank': 3401, 'rank_categoria': 'Salud y cuidado personal', 'resultado': 'dato', 'precio': 15.90, 'canal': 'BB-FBA',
        'ref_eur': 2.39, 'fee_fba': 3.60}
v = op.n_del_pack(PACK, SUELTO)
eq('B1 · Kukident 2 × 99: pack de 2 por tres señales (paquete, contenido y título; el «99» del título es del suelto)',
   [v['estado'], v['n'], v['senales']],
   [eo.PACK_SI, 2, 'paquete 2 (suelto 1) → 2 · contenido 198 Count (suelto 99) → 2 · título «2 x 99» → 2'])
eq('B2 · una sola señal (solo el paquete): posible pack, con su N',
   [op.n_del_pack({'paquete': 2, 'titulo': 'Kukident Aktiv Plus Tabletas'}, SUELTO)[k] for k in ('estado', 'n')], [eo.PACK_DUDOSO, 2])
eq('B3 · lleva lo mismo que el suelto (contenido 99): no es pack, diga lo que diga el paquete',
   op.n_del_pack({'paquete': 2, 'valor_ud': 99, 'tipo_ud': 'Count', 'titulo': ''}, SUELTO)['estado'], eo.PACK_NO)
eq('B4 · señales que se contradicen (paquete 2, contenido 3): posible pack con el N MAYOR (el coste más prudente)',
   [op.n_del_pack({'paquete': 2, 'valor_ud': 297, 'tipo_ud': 'Count', 'titulo': ''}, SUELTO)[k] for k in ('estado', 'n')], [eo.PACK_DUDOSO, 3])
eq('B5 · sin ninguna señal ≥ 2: no es pack', op.n_del_pack({'paquete': 1, 'titulo': 'Kukident Aktiv Plus'}, SUELTO)['estado'], eo.PACK_NO)
eq('B6 · sin las lecturas del suelto, el paquete se compara con 1 y se dice', op.n_del_pack({'paquete': 2, 'titulo': 'pack de 2'}, None)['senales'].startswith('sin las lecturas del suelto'), True)

# ── C · LA VALORACION ───────────────────────────────────────────────────────────────────
M = e2.cargar_motor()
M.poner_catalogo_propio([p for p in PRODUCTOS if p['activo']])
import en_mi_bd  # noqa: E402
M.poner_foto_fba(en_mi_bd.foto_de_filas([], None, None, error='sin foto en el test'))
COLA = [dict(por['4213'], id='c-kuki', estado='buscada'), dict(por['18459'], id='c-lenor', estado='buscada')]
UNA = {'asin': 'B0UNASENAL', 'clase': 'mismo', 'porque': 'x', 'paquete': 2, 'titulo': 'Kukident Aktiv Plus Tabletas', 'resultado': 'dato',
       'precio': 13.0, 'canal': 'SIN BB', 'ref_eur': 1.95, 'fee_fba': 3.6, 'rank': None}
PARECIDO = dict(PACK, asin='B0KUKI0066', clase='parecido', porque='parecido: revisar (comparte «aktiv plus»; le falta «99»)', titulo='Kukident Aktiv Plus 66, pack de 2')
NOPACK = dict(PACK, asin='B0MISMO001', paquete=1, valor_ud=99, titulo='Kukident Aktiv Plus 99')
CAND = [dict(SUELTO, cola_id='c-kuki'), dict(PACK, cola_id='c-kuki'), dict(UNA, cola_id='c-kuki'), dict(PARECIDO, cola_id='c-kuki'),
        dict(NOPACK, cola_id='c-kuki'),
        {'cola_id': 'c-lenor', 'asin': 'B07HCJQ45L', 'clase': 'nuestro', 'porque': 'es de una ficha nuestra: no se valora'}]
NOV_KEEPA = [{'pais': 'ES', 'consultada_en': '2026-09-30T07:00:00Z', 'fichas': [{'asin': 'B08629V7YQ', 'caidas_30d': 30}, {'asin': 'B0KUKI0066', 'caidas_30d': 40},
                                                                                 {'asin': 'B0UNASENAL', 'caidas_30d': 50}]}]
packs, cuentas = op.valorar(COLA, CAND, M, 6, NOV_KEEPA, None, AHORA)
pp = {p['c']['asin']: p for p in packs}
k = pp['B08629V7YQ']
calc = e2.calcular_pais('ES', op.rec_de_candidato(PACK), 7.04, '4002448039440', M)
eq('C1 · Kukident: coste 2 × 3,52 = 7,04 € (la PA de una unidad con porte, ya redondeada)', [k['n'], k['n_osma'], k['pa']], [2, 2, 7.04])
eq('C2 · la cuenta es la de calcular_pais (el viejo), con la comisión de Amazon en % del precio y la FBA tal cual',
   [k['calc']['margen'], k['calc']['beneficio'], k['calc']['ref_pct'], k['calc']['fee_fba'], k['calc']['precio_venta'], k['calc']['canal']],
   [calc['margen'], calc['beneficio'], 2.39 / 15.9 * 100, 3.6, 15.9, 'BB-FBA'])
eq('C3 · con ventas de Keepa (30 caídas, guardadas hace 3 días), la decisión es la del viejo, sin tope',
   [k['decision'], k['tope'], k['caidas'], k['ventas_de']], [M.decision_de(calc['margen']), False, 30, 'Keepa (novedades, 30/09)'])
eq('C4 · el posible pack (una señal) y el parecido: con tope (nunca COMPRAR) y su porqué',
   [pp['B0UNASENAL']['tope'], pp['B0UNASENAL']['decision'] != 'COMPRAR', pp['B0KUKI0066']['tope'], pp['B0KUKI0066']['decision'] != 'COMPRAR',
    any(m.startswith(eo.TEXTO_POSIBLE_PACK) for m in pp['B0UNASENAL']['motivos']), op.TEXTO_PARECIDO in pp['B0KUKI0066']['motivos']],
   [True, True, True, True, True, True])
eq('C5 · el que lleva lo mismo que el suelto no entra; el NUESTRO (Lenor B07HCJQ45L) queda fuera y se dice',
   ['B0MISMO001' in pp, cuentas['no_pack'], cuentas['nuestros']], [False, 1, [('18459', 'B07HCJQ45L')]])
packs_sv, cuentas_sv = op.valorar(COLA, CAND, M, 6, [], None, AHORA)
eq('C6 · SIN ventas de Keepa: ninguno es COMPRAR, todos con «sin ventas de Keepa: entra en la lista del próximo PRO»',
   [any(p['decision'] == 'COMPRAR' for p in packs_sv), all(op.TEXTO_SIN_VENTAS in p['motivos'] for p in packs_sv), cuentas_sv['sin_ventas']],
   [False, True, 3])
viejo = [dict(NOV_KEEPA[0], consultada_en='2026-09-10T07:00:00Z')]
eq('C7 · unas ventas de hace más de 15 días no valen (como sin ventas)', op.ventas_de('B08629V7YQ', viejo, None, AHORA), (None, None))
eq('C8 · las del CSV del PRO de esta pasada valen si se subió hace 15 días o menos',
   op.ventas_de('B08629V7YQ', [], {'caidas': {'B08629V7YQ': 12}, 'fecha': '2026-10-02T09:00:00Z'}, AHORA), (12, 'Keepa (CSV del PRO, 02/10)'))
bajo = [{'pais': 'ES', 'consultada_en': '2026-09-30T07:00:00Z', 'fichas': [{'asin': 'B08629V7YQ', 'caidas_30d': 6}]}]
_p, _c = op.valorar(COLA, CAND, M, 6, bajo, None, AHORA)
eq('C9 · con 6 caídas (el corte es 7 o más) no se vende: fuera de «Análisis», contado', ['B08629V7YQ' in {p['c']['asin'] for p in _p}, _c['no_se_vende']],
   [False, [('4213', 'B08629V7YQ', 6)]])
eq('C10 · la comisión: ReferralFee ÷ precio, como rec_de_fila de las novedades',
   op.rec_de_candidato(PACK), {'asin': 'B08629V7YQ', 'titulo': PACK['titulo'], 'rank': 3401, 'rank90': None, 'buybox': 15.9, 'es_fba': True,
                               'nuevo': None, 'compct': 2.39 / 15.9 * 100, 'fba': 3.6})

# ── D · EL EXCEL ────────────────────────────────────────────────────────────────────────
from openpyxl import load_workbook  # noqa: E402
INFO = {'cruce': 'cruce-1', 'pasada': 'pasada-1', 'porte': {'pedido_previsto': 2500.0, 'gastos_envio': 168.19}, 'umbral': 6, 'avisos': []}
wb = load_workbook(io.BytesIO(op.escribir_excel(packs, cuentas, M, INFO)))
an = wb['Análisis']
cab = [c.value for c in an[1]]
# La cabecera del Análisis de OSMA (solo ES), de verdad: la de una fila de OSMA por su propio código.
_fo = [{'id': 'x', 'ean_original': '4002448039440', 'ean_core': '4002448039440', 'nombre': 'n', 'marca': 'm', 'precio_unidad': 3.52,
        'precio_catalogo': 3.299, 'fin_de_vida': False, 'aviso_caja': None, '_factor': 1, '_sin_porte': 3.299}]
_re = [{'foto_id': 'x', 'puerta': 'f', 'motivo': 'f', 'detalle': '', 'asin': 'B001PASC5E', 'paises': {'ES': dict(calc, caidas_30d=30)}, 'fichas': None}]
cab_osma = [c.value for c in eo.excel_como_el_viejo_osma(_fo, _re, [], M, paises=['ES'])['Análisis'][1]]
eq('D1 · solo dos hojas: «Análisis» y «Resumen»', wb.sheetnames, ['Análisis', 'Resumen'])
eq('D2 · «Análisis»: las MISMAS columnas que el Análisis de OSMA (solo ES) y detrás las cuatro propias', cab, cab_osma + op.COLUMNAS_EXTRA)
fila = next(r for r in range(2, an.max_row + 1) if an.cell(row=r, column=cab.index('ASIN') + 1).value == 'B08629V7YQ')
celda = lambda nombre, r=fila: an.cell(row=r, column=cab.index(nombre) + 1).value  # noqa: E731
from openpyxl.utils import get_column_letter  # noqa: E402
l_sp = get_column_letter(cab.index(eo.COLUMNA_SIN_PORTE) + 1)
eq('D3 · la fila de Kukident: PA viva (2 × la unidad con porte redondeada), 2 unidades, sus señales, el suelto y el puesto',
   [celda('PA (€)'), celda(eo.COLUMNA_SIN_PORTE), celda('Unidades del pack'), celda('ASIN suelto de referencia'), celda('Puesto de ventas de Amazon'),
    celda('Señales del pack').startswith('pack de 2 · paquete 2 (suelto 1) → 2')],
   ['=2*ROUND(%s%d/2*(1+IF(N(PedidoPrevisto)>0,N(PorteEnvio)/PedidoPrevisto,0)),2)' % (l_sp, fila), 6.598, 2, 'B001PASC5E',
    '3401 · Salud y cuidado personal', True])
fila_u = next(r for r in range(2, an.max_row + 1) if an.cell(row=r, column=cab.index('ASIN') + 1).value == 'B0UNASENAL')
eq('D4 · la Decisión viva: la de Kukident con COMPRAR; la del posible pack SIN la rama de COMPRAR',
   ['"COMPRAR"' in str(celda('Decisión')), '"COMPRAR"' in str(celda('Decisión', fila_u))], [True, False])
wr = wb['Resumen']
eq('D5 · «Resumen»: pedido 2.500 y porte 168,19 en las dos casillas amarillas, con nombre; sin línea final de control',
   [wr['B2'].value, wr['B3'].value, wr['B2'].fill.start_color.rgb[-6:], sorted(wb.defined_names), any('control' in str(c.value).lower() for c in wr['A'])],
   [2500.0, 168.19, 'FFFF00', ['PedidoPrevisto', 'PorteEnvio'], False])
eq('D6 · «Resumen» dice el nuestro encontrado y excluido', any('18459 · B07HCJQ45L' in str(c.value) for c in wr['B']), True)
eq('D7 · sale aunque no haya ningún COMPRAR (sin ventas)', load_workbook(io.BytesIO(op.escribir_excel(packs_sv, cuentas_sv, M, INFO))).sheetnames, ['Análisis', 'Resumen'])
eq('D8 · la ruta: osma/<pasada>/<id de SU fila>/Escaner2_OSMA_Packs_<AAAAMMDD_HHMM>.xlsx',
   op.ruta_excel('655b3e06-7637-4cd5-92de-789a094a504c', 'e0000000-0000-0000-0000-0000000000e1', '20261003_1600'),
   'osma/655b3e06-7637-4cd5-92de-789a094a504c/e0000000-0000-0000-0000-0000000000e1/Escaner2_OSMA_Packs_20261003_1600.xlsx')

# ── E · LA LISTA DEL PROXIMO BARRIDO ─────────────────────────────────────────────────────
eq('E1 · los EAN de los packs que la lista no trae (sin ceros delante, sin repetir); el nuestro no; el que no tiene EAN, contado',
   op.eans_de_packs([{'clase': 'mismo', 'eans': ['4002448039440', '04002448126515']}, {'clase': 'parecido', 'eans': []},
                     {'clase': 'nuestro', 'eans': ['8001090999999']}, {'clase': 'mismo', 'eans': ['4002448126515']}], ['4002448039440'], M),
   (['04002448126515'], 1))

# ── F · LOS PROGRAMAS, CONTRA UNA BASE DE MENTIRA ────────────────────────────────────────
import uuid  # noqa: E402
DEFAULTS = {'osma_packs_cola': {'estado': 'pendiente', 'intentos': 0, 'id': lambda: str(uuid.uuid4())}}


class _R:
    def __init__(self, data, count=None):
        self.data, self.count = data, count


class _Q:
    def __init__(self, bd, t):
        self.bd, self.t, self.f, self.op, self.cnt, self.o, self.rg, self.lim, self.payload = bd, t, [], 'select', False, None, None, None, None

    def select(self, _c='*', count=None):
        self.cnt = count == 'exact'
        return self

    def insert(self, filas):
        self.op, self.payload = 'insert', filas if isinstance(filas, list) else [filas]
        return self

    def eq(self, k, v):
        self.f.append(lambda x, k=k, v=v: str(x.get(k)) == str(v))
        return self

    def gte(self, k, v):
        self.f.append(lambda x, k=k, v=v: str(x.get(k)) >= str(v))
        return self

    def in_(self, k, vs):
        self.f.append(lambda x, k=k, vs=vs: x.get(k) in vs)
        return self

    def is_(self, k, v):
        self.f.append(lambda x, k=k: x.get(k) is None)
        return self

    def order(self, c, desc=False):
        self.o = (c, desc)
        return self

    def range(self, a, b):
        self.rg = (a, b)
        return self

    def limit(self, n):
        self.lim = n
        return self

    def execute(self):
        filas = self.bd['t'].setdefault(self.t, [])
        if self.op == 'insert':
            # Los DEFAULT de las tablas, como los pone Postgres (osma_packs_cola: estado 'pendiente', su id).
            nuevas = json.loads(json.dumps(self.payload, default=str))
            for x in nuevas:
                for col, v in DEFAULTS.get(self.t, {}).items():
                    x.setdefault(col, v() if callable(v) else v)
            filas.extend(nuevas)
            return _R(self.payload)
        sel = [x for x in filas if all(g(x) for g in self.f)]
        if self.o:
            sel = sorted(sel, key=lambda x: str(x.get(self.o[0])), reverse=self.o[1])
        n = len(sel)
        if self.rg:
            sel = sel[self.rg[0]:self.rg[1] + 1]
        if self.lim is not None:
            sel = sel[:self.lim]
        return _R(sel, n if self.cnt else None)


class _Cubo:
    def __init__(self, bd):
        self.bd = bd

    def upload(self, ruta, datos, _o=None):
        self.bd['s'][ruta] = datos
        return {'Key': ruta}

    def download(self, ruta):
        if ruta not in self.bd['s']:
            raise RuntimeError('404')
        return self.bd['s'][ruta]

    def list(self, carpeta=None, opciones=None):
        op_ = dict({'limit': 100, 'offset': 0}, **(opciones or {}))
        pref = carpeta.rstrip('/') + '/'
        nombres = sorted(r[len(pref):] for r in self.bd['s'] if r.startswith(pref) and '/' not in r[len(pref):])
        return [{'name': x, 'created_at': '2026-10-02T09:00:00Z'} for x in nombres[op_['offset']:op_['offset'] + op_['limit']]]


class _Sb:
    def __init__(self, bd):
        self.bd = bd
        self.storage = types.SimpleNamespace(from_=lambda _b: _Cubo(bd))

    def table(self, t):
        return _Q(self.bd, t)


PASADA = '655b3e06-7637-4cd5-92de-789a094a504c'
CRUCE = '899bd29e-cf2d-47f1-8898-9583d181bc08'
_csv = ('ASIN,Título,Marca,Número de artículos,Detalles de la unidad: Valor de la unidad,Detalles de la unidad: Tipo de unidad,Paquete: Cantidad,Tamaño\n'
        'B001PASC5E,"%s",Kukident,,99,unidad,1,\nB0NUEVA001,"Corega Tabs Intensiv limpiador, 108 tabletas",Corega,,108,unidad,1,\n' % TIT_KUKI)
BD = {'t': {
    'escaner2_cruce': [{'id': CRUCE, 'estado': 'lista', 'pasada_id': PASADA, 'run_id': 4242, 'creado_en': '2026-10-03T10:00:00Z'}],
    'escaner2_pasada': [{'id': PASADA, 'proveedor': 'OSMA'}],
    'escaner2_foto': [dict(f, pasada_id=PASADA) for f in FOTO],
    'escaner2_resultado_ean': [dict(r, cruce_id=CRUCE, id='r%d' % i) for i, r in enumerate(RESULTADOS)],
    'productos': PRODUCTOS, 'amz_ficha': [dict(asin='B014DGG0OQ', pais='ES', **AMZ['B014DGG0OQ'])], 'keepa_escaparate': [],
    'escaner2_parametros': [{'proveedor': 'OSMA', 'umbral_caidas_30d': 6}],
    'nov_keepa': [dict(NOV_KEEPA[0])], 'inventario_fba': [],
}, 's': {'osma/%s/barrido.json' % PASADA: json.dumps({'enlaces': ENLACES, 'porte': INFO['porte']}).encode(),
         'osma/%s/csv/20261002_0900_es.csv' % PASADA: _csv.encode('utf-8')}}
sb = _Sb(BD)
log = []
imp = lambda *a, **k: log.append(' '.join(str(x) for x in a))  # noqa: E731
n = op.encolar(sb, PASADA, '4242', imprimir=imp)
eq('F1 · encolar: las familias de SU cruce, con el título y la marca del CSV de la pasada', [n, sorted(x['codigo_osma'] for x in BD['t']['osma_packs_cola'])],
   [3, ['18459', '1929', '4213']])
eq('F2 · encolar dos veces el mismo cruce no apunta otra cola', [op.encolar(sb, PASADA, '4242', imprimir=imp), len(BD['t']['osma_packs_cola'])], [0, 3])
eq('F3 · 🔴 el registro (repo público) solo lleva recuentos: ni EAN, ni ASIN, ni nombres',
   [x for x in ('4002448039440', 'B001PASC5E', 'Kukident', 'Corega', '3.52') if any(x in linea for linea in log)], [])
try:
    op.valorar_cruce(sb, CRUCE, '5151', ahora=AHORA, imprimir=imp, enviar=lambda *a: None)
    eq('F4 · con familias esperando la búsqueda, NO se valora', 'valoró', 'no valora')
except op.FalloPacks as ex:
    eq('F4 · con familias esperando la búsqueda, NO se valora', str(ex), 'aún hay familias esperando la búsqueda en Amazon')
# El cartero ha buscado (lo que guardaría osma_packs_guardar).
_ids = {x['codigo_osma']: x for x in BD['t']['osma_packs_cola']}
for x in BD['t']['osma_packs_cola']:
    x['id'], x['estado'] = 'cola-' + x['codigo_osma'], 'buscada'
BD['t']['osma_packs_candidato'] = [dict(SUELTO, cola_id='cola-4213'), dict(PACK, cola_id='cola-4213'), dict(UNA, cola_id='cola-4213'),
                                   {'cola_id': 'cola-18459', 'asin': 'B07HCJQ45L', 'clase': 'nuestro', 'porque': 'nuestro'}]
enviados = []
fila_x = op.valorar_cruce(sb, CRUCE, '5151', ahora=AHORA, imprimir=imp, enviar=lambda texto, env, post, imprimir: enviados.append(texto))
import re as _re_  # noqa: E402
eq('F5 · valorar: la fila de osma_packs_excel con la ruta de SU fila (la forma de la papelera) y sus cuentas',
   [bool(_re_.fullmatch(r'osma/%s/%s/Escaner2_OSMA_Packs_20261003_1600\.xlsx' % (PASADA, fila_x['id']), fila_x['ruta_excel'])),
    fila_x['n_familias'], fila_x['n_packs'], fila_x['cruce_id'], fila_x['run_id'], fila_x['ruta_excel'] in BD['s']],
   [True, 3, 2, CRUCE, 5151, True])
eq('F6 · el Telegram sale (con el aviso que ya existe) y dice cuántos', [len(enviados), enviados[0].startswith('🟢 <b>Packs de OSMA en Amazon</b>: 2 que no vendemos')],
   [1, M.decision_de(calc['margen']) == 'COMPRAR'])
eq('F7 · no se valora dos veces el mismo cruce (con su Excel en la biblioteca)', op.valorar_cruce(sb, CRUCE, '5151', ahora=AHORA, imprimir=imp,
                                                                                               enviar=lambda *a: None), None)
eq('F8 · 🔴 el registro de valorar tampoco lleva EAN, ASIN ni nombres',
   [x for x in ('4002448039440', 'B08629V7YQ', 'Kukident', '7.04') if any(x in linea for linea in log)], [])

# El barrido (escaner2_osma_barrido.py) con la misma base de mentira: los EAN de los packs del ultimo cruce buscado, detras.
for x in BD['t']['osma_packs_cola']:
    x['buscada_en'] = '2026-10-03T13:20:00Z'
BD['t']['osma_packs_candidato'].append({'cola_id': 'cola-4213', 'asin': 'B0PARECIDO', 'clase': 'parecido', 'eans': []})
_ent = {k: os.environ.get(k) for k in ('SUPABASE_SERVICE_KEY', 'SUPABASE_URL')}
os.environ.update(SUPABASE_SERVICE_KEY='svc-de-mentira', SUPABASE_URL='https://doble.invalid')
_sup = types.ModuleType('supabase')
_sup.create_client = lambda url, llave: sb
_previo = sys.modules.get('supabase')
sys.modules['supabase'] = _sup
try:
    import escaner2_osma_barrido as barrido  # noqa: E402
    eq('E2 · el barrido añade detrás los EAN de los packs que encontró el cartero en el último PRO (el parecido sin EAN, contado)',
       barrido.packs_para_la_lista(['4002448039440'], M), (['4002448126515'], 2, None))

    class _Roto:
        def table(self, _t):
            raise RuntimeError('relation "osma_packs_cola" does not exist')
    barrido.sb = _Roto()
    eq('E3 · si las tablas de los packs no existen (o fallan), el barrido sigue con su lista y lo dice', barrido.packs_para_la_lista(['1'], M),
       ([], 0, 'RuntimeError'))
finally:
    if _previo is None:
        sys.modules.pop('supabase', None)
    else:
        sys.modules['supabase'] = _previo
    for _k, _v in _ent.items():
        if _v is None:
            os.environ.pop(_k, None)
        else:
            os.environ[_k] = _v

# ── G · QUE LOS TESTS MUERDEN ───────────────────────────────────────────────────────────
_orig = op.n_del_pack


def _una_basta(pack, suelto):
    r = _orig(pack, suelto)
    return dict(r, estado=eo.PACK_SI) if r['estado'] == eo.PACK_DUDOSO else r


op.n_del_pack = _una_basta
try:
    _pm, _ = op.valorar(COLA, CAND, M, 6, NOV_KEEPA, None, AHORA)
    eq('G1 · con «una señal basta», el posible pack pierde su tope (la prueba C4 se pondría roja)',
       {p['c']['asin']: p['tope'] for p in _pm}['B0UNASENAL'], False)
finally:
    op.n_del_pack = _orig
_tv = op.TEXTO_SIN_VENTAS
_orig_v = op.ventas_de
op.ventas_de = lambda *a, **k: (99, 'inventado')
try:
    _pm, _ = op.valorar(COLA, CAND, M, 6, [], None, AHORA)
    eq('G2 · si las ventas se inventaran, «sin ventas» no pondría tope (la prueba C6 se pondría roja)', any(_tv in p['motivos'] for p in _pm), False)
finally:
    op.ventas_de = _orig_v

print('\n%d OK · %d fallos' % (N_OK[0], len(FALLOS)))
if FALLOS:
    print('FALLAN: ' + ' | '.join(FALLOS))
    sys.exit(1)
