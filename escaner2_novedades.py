# -*- coding: utf-8 -*-
"""ESCANER 2 · NOVEDADES DE FUNKO, TRAMO 2: VALORARLAS (encargo I, 29-sep-2026; encargo V, 30-sep-2026).

Lo llama el programa de la foto de HEO (escaner2_heo_disponibilidad.py) en cada pasada, DESPUES de aplicar la
pasada y de seleccionar sus novedades (tramo 1), como PASO APARTE: si algo de aqui falla, la pasada y la
seleccion siguen aplicadas, el fallo queda apuntado (nov_pasada.valoracion_*) y el run acaba en ROJO al final.
`valorar_pasada` NO LANZA NUNCA.

QUE HACE, EN ORDEN:
  0. EL INTERRUPTOR (nov_parametros.valorar). Apagado: ni una llamada a Keepa, ni una cuenta; solo se cierra la
     valoracion de la pasada como 'apagada' (con la foto de todas las novedades por estado, que cuenta la base).
  1. LA CUENTA de las novedades «lista» (el cartero de novedades, a y 20, ya les dejo precio y tarifa de Amazon de
     cada pais tras sus reintentos, o el pais marcado «no esta en Amazon», «nadie lo vende ahora» o «Amazon fallo»):
     🔑 (añadido 2 del encargo V) los paises en que Amazon fallo se piden a Keepa EN ESE MOMENTO (una peticion nueva
     por ASIN y pais, con `buybox`; nunca la cache): la logistica (pickAndPackFee) y el % de comision; la comision =
     ese % × el precio de Amazon de esa corrida (si Amazon no dio precio, el de Keepa: caja de compra, si no «nuevo»).
     🔴 EL ESCALON DE LOS 20 € (Identidad 10): Keepa no documenta con que precio calcula su tarifa; medido el 30-sep
     (40 de 40 pares ASIN-pais) coincide al centimo con la de Amazon al precio de ese momento, asi que el precio base
     es el precio actual de Keepa en la misma peticion. Si ese precio y el de venta quedan a lados distintos de 20 €,
     o Keepa no trae precio actual (no se sabe), ese pais sale «pendiente: cambia de escalon» y el cartero lo vuelve
     a pedir a Amazon en su corrida siguiente. Con los paises que tienen cifra se decide igual:
     con el MISMO codigo con el que el Escaneo PRO decide COMPRAR / VALORAR (escaner2_motor: `_decidir_con_asin`,
     que llama a `calcular_pais` -> `calc_rentabilidad` y `decision_de` del viejo). Las tarifas previstas (API o
     Keepa) se restan TAL CUAL, sin dividir entre 1,21 (Identidad 9). pa = el precio por unidad de HEO.
     🔴 (encargo V) Precio, tarifa y comision, de Amazon; y, solo si Amazon fallo, del respaldo de Keepa del momento,
     marcado. La cache de Keepa y el Escaneo PRO nunca dan tarifa ni comision.
  2. KEEPA, recorriendo `nov_cola` en su orden (bajadas por % primero; luego nuevo y vuelve; al final las subidas
     que quedaron dentro). Por cada novedad y pais (ES, IT, FR, DE), el dato de ventas sale de, por este orden:
       · la CACHE (nov_keepa): una respuesta de Keepa de menos de 15 dias -> «keepa_cache», 0 tokens;
       · el ESCANEO PRO de menos de 15 dias (un cruce 'lista' con los cuatro paises, y ese EAN con ficha): su
         dato de ventas -> «escaneo_pro», 0 tokens (si es mas reciente que la cache, manda el);
         🔑 (encargo V) UNA SOLA VENTANA, de 15 dias, para las dos (Fernando: «es un dato medio estable»): la base
         obliga a que keepa_horas_nuevo, keepa_dias_precio y escaneo_pro_dias sean la misma. De lo guardado solo se
         usan las VENTAS (caidas, puestos): ni la tarifa ni la comision.
       · KEEPA: una peticion por EAN y pais (`/product`, `code`, `stats=90`, sin historial ni caja de compra).
         🔴 Se lee el saldo ANTES de pedir y cada peticion exige saldo - reserva (20) >= el tope de una peticion
         (3): la reserva no se cruza y nunca se entra en negativo. Lo que no cabe espera a la pasada siguiente
         (espera_keepa). El saldo es compartido con el director viejo de HEO (misma llave, a y 30).
     Con el dato de los cuatro paises, las PUERTAS del Escaneo PRO (`escaner2_motor.decidir`, con la regla del
     viejo para elegir ficha si un EAN tiene varias): no esta en Amazon o sin dato -> SIN HISTORIAL; no se vende
     -> NO SE VENDE; varias fichas sin poder elegir -> Sin datos; se vende -> a esperar a Amazon (el cartero).
     🔴 (encargo T, 30-sep-2026) ANTES de las puertas, la ficha de un pais cuyo titulo no es de la marca (Funko: sin
     «Funko» ni «Pop» como PALABRA, encargo V) se aparta: ese pais cuenta como sin dato y el motivo dice «ficha
     dudosa: <titulo>».
     🔑 (encargo V) LAS NUESTRAS van por el mismo camino, con SU ficha: la cola trae sus ASIN (productos.asin, que la
     base guarda en nov_novedad.asins_nuestros); de la cache y del Escaneo PRO solo valen esas fichas, y a Keepa se
     le pregunta por ASIN, no por EAN. Nunca se adivina el ASIN de una nuestra por su EAN.
  3. EL EXCEL (encargo V): con las novedades valoradas en ESTA ejecucion, si al menos una sale COMPRAR, el Excel del
     escaner nuevo de HEO (escaner2_motor.excel_como_el_viejo, con la columna «Ventas») y detras una hoja
     «Novedades» con el origen de cada uno de sus 4 paises (Amazon · de Keepa (Amazon fallo) · pendiente: cambia de
     escalon · no esta en Amazon <pais> · nadie lo vende ahora en <pais>). Al bucket escaner2 (heo/novedades/<dia>/)
     y a nov_excel, que es lo que lee la biblioteca de escaneos de la v2. Sin COMPRAR, no se deja nada. Con mas de
     AVALANCHA novedades en la pasada, la hoja «Novedades» se abre la primera con «Avalancha: N cambios — lanza un
     Escaneo PRO» arriba.
  3 bis. EL TELEGRAM (añadido 3), como el escaner viejo: con Excel, «🟢 Novedades HEO Funko: N para COMPRAR» y una
     linea por COMPRAR (tope 20); con avalancha, aunque no haya COMPRAR; y en 🔴, diciendolo, si algun pais quedo
     pendiente o Keepa no llego a todo. Sin TELEGRAM_TOKEN/TELEGRAM_CHAT_ID en el paso, no se envia; si Telegram
     falla, la corrida sigue.
  4. EL CIERRE: nov_cerrar_valoracion con el flujo del paso (lo cuadra la base) y los tokens gastados.

LAS CUENTAS SUELTAS (encargo V): `solo_cuentas` hace el paso 1 (con el respaldo de Keepa de los paises en que Amazon
fallo), el Excel y el Telegram, sin ventas de Keepa y sin HEO. La llama escaner2_heo_novedades_cuentas.py (workflow
escaner2-heo-novedades-cuentas.yml, a y 40: despues de los reintentos del cartero de y 20).

🔑 LA TARIFA DE KEEPA (fbaFees.pickAndPackFee) viene en el mismo objeto de producto que las caidas y se guarda en la
   pelicula (nov_keepa), pero la de la CACHE no se usa nunca (nov_guardar_keepa rechaza una fila que la traiga): solo
   la de una peticion del momento, en el respaldo, marcada (tarifa_origen = keepa_respaldo).
🔑 LO QUE CUESTA KEEPA: 1 token por producto con stats y sin historia (medido: 124 peticiones = 124 tokens el 30-sep);
   `buybox` suma 2 por producto (documentacion de Keepa); `stats` no cuesta: stats.current trae el precio «nuevo»
   (indice 1) sin coste; la caja de compra (indice 18) solo con `buybox`. El respaldo cuesta 3 por ASIN y pais.
🔒 NI escaner_resultados NI el buzon `informes`: el Excel de novedades vive en el bucket escaner2 y en nov_excel.
🔒 LA LLAVE DE KEEPA no se imprime nunca: los mensajes de error se limpian de ella.
"""
import html
import io
import os
import re
import time
from datetime import datetime, timedelta, timezone

import en_mi_bd
import escaner2_motor as e2

PROVEEDOR = 'HEO'
PAISES = ('ES', 'IT', 'FR', 'DE')
# Los dominios de Keepa (keepa.DCODES): de 3, fr 4, it 8, es 9.
DOMINIO_KEEPA = {'ES': 9, 'IT': 8, 'FR': 4, 'DE': 3}
URL_KEEPA = 'https://api.keepa.com'
# La serie del puesto de ventas (SALES) en `stats.current` / `stats.avg90`, como el director (IDX_RANK = 3).
IDX_RANK = 3
LOTE_IN = 100
# El Excel de novedades (encargo V): el bucket del escaner 2 y su carpeta, que la biblioteca de la v2 firma por la fila
# de nov_excel (lib/escaner2/build.ts: esRutaDelBucket exige «heo/»).
BUCKET_EXCEL = 'escaner2'
# Mas de estas novedades en una pasada es una AVALANCHA (Fernando: «si hay mas de 100 cambios aviso para hacerlo Pro»).
AVALANCHA = 100
# Identidad 10: el escalon de la tarifa FBA en los cuatro paises. 20,00 € exactos pagan la alta.
ESCALON_EUR = 20.0
# Los indices de stats.current de Keepa (Product.CsvType): NEW 1, BUY_BOX_SHIPPING 18.
IDX_NUEVO, IDX_CAJA = 1, 18
CARPETA_EXCEL = 'heo/novedades'
# Lo que se lee de una novedad «lista» (la cuenta y la hoja «Novedades» del Excel).
COLUMNAS_NOVEDAD = ('id,producto_prov,ean_norm,nombre,motivo,precio_antes,precio_ahora,cambio_pct,es_chase,es_caja,uds_caja,nuestro,'
                    'decision_anterior')


def _ahora():
    return datetime.now(timezone.utc)


def _fecha(s):
    return datetime.fromisoformat(str(s).replace('Z', '+00:00').replace(' ', 'T')) if s else None


def _entero(x):
    return None if x is None else int(round(float(x)))


# ═══════════════════════════════════════════════════════════════════════════════
# 1 · KEEPA: el cliente (saldo y producto por codigo) y la ficha de cada producto
# ═══════════════════════════════════════════════════════════════════════════════
class KeepaFalla(RuntimeError):
    """Keepa no contesto (red, 5xx, 429, respuesta rara): lo que faltaba espera a la pasada siguiente."""


def _http_de_verdad(url, params, timeout):
    import requests  # tardio: el banco de pruebas no lo necesita
    r = requests.get(url, params=params, timeout=timeout)
    try:
        cuerpo = r.json()
    except ValueError:
        cuerpo = None
    return r.status_code, cuerpo


class Keepa:
    """La API de Keepa, lo justo: el SALDO (`/token`, 0 tokens) y los productos de unos codigos en un pais
    (`/product`). Lleva la cuenta de las peticiones de producto y de los tokens que Keepa dice haber gastado
    (`tokensConsumed`), y el saldo de la ultima respuesta (`tokensLeft`)."""

    def __init__(self, llave, http=None, dormir=time.sleep, timeout=60, intentos=2):
        self.llave = llave or ''
        self._http = http or _http_de_verdad
        self._dormir = dormir
        self.timeout, self.intentos = timeout, intentos
        self.peticiones = 0
        self.tokens = 0
        self.saldo = None
        self.con_tarifa = 0
        self.fichas = 0
        # 🔴 (2.ª auditoria) Keepa no contesta: se recuerda en la corrida, y nadie mas lo vuelve a intentar (cada intento
        #    son hasta 2 × 60 s de espera y el paso tiene un tope).
        self.caido = None

    def _limpio(self, texto):
        texto = str(texto)
        return texto.replace(self.llave, '***') if self.llave else texto

    def _get(self, ruta, params):
        if not self.llave:
            raise KeepaFalla('sin KEEPA_API_KEY en el entorno')
        ultimo = None
        for intento in range(self.intentos):
            try:
                estado, cuerpo = self._http(URL_KEEPA + ruta, dict(params, key=self.llave), self.timeout)
            except Exception as ex:
                ultimo = '%s: %s' % (type(ex).__name__, self._limpio(ex)[:200])
            else:
                if isinstance(cuerpo, dict) and cuerpo.get('tokensLeft') is not None:
                    self.saldo = int(cuerpo['tokensLeft'])
                if estado == 200 and isinstance(cuerpo, dict) and not cuerpo.get('error'):
                    return cuerpo
                ultimo = 'HTTP %s en %s%s' % (estado, ruta, (' · ' + self._limpio(cuerpo.get('error'))[:200])
                                              if isinstance(cuerpo, dict) and cuerpo.get('error') else '')
                # 🔴 Un 429 (sin tokens) o un 4xx no se reintenta: no va a cambiar en 5 segundos.
                if estado is not None and 400 <= estado < 500:
                    break
            if intento + 1 < self.intentos:
                self._dormir(5)
        raise KeepaFalla(ultimo or 'sin respuesta')

    def leer_saldo(self):
        self._get('/token', {})
        if self.saldo is None:
            raise KeepaFalla('la respuesta de /token no trae tokensLeft')
        return self.saldo

    def productos(self, pais, codigos, por='code', buybox=False):
        """Los productos de unos codigos (EAN/UPC, `por='code'`) o de unos ASIN (`por='asin'`: las nuestras y el
        respaldo). `buybox` (el respaldo): la caja de compra en stats.current, +2 tokens por producto."""
        if por not in ('code', 'asin'):
            raise ValueError('por: code o asin')
        params = {'domain': DOMINIO_KEEPA[pais], por: ','.join(codigos), 'stats': 90, 'history': 0}
        if buybox:
            params['buybox'] = 1
        cuerpo = self._get('/product', params)
        self.peticiones += 1
        self.tokens += int(cuerpo.get('tokensConsumed') or 0)
        productos = cuerpo.get('products') or []
        if not isinstance(productos, list):
            raise KeepaFalla('la respuesta de /product no trae la lista de productos')
        return productos


def _serie(lista):
    v = lista[IDX_RANK] if isinstance(lista, list) and len(lista) > IDX_RANK else None
    return int(v) if isinstance(v, (int, float)) and v > 0 else None


def ficha_de_keepa(p):
    """Un producto de Keepa -> la ficha que se guarda en nov_keepa y con la que se decide. Puro.
    🔴 Sin dato no es cero: caidas -1 (Keepa no lo sabe) -> None; la tarifa que no viene o viene a 0 -> None."""
    st = p.get('stats') or {}
    caidas = st.get('salesRankDrops30')
    fba = (p.get('fbaFees') or {}).get('pickAndPackFee')
    ref = p.get('referralFeePercentage')
    if ref is None:
        ref = p.get('referralFeePercent')
    return {'asin': p.get('asin'), 'titulo': p.get('title') or '',
            'rank': _serie(st.get('current')), 'rank_90d': _serie(st.get('avg90')),
            'caidas_30d': int(caidas) if isinstance(caidas, (int, float)) and caidas >= 0 else None,
            'fee_fba': round(fba / 100.0, 2) if isinstance(fba, (int, float)) and fba > 0 else None,
            'ref_pct': float(ref) if isinstance(ref, (int, float)) and ref >= 0 else None,
            'eans': sorted({str(x) for x in (p.get('eanList') or []) + (p.get('upcList') or []) if x})}


# 🔴 LA FICHA TIENE QUE SER DE LA MARCA (encargo T, 30-sep-2026). FK93061 (NFL Saquon Barkley): en IT, Keepa cruzo su
#    EAN con B08HH6GYRP, un casco de moto que lleva ese EAN en Amazon, y la novedad salio «NO SE VENDE» con las
#    caidas del casco. Antes de usar la ficha de un pais, su titulo tiene que llevar alguna de estas palabras (sin
#    distinguir mayusculas); si no, ese pais cuenta como SIN DATO y el motivo lo dice.
#    🔑 (encargo V) Como PALABRA ENTERA, no como texto dentro: «Popcorn» o «Lollipop» ya no pasan; «Pop!», «POP
#    Vinyl» o «Funko's» si. Medido el 30-sep-2026 sobre las fichas de nov_keepa (78) y del ultimo Escaneo PRO (4.877):
#    ninguna cambia de lado.
#    El Escaneo PRO no tiene una comprobacion asi que reutilizar: su «Coincide» (moloka_escaner_nube._coincide_titulo)
#    compara con el NOMBRE del proveedor y solo marca, y el cotejo del viejo (`cotejar`) solo elige entre dos o mas
#    fichas; una sola ficha pasa sin mirar su titulo.
PALABRAS_DE_LA_MARCA = ('funko', 'pop')


_RE_MARCA = re.compile(r'\b(?:%s)\b' % '|'.join(PALABRAS_DE_LA_MARCA), re.I)


def ficha_de_la_marca(titulo):
    """True si el titulo de la ficha es de un producto de la marca (Funko: lleva «Funko» o «Pop» como palabra). Puro."""
    return bool(_RE_MARCA.search(str(titulo or '')))


def apartar_fichas_dudosas(fichas_por_pais):
    """({pais: [fichas de la marca]}, {pais: 'ficha dudosa: <titulo>'}). Un pais que solo tenia fichas dudosas se queda
    SIN fichas: para las puertas del Escaneo PRO es «sin dato» ahi, como si Keepa no lo hubiera encontrado. Puro."""
    buenas, dudosas = {}, {}
    for pais, fichas in fichas_por_pais.items():
        buenas[pais] = [f for f in fichas if ficha_de_la_marca(f.get('titulo'))]
        malas = [f for f in fichas if not ficha_de_la_marca(f.get('titulo'))]
        if malas:
            dudosas[pais] = '; '.join('ficha dudosa: %s' % ((f.get('titulo') or '(sin título)')[:120]) for f in malas)
    return buenas, dudosas


def rec_de_ficha(f):
    """Una ficha (de Keepa o del Escaneo PRO) con la forma de una fila del CSV del Visualizador, que es la que leen
    las puertas del Escaneo PRO. SIN precio, SIN comision y SIN tarifa (encargo V): de lo guardado solo valen las
    VENTAS; el precio, la tarifa y la comision los pone Amazon (el cartero)."""
    return {'asin': f.get('asin'), 'titulo': f.get('titulo') or '', 'rank': f.get('rank'), 'rank90': f.get('rank_90d'),
            'buybox': None, 'es_fba': False, 'nuevo': None, 'compct': None, 'fba': None}


def rec_de_fila(fila):
    """Una fila de valoracion (nov_valoracion, o escaner2_resultado_pais del Escaneo PRO) -> la fila del CSV del
    Visualizador que `calcular_pais` lee: el precio va como caja de compra (BB-FBA / BB-FBM) o como «nuevo» (SIN BB),
    segun su canal, para que `precio_de_venta` lo devuelva igual. Puro."""
    precio, canal = fila.get('precio_venta'), fila.get('canal')
    return {'asin': fila.get('asin'), 'titulo': fila.get('titulo') or '', 'rank': fila.get('rank'),
            'rank90': fila.get('rank_90d'),
            'buybox': precio if canal in ('BB-FBA', 'BB-FBM') else None, 'es_fba': canal == 'BB-FBA',
            'nuevo': precio if canal == 'SIN BB' else None, 'compct': fila.get('ref_pct'), 'fba': fila.get('fee_fba')}


def cuenta_de_pais(pais, fila, pa, core, M):
    """LA CUENTA de un pais, con el codigo del Escaneo PRO (`escaner2_motor.calcular_pais`). Puro."""
    return e2.calcular_pais(pais, rec_de_fila(fila), pa, core, M)


def decision_de_resultado(r):
    """(decision, mejor pais) de lo que devuelven las puertas del Escaneo PRO, con el dominio de nov_novedad
    (comparable con escaner2_resultado_ean: f COMPRAR, e VALORAR, d NO COMPRAR, c NO SE VENDE)."""
    motivo = r.get('motivo')
    if r['puerta'] == 'f':
        return 'COMPRAR', (r.get('mejor') or {}).get('pais')
    if r['puerta'] == 'e':
        return 'VALORAR', (r.get('mejor') or {}).get('pais')
    if motivo == 'd_sin_margen':
        return 'NO COMPRAR', None
    if motivo == 'c_pocas_caidas':
        return 'NO SE VENDE', None
    if motivo in ('c_sin_dato', 'a_no_aparece', 'a_sin_asin'):
        return 'SIN HISTORIAL', None
    return 'Sin datos', None      # b_varias_fichas y d_sin_datos


# ═══════════════════════════════════════════════════════════════════════════════
# 2 · LA BASE: lecturas paginadas
# ═══════════════════════════════════════════════════════════════════════════════
def _todas(sb, tabla, columnas, orden, filtros=()):
    """Todas las filas, de 1.000 en 1.000. `filtros`: (op, columna, valor) con op en eq, in, gte."""
    salida, desde = [], 0
    while True:
        q = sb.table(tabla).select(columnas)
        for op, col, val in filtros:
            q = q.eq(col, val) if op == 'eq' else q.in_(col, list(val)) if op == 'in' else q.gte(col, val)
        pagina = q.order(orden).range(desde, desde + 999).execute().data or []
        salida.extend(pagina)
        if len(pagina) < 1000:
            return salida
        desde += 1000


def _en_trozos(sb, tabla, columnas, orden, col_in, valores, filtros=()):
    salida, valores = [], sorted(set(valores))
    for i in range(0, len(valores), LOTE_IN):
        salida += _todas(sb, tabla, columnas, orden, list(filtros) + [('in', col_in, valores[i:i + LOTE_IN])])
    return salida


def leer_parametros(sb):
    par = (sb.table('nov_parametros').select('*').eq('proveedor', PROVEEDOR).execute().data or [None])[0]
    if not par:
        raise RuntimeError('nov_parametros no tiene la fila de %s' % PROVEEDOR)
    return par


def leer_params_escaner2(sb):
    par = (sb.table('escaner2_parametros').select('*').eq('proveedor', PROVEEDOR).execute().data or [None])[0]
    if not par:
        raise RuntimeError('escaner2_parametros no tiene la fila de %s (umbral y paises)' % PROVEEDOR)
    return {'umbral': int(par['umbral_caidas_30d']), 'paises_filtro': list(par['paises_filtro']),
            'paises_calculo': list(par['paises_calculo'])}


def estados_de_heo(sb, productos):
    """{producto_prov: fila de disp_estado} (EAN de cruce y nombre) de los productos de las novedades."""
    filas = _en_trozos(sb, 'disp_estado', 'producto_prov,ean_core,ean_norm,nombre', 'producto_prov', 'producto_prov',
                       productos, [('eq', 'proveedor', PROVEEDOR)])
    return {f['producto_prov']: f for f in filas}


def fila_foto(nov, estado, M):
    """La fila de la foto con la que deciden las puertas del Escaneo PRO: las variantes del EAN (para cruzar con
    las fichas), el precio por unidad de HEO, el EAN de cruce (el IVA de la ficha) y el nombre (elegir ficha)."""
    core = (estado or {}).get('ean_core') or nov['ean_norm'].zfill(13)
    return {'variantes': sorted({M.norm(v) for v in M.variantes_ean(core)} - {''}),
            'codigos_keepa': e2.codigos_para_keepa(core, M), 'precio_unidad': nov['precio_ahora'],
            'ean_core': core, 'nombre': nov.get('nombre') or (estado or {}).get('nombre') or ''}


class _EleccionPerezosa:
    """La regla del viejo para elegir ficha (`escaner2_motor.EleccionViejo`), que solo carga su corpus (los nombres
    del catalogo disponible de HEO en la foto) la PRIMERA vez que un EAN tiene dos o mas fichas."""

    def __init__(self, sb):
        self._sb, self._e = sb, None

    def elegir(self, *a, **k):
        if self._e is None:
            filas = _todas(self._sb, 'disp_estado', 'ean_norm,es_chase,nombre,precio_catalogo', 'producto_prov',
                           [('eq', 'proveedor', PROVEEDOR), ('eq', 'disponible', True), ('eq', 'ausencias', 0)])
            mejor = {}
            for f in filas:
                if not f.get('ean_norm'):
                    continue
                k2 = (f['ean_norm'], bool(f.get('es_chase')))
                p = f.get('precio_catalogo')
                if k2 not in mejor or (p is not None and (mejor[k2][0] is None or p < mejor[k2][0])):
                    mejor[k2] = (p, f.get('nombre') or '')
            self._e = e2.cargar_eleccion_viejo([v[1] for v in mejor.values()])
        return self._e.elegir(*a, **k)


# ═══════════════════════════════════════════════════════════════════════════════
# 3 · EL DATO DE VENTAS REUTILIZABLE: la cache de Keepa y el Escaneo PRO
# ═══════════════════════════════════════════════════════════════════════════════
def cache_keepa(sb, eans, desde):
    """{(ean_norm, pais): la respuesta de Keepa mas reciente desde `desde`} (nov_keepa)."""
    filas = _en_trozos(sb, 'nov_keepa', 'ean_norm,pais,fichas,consultada_en', 'consultada_en', 'ean_norm', eans,
                       [('gte', 'consultada_en', desde.isoformat())])
    salida = {}
    for f in filas:
        k = (f['ean_norm'], f['pais'])
        if k not in salida or _fecha(f['consultada_en']) > _fecha(salida[k]['consultada_en']):
            salida[k] = f
    return salida


def escaneo_pro_reciente(sb, claves, desde):
    """{(ean_norm, es_chase): {'fecha', 'fichas': {pais: [ficha]}}} del ULTIMO Escaneo PRO (cruce 'lista') con los
    cuatro paises y de menos de `desde`, en el que el EAN salio con ficha (puertas c, d, e, f). La fecha es la del
    DATO (el CSV mas viejo del cruce), no la del cruce."""
    if not claves:
        return {}
    cruces = [c for c in (sb.table('escaner2_cruce').select('id,pasada_id,creado_en,fecha_datos,paises_usados')
                          .eq('estado', 'lista').order('creado_en', desc=True).limit(60).execute().data or [])
              if _fecha(c.get('fecha_datos') or c['creado_en']) >= desde and set(PAISES) <= set(c.get('paises_usados') or [])]
    if not cruces:
        return {}
    eans = {e for e, _ch in claves}
    fotos = _en_trozos(sb, 'escaner2_foto', 'id,pasada_id,ean_norm,es_chase', 'id', 'ean_norm', eans,
                       [('in', 'pasada_id', {c['pasada_id'] for c in cruces})])
    fotos = {f['id']: f for f in fotos if (f['ean_norm'], bool(f['es_chase'])) in claves}
    if not fotos:
        return {}
    por_cruce = {c['id']: c for c in cruces}
    resultados = [r for r in _en_trozos(sb, 'escaner2_resultado_ean', 'id,cruce_id,foto_id,puerta,asin', 'id', 'foto_id',
                                        list(fotos), [('in', 'cruce_id', list(por_cruce))])
                  if r['puerta'] in ('c', 'd', 'e', 'f') and r.get('asin')]
    elegido = {}
    for r in resultados:
        f = fotos[r['foto_id']]
        k = (f['ean_norm'], bool(f['es_chase']))
        fecha = _fecha(por_cruce[r['cruce_id']].get('fecha_datos') or por_cruce[r['cruce_id']]['creado_en'])
        if k not in elegido or fecha > elegido[k][0]:
            elegido[k] = (fecha, r)
    if not elegido:
        return {}
    paises = _en_trozos(sb, 'escaner2_resultado_pais', 'resultado_ean_id,pais,asin,titulo,caidas_30d,rank,rank_90d',
                        'resultado_ean_id', 'resultado_ean_id', [r['id'] for _f, r in elegido.values()])
    por_res = {}
    for p in paises:
        por_res.setdefault(p['resultado_ean_id'], []).append(p)
    salida = {}
    for k, (fecha, r) in elegido.items():
        fichas = {pais: [] for pais in PAISES}
        for p in por_res.get(r['id'], []):
            if p['pais'] in fichas:
                # Del Escaneo PRO, SOLO las ventas (encargo V): su tarifa y su comision no se leen.
                fichas[p['pais']].append({'asin': p['asin'], 'titulo': p.get('titulo') or '', 'rank': p.get('rank'),
                                          'rank_90d': p.get('rank_90d'), 'caidas_30d': p.get('caidas_30d')})
        salida[k] = {'fecha': fecha, 'fichas': fichas}
    return salida


def _cent(x):
    return round(x / 100.0, 2) if isinstance(x, (int, float)) and x > 0 else None


def respaldo_de_producto(p, precio_amazon, canal_amazon):
    """(añadido 2 del encargo V) De UN producto de Keepa pedido en ese momento (con `buybox`), el respaldo de un pais en
    que Amazon fallo. Puro. Devuelve {'estado': 'keepa', precio_venta, canal, ref_pct, fee_fba, keepa_precio_base} o
    {'estado': 'repedir', 'motivo', 'tipo'}:
      · la logistica: fbaFees.pickAndPackFee (en centimos: 0 o sin dato = Keepa no la tiene); el %: referralFeePercentage;
      · el PRECIO BASE de su tarifa: su precio actual en esa misma peticion, la caja de compra (stats.current[18]) o, si
        no hay, «nuevo» (stats.current[1]). Sin ninguno, no se sabe con que precio la calculo: se trata como cruce;
      · el precio de venta: el de Amazon de esa corrida si lo dio (la comision = el % × ese precio); si no, el de Keepa;
      · 🔴 si el precio base y el de venta quedan a lados distintos de 20 €, CAMBIA DE ESCALON: repedir."""
    if not p:
        return {'estado': 'repedir', 'motivo': 'Keepa no tiene este ASIN en ese país', 'tipo': 'keepa_sin_dato'}
    fee = _cent((p.get('fbaFees') or {}).get('pickAndPackFee'))
    ref = p.get('referralFeePercentage')
    if ref is None:
        ref = p.get('referralFeePercent')
    if fee is None or not isinstance(ref, (int, float)) or ref < 0:
        return {'estado': 'repedir', 'tipo': 'keepa_sin_dato',
                'motivo': 'Keepa no trae la tarifa o la comisión (logística %s, comisión %s)' % (fee, ref)}
    st = p.get('stats') or {}
    cur = st.get('current') if isinstance(st.get('current'), list) else []
    caja = _cent(cur[IDX_CAJA]) if len(cur) > IDX_CAJA else None
    nuevo = _cent(cur[IDX_NUEVO]) if len(cur) > IDX_NUEVO else None
    base = caja or nuevo
    if base is None:
        return {'estado': 'repedir', 'tipo': 'escalon',
                'motivo': 'cambia de escalón: no se sabe con qué precio calculó Keepa su tarifa (sin precio actual en Keepa)'}
    if precio_amazon:
        precio, canal = float(precio_amazon), canal_amazon
    else:
        precio = base
        canal = ('BB-FBA' if st.get('buyBoxIsFBA') else 'BB-FBM') if caja else 'SIN BB'
    if (base >= ESCALON_EUR) != (precio >= ESCALON_EUR):
        return {'estado': 'repedir', 'tipo': 'escalon',
                'motivo': ('cambia de escalón: Keepa calculó su tarifa a %.2f € y el precio es %.2f €' % (base, precio)).replace('.', ',')}
    return {'estado': 'keepa', 'precio_venta': precio, 'canal': canal, 'ref_pct': float(ref), 'fee_fba': fee,
            'keepa_precio_base': base}


def respaldo_keepa(keepa, pais, asin, precio_amazon, canal_amazon, reserva, tope):
    """El respaldo de UN pais: una peticion NUEVA a Keepa (por ASIN, con `buybox`; nunca la cache) y
    `respaldo_de_producto`. Sin llave, sin saldo por encima de la reserva o con Keepa caido: repedir. NUNCA LANZA."""
    if keepa is None or not keepa.llave:
        return {'estado': 'repedir', 'motivo': 'sin KEEPA_API_KEY: el respaldo de Keepa no se ha podido pedir', 'tipo': 'keepa_caido'}
    if keepa.caido:
        return {'estado': 'repedir', 'motivo': 'Keepa no contesta en esta corrida: %s' % keepa.caido[:200], 'tipo': 'keepa_caido'}
    try:
        if keepa.saldo is None:
            keepa.leer_saldo()
        if keepa.saldo - reserva < tope:
            return {'estado': 'repedir', 'tipo': 'reserva',
                    'motivo': 'Keepa sin saldo: quedaban %d tokens y la reserva es %d' % (keepa.saldo, reserva)}
        prods = keepa.productos(pais, [asin], por='asin', buybox=True)
    except KeepaFalla as ex:
        keepa.caido = str(ex)
        return {'estado': 'repedir', 'motivo': 'Keepa falló: %s' % str(ex)[:200], 'tipo': 'keepa_caido'}
    return respaldo_de_producto(next((x for x in prods if x.get('asin') == asin), None), precio_amazon, canal_amazon)


def con_respaldo(fila, respaldo):
    """La fila de valoracion con el respaldo puesto (para la cuenta): de Keepa, su precio, canal, % y logistica; repedir,
    sin tarifa (ese pais sale «Sin datos»). Sin respaldo, la fila tal cual. Puro."""
    if not respaldo:
        return fila
    if respaldo['estado'] == 'keepa':
        return dict(fila, precio_venta=respaldo['precio_venta'], canal=respaldo['canal'], ref_pct=respaldo['ref_pct'],
                    fee_fba=respaldo['fee_fba'], amazon_estado='keepa', amazon_motivo=None)
    return dict(fila, ref_pct=None, fee_fba=None, amazon_estado='repedir', amazon_motivo=respaldo['motivo'])


def solo_suyas(fichas, asins):
    """(encargo V) De unas fichas (de la cache o del Escaneo PRO), SOLO las de nuestra ficha (sus ASIN). Puro."""
    return [f for f in (fichas or []) if f.get('asin') in asins]


# ═══════════════════════════════════════════════════════════════════════════════
# 4 · EL PASO ENTERO
# ═══════════════════════════════════════════════════════════════════════════════
def _decidir_ventas(nov, foto, fichas_por_pais, params, M, eleccion):
    """Las puertas del Escaneo PRO con el dato de ventas de los cuatro paises (sin precio) -> (destino, decision, r).
    (AA, 1-oct-2026) El umbral es el de escaner2_parametros, como en el Escaneo PRO. La FICHA COMPARTIDA no se reparte
    aqui: lo guardado (nov_keepa.fichas, `ficha_de_keepa`; escaner2_resultado_pais) no trae ASIN padre ni recuento de
    variaciones, y sin ellos no se puede detectar sin pedir otra vez a Keepa."""
    cands = {p: [rec_de_ficha(f) for f in fichas_por_pais.get(p, [])] for p in PAISES}
    caidas = {p: {f['asin']: f.get('caidas_30d') for f in fichas_por_pais.get(p, []) if f.get('asin')} for p in PAISES}
    r = e2.decidir(foto, cands, caidas, params, M, eleccion)
    if r['puerta'] in ('d', 'e', 'f'):
        return 'espera_amazon', None, r
    return 'valorada', decision_de_resultado(r)[0], r


def _filas_de_ventas(r, origen, fecha):
    """Las filas del dato de ventas de cada pais donde esta la ficha elegida (lo que guarda nov_guardar_keepa)."""
    filas = []
    for pais, c in (r.get('paises') or {}).items():
        filas.append({'pais': pais, 'asin': c['asin'], 'titulo': c.get('titulo') or None,
                      'caidas_30d': _entero(c.get('caidas_30d')), 'rank': _entero(c.get('rank')),
                      'rank_90d': _entero(c.get('rank_90d')), 'vende_aqui': bool(c.get('vende_aqui')),
                      'ventas_origen': origen[pais], 'ventas_de': fecha[pais].isoformat()})
    # 🔴 Sin 'fee_fba' ni 'ref_pct' (encargo V): la base las rechaza (KEEPA_SIN_TARIFA).
    return filas


def cuentas(sb, M, params, imprimir=print, valoradas=None, keepa=None, reserva=20, tope=3, incompletos=None):
    """LA CUENTA de las novedades «lista» (tambien las nuestras, encargo V). Devuelve (hechas, fallos, listas, avisos);
    si se le da `valoradas` (una lista), le anade cada una que ha contado, para el Excel y el Telegram; y en
    `incompletos`, cada pais que quedo sin cifra (cambia de escalon, Keepa sin dato, sin saldo o caido).
    🔑 (añadido 2) Los paises «falta_amazon» se piden a Keepa en ese momento (`respaldo_keepa`, con `keepa`)."""
    listas = _todas(sb, 'nov_novedad', COLUMNAS_NOVEDAD, 'creada_en',
                    [('eq', 'proveedor', PROVEEDOR), ('eq', 'estado', 'lista')])
    if not listas:
        return 0, 0, 0, []
    productos = _todas(sb, 'productos', 'ean,asin,iva_pct,stock_moloka,es_chase', 'id', [('eq', 'activo', True)])
    if not productos:
        raise RuntimeError('productos devolvió 0 filas con activo=true: el IVA de la ficha no se puede leer')
    M.poner_catalogo_propio(productos)
    # (encargo E, 30-sep-2026) «En mi BD» del Excel: almacen de la ficha y FBA de la ultima foto de inventario_fba
    # (hasta hoy esta consulta no pedia `stock_moloka` y la columna decia «OK Alm:0 FBA:0» de TODAS las nuestras).
    M.poner_foto_fba(en_mi_bd.leer_foto_fba(sb, imprimir=lambda linea: imprimir(linea, flush=True)))
    estados = estados_de_heo(sb, [n['producto_prov'] for n in listas])
    filas = _en_trozos(sb, 'nov_valoracion', '*', 'pais', 'novedad_id', [n['id'] for n in listas])
    por_nov = {}
    for f in filas:
        por_nov.setdefault(f['novedad_id'], []).append(f)
    hechas, fallos, avisos = 0, 0, []
    for n in listas:
        try:
            suyas = por_nov.get(n['id']) or []
            asins = {f['asin'] for f in suyas}
            if not suyas or len(asins) != 1:
                raise RuntimeError('la novedad lista no tiene sus filas con UNA ficha (%d filas, %d ASIN)' % (len(suyas), len(asins)))
            asin = asins.pop()
            foto = fila_foto(n, estados.get(n['producto_prov']), M)
            # 🔑 EL RESPALDO de los paises en que Amazon fallo: Keepa del momento, o «repedir» con su motivo.
            respaldos = {}
            for f in suyas:
                if f.get('amazon_estado') == 'falta_amazon':
                    respaldos[f['pais']] = respaldo_keepa(keepa, f['pais'], asin, f.get('precio_venta'), f.get('canal'),
                                                          reserva, tope)
                    if respaldos[f['pais']]['estado'] == 'repedir' and incompletos is not None:
                        incompletos.append({'novedad': n['id'], 'pais': f['pais'], 'motivo': respaldos[f['pais']]['motivo'],
                                            'tipo': respaldos[f['pais']].pop('tipo', 'otro')})
                    respaldos[f['pais']].pop('tipo', None)
            suyas = [con_respaldo(f, respaldos.get(f['pais'])) for f in suyas]
            cands = {f['pais']: [rec_de_fila(f)] for f in suyas}
            caidas = {f['pais']: {asin: f.get('caidas_30d')} for f in suyas}
            base = {'asin': None, 'fichas': None, 'paises': {}, 'caidas': {}, 'mejor': None}
            r = e2._decidir_con_asin(foto, cands, caidas, params, M, asin, base)
            decision, mejor = decision_de_resultado(r)
            paises = [{'pais': p, 'iva': c['iva'], 'iva_origen': c['iva_origen'], 'almacen': c['almacen'],
                       'com_digitales': c['com_digitales'], 'isd_pct': c['isd_pct'], 'isd_incluye_fba': c['isd_incluye_fba'],
                       'pa': c['pa'], 'com_amazon': c['com_amazon'], 'beneficio': c['beneficio'], 'roi': c['roi'],
                       'margen': c['margen'], 'decision': c['decision'], 'vende_aqui': bool(c.get('vende_aqui'))}
                      for p, c in r['paises'].items()]
            for fila in paises:
                if fila['pais'] in respaldos:
                    fila['respaldo'] = respaldos[fila['pais']]
            sb.rpc('nov_guardar_cuenta', {'p_novedad': n['id'], 'p_decision': decision, 'p_mejor_pais': mejor,
                                          'p_motivo': '%s: %s' % (r['motivo'], r['detalle']), 'p_paises': paises}).execute()
            hechas += 1
            # 🔑 (2.ª auditoria) SIN DUPLICADOS: una valorada que vuelve a la cuenta (un pais «repedir» llego) solo sale
            #    otra vez en el Excel y en el Telegram si su decision CAMBIA.
            repetida = n.get('decision_anterior') is not None and n.get('decision_anterior') == decision
            if repetida:
                imprimir('    cuenta %s (%s): vuelve a la cuenta con la MISMA decisión (%s): no sale otra vez en el Excel'
                         % (n['ean_norm'], asin, decision), flush=True)
            if valoradas is not None and not repetida:
                valoradas.append({'nov': n, 'foto': foto, 'r': r, 'decision': decision, 'mejor': mejor, 'M': M,
                                  'amazon': {f['pais']: (f.get('amazon_estado'), f.get('amazon_motivo') or f.get('error_amazon'))
                                             for f in suyas},
                                  'motivo': '%s: %s' % (r['motivo'], r['detalle'])})
            imprimir('    cuenta %s (%s): %s%s' % (n['ean_norm'], asin, decision, (' en %s' % mejor) if mejor else ''), flush=True)
        except Exception as ex:
            fallos += 1
            avisos.append('la cuenta de %s falló: %s: %s' % (n['id'], type(ex).__name__, str(ex)[:300]))
    return hechas, fallos, len(listas), avisos


def valorar_pasada(sb, pasada, keepa_llave=None, http=None, dormir=time.sleep, ahora=_ahora, imprimir=print, run_id=None,
                   novedades_pasada=None, env=None, post=None):
    """EL PASO DE VALORACION de una pasada. NUNCA LANZA. Devuelve (ok, resumen): ok=False -> el run en rojo al final.
    `novedades_pasada`: las novedades que saco la seleccion (para la AVALANCHA). `env`/`post`: el Telegram (pruebas)."""
    avisos, datos = [], {'estado': 'hecha'}
    avalancha = novedades_pasada if (novedades_pasada or 0) > AVALANCHA else None
    try:
        par = leer_parametros(sb)
        if par.get('valorar') is not True:
            # 🔴 APAGADO: ni Keepa, ni cuentas. Solo se cierra (la base cuenta y cuadra la foto de las novedades).
            datos = {'estado': 'apagada'}
        else:
            valoradas, incompletos = [], []
            if avalancha:
                imprimir('>>> AVALANCHA: %d cambios en HEO Funko (más de %d) — lanza un Escaneo PRO. El automático sigue con '
                         'las que quepan por prioridad sin cruzar la reserva de Keepa.' % (avalancha, AVALANCHA), flush=True)
                _resumen_del_run('⚠️ **Avalancha: %d cambios en HEO Funko — lanza un Escaneo PRO**' % avalancha, env)
            _valorar(sb, pasada, par, datos, avisos, keepa_llave, http, dormir, ahora, imprimir, valoradas, incompletos)
            # 🔑 (encargo V) EL EXCEL, con lo valorado en esta pasada, si hay un COMPRAR. Un fallo aqui es un aviso (rojo
            #    al final), no tumba la valoracion.
            aviso = excel_de_novedades(sb, valoradas, 'pasada', pasada, run_id, ahora, imprimir, avalancha)
            if aviso:
                avisos.append(aviso)
            aviso_telegram(valoradas, avalancha, incompletos, imprimir, env, post)
    except Exception as ex:
        datos = {'estado': 'fallida'}
        avisos.append('la valoración no terminó: %s: %s' % (type(ex).__name__, str(ex)[:500]))
    if avisos:
        datos['motivo'] = ' · '.join(avisos)[:4000]
    if datos['estado'] == 'fallida' and not datos.get('motivo'):
        datos['motivo'] = 'la valoración no terminó'
    try:
        res = sb.rpc('nov_cerrar_valoracion', {'p_pasada': pasada, 'p_datos': datos}).execute().data
    except Exception as ex:
        imprimir('NOVEDADES_VALORACION_SIN_CERRAR: %s: %s' % (type(ex).__name__, str(ex)[:500]), flush=True)
        return False, {'estado': 'sin cerrar', 'avisos': avisos}
    if datos['estado'] == 'apagada':
        imprimir('>>> VALORACION DE NOVEDADES: APAGADA (nov_parametros.valorar): 0 llamadas a Keepa. %s' % (res,), flush=True)
    else:
        imprimir('>>> VALORACION DE NOVEDADES: %s' % (res,), flush=True)
    for a in avisos:
        imprimir('    AVISO: %s' % a, flush=True)
    ok = (res or {}).get('estado') == datos['estado'] and datos['estado'] in ('hecha', 'apagada') and not avisos
    return ok, res


def _valorar(sb, pasada, par, datos, avisos, keepa_llave, http, dormir, ahora, imprimir, valoradas=None, incompletos=None):
    """El paso con el interruptor encendido: la cuenta de las «lista» y Keepa para la cola. Deja en `datos` el flujo y
    los tokens, en `avisos` lo que haya fallado sin tumbar el paso (Keepa caído, una cuenta, una novedad) y en
    `valoradas` (encargo V) cada novedad valorada en esta pasada, para el Excel."""
    params = leer_params_escaner2(sb)
    M = e2.cargar_motor()
    reserva, tope = int(par['keepa_reserva']), int(par['keepa_tope_peticion'])
    # 🔑 El cliente de Keepa nace ANTES de las cuentas: el respaldo de los paises en que Amazon fallo lo usa, con la
    #    misma reserva y el mismo saldo que el paso de ventas de despues.
    keepa = Keepa(keepa_llave, http=http, dormir=dormir)
    t = {k: 0 for k in ('v_en_cola', 'v_keepa', 'v_keepa_cache', 'v_escaneo_pro', 'v_espera_saldo', 'v_espera_fallo',
                        'v_no_se_vende', 'v_sin_historial', 'v_varias_fichas', 'v_a_amazon')}

    # ── 1 · LA CUENTA de las «lista» (el cartero ya dejó precio y tarifa) ──
    try:
        hechas, fallos_cuenta, n_listas, av = cuentas(sb, M, params, imprimir, valoradas, keepa, reserva, tope, incompletos)
    except Exception as ex:
        hechas, fallos_cuenta, n_listas, av = 0, 0, 0, ['las cuentas no se pudieron hacer: %s: %s' % (type(ex).__name__, str(ex)[:300])]
        datos['estado'] = 'fallida'
    avisos += av
    datos.update(v_listas=n_listas, v_cuentas=hechas, v_cuenta_fallo=fallos_cuenta)

    # ── 2 · KEEPA, en el orden de la cola ──
    cola = _todas(sb, 'nov_cola', 'puesto,id,producto_prov,ean_norm,nombre,motivo,precio_ahora,precio_antes,cambio_pct,'
                  'es_chase,es_caja,uds_caja,estado,nuestro,asins_nuestros', 'puesto', [('eq', 'proveedor', PROVEEDOR)])
    t['v_en_cola'] = len(cola)
    keepa_caido = keepa.caido
    if cola:
        momento = ahora()
        estados = estados_de_heo(sb, [n['producto_prov'] for n in cola])
        horas_nuevo, dias_precio = int(par['keepa_horas_nuevo']), int(par['keepa_dias_precio'])
        ventana = lambda n: timedelta(hours=horas_nuevo) if n['motivo'] in ('nuevo', 'vuelve') else timedelta(days=dias_precio)
        cache = cache_keepa(sb, [n['ean_norm'] for n in cola], momento - max(timedelta(hours=horas_nuevo), timedelta(days=dias_precio)))
        pro = escaneo_pro_reciente(sb, {(n['ean_norm'], bool(n['es_chase'])) for n in cola},
                                   momento - timedelta(days=int(par['escaneo_pro_dias'])))
        eleccion = _EleccionPerezosa(sb)
        for n in cola:
            foto = fila_foto(n, estados.get(n['producto_prov']), M)
            desde = momento - ventana(n)
            # 🔑 (encargo V) Una nuestra: solo SUS fichas (productos.asin), y a Keepa se le pregunta por ASIN.
            suyos = list(n.get('asins_nuestros') or []) if n.get('nuestro') else None
            fichas, origen, fecha, faltan = {}, {}, {}, []
            for p in PAISES:
                c = cache.get((n['ean_norm'], p))
                utiles = solo_suyas(c['fichas'], suyos) if (c and suyos is not None) else (list(c['fichas']) if c else None)
                if c and _fecha(c['consultada_en']) >= desde and (suyos is None or utiles):
                    fichas[p], origen[p], fecha[p] = utiles, 'keepa_cache', _fecha(c['consultada_en'])
                else:
                    faltan.append(p)
            ep = pro.get((n['ean_norm'], bool(n['es_chase'])))
            if ep and suyos is not None:
                # Del Escaneo PRO de una nuestra, solo los paises donde esta SU ficha; los demas se preguntan por ASIN.
                ep = {'fecha': ep['fecha'], 'fichas': {p: solo_suyas(v, suyos) for p, v in ep['fichas'].items()}}
                if not any(ep['fichas'].values()):
                    ep = None
            # El Escaneo PRO, si la cache no está entera o si es más reciente que ella.
            if ep and (faltan or ep['fecha'] > min(fecha.values())):
                if suyos is None:
                    fichas, origen, fecha, faltan = dict(ep['fichas']), {p: 'escaneo_pro' for p in PAISES}, {p: ep['fecha'] for p in PAISES}, []
                else:
                    for p in PAISES:
                        if ep['fichas'].get(p):
                            fichas[p], origen[p], fecha[p] = ep['fichas'][p], 'escaneo_pro', ep['fecha']
                    faltan = [p for p in PAISES if p not in fichas]
            preguntadas = []
            espera = None
            for p in faltan:
                if keepa_caido:
                    espera = ('fallo', 'Keepa falló en esta pasada: %s' % keepa_caido)
                    break
                try:
                    if keepa.saldo is None:
                        keepa.leer_saldo()
                        datos['keepa_saldo_antes'] = keepa.saldo
                    # 🔴 LA RESERVA: cada petición exige saldo − reserva ≥ lo más que puede costar una.
                    if keepa.saldo - reserva < tope:
                        espera = ('saldo', 'sin saldo de Keepa: quedaban %d tokens, reserva %d y una petición puede costar %d'
                                  % (keepa.saldo, reserva, tope))
                        if incompletos is not None:
                            incompletos.append({'novedad': n['id'], 'pais': p, 'motivo': espera[1], 'tipo': 'reserva'})
                        break
                    prods = (keepa.productos(p, suyos, por='asin') if suyos is not None
                             else keepa.productos(p, foto['codigos_keepa']))
                except KeepaFalla as ex:
                    keepa_caido = keepa.caido = str(ex)
                    espera = ('fallo', 'Keepa falló: %s' % keepa_caido)
                    break
                fichas_p = []
                for prod in prods:
                    f = ficha_de_keepa(prod)
                    if f['asin'] and all(f['asin'] != x['asin'] for x in fichas_p) and (suyos is None or f['asin'] in suyos):
                        fichas_p.append(f)
                keepa.fichas += len(fichas_p)
                keepa.con_tarifa += sum(1 for f in fichas_p if f['fee_fba'] is not None)
                sb.table('nov_keepa').insert({'proveedor': PROVEEDOR, 'pasada_id': pasada, 'ean_norm': n['ean_norm'], 'pais': p,
                                              'codigos': suyos if suyos is not None else foto['codigos_keepa'],
                                              'fichas': fichas_p}).execute()
                fichas[p], origen[p], fecha[p] = fichas_p, 'keepa', ahora()
                preguntadas.append(p)
                if keepa.saldo is not None and keepa.saldo < reserva:
                    datos['keepa_reserva_cruzada'] = True
            if espera:
                if n['estado'] == 'pendiente':
                    try:
                        sb.rpc('nov_guardar_keepa', {'p_novedad': n['id'], 'p_destino': 'espera_keepa', 'p_decision': None,
                                                     'p_motivo': espera[1], 'p_filas': None}).execute()
                    except Exception as ex:
                        avisos.append('no se pudo dejar %s esperando Keepa: %s' % (n['id'], str(ex)[:200]))
                t['v_espera_saldo' if espera[0] == 'saldo' else 'v_espera_fallo'] += 1
                # (2.ª auditoria) Keepa caido en el paso de ventas: la corrida queda INCOMPLETA (🔴 en el Telegram).
                if espera[0] == 'fallo' and incompletos is not None:
                    incompletos.append({'novedad': n['id'], 'pais': None, 'motivo': espera[1], 'tipo': 'keepa_caido'})
                continue
            try:
                # 🔴 Encargo T: la ficha que no es de la marca no se usa; ese pais, sin dato, y el motivo lo dice. Una
                #    nuestra no pasa por aqui: su ficha es la nuestra (productos.asin), no una adivinada por EAN.
                fichas, dudosas = apartar_fichas_dudosas(fichas) if suyos is None else (fichas, {})
                destino, decision, r = _decidir_ventas(n, foto, fichas, params, M, eleccion)
                filas = _filas_de_ventas(r, origen, fecha) if r.get('asin') and r['puerta'] != 'b' else []
                fuentes = sorted(set(origen.values()))
                motivo = '%s: %s%s · dato de ventas: %s' % (
                    r['motivo'], r['detalle'],
                    ''.join(' · %s sin dato (%s)' % (p, dudosas[p]) for p in PAISES if p in dudosas), ', '.join(fuentes))
                sb.rpc('nov_guardar_keepa', {'p_novedad': n['id'], 'p_destino': destino, 'p_decision': decision,
                                             'p_motivo': motivo[:2000], 'p_filas': filas}).execute()
            except Exception as ex:
                avisos.append('la novedad %s no avanzó: %s: %s' % (n['id'], type(ex).__name__, str(ex)[:300]))
                t['v_espera_fallo'] += 1
                continue
            t['v_keepa' if preguntadas else 'v_escaneo_pro' if 'escaneo_pro' in fuentes else 'v_keepa_cache'] += 1
            if destino == 'valorada' and valoradas is not None:
                valoradas.append({'nov': n, 'foto': foto, 'r': r, 'decision': decision, 'mejor': None, 'M': M,
                                  'amazon': {}, 'motivo': motivo[:2000]})
            if destino == 'espera_amazon':
                t['v_a_amazon'] += 1
            elif r['puerta'] == 'b':
                t['v_varias_fichas'] += 1
            elif decision == 'NO SE VENDE':
                t['v_no_se_vende'] += 1
            else:
                t['v_sin_historial'] += 1
            imprimir('    %s %s (%s): %s%s' % (n['motivo'], n['ean_norm'], ', '.join(fuentes), destino,
                                               (' · %s' % decision) if decision else ''), flush=True)
    if keepa_caido:
        avisos.append('Keepa falló: %s' % keepa_caido)
    datos.update(t)
    datos.update(keepa_saldo_despues=keepa.saldo, keepa_peticiones=keepa.peticiones, keepa_tokens=keepa.tokens,
                 keepa_reserva=reserva, keepa_con_tarifa=keepa.con_tarifa, keepa_fichas=keepa.fichas)
    datos.setdefault('keepa_reserva_cruzada', False)
    if datos['keepa_reserva_cruzada']:
        avisos.append('🔴 la reserva de %d tokens se ha cruzado: quedan %s' % (reserva, keepa.saldo))


# ═══════════════════════════════════════════════════════════════════════════════
# 5 · EL EXCEL DE NOVEDADES Y LAS CUENTAS SUELTAS (encargo V, 30-sep-2026)
# ═══════════════════════════════════════════════════════════════════════════════
# Fernando: el Excel en la biblioteca de escaneos «tal cual hacia el escaner viejo», y «cuando salga al menos uno
# para comprar». Es EXACTAMENTE el del escaner nuevo de HEO: las seis hojas del viejo escritas por su codigo
# (escaner2_motor.excel_como_el_viejo, con la columna «Ventas» de la excepcion del 28-sep), con las novedades valoradas
# en ESA ejecucion; y detras, como el escaner 2 pone las suyas, una hoja «Novedades» con lo que el Excel del viejo no
# tiene donde poner: el motivo de la novedad, los dos precios y lo que dijo Amazon de cada pais («no se vende aqui»,
# «Amazon no ha dado la tarifa»…).
# 🔒 Donde vive: el bucket escaner2 (heo/novedades/<dia>/) y una fila en nov_excel, que es lo que lee la biblioteca de
#    la v2 (lib/escaner2/query.ts: cargarNovedadesBiblioteca). NUNCA escaner_resultados ni el buzon `informes`: esa
#    tabla la leen como «Excel del escaner viejo» el cruce del escaner 2, la espera de BEMS y el centinela.
XLSX = 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
# El origen de cada pais en el Excel (añadido 2 del encargo V, con las palabras de Fernando y Cowork).
QUE_DIJO_AMAZON = {'dato': 'Amazon', 'keepa': 'de Keepa (Amazon falló)', 'no_dado': 'Amazon no ha dado la tarifa',
                   'falta_amazon': 'Amazon falló (sin respaldo)'}
_ORDEN_DECISION = {'COMPRAR': 0, 'VALORAR': 1, 'NO COMPRAR': 2, 'NO SE VENDE': 3, 'SIN HISTORIAL': 4, 'Sin datos': 5}


def _madrid(d):
    from zoneinfo import ZoneInfo
    return d.astimezone(ZoneInfo('Europe/Madrid'))


def ruta_del_excel(momento):
    """heo/novedades/<AAAA-MM-DD>/Novedades_HEO_Funko_<AAAA-MM-DD_HHMM>.xlsx, en hora de Madrid. Puro."""
    m = _madrid(momento)
    return '%s/%s/Novedades_HEO_Funko_%s.xlsx' % (CARPETA_EXCEL, m.strftime('%Y-%m-%d'), m.strftime('%Y-%m-%d_%H%M'))


def datos_del_excel(valoradas):
    """(foto, resultados) con la forma que lee escaner2_motor.excel_como_el_viejo: una fila de foto por novedad (su EAN
    de cruce, su nombre y su precio por unidad de HEO) y el resultado de sus puertas. Puro."""
    foto, resultados = [], []
    for v in valoradas:
        n, f = v['nov'], v['foto']
        foto.append({'id': n['id'], 'ean_original': f['ean_core'], 'ean_core': f['ean_core'],
                     'nombre': f.get('nombre') or n.get('nombre') or '', 'marca': 'Funko',
                     'precio_unidad': f['precio_unidad'], 'aviso_caja': None})
        resultados.append(dict(v['r'], foto_id=n['id']))
    return foto, resultados


def texto_amazon(v, pais):
    """El origen de un pais, para la hoja «Novedades»: Amazon · de Keepa (Amazon fallo) · pendiente: <motivo> · no esta
    en Amazon <pais> · nadie lo vende ahora en <pais> · Amazon no ha dado la tarifa. Puro."""
    if not v['amazon']:
        return 'no se pidió (sin Amazon: %s)' % v['decision']
    if pais not in v['amazon']:
        return 'no está en Amazon %s' % pais
    estado, texto = v['amazon'][pais]
    if estado == 'no_se_vende':
        return 'no está en Amazon %s' % pais
    if estado == 'sin_ofertas':
        return 'nadie lo vende ahora en %s' % pais
    if estado == 'repedir':
        return 'pendiente: %s' % (texto or 'se vuelve a pedir a Amazon')
    return QUE_DIJO_AMAZON.get(estado, estado or '—') + (' (%s)' % texto if estado in ('no_dado', 'falta_amazon') and texto else '')


def hoja_novedades(wb, valoradas, avalancha=None):
    """La hoja «Novedades», detras de las del viejo: una fila por novedad, las COMPRAR primero. Con AVALANCHA, arriba
    «Avalancha: N cambios — lanza un Escaneo PRO» y es la hoja que se abre."""
    from openpyxl.styles import Font
    ws = wb.create_sheet('Novedades')
    if avalancha:
        ws.append(['Avalancha: %d cambios — lanza un Escaneo PRO' % avalancha])
        ws['A1'].font = Font(bold=True, color='C00000', size=14)
        wb.active = wb.sheetnames.index('Novedades')
    ws.append(['EAN', 'Nombre', 'Novedad', 'Precio antes (ud)', 'Precio ahora (ud)', 'Cambio %', 'Nuestra', 'Decisión',
               'Mejor país'] + ['Amazon ' + p for p in PAISES] + ['Por qué'])
    for c in ws[ws.max_row]:
        c.font = Font(bold=True)
    for v in sorted(valoradas, key=lambda x: (_ORDEN_DECISION.get(x['decision'], 9), str(x['nov'].get('nombre') or ''))):
        n = v['nov']
        ws.append([v['foto']['ean_core'], n.get('nombre') or v['foto'].get('nombre'), n.get('motivo'), n.get('precio_antes'),
                   n.get('precio_ahora'), n.get('cambio_pct'), 'sí' if n.get('nuestro') else '', v['decision'], v['mejor']]
                  + [texto_amazon(v, p) for p in PAISES] + [v['motivo']])
    for letra, ancho in {'A': 16, 'B': 48, 'C': 12, 'H': 14, 'J': 28, 'K': 28, 'L': 28, 'M': 28, 'N': 70}.items():
        ws.column_dimensions[letra].width = ancho
    ws.freeze_panes = 'A3' if avalancha else 'A2'
    return ws


def excel_de_novedades(sb, valoradas, origen, pasada, run_id=None, ahora=_ahora, imprimir=print, avalancha=None):
    """EL EXCEL de las novedades valoradas en esta ejecucion, si al menos una sale COMPRAR: al bucket escaner2 y a
    nov_excel. Devuelve None (hecho, o no tocaba) o el aviso de por que no se pudo. NUNCA LANZA."""
    comprar = [v for v in valoradas if v['decision'] == 'COMPRAR']
    if not comprar:
        imprimir('    EXCEL DE NOVEDADES: ninguna COMPRAR entre las %d valoradas en esta ejecución: no se deja Excel'
                 % len(valoradas), flush=True)
        return None
    try:
        foto, resultados = datos_del_excel(valoradas)
        wb = e2.excel_como_el_viejo(foto, resultados, [], comprar[0]['M'])
        hoja_novedades(wb, valoradas, avalancha)
        buf = io.BytesIO()
        wb.save(buf)
        ruta = ruta_del_excel(ahora())
        run_id = run_id or os.environ.get('GITHUB_RUN_ID')
        sb.storage.from_(BUCKET_EXCEL).upload(ruta, buf.getvalue(), {'content-type': XLSX, 'upsert': 'true'})
        sb.table('nov_excel').insert({
            'proveedor': PROVEEDOR, 'origen': origen, 'pasada_id': pasada if origen == 'pasada' else None,
            'run_id': int(run_id) if run_id else None, 'ruta_excel': ruta, 'n_novedades': len(valoradas),
            'n_comprar': len(comprar), 'n_valorar': sum(1 for v in valoradas if v['decision'] == 'VALORAR'),
            'novedades': [v['nov']['id'] for v in valoradas]}).execute()
        imprimir('>>> EXCEL DE NOVEDADES: %s (%d valoradas, %d COMPRAR) · en la biblioteca de escaneos'
                 % (ruta, len(valoradas), len(comprar)), flush=True)
        return None
    except Exception as ex:
        return 'el Excel de novedades no se pudo dejar: %s: %s' % (type(ex).__name__, str(ex)[:300])


# ── EL TELEGRAM (añadido 3 del encargo V): «tal cual hacia el escaner viejo» (moloka_escaner_nube.py, el aviso de los
#    directores): una linea por COMPRAR, tope 20; en 🔴 y diciendolo si la ejecucion quedo incompleta, y enviado aunque
#    no haya COMPRAR (un escaneo roto no se queda en silencio); la avalancha, aunque no haya COMPRAR.
_ORIGEN_TG = {'keepa': ' de Keepa'}


def mensaje_telegram(valoradas, avalancha=None, incompletos=()):
    """El texto del Telegram de una ejecucion, o None si no toca mandarlo (sin COMPRAR, sin avalancha y completa). Puro."""
    compras = []
    for v in valoradas:
        if v['decision'] != 'COMPRAR' or not v.get('mejor'):
            continue
        c = (v['r'].get('paises') or {}).get(v['mejor']) or {}
        compras.append((v, c.get('margen'), c.get('precio_venta'), v['mejor'], (v.get('amazon') or {}).get(v['mejor'], (None, None))[0]))
    incompletos = list(incompletos or [])
    if not compras and not avalancha and not incompletos:
        return None
    lineas = []
    if compras or incompletos:
        lineas.append('%s <b>Novedades HEO Funko</b>: %d para COMPRAR' % ('🔴' if incompletos else '🟢', len(compras)))
    if avalancha:
        lineas.append('⚠️ <b>Avalancha</b>: %d cambios en HEO Funko — lanza un Escaneo PRO' % avalancha)
    if incompletos:
        escalon = sum(1 for x in incompletos if x['tipo'] == 'escalon')
        reserva = sum(1 for x in incompletos if x['tipo'] == 'reserva')
        caido = sum(1 for x in incompletos if x['tipo'] == 'keepa_caido')
        otros = len(incompletos) - escalon - reserva - caido
        trozos = []
        if escalon:
            trozos.append('%d país(es) pendiente(s): cambia de escalón' % escalon)
        if otros:
            trozos.append('%d país(es) sin dato (Amazon falló y Keepa tampoco)' % otros)
        if reserva:
            trozos.append('%d consulta(s) que Keepa no hizo por la reserva' % reserva)
        if caido:
            trozos.append('%d consulta(s) sin hacer porque Keepa no contesta' % caido)
        lineas.append('⚠️ <b>NOVEDADES INCOMPLETAS</b>: ' + ' y '.join(trozos) + '. Los COMPRAR de abajo pueden no ser todos; '
                      'lo pendiente se vuelve a pedir en la corrida siguiente.')
    for v, margen, precio, pais, origen in compras[:20]:
        # 🔒 El nombre viene de HEO: se escapa para el HTML de Telegram (un «&» o un «<» rompen el mensaje entero).
        nombre = html.escape(str(v['nov'].get('nombre') or v['foto'].get('nombre') or '')[:45], quote=False)
        lineas.append('• %s — %s — %s (%s)%s' % (nombre, ('%.0f%%' % (margen * 100)) if margen is not None else 's/margen',
                                                 ('%.2f€' % precio) if precio else 's/precio', pais, _ORIGEN_TG.get(origen, '')))
    if len(compras) > 20:
        lineas.append('…y %d más (mira el Excel de la Biblioteca).' % (len(compras) - 20))
    return '\n'.join(lineas)


def novedades_a_valorar(seleccion):
    """Las novedades de una pasada que entran a valorar (2.ª auditoria, Fernando: la avalancha cuenta «solo lo que entra
    a valorar»): la cifra `novedades` de nov_seleccionar_pasada (que cuenta las subidas, medido por SQL el 30-sep) menos
    `subidas_fuera` (las que ya no se insertan). Puro."""
    s = seleccion or {}
    return max(int(s.get('novedades') or 0) - int(s.get('subidas_fuera') or 0), 0)


def _post_de_verdad(url, data, timeout):
    import requests  # tardio: el banco de pruebas no lo necesita
    return requests.post(url, data=data, timeout=timeout)


def enviar_telegram(texto, env=None, post=None, imprimir=print):
    """Manda el texto con TELEGRAM_TOKEN y TELEGRAM_CHAT_ID. Sin claves, no manda. Si Telegram falla, lo dice y sigue
    (como el viejo). 🔒 La llave no sale en ningun mensaje. Devuelve True si se envio."""
    env = os.environ if env is None else env
    token, chat = env.get('TELEGRAM_TOKEN'), env.get('TELEGRAM_CHAT_ID')
    if not token or not chat:
        imprimir('>>> Telegram: sin claves en este paso -> no se envía.', flush=True)
        return False
    try:
        r = (post or _post_de_verdad)('https://api.telegram.org/bot%s/sendMessage' % token,
                                      {'chat_id': chat, 'text': texto, 'parse_mode': 'HTML', 'disable_web_page_preview': 'true'}, 20)
        codigo = getattr(r, 'status_code', 200)
        if codigo != 200:
            raise RuntimeError('HTTP %s' % codigo)
        imprimir('>>> Telegram enviado.', flush=True)
        return True
    except Exception as ex:
        imprimir('AVISO Telegram (no se envió, la corrida sigue): %s: %s' % (type(ex).__name__, str(ex).replace(token, '***')[:200]),
                 flush=True)
        return False


def aviso_telegram(valoradas, avalancha, incompletos, imprimir=print, env=None, post=None):
    """El aviso de una ejecucion: compone y, si toca, manda. NUNCA LANZA."""
    try:
        texto = mensaje_telegram(valoradas, avalancha, incompletos)
    except Exception as ex:
        imprimir('AVISO Telegram (no se pudo componer): %s: %s' % (type(ex).__name__, str(ex)[:200]), flush=True)
        return False
    if texto is None:
        imprimir('>>> Telegram: 0 COMPRAR, sin avalancha y completo -> no se envía aviso.', flush=True)
        return False
    return enviar_telegram(texto, env, post, imprimir)


def _resumen_del_run(linea, env=None):
    """Una linea al resumen de la pagina del run (GITHUB_STEP_SUMMARY), si lo hay. NUNCA LANZA."""
    ruta = (os.environ if env is None else env).get('GITHUB_STEP_SUMMARY')
    if not ruta:
        return
    try:
        with io.open(ruta, 'a', encoding='utf-8') as fh:
            fh.write(linea + '\n')
    except OSError:
        pass


def solo_cuentas(sb, run_id=None, ahora=_ahora, imprimir=print, keepa_llave=None, http=None, dormir=time.sleep, env=None,
                 post=None):
    """LAS CUENTAS SUELTAS (workflow escaner2-heo-novedades-cuentas.yml, a y 40): la cuenta de las «lista» (con el
    respaldo de Keepa de los paises en que Amazon fallo), su Excel y su Telegram. Sin ventas de Keepa, sin HEO y sin
    tocar nov_pasada (la pasada en punto siguiente cuadra la foto de todas las novedades).
    NUNCA LANZA. Devuelve (ok, resumen): ok=False -> el run en rojo."""
    avisos, res = [], {'estado': 'hecha'}
    try:
        par = leer_parametros(sb)
        if par.get('valorar') is not True:
            imprimir('>>> CUENTAS DE NOVEDADES: el interruptor está APAGADO (nov_parametros.valorar): no se cuenta nada',
                     flush=True)
            return True, {'estado': 'apagada'}
        params = leer_params_escaner2(sb)
        M = e2.cargar_motor()
        valoradas, incompletos = [], []
        keepa = Keepa(keepa_llave, http=http, dormir=dormir)
        hechas, fallos, n_listas, av = cuentas(sb, M, params, imprimir, valoradas, keepa, int(par['keepa_reserva']),
                                               int(par['keepa_tope_peticion']), incompletos)
        avisos += av
        aviso = excel_de_novedades(sb, valoradas, 'cuentas', None, run_id, ahora, imprimir)
        if aviso:
            avisos.append(aviso)
        aviso_telegram(valoradas, None, incompletos, imprimir, env, post)
        res.update(keepa_peticiones=keepa.peticiones, keepa_tokens=keepa.tokens, incompletos=len(incompletos))
        res.update(listas=n_listas, cuentas=hechas, fallos=fallos,
                   comprar=sum(1 for v in valoradas if v['decision'] == 'COMPRAR'),
                   valorar=sum(1 for v in valoradas if v['decision'] == 'VALORAR'))
    except Exception as ex:
        res = {'estado': 'fallida'}
        avisos.append('las cuentas no terminaron: %s: %s' % (type(ex).__name__, str(ex)[:500]))
    imprimir('>>> CUENTAS DE NOVEDADES: %s' % (res,), flush=True)
    for a in avisos:
        imprimir('    AVISO: %s' % a, flush=True)
    return res['estado'] == 'hecha' and not avisos, dict(res, avisos=avisos)
