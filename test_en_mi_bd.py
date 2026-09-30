# -*- coding: utf-8 -*-
"""Banco de «En mi BD» con el stock de verdad (encargo E, 30-sep-2026): en_mi_bd.py y sus TRES usuarios.

EL FALLO QUE CIERRA. La columna «En mi BD» de los escaneres decia «OK Alm:X FBA:Y» con `productos.stock_fba`,
que esta congelado (lo escribia moloka_actualizar_nube.py; ultima corrida, 22-jul-2026). Caso real, medido en
produccion el 30-sep-2026: Claptrap (B0DP7C737P, EAN 889698882774) con `stock_fba` 0 desde el 16-ago y 18
disponibles en la foto de inventario_fba de ese dia. El Excel decia «OK Alm:0 FBA:0».

SIN RED, SIN SECRETOS Y SIN BASE. QUE PRUEBA:
  (A) EL TEXTO, con las filas REALES de la foto del 30-sep-2026 (ASIN y cantidades; los SKU van inventados):
      Claptrap → «FBA:18»; un ASIN que no esta en la foto → «FBA: —», nunca 0; sin ASIN o chase → «FBA: —»;
      dos vidas de SKU del mismo ASIN se suman; «+N en camino»; foto de mas de dos dias → «foto vieja del DD-MM»
      y sin cifra (con dos dias justos, la cifra); sin foto → «sin foto»; almacen sin dato → «Alm: —».
  (B) LA LECTURA (`leer_foto_fba`) con un doble de Supabase: solo la ULTIMA fecha, paginando de 1.000 en 1.000;
      la linea FOTO_FBA al log y el AVISO si es vieja o ilegible; y NUNCA LANZA.
  (C) LOS TRES USUARIOS, EJECUTADOS, con el caso de Claptrap:
      · el escaner viejo: el `en_bd_txt` de moloka_escaner_nube.py, sacado de su fichero;
      · el Pro: `escanear_pro` de moloka_escaner_pro.py de verdad, hasta su Excel;
      · el escaner 2: `Motor.en_bd_txt` y la Celda 9 (`escribir_celda9`), hasta su Excel.
      Y el mismo caso con el codigo de ANTES (copiado tal cual estaba en main el 30-sep) da «FBA:0»: el banco
      distingue.
  (D) POR ESTRUCTURA: ninguno de los programas que escriben la columna lee `stock_fba` de productos; el heredado
      del escaner 2 ya no trae `en_bd_txt`; la Celda 9 no la saca de el; y los tres llaman a `texto_en_mi_bd`.
"""
import ast
import copy
import io
import os
import sys
import tempfile
from contextlib import redirect_stdout
from datetime import date

AQUI = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, AQUI)
import en_mi_bd as E  # noqa: E402

fallos = []


def eq(nombre, obtenido, esperado):
    ok = obtenido == esperado
    if not ok:
        fallos.append(nombre)
    print(('OK ' if ok else 'XX ') + nombre + ('' if ok else '   got=%r exp=%r' % (obtenido, esperado)))


HOY = date(2026, 9, 30)
# Filas REALES de inventario_fba, foto del 30-sep-2026 (conector de lectura, 30-sep ~16:20): ASIN y cantidades
# tal cual; los SKU, inventados (el repo es publico).
FILAS_30SEP = [
    {'sku': 'SKU-CLAPTRAP', 'asin': 'B0DP7C737P', 'available': 18, 'fc_transfer': 0, 'inbound_shipped': 0,
     'inbound_receiving': 0, 'fecha_foto': '2026-09-30'},
    # dos vidas de SKU del mismo ASIN: 13 + 5
    {'sku': 'SKU-VIDA1', 'asin': 'B07GRRYFL1', 'available': 13, 'fc_transfer': 0, 'inbound_shipped': 0,
     'inbound_receiving': 0, 'fecha_foto': '2026-09-30'},
    {'sku': 'SKU-VIDA2', 'asin': 'B07GRRYFL1', 'available': 5, 'fc_transfer': 0, 'inbound_shipped': 0,
     'inbound_receiving': 0, 'fecha_foto': '2026-09-30'},
    # con envio en camino: 23 disponibles y 36 enviados
    {'sku': 'SKU-CAMINO', 'asin': 'B07T9C5CQR', 'available': 23, 'fc_transfer': 0, 'inbound_shipped': 36,
     'inbound_receiving': 0, 'fecha_foto': '2026-09-30'},
    # todo en camino, nada disponible: el 0 es un cero MEDIDO
    {'sku': 'SKU-SOLOCAMINO', 'asin': 'B0BZJYGWKT', 'available': 0, 'fc_transfer': 0, 'inbound_shipped': 73,
     'inbound_receiving': 0, 'fecha_foto': '2026-09-30'},
    # fc_transfer cuenta como disponible (criterio de la casa), y recibiendo cuenta como en camino
    {'sku': 'SKU-FCT', 'asin': 'B0FCTRANSF', 'available': 4, 'fc_transfer': 3, 'inbound_shipped': 0,
     'inbound_receiving': 2, 'fecha_foto': '2026-09-30'},
]
# La ficha de Claptrap, como esta en `productos` el 30-sep-2026.
CLAPTRAP = {'ean': '889698882774', 'asin': 'B0DP7C737P', 'iva_pct': 0.21, 'stock_moloka': 0, 'stock_fba': 0,
            'es_chase': False}
ESPERADO_CLAPTRAP = 'OK Alm:0 FBA:18 (foto 2026-09-30)'
FOTO = E.foto_de_filas(FILAS_30SEP, '2026-09-30', hoy=HOY)


def ficha(asin, alm=0, **extra):
    return dict({'ean': '8400000000017', 'asin': asin, 'stock_moloka': alm, 'stock_fba': 0}, **extra)


# ═══════════════════════════════════════════════════════════════════════════════
print('(A) el texto de la columna')
eq('(A) 🔴 Claptrap: stock_fba 0 en la ficha y 18 en la foto → «FBA:18», con la fecha de la foto',
   E.texto_en_mi_bd(CLAPTRAP, FOTO), ESPERADO_CLAPTRAP)
eq('(A) 🔴 ASIN que no está en la foto → «FBA: —» con la fecha (ausencia de dato no es cero)',
   E.texto_en_mi_bd(ficha('B0NOESTAAA', 2), FOTO), 'OK Alm:2 FBA: — (foto 2026-09-30)')
eq('(A) ficha sin ASIN → «FBA: —» (no se ha preguntado a la foto: sin fecha)',
   E.texto_en_mi_bd(ficha(None, 5), FOTO), 'OK Alm:5 FBA: —')
eq('(A) ASIN en blanco, igual que sin ASIN', E.texto_en_mi_bd(ficha('  ', 5), FOTO), 'OK Alm:5 FBA: —')
eq('(A) 🔴 un chase nunca cruza por ASIN, aunque alguien le hubiera clavado uno',
   E.texto_en_mi_bd(ficha('B0DP7C737P', 3, es_chase=True), FOTO), 'OK Alm:3 FBA: —')
eq('(A) dos vidas de SKU del mismo ASIN se SUMAN (13 + 5)',
   E.texto_en_mi_bd(ficha('B07GRRYFL1'), FOTO), 'OK Alm:0 FBA:18 (foto 2026-09-30)')
eq('(A) con envío en camino: «+N en camino»',
   E.texto_en_mi_bd(ficha('B07T9C5CQR', 1), FOTO), 'OK Alm:1 FBA:23 +36 en camino (foto 2026-09-30)')
eq('(A) todo en camino: el 0 de la foto es un cero MEDIDO, y se dice',
   E.texto_en_mi_bd(ficha('B0BZJYGWKT'), FOTO), 'OK Alm:0 FBA:0 +73 en camino (foto 2026-09-30)')
eq('(A) disponible = available + fc_transfer; en camino = shipped + receiving',
   E.texto_en_mi_bd(ficha('B0FCTRANSF'), FOTO), 'OK Alm:0 FBA:7 +2 en camino (foto 2026-09-30)')
eq('(A) almacén sin dato (la consulta no pidió stock_moloka) → «Alm: —», no 0',
   E.texto_en_mi_bd({'ean': '1', 'asin': 'B0DP7C737P'}, FOTO), 'OK Alm: — FBA:18 (foto 2026-09-30)')
eq('(A) almacén nulo → «Alm: —»', E.texto_en_mi_bd(ficha('B0DP7C737P', None), FOTO),
   'OK Alm: — FBA:18 (foto 2026-09-30)')
eq('(A) sin ficha, la columna calla', (E.texto_en_mi_bd(None, FOTO), E.texto_en_mi_bd({}, FOTO)), ('', ''))
_vieja = E.foto_de_filas(FILAS_30SEP, '2026-09-27', hoy=HOY)
eq('(A) 🔴 foto de hace 3 días (más de 2) → «foto vieja del DD-MM» y SIN cifra',
   E.texto_en_mi_bd(CLAPTRAP, _vieja), 'OK Alm:0 FBA: foto vieja del 27-09')
eq('(A) …y la foto dice su edad', (_vieja['edad'], _vieja['vieja']), (3, True))
_justa = E.foto_de_filas(FILAS_30SEP, '2026-09-28', hoy=HOY)
eq('(A) con 2 días justos todavía vale: la cifra', E.texto_en_mi_bd(CLAPTRAP, _justa),
   'OK Alm:0 FBA:18 (foto 2026-09-28)')
eq('(A) foto vieja y ficha sin ASIN: raya (la foto no pinta nada)', E.texto_en_mi_bd(ficha(None, 1), _vieja),
   'OK Alm:1 FBA: —')
eq('(A) 🔴 sin foto (no se leyó) → «sin foto», nunca 0',
   (E.texto_en_mi_bd(CLAPTRAP, None), E.texto_en_mi_bd(CLAPTRAP, E.foto_de_filas([], None, hoy=HOY))),
   ('OK Alm:0 FBA: sin foto', 'OK Alm:0 FBA: sin foto'))
eq('(A) la foto cuenta filas y ASIN distintos', (FOTO['filas'], len(FOTO['por_asin'])), (6, 5))


# ═══════════════════════════════════════════════════════════════════════════════
print('\n(B) la lectura de la foto, con un doble de Supabase')


class _Resp:
    def __init__(self, data):
        self.data = data


class _Q:
    def __init__(self, bd, tabla):
        self.bd, self.tabla, self.f, self.o, self.r, self.t = bd, tabla, [], None, None, None

    def select(self, _c):
        return self

    def eq(self, k, v):
        self.f.append((k, v))
        return self

    def order(self, c, desc=False):
        self.o = (c, desc)
        return self

    def range(self, a, b):
        self.r = (a, b)
        return self

    def limit(self, n):
        self.t = n
        return self

    def execute(self):
        self.bd['consultas'].append((self.tabla, tuple(self.f), self.r))
        if self.bd.get('lanza'):
            raise RuntimeError('conexion caida')
        filas = [x for x in self.bd['tablas'].get(self.tabla, []) if all(str(x.get(k)) == str(v) for k, v in self.f)]
        if self.o:
            filas = sorted(filas, key=lambda x: str(x.get(self.o[0])), reverse=self.o[1])
        if self.r:
            filas = filas[self.r[0]:self.r[1] + 1]
        if self.t is not None:
            filas = filas[:self.t]
        return _Resp(filas)


class _SB:
    def __init__(self, filas, lanza=False):
        self.bd = {'tablas': {'inventario_fba': filas}, 'consultas': [], 'lanza': lanza}

    def table(self, t):
        return _Q(self.bd, t)


def leer(sb, hoy=HOY):
    log = []
    foto = E.leer_foto_fba(sb, hoy=hoy, imprimir=log.append)
    return foto, log


# Una foto de otro dia debajo (la tabla es FOTO y solo guarda una; si algun dia guardara dos, manda la ultima).
_otra = [dict(f, fecha_foto='2026-09-29', available=999) for f in FILAS_30SEP]
_f, _log = leer(_SB(_otra + FILAS_30SEP))
eq('(B) 🔴 solo la ÚLTIMA fecha: Claptrap 18, no 999 + 18', E.texto_en_mi_bd(CLAPTRAP, _f), ESPERADO_CLAPTRAP)
eq('(B) la línea FOTO_FBA al log, y sin aviso', _log, ['FOTO_FBA: fecha=2026-09-30 | filas=6 | asins=5 | edad=0 días'])
_mucha = [{'sku': 'S%05d' % i, 'asin': 'B0DP7C737P' if i % 1000 == 7 else 'B%09d' % i, 'available': 1,
           'fc_transfer': 0, 'inbound_shipped': 0, 'inbound_receiving': 0, 'fecha_foto': '2026-09-30'}
          for i in range(2500)]
_sb = _SB(_mucha)
_f, _log = leer(_sb)
eq('(B) 🔴 pagina de 1.000 en 1.000: las 2.500 filas, y Claptrap suma sus 3 (una en cada página)',
   (_f['filas'], _f['por_asin']['B0DP7C737P']['disponible']), (2500, 3))
eq('(B) …con tres páginas pedidas', [c[2] for c in _sb.bd['consultas'] if c[2]], [(0, 999), (1000, 1999), (2000, 2999)])
_f, _log = leer(_SB(FILAS_30SEP), hoy=date(2026, 10, 4))
eq('(B) 🔴 foto vieja: el AVISO al log, con la fecha, y la columna lo dice',
   (len(_log), _log[1].startswith('AVISO FOTO_FBA VIEJA: la última foto de inventario_fba es del 2026-09-30 (4 días'),
    E.texto_en_mi_bd(CLAPTRAP, _f)), (2, True, 'OK Alm:0 FBA: foto vieja del 30-09'))
_f, _log = leer(_SB(FILAS_30SEP, lanza=True))
eq('(B) 🔴 si la lectura LANZA, no revienta: «sin foto» y el AVISO con el motivo',
   (E.texto_en_mi_bd(CLAPTRAP, _f), 'ILEGIBLE (RuntimeError: conexion caida)' in _log[1]), ('OK Alm:0 FBA: sin foto', True))
_f, _log = leer(_SB([]))
eq('(B) tabla vacía: «sin foto» y el AVISO', (E.texto_en_mi_bd(CLAPTRAP, _f), _log[0], 'ILEGIBLE' in _log[1]),
   ('OK Alm:0 FBA: sin foto', 'FOTO_FBA: fecha=NINGUNA | filas=0', True))


# ═══════════════════════════════════════════════════════════════════════════════
print('\n(C) los tres usuarios, ejecutados con Claptrap')
import escaner2_motor as e2  # noqa: E402

# · El escaner viejo: SU `en_bd_txt`, sacado de moloka_escaner_nube.py y ejecutado con sus piezas (`_sup`, `norm`).
VIEJO = os.path.join(AQUI, 'moloka_escaner_nube.py')
M_v = e2.cargar_motor(VIEJO)
M_v.poner_catalogo_propio([CLAPTRAP])
M_v.poner_foto_fba(FOTO)
_ns_v = e2.sacar_piezas(VIEJO, ('en_bd_txt',), (), base=M_v._ns)
eq('(C) 🔴 escáner viejo (moloka_escaner_nube.py · en_bd_txt): «FBA:18»', _ns_v['en_bd_txt']('889698882774'),
   ESPERADO_CLAPTRAP)
eq('(C) …y con el EAN con su cero delante, lo mismo (lo busca `_sup`)', _ns_v['en_bd_txt']('0889698882774'),
   ESPERADO_CLAPTRAP)
eq('(C) …y un EAN que no es nuestro, vacío', _ns_v['en_bd_txt']('8400000000999'), '')

# · El codigo de ANTES, copiado tal cual estaba en main el 30-sep-2026 (moloka_escaner_nube.py, lineas 1465-1468):
#   con la MISMA ficha da «FBA:0». Si el de ahora diera lo mismo, este banco no distinguiria nada.
_ANTES = '''
def en_bd_txt(core):
    s = _sup(core)
    if not s: return ''
    return f"OK Alm:{s.get('stock_moloka',0)} FBA:{s.get('stock_fba',0)}"
'''
_ns_a = dict(M_v._ns)
exec(compile(_ANTES, 'antes', 'exec'), _ns_a)
eq('(C) 🔴 el código de ANTES, con la misma ficha, decía «FBA:0»: el banco distingue',
   _ns_a['en_bd_txt']('889698882774'), 'OK Alm:0 FBA:0')

# · El escaner 2: el motor de verdad (con las piezas del heredado) y su Celda 9.
M = e2.cargar_motor()
M.poner_catalogo_propio([CLAPTRAP])
eq('(C) escáner 2 sin foto puesta: «sin foto», nunca un cero', M.en_bd_txt('889698882774'), 'OK Alm:0 FBA: sin foto')
M.poner_foto_fba(FOTO)
eq('(C) 🔴 escáner 2 (Motor.en_bd_txt): «FBA:18»', M.en_bd_txt('889698882774'), ESPERADO_CLAPTRAP)
_reg = {'nombre': 'POP! Vinyl - Claptrap', 'ean': '889698882774', 'asin': 'B0DP7C737P', 'marca': 'Funko',
        'core': '889698882774', '_pa_efectivo': 5.0, 'ambiguo': False, 'titulo_amz': 'Claptrap', 'coincide': 'SI',
        'coherencia_caja': None, 'url': '', 'volumen': None, '_margen_es': None,
        '_paises_calc': {'ES': {'rank_act': 1000, 'rank90': 1200, 'vendidos': 10, 'precio': None, 'canal': 'sin precio',
                                'n_of': 3, 'ref_pct': None, 'fee': None, 'iva': 0.21, 'decision': 'Sin datos',
                                'margen': None, 'caidas_30d': 9}}}
_datos = {'registros': [_reg], 'problematicos': [], 'no_encontrados': [], 'chase_sueltos': [], '_dups': [],
          'ambiguos': [], 'sin_rank': [], 'chase_pendientes': [], 'cotejo_info': {}, 'PROVEEDOR': 'HEO'}
_wb = e2.escribir_celda9(copy.deepcopy(_datos), M)
_ws = _wb['Análisis']
_cab = [c.value for c in _ws[1]]
eq('(C) 🔴 …y llega al Excel del escáner 2 (hoja «Análisis», columna «En mi BD»)',
   _ws.cell(row=2, column=_cab.index('En mi BD') + 1).value, ESPERADO_CLAPTRAP)

# · El Pro: `escanear_pro` de verdad, con el catalogo del proveedor y el CSV del Visualizador sustituidos (lo unico
#   que se sustituye: la lectura de los ficheros), hasta su Excel.
import moloka_escaner_pro as PRO  # noqa: E402
_fila = {'ean_in': '0889698882774', 'core': '889698882774', 'nombre': 'POP! Vinyl - Claptrap', 'marca': 'Funko',
         'pa': 5.0, 'es_chase': False, 'variantes': ['889698882774', '0889698882774'], 'volumen': None}
_rec = {'asin': 'B0DP7C737P', 'rank': 1000, 'rank90': 1200, 'nuevo': 15.0, 'buybox': 15.0, 'es_fba': True,
        'fba': 3.5, 'compct': 15.0, 'nof': 3, 'vendidos': 10, 'nvar': 0, 'titulo': 'Funko POP Claptrap'}
_orig = PRO.leer_proveedor, PRO.leer_csv_visualizador
PRO.leer_proveedor = lambda prov, marca, ruta: ([_fila], [])
PRO.leer_csv_visualizador = lambda ruta: {PRO.norm('889698882774'): [_rec]}
try:
    _sup_pro = {PRO.norm(CLAPTRAP['ean']): CLAPTRAP}
    with redirect_stdout(io.StringIO()):
        _res = PRO.escanear_pro('HEO', 'Funko', 'catalogo.xlsx', {'ES': 'es.csv'}, sup=_sup_pro, foto_fba=FOTO)
        _res_sin = PRO.escanear_pro('HEO', 'Funko', 'catalogo.xlsx', {'ES': 'es.csv'}, sup=_sup_pro)
        _xlsx = os.path.join(tempfile.mkdtemp(), 'pro.xlsx')
        PRO.escribir_excel(_res, _xlsx)
finally:
    PRO.leer_proveedor, PRO.leer_csv_visualizador = _orig
eq('(C) 🔴 Pro (escanear_pro): «FBA:18»', [r['en_bd'] for r in _res['registros']], [ESPERADO_CLAPTRAP])
eq('(C) Pro sin foto (lanzado a mano, sin base): «sin foto»', [r['en_bd'] for r in _res_sin['registros']],
   ['OK Alm:0 FBA: sin foto'])
from openpyxl import load_workbook  # noqa: E402
_wp = load_workbook(_xlsx)['Análisis']
_cp = [c.value for c in _wp[1]]
eq('(C) 🔴 …y llega al Excel del Pro', _wp.cell(row=2, column=_cp.index('En mi BD') + 1).value, ESPERADO_CLAPTRAP)


# ═══════════════════════════════════════════════════════════════════════════════
print('\n(D) por estructura: nadie escribe la columna con stock_fba, y los tres llaman a la misma función')


def _arbol(ruta):
    with io.open(os.path.join(AQUI, ruta), encoding='utf-8') as fh:
        return ast.parse(fh.read(), ruta)


def _textos_con(ruta, que):
    """Las cadenas del CODIGO (no comentarios ni docstrings) que contienen `que`: lo que se pediria a la base."""
    arbol = _arbol(ruta)
    docs = {id(n.value) for n in ast.walk(arbol) if isinstance(n, ast.Expr) and isinstance(n.value, ast.Constant)}
    return [n.value for n in ast.walk(arbol) if isinstance(n, ast.Constant) and isinstance(n.value, str)
            and id(n) not in docs and que in n.value]


def _llama(nodo, modulo, funcion):
    return any(isinstance(c, ast.Call) and isinstance(c.func, ast.Attribute) and c.func.attr == funcion
               and isinstance(c.func.value, ast.Name) and c.func.value.id == modulo for c in ast.walk(nodo))


def _def(arbol, nombre, clase=None):
    cuerpo = arbol.body
    if clase:
        cuerpo = [n for n in arbol.body if isinstance(n, ast.ClassDef) and n.name == clase][0].body
    defs = [n for n in cuerpo if isinstance(n, ast.FunctionDef) and n.name == nombre]
    return defs[0] if len(defs) == 1 else None


PROGRAMAS = ('moloka_escaner_nube.py', 'moloka_escaner_pro.py', 'moloka_escaner_pro_nube.py', 'escaner2_motor.py',
             'escaner2_heo_cruce.py', 'escaner2_novedades.py', 'en_mi_bd.py')
eq('(D) 🔴 ninguno de los programas de la columna pide ni lee `stock_fba` (en el código, no en los comentarios)',
   {p: _textos_con(p, 'stock_fba') for p in PROGRAMAS if _textos_con(p, 'stock_fba')}, {})
eq('(D) …y la comprobación muerde: con el código de antes lo habría cazado',
   bool([n for n in ast.walk(ast.parse(_ANTES)) if isinstance(n, ast.Constant) and 'stock_fba' in str(n.value)]), True)
eq('(D) 🔴 el escáner viejo: `en_bd_txt` llama a en_mi_bd.texto_en_mi_bd',
   _llama(_def(_arbol('moloka_escaner_nube.py'), 'en_bd_txt'), 'en_mi_bd', 'texto_en_mi_bd'), True)
eq('(D) …y la foto se lee a nivel superior, en la Celda 5 (FOTO_FBA = en_mi_bd.leer_foto_fba(sb))',
   [ast.unparse(n) for n in _arbol('moloka_escaner_nube.py').body
    if isinstance(n, ast.Assign) and [getattr(t, 'id', None) for t in n.targets] == ['FOTO_FBA']],
   ['FOTO_FBA = en_mi_bd.leer_foto_fba(sb)'])
eq('(D) 🔴 el Pro: `escanear_pro` llama a en_mi_bd.texto_en_mi_bd',
   _llama(_def(_arbol('moloka_escaner_pro.py'), 'escanear_pro'), 'en_mi_bd', 'texto_en_mi_bd'), True)
_main_pro = _def(_arbol('moloka_escaner_pro_nube.py'), 'main')
eq('(D) …y el Pro de la nube le pasa la foto que lee (foto_fba=… de en_mi_bd.leer_foto_fba)',
   (_llama(_main_pro, 'en_mi_bd', 'leer_foto_fba'),
    any(isinstance(c, ast.Call) and getattr(c.func, 'id', None) == 'escanear_pro' and
        any(k.arg == 'foto_fba' for k in c.keywords) for c in ast.walk(_main_pro))), (True, True))
eq('(D) 🔴 el escáner 2: Motor.en_bd_txt llama a en_mi_bd.texto_en_mi_bd',
   _llama(_def(_arbol('escaner2_motor.py'), 'en_bd_txt', clase='Motor'), 'en_mi_bd', 'texto_en_mi_bd'), True)
eq('(D) 🔴 el heredado del escáner 2 ya no trae el `en_bd_txt` del viejo',
   [n.name for n in _arbol('escaner2_heredado_nube.py').body if isinstance(n, ast.FunctionDef) and n.name == 'en_bd_txt'], [])
_celda9 = _def(_arbol('escaner2_motor.py'), 'escribir_celda9')
eq('(D) …y la Celda 9 no se la pide al heredado: se la inyecta el motor (en_bd_txt=M.en_bd_txt)',
   ('en_bd_txt' in [e.value for c in ast.walk(_celda9) if isinstance(c, ast.Tuple) for e in c.elts
                    if isinstance(e, ast.Constant)],
    any(isinstance(k, ast.keyword) and k.arg == 'en_bd_txt' for k in ast.walk(_celda9))), (False, True))
for _p in ('escaner2_heo_cruce.py', 'escaner2_novedades.py'):
    _a = _arbol(_p)
    eq('(D) %s pone la foto en el motor (M.poner_foto_fba(en_mi_bd.leer_foto_fba(sb, …)))' % _p,
       any(isinstance(c, ast.Call) and isinstance(c.func, ast.Attribute) and c.func.attr == 'poner_foto_fba'
           and _llama(c, 'en_mi_bd', 'leer_foto_fba') for c in ast.walk(_a)), True)
    eq('(D) …y pide stock_moloka y es_chase a productos',
       any('stock_moloka' in t and 'es_chase' in t and t.startswith('ean,') for t in _textos_con(_p, 'stock_moloka')), True)

print()
if fallos:
    print('ROJO: %d comprobaciones fallan: %s' % (len(fallos), ' | '.join(fallos)))
    sys.exit(1)
print('VERDE: «En mi BD» dice el almacén de la ficha y el FBA de la última foto de inventario_fba, con su fecha, '
      'en el escáner viejo, el Pro y el escáner 2.')
