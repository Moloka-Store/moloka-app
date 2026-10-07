# -*- coding: utf-8 -*-
"""Banco del ESCANEO PRO DE OCIOSTOCK (encargo OC3, 07-oct-2026): la parte pura y los DOS programas de verdad, con
dobles. Sin red, sin secretos, sin produccion. 🔴 Repo PUBLICO (Seguridad 21): la foto de OcioStock de aqui es
INVENTADA (ids 9000xx, precios redondos, EAN de mentira con su digito de control bueno, marcas reales sin precio).

PARTE A · LO PURO (escaner2_ociostock_pro.py), en este proceso:
  A1  el filtro de marca por modo: el del director (el nombre de la regla DENTRO del de la fila, como el viejo), todas
      (las sin marca tambien) y elegidas (coincidencia EXACTA, como HEO);
  A2  la foto: solo lo que se puede comprar ya (sin preventa), sin chase suelto (por la regla de la foto y por chase
      que no es caja), la caja «5 + 1» SI (es_chase y es_caja), una fila por EAN base con el menor PA (a igualdad, la
      suelta; luego el id menor), «caja de 6» marcada, PA = precio_pa y el precio de la ficha aparte; crudo = previas
      + foto, y si no, lo dice;
  A3  🔴 mutantes de lo puro: cada uno, A2 en rojo.

PARTE B · LOS PROGRAMAS (escaner2_ociostock_barrido.py y escaner2_ociostock_cruce.py), cada uno en su proceso, contra un
`supabase` en memoria:
  B1  [sin_llave] sin la llave de servicio no nace ningun cliente; [ofertas] el modo elegidas con ofertas, tampoco;
  B2  [marcas] foto → barrido → 'esperando_csv' con su foto, eans.txt y barrido.json → los CUATRO CSV (ES, IT, FR, DE,
      con la CABECERA REAL del Visualizador) → cruce → 'lista', cuadra en los cuatro, PA = precio_pa, y el Excel: las
      hojas del PRO de HEO en su orden (sin «Chase_manual»), «Análisis» con «Uds. escalón» y «Precio unidad» AL FINAL,
      «caja de 6» en «Coherencia caja», y la biblioteca puede leer su ruta;
  B3  🔴 el registro (repo PUBLICO) no lleva ni un EAN, precio, id de OcioStock ni nombre;
  B4  nada escribe fuera de escaner2_* ni del almacen ociostock/<pasada>/; disp_*, reglas_director y productos acaban
      como empezaron;
  B5  [todas] y [elegidas]: su foto; [ajena] una foto que no es la de su ultima pasada aplicada: ROJO y la pasada
      'fallida' sin foto.
"""
import base64
import csv
import io
import json
import os
import subprocess
import sys
import tempfile
import types

AQUI = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, AQUI)
fallos = []


def eq(nombre, obtenido, esperado):
    ok = obtenido == esperado
    if not ok:
        fallos.append(nombre)
    print(('OK ' if ok else 'XX ') + nombre + ('' if ok else '   got=%r exp=%r' % (obtenido, esperado)))


# ═══════════════════════════════════════════════════════════════════════════════
# EL DOBLE DE SUPABASE (el de test_escaner2_zentrada_pro.py, copiado: cada banco, el suyo)
# ═══════════════════════════════════════════════════════════════════════════════
class _Resp:
    def __init__(self, data, count=None):
        self.data, self.count = data, count


class _Consulta:
    def __init__(self, bd, tabla):
        self.bd, self.tabla = bd, tabla
        self.op, self.filtros, self.payload = 'select', [], None
        self.cuenta, self.orden, self.rango, self.tope = False, None, None, None

    def select(self, _cols='*', count=None):
        self.op, self.cuenta = 'select', (count == 'exact')
        return self

    def insert(self, filas):
        self.op, self.payload = 'insert', filas if isinstance(filas, list) else [filas]
        return self

    def update(self, cambios):
        self.op, self.payload = 'update', cambios
        return self

    def eq(self, k, v):
        self.filtros.append((k, v))
        return self

    def order(self, col, desc=False):
        self.orden = (col, desc)
        return self

    def range(self, a, b):
        self.rango = (a, b)
        return self

    def limit(self, n):
        self.tope = n
        return self

    def execute(self):
        self.bd.setdefault('ops', []).append([self.bd.get('programa'), self.op, self.tabla])
        filas = self.bd['tablas'].setdefault(self.tabla, [])
        if self.op == 'insert':
            for f in self.payload:
                filas.append(json.loads(json.dumps(f, default=str)))
            return _Resp(self.payload)
        sel = [f for f in filas if all(str(f.get(k)) == str(v) for k, v in self.filtros)]
        if self.op == 'update':
            for f in sel:
                f.update(json.loads(json.dumps(self.payload, default=str)))
            return _Resp(sel)
        if self.orden:
            sel = sorted(sel, key=lambda f: (f.get(self.orden[0]) is None, str(f.get(self.orden[0]))),
                         reverse=self.orden[1])
        n = len(sel)
        if self.rango:
            sel = sel[self.rango[0]:self.rango[1] + 1]
        if self.tope is not None:
            sel = sel[:self.tope]
        return _Resp(sel, n if self.cuenta else None)


class _Cubo:
    def __init__(self, bd, cubo):
        self.bd, self.cubo = bd, cubo

    def _objs(self):
        return self.bd['storage'].setdefault(self.cubo, {})

    def upload(self, ruta, datos, _opciones=None):
        self.bd.setdefault('subidas', []).append(ruta)
        self._objs()[ruta] = base64.b64encode(datos if isinstance(datos, bytes) else datos.encode()).decode()
        return {'Key': ruta}

    def download(self, ruta):
        if ruta not in self._objs():
            raise RuntimeError('404 Object not found: %s/%s' % (self.cubo, ruta))
        return base64.b64decode(self._objs()[ruta])

    def list(self, carpeta=None, opciones=None):
        op = dict({'limit': 100, 'offset': 0}, **(opciones or {}))
        pref = (carpeta or '').rstrip('/') + '/'
        nombres = sorted(r[len(pref):] for r in self._objs() if r.startswith(pref) and '/' not in r[len(pref):])
        trozo = nombres[op['offset']:op['offset'] + op['limit']]
        return [{'name': n, 'created_at': '2026-10-07T09:%02d:00Z' % i, 'metadata': {'size': 1}}
                for i, n in enumerate(trozo)]


class _Cliente:
    def __init__(self, bd):
        self.bd = bd
        self.storage = types.SimpleNamespace(from_=lambda cubo: _Cubo(bd, cubo))

    def table(self, nombre):
        return _Consulta(self.bd, nombre)


def hijo(ruta_estado, programa):
    with open(ruta_estado, encoding='utf-8') as fh:
        bd = json.load(fh)
    bd['programa'] = programa
    clientes = []

    def crear(url, llave):
        clientes.append(llave)
        return _Cliente(bd)

    sys.modules['supabase'] = types.ModuleType('supabase')
    sys.modules['supabase'].create_client = crear
    import runpy
    sys.argv = [programa]
    codigo = 0
    try:
        runpy.run_path(os.path.join(AQUI, programa), run_name='__main__')
    except SystemExit as ex:
        codigo = ex.code if isinstance(ex.code, int) else (0 if ex.code is None else 1)
    finally:
        with open(ruta_estado, 'w', encoding='utf-8') as fh:
            json.dump(bd, fh, default=str)
        print('CLIENTES_CREADOS=%d' % len(clientes))
    sys.exit(codigo)


if __name__ == '__main__' and len(sys.argv) >= 4 and sys.argv[1] == '--hijo':
    hijo(sys.argv[2], sys.argv[3])


def correr(ruta_estado, programa, env_extra):
    env = {k: v for k, v in os.environ.items()
           if k not in ('SUPABASE_SERVICE_KEY', 'PASADA', 'GITHUB_RUN_ID', 'MODO_BARRIDO', 'MARCAS_ELEGIDAS',
                        'OFERTAS_ELEGIDAS')}
    env.update(env_extra, PYTHONIOENCODING='utf-8', SUPABASE_URL='https://doble.invalid')
    p = subprocess.run([sys.executable, '-u', os.path.abspath(__file__), '--hijo', ruta_estado, programa],
                       capture_output=True, text=True, encoding='utf-8', errors='replace', env=env, cwd=AQUI)
    return p.returncode, (p.stdout or '') + (p.stderr or '')


import escaner2_motor as e2  # noqa: E402
import escaner2_ociostock_pro as eoc  # noqa: E402
import escaner2_heredado_pro as pro  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

M = e2.cargar_motor()


def ean13(cuerpo):
    return cuerpo + M._chk13(cuerpo)


# ═══════════════════════════════════════════════════════════════════════════════
# LA FOTO DE OCIOSTOCK, INVENTADA (la forma de disp_estado tras la pieza 1)
# ═══════════════════════════════════════════════════════════════════════════════
E1 = ean13('840000000101')    # FUNKO suelta 2,97 con escalon de 12; su caja «5 + 1», mas cara: gana la suelta
E2 = ean13('840000000102')    # FUNKO: la caja «5 + 1» a 1,98 y la suelta a 3,069: gana la CAJA («caja de 6»)
E3 = ean13('840000000103')    # FUNKO chase suelto (regla de la foto)
E4 = ean13('840000000104')    # FUNKO chase que no es caja y sin regla: fuera igual
E5 = ean13('840000000105')    # KIDS LICENSING
E6 = ean13('840000000106')    # CERDÁ: fuera de la lista del director
E7 = ean13('840000000107')    # POKEMON JUEGO DE CARTAS («Pokemon Juego De Cartas» en la regla)
E8 = ean13('840000000108')    # sin marca
E9 = ean13('840000000109')    # dos sueltas al mismo PA: el id menor (marca «FUNKO EUROPE»: entra por trozo)
E10 = ean13('840000000110')   # suelta y caja al mismo PA: la suelta
EG = ean13('840000000111')    # FUNKO agotado
EP = ean13('840000000112')    # FUNKO en preventa (disponible = true a proposito: la preventa manda)
DISP = '00000000-0000-4000-8000-0000000d1590'
OTRA = '00000000-0000-4000-8000-0000000d1591'


def fila(id_, ean, marca, unidad, escalon, uds, pa, caja=False, chase=False, disp=True, pre=False, regla=None,
         core=None, nombre=None, pasada=DISP):
    return {'proveedor': 'OCIOSTOCK', 'producto_prov': str(id_), 'ean_original': ean + (' C6' if caja else ''),
            'ean_core': core if core is not None else ean, 'ean_norm': (core or ean).lstrip('0'), 'marca': marca,
            'nombre': nombre or 'Figura inventada %d' % id_, 'categoria': 'Figuras', 'precio_unidad': unidad,
            'precio_escalon': escalon, 'uds_escalon': uds, 'precio_pa': pa, 'es_caja': caja,
            'uds_caja': 6 if caja else None, 'es_chase': chase, 'disponible': disp, 'preorder': pre, 'regla': regla,
            'pasada_id': pasada, 'ausencias': 0}


def foto_disp(pasada_ajena=False):
    return [
        fila(900001, E1, 'FUNKO', 4.0, 3.0, 12, 2.97),
        fila(900002, E1, 'FUNKO', 4.5, 4.5, 6, 4.455, caja=True, chase=True, nombre='Figura inventada 5 + 1 Chase'),
        fila(900003, E2, 'FUNKO', 2.5, 2.0, 6, 1.98, caja=True, chase=True),
        fila(900004, E2, 'FUNKO', 3.5, 3.1, 6, 3.069),
        fila(900005, E3, 'FUNKO', 9.0, 9.0, 1, 8.91, chase=True, regla='chase_suelto'),
        fila(900006, E4, 'FUNKO', 9.0, 9.0, 1, 8.91, chase=True),
        fila(900007, E5, 'KIDS LICENSING', 6.0, 5.0, 24, 4.95),
        fila(900008, E6, 'CERDÁ', 3.0, 3.0, 1, 2.97),
        fila(900009, EG, 'FUNKO', 4.0, 3.0, 12, 2.97, disp=False),
        fila(900010, EP, 'FUNKO', 4.0, 3.0, 12, 2.97, pre=True),
        fila(900011, '840000000', 'FUNKO', 4.0, 3.0, 12, 2.97, core='84000000011'),
        fila(900012, E7, 'POKEMON JUEGO DE CARTAS', 5.0, 4.0, 12, 3.96),
        fila(900013, E8, None, 2.0, 2.0, 1, 1.98),
        # «FUNKO EUROPE» (inventada): la regla del viejo la coge por trozo («funko» dentro).
        fila(900015, E9, 'FUNKO EUROPE', 4.0, 3.0, 12, 2.97),
        fila(900014, E9, 'FUNKO EUROPE', 4.0, 3.0, 12, 2.97),
        fila(900017, E10, 'FUNKO', 4.0, 3.0, 6, 2.97, caja=True, chase=True),
        fila(900016, E10, 'FUNKO', 4.0, 3.0, 12, 2.97, pasada=OTRA if pasada_ajena else DISP),
    ]


N_SIN_EAN = 3
N_CRUDO = len(foto_disp()) + N_SIN_EAN          # 20
DIRECTOR = ['Funko', 'Kids Licensing', 'Canenco', 'Jazwares', 'Pokemon Juego De Cartas']

# ═══════════════════════════════════════════════════════════════════════════════
# PARTE A · LO PURO
# ═══════════════════════════════════════════════════════════════════════════════
print('A1 · el filtro de marca, por modo')
_q, _i = eoc.filtro_marca('marcas', DIRECTOR)
eq('A1 · marcas: la regla del viejo (el nombre de la regla dentro del de la fila, sin mayúsculas)',
   [_q(m) for m in ('FUNKO', 'KIDS LICENSING', 'POKEMON JUEGO DE CARTAS', 'CERDÁ', None, 'FUNKO POP')],
   [True, True, True, False, False, True])
eq('A1 · …y guarda las marcas de la regla', _i['marcas_reales'], DIRECTOR)
_q, _i = eoc.filtro_marca('todas')
eq('A1 · todas: también las sin marca, y sin marcas que guardar', ([_q(m) for m in ('CERDÁ', None)], _i['marcas_reales']),
   ([True, True], None))
_q, _i = eoc.filtro_marca('elegidas', elegidas=['Kids Licensing'])
eq('A1 · elegidas: coincidencia EXACTA (sin mayúsculas), como HEO', [_q(m) for m in ('KIDS LICENSING', 'KIDS LICENSING X',
                                                                                    ' kids licensing ')],
   [True, False, True])
for _modo, _args in (('marcas', ([],)), ('elegidas', (None, [])), ('raro', ())):
    try:
        eoc.filtro_marca(_modo, *_args)
        _r = 'pasa'
    except eoc.FalloOcioStock:
        _r = 'para'
    eq('A1 · %s sin marcas (o desconocido): para' % _modo, _r, 'para')


def comprobar_foto(mod):
    """A2 entero contra un modulo (el de verdad o un mutante). Devuelve los nombres de lo que falla."""
    rotas = []

    def chk(nombre, obtenido, esperado):
        if obtenido != esperado:
            rotas.append((nombre, obtenido, esperado))
    quiere, _ = mod.filtro_marca('marcas', DIRECTOR)
    foto, apartados, cuentas, escalon = mod.construir_foto(foto_disp(), quiere, M, N_CRUDO, N_SIN_EAN, 'marcas')
    por = {f['ean_core']: f for f in foto}
    chk('foto', sorted(por), sorted([E1, E2, E5, E7, E9, E10]))
    chk('ids', {k: por[k]['producto_heo'] for k in por},
        {E1: '900001', E2: '900003', E5: '900007', E7: '900012', E9: '900014', E10: '900016'})
    chk('PA = precio_pa y el de la ficha aparte', (por[E1]['precio_unidad'], por[E1]['precio_catalogo']), (2.97, 4.0))
    chk('caja', [(por[k]['es_caja'], por[k]['es_chase'], por[k]['uds_caja'], por[k]['aviso_caja']) for k in (E1, E2)],
        [(False, False, None, None), (True, True, 6, 'caja de 6')])
    chk('códigos', (por[E2]['ean_original'], por[E2]['codigos_keepa'][0]), (E2 + ' C6', E2))
    previas = cuentas['previas']
    chk('previas', previas, {'chase_funko': 0, 'sin_gtin': 3, 'no_disponible': 2, 'marca_fuera': 2,
                             'estado_no_servible': 0, 'chase_suelto': 2, 'ean_forma_rara': 1, 'duplicado_proveedor': 4})
    chk('apartados', sorted((a['producto_heo'], a['motivo']) for a in apartados),
        sorted([('900008', 'marca_fuera'), ('900013', 'marca_fuera'), ('900005', 'chase_suelto'),
                ('900006', 'chase_suelto'), ('900011', 'ean_forma_rara'), ('900002', 'duplicado_proveedor'),
                ('900004', 'duplicado_proveedor'), ('900015', 'duplicado_proveedor'),
                ('900017', 'duplicado_proveedor')]))
    chk('cuadra', (cuentas['n_foto'], cuentas['n_previas'], cuentas['cuadra_previo']), (6, 14, True))
    chk('escalón', escalon.get('900001'), {'uds_escalon': 12, 'precio_unidad': 4.0, 'precio_escalon': 3.0,
                                           'precio_pa': 2.97})
    chk('escalón solo de la foto', sorted(escalon), sorted(por[k]['producto_heo'] for k in por))
    _f, _a, c2, _e = mod.construir_foto(foto_disp(), quiere, M, N_CRUDO + 1, N_SIN_EAN, 'marcas')
    chk('fichero que no casa: no cuadra y lo dice', (c2['cuadra_previo'], 'el fichero tenía 21 filas' in
                                                     (c2['motivo_previo'] or '')), (False, True))
    return rotas


print('A2 · la foto de la pasada')
_rotas = comprobar_foto(eoc)
for _n, _o, _e in _rotas:
    eq('A2 · ' + _n, _o, _e)
eq('A2 · las once comprobaciones de la foto, en verde', len(_rotas), 0)
quiere_t, _ = eoc.filtro_marca('todas')
_ft, _at, _ct, _et = eoc.construir_foto(foto_disp(), quiere_t, M, N_CRUDO, N_SIN_EAN, 'todas')
eq('A2 · todas: entran CERDÁ y la sin marca (8), y cuadra', (len(_ft), _ct['cuadra_previo'], _ct['previas']['marca_fuera']),
   (8, True, 0))
_qe, _ = eoc.filtro_marca('elegidas', elegidas=['KIDS LICENSING'])
_fe, _ae, _ce, _ee = eoc.construir_foto(foto_disp(), _qe, M, N_CRUDO, N_SIN_EAN, 'elegidas')
eq('A2 · elegidas: una, y la marca fuera dice «no elegida»', (
    [f['ean_core'] for f in _fe], sorted({a['detalle'].split("'")[0] for a in _ae if a['motivo'] == 'marca_fuera'})),
   ([E5], ['Marca ']))
eq('A2 · …y su puerta previa se llama «Marca no elegida»', eoc.nombre_previa('marca_fuera', 'elegidas'), 'Marca no elegida')
eq('A2 · las rutas que admiten los checks de la v2',
   (eoc.ruta_lista(DISP), eoc.ruta_excel(DISP, OTRA, '20261007_1530')),
   ('ociostock/%s/eans.txt' % DISP, 'ociostock/%s/%s/Escaner2_OCIOSTOCK_20261007_1530.xlsx' % (DISP, OTRA)))

print('A3 · 🔴 mutantes de lo puro: cada uno pone A2 en rojo')
with io.open(os.path.join(AQUI, 'escaner2_ociostock_pro.py'), encoding='utf-8') as _fh:
    FUENTE = _fh.read()
MUTANTES = [
    ('el PA es el precio de la ficha', "'precio_unidad': _num(f.get('precio_pa'))", "'precio_unidad': _num(f.get('precio_unidad'))"),
    ('la preventa entra', "if not f.get('disponible') or f.get('preorder'):", "if not f.get('disponible'):"),
    ('el chase que no es caja entra', "or (f.get('es_chase') and not f.get('es_caja'))", ''),
    ('fuera toda caja con chase', "or (f.get('es_chase') and not f.get('es_caja'))", "or f.get('es_chase')"),
    ('gana la más cara', "return clave(a) < clave(b)", "return clave(a) > clave(b)"),
    ('a igualdad, la caja', "bool(f.get('es_caja')), _clave_id", "not f.get('es_caja'), _clave_id"),
    ('sin «caja de 6»', "'aviso_caja': TEXTO_CAJA if caja else None", "'aviso_caja': None"),
    ('sin las sin EAN en el cuadre', "if cuentas['n_filas'] + previas['sin_gtin'] != cuentas['n_crudo']:",
     "if cuentas['n_filas'] != cuentas['n_crudo']:"),
    ('la forma rara entra', "if regla or not core.isdigit() or len(core) not in (12, 13):", "if regla:"),
    ('el director por igualdad', "any(b in str(m or '').lower() for b in bajas)", "str(m or '').lower() in bajas"),
]
for _nombre, _antes, _despues in MUTANTES:
    if FUENTE.count(_antes) != 1:
        eq('A3 · el ancla del mutante «%s» es única' % _nombre, FUENTE.count(_antes), 1)
        continue
    _mod = types.ModuleType('mutante')
    exec(compile(FUENTE.replace(_antes, _despues), 'mutante', 'exec'), _mod.__dict__)
    eq('A3 · mutante «%s»: A2 en rojo' % _nombre, len(comprobar_foto(_mod)) > 0, True)

# ═══════════════════════════════════════════════════════════════════════════════
# PARTE B · LOS DOS PROGRAMAS, DE PUNTA A PUNTA
# ═══════════════════════════════════════════════════════════════════════════════
with io.open(os.path.join(AQUI, 'test_escaner2_osma_cabecera.json'), encoding='utf-8') as fh:
    CABECERA = json.load(fh)['cabecera']
COL_PAIS, COL_CAIDAS = e2.columnas_keepa()
COL_PADRE, COL_NVAR, COL_RANK = e2.columnas_ficha_compartida()
C = pro.CSV_COLS


def csv_real(pais, filas):
    """Un export del Visualizador con la CABECERA REAL (561 columnas, la del banco de OSMA) y solo lo que se lee."""
    buf = io.StringIO()
    w = csv.writer(buf)
    w.writerow(CABECERA)
    ix = {h: i for i, h in enumerate(CABECERA)}
    for asin, eans, caidas, precio, titulo in filas:
        fila_ = [''] * len(CABECERA)
        for h, v in ((C['asin'], asin), (COL_PAIS, pais), ('Título', titulo), (C['ean'], eans),
                     (C['rank'], '9000'), (C['rank90'], '11000'), (COL_CAIDAS, caidas), (C['buybox'], precio),
                     (C['es_fba'], 'yes'), (C['nuevo'], precio), (C['fba'], '3.10'), (C['compct'], '15.01 %'),
                     (C['nvar'], ''), (COL_PADRE, '')):
            fila_[ix[h]] = v
        w.writerow(fila_)
    return ('﻿' + buf.getvalue()).encode('utf-8')


ASIN1, ASIN2, ASIN5, ASIN9 = 'B0OCIO0001', 'B0OCIO0002', 'B0OCIO0005', 'B0OCIO0009'
PRODUCTOS = [{'id': 1, 'ean': E1, 'asin': ASIN1, 'nombre': 'Figura de Moloka', 'activo': True, 'es_chase': False,
              'iva_pct': 0.21, 'stock_moloka': 3}]
SEMBRADAS = ('productos', 'inventario_fba', 'disp_pasada', 'disp_estado', 'reglas_director')


def estado_inicial(escena):
    return {'tablas': {
        'escaner2_parametros': [{'proveedor': 'HEO', 'umbral_caidas_30d': 6, 'paises_filtro': ['ES', 'IT', 'FR', 'DE'],
                                 'paises_calculo': ['ES', 'IT', 'FR', 'DE']},
                                # La fila de OcioStock que deja la migracion de la v2: el corte y los cuatro de HEO.
                                {'proveedor': 'OCIOSTOCK', 'umbral_caidas_30d': 6,
                                 'paises_filtro': ['ES', 'IT', 'FR', 'DE'], 'paises_calculo': ['ES', 'IT', 'FR', 'DE']}],
        'escaner2_pasada': [],
        'disp_pasada': [{'id': OTRA, 'proveedor': 'OCIOSTOCK', 'estado': 'aplicada', 'creada_en': '2026-10-06T07:00:00Z',
                         'n_leidas': 1, 'n_crudo': 2, 'n_sin_gtin': 1},
                        {'id': DISP, 'proveedor': 'OCIOSTOCK', 'estado': 'aplicada', 'creada_en': '2026-10-07T07:00:00Z',
                         'terminada_en': '2026-10-07T07:01:00Z', 'n_leidas': len(foto_disp()), 'n_crudo': N_CRUDO,
                         'n_sin_gtin': N_SIN_EAN, 'fichero_md5': 'f' * 32},
                        {'id': 'rechazada-al-dia', 'proveedor': 'OCIOSTOCK', 'estado': 'rechazada',
                         'creada_en': '2026-10-07T11:00:00Z'}],
        'disp_estado': foto_disp(pasada_ajena=(escena == 'ajena')),
        'reglas_director': [{'proveedor': 'OCIOSTOCK', 'activo': True, 'marcas': DIRECTOR},
                            {'proveedor': 'HEO', 'activo': True, 'marcas': ['Otra']}],
        'productos': json.loads(json.dumps(PRODUCTOS)),
        'inventario_fba': [],
    }, 'storage': {'escaner2': {}}}


def caso(escena, env=None):
    tmp = tempfile.mkdtemp(prefix='e2oc_')
    ruta = os.path.join(tmp, 'estado.json')
    inicial = estado_inicial(escena)
    with open(ruta, 'w', encoding='utf-8') as fh:
        json.dump(inicial, fh)
    cod, log = correr(ruta, 'escaner2_ociostock_barrido.py', dict({'SUPABASE_SERVICE_KEY': 'svc-de-mentira',
                                                                   'GITHUB_RUN_ID': '717171'}, **(env or {})))
    bd = json.load(open(ruta, encoding='utf-8'))
    return cod, log, bd, ruta, inicial


print('B1 · sin llave, o elegidas con ofertas: no se corre')
_tmp = tempfile.mkdtemp(prefix='e2oc_')
_ruta = os.path.join(_tmp, 'estado.json')
json.dump(estado_inicial('marcas'), open(_ruta, 'w', encoding='utf-8'))
for _prog in ('escaner2_ociostock_barrido.py', 'escaner2_ociostock_cruce.py'):
    _cod, _log = correr(_ruta, _prog, {'PASADA': DISP})
    eq('B1 · %s: ROJO, la línea exacta y NINGÚN cliente' % _prog,
       (_cod, 'ESCANER2_NO_EJECUTADO: sin llave de servicio' in _log, 'CLIENTES_CREADOS=0' in _log), (1, True, True))
_cod, _log = correr(_ruta, 'escaner2_ociostock_barrido.py', {'SUPABASE_SERVICE_KEY': 'x', 'MODO_BARRIDO': 'elegidas',
                                                             'MARCAS_ELEGIDAS': '["FUNKO"]', 'OFERTAS_ELEGIDAS': 'true'})
eq('B1 · elegidas con ofertas: ROJO, lo dice y NINGÚN cliente',
   (_cod, 'OcioStock no tiene ofertas' in _log, 'CLIENTES_CREADOS=0' in _log), (1, True, True))
_cod, _log = correr(_ruta, 'escaner2_ociostock_barrido.py', {'SUPABASE_SERVICE_KEY': 'x', 'MODO_BARRIDO': 'elegidas',
                                                             'MARCAS_ELEGIDAS': '["FUNKO", "a$b"]',
                                                             'OFERTAS_ELEGIDAS': 'false'})
eq('B1 · una marca con «$»: ROJO y NINGÚN cliente (la validación de HEO)',
   (_cod, 'selección de marcas no válida' in _log, 'CLIENTES_CREADOS=0' in _log), (1, True, True))

print('B2 · [marcas] la foto → barrido → la lista → los 4 CSV → cruce → Excel')
cod, log, bd, ruta, inicial = caso('marcas')
eq('B2 · el barrido sale en VERDE', cod, 0)
pas = [p for p in bd['tablas']['escaner2_pasada'] if p.get('run_id') == 717171][0]
PASADA = pas['id']
eq('B2 · 🔴 la pasada: OCIOSTOCK, modo marcas, esperando_csv, fichero 20 = previas 14 + foto 6, las marcas del director',
   (pas['proveedor'], pas['modo'], pas['estado'], pas['n_crudo'], pas['n_foto'],
    sum(pas['p_' + p] for p in e2.PUERTAS_PREVIAS), pas['n_tandas'], pas['ruta_lista'], pas['marcas'], pas['ofertas']),
   ('OCIOSTOCK', 'marcas', 'esperando_csv', 20, 6, 14, 1, 'ociostock/%s/eans.txt' % PASADA, DIRECTOR, False))
_foto = [f for f in bd['tablas']['escaner2_foto'] if f['pasada_id'] == PASADA]
eq('B2 · la foto en la base: 6, y las listas de las puertas previas: 9',
   (len(_foto), len([a for a in bd['tablas']['escaner2_apartado'] if a['pasada_id'] == PASADA])), (6, 9))
_alm = bd['storage']['escaner2']
_txt = lambda k: base64.b64decode(_alm[k]).decode('utf-8')  # noqa: E731
eq('B2 · en el almacén: eans.txt y barrido.json', sorted(k.split('/')[-1] for k in _alm), ['barrido.json', 'eans.txt'])
eq('B2 · eans.txt: los seis EAN (el de la caja, sin su « C6»)',
   sorted(_txt('ociostock/%s/eans.txt' % PASADA).split('\n')), sorted([E1, E2, E5, E7, E9, E10]))
_side = json.loads(_txt('ociostock/%s/barrido.json' % PASADA))
eq('B2 · barrido.json: de qué foto sale y el escalón de cada fila',
   (_side['disp_pasada'], sorted(_side['escalon']), _side['escalon']['900003']),
   (DISP, ['900001', '900003', '900007', '900012', '900014', '900016'],
    {'uds_escalon': 6, 'precio_unidad': 2.5, 'precio_escalon': 2.0, 'precio_pa': 1.98}))

_carpeta = 'ociostock/%s/csv/' % PASADA
# E1 y E2 se venden y dan COMPRAR; E5 se vende con poco margen; E9 no se vende; E7 y E10 no están en Keepa.
_filas = [(ASIN1, E1, '30', '12.95', 'Figura uno'), (ASIN2, E2, '30', '12.95', 'Figura dos'),
          (ASIN5, E5, '30', '9.95', 'Peluche cinco'), (ASIN9, E9, '2', '12.95', 'Figura nueve')]
for _i, _p in enumerate(('es', 'it', 'fr', 'de')):
    _alm[_carpeta + '20261007-1000%02d-KeepaExport-2026-10-07-VisualizadorDeProductos (%d).csv' % (_i, _i)] = \
        base64.b64encode(csv_real(_p, _filas)).decode()
bd.pop('programa', None)
json.dump(bd, open(ruta, 'w', encoding='utf-8'), default=str)
cod2, log2 = correr(ruta, 'escaner2_ociostock_cruce.py', {'SUPABASE_SERVICE_KEY': 'svc-de-mentira', 'PASADA': PASADA,
                                                           'GITHUB_RUN_ID': '727272'})
bd = json.load(open(ruta, encoding='utf-8'))
T = bd['tablas']
_alm = bd['storage']['escaner2']
eq('B2 · el cruce sale en VERDE', cod2, 0)
cr = [c for c in T['escaner2_cruce'] if c.get('run_id') == 727272][0]
eq('B2 · 🔴 el cruce: lista, cuadra, fichero 20 = previas 14 + 6 puertas, en los CUATRO países',
   (cr['estado'], cr['cuadra'], cr['n_crudo'], cr['n_previas'], cr['n_entradas'], sum(cr['n_' + x] for x in 'abcdef'),
    cr['paises_filtro'], cr['paises_calculo'], cr['paises_usados'], cr['aviso']),
   ('lista', True, 20, 14, 6, 6, ['ES', 'IT', 'FR', 'DE'], ['ES', 'IT', 'FR', 'DE'], ['ES', 'IT', 'FR', 'DE'], None))
_res = {f['ean_core']: r for r in T['escaner2_resultado_ean'] for f in T['escaner2_foto'] if f['id'] == r['foto_id']}
_pais = {}
for _r in T['escaner2_resultado_pais']:
    _pais.setdefault(_r['resultado_ean_id'], {})[_r['pais']] = _r
eq('B2 · 🔴 E1: COMPRAR con PA = precio_pa (2,97) en los cuatro países',
   (_res[E1]['puerta'], sorted(_pais[_res[E1]['id']]), {r['pa'] for r in _pais[_res[E1]['id']].values()}),
   ('f', ['DE', 'ES', 'FR', 'IT'], {2.97}))
eq('B2 · 🔴 E2: la caja «5 + 1», con su PA por figura (1,98), sin dividir entre 6',
   (_res[E2]['puerta'], _pais[_res[E2]['id']]['ES']['pa']), ('f', 1.98))
eq('B2 · E9 no se vende (puerta c) y E7/E10 no están en Keepa (puerta a)',
   (_res[E9]['puerta'], _res[E7]['puerta'], _res[E10]['puerta']), ('c', 'a', 'a'))

_xl = [k for k in _alm if k.startswith('ociostock/%s/%s/Escaner2_OCIOSTOCK_' % (PASADA, cr['id']))]
eq('B2 · el Excel, en su carpeta y con su nombre (el que guarda el cruce)', (len(_xl), cr['ruta_excel'] == _xl[0]),
   (1, True))
_wb = load_workbook(io.BytesIO(base64.b64decode(_alm[_xl[0]])))
eq('B2 · las hojas del PRO de HEO, en su orden (sin «Chase_manual», que es de HEO)', _wb.sheetnames,
   ['Análisis', 'Descartados', 'Ambiguos', 'Sin_rank', 'Precio por lote', 'Resumen', 'Comparación', 'Varias fichas',
    'Puertas', 'Puertas previas'])
_cab = [c.value for c in _wb['Análisis'][1]]
eq('🔴 B2 · «Análisis»: las del PRO de HEO y, AL FINAL, «Ficha compartida», «Uds. escalón» y «Precio unidad»',
   _cab, e2.columnas_analisis() + [e2.COLUMNA_FICHA_COMPARTIDA, 'Uds. escalón', 'Precio unidad'])
_ix = {n: _cab.index(n) for n in _cab}
_an = list(_wb['Análisis'].iter_rows(min_row=2, values_only=True))
_por = {}
for _r in _an:
    _por.setdefault(_r[_ix['EAN']], []).append(_r)
eq('B2 · «Análisis»: los que se venden, una fila por país (los cuatro)',
   sorted((k, sorted(r[_ix['País']] for r in v)) for k, v in _por.items()),
   sorted((k, ['DE', 'ES', 'FR', 'IT']) for k in (E1, E2 + ' C6', E5)))
eq('🔴 B2 · E1: PA 2,97, Uds. escalón 12 y Precio unidad 4,00 (sin escalón ni 0,99), en las cuatro filas',
   {(r[_ix['PA (€)']], r[_ix['Uds. escalón']], r[_ix['Precio unidad']], r[_ix['Decisión']]) for r in _por[E1]},
   {(2.97, 12, 4.0, 'COMPRAR')})
eq('🔴 B2 · la caja: «caja de 6» en «Coherencia caja», PA 1,98, 6 uds de escalón y 2,50 de precio de la ficha',
   {(r[_ix['Coherencia caja']], r[_ix['PA (€)']], r[_ix['Uds. escalón']], r[_ix['Precio unidad']])
    for r in _por[E2 + ' C6']}, {('caja de 6', 1.98, 6, 2.5)})
_tabla = list(_wb['Análisis'].tables.values())[0]
eq('B2 · la tabla de «Análisis» cubre las dos columnas nuevas',
   (_tabla.ref.split(':')[1].rstrip('0123456789'), [c.name for c in _tabla.tableColumns][-2:]),
   (_wb['Análisis'].cell(row=1, column=len(_cab)).column_letter, ['Uds. escalón', 'Precio unidad']))
_pu = list(_wb['Puertas'].iter_rows(values_only=True))
_pcab = list(_pu[0])
_pcaja = [r for r in _pu[1:] if r[0] == E2 + ' C6'][0]
eq('B2 · «Puertas»: la caja con su rótulo y el precio de la caja (6 × 2,50)',
   (_pcaja[_pcab.index('Caja')], _pcaja[_pcab.index('Precio caja (€)')], len(_pu) - 1),
   ('caja con chase · 6 uds', 15.0, 6))
_resu = {r[0]: r[1] for r in _wb['Resumen'].iter_rows(values_only=True)}
eq('B2 · el Resumen: modo, marcas del director, cuadre y las puertas previas con nombre de OcioStock',
   (_resu.get('Modo del barrido'), _resu.get('Marcas del director'), _resu.get('Cuadra'),
    _resu.get('Puerta previa · Chase suelto (no se compra nunca)'), _resu.get('Puerta previa · Sin stock o en preventa'),
    _resu.get('Catálogo crudo de OcioStock (el fichero)')),
   ('marcas de siempre', ', '.join(DIRECTOR), 'SÍ', 2, 2, 20))

print('B3 · 🔴 el registro (repo PÚBLICO) no lleva ni un EAN, precio, id de OcioStock ni nombre')
_prohibido = [E1, E2, E3, E4, E5, E6, E7, E8, E9, E10, EG, EP, ASIN1, ASIN2, ASIN5, ASIN9, '2.97', '1.98', '4.455',
              '3.069', '12.95', '900001', '900003', 'Figura inventada', 'Figura uno', 'Peluche']
eq('B3 · ni en el barrido ni en el cruce (si sale algo, la línea que lo lleva)',
   [(x, [ln for ln in (log + log2).splitlines() if x in ln][:2]) for x in _prohibido if x in log or x in log2], [])

print('B4 · solo se escribe en escaner2_* y en la carpeta de la pasada')
eq('B4 · las tablas escritas', sorted({t for _p, op, t in bd['ops'] if op != 'select'}),
   ['escaner2_apartado', 'escaner2_cruce', 'escaner2_foto', 'escaner2_pasada', 'escaner2_resultado_ean',
    'escaner2_resultado_pais'])
eq('B4 · y las sembradas acaban como empezaron',
   [t for t in SEMBRADAS if bd['tablas'][t] != json.loads(json.dumps(inicial['tablas'][t]))], [])
eq('B4 · en el almacén, solo ociostock/<pasada>/', sorted({'/'.join(s.split('/')[:2]) for s in bd.get('subidas', [])}),
   ['ociostock/' + PASADA])

print('B5 · los otros modos, y una foto que no es la suya')
cod, log, bd, _r, _i0 = caso('todas', {'MODO_BARRIDO': 'todas'})
_p = bd['tablas']['escaner2_pasada'][0]
eq('B5 · [todas] VERDE, 8 en la foto (CERDÁ y la sin marca), sin marcas guardadas',
   (cod, _p['modo'], _p['n_foto'], _p['p_marca_fuera'], _p['marcas']), (0, 'todas', 8, 0, None))
cod, log, bd, _r, _i0 = caso('elegidas', {'MODO_BARRIDO': 'elegidas', 'MARCAS_ELEGIDAS': '["Kids Licensing"]',
                                          'OFERTAS_ELEGIDAS': 'false'})
_p = bd['tablas']['escaner2_pasada'][0]
eq('B5 · [elegidas] VERDE, una en la foto, con su lista y ofertas a false',
   (cod, _p['modo'], _p['n_foto'], _p['marcas'], _p['ofertas']), (0, 'elegidas', 1, ['Kids Licensing'], False))
cod, log, bd, _r, _i0 = caso('ajena')
_p = bd['tablas']['escaner2_pasada'][0]
eq('B5 · 🔴 [ajena] una fila de otra pasada de disponibilidad: ROJO, la pasada fallida con el motivo y SIN foto',
   (cod, _p['estado'], 'no es la de su última pasada aplicada' in (_p['motivo_fallo'] or ''),
    bd['tablas'].get('escaner2_foto', [])), (1, 'fallida', True, []))

print()
if fallos:
    print('FALLAN %d: %s' % (len(fallos), fallos))
    sys.exit(1)
print('TODO OK')
