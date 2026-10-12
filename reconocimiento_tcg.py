# ============================================================
# RECONOCIMIENTO TCG (encargo T1, 12-oct-2026) - SOLO LECTURA
# ------------------------------------------------------------
# Baja el catalogo de TCG por el MISMO camino que director_tcg_prep.py (login
# PrestaShop por servidor + smcatalog downloadcatalog?format=excel) y lo deja en
# la carpeta `salida/` para que el workflow lo suba como artefacto CIFRADO.
#
# Nada de Supabase, nada de escribir en ningun sitio: solo lee la web de TCG.
# El repo es PUBLICO: el log solo lleva recuentos y huellas, nunca precios.
#
# Huella de CONTENIDO (no de bytes): el .xlsx lleva fechas internas, asi que dos
# descargas identicas pueden pesar distinto. La huella es el sha256 de las celdas
# de la hoja, fila a fila, como texto. Se guarda tambien el catalogo en CSV.
#
# Credenciales: SOLO en GitHub Secrets (TCG_USER / TCG_PASS).
# ============================================================
import os, sys, io, csv, json, time, random, hashlib
from datetime import datetime, timezone
from curl_cffi import requests as cr
from bs4 import BeautifulSoup
import openpyxl

BASE = 'https://tcgfactory.com'
LOGIN_URLS = [f'{BASE}/es/iniciar-sesion', f'{BASE}/es/mi-cuenta']
DOWNLOAD_URL = f'{BASE}/es/module/smcatalog/downloadcatalog?format=excel'
SALIDA = 'salida'

USER, PASS = os.environ.get('TCG_USER'), os.environ.get('TCG_PASS')
if not USER or not PASS:
    print("ERROR: faltan TCG_USER / TCG_PASS."); sys.exit(1)

VECES = max(1, int(os.environ.get('VECES') or '1'))
CADA_MIN = max(1, int(os.environ.get('CADA_MINUTOS') or '60'))
os.makedirs(SALIDA, exist_ok=True)


def pausa(a=1.5, b=4.0):
    time.sleep(random.uniform(a, b))


def entrar():
    s = cr.Session(impersonate='chrome120')
    form = None
    for url in LOGIN_URLS:
        try:
            r = s.get(url, timeout=60)
            soup = BeautifulSoup(r.text, 'html.parser')
            for f in soup.find_all('form'):
                if f.find('input', {'type': 'password'}):
                    form = f; break
            if form:
                break
        except Exception as e:
            print(f"   AVISO {url}: {e}")
        pausa()
    if not form:
        print("ERROR: no encuentro el login de TCG (¿Cloudflare/captcha?)."); sys.exit(1)
    action = form.get('action') or LOGIN_URLS[0]
    if action.startswith('/'):
        action = BASE + action
    data = {i.get('name'): i.get('value', '') for i in form.find_all('input') if i.get('name')}
    data['email'] = USER; data['password'] = PASS; data['submitLogin'] = '1'
    pausa()
    r2 = s.post(action, data=data, timeout=60)
    h = r2.text.lower()
    ok = ('cerrar sesión' in h) or ('cerrar sesion' in h) or ('logout' in h)
    print(f">>> login -> {r2.status_code} | sesion: {ok}", flush=True)
    if not ok:
        print("ERROR: el login no entro. Paro sin reintentar."); sys.exit(1)
    return s


def huella_contenido(contenido):
    wb = openpyxl.load_workbook(io.BytesIO(contenido), read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]
    h = hashlib.sha256()
    filas = []
    for fila in ws.iter_rows(values_only=True):
        txt = ['' if v is None else str(v) for v in fila]
        filas.append(txt)
        h.update(('\x1f'.join(txt) + '\x1e').encode('utf-8'))
    return wb.sheetnames, filas, h.hexdigest()


def descargar(s, n):
    sello = datetime.now(timezone.utc).strftime('%Y%m%dT%H%M%SZ')
    t0 = time.time()
    r = s.get(DOWNLOAD_URL, timeout=180)
    contenido = r.content
    if contenido[:2] != b'PK':
        print(f"ERROR descarga {n}: no es un Excel ({r.status_code}). Paro."); return False
    hojas, filas, huella = huella_contenido(contenido)
    with open(f'{SALIDA}/catalogo_{sello}.xlsx', 'wb') as fp:
        fp.write(contenido)
    with open(f'{SALIDA}/catalogo_{sello}.csv', 'w', newline='', encoding='utf-8') as fp:
        csv.writer(fp).writerows(filas)
    meta = {'n': n, 'utc': sello, 'bytes': len(contenido),
            'sha256_bytes': hashlib.sha256(contenido).hexdigest(),
            'sha256_contenido': huella, 'hojas': hojas, 'filas': len(filas),
            'segundos': round(time.time() - t0, 1)}
    with open(f'{SALIDA}/huellas.jsonl', 'a', encoding='utf-8') as fp:
        fp.write(json.dumps(meta) + '\n')
    # Al log: solo tamano, filas y huellas (el repo es publico; nada de precios)
    print(f">>> descarga {n}/{VECES} {sello}: {len(contenido)} bytes, {len(filas)} filas, "
          f"hojas {hojas}, contenido {huella[:16]}, bytes {meta['sha256_bytes'][:16]}", flush=True)
    return True


for n in range(1, VECES + 1):
    s = entrar()
    pausa()
    if not descargar(s, n):
        sys.exit(1)
    if n < VECES:
        time.sleep(CADA_MIN * 60 + random.uniform(-90, 90))
print(">>> FIN: solo lectura, nada escrito fuera de salida/.")
