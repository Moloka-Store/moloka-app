#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""ESCANER 2 · LA PASADA DE DISPONIBILIDAD DE OSMA (encargo AE, tramo 2 del plano Y, 01-oct-2026).

La lanza escaner2-osma-disponibilidad.yml. HOY A MANO: el reloj (cron-job.org, todos los dias a las 07:15
de Madrid) lo pone Fernando cuando esto este fusionado y la migracion 20261001124000 de la v2 aplicada.
Fernando, 01-oct-2026: «Quiero que se baje solo» y «Vale, a diario».

QUE HACE, EN ORDEN (el esqueleto de escaner2_heo_disponibilidad.py, con la descarga de OSMA):
  1. abre una pasada en `disp_pasada` (proveedor OSMA, estado 'leyendo', con el id del run), y lee en la
     base el minimo de filas (`disp_parametros.crudo_minimo`, 2.500) y la ultima pasada APLICADA;
  2. ENTRA en la web de OSMA (tienda Shopware: formulario de correo y contraseña, plano Y 1.1) con
     OSMA_USER / OSMA_PASS. 🔴 UN SOLO INTENTO: si al volver sigue en la pagina de entrada, la pasada se
     cierra 'fallida' («entrada rechazada») y no se reintenta: no bloquear la cuenta de Fernando;
  3. localiza en /en/downloads el enlace de `osma_articles.xlsx` de la carpeta `Ordersatz_Updateliste`
     (si la carpeta es otra pagina, la abre). 0 o mas de 1 → 'fallida', sin bajar nada;
  4. lo baja a memoria (la copia al almacen va al final, paso 8: solo la de un fichero bueno);
  5. VALIDA antes de subir nada a la base: abre, hoja `Data`, la cabecera (las columnas del plano, con su
     nombre exacto y en su orden; si falta o cambia una → 'fallida' con su nombre), filas >= el minimo
     (si no → 'rechazada_vaciado'), el codigo de articulo (`Item-no_BAN`, la llave) unico y no vacio
     (si no → 'fallida'), y la fecha de creacion del fichero (docProps/core.xml), con dos casos (AE2):
       · «YA ESTABA AL DIA»: la ultima pasada aplicada es de HOY (dia de Madrid) y el fichero no es mas nuevo
         que ella (es el que ya entro esta mañana, y alguien ha relanzado) → la pasada se cierra
         'rechazada' con motivo «al día: …» y el run sale en VERDE (OSMA_AL_DIA). Una 'rechazada' no cuenta
         en ningun sitio (medido el 01-oct-2026: el freno de disp_aplicar_pasada solo cuenta
         'rechazada_vaciado'; la frescura, la huella de la pantalla y las novedades solo miran 'aplicada');
       · «OSMA NO HA RENOVADO»: la ultima aplicada es de un dia anterior y el fichero no es mas nuevo que
         ella, o el fichero tiene mas de 48 h → 'rechazada', en ROJO entre semana y en VERDE el sabado y el
         domingo (dia de Madrid), que no son horas previstas (disp_parametros.horario_finde vacio);
  6. convierte cada fila con la tabla del apartado 2 del plano (llave = codigo OSMA; precio neto por
     unidad; disponible = stock > 0 o marcado de mas de 10.000), sube las filas a `disp_lectura` en lotes
     de 500 y deja en la pasada sus recuentos: declarado = llegados = filas de la hoja (OSMA no tiene «tres
     endpoints»; la integridad la dan las guardas del paso 5);
  7. llama a `disp_aplicar_pasada` (la funcion comun, SIN TOCAR) y relee la pasada en la base: lo que vale
     es lo que hay alli;
  8. SOLO SI HA QUEDADO 'aplicada', guarda la copia ENTERA del Excel en el almacen privado `escaner2`, en una
     ruta FIJA, `osma/ultimo/osma_articles.xlsx` (se pisa cada dia: Fernando, 01-oct-2026, «con tener la
     ultima foto ya me valdría de sobre»). Un fichero rechazado o roto no pisa la ultima copia buena. Si la
     subida falla, la pasada sigue aplicada (no se toca) y el run sale en ROJO con una linea que lo dice.
  9. (04-oct-2026) SOLO SI HA QUEDADO 'aplicada', y con la MISMA sesion: las FECHAS DE VUELTA de los agotados de
     OSMA que alimentan una ficha nuestra (escaner2_osma_fecha_vuelta.py: su cabecera dice cuales, como y con que
     frenos — 2 s entre peticiones, 40 como mucho, parada en seco al primer 403/429). Borra la de los que vuelven
     a estar disponibles. Si falla, la pasada sigue aplicada y la copia guardada; el run sale en ROJO diciendolo.
     Fernando, 04-oct-2026: «que reponer identifique que un producto de OSMA no esta disponible hoy pero lo
     estará pronto y lo enseñe».

🔴 REPO PUBLICO: LOS REGISTROS DE EJECUCION LOS VE CUALQUIERA. Este programa SOLO imprime estados y
   recuentos (filas, aplicadas, motivos que escribe el mismo). NUNCA: precios, stock, nombres o codigos de
   producto, EAN, el usuario, cookies, cabeceras, el HTML de una respuesta, la URL del fichero con su id ni
   el cuerpo de un error de la web. Un error inesperado se imprime SOLO por su tipo; su detalle va a
   `disp_pasada.motivo` (la base es privada). Sin modo verboso. Lo comprueba
   test_escaner2_osma_disponibilidad.py ejecutando el programa contra respuestas falsas.

🔑 EN SOMBRA: el interruptor de OSMA (disp_fuente) nace en 'memoria'; Reponer sigue leyendo escaner_memoria
   hasta que Fernando lo pase a 'disp'. Este programa no lo toca.
🔒 SOLO TOCA: `disp_pasada`, `disp_lectura`, la funcion `disp_aplicar_pasada`, y para LEER `disp_parametros`;
   y el almacen `escaner2` en `osma/ultimo/`. Ni escaner_memoria, ni escaner2_*, ni productos. Cero Keepa, cero
   Amazon. El paso 9 (otro fichero) escribe ADEMAS las dos columnas fecha_vuelta y fecha_vuelta_leida_en de
   `disp_estado`, y lee codigos_proveedor, productos y osma_mapa_ficha.
🔒 SIN LOS SECRETOS, NO SE CORRE: sin OSMA_USER/OSMA_PASS (o sin la base) aborta antes de abrir ninguna
   pasada ni ningun cliente.

Uso:  python escaner2_osma_disponibilidad.py            (la pasada)
      python escaner2_osma_disponibilidad.py --rescate  (ultimo paso del workflow si el run fallo)
      python escaner2_osma_disponibilidad.py --solo-fechas [--sin-escribir]
            (solo el paso 9, contra la ultima foto aplicada: entra, no baja el Excel ni abre pasada; con
             --sin-escribir no toca la base y solo imprime los recuentos)
"""
import io
import os
import re
import sys
import zipfile
from datetime import datetime, timedelta, timezone
from html.parser import HTMLParser
from urllib.parse import urljoin, urlparse

PROVEEDOR = 'OSMA'
LOTE = 500
BUCKET = 'escaner2'
WEB = 'https://osma-werm.com'
RUTA_ENTRADA = '/en/account/login'
RUTA_DESCARGAS = '/en/downloads'
FICHERO = 'osma_articles.xlsx'
CARPETA = 'Ordersatz_Updateliste'
HOJA = 'Data'
MAX_EDAD = timedelta(hours=48)
RUTA_COPIA = 'osma/ultimo/' + FICHERO
TAM_MAXIMO = 50 * 1024 * 1024       # el tope del almacen escaner2 (medido: 52.428.800)
ESPERA_S = 60
AGENTE = 'Mozilla/5.0 (compatible; Moloka-escaner2/1.0)'

# 🔑 LA CABECERA: las columnas del plano Y con su nombre exacto, EN SU ORDEN (lista de Cowork, 30-sep y
#    01-oct-2026, leida del fichero). Entre medias puede haber otras (grupo y subgrupo en aleman e ingles,
#    descripciones largas: sus nombres exactos no estan medidos y no se usan). Si una de estas falta,
#    cambia de nombre o de orden, la pasada es 'fallida' y dice cual.
COLUMNAS = (
    'BRAND', 'Description_short_ger_SDESC', 'Item-no_BAN', 'EAN_GTIN', 'MOQ', 'Price_net', 'Last price',
    'SALE', 'New', 'Stock_available_pcs', 'Stock_over10k_pcs', 'Ordered_pcs', 'Discontinued', 'MHD_EXP',
    'Inner_Pack_pcs_IP', 'Outer_Case_pcs_OC', 'GTIN_Case', 'Layer_pcs', 'PAL_pcs', 'PAL_cm_HGT',
    'Weight_gram', 'VAT',
)
_NOMBRE_RE = re.compile(r'(?<![\w.])' + re.escape(FICHERO) + r'(?![\w.])', re.I)
_CARPETA_RE = re.compile(r'(?<![\w.])' + re.escape(CARPETA) + r'(?![\w.])', re.I)
_ENLACE_FICHERO_RE = re.compile(r'/downloads/file/([^/?#]+)')
_CREADO_RE = re.compile(rb'<dcterms:created[^>]*>([^<]+)</dcterms:created>')
_NO_MARCA = {'', '0', 'no', 'nein', 'n', 'false', 'falsch', '-', 'none'}


def abortar(motivo):
    """Un run que no hace su trabajo sale en ROJO y con una linea que se puede buscar. `motivo` es SIEMPRE
    un texto escrito por este programa: nunca el de un error ni nada de la web."""
    print(f"OSMA_NO_EJECUTADA: {motivo}")
    sys.exit(1)


class Rechazo(Exception):
    """La pasada no se aplica: `estado` ('fallida', 'rechazada' o 'rechazada_vaciado') y un `motivo` que
    escribe este programa (estados y recuentos; se puede imprimir)."""

    def __init__(self, estado, motivo):
        super().__init__(motivo)
        self.estado, self.motivo = estado, motivo


class AlDia(Rechazo):
    """El fichero de hoy ya esta aplicado (relanzar el mismo dia no es un fallo): 'rechazada' con motivo «al día»,
    que no cuenta en ningun sitio, y el run en verde."""

    def __init__(self, pasada_aplicada):
        super().__init__('rechazada', f'al día: el fichero de hoy ya está aplicado (pasada {pasada_aplicada}); nada nuevo')
        self.pasada_aplicada = pasada_aplicada


def _ahora():
    return datetime.now(timezone.utc)


def _domingo_ultimo(anio, mes):
    """El ultimo domingo del mes (marzo u octubre), a la 01:00 UTC: cuando cambia la hora en Europa."""
    d = datetime(anio, mes, 31, 1, tzinfo=timezone.utc)
    return d - timedelta(days=(d.weekday() + 1) % 7)


def hora_de_madrid(t):
    """La hora de Madrid de un instante: UTC+2 del ultimo domingo de marzo al ultimo de octubre (a la 01:00 UTC,
    la regla europea) y UTC+1 el resto. Sin zoneinfo: no depende de los datos de zonas de la maquina."""
    t = t.astimezone(timezone.utc)
    verano = _domingo_ultimo(t.year, 3) <= t < _domingo_ultimo(t.year, 10)
    return (t + timedelta(hours=2 if verano else 1)).replace(tzinfo=None)


# ── Lectura de celdas ─────────────────────────────────────────────────────────────────────
def marcado(v):
    """Una casilla de marca de OSMA (SALE, Discontinued, Stock_over10k_pcs…): vacia, 0 o «no» = no; lo demas, si."""
    if v is None or v is False:
        return False
    if v is True:
        return True
    if isinstance(v, (int, float)):
        return v != 0
    return str(v).strip().lower() not in _NO_MARCA


def numero(v):
    """Un numero de la hoja: None si la celda esta vacia; ValueError si no es un numero (no se inventa)."""
    if v is None or isinstance(v, bool):
        if v is None:
            return None
        raise ValueError('booleano')
    if isinstance(v, (int, float)):
        return float(v)
    s = str(v).strip().replace(' ', '')
    if not s:
        return None
    if re.fullmatch(r'-?\d+([.,]\d+)?', s):
        return float(s.replace(',', '.'))
    raise ValueError('no numerico')


def texto_codigo(v):
    """El codigo de articulo como texto: 1538 y 1538.0 son '1538'."""
    if v is None:
        return ''
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v).strip()


def ean_core(v):
    """(ean_original, ean_core, regla). El EAN con el que se cruza: 12 o 13 cifras (14 con un 0 delante se
    queda en 13; si Excel lo guardo como numero y perdio los ceros de delante, se rellenan). Vacio: sin EAN
    y sin regla. Otra forma: ean_core vacio y regla 'ean_forma_rara' (se guarda igual: la llave es el codigo)."""
    if v is None or (isinstance(v, str) and not v.strip()):
        return None, None, None
    numerico = isinstance(v, (int, float)) and not isinstance(v, bool)
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    original = str(v).strip()
    s = original
    if len(s) == 14 and s.isdigit() and s.startswith('0'):
        s = s[1:]
    if numerico and s.isdigit() and 9 <= len(s) < 12:
        s = s.zfill(13)
    if s.isdigit() and len(s) in (12, 13):
        return original, s, None
    return original, None, 'ean_forma_rara'


# ── La pagina de descargas, sin dependencias: un arbol minimo con html.parser ────────────
_VACIOS = {'area', 'base', 'br', 'col', 'embed', 'hr', 'img', 'input', 'link', 'meta', 'param', 'source',
           'track', 'wbr'}


class _Nodo:
    __slots__ = ('tag', 'attrs', 'hijos', 'padre', 'trozos')

    def __init__(self, tag, attrs, padre):
        self.tag, self.attrs, self.hijos, self.padre, self.trozos = tag, dict(attrs), [], padre, []

    def texto(self):
        partes = list(self.trozos)
        for h in self.hijos:
            partes.append(h.texto())
        return ' '.join(p for p in partes if p)


class _Arbol(HTMLParser):
    def __init__(self):
        super().__init__(convert_charrefs=True)
        self.raiz = _Nodo('raiz', {}, None)
        self.pila = [self.raiz]
        self.enlaces = []

    def handle_starttag(self, tag, attrs):
        n = _Nodo(tag, attrs, self.pila[-1])
        self.pila[-1].hijos.append(n)
        if tag == 'a':
            self.enlaces.append(n)
        if tag not in _VACIOS:
            self.pila.append(n)

    def handle_startendtag(self, tag, attrs):
        n = _Nodo(tag, attrs, self.pila[-1])
        self.pila[-1].hijos.append(n)
        if tag == 'a':
            self.enlaces.append(n)

    def handle_endtag(self, tag):
        for i in range(len(self.pila) - 1, 0, -1):
            if self.pila[i].tag == tag:
                del self.pila[i:]
                return

    def handle_data(self, data):
        d = data.strip()
        if d:
            self.pila[-1].trozos.append(d)


def _id_fichero(href):
    m = _ENLACE_FICHERO_RE.search(urlparse(href or '').path)
    return m.group(1) if m else None


def _buscar_en_pagina(html, base_url, exigir_carpeta):
    """Los enlaces de `osma_articles.xlsx` de una pagina. Cada fichero (por su id) es la «ficha» mas grande
    que lo contiene solo a el; su nombre se busca en el texto de esa ficha. Con `exigir_carpeta`, la ficha
    tiene que estar dentro del primer bloque que nombra la carpeta, y ese bloque no puede contener otro
    `osma_articles.xlsx` (si no, no se sabe de que carpeta es: no se elige). Devuelve (urls, cuentas)."""
    a = _Arbol()
    a.feed(html or '')
    ids = {}
    for e in a.enlaces:
        i = _id_fichero(e.attrs.get('href'))
        if i:
            ids.setdefault(i, []).append(e)

    memo = {}

    def ids_de(nodo):
        if id(nodo) not in memo:
            s = set()
            if nodo.tag == 'a':
                i = _id_fichero(nodo.attrs.get('href'))
                if i:
                    s.add(i)
            for h in nodo.hijos:
                s |= ids_de(h)
            memo[id(nodo)] = s
        return memo[id(nodo)]

    fichas = {}
    for i, enlaces in ids.items():
        n = enlaces[0]
        while n.padre is not None and ids_de(n.padre) == {i}:
            n = n.padre
        fichas[i] = n
    osma = {i for i, n in fichas.items() if _NOMBRE_RE.search(n.texto())}
    elegidos = []
    for i in sorted(osma):
        if not exigir_carpeta:
            elegidos.append(i)
            continue
        n = fichas[i]
        while n is not None and not _CARPETA_RE.search(n.texto()):
            n = n.padre
        if n is not None and len(ids_de(n) & osma) == 1:
            elegidos.append(i)
    urls = [urljoin(base_url, ids[i][0].attrs['href']) for i in elegidos]
    carpetas = sorted({urljoin(base_url, e.attrs.get('href') or '') for e in a.enlaces
                       if _CARPETA_RE.search(e.texto()) and e.attrs.get('href') and not _id_fichero(e.attrs.get('href'))})
    return urls, {'ficheros': len(ids), 'con_el_nombre': len(osma), 'en_la_carpeta': len(elegidos),
                  'enlaces_a_la_carpeta': carpetas}


def _misma_web(url):
    return urlparse(url).netloc == urlparse(WEB).netloc


def localizar_enlace(sesion, html, base_url):
    """La URL del fichero del dia, o Rechazo('fallida'). Solo imprime recuentos."""
    urls, c = _buscar_en_pagina(html, base_url, exigir_carpeta=True)
    print(f">>> Página de descargas: {c['ficheros']} fichero(s), {c['con_el_nombre']} con el nombre {FICHERO}, "
          f"{c['en_la_carpeta']} en la carpeta {CARPETA}, {len(c['enlaces_a_la_carpeta'])} enlace(s) a la carpeta.",
          flush=True)
    if len(urls) == 1:
        return urls[0]
    if len(urls) > 1:
        raise Rechazo('fallida', f'enlace ambiguo: {len(urls)} {FICHERO} en la carpeta {CARPETA}; no se baja nada')
    if c['con_el_nombre']:
        raise Rechazo('fallida', f'enlace no encontrado o ambiguo: hay {c["con_el_nombre"]} {FICHERO} y ninguno está solo '
                                 f'en la carpeta {CARPETA}; no se baja nada')
    carpetas = [u for u in c['enlaces_a_la_carpeta'] if _misma_web(u)]
    if len(carpetas) != 1:
        raise Rechazo('fallida', f'enlace no encontrado: ningún {FICHERO} en la página de descargas y '
                                 f'{len(carpetas)} enlace(s) a la carpeta {CARPETA}; no se baja nada')
    r = sesion.get(carpetas[0], timeout=ESPERA_S, allow_redirects=True)
    if r.status_code != 200:
        raise Rechazo('fallida', f'la carpeta {CARPETA} respondió {r.status_code}; no se baja nada')
    urls, c = _buscar_en_pagina(r.text, r.url, exigir_carpeta=False)
    print(f">>> Página de la carpeta: {c['ficheros']} fichero(s), {c['con_el_nombre']} con el nombre {FICHERO}.", flush=True)
    if len(urls) != 1:
        raise Rechazo('fallida', f'enlace {"no encontrado" if not urls else "ambiguo"}: {len(urls)} {FICHERO} en la '
                                 f'carpeta {CARPETA}; no se baja nada')
    return urls[0]


# ── La web ────────────────────────────────────────────────────────────────────────────────
def _en_la_entrada(url):
    return urlparse(url or '').path.rstrip('/').endswith(RUTA_ENTRADA)


def entrar(sesion, usuario, clave):
    """UN SOLO intento. Si al volver sigue en la pagina de entrada → Rechazo «entrada rechazada»."""
    r = sesion.get(WEB + RUTA_ENTRADA, timeout=ESPERA_S, allow_redirects=True)
    if r.status_code != 200:
        raise Rechazo('fallida', f'la página de entrada respondió {r.status_code}; no se ha intentado entrar')
    r = sesion.post(WEB + RUTA_ENTRADA, data={'username': usuario, 'password': clave,
                                              'redirectTo': 'frontend.account.home.page', 'redirectParameters': '[]'},
                    timeout=ESPERA_S, allow_redirects=True)
    if r.status_code >= 400:
        raise Rechazo('fallida', f'entrada rechazada: la web respondió {r.status_code}; no se reintenta')
    if _en_la_entrada(r.url):
        raise Rechazo('fallida', 'entrada rechazada: la web ha vuelto a la página de entrada; no se reintenta '
                                 '(no bloquear la cuenta)')


def bajar(sesion, url):
    r = sesion.get(url, timeout=ESPERA_S, allow_redirects=True)
    if r.status_code != 200:
        raise Rechazo('fallida', f'la descarga respondió {r.status_code}')
    contenido = r.content or b''
    if _en_la_entrada(r.url):
        raise Rechazo('fallida', 'la descarga ha vuelto a la página de entrada: la sesión no vale')
    if len(contenido) > TAM_MAXIMO:
        raise Rechazo('fallida', f'lo bajado pesa {len(contenido)} bytes, más que el tope del almacén')
    if not contenido.startswith(b'PK\x03\x04'):
        raise Rechazo('fallida', 'lo bajado no es un xlsx (¿la web ha devuelto una página?); no se guarda')
    return contenido


# ── El fichero ────────────────────────────────────────────────────────────────────────────
def fecha_de_creacion(contenido):
    """La fecha «created» de docProps/core.xml, en UTC; Rechazo si no esta."""
    try:
        with zipfile.ZipFile(io.BytesIO(contenido)) as z:
            core = z.read('docProps/core.xml')
    except Exception:
        raise Rechazo('fallida', 'el fichero no dice cuándo se creó (sin docProps/core.xml)') from None
    m = _CREADO_RE.search(core)
    if not m:
        raise Rechazo('fallida', 'el fichero no dice cuándo se creó (sin dcterms:created)')
    try:
        t = datetime.fromisoformat(m.group(1).decode('ascii').strip().replace('Z', '+00:00'))
    except ValueError:
        raise Rechazo('fallida', 'el fichero dice una fecha de creación que no se entiende') from None
    return t if t.tzinfo else t.replace(tzinfo=timezone.utc)


def comprobar_fecha(creado, ultima_aplicada, ahora, ultima_id=None):
    """🔴 OSMA regenera el fichero de madrugada. Uno que no es mas nuevo que la ultima pasada aplicada ya esta
    aplicado: si esa pasada es de HOY (dia de Madrid), es un relanzamiento → AlDia (verde); si es de otro dia,
    OSMA no lo ha renovado → 'rechazada'. Uno de mas de 48 h esta parado → 'rechazada'. La foto no se toca."""
    if ultima_aplicada is not None and creado <= ultima_aplicada:
        if hora_de_madrid(ultima_aplicada).date() == hora_de_madrid(ahora).date():
            raise AlDia(ultima_id)
        raise Rechazo('rechazada', f'OSMA no ha regenerado el fichero: creado el {creado:%Y-%m-%d %H:%M} UTC, antes de la '
                                   f'última pasada aplicada ({ultima_aplicada:%Y-%m-%d %H:%M} UTC, de otro día)')
    if ahora - creado > MAX_EDAD:
        raise Rechazo('rechazada', f'OSMA no ha regenerado el fichero: creado el {creado:%Y-%m-%d %H:%M} UTC, hace más '
                                   f'de 48 h')


def comprobar_cabecera(cabecera):
    """{columna: indice}. Rechazo('fallida') con el nombre de la que falta, se repite o esta fuera de su orden."""
    vistas = {}
    for i, nombre in enumerate(cabecera):
        if not nombre:
            continue
        if nombre in vistas:
            raise Rechazo('fallida', f'cabecera: la columna «{nombre}» está repetida')
        vistas[nombre] = i
    indices, anterior = {}, -1
    for col in COLUMNAS:
        pos = vistas.get(col)
        if pos is None:
            raise Rechazo('fallida', f'cabecera: falta la columna «{col}» (o ha cambiado de nombre)')
        if pos <= anterior:
            raise Rechazo('fallida', f'cabecera: la columna «{col}» está fuera de su orden')
        indices[col], anterior = pos, pos
    return indices


def leer_hoja(contenido):
    """(cabecera, filas) de la hoja Data: la cabecera como textos y las filas no vacias como tuplas."""
    import openpyxl
    try:
        libro = openpyxl.load_workbook(io.BytesIO(contenido), read_only=True, data_only=True)
    except Exception:
        raise Rechazo('fallida', 'el fichero no abre como xlsx') from None
    try:
        if HOJA not in libro.sheetnames:
            raise Rechazo('fallida', f'el fichero no tiene la hoja {HOJA}')
        filas = libro[HOJA].iter_rows(values_only=True)
        primera = next(filas, None)
        if primera is None:
            raise Rechazo('fallida', f'la hoja {HOJA} está vacía')
        cabecera = ['' if c is None else str(c).strip() for c in primera]
        datos = [f for f in filas if any(c is not None and str(c).strip() != '' for c in f)]
    finally:
        libro.close()
    return cabecera, datos


def comprobar_codigos(datos, ix):
    vacios, vistos, repetidos = 0, set(), set()
    for f in datos:
        c = texto_codigo(f[ix] if ix < len(f) else None)
        if not c:
            vacios += 1
        elif c in vistos:
            repetidos.add(c)
        else:
            vistos.add(c)
    if vacios:
        raise Rechazo('fallida', f'{vacios} fila(s) sin código de artículo (Item-no_BAN, la llave)')
    if repetidos:
        raise Rechazo('fallida', f'{len(repetidos)} código(s) de artículo repetido(s) (Item-no_BAN, la llave)')


def convertir(datos, ix):
    """Las filas de disp_lectura (tabla del apartado 2 del plano Y) y sus recuentos. Rechazo('fallida') si un
    precio o un stock no son numeros (se cuentan, no se dicen)."""
    def celda(f, col):
        i = ix[col]
        return f[i] if i < len(f) else None

    filas = []
    c = {'n_leidas': 0, 'n_disponibles': 0, 'n_agotados': 0, 'n_sin_dato_precio': 0, 'sin_ean': 0,
         'ean_forma_rara': 0, 'sin_marca': 0, 'mas_de_10000': 0}
    precio_malo = stock_malo = 0
    for f in datos:
        over = marcado(celda(f, 'Stock_over10k_pcs'))
        try:
            stock = numero(celda(f, 'Stock_available_pcs'))
        except ValueError:
            stock = None
            if not over:
                stock_malo += 1
        try:
            precio = numero(celda(f, 'Price_net'))
        except ValueError:
            precio = None
            precio_malo += 1
        if precio is not None and precio < 0:
            precio_malo += 1
            precio = None
        disponible = over or (stock is not None and stock > 0)
        original, core, regla = ean_core(celda(f, 'EAN_GTIN'))
        marca = celda(f, 'BRAND')
        marca = str(marca).strip() if marca is not None and str(marca).strip() else None
        nombre = celda(f, 'Description_short_ger_SDESC')
        nombre = str(nombre).strip() if nombre is not None and str(nombre).strip() else None
        filas.append({
            'producto_prov': texto_codigo(celda(f, 'Item-no_BAN')),
            'ean_original': original, 'ean_core': core,
            'marca': marca, 'nombre': nombre, 'categoria': None,
            'es_caja': False, 'uds_caja': None, 'es_chase': False,
            'precio_catalogo': None if precio is None else round(precio, 6),
            'precio_unidad': None if precio is None else round(precio, 6),
            'disponible': disponible, 'disponibilidad': None,
            'en_oferta': marcado(celda(f, 'SALE')), 'preorder': None,
            'fin_de_vida': marcado(celda(f, 'Discontinued')),
            'regla': regla, 'aviso': None,
            'sin_dato_disponibilidad': False, 'sin_dato_precio': precio is None,
        })
        c['n_leidas'] += 1
        c['n_disponibles' if disponible else 'n_agotados'] += 1
        c['n_sin_dato_precio'] += precio is None
        c['sin_ean'] += original is None
        c['ean_forma_rara'] += regla == 'ean_forma_rara'
        c['sin_marca'] += marca is None
        c['mas_de_10000'] += over
    if precio_malo:
        raise Rechazo('fallida', f'{precio_malo} fila(s) con Price_net que no es un número válido')
    if stock_malo:
        raise Rechazo('fallida', f'{stock_malo} fila(s) con Stock_available_pcs que no es un número')
    return filas, c


# ── La pasada ─────────────────────────────────────────────────────────────────────────────
def _cerrar(sb, pasada, estado, motivo, extra=None):
    """La pasada queda cerrada con su estado y su motivo, y sin nada leido colgando (disp_lectura es de paso)."""
    sb.table('disp_lectura').delete().eq('pasada_id', pasada).execute()
    datos = dict(extra or {}, estado=estado, motivo=motivo[:1000], terminada_en=_ahora().isoformat())
    sb.table('disp_pasada').update(datos).eq('id', pasada).eq('estado', 'leyendo').execute()


def _fecha_base(v):
    if not v:
        return None
    t = datetime.fromisoformat(str(v).replace('Z', '+00:00'))
    return t if t.tzinfo else t.replace(tzinfo=timezone.utc)


def pasada(sb, sesion, usuario, clave, run_id, ahora=None):
    """La pasada entera. Devuelve el codigo de salida del programa."""
    abrir = {'proveedor': PROVEEDOR, 'estado': 'leyendo', 'run_id': int(run_id) if run_id.isdigit() else None}
    pid = sb.table('disp_pasada').insert(abrir).execute().data[0]['id']
    print(f">>> Pasada de disponibilidad de OSMA {pid} abierta (leyendo).", flush=True)
    rec = contenido = None
    try:
        par = sb.table('disp_parametros').select('crudo_minimo').eq('proveedor', PROVEEDOR).execute().data or [{}]
        minimo = par[0].get('crudo_minimo')
        if not isinstance(minimo, int) or isinstance(minimo, bool) or minimo <= 0:
            raise Rechazo('fallida', 'sin mínimo: disp_parametros no da un crudo_minimo válido para OSMA; no se baja nada')
        ult = (sb.table('disp_pasada').select('id, creada_en').eq('proveedor', PROVEEDOR).eq('estado', 'aplicada')
                 .order('creada_en', desc=True).limit(1).execute().data or [{}])
        ultima_aplicada, ultima_id = _fecha_base(ult[0].get('creada_en')), ult[0].get('id')

        entrar(sesion, usuario, clave)
        print(">>> Entrada en la web de OSMA: aceptada.", flush=True)
        r = sesion.get(WEB + RUTA_DESCARGAS, timeout=ESPERA_S, allow_redirects=True)
        if r.status_code != 200 or _en_la_entrada(r.url):
            raise Rechazo('fallida', f'la página de descargas no se ha abierto con la sesión (respondió {r.status_code})')
        url = localizar_enlace(sesion, r.text, r.url)
        contenido = bajar(sesion, url)
        print(f">>> Fichero bajado: {len(contenido)} bytes.", flush=True)

        cabecera, datos = leer_hoja(contenido)
        ix = comprobar_cabecera(cabecera)
        n = len(datos)
        rec = {'n_declarado': n, 'n_crudo': n, 'n_declarado_precios': n, 'n_precios': n,
               'n_declarado_disponibilidades': n, 'n_disponibilidades': n, 'n_sin_gtin': 0,
               'n_duplicados': 0, 'n_duplicados_precios': 0, 'n_duplicados_disponibilidades': 0}
        print(f">>> Hoja {HOJA}: {n} filas; cabecera conforme ({len(COLUMNAS)} columnas comprobadas de "
              f"{sum(1 for x in cabecera if x)}).", flush=True)
        if n < minimo:
            raise Rechazo('rechazada_vaciado', f'vaciado: {n} filas en la hoja, por debajo del mínimo {minimo}')
        comprobar_codigos(datos, ix['Item-no_BAN'])
        creado = fecha_de_creacion(contenido)
        ahora = ahora or _ahora()
        print(f">>> Fichero creado el {creado:%Y-%m-%d %H:%M} UTC (hace {(ahora - creado).total_seconds() / 3600:.1f} h).",
              flush=True)
        comprobar_fecha(creado, ultima_aplicada, ahora, ultima_id)

        filas, cuentas = convertir(datos, ix)
        print(f">>> Leídas {cuentas['n_leidas']} (disponibles {cuentas['n_disponibles']}, agotados {cuentas['n_agotados']}; "
              f"más de 10.000 {cuentas['mas_de_10000']}) · sin precio {cuentas['n_sin_dato_precio']} · sin EAN "
              f"{cuentas['sin_ean']} · EAN de forma rara {cuentas['ean_forma_rara']} · sin marca {cuentas['sin_marca']}",
              flush=True)
        sb.table('disp_pasada').update(rec).eq('id', pid).execute()
        for i in range(0, len(filas), LOTE):
            sb.table('disp_lectura').insert([dict(f, pasada_id=pid) for f in filas[i:i + LOTE]]).execute()
        sb.table('disp_pasada').update({
            'n_leidas': cuentas['n_leidas'], 'n_disponibles': cuentas['n_disponibles'],
            'n_agotados': cuentas['n_agotados'], 'n_sin_dato_disponibilidad': 0,
            'n_sin_dato_precio': cuentas['n_sin_dato_precio'],
        }).eq('id', pid).execute()
        sb.rpc('disp_aplicar_pasada', {'p_pasada': pid}).execute()
    except Rechazo as ex:
        _cerrar(sb, pid, ex.estado, ex.motivo, rec)
        if isinstance(ex, AlDia):
            # 🔑 Relanzar el mismo día no es un fallo (Fernando, 01-oct-2026, del rojo: «eso no me gusta»).
            print(f"OSMA_AL_DIA: el fichero de hoy ya está aplicado (pasada {ex.pasada_aplicada}); nada nuevo.", flush=True)
            return 0
        if ex.estado == 'rechazada' and hora_de_madrid(ahora or _ahora()).weekday() >= 5:
            # 🔑 El fin de semana no es hora prevista (disp_parametros.horario_finde vacío): un fichero sin
            #    regenerar el sábado o el domingo no es un fallo. Se dice y el run sale bien.
            print(f"OSMA_SIN_FICHERO_NUEVO: pasada {pid} {ex.estado}: {ex.motivo} · fin de semana: no es un fallo.", flush=True)
            return 0
        print(f"OSMA_NO_APLICADA: pasada {pid} {ex.estado}: {ex.motivo}", flush=True)
        return 1
    except Exception as ex:
        # 🔴 REPO PÚBLICO: el texto de un error puede llevar datos (la URL del fichero, una fila de la base…).
        #    Al registro, SOLO su tipo; el detalle, a la base, que es privada.
        motivo_log = f'error inesperado ({type(ex).__name__}); el detalle queda en disp_pasada.motivo'
        try:
            _cerrar(sb, pid, 'fallida', f'{type(ex).__name__}: {ex}', rec)
        except Exception as ex2:
            motivo_log += f'; tampoco se ha podido cerrar la pasada ({type(ex2).__name__})'
        print(f"OSMA_NO_APLICADA: pasada {pid} fallida: {motivo_log}", flush=True)
        return 1

    # Lo que vale es lo que hay en la base.
    fila = sb.table('disp_pasada').select('*').eq('id', pid).execute().data[0]
    print(">>> EN LA BASE: " + ' · '.join(f'{k} {fila.get(k)}' for k in (
        'estado', 'primera', 'caida_aceptada', 'n_crudo', 'n_leidas', 'n_disponibles', 'n_agotados',
        'n_sin_dato_precio', 'n_entran', 'n_vuelven', 'n_salen', 'n_a_disponible', 'n_a_agotado',
        'n_cambio_precio', 'n_ausentes', 'n_en_catalogo', 'n_disponibles_estado')), flush=True)
    if fila.get('estado') != 'aplicada':
        # Los motivos de disp_aplicar_pasada son recuentos (los escribe la función, no la web).
        print(f"OSMA_NO_APLICADA: pasada {pid} {fila.get('estado')}: {fila.get('motivo')}", flush=True)
        return 1
    print(f">>> PASADA APLICADA: {fila.get('n_en_catalogo')} artículos de OSMA en el catálogo, "
          f"{fila.get('n_disponibles_estado')} disponibles.", flush=True)
    codigo = 0
    # 🔑 LA COPIA, SOLO DE UNA PASADA APLICADA y en la ruta fija: un fichero rechazado no pisa la última buena.
    try:
        guardar_copia(sb, contenido)
        print(f">>> Copia entera del Excel guardada en el almacén {BUCKET}: {RUTA_COPIA}", flush=True)
    except Exception as ex:
        print(f"OSMA_COPIA_NO_GUARDADA: la pasada {pid} está aplicada (no se toca); la copia del Excel no se ha "
              f"guardado ({type(ex).__name__}).", flush=True)
        codigo = 1  # la copia no se ha guardado: que se vea
    # 🔑 LAS FECHAS DE VUELTA (04-oct-2026), SOLO DE UNA PASADA APLICADA y con la misma sesión. Si fallan, la
    #    pasada sigue aplicada y la copia guardada: no se toca nada de lo de arriba.
    codigo = max(codigo, fechas_de_vuelta(sb, sesion, ahora or _ahora(), escribir=True))
    if fila.get('caida_aceptada'):
        print(f"CAIDA_ACEPTADA: la pasada {pid} se ha aplicado como NUEVA REFERENCIA tras varios rechazos estables por "
              f"el freno del 90 %: hay que mirar si OSMA ha caído de verdad.", flush=True)
        return 1
    return codigo


def inicio_del_dia_de_madrid(t):
    """Las 00:00 del día de Madrid de `t`, como instante UTC."""
    local = hora_de_madrid(t)
    desfase = local - t.astimezone(timezone.utc).replace(tzinfo=None)
    return (local.replace(hour=0, minute=0, second=0, microsecond=0) - desfase).replace(tzinfo=timezone.utc)


def fechas_de_vuelta(sb, sesion, ahora, escribir):
    """El repaso de escaner2_osma_fecha_vuelta.py. Devuelve el código de salida: 1 si la web ha parado (403/429 o
    sin sesión), si ninguna ficha abierta se ha podido leer (la web ha cambiado) o si algo revienta."""
    import escaner2_osma_fecha_vuelta as fv
    try:
        cuentas, parada = fv.repasar(sb, sesion, inicio_del_dia_de_madrid(ahora), escribir=escribir)
    except Exception as ex:
        # 🔴 REPO PÚBLICO: solo el tipo (el texto puede llevar una fila de la base o una URL).
        print(f"OSMA_FECHAS_NO_LEIDAS: el repaso de las fechas de vuelta ha fallado ({type(ex).__name__}); la foto "
              f"no se toca.", flush=True)
        return 1
    print(fv.resumen(cuentas, escribir), flush=True)
    if parada:
        print(f"OSMA_FECHAS_PARADAS: {parada}; no se ha reintentado y quedan {cuentas['sin_turno']} sin leer (mañana).",
              flush=True)
        return 1
    if cuentas['abiertas'] and not (cuentas['con_fecha'] + cuentas['sin_fecha'] + cuentas['fecha_rara']):
        print("OSMA_FECHAS_WEB_CAMBIADA: ninguna ficha abierta se ha podido leer (sin el bloque del producto, en otro "
              "idioma o de otro artículo); no se ha escrito ninguna fecha.", flush=True)
        return 1
    return 0


def solo_fechas(sb, sesion, usuario, clave, escribir):
    """--solo-fechas: entra y repasa las fechas contra la ÚLTIMA foto aplicada de OSMA, sin bajar el Excel ni
    abrir pasada. Con --sin-escribir no toca la base (ni borra ni apunta): solo los recuentos."""
    ult = (sb.table('disp_pasada').select('id').eq('proveedor', PROVEEDOR).eq('estado', 'aplicada')
             .order('creada_en', desc=True).limit(1).execute().data or [])
    if not ult:
        abortar('no hay ninguna pasada de OSMA aplicada: no hay foto sobre la que leer fechas')
    try:
        entrar(sesion, usuario, clave)
    except Rechazo as ex:
        abortar(ex.motivo)
    print(">>> Entrada en la web de OSMA: aceptada (solo fechas"
          + ("" if escribir else ", sin escribir") + ").", flush=True)
    return fechas_de_vuelta(sb, sesion, _ahora(), escribir)


def guardar_copia(sb, contenido):
    """La copia entera del Excel, en la ruta fija (se pisa: solo la última buena)."""
    sb.storage.from_(BUCKET).upload(RUTA_COPIA, contenido, {
        'content-type': 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet', 'upsert': 'true'})


def rescatar(sb, run_id):
    """🔴 Si el run muere a medias (tope de tiempo, cancelacion), su pasada no se queda 'leyendo' para siempre:
    SOLO la de ESTE run, de OSMA, y solo si sigue leyendo."""
    if not run_id.isdigit():
        abortar('sin GITHUB_RUN_ID: no se sabe qué pasada es de este run')
    res = (sb.table('disp_pasada').select('id').eq('run_id', int(run_id)).eq('proveedor', PROVEEDOR)
             .eq('estado', 'leyendo').execute())
    for f in res.data or []:
        _cerrar(sb, f['id'], 'fallida', f'el run {run_id} terminó sin cerrar la pasada (tope de tiempo, cancelación o '
                                         f'caída); mira su log en Actions')
    print(f">>> RESCATE disp_pasada: {len(res.data or [])} pasada(s) de OSMA de este run pasan de 'leyendo' a 'fallida'.")


def principal(argv):
    rescate = argv == ['--rescate']
    fechas = argv[:1] == ['--solo-fechas']
    if argv not in ([], ['--rescate'], ['--solo-fechas'], ['--solo-fechas', '--sin-escribir']):
        abortar('uso: escaner2_osma_disponibilidad.py [--rescate | --solo-fechas [--sin-escribir]]')
    if not rescate and not (os.environ.get('OSMA_USER') and os.environ.get('OSMA_PASS')):
        abortar('faltan OSMA_USER/OSMA_PASS (secretos del repo); no se abre ninguna pasada')
    if not (os.environ.get('SUPABASE_URL') and os.environ.get('SUPABASE_SERVICE_KEY')):
        abortar('faltan SUPABASE_URL/SUPABASE_SERVICE_KEY; no se abre ninguna pasada')
    run_id = os.environ.get('GITHUB_RUN_ID') or ''
    from supabase import create_client
    sb = create_client(os.environ['SUPABASE_URL'], os.environ['SUPABASE_SERVICE_KEY'])
    if rescate:
        try:
            rescatar(sb, run_id)
        except Exception as ex:
            abortar(f'el rescate ha fallado ({type(ex).__name__})')
        return 0
    import requests
    sesion = requests.Session()
    sesion.headers['User-Agent'] = AGENTE
    if fechas:
        return solo_fechas(sb, sesion, os.environ['OSMA_USER'], os.environ['OSMA_PASS'], escribir='--sin-escribir' not in argv)
    return pasada(sb, sesion, os.environ['OSMA_USER'], os.environ['OSMA_PASS'], run_id)


if __name__ == '__main__':
    try:
        codigo = principal(sys.argv[1:])
    except SystemExit:
        raise
    except BaseException as ex:  # noqa: BLE001 — 🔴 nada de trazas: pueden llevar datos de la petición
        print(f"OSMA_NO_EJECUTADA: error inesperado ({type(ex).__name__})")
        sys.exit(1)
    sys.exit(codigo)
