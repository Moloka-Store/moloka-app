# -*- coding: utf-8 -*-
"""Banco del encargo TL (24-oct-2026): la columna «También lo tiene» de «Análisis» en los Excel del escaner 2.

SIN RED, SIN SECRETOS Y SIN BASE: una base de mentira en memoria con la forma que lee escaner2_tambien.py
(table().select().eq().in_().order().range().limit().execute()), filas de disp_estado inventadas para cada proveedor y
el motor de verdad. Los precios y los EAN son INVENTADOS (repo publico).

QUE PRUEBA:
  (A) LA CELDA del ejemplo de Fernando, literal: «TCG 6,95 € oferta · OcioStock 7,20 € · BEMS 7,40 € (10-oct)», en un
      Excel de DBLine: el propio proveedor (DBLine, mas barato que todos) NO sale; del mas barato al mas caro; la fecha
      solo si la foto no es de hoy.
  (B) EL PRECIO ES EL DE SU ESCANEO: TCG el tramo mas barato (`precio_escalon`, no el suelto); OcioStock `precio_pa`;
      OSMA con su porte (facturas + pedido previsto); DBLine la promo vigente (la caducada, Prezzo y sin «oferta»);
      la caja de HEO, entre sus unidades; la caja «5 + 1» de OcioStock, por figura; dos codigos del mismo EAN, el mas
      barato; Zentrada, la foto de su ultimo Escaneo PRO con cruce «lista».
  (C) FUERA: chase suelto, preventa, no disponible (no se piden), foto de mas de 30 dias (31 fuera, 30 dentro: la regla
      de Reponer para BEMS); nadie mas → celda VACIA.
  (D) LA LECTURA CAIDA: toda la columna «sin consultar», el aviso, y el Excel se guarda.
  (E) LA HOJA: la columna es la ULTIMA, detras de «Amazon», dentro de la tabla; las demas celdas no cambian; sin lector
      no hay columna. Solo lee: ni una escritura en la base.
  (F) POR ESTRUCTURA: los siete cruces le pasan el lector a su Excel, los siete Excel y el de novedades ponen la columna,
      y el escaneo de facturas no.
"""
import io
import os
from datetime import date

from openpyxl import Workbook, load_workbook
from openpyxl.worksheet.table import Table, TableStyleInfo

import escaner2_motor as e2
import escaner2_tambien as etl

AQUI = os.path.dirname(os.path.abspath(__file__))
fallos = []


def eq(nombre, obtenido, esperado):
    ok = obtenido == esperado
    if not ok:
        fallos.append(nombre)
    print(('OK ' if ok else 'XX ') + nombre + ('' if ok else '   got=%r exp=%r' % (obtenido, esperado)))


M = e2.cargar_motor()
HOY = date(2026, 10, 24)
HOY_I = '2026-10-24T07:10:00+00:00'
UPC9 = '889698000097'


def ean(n):
    cuerpo = '84355%07d' % n
    return cuerpo + M._chk13(cuerpo)


def d(proveedor, codigo, n, precio, **k):
    """Una fila de disp_estado de mentira: vista hoy, en catalogo, disponible, suelta y sin chase."""
    core = k.pop('core', ean(n))
    f = {'proveedor': proveedor, 'producto_prov': codigo, 'ean_original': core, 'ean_core': core, 'ean_norm': M.norm(core),
         'marca': 'Funko', 'nombre': 'Funko Pop %s' % codigo, 'categoria': None, 'precio_unidad': precio,
         'precio_catalogo': precio, 'precio_escalon': None, 'uds_escalon': None, 'precio_pa': None, 'es_caja': False,
         'uds_caja': None, 'es_chase': False, 'en_oferta': False, 'fin_oferta': None, 'disponible': True,
         'preorder': False, 'fin_de_vida': False, 'regla': None, 'aviso': None, 'en_catalogo': True, 'ausencias': 0,
         'visto_en': HOY_I, 'pasada_id': 'P-' + proveedor}
    f.update(k)
    return f


# ── LA BASE DE MENTIRA ────────────────────────────────────────────────────────────
DISP = [
    # (A) el ejemplo de Fernando, EAN 1: DBLine (el propio) mas barato que todos; TCG en oferta con el tramo; OcioStock
    #     con su precio_pa; BEMS con la foto del 10-oct.
    d('DBLINE', 'DB1', 1, 5.00),
    d('TCG', 'T1', 1, 7.50, precio_escalon=6.95, uds_escalon=12, en_oferta=True),
    d('OCIOSTOCK', 'O1', 1, 8.00, precio_escalon=7.2727, uds_escalon=6, precio_pa=7.20),
    d('BEMS', 'B1', 1, 7.40, visto_en='2026-10-10T14:47:13+00:00'),
    # EAN 2: la caja de HEO (30 € la caja de 6 → 5,00 por unidad) y un chase suelto de HEO mas barato (fuera); una
    #        preventa de TCG (fuera); la caja «5 + 1» de OcioStock, por figura (entra); un chase suelto de OcioStock (fuera).
    d('HEO', 'H2', 2, 5.00, precio_catalogo=30.00, es_caja=True, uds_caja=6),
    d('HEO', 'H2C', 2, 1.00, regla='chase_suelto'),  # asi llegan en la base: es_chase falso y la regla
    d('TCG', 'T2', 2, 2.00, precio_escalon=2.00, preorder=True),
    d('OCIOSTOCK', 'O2', 2, 6.00, precio_pa=5.94, es_caja=True, uds_caja=6, es_chase=True),
    d('OCIOSTOCK', 'O2C', 2, 1.50, precio_pa=1.49, es_chase=True),
    # EAN 3: BEMS con la foto de hace 31 dias (fuera): nadie → vacia. EAN 4: BEMS de hace 30 dias justos (dentro).
    d('BEMS', 'B3', 3, 8.00, visto_en='2026-09-23T14:47:13+00:00'),
    d('BEMS', 'B4', 4, 8.00, visto_en='2026-09-24T14:47:13+00:00'),
    #        y dos de HEO mas baratos que no cuentan: una preventa que HEO da por disponible (en la base hay) y un chase
    #        que no es caja (en la base llegan con regla 'chase_suelto'; sin ella, fuera igual).
    d('HEO', 'H4P', 4, 1.00, preorder=True),
    d('HEO', 'H4C', 4, 1.50, es_chase=True),
    # EAN 5: OSMA en SALE, 10,00 sin porte → 10,50 con el 5 % (50 € de porte de la ultima factura ÷ 1.000 de pedido);
    #        DBLine con dos codigos (el mas barato gana) y la promo vigente.
    d('OSMA', '9001', 5, 10.00, en_oferta=True),
    d('DBLINE', 'DB5a', 5, 9.00, precio_catalogo=11.00, en_oferta=True, fin_oferta='2026-10-31'),
    d('DBLINE', 'DB5b', 5, 12.00),
    # EAN 6: DBLine con la promo que ya acabo (la foto la dio por buena): vale Prezzo, sin «oferta».
    d('DBLINE', 'DB6', 6, 4.00, precio_catalogo=6.00, en_oferta=True, fin_oferta='2026-10-20'),
    # EAN 8: un HEO agotado (no se pide de la base: `disponible` = true en la consulta; si llegara, fuera igual).
    d('HEO', 'H8', 8, 3.00, disponible=False),
    # EAN 9: OcioStock lo da en UPC-12 y la hoja con el 0 delante (EAN-13): casa por el EAN normalizado.
    d('OCIOSTOCK', 'O9', 9, 4.00, precio_pa=3.96, core=UPC9),
]
ZENTRADA_PASADAS = [{'id': 'Z-nueva', 'creada_en': '2026-10-20T10:00:00+00:00'},
                    {'id': 'Z-vieja', 'creada_en': '2026-10-11T16:10:57+00:00'}]
ZENTRADA_CRUCES = [{'pasada_id': 'Z-vieja', 'estado': 'lista'}, {'pasada_id': 'Z-nueva', 'estado': 'fallida'}]
ZENTRADA_FOTO = [{'id': 'zf1', 'pasada_id': 'Z-vieja', 'ean_core': ean(6), 'ean_norm': M.norm(ean(6)), 'precio_unidad': 3.99,
                  'es_chase': False, 'preorder': False, 'en_oferta': False},
                 {'id': 'zf2', 'pasada_id': 'Z-nueva', 'ean_core': ean(7), 'ean_norm': M.norm(ean(7)), 'precio_unidad': 1.00,
                  'es_chase': False, 'preorder': False, 'en_oferta': False}]
TABLAS = {
    'disp_estado': DISP,
    'disp_parametros': [{'proveedor': 'OSMA', 'pedido_previsto_eur': 1000}],
    'facturas': [{'id': 1, 'proveedor': 'OSMA', 'fecha': '2026-09-01', 'created_at': '2026-09-01', 'gastos_envio': 80},
                 {'id': 2, 'proveedor': 'OSMA', 'fecha': '2026-10-01', 'created_at': '2026-10-01', 'gastos_envio': 50}],
    'codigos_proveedor': [], 'productos': [],
    'escaner2_pasada': [dict(p, proveedor='ZENTRADA') for p in ZENTRADA_PASADAS],
    'escaner2_cruce': ZENTRADA_CRUCES,
    'escaner2_foto': ZENTRADA_FOTO,
}


class SbFalso:
    """Solo lo que lee escaner2_tambien: select con eq / in_ / order / range / limit. Cualquier escritura se anota."""

    def __init__(self, tablas, rompe=False):
        self.tablas, self.rompe, self.pedidas, self.escrituras = tablas, rompe, [], []

    def table(self, nombre):
        if self.rompe:
            raise RuntimeError("Could not find the table 'public.disp_estado'")
        return _Q(self, nombre)

    def rpc(self, *a, **k):
        self.escrituras.append(('rpc', a))
        raise RuntimeError('rpc')


class _Q:
    def __init__(self, sb, tabla):
        self.sb, self.tabla, self.filtros, self.rango, self.tope = sb, tabla, [], None, None

    def select(self, cols):
        self.cols = cols
        return self

    def eq(self, k, v):
        self.filtros.append((k, lambda x, v=v: str(x) == str(v)))
        return self

    def in_(self, k, vs):
        vs = {str(x) for x in vs}
        self.filtros.append((k, lambda x, vs=vs: str(x) in vs))
        return self

    def order(self, *_a, **_k):
        return self

    def range(self, a, b):
        self.rango = (a, b)
        return self

    def limit(self, n):
        self.tope = n
        return self

    def _escribe(self, *a, **k):
        self.sb.escrituras.append((self.tabla, a))
        return self

    insert = update = upsert = delete = _escribe

    def execute(self):
        self.sb.pedidas.append((self.tabla, tuple(k for k, _ in self.filtros)))
        filas = [dict(f) for f in self.sb.tablas.get(self.tabla, [])
                 if all(k in f and ok(f[k]) for k, ok in self.filtros)]
        if self.rango:
            filas = filas[self.rango[0]:self.rango[1] + 1]
        if self.tope is not None:
            filas = filas[:self.tope]
        return type('R', (), {'data': filas})()


def hoja(eans, con_amazon=True):
    """Un libro con «Análisis» como el del escaner 2: cabecera con «EAN» y «Amazon» al final, y la tabla T_Analisis."""
    wb = Workbook()
    ws = wb.active
    ws.title = 'Análisis'
    cab = ['Nombre', 'EAN', 'ASIN', 'PA (€)', 'Decisión'] + (['Amazon'] if con_amazon else [])
    ws.append(cab)
    for i, x in enumerate(eans):
        ws.append(['Funko %d' % i, x, 'B0TEST%04d' % i, 9.99, 'COMPRAR'] + (['sin consultar'] if con_amazon else []))
    t = Table(displayName='T_Analisis', ref='A1:%s%d' % ('F' if con_amazon else 'E', len(eans) + 1))
    t.tableStyleInfo = TableStyleInfo(name='TableStyleMedium2', showRowStripes=True)
    ws.add_table(t)
    return wb


def celdas(wb):
    ws = wb['Análisis']
    cab = [c.value for c in ws[1]]
    return cab, [dict(zip(cab, r)) for r in ws.iter_rows(min_row=2, values_only=True)]


def guardar_y_abrir(wb):
    buf = io.BytesIO()
    wb.save(buf)
    return load_workbook(io.BytesIO(buf.getvalue()))


# ═══════════════════════════════════════════════════════════════════════════════
print('(A) la celda del ejemplo de Fernando, en un Excel de DBLine')
EANS = [ean(n) for n in range(1, 9)] + ['0' + UPC9]
sb = SbFalso(TABLAS)
avisos = []
wb = hoja(EANS)
_antes = celdas(guardar_y_abrir(hoja(EANS)))
n = e2.poner_columna_tambien(wb, etl.lector_tambien(sb, 'DBLINE', M, avisos, hoy=HOY))
wb = guardar_y_abrir(wb)
cab, filas = celdas(wb)
T = {f['EAN']: f['También lo tiene'] for f in filas}
eq('(A) 🔴 el ejemplo, literal: el propio (DBLine, 5,00) no sale; del más barato al más caro; fecha solo si no es de hoy',
   T[ean(1)], 'TCG 6,95 € oferta · OcioStock 7,20 € · BEMS 7,40 € (10-oct)')

print('\n(B) el precio es el de su escaneo')
eq('(B) TCG con el tramo más barato (6,95, no el suelto 7,50) y OcioStock con precio_pa (7,20, no 8,00)',
   ('TCG 6,95' in T[ean(1)], 'OcioStock 7,20' in T[ean(1)], '7,50' in T[ean(1)], '8,00' in T[ean(1)]), (True, True, False, False))
eq('(B) 🔴 la caja de HEO, entre sus unidades (30 ÷ 6 = 5,00); la «5 + 1» de OcioStock, por figura; el chase suelto '
   '(HEO 1,00 y OcioStock 1,49) y la preventa de TCG (2,00), fuera', T[ean(2)], 'HEO 5,00 € · OcioStock 5,94 €')
eq('(B) OSMA con su porte (10,00 × 1,05 = 10,50, la última factura con porte) y «oferta» (SALE); el propio DBLine no',
   T[ean(5)], 'OSMA 10,50 € oferta')
eq('(B) Zentrada: la foto del último escaneo con cruce «lista» (la del 11-oct), con su fecha', T[ean(6)], 'Zentrada 3,99 € (11-oct)')
eq('(B) y la del escaneo cuyo cruce falló no cuenta', T[ean(7)], None)
eq('(B) UPC/EAN: casa por el EAN normalizado', T['0' + UPC9], 'OcioStock 3,96 €')

print('\n(C) fuera de la celda')
eq('(C) 🔴 foto de BEMS de hace 31 días: fuera, y sin nadie más la celda VACÍA', T[ean(3)], None)
eq('(C) foto de BEMS de hace 30 días justos: dentro (la regla de Reponer es «más de 30»); la preventa y el chase suelto '
   'de HEO, fuera', T[ean(4)], 'BEMS 8,00 € (24-sep)')
eq('(C) un agotado no se pide a la base, y sin nadie más la celda vacía', T[ean(8)], None)
eq('(C) se cuentan solo las filas con algo', n, sum(1 for v in T.values() if v))

print('\n(B bis) en un Excel de TCG: DBLine sale, con su promo vigente y su Prezzo cuando la promo ya acabó')
wb2 = hoja(EANS)
e2.poner_columna_tambien(wb2, etl.lector_tambien(SbFalso(TABLAS), 'TCG', M, [], hoy=HOY))
T2 = {f['EAN']: f['También lo tiene'] for f in celdas(wb2)[1]}
eq('(B) el propio (TCG) no sale; DBLine 5,00 el más barato', T2[ean(1)], 'DBLine 5,00 € · OcioStock 7,20 € · BEMS 7,40 € (10-oct)')
eq('(B) 🔴 DBLine: dos códigos del mismo EAN, el más barato (9,00 de la promo vigente, «oferta»), y OSMA detrás',
   T2[ean(5)], 'DBLine 9,00 € oferta · OSMA 10,50 € oferta')
eq('(B) 🔴 DBLine con la promo acabada: Prezzo (6,00) y sin «oferta»; Zentrada más barata delante',
   T2[ean(6)], 'Zentrada 3,99 € (11-oct) · DBLine 6,00 €')
wb3 = hoja(EANS)
e2.poner_columna_tambien(wb3, etl.lector_tambien(SbFalso(TABLAS), 'ZENTRADA', M, [], hoy=HOY))
eq('(B) en un Excel de Zentrada, Zentrada no sale', {f['EAN']: f['También lo tiene'] for f in celdas(wb3)[1]}[ean(6)],
   'DBLine 6,00 €')

print('\n(D) la lectura caída')
avisos_c = []
wbc = hoja(EANS)
e2.poner_columna_tambien(wbc, etl.lector_tambien(SbFalso(TABLAS, rompe=True), 'DBLINE', M, avisos_c, hoy=HOY))
wbc = guardar_y_abrir(wbc)
eq('(D) 🔴 toda la columna «sin consultar», el aviso con el porqué, y el Excel se guarda y se abre',
   ({f['También lo tiene'] for f in celdas(wbc)[1]}, len(avisos_c), 'Columna «También lo tiene»' in (avisos_c or [''])[0]),
   ({'sin consultar'}, 1, True))
eq('(D) y con la lectura buena, sin aviso', avisos, [])

print('\n(E) la hoja')
eq('(E) 🔴 «También lo tiene» es la ÚLTIMA columna, detrás de «Amazon»', cab[-2:], ['Amazon', 'También lo tiene'])
_ws = wb['Análisis']
eq('(E) la tabla T_Analisis la cubre (y su columna nueva)',
   [(t.displayName, t.ref, t.tableColumns[-1].name if t.tableColumns else None) for t in _ws.tables.values()],
   [('T_Analisis', 'A1:G10', 'También lo tiene')])
eq('(E) 🔴 las columnas de antes, idénticas (mismas filas, mismos valores)',
   ([c for c in cab if c != 'También lo tiene'], [{k: v for k, v in f.items() if k != 'También lo tiene'} for f in filas]),
   _antes)
eq('(E) sin lector, no hay columna: el libro es el de antes',
   (e2.poner_columna_tambien(hoja(EANS), None), celdas(hoja(EANS))[0][-1]), (0, 'Amazon'))
eq('(E) 🔒 solo lee: ni una escritura en la base', sb.escrituras, [])
eq('(E) lee disp_estado vistas, en catálogo y disponibles, de los OTROS proveedores',
   sorted({p for p in sb.pedidas if p[0] == 'disp_estado'}),
   [('disp_estado', ('proveedor', 'ean_norm', 'ausencias', 'en_catalogo', 'disponible'))])
_prov = []
_orig = etl.leer_disp
etl.leer_disp = lambda s, provs, es: (_prov.append(tuple(provs)) or _orig(s, provs, es))
etl.lector_tambien(SbFalso(TABLAS), 'DBLINE', M, [], hoy=HOY)(EANS)
etl.leer_disp = _orig
eq('(E) 🔴 el propio proveedor ni se pide', _prov, [('BEMS', 'HEO', 'OCIOSTOCK', 'OSMA', 'TCG')])

print('\n(G) el texto, puro')
O = lambda p, x, of=False, dia=HOY: {'proveedor': p, 'precio': x, 'oferta': of, 'dia': dia}
eq('(G) orden por precio y, a igual precio, por nombre; dos decimales con coma y redondeo comercial',
   etl.texto_tambien([O('TCG', 7.0), O('BEMS', 7.0), O('HEO', 6.945, True, date(2026, 1, 3))], HOY),
   'HEO 6,95 € oferta (3-ene) · BEMS 7,00 € · TCG 7,00 €')
eq('(G) nadie: None', (etl.texto_tambien([], HOY), etl.texto_tambien(None, HOY)), (None, None))
eq('(G) la regla de los 30 días: el día UTC contra hoy de Madrid',
   [etl.caducada('2026-09-24T23:30:00+00:00', HOY), etl.caducada('2026-09-23T23:30:00+00:00', HOY), etl.caducada(None, HOY)],
   [False, True, False])

print('\n(F) por estructura')
for _f in ('escaner2_heo_cruce.py', 'escaner2_bems_cruce.py', 'escaner2_dbline_cruce.py', 'escaner2_ociostock_cruce.py',
           'escaner2_osma_cruce.py', 'escaner2_tcg_cruce.py', 'escaner2_zentrada_cruce.py'):
    _t = io.open(os.path.join(AQUI, _f), encoding='utf-8').read()
    eq('(F) %s pasa el lector a su Excel, con su proveedor' % _f, _t.count("'tambien': etl.lector_tambien(sb, "), 1)
for _f in ('escaner2_heo_cruce.py', 'escaner2_bems_pro.py', 'escaner2_dbline_pro.py', 'escaner2_ociostock_pro.py',
           'escaner2_osma.py', 'escaner2_tcg_pro.py', 'escaner2_zentrada.py'):
    _t = io.open(os.path.join(AQUI, _f), encoding='utf-8').read()
    _i, _j = _t.find("e2.poner_columna_amazon(wb, info.get('restricciones'))"), _t.find("e2.poner_columna_tambien(wb, info.get('tambien'))")
    eq('(F) %s pone «También lo tiene» una vez y detrás de «Amazon»' % _f,
       (_t.count("e2.poner_columna_tambien(wb, info.get('tambien'))"), 0 <= _i < _j), (1, True))
_t = io.open(os.path.join(AQUI, 'escaner2_novedades.py'), encoding='utf-8').read()
eq('(F) el Excel de novedades pone la columna', _t.count("e2.poner_columna_tambien(wb, etl.lector_tambien(sb, proveedor, "), 1)
eq('(F) 🔴 el escaneo de facturas, NO («no se mira»)',
   'poner_columna_tambien' in io.open(os.path.join(AQUI, 'escaner2_factura.py'), encoding='utf-8').read(), False)

print()
if fallos:
    print('FALLAN %d:' % len(fallos))
    for f in fallos:
        print('  - ' + f)
    raise SystemExit(1)
print('TODO OK')
