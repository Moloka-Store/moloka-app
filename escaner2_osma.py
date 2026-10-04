# -*- coding: utf-8 -*-
"""ESCANER 2 · ESCANEO PRO DE OSMA · LA PARTE PURA (encargo AG, 02-oct-2026; AG2, mismo dia: solo España y una lista)

QUE ES. El calculo del Escaneo PRO de OSMA, sin red y sin base: la foto desde la descarga diaria, la lista
para el Visualizador, el cruce de cada articulo con sus fichas de Keepa y el Excel. La red vive en los
dos programas que lo usan:
  · escaner2_osma_barrido.py (workflow escaner2-osma-barrido.yml): de la descarga a la foto y la lista.
  · escaner2_osma_cruce.py   (workflow escaner2-osma-cruce.yml):   los CSV del Visualizador, las puertas y el Excel.

🔑 CADA ESCANER, A LA MEDIDA DE SU PROVEEDOR (Fernando, 02-oct-2026): nada de tronco comun con adaptadores. Leer
   el catalogo, filtrar y cruzar es codigo de OSMA y vive aqui. Lo que se REUTILIZA del escaner 2 de HEO, sin
   tocarlo, son las REGLAS DE COMPRA (las mismas que aplica hoy su PRO) y el formato del Excel:
     · `escaner2_motor.decidir` (las seis puertas, el corte de ventas de escaner2_parametros, el calculo en los
       paises de escaner2_parametros venda o no en cada uno, la eleccion de ficha del viejo y la FICHA COMPARTIDA
       del encargo AA; en OSMA, desde el AG2, los paises son solo ES),
       que por debajo usa `calc_rentabilidad` y `decision_de` del viejo, copiados en escaner2_heredado_nube.py;
     · `escaner2_motor.examinar_csv`, `fichas_compartidas` y el lector del CSV del Visualizador
       (escaner2_heredado_pro.py);
     · la Celda 9 del viejo (`escaner2_motor.datos_como_el_viejo` + `escribir_celda9`) y la columna «Ficha
       compartida» (`poner_columna_ficha_compartida`).

LO PROPIO DE OSMA:
  1. LA FOTO SALE DE LA DESCARGA DIARIA (disp_estado, filas vistas en la ultima pasada aplicada), no de la web.
     Filtro de Fernando (01-oct-2026, «las cosas de marca»): con marca y disponible. Los descatalogados con
     existencias ENTRAN, con la marca «no habrá más» en el Excel.
  2. EL PRECIO LLEVA EL PORTE, UNA SOLA VEZ: el de la puerta comun (v_escaner_fuente, encargo AF): Price_net ×
     (1 + porte de la ultima factura de OSMA ÷ pedido previsto), redondeado a 2 decimales DESPUES. La puerta comun
     solo trae las ~30 filas de OSMA que casan con una ficha; para las demas, la MISMA cuenta (`pa_con_porte`), y
     el barrido comprueba que en las filas que estan en los dos sitios el precio es EL MISMO (si no, falla).
  3. UNA SOLA LISTA Y UN SOLO CSV, EL DE ESPAÑA (encargo AG2, Fernando, 02-oct-2026: «Las cosas de Osma es
     dificil venderlas en Pan EU»). La lista lleva los EAN de OSMA del filtro + el EAN de NUESTRA ficha de cada
     codigo enlazado que esta en la foto, sin repetir (`lista_para_keepa_osma`). La ficha cuyo EAN no conoce Keepa
     NO se rescata (el Lenor 18459, medido el 02-oct-2026): se dice en el Resumen, calculado en cada pasada con
     keepa_escaparate (`fuera_de_keepa`). (Mapa de OSMA, 04-oct-2026) Y una segunda lista, SOLO de ASIN: la del mapa
     fijo (punto 10).
  4. EL CRUCE POR CODIGO: el articulo de OSMA de un codigo enlazado (codigos_proveedor, lista CERRADA de 20) se
     decide con SUS fichas (las de su familia que traiga el CSV), no con lo que diga el EAN nuevo de OSMA. Una fila
     del CSV con el EAN de una ficha enlazada es de la fila de ESE codigo; la misma ficha por los dos caminos: una,
     la del codigo. Una fila de Keepa de una ficha enlazada no se cuelga de OTRO articulo de OSMA aunque comparta EAN.
  5. NUESTROS PACKS: coste = unidades × precio (como Reponer desde el encargo AC; el factor, como lib/packs de la
     v2, portado aqui: `factores_por_ficha`).
  6. UN SOLO AÑADIDO EN EL EXCEL: la columna «No habrá más» al final de «Análisis» (descatalogados con existencias).
     Descartado por Fernando (02-oct-2026): ni la marca «solo si va en el palé» ni la linea de control final.
  7. (AM, 02-oct-2026) LOS PACKS DE AMAZON: si el ASIN de una fila que se vende es un multipack de la unidad de OSMA,
     el coste es N × precio, con la regla de `factor_pack_amazon` (dos señales que dicen lo mismo); si las señales no
     bastan o se contradicen, «posible pack en Amazon: revisar» y nunca COMPRAR. Va en «Coherencia caja».
  8. (AM) EL PEDIDO PREVISTO Y EL PORTE, VIVOS EN EL EXCEL: dos celdas amarillas en «Resumen», «PA (€)» y «Decisión»
     de «Análisis» como formulas que apuntan a ellas, y la columna «Precio OSMA sin porte (€)» al final. La base no se
     toca (disp_parametros, la puerta comun y Reponer siguen con su pedido).
  9. (AM) «Análisis» pinta solo los paises que se calculan cuando es uno (OSMA: ES); HEO sigue con sus cuatro.
 10. (Mapa de OSMA, Fernando, 04-oct-2026) EL MAPA FIJO `osma_mapa_ficha` (codigo OSMA → ASIN → unidades), hecho a
     mano: sus ASIN van en una lista aparte (modo «ASIN» del Visualizador), cada ficha se valora con SU coste
     (unidades × precio × (1 + porte)) y sale en el Excel con su marca al final; lo nuestro agotado en OSMA, en su
     hoja. Ver la seccion 4 ter. Sustituye a la busqueda automatica de packs del encargo AN, retirada.
"""
import io
import re
from decimal import Decimal, ROUND_HALF_UP

import escaner2_motor as e2

PROVEEDOR = 'OSMA'
CARPETA = 'osma'
# El modo de la pasada (check de escaner2_pasada.modo): todas las marcas, solo lo disponible. Para OSMA, «marca
# fuera» es «sin marca».
MODO = 'todas'
# El Visualizador no admite mas de 10.000 codigos por lista; por encima de 5.000 se avisa (encargo AG).
TOPE_LISTA = e2.TOPE_VISUALIZADOR
AVISO_LISTA = 5000

# Los nombres de las puertas previas EN OSMA. Las mismas ocho del escaner 2 (el cuadre crudo = previas + foto es
# el de siempre); en OSMA solo se usan cinco y las otras tres salen a 0.
NOMBRE_PREVIA = dict(e2.NOMBRE_PUERTA_PREVIA,
                     no_disponible='Sin existencias en OSMA',
                     marca_fuera='Sin marca',
                     sin_gtin='Sin EAN',
                     ean_forma_rara='EAN con forma rara o de relleno',
                     duplicado_proveedor='Mismo EAN que otro artículo de OSMA')

COLUMNA_NO_HABRA_MAS = 'No habrá más'
TEXTO_NO_HABRA_MAS = 'no habrá más'
ANCHO_NO_HABRA_MAS = 14


class FalloOsma(Exception):
    """Algo que impide seguir: la pasada o el cruce quedan 'fallida' con este motivo."""


# ═══════════════════════════════════════════════════════════════════════════════
# 1 · EL PRECIO CON EL PORTE: LA CUENTA DE LA PUERTA COMUN
# ═══════════════════════════════════════════════════════════════════════════════
def _dec(x):
    return None if x is None else Decimal(str(x))


def _desc_nulos_primero(v):
    """Una clave de orden DESCENDENTE con los NULL delante, como `ORDER BY … DESC` de Postgres."""
    return (v is None, '' if v is None else str(v))


def porte_de(facturas, pedido_previsto):
    """El porte de OSMA con la MISMA regla que v_escaner_fuente (migracion 20261001154500, encargo AF):
    `gastos_envio` de la factura de OSMA mas reciente que lo tenga (`fecha DESC, created_at DESC, id DESC`, con
    los NULL delante como en Postgres) ÷ `disp_parametros.pedido_previsto_eur`. Sin pedido previsto o sin factura
    con porte, `pct` es None y el precio va sin porte, como en la puerta comun.

    `facturas`: filas de `facturas` de OSMA con `id, fecha, created_at, gastos_envio`."""
    pedido = _dec(pedido_previsto)
    con = [f for f in facturas or [] if f.get('gastos_envio') is not None]
    if pedido is None or pedido <= 0 or not con:
        return {'pct': None, 'factura': None, 'fecha': None, 'gastos_envio': None,
                'pedido_previsto': None if pedido is None else float(pedido)}
    ultima = sorted(con, key=lambda f: (_desc_nulos_primero(f.get('fecha')), _desc_nulos_primero(f.get('created_at')),
                                        _desc_nulos_primero(f.get('id'))), reverse=True)[0]
    gastos = _dec(ultima['gastos_envio'])
    return {'pct': gastos / pedido, 'factura': ultima.get('id'), 'fecha': ultima.get('fecha'),
            'gastos_envio': float(gastos), 'pedido_previsto': float(pedido)}


def pa_con_porte(precio, pct):
    """El precio de compra de UNA unidad de OSMA con el porte, como la puerta comun: `round(precio × (1 + pct), 2)`
    (sin porte, `precio` a 2 decimales: la vista lo saca como numeric(10,2)). Redondeo de Postgres (la mitad, lejos
    del cero), con Decimal: nada de coma flotante en medio."""
    if precio is None:
        return None
    p = _dec(precio)
    v = p if pct is None else p * (1 + pct)
    return float(v.quantize(Decimal('0.01'), rounding=ROUND_HALF_UP))


def texto_pct(pct):
    return '—' if pct is None else ('%.2f %%' % (float(pct) * 100)).replace('.', ',')


# ═══════════════════════════════════════════════════════════════════════════════
# 2 · LOS ENLACES POR CODIGO Y NUESTROS PACKS
# ═══════════════════════════════════════════════════════════════════════════════
def ean_norm(v):
    """`moloka_ean_norm` de la base y `eanNorm` de la v2: solo cifras y sin ceros delante; vacio → None."""
    s = re.sub(r'[^0-9]', '', str(v or '')).lstrip('0')
    return s or None


def _clave_codigo(c):
    s = str(c or '')
    return (0, int(s), s) if s.isdigit() else (1, 0, s)


def _ficha_viva(p):
    """Una ficha que se vende en Amazon por ASIN: activa, no chase (el chase NUNCA lleva ASIN) y con ASIN."""
    return bool(p.get('activo')) and not p.get('es_chase') and bool(p.get('asin'))


def enlaces_de(codigos, productos):
    """{codigo: {'producto_id', 'ean_ficha', 'asins'}} de los enlaces de OSMA (`codigos_proveedor`), con las fichas
    de la MISMA familia que la puerta comun: la ficha del codigo y todas las fichas vivas con su EAN normalizado
    (asi el 18459 trae el Lenor suelto y el Pack 2). Y los avisos de lo que no se puede usar (un codigo sin ficha,
    o sin ninguna ficha viva). Puro.

    `codigos`: filas de codigos_proveedor (proveedor OSMA) con `codigo_proveedor, producto_id`;
    `productos`: filas con `id, asin, ean, activo, es_chase`."""
    por_id = {p['id']: p for p in productos or []}
    familia = {}
    for p in productos or []:
        k = ean_norm(p.get('ean'))
        if k and _ficha_viva(p):
            familia.setdefault(k, []).append(p)
    enlaces, avisos = {}, []
    for c in sorted(codigos or [], key=lambda x: _clave_codigo(x.get('codigo_proveedor'))):
        codigo = str(c.get('codigo_proveedor') or '').strip()
        ficha = por_id.get(c.get('producto_id'))
        if not codigo:
            continue
        if ficha is None:
            avisos.append('El código %s de OSMA apunta a una ficha que no está en productos' % codigo)
            continue
        k = ean_norm(ficha.get('ean'))
        fichas = list(familia.get(k, [])) if k else ([ficha] if _ficha_viva(ficha) else [])
        asins = sorted({p['asin'] for p in fichas})
        if not asins:
            avisos.append('El código %s de OSMA no tiene ninguna ficha activa con ASIN' % codigo)
            continue
        enlaces[codigo] = {'producto_id': ficha['id'], 'ean_ficha': ficha.get('ean'), 'asins': asins}
    return enlaces, avisos


# ── EL FACTOR DE UN PACK: lib/packs/factores.ts y lib/packs/grupos.ts de la v2, PORTADOS ─────────────────────
# 🔑 La MISMA regla que Reponer (encargo AC) y la casilla «Unidades por pack» (encargo AD). Si se cambia alli, se
#    cambia aqui: test_escaner2_osma.py coteja los tres grupos reales que mide la v2 (Ultra Pro x1-x4, Protefix
#    x1-x2-x5, Lenor x1-x2) y la casilla. 🔴 El marcador del nombre SOLO decide dentro de un grupo de fichas que
#    comparten EAN: fuera de uno, factor 1 (los Funko «2-PACK», el Kleenex «Maxi pack 128»).
#    re.ASCII: `\b` y `\d` como en JavaScript.
_RE_PACK = re.compile(r'\bpack\s*(?:de\s*)?(\d{1,4})\b', re.I | re.A)
_RE_PACK_INV = re.compile(r'\b(\d{1,4})\s*-?\s*pack\b', re.I | re.A)


def unidades_por_pack_de(v):
    """`unidadesPorPackDe`: la casilla vale solo si es un entero > 1."""
    try:
        if isinstance(v, bool):
            return None
        n = float(v) if isinstance(v, (int, float)) else (float(v) if isinstance(v, str) and v.strip() else float('nan'))
    except ValueError:
        return None
    return int(n) if (n == n and n.is_integer() and n > 1) else None


def tamano_pack_por_marcador(nombre):
    s = str(nombre or '')
    m = _RE_PACK.search(s) or _RE_PACK_INV.search(s)
    if not m:
        return None
    n = int(m.group(1))
    return n if n > 0 else None


def _esqueleto(nombre):
    nums = []

    def poner(m):
        nums.append(int(m.group(0)))
        return '#'
    return re.sub(r'\d+', poner, str(nombre or ''), flags=re.A), nums


def tamanos_por_plantilla(nombres):
    if len(nombres) < 2:
        return None
    partes = [_esqueleto(n) for n in nombres]
    if any(p[0] != partes[0][0] for p in partes):
        return None
    n = len(partes[0][1])
    if n == 0 or any(len(p[1]) != n for p in partes):
        return None
    variables = [i for i in range(n) if any(p[1][i] != partes[0][1][i] for p in partes)]
    if len(variables) != 1:
        return None
    tam = [p[1][variables[0]] for p in partes]
    return tam if all(t > 0 for t in tam) else None


def _normalizar(tam):
    m = min(tam)
    if not m > 0:
        return None
    f = [t / m for t in tam]
    return [int(x) for x in f] if all(float(x).is_integer() and x >= 1 for x in f) else None


def factores_de_grupo(nombres):
    if len(nombres) < 2:
        return None
    por_marcador = [tamano_pack_por_marcador(n) or 1 for n in nombres]
    if any(t > 1 for t in por_marcador) and len(set(por_marcador)) > 1:
        f = _normalizar(por_marcador)
        if f:
            return f
    por_plantilla = tamanos_por_plantilla(nombres)
    if por_plantilla and len(set(por_plantilla)) > 1:
        f = _normalizar(por_plantilla)
        if f:
            return f
    return None


def factores_por_ficha(fichas):
    """`factoresPorFicha`: {id: factor} SOLO de las fichas con factor deducido (las de factor 1 de un grupo con
    factores tambien entran); quien lo lee usa `.get(id, 1)`. La casilla de la ficha manda sobre todo."""
    out = {}
    por_ean = {}
    for f in fichas or []:
        if f.get('es_chase'):
            continue
        k = ean_norm(f.get('ean'))
        if k:
            por_ean.setdefault(k, []).append(f)
    for lista in por_ean.values():
        if len(lista) < 2:
            continue
        activas = [f for f in lista if f.get('activo') is not False]
        if len(activas) >= 2:
            fx = factores_de_grupo([f.get('nombre') for f in activas])
            if fx:
                for f, x in zip(activas, fx):
                    out[f['id']] = x
        for f in lista:
            if f.get('activo') is not False:
                continue
            grupo = activas + [f]
            if len(grupo) < 2:
                continue
            fx = factores_de_grupo([g.get('nombre') for g in grupo])
            if fx:
                out[f['id']] = fx[-1]
    for f in fichas or []:
        if f.get('es_chase'):
            continue
        n = unidades_por_pack_de(f.get('unidades_por_pack'))
        if n is not None:
            out[f['id']] = n
    return out


def factor_por_asin(productos):
    """{asin: factor > 1} de nuestras fichas vivas que son pack, y los avisos (un ASIN en dos fichas vivas con
    factores distintos no se adivina: factor 1 y se dice)."""
    por_id = factores_por_ficha(productos)
    vistos, avisos = {}, []
    for p in productos or []:
        if not _ficha_viva(p):
            continue
        vistos.setdefault(p['asin'], set()).add(por_id.get(p['id'], 1))
    salida = {}
    for asin, fs in vistos.items():
        if len(fs) > 1:
            avisos.append('El ASIN %s está en dos fichas con factores de pack distintos (%s): se valora como 1'
                          % (asin, ', '.join(str(x) for x in sorted(fs))))
        elif next(iter(fs)) > 1:
            salida[asin] = next(iter(fs))
    return salida, avisos


# ═══════════════════════════════════════════════════════════════════════════════
# 3 · LA FOTO, DESDE LA DESCARGA DIARIA
# ═══════════════════════════════════════════════════════════════════════════════
def ean_valido(core, M):
    """Un EAN-13 (o UPC-12) con su digito de control bueno (`_ean_ok` del viejo). 🔑 Regla de OSMA, medida el
    02-oct-2026: 76 articulos con marca y existencias traen EAN de RELLENO (1111111111111 ×69, 9999999999999 ×7,
    de displays y cajas surtidas) y 4 con el control mal; ninguno identifica un producto en Amazon."""
    c = str(core or '')
    if not c.isdigit():
        return False
    if len(c) == 13:
        return M._ean_ok(c)
    if len(c) == 12:
        return M._ean_ok('0' + c)
    return False


def _num(x):
    return None if x is None else float(x)


def _apartado(f, motivo, detalle):
    return {'producto_heo': str(f.get('producto_prov') or ''), 'ean_original': str(f.get('ean_original') or ''),
            'nombre': f.get('nombre') or '', 'marca': (f.get('marca') or '').strip(),
            'precio_catalogo': _num(f.get('precio_unidad')), 'motivo': motivo, 'detalle': detalle,
            'en_oferta': bool(f.get('en_oferta')) if motivo == 'marca_fuera' else None}


def construir_foto(filas, enlaces, M, pct):
    """De la descarga de OSMA a la foto de la pasada, sin perder a nadie: cada articulo sale por UNA puerta previa
    o entra en la foto, y crudo = previas + foto. En este orden:
      1. no disponible (`disponible` de la descarga: stock > 0 o mas de 10.000) → `no_disponible` (se cuenta);
      2. sin marca → `marca_fuera` (se lista: «Sin marca»);
      3. sin EAN → `sin_gtin` (se cuenta);
      4. EAN con forma rara (`regla` de la descarga) o con el digito de control mal → `ean_forma_rara` (se lista);
      5. dos articulos con el mismo EAN → uno: el de un codigo ENLAZADO si lo hay y si no el mas barato (a igual
         precio, el de codigo menor); el otro → `duplicado_proveedor` (se lista).
    `filas`: disp_estado de OSMA vistas en la ultima pasada aplicada; `enlaces`: los de `enlaces_de`; `pct`: el
    porte de `porte_de`. Devuelve (foto, apartados, cuentas)."""
    previas = {p: 0 for p in e2.PUERTAS_PREVIAS}
    apartados, sirven = [], []
    for f in sorted(filas or [], key=lambda x: _clave_codigo(x.get('producto_prov'))):
        if not f.get('disponible'):
            previas['no_disponible'] += 1
            continue
        if not (f.get('marca') or '').strip():
            apartados.append(_apartado(f, 'marca_fuera', 'Sin marca en el catálogo de OSMA'))
            continue
        if not str(f.get('ean_original') or '').strip():
            previas['sin_gtin'] += 1
            continue
        core = str(f.get('ean_core') or '').strip()
        if f.get('regla') or not core:
            apartados.append(_apartado(f, 'ean_forma_rara', 'EAN con forma rara (%s)' % (f.get('ean_original') or '')))
            continue
        if not ean_valido(core, M):
            apartados.append(_apartado(f, 'ean_forma_rara',
                                       'EAN %s con el dígito de control mal: de relleno o mal escrito' % core))
            continue
        sirven.append(f)

    uno = {}
    for f in sirven:
        k = M.norm(f['ean_core'])
        prev = uno.get(k)
        if prev is None:
            uno[k] = f
            continue
        enl_f, enl_p = str(f['producto_prov']) in enlaces, str(prev['producto_prov']) in enlaces
        if enl_f != enl_p:
            gana, pierde = (f, prev) if enl_f else (prev, f)
            porque = 'el código %s está enlazado a nuestra ficha' % gana['producto_prov']
        else:
            pf, pp = f.get('precio_unidad'), prev.get('precio_unidad')
            gana, pierde = (f, prev) if (pf is not None and (pp is None or pf < pp)) else (prev, f)
            porque = 'el más barato'
        uno[k] = gana
        apartados.append(_apartado(pierde, 'duplicado_proveedor', 'Mismo EAN que el código %s: me quedo con %s (%s)'
                                   % (gana['producto_prov'], gana['producto_prov'], porque)))

    foto = []
    for f in sorted(uno.values(), key=lambda x: _clave_codigo(x.get('producto_prov'))):
        core = f['ean_core'].strip()
        foto.append({
            'producto_heo': str(f['producto_prov']), 'ean_original': str(f['ean_original']).strip(), 'ean_core': core,
            'variantes': sorted({M.norm(v) for v in M.variantes_ean(core)} - {''}),
            'codigos_keepa': e2.codigos_para_keepa(core, M),
            'nombre': f.get('nombre') or '', 'marca': f['marca'].strip(), 'categoria': f.get('categoria') or '',
            'precio_catalogo': _num(f.get('precio_unidad')), 'precio_unidad': pa_con_porte(f.get('precio_unidad'), pct),
            'es_caja': False, 'uds_caja': None, 'es_chase': False,
            'en_oferta': bool(f.get('en_oferta')), 'campana': '', 'disponibilidad': '',
            'fin_de_vida': bool(f.get('fin_de_vida')), 'preorder': False, 'imagen': '', 'aviso_caja': None,
            'origen_ean': None, 'aviso_ean': None,
        })
    for m in e2.MOTIVOS_APARTADO:
        previas[m] = sum(1 for a in apartados if a['motivo'] == m)
    n_previas = sum(previas.values())
    cuentas = {'n_crudo': len(filas or []), 'n_foto': len(foto), 'previas': previas, 'n_previas': n_previas,
               'cuadra_previo': len(filas or []) == n_previas + len(foto)}
    return foto, apartados, cuentas


def comprobar_porte(foto, enlaces, puerta):
    """🔴 EL PORTE, UNA SOLA VEZ. Cada fila de OSMA de la puerta comun (`v_escaner_fuente`: ean de la ficha, pa,
    presente) cuyo articulo esta en la foto tiene que traer EL MISMO precio que la foto. Si alguien sumara el porte
    otra vez (o la puerta comun dejara de sumarlo), aqui no cuadra y la pasada falla.
    → (comprobadas, [descuadres]). El articulo de cada fila de la puerta: el de su codigo enlazado (por el EAN de
    la ficha) o el de su mismo EAN."""
    por_clave = {}
    for f in foto:
        por_clave.setdefault(ean_norm(f['ean_core']), []).append(f)
    for codigo, e in enlaces.items():
        for f in foto:
            if f['producto_heo'] == codigo:
                por_clave.setdefault(ean_norm(e['ean_ficha']), []).append(f)
    comprobadas, mal = 0, []
    for v in puerta or []:
        if v.get('es_case') or not v.get('presente'):
            continue
        filas = por_clave.get(ean_norm(v.get('ean')))
        if not filas:
            continue
        comprobadas += 1
        pa = _num(v.get('pa'))
        if not any(f['precio_unidad'] is not None and pa is not None and abs(f['precio_unidad'] - pa) < 0.005
                   for f in filas):
            mal.append('EAN %s: la puerta común dice %s y la foto %s'
                       % (v.get('ean'), v.get('pa'), ', '.join(str(f['precio_unidad']) for f in filas)))
    return comprobadas, mal


def lista_para_keepa_osma(foto, enlaces, M):
    """(AG2) LA lista para el Visualizador → (codigos, eans_fichas):
      · los EAN de la foto (uno por linea, sin repetir, con el rescate de GTIN del viejo como HEO), y
      · el EAN de NUESTRA ficha de cada codigo enlazado QUE ESTA EN LA FOTO (un codigo sin existencias o sin marca
        no tiene a que fila colgarse), detras y solo si no estaba ya (sin ceros delante, como compara Keepa).
    La familia de un codigo (activas, no chase, `enlaces_de`) comparte EAN: es UN EAN por codigo. `eans_fichas`:
    {codigo: EAN de la ficha} de los codigos de la foto que tienen EAN (este o no ya en la lista)."""
    codigos = e2.lista_para_keepa(foto)
    vistos = {M.norm(c) for c in codigos}
    en_foto = {f['producto_heo'] for f in foto}
    eans_fichas = {}
    for c in sorted(enlaces, key=_clave_codigo):
        ean = str(enlaces[c].get('ean_ficha') or '').strip()
        if c not in en_foto or not M.norm(ean):
            continue
        eans_fichas[c] = ean
        if M.norm(ean) not in vistos:
            vistos.add(M.norm(ean))
            codigos.append(ean)
    return codigos, eans_fichas


def _codigos_keepa(fila):
    """Los codigos que Keepa tiene de una ficha (keepa_escaparate: `ean_keepa_crudo` y `upc_keepa`), normalizados."""
    texto = '%s,%s' % (fila.get('ean_keepa_crudo') or '', fila.get('upc_keepa') or '')
    return {k for k in (ean_norm(x) for x in re.split(r'[^0-9]+', texto)) if k}


def fuera_de_keepa(enlaces, eans_fichas, keepa, dominios, mapa_por_codigo=None):
    """(AG2) Nuestras fichas de los codigos de la foto que la lista NO puede traer porque Keepa no conoce su EAN:
    NO se rescatan (Fernando, 02-oct-2026). Se calcula en cada cruce con keepa_escaparate (FOTO: se dice su fecha),
    en los paises del cruce (`dominios`, en minusculas: 'es'). → {'fuera': [(codigo, asin, por que)],
    'sin_dato': [(codigo, asin)], 'fecha': la mas reciente de las filas miradas}.
      · el codigo sin EAN en la ficha → fuera ('la ficha no tiene EAN');
      · la ficha con fila en keepa_escaparate en algun pais del cruce y en NINGUNO lista nuestro EAN → fuera;
      · sin fila en ningun pais del cruce → 'sin_dato' (no se sabe: se dice aparte, no se da por fuera).
    (Mapa de OSMA, Cowork 04-oct-2026) Una ficha que el MAPA ya valora para ese codigo (`mapa_por_codigo`: {codigo:
    {asin: unidades}}, el mapa de los codigos disponibles) no queda fuera: entra por su ASIN en la lista del mapa. No se
    mira ni se lista (el Lenor 18459: B014DGG0OQ y B07HCJQ45L)."""
    mapa_por_codigo = mapa_por_codigo or {}
    por_asin = {}
    for k in keepa or []:
        if str(k.get('dominio') or '').lower() in dominios:
            por_asin.setdefault(k.get('asin'), []).append(k)
    fuera, sin_dato, fechas = [], [], []
    for c in sorted(enlaces, key=_clave_codigo):
        clave = ean_norm(eans_fichas.get(c))
        for a in enlaces[c]['asins']:
            if a in (mapa_por_codigo.get(c) or {}):
                continue
            filas = por_asin.get(a) or []
            fechas += [str(k['fecha_foto']) for k in filas if k.get('fecha_foto')]
            if not clave:
                fuera.append((c, a, 'la ficha no tiene EAN'))
            elif not filas:
                sin_dato.append((c, a))
            elif not any(clave in _codigos_keepa(k) for k in filas):
                fuera.append((c, a, 'Keepa no conoce el EAN %s de la ficha (%s)'
                              % (eans_fichas[c], ', '.join(sorted(str(k['dominio']).upper() for k in filas)))))
    return {'fuera': fuera, 'sin_dato': sin_dato, 'fecha': max(fechas) if fechas else None}


def texto_fuera_de_keepa(fk):
    """La linea del Resumen: las fichas que quedan fuera, por codigo, y las que no se saben."""
    if not fk:
        return 'no se ha podido mirar'
    partes = ['%s · %s: %s' % (c, a, porque) for c, a, porque in fk['fuera']] or ['ninguna']
    if fk['sin_dato']:
        partes.append('sin fila en keepa_escaparate (no se sabe): %s'
                      % ', '.join('%s · %s' % (c, a) for c, a in fk['sin_dato']))
    return '%s (keepa_escaparate del %s)' % (' | '.join(partes), fk['fecha'] or '—')


# ═══════════════════════════════════════════════════════════════════════════════
# 4 · EL CRUCE: LAS FICHAS DE CADA ARTICULO
# ═══════════════════════════════════════════════════════════════════════════════
def indice_por_asin(datos_por_ean):
    """{asin: registro} a partir de lo que devuelve `leer_csv_visualizador` (indexado por EAN): el primer
    registro de cada ASIN (el lector recorre los CSV del mas nuevo al mas viejo, como el cruce de HEO)."""
    salida = {}
    for recs in (datos_por_ean or {}).values():
        for r in recs:
            if r.get('asin') and r['asin'] not in salida:
                salida[r['asin']] = r
    return salida


def de_las_fichas(datos, eans_fichas, M):
    """(AG2) Los registros del CSV que llegan por el EAN de una ficha enlazada (por identidad: el lector cuelga el
    mismo registro de cada uno de sus EAN). Son de la fila de ESE codigo, nunca de otro articulo de OSMA."""
    return {id(r) for e in (eans_fichas or {}).values() for r in (datos or {}).get(M.norm(e), [])}


def candidatos_osma(fila, datos, por_asin, enlace, asins_enlazados, M, ean_ficha=None, ids_fichas=frozenset()):
    """Las fichas de UN pais para una fila de la foto → (candidatas, camino, apartadas). (AG2) UN solo CSV, con los
    EAN de OSMA y los de nuestras fichas:
      · fila de un codigo ENLAZADO: las de nuestras fichas que trae el CSV, por cualquiera de sus EAN (`camino`
        'codigo'); lo que traen su EAN de OSMA o el de nuestra ficha con OTRO ASIN se aparta (la misma ficha por los
        dos caminos: una, la del codigo). Si el CSV no trae ninguna de nuestras fichas, las de su EAN de OSMA, como
        cualquier otra fila ('ean_sin_codigo').
      · las demas: las de su EAN ('ean'), menos las de una ficha enlazada y las que llegan por el EAN de una ficha
        enlazada (`ids_fichas`, de `de_las_fichas`), que se apartan: son de la fila de su codigo.
    `datos`: el lector del CSV; `por_asin`: `indice_por_asin` del mismo CSV."""
    por_ean = e2.candidatos(fila, datos or {})
    if enlace:
        por_codigo = [por_asin[a] for a in enlace['asins'] if a in (por_asin or {})]
        if por_codigo:
            por_ficha = e2.candidatos({'variantes': [M.norm(ean_ficha)]}, datos or {}) if M.norm(ean_ficha or '') else []
            vistos, apartadas = set(), []
            for r in por_ean + por_ficha:
                if r.get('asin') and r['asin'] not in enlace['asins'] and r['asin'] not in vistos:
                    vistos.add(r['asin'])
                    apartadas.append(r)
            return por_codigo, 'codigo', apartadas
        return por_ean, 'ean_sin_codigo', []
    fuera = [r for r in por_ean if r.get('asin') in asins_enlazados or id(r) in ids_fichas]
    return [r for r in por_ean if not (r.get('asin') in asins_enlazados or id(r) in ids_fichas)], 'ean', fuera


def decidir_osma(fila, cands_por_pais, caidas_por_pais, params, M, eleccion, compartidas, enlace, factores,
                 caminos=None, apartadas=None, senales_pack=None, mapa_codigo=None, pct=None):
    """La puerta de UNA fila de OSMA con las reglas de compra del PRO de HEO (`escaner2_motor.decidir`, sin
    tocarlas), y lo propio de OSMA alrededor:
      · fila enlazada: el IVA (y «En mi BD») se leen de NUESTRA ficha (su EAN), no del EAN nuevo de OSMA;
      · si la ficha que sale es uno de nuestros PACKS, el coste es unidades × precio: se decide OTRA VEZ con ese
        coste (la eleccion de ficha no depende del precio: sale la misma, y se comprueba);
      · (AM) si no lo es y la fila se vende (puertas d, e y f), el PACK DE AMAZON (`factor_pack_amazon`, con las
        señales de `senales_pack`: {pais: {asin: señal}} de `senales_pack_csv`): con pack, igual que el nuestro
        (N × precio y se decide otra vez); con «posible pack», la fila nunca es COMPRAR (baja a VALORAR);
      · el detalle dice por donde llego la ficha y lo apartado.
      · (Mapa de OSMA, Fernando, 04-oct-2026) si el ASIN que sale esta en el MAPA de este codigo (`mapa_codigo`:
        {asin: unidades}), MANDAN LAS UNIDADES DEL MAPA, sobre nuestro factor y sobre las señales del AM (que ni se
        miran), y el coste es el del mapa (`pa_mapa` con `pct`, un solo redondeo).
    Devuelve el resultado de `decidir` con 'factor', 'pa' (el coste con el que se ha calculado), 'pack' (None,
    'nuestro', 'amazon' o 'mapa') y 'pack_amazon' (el veredicto, o None si no se ha mirado)."""
    calc = dict(fila, ean_core=str(enlace['ean_ficha']).strip()) if (enlace and enlace.get('ean_ficha')) else dict(fila)
    r = e2.decidir(calc, cands_por_pais, caidas_por_pais, params, M, eleccion, compartidas)
    factor = (factores or {}).get(r.get('asin'), 1) if r.get('asin') else 1
    pack = 'nuestro' if factor > 1 else None
    del_mapa = (mapa_codigo or {}).get(r.get('asin')) if r.get('asin') else None
    if del_mapa is not None:
        factor, pack = del_mapa, ('mapa' if del_mapa > 1 else None)
    pack_amz = None
    if (del_mapa is None and factor == 1 and senales_pack is not None and r.get('asin')
            and r['puerta'] in e2.PUERTAS_ANALISIS):
        senal = next((senales_pack[p][r['asin']] for p in params['paises_calculo']
                      if r['asin'] in (senales_pack.get(p) or {})), None)
        pack_amz = factor_pack_amazon(senal, cantidad_osma(fila.get('nombre')))
        if pack_amz['estado'] == PACK_SI:
            factor, pack = pack_amz['factor'], 'amazon'
    pa = fila['precio_unidad']
    if factor > 1 and pa is not None:
        if pack == 'mapa' and fila.get('precio_catalogo') is not None:
            pa = pa_mapa(fila['precio_catalogo'], factor, pct)
        else:
            pa = float((_dec(pa) * factor).quantize(Decimal('0.01'), rounding=ROUND_HALF_UP))
        r2 = e2.decidir(dict(calc, precio_unidad=pa), cands_por_pais, caidas_por_pais, params, M, eleccion, compartidas)
        if r2.get('asin') != r.get('asin'):
            raise FalloOsma('la ficha elegida cambia con el coste del pack (%s → %s): no debería'
                            % (r.get('asin'), r2.get('asin')))
        r = r2
    if pack_amz and pack_amz['estado'] == PACK_DUDOSO and r['puerta'] == 'f':
        r = _sin_comprar(r, params)
    notas = []
    caminos = caminos or {}
    if enlace and 'codigo' in caminos.values():
        notas.append('Por el código %s de OSMA (nuestra ficha %s)' % (fila['producto_heo'], ', '.join(enlace['asins'])))
    elif enlace:
        notas.append('Código %s enlazado, pero el CSV no trae nuestra ficha: decidido por el EAN de OSMA'
                     % fila['producto_heo'])
    if factor > 1:
        notas.append(texto_pack(factor, fila['precio_unidad'], pa, pack))
    if pack == 'amazon':
        notas.append('señales: %s' % pack_amz['senales'])
    if pack_amz and pack_amz['estado'] == PACK_DUDOSO:
        notas.append('%s (%s)' % (TEXTO_POSIBLE_PACK, pack_amz['senales']))
    aparte = sorted({x.get('asin') for xs in (apartadas or {}).values() for x in xs if x.get('asin')})
    if aparte:
        notas.append('apartadas (de otra ficha enlazada o del EAN nuevo): %s' % ', '.join(aparte))
    if notas:
        r = dict(r, detalle='%s · %s' % (r['detalle'], ' · '.join(notas)))
    return dict(r, factor=factor, pa=pa, pack=pack, pack_amazon=pack_amz)


def texto_pack(factor, precio, pa, pack):
    """«pack de 3: coste 3 × 2,13 € = 6,39 €» (nuestro) o «pack de 48 en Amazon: coste 48 × 0,32 € = 15,36 €»."""
    if pack == 'mapa':
        # Sin «N × precio con porte»: el del mapa redondea una vez, al final (2 × 1,599 × 1,067276 = 3,41, no 2 × 1,71).
        return 'pack de %d del mapa: coste %d × precio de OSMA × (1 + porte) = %s €' % (factor, factor, _coma(pa))
    return 'pack de %d%s: coste %d × %s € = %s €' % (factor, ' en Amazon' if pack == 'amazon' else '', factor,
                                                     _coma(precio), _coma(pa))


def _sin_comprar(r, params):
    """(AM) Un «posible pack en Amazon» nunca es COMPRAR: cada pais que decia COMPRAR pasa a VALORAR y la puerta se
    saca OTRA VEZ con `escaner2_motor._puerta_def` (la e, con su mejor pais). En el detalle se cambia SOLO el trozo
    que escribe `_puerta_def` (lo de delante, la eleccion de ficha, y lo de detras, la ficha compartida, se quedan)."""
    paises = [p for p in params['paises_calculo'] if p in (r.get('paises') or {})]
    antes = e2._puerta_def(dict(r), paises, r['paises'])['detalle']
    calculos = {p: dict(c, decision='VALORAR') if c.get('decision') == 'COMPRAR' else c for p, c in r['paises'].items()}
    nuevo = e2._puerta_def(dict(r, paises=calculos), paises, calculos)
    if r['detalle'].count(antes) != 1:
        raise FalloOsma('el detalle de la puerta no se puede rehacer (posible pack): %r no está una vez' % antes)
    return dict(nuevo, detalle=r['detalle'].replace(antes, nuevo['detalle']))


def _coma(x):
    return '—' if x is None else ('%.2f' % x).replace('.', ',')


# ═══════════════════════════════════════════════════════════════════════════════
# 4 bis · (AM, 02-oct-2026) LOS PACKS DE AMAZON
# ═══════════════════════════════════════════════════════════════════════════════
# OSMA vende la unidad suelta, y el cruce por EAN cae a veces en un ASIN que es un MULTIPACK con el EAN de la unidad
# (las Melody Pops de 15 g en el «Pack 48»): con el coste de UNA unidad salian COMPRAR falsos (primera pasada real,
# 655b3e06). Aqui se lee, del CSV del Visualizador, cuantas unidades de OSMA lleva el ASIN, y se multiplica el coste
# SOLO cuando lo sostienen dos señales independientes que dicen lo mismo.
#
# LA REGLA. Cada señal es una LECTURA de N = unidades de OSMA que lleva el ASIN, comparando lo que dice Amazon con la
# cantidad del nombre de OSMA (`cantidad_osma`: el recuento «24er», «8 Stück», «56x10», y la medida «45g», «50ml»,
# «100ml + 250ml»; sin recuento en el nombre, el articulo de OSMA es UNA unidad):
#   · contenido   «Detalles de la unidad»: valor y tipo. En g/ml, ÷ la medida de OSMA; en unidades, ÷ su recuento.
#   · tamaño      «Tamaño» («45 g (Paquete de 24)» = 1080 g; «24 unità (Confezione da 1)» = 24 ud), igual.
#   · recuento    «Número de artículos» (solo si es 2 o más: Amazon pone 1 casi siempre), ÷ el recuento de OSMA.
#   · paquete     «Paquete: Cantidad» (solo si es 2 o más), igual.
#   · título      «Pack 48», «paquete de 4», «lote de 12», «3 x 50 ml», «10 piezas», «24 unidades» (2 o más), igual.
#   El contenido en unidades y el recuento son, los dos, el recuento que escribe el vendedor: cuentan como UNA señal.
# Y el veredicto, por este orden:
#   1. Si las lecturas de CONTENIDO y TAMAÑO que hay dan todas 1 contra una cantidad LEIDA en el nombre de OSMA, el
#      ASIN lleva justo lo que vende OSMA: no es pack, diga lo que diga el recuento (el set de Adidas: 100 + 250 ml =
#      350 ml, aunque Amazon diga «6 artículos»).
#   2. Si ninguna lectura da 2 o más: no es pack.
#   3. Se MULTIPLICA por N si todas las lecturas dan N (un entero ≥ 2), vienen de DOS señales distintas como poco, y
#      el nombre de OSMA se ha leido sin ambigüedad.
#   4. Si no (una sola señal, o señales que se contradicen, o el nombre de OSMA con dos cantidades distintas): NO se
#      multiplica y la fila nunca es COMPRAR: baja a VALORAR con «posible pack en Amazon: revisar (…las señales…)».
# 🔒 Si el ASIN es de una ficha nuestra que ya se valora como pack (`factor_por_asin`), esto no se mira: no se
#    multiplica dos veces.
COLUMNAS_PACK = {'n_art': 'Número de artículos', 'valor_ud': 'Detalles de la unidad: Valor de la unidad',
                 'tipo_ud': 'Detalles de la unidad: Tipo de unidad', 'paquete': 'Paquete: Cantidad',
                 'tamano': 'Tamaño'}
# La columna del titulo, como la busca el lector del CSV del viejo (escaner2_heredado_pro.leer_csv_visualizador).
COLUMNAS_TITULO = ('Título', 'Titulo', 'Title', 'Título principal', 'Titulo principal')
PACK_SI, PACK_DUDOSO, PACK_NO = 'pack', 'dudoso', 'no'
# Lo que se admite como «el mismo numero» (58,33 ml × 6 = 349,98 ml frente a 350 ml).
TOLERANCIA_PACK = 0.02
TEXTO_POSIBLE_PACK = 'posible pack en Amazon: revisar'

# Unidades, a g, ml o 'ud' (recuento). Lo que no esta aqui (cm, metro, onzas…) no es una cantidad que se compare.
_UNIDAD = {'ml': ('ml', 1), 'mililitro': ('ml', 1), 'millilitro': ('ml', 1), 'milliliter': ('ml', 1),
           'milliliters': ('ml', 1), 'mililitros': ('ml', 1), 'cl': ('ml', 10), 'l': ('ml', 1000),
           'ltr': ('ml', 1000), 'liter': ('ml', 1000), 'litro': ('ml', 1000), 'litros': ('ml', 1000),
           'g': ('g', 1), 'gr': ('g', 1), 'gramo': ('g', 1), 'gramos': ('g', 1), 'gram': ('g', 1), 'grams': ('g', 1),
           'kg': ('g', 1000), 'kilogramo': ('g', 1000), 'kilogramos': ('g', 1000),
           'unidad': ('ud', 1), 'unidades': ('ud', 1), 'unità': ('ud', 1), 'unité': ('ud', 1), 'unités': ('ud', 1),
           'stück': ('ud', 1), 'stückzahl': ('ud', 1), 'count': ('ud', 1), 'pieza': ('ud', 1), 'piezas': ('ud', 1)}
_MEDIDA = r'(ml|cl|ltr|liter|l|kg|gr|g)'
_FIN_PALABRA = r'(?![a-zäöüßà-ÿ])'
_RE_N_X_MEDIDA = re.compile(r'(?<![\d.,])(\d{1,3})\s*[xX]\s*(\d+(?:[.,]\d+)?)\s*' + _MEDIDA + _FIN_PALABRA, re.I)
_RE_MEDIDA_OSMA = re.compile(r'(?<![\d.,])(\d+(?:[.,]\d+)?)\s*' + _MEDIDA + r'\.?' + _FIN_PALABRA, re.I)
_RE_ER = re.compile(r'(?<![\d.,])(\d+)\s*er' + _FIN_PALABRA, re.I)
_RE_STUECK = re.compile(r'(?<![\d.,])(\d+)\s*-?\s*(?:stück|stck|stk|st|teilig|teil|tlg)\.?' + _FIN_PALABRA, re.I)
# «56x10» (Tempo: 56 paquetes de 10) es un recuento; «190x68» (una vela) o «20x15cm» son medidas: solo cifras de una
# o dos y sin unidad ni otra «x» detras.
_RE_NXM = re.compile(r'(?<![\d.,])(\d{1,2})\s*[xX]\s*(\d{1,2})(?![\d.,])(?!\s*(?:[xX]|cm|mm|m' + _FIN_PALABRA + r'|'
                     + _MEDIDA + _FIN_PALABRA + r'))', re.I)
# «im 18er Tray», «12 St. im Aufsteller»: el expositor, no lo que se vende.
_RE_EXPOSITOR = re.compile(r'^\W*(?:im\s+)?(?:tray|display|aufsteller|karton)' + _FIN_PALABRA, re.I)
_RE_TAMANO = re.compile(r'^\s*(?:(\d{1,3})\s*x\s*)?(\d+(?:[.,]\d+)?)\s*([a-zà-ÿ]+)\b(?:\s*\((?:paquete de|pack of|'
                        r'confezione da|lot de|packung mit)\s*(\d+)\)|\s*\((\d+)er[\s-]*pack\))?', re.I)
_RES_TITULO = [re.compile(p, re.I) for p in (
    r'\bpack\s*(?:de\s*)?(\d+)\b', r'\b(\d+)\s*-?\s*pack\b', r'\blote\s+de\s+(\d+)\b', r'\bpaquete\s+de\s+(\d+)\b',
    r'\bcaja\s+de\s+(\d+)\b', r'\b(\d+)\s*x\s*\d+(?:[.,]\d+)?\s*' + _MEDIDA + r'\b', r'\b(\d+)\s+(?:piezas|unidades)\b')]


def _numero(s):
    try:
        return float(str(s).strip().replace(',', '.'))
    except (TypeError, ValueError):
        return None


def senales_pack_csv(rutas):
    """{asin: {'n_art', 'valor_ud', 'tipo_ud', 'paquete', 'tamano', 'titulo'}} del CSV del Visualizador, tal cual
    (texto). `rutas`: una o varias, del mas nuevo al mas viejo, y EL MAS NUEVO MANDA (como el lector del viejo). Una
    columna que el CSV no trae sale vacia."""
    import csv
    salida = {}
    for ruta in ([rutas] if isinstance(rutas, str) else rutas):
        with open(ruta, encoding='utf-8-sig', newline='') as fh:
            filas = csv.reader(fh)
            cab = next(filas, [])
            ix = {h: i for i, h in enumerate(cab)}
            if 'ASIN' not in ix:
                continue
            col_tit = next((c for c in COLUMNAS_TITULO if c in ix), None)

            def celda(fila, col):
                return fila[ix[col]].strip() if (col in ix and ix[col] < len(fila)) else ''
            for fila in filas:
                asin = celda(fila, 'ASIN')
                if asin and asin not in salida:
                    salida[asin] = dict({k: celda(fila, c) for k, c in COLUMNAS_PACK.items()},
                                        titulo=celda(fila, col_tit) if col_tit else '')
    return salida


def cantidad_osma(nombre):
    """Lo que dice el nombre de OSMA de su cantidad → {'n': recuento o None, 'medida': (total, 'g'|'ml') o None,
    'ambigua': bool, 'texto': lo leido}. El recuento: «24er», «8 Stück», «132 St.», «77Stk», «32 teil.», «4tlg»,
    «56x10», o el N de «6x50g»; el de un expositor («im 18er Tray») no cuenta. La medida, en g o ml: «45g», «1,35 Liter», «6x50g» = 300 g,
    y las de un set con «+» se suman («100ml + 250ml» = 350 ml). Dos recuentos distintos, o dos medidas distintas sin
    «+», no se adivinan: 'ambigua'."""
    s = str(nombre or '')
    recuentos, medidas, tapado = [], [], []
    for m in _RE_N_X_MEDIDA.finditer(s):
        dim, k = _UNIDAD[m.group(3).lower()]
        recuentos.append(int(m.group(1)))
        medidas.append((m.start(), m.end(), int(m.group(1)) * _numero(m.group(2)) * k, dim))
        tapado.append((m.start(), m.end()))
    for m in _RE_MEDIDA_OSMA.finditer(s):
        if any(a <= m.start() < b for a, b in tapado):
            continue
        dim, k = _UNIDAD[m.group(2).lower()]
        medidas.append((m.start(), m.end(), _numero(m.group(1)) * k, dim))
    for rx in (_RE_ER, _RE_STUECK):
        for m in rx.finditer(s):
            if not _RE_EXPOSITOR.match(s[m.end():]):
                recuentos.append(int(m.group(1)))
    for m in _RE_NXM.finditer(s):
        recuentos.append(int(m.group(1)) * int(m.group(2)))
    recuentos = [n for n in recuentos if n > 0]
    ambigua = len(set(recuentos)) > 1
    n = recuentos[0] if (recuentos and not ambigua) else None
    medida = None
    if medidas:
        medidas.sort()
        dims = {x[3] for x in medidas}
        unidas = all('+' in s[medidas[i][1]:medidas[i + 1][0]] for i in range(len(medidas) - 1))
        if len(dims) == 1 and len(medidas) > 1 and unidas:
            medida = (sum(x[2] for x in medidas), medidas[0][3])
        elif len(dims) == 1 and len({round(x[2], 6) for x in medidas}) == 1:
            medida = (medidas[0][2], medidas[0][3])
        else:
            ambigua = True
    partes = ([('%d ud' % n)] if n else []) + ([_cifra_pack(medida[0]) + ' ' + medida[1]] if medida else [])
    return {'n': n, 'medida': medida, 'ambigua': ambigua,
            'texto': ' · '.join(partes) or ('cantidad ambigua' if ambigua else 'sin cantidad: 1 unidad')}


def _cifra_pack(x):
    return ('%d' % round(x)) if abs(x - round(x)) < 1e-9 else ('%.2f' % x).rstrip('0').rstrip('.').replace('.', ',')


def _entero_pack(v):
    """La lectura como entero si lo es (con la tolerancia), y si no None."""
    if v is None or v <= 0:
        return None
    n = round(v)
    return n if (n >= 1 and abs(v - n) <= TOLERANCIA_PACK * n) else None


def lecturas_pack(senal, cant):
    """Las lecturas de N (unidades de OSMA que lleva el ASIN) de cada señal del CSV → [{'fuente', 'grupo', 'texto',
    'valor', 'trivial'}]. `trivial`: comparada contra el «1 unidad» que se supone a un nombre de OSMA sin recuento
    (no contra una cantidad leida)."""
    if not senal:
        return []
    recuento = cant.get('n') or 1
    sin_recuento = cant.get('n') is None
    medida = cant.get('medida')
    salida = []

    def contra(total, dim, fuente, texto, grupo=None):
        if dim == 'ud':
            salida.append({'fuente': fuente, 'grupo': grupo or fuente, 'texto': texto, 'valor': total / recuento,
                           'trivial': sin_recuento})
        elif medida and medida[1] == dim and medida[0] > 0:
            salida.append({'fuente': fuente, 'grupo': grupo or fuente, 'texto': texto, 'valor': total / medida[0],
                           'trivial': False})

    v, tipo = _numero(senal.get('valor_ud')), _UNIDAD.get(str(senal.get('tipo_ud') or '').strip().lower())
    if v and v > 0 and tipo:
        contra(v * tipo[1], tipo[0], 'contenido', 'contenido %s %s' % (senal['valor_ud'], senal['tipo_ud']),
               grupo='recuento' if tipo[0] == 'ud' else 'contenido')
    m = _RE_TAMANO.match(str(senal.get('tamano') or ''))
    if m and m.group(3).lower() in _UNIDAD:
        dim, k = _UNIDAD[m.group(3).lower()]
        por = int(m.group(1) or 1) * int(m.group(4) or m.group(5) or 1)
        total = _numero(m.group(2)) * k * por
        texto = 'tamaño «%s»' % senal['tamano']
        if dim == 'ud' or (medida and medida[1] == dim):
            contra(total, dim, 'tamaño', texto)
        elif por >= 2:
            # «20 g (Paquete de 3)» sin medida en el nombre de OSMA: cuenta el «Paquete de».
            salida.append({'fuente': 'tamaño', 'grupo': 'tamaño', 'texto': texto, 'valor': por / recuento,
                           'trivial': sin_recuento})
    for clave, fuente, grupo, rotulo in (('n_art', 'recuento', 'recuento', 'n.º de artículos'),
                                         ('paquete', 'paquete', 'paquete', 'paquete: cantidad')):
        x = _numero(senal.get(clave))
        if x is not None and x >= 2 and x == int(x):
            salida.append({'fuente': fuente, 'grupo': grupo, 'texto': '%s %d' % (rotulo, x), 'valor': x / recuento,
                           'trivial': sin_recuento})
    vistos = set()
    for rx in _RES_TITULO:
        for mt in rx.finditer(str(senal.get('titulo') or '')):
            x = int(mt.group(1))
            if x >= 2 and x not in vistos:
                vistos.add(x)
                salida.append({'fuente': 'título', 'grupo': 'título', 'texto': 'título «%s»' % mt.group(0).strip(),
                               'valor': x / recuento, 'trivial': sin_recuento})
    return salida


def factor_pack_amazon(senal, cant):
    """El veredicto de la regla de arriba → {'estado': PACK_SI | PACK_DUDOSO | PACK_NO, 'factor', 'senales': texto}."""
    lecturas = lecturas_pack(senal, cant)
    texto = 'OSMA %s; Amazon: %s' % (cant['texto'], ' · '.join(
        '%s → %s' % (x['texto'], _cifra_pack(x['valor'])) for x in lecturas) or 'sin señales')
    no = {'estado': PACK_NO, 'factor': 1, 'senales': texto}
    totales = [x for x in lecturas if x['fuente'] in ('contenido', 'tamaño') and not x['trivial']]
    if totales and all(_entero_pack(x['valor']) == 1 for x in totales):
        return no
    packs = [x for x in lecturas if (_entero_pack(x['valor']) or 0) >= 2]
    if not packs:
        return no
    ns = {_entero_pack(x['valor']) for x in packs}
    contra = [x for x in lecturas if _entero_pack(x['valor']) not in ns]
    if len(ns) == 1 and not contra and not cant['ambigua'] and len({x['grupo'] for x in packs}) >= 2:
        return {'estado': PACK_SI, 'factor': ns.pop(), 'senales': texto}
    return {'estado': PACK_DUDOSO, 'factor': 1, 'senales': texto}


# ═══════════════════════════════════════════════════════════════════════════════
# 4 ter · (Mapa de OSMA, Fernando, 04-oct-2026) EL MAPA FIJO: CODIGO OSMA → ASIN → UNIDADES
# ═══════════════════════════════════════════════════════════════════════════════
# El catalogo de OSMA es estable: en vez de buscar packs en Amazon cada vez, Cowork hace a mano un MAPA
# (`osma_mapa_ficha`, migracion 20261004113000 de la v2) con las fichas de Amazon de cada codigo de OSMA y cuantas
# unidades de OSMA lleva cada una (el Lenor 18459: B014DGG0OQ ×1, B07HCJQ45L ×2, B0794VHRVZ ×1). El PRO lo usa como
# lista ADICIONAL de fichas a valorar, sin tocar nada de lo de siempre (la foto, las puertas, el cuadre):
#   · BARRIDO: los ASIN del mapa de los codigos DISPONIBLES en la descarga van a `asins.txt`, una lista aparte que se
#     pega en el Visualizador de España en el modo «ASIN» (mezclados con los EAN, Keepa los descarta sin avisar). Los
#     codigos del mapa AGOTADOS en la descarga no van a Keepa: van a la hoja «Nuestros agotados en OSMA».
#   · CRUCE: cada fila del CSV de ASIN es del codigo y las unidades que dice el MAPA, nunca de lo que diga su EAN (la
#     ficha de un pack lleva su propio EAN). Coste = unidades × precio de OSMA × (1 + porte), redondeado al final
#     (Lenor 2 × 1,599 × 1,067276 = 3,41). La rentabilidad, la de siempre (`escaner2_motor.decidir`, sin tocarla).
#   · EXCEL: las fichas del mapa que se venden salen en «Análisis» como una fila mas, con la marca en una columna NUEVA AL
#     FINAL («Ficha del mapa»: hay quien lee por letra); si su ASIN ya estaba en una fila normal, MANDA la del mapa y la
#     normal no sale en «Análisis» (no se duplica; la base y el cuadre no cambian). Todas, en «Puertas».
COLUMNA_MAPA = 'Ficha del mapa'
ANCHO_MAPA = 24
HOJA_AGOTADOS = 'Nuestros agotados en OSMA'
TEXTO_AGOTADO = 'agotado en OSMA'
COLUMNAS_AGOTADOS = ['Código OSMA', 'Nombre OSMA', 'ASIN', 'Nuestra ficha', 'Unidades por ficha', 'Precio OSMA (€)',
                     'Estado']
# El CSV de ASIN se reconoce por su contenido: casi todos sus ASIN estan en la lista del mapa (el de EAN trae miles que no).
PARTE_CSV_MAPA = 0.8


def mapa_de_la_descarga(mapa, filas):
    """Las filas de `osma_mapa_ficha` de los codigos que estan en la descarga → (disponibles, agotados, codigos, fuera).
    `disponibles`/`agotados`: [{'codigo', 'asin', 'unidades', 'es_nuestra'}], por codigo y ASIN; `codigos`: {codigo:
    {'nombre', 'precio_catalogo', 'ean_original', 'ean_core', 'fin_de_vida', 'disponible'}} de la descarga; `fuera`:
    cuantas filas del mapa son de un codigo que la descarga no trae. Puro.
    `mapa`: filas de osma_mapa_ficha; `filas`: disp_estado de OSMA vistas en la ultima pasada aplicada."""
    por_cod = {str(f.get('producto_prov')): f for f in filas or []}
    disponibles, agotados, codigos, fuera = [], [], {}, 0
    for m in sorted(mapa or [], key=lambda x: (_clave_codigo(x.get('codigo_osma')), str(x.get('asin') or ''))):
        c = str(m.get('codigo_osma') or '').strip()
        f = por_cod.get(c)
        if f is None:
            fuera += 1
            continue
        fila = {'codigo': c, 'asin': str(m['asin']).strip(), 'unidades': int(m['unidades']),
                'es_nuestra': bool(m.get('es_nuestra'))}
        (disponibles if f.get('disponible') else agotados).append(fila)
        codigos[c] = {'nombre': f.get('nombre') or '', 'precio_catalogo': _num(f.get('precio_unidad')),
                      'ean_original': str(f.get('ean_original') or '').strip(),
                      'ean_core': str(f.get('ean_core') or '').strip(), 'fin_de_vida': bool(f.get('fin_de_vida')),
                      'disponible': bool(f.get('disponible'))}
    return disponibles, agotados, codigos, fuera


def lista_asin_mapa(disponibles):
    """La lista de ASIN para el Visualizador (modo «ASIN»): los del mapa de los codigos disponibles, sin repetir."""
    return list(dict.fromkeys(m['asin'] for m in disponibles or []))


def pa_mapa(precio, unidades, pct):
    """El coste de una ficha del mapa: unidades × precio de OSMA × (1 + porte), redondeado a 2 decimales UNA vez, al
    final (Fernando, 04-oct-2026: Lenor pack de 2 = 2 × 1,599 × 1,067276 = 3,41). Sin porte, unidades × precio."""
    if precio is None:
        return None
    v = _dec(precio) * int(unidades) * (1 + (_dec(pct) if pct is not None else 0))
    return float(v.quantize(Decimal('0.01'), rounding=ROUND_HALF_UP))


def es_csv_del_mapa(asins_csv, asins_mapa):
    """¿Es este CSV el de la lista de ASIN del mapa? Si casi todos sus ASIN (`PARTE_CSV_MAPA`) son del mapa."""
    asins_csv, asins_mapa = set(asins_csv or ()), set(asins_mapa or ())
    if not asins_csv or not asins_mapa:
        return False
    return len(asins_csv & asins_mapa) >= PARTE_CSV_MAPA * len(asins_csv)


def leer_csv_mapa(rutas, pro):
    """El CSV de la lista de ASIN, leido con el MISMO lector del Escaner Pro (`leer_csv_visualizador`, sin tocarlo) pero
    por ASIN: en una copia temporal, la columna de los EAN lleva el ASIN de la fila. Asi cada ficha sale igual que en
    el CSV de EAN, y una ficha sin EAN en Keepa (la de un pack) tampoco se pierde. → {asin: registro}, y el mas
    nuevo manda (`rutas` del mas nuevo al mas viejo, como el cruce)."""
    import csv
    import tempfile
    col_ean, col_asin = pro.CSV_COLS['ean'], pro.CSV_COLS['asin']
    copias = []
    for ruta in ([rutas] if isinstance(rutas, str) else rutas):
        with open(ruta, encoding='utf-8-sig', newline='') as fh:
            filas = list(csv.reader(fh))
        if not filas or col_asin not in filas[0]:
            continue
        cab = filas[0]
        i_asin = cab.index(col_asin)
        if col_ean not in cab:
            cab = cab + [col_ean]
        i_ean = cab.index(col_ean)
        fd, copia = tempfile.mkstemp(suffix='.csv')
        with open(fd, 'w', encoding='utf-8', newline='') as fh:
            w = csv.writer(fh)
            w.writerow(cab)
            for fila in filas[1:]:
                fila = fila + [''] * (len(cab) - len(fila))
                fila[i_ean] = fila[i_asin].strip()
                w.writerow(fila)
        copias.append(copia)
    salida = {}
    for recs in (pro.leer_csv_visualizador(copias) if copias else {}).values():
        for r in recs:
            if r.get('asin') and r['asin'] not in salida:
                salida[r['asin']] = r
    return salida


def texto_mapa(unidades):
    """La marca de la columna «Ficha del mapa»."""
    return 'pack ×%d (ficha del mapa)' % unidades if unidades > 1 else 'ficha del mapa'


def _coma3(x):
    return '—' if x is None else ('%.3f' % x).replace('.', ',')


def valorar_mapa(disponibles, codigos, por_asin_mapa, por_asin_csv, caidas_por_pais, params, usados, M, compartidas,
                 pct, fichas_nuestras):
    """Cada fila del mapa de un codigo disponible, valorada con las reglas de compra del PRO (`escaner2_motor.decidir`,
    sin tocarlas) y SU coste (`pa_mapa`) → [resultado de `decidir` + lo del mapa]. Puro.
      · la ficha, por su ASIN: la del CSV de ASIN (`por_asin_mapa`: {pais: {asin: registro}}) y, si no la trae, la del
        CSV de EAN (`por_asin_csv`, igual); nunca por un EAN;
      · el IVA (y «En mi BD»), de NUESTRA ficha si el ASIN es nuestro (`fichas_nuestras`: {asin: {'ean', 'nombre'}});
        si no, del EAN de OSMA.
    `usados`: los paises con CSV; `pct`: el porte (Decimal o texto) del barrido."""
    salida = []
    for m in disponibles or []:
        cod = codigos.get(m['codigo']) or {}
        ficha = (fichas_nuestras or {}).get(m['asin'])
        core = str((ficha or {}).get('ean') or '').strip() or cod.get('ean_core') or ''
        pa = pa_mapa(cod.get('precio_catalogo'), m['unidades'], pct)
        cands, fuente = {}, None
        for p in usados:
            rec = ((por_asin_mapa or {}).get(p) or {}).get(m['asin'])
            if rec is not None:
                fuente = fuente or 'asin'
            else:
                rec = ((por_asin_csv or {}).get(p) or {}).get(m['asin'])
                if rec is not None:
                    fuente = fuente or 'ean'
            cands[p] = [rec] if rec is not None else []
        r = e2.decidir({'nombre': cod.get('nombre') or '', 'ean_core': core, 'precio_unidad': pa}, cands,
                       caidas_por_pais, params, M, None, compartidas)
        coste = '%d × %s € × (1 + porte) = %s €' % (m['unidades'], _coma3(cod.get('precio_catalogo')), _coma(pa))
        detalle = 'Ficha del mapa (código %s, %d ud de OSMA en la ficha; coste %s)%s · %s' % (
            m['codigo'], m['unidades'], coste, '' if fuente else ' · ningún CSV la trae', r['detalle'])
        salida.append(dict(r, detalle=detalle, codigo=m['codigo'], asin_mapa=m['asin'], unidades=m['unidades'],
                           es_nuestra=m['es_nuestra'], nuestra=ficha is not None, pa=pa,
                           precio_catalogo=cod.get('precio_catalogo'), nombre_osma=cod.get('nombre') or '',
                           ean_original=cod.get('ean_original') or '', ean_core=core,
                           fin_de_vida=bool(cod.get('fin_de_vida')), fuente=fuente, coste=coste,
                           marca_mapa=texto_mapa(m['unidades'])))
    return salida


def agotados_nuestros(agotados, codigos, fichas_nuestras):
    """La hoja «Nuestros agotados en OSMA» (Fernando, 04-oct-2026): una fila por cada fila del mapa de un codigo NO
    disponible en la descarga del dia que tengamos a la venta (el ASIN en `productos`, vivo, o `es_nuestra` en el
    mapa). Sin Keepa ni margen: que no se esconda (el Dr. Beckmann 2346, B003U1NE40). Puro."""
    filas = []
    for m in agotados or []:
        ficha = (fichas_nuestras or {}).get(m['asin'])
        if ficha is None and not m.get('es_nuestra'):
            continue
        cod = codigos.get(m['codigo']) or {}
        filas.append([m['codigo'], cod.get('nombre') or '', m['asin'], (ficha or {}).get('nombre') or '—',
                      m['unidades'], cod.get('precio_catalogo'), TEXTO_AGOTADO])
    return filas


def fichas_nuestras_de(productos):
    """{asin: {'ean', 'nombre'}} de nuestras fichas vivas (activas, no chase, con ASIN; el chase nunca lleva ASIN)."""
    salida = {}
    for p in productos or []:
        if _ficha_viva(p) and p['asin'] not in salida:
            salida[p['asin']] = {'ean': p.get('ean'), 'nombre': p.get('nombre') or ''}
    return salida


def filas_mapa_excel(mapa):
    """Las fichas del mapa que se venden (puertas d, e y f), como filas de la foto y resultados para las hojas del viejo
    → (foto, resultados). Cada una con su id propio ('mapa:<codigo>:<asin>'), el EAN de OSMA en la columna EAN (o el
    ASIN, si el codigo no tiene EAN), el coste del mapa ya calculado y «PA (€)» vivo como el de UNA unidad sobre
    «unidades × precio sin porte» (asi da lo mismo que `pa_mapa`, con un solo redondeo)."""
    foto, resultados = [], []
    for r in mapa or []:
        if r['puerta'] not in e2.PUERTAS_ANALISIS:
            continue
        fid = 'mapa:%s:%s' % (r['codigo'], r['asin_mapa'])
        foto.append({
            'id': fid, 'producto_heo': r['codigo'], 'ean_original': r['ean_original'] or r['asin_mapa'],
            'ean_core': r['ean_core'], 'nombre': '%s · %s' % (r['nombre_osma'], r['marca_mapa']), 'marca': '',
            'precio_catalogo': r['precio_catalogo'], 'precio_unidad': r['pa'], 'fin_de_vida': r['fin_de_vida'],
            'aviso_caja': 'coste del mapa: %s' % r['coste'], 'asin': r['asin_mapa'], '_factor': 1,
            '_sin_porte': (None if r['precio_catalogo'] is None
                           else float(_dec(r['precio_catalogo']) * r['unidades'])),
            '_mapa': r['marca_mapa']})
        resultados.append(dict(r, foto_id=fid, id=fid, fichas=None, compartida={}, eleccion=None,
                               pack=None, pack_amazon=None, factor=1))
    return foto, resultados


def poner_columna_mapa(wb, foto_excel):
    """«Ficha del mapa» detras de la ultima columna de «Análisis» (columna nueva → al final, SIEMPRE), fila a fila por
    (EAN, ASIN), con la marca de las filas del mapa. Alarga la tabla. Devuelve cuantas filas."""
    from openpyxl.styles import Font
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.table import TableColumn
    ws = wb['Análisis']
    cab = [c.value for c in ws[1]]
    i_ean, i_asin = cab.index('EAN'), cab.index('ASIN')
    col = len(cab) + 1
    ws.cell(row=1, column=col, value=COLUMNA_MAPA).font = Font(bold=True)
    ws.column_dimensions[get_column_letter(col)].width = ANCHO_MAPA
    marcas = {(str(f['ean_original']), f.get('asin')): f['_mapa'] for f in foto_excel if f.get('_mapa')}
    n = 0
    for fila in range(2, ws.max_row + 1):
        m = marcas.get((str(ws.cell(row=fila, column=i_ean + 1).value), ws.cell(row=fila, column=i_asin + 1).value))
        if m:
            ws.cell(row=fila, column=col, value=m)
            n += 1
    for tabla in ws.tables.values():
        ini, fin = tabla.ref.split(':')
        tabla.ref = '%s:%s%s' % (ini, get_column_letter(col), re.sub(r'^[A-Z]+', '', fin))
        if tabla.tableColumns:
            tabla.tableColumns.append(TableColumn(id=len(tabla.tableColumns) + 1, name=COLUMNA_MAPA))
        if tabla.autoFilter is not None:
            tabla.autoFilter.ref = tabla.ref
    return n


# ═══════════════════════════════════════════════════════════════════════════════
# 5 · EL EXCEL: EL DEL ESCANER 2 DE HEO, CON UN SOLO AÑADIDO
# ═══════════════════════════════════════════════════════════════════════════════
def foto_para_excel(foto, resultados, enlaces):
    """La foto con la que se escribe la Celda 9, fila a fila igual que la guardada salvo dos cosas, para que las
    formulas vivas de «Análisis» den lo mismo que la decision:
      · si la ficha que sale es uno de nuestros packs, «PA (€)» es el coste del pack (unidades × precio) y el
        nombre lo dice (« · pack de 2 (2 × 1,71 €)»);
      · si la fila es de un codigo enlazado, `ean_core` es el EAN de nuestra ficha (el IVA y «En mi BD» de la
        ficha); la columna «EAN» sigue diciendo el de OSMA."""
    por_foto = {r['foto_id']: r for r in resultados}
    salida = []
    for f in foto:
        r = por_foto.get(f['id']) or {}
        g = dict(f)
        e = enlaces.get(f['producto_heo'])
        if e and e.get('ean_ficha'):
            g['ean_core'] = str(e['ean_ficha']).strip()
        factor = r.get('factor', 1)
        # (AM) Lo que pide la celda viva de «PA (€)»: el precio de OSMA sin porte (× N si es pack) y el factor.
        g['_factor'] = factor
        g['_sin_porte'] = (None if f.get('precio_catalogo') is None
                           else float(_dec(f['precio_catalogo']) * factor))
        if factor > 1:
            g['precio_unidad'] = r['pa']
            g['nombre'] = '%s · pack de %d%s (%d × %s €)' % (f.get('nombre') or '', factor,
                                                           {'amazon': ' en Amazon', 'mapa': ' del mapa'}.get(
                                                               r.get('pack'), ''), factor,
                                                           _coma(f['precio_unidad']))
            # (AM) «Coherencia caja», que en OSMA iba vacia: el pack, nuestro o de Amazon.
            g['aviso_caja'] = texto_pack(factor, f['precio_unidad'], r['pa'], r.get('pack'))
        elif (r.get('pack_amazon') or {}).get('estado') == PACK_DUDOSO:
            g['aviso_caja'] = '%s (%s)' % (TEXTO_POSIBLE_PACK, r['pack_amazon']['senales'])
        # (Mapa de OSMA) El ASIN de la fila: con el EAN, la llave de cada fila de «Análisis» (una ficha del mapa puede
        # llevar el mismo EAN de OSMA que la fila normal de su código).
        g['asin'] = r.get('asin')
        salida.append(g)
    return salida


def paises_de_la_hoja(params, M):
    """(AM) Los paises que pinta «Análisis»: con UN solo pais calculado (OSMA, solo ES), ese; con mas, los cuatro del
    viejo, como HEO (las filas de un pais sin CSV salen «Sin datos»)."""
    calc = list(params.get('paises_calculo') or [])
    return [p for p in M.PAISES if p in calc] if len(calc) == 1 else list(M.PAISES)


def excel_como_el_viejo_osma(foto_excel, resultados, apartados, M, eleccion=None, paises=None, mapa=None):
    """Las hojas del viejo con SU codigo (la Celda 9, `escaner2_motor.escribir_celda9`) y la columna «Ficha
    compartida», como el PRO de HEO (`escaner2_motor.excel_como_el_viejo`), con `PROVEEDOR` = 'OSMA': la Celda 9
    solo escribe «Chase_manual» para HEO (son los Funko chase sin ASIN), igual que hacia el escaner viejo con
    OSMA. Y DETRAS de todo, en «Análisis», la columna «No habrá más» y (AM) «Precio OSMA sin porte (€)», con «PA (€)»
    y «Decisión» vivas (`poner_pedido_vivo`). (AM) `paises`: los de `paises_de_la_hoja`.
    (Mapa de OSMA) `mapa` = (resultados de las fichas del mapa que se venden, ASIN de TODAS las fichas del mapa
    valoradas): sus filas entran en «Análisis» con las de siempre y ordenadas igual (por margen en ES); la fila normal
    cuyo ASIN es del mapa no sale (manda la del mapa). Detras de todo, la columna «Ficha del mapa»."""
    res_mapa, asins_mapa = mapa or ([], set())
    datos = e2.datos_como_el_viejo(foto_excel, resultados, apartados, M, eleccion)
    if res_mapa or asins_mapa:
        extra = e2.datos_como_el_viejo(foto_excel, res_mapa, [], M, None)['registros']
        regs = [x for x in datos['registros'] if x['asin'] not in asins_mapa] + extra
        regs.sort(key=lambda x: x['_margen_es'] if x['_margen_es'] is not None else -10 ** 9, reverse=True)
        datos['registros'] = regs
        for x in extra:
            datos['cotejo_info'].setdefault(x['ean'], {'veredicto': None, 'detalle': None})
    datos['PROVEEDOR'] = PROVEEDOR
    wb = e2.escribir_celda9(datos, M, paises=paises)
    e2.poner_columna_ficha_compartida(wb, foto_excel, resultados)
    poner_columna_no_habra_mas(wb, foto_excel)
    poner_pedido_vivo(wb, foto_excel, list(resultados) + list(res_mapa))
    poner_columna_mapa(wb, foto_excel)
    return wb


# ── (AM) EL PEDIDO PREVISTO Y EL PORTE, VIVOS EN EL EXCEL ────────────────────────────────────────────────────────
# Fernando quiere ver que pasa si pide mas o menos. NO se toca la base (disp_parametros.pedido_previsto_eur, la
# puerta comun ni Reponer): en «Resumen» van dos celdas amarillas, el pedido previsto y el porte de la pasada, y en
# «Análisis» «PA (€)» y «Decisión» son formulas que apuntan a ellas (por su nombre: PedidoPrevisto y PorteEnvio).
#   · PA (€) = REDONDEAR(sin porte × (1 + porte ÷ pedido); 2), la cuenta de `pa_con_porte` (el redondeo DESPUES del
#     porte, como la puerta comun); en un pack, N × REDONDEAR(sin porte ÷ N × (1 + porte ÷ pedido); 2), como
#     `decidir_osma` (N × el precio de la unidad ya redondeado). Sin pedido (o a 0), sin porte, como `porte_de`.
#   · Decisión = la de `decision_de` del viejo con el margen de la fila: COMPRAR si margen × 100 ≥ 10, VALORAR si ≥ 1,
#     y si no NO COMPRAR; y en un «posible pack en Amazon», nunca COMPRAR. Solo donde la fila tiene cuenta (si no,
#     «Sin datos» escrito, como hoy). Beneficio, ROI y Margen no se tocan.
COLUMNA_SIN_PORTE = 'Precio OSMA sin porte (€)'
ANCHO_SIN_PORTE = 16
NOMBRE_PEDIDO, NOMBRE_PORTE = 'PedidoPrevisto', 'PorteEnvio'
CELDA_PEDIDO, CELDA_PORTE = '$B$2', '$B$3'
UMBRAL_COMPRAR, UMBRAL_VALORAR = 10, 1   # los de decision_de (escaner2_heredado_nube.py), en % de margen


def formula_pa(letra_sin_porte, fila, factor):
    pct = 'IF(N(%s)>0,N(%s)/%s,0)' % (NOMBRE_PEDIDO, NOMBRE_PORTE, NOMBRE_PEDIDO)
    if factor > 1:
        return '=%d*ROUND(%s%d/%d*(1+%s),2)' % (factor, letra_sin_porte, fila, factor, pct)
    return '=ROUND(%s%d*(1+%s),2)' % (letra_sin_porte, fila, pct)


def formula_decision(letra_margen, fila, posible_pack):
    m = '%s%d*100' % (letra_margen, fila)
    if posible_pack:
        return '=IF(%s>=%d,"VALORAR","NO COMPRAR")' % (m, UMBRAL_VALORAR)
    return '=IF(%s>=%d,"COMPRAR",IF(%s>=%d,"VALORAR","NO COMPRAR"))' % (m, UMBRAL_COMPRAR, m, UMBRAL_VALORAR)


def poner_pedido_vivo(wb, foto_excel, resultados):
    """«Precio OSMA sin porte (€)» detras de la ultima columna de «Análisis» (alarga la tabla), y «PA (€)» y
    «Decisión» como formulas (ver arriba). Devuelve cuantas filas llevan la Decisión viva."""
    from openpyxl.styles import Font
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.table import TableColumn
    ws = wb['Análisis']
    cab = [c.value for c in ws[1]]
    i = {n: cab.index(n) + 1 for n in ('EAN', 'ASIN', 'PA (€)', 'Beneficio (€)', 'Margen', 'Decisión')}
    col = len(cab) + 1
    letra = get_column_letter(col)
    ws.cell(row=1, column=col, value=COLUMNA_SIN_PORTE).font = Font(bold=True)
    ws.column_dimensions[letra].width = ANCHO_SIN_PORTE
    # (Mapa de OSMA) Cada fila, por (EAN, ASIN): la ficha del mapa y la fila normal de su código pueden compartir EAN.
    por_clave = {(str(f['ean_original']), f.get('asin')): f for f in foto_excel}
    por_ean = {}
    for f in foto_excel:
        por_ean.setdefault(str(f['ean_original']), f)
    por_foto = {r['foto_id']: r for r in resultados}
    n = 0
    for fila in range(2, ws.max_row + 1):
        ean = str(ws.cell(row=fila, column=i['EAN']).value)
        f = por_clave.get((ean, ws.cell(row=fila, column=i['ASIN']).value)) or por_ean[ean]
        r = por_foto.get(f['id']) or {}
        if f.get('_sin_porte') is None:
            continue
        c = ws.cell(row=fila, column=col, value=f['_sin_porte'])
        c.number_format = '0.000'
        if ws.cell(row=fila, column=i['PA (€)']).value is not None:
            ws.cell(row=fila, column=i['PA (€)']).value = formula_pa(letra, fila, f.get('_factor', 1))
        # La Decision, viva solo donde la hoja echa la cuenta (el Beneficio es formula); si no, «Sin datos» escrito.
        if str(ws.cell(row=fila, column=i['Beneficio (€)']).value or '').startswith('='):
            dudoso = (r.get('pack_amazon') or {}).get('estado') == PACK_DUDOSO
            ws.cell(row=fila, column=i['Decisión']).value = formula_decision(
                get_column_letter(i['Margen']), fila, dudoso)
            n += 1
    for tabla in ws.tables.values():
        ini, fin = tabla.ref.split(':')
        tabla.ref = '%s:%s%s' % (ini, letra, re.sub(r'^[A-Z]+', '', fin))
        if tabla.tableColumns:
            tabla.tableColumns.append(TableColumn(id=len(tabla.tableColumns) + 1, name=COLUMNA_SIN_PORTE))
        if tabla.autoFilter is not None:
            tabla.autoFilter.ref = tabla.ref
    return n


def poner_celdas_pedido(wb, ws_resumen):
    """Los nombres PedidoPrevisto y PorteEnvio, a las dos celdas de entrada del Resumen."""
    from openpyxl.workbook.defined_name import DefinedName
    hoja = "'%s'" % ws_resumen.title
    for nombre, celda in ((NOMBRE_PEDIDO, CELDA_PEDIDO), (NOMBRE_PORTE, CELDA_PORTE)):
        wb.defined_names[nombre] = DefinedName(nombre, attr_text='%s!%s' % (hoja, celda))


def poner_columna_no_habra_mas(wb, foto):
    """«No habrá más» detras de la ultima columna de «Análisis» (columna nueva → al final, SIEMPRE: hay quien lee
    la hoja por letra), fila a fila por EAN, en los descatalogados de OSMA (Discontinued) que aun tienen
    existencias (todos los de la foto lo son: el filtro exige disponible). Alarga la tabla. Devuelve cuantas filas."""
    from openpyxl.styles import Font
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.table import TableColumn
    ws = wb['Análisis']
    cab = [c.value for c in ws[1]]
    i_ean = cab.index('EAN')
    col = len(cab) + 1
    ws.cell(row=1, column=col, value=COLUMNA_NO_HABRA_MAS).font = Font(bold=True)
    ws.column_dimensions[get_column_letter(col)].width = ANCHO_NO_HABRA_MAS
    fdv = {str(f['ean_original']) for f in foto if f.get('fin_de_vida')}
    n = 0
    for fila in range(2, ws.max_row + 1):
        if str(ws.cell(row=fila, column=i_ean + 1).value) in fdv:
            ws.cell(row=fila, column=col, value=TEXTO_NO_HABRA_MAS)
            n += 1
    for tabla in ws.tables.values():
        ini, fin = tabla.ref.split(':')
        tabla.ref = '%s:%s%s' % (ini, get_column_letter(col), re.sub(r'^[A-Z]+', '', fin))
        if tabla.tableColumns:
            tabla.tableColumns.append(TableColumn(id=len(tabla.tableColumns) + 1, name=COLUMNA_NO_HABRA_MAS))
        if tabla.autoFilter is not None:
            tabla.autoFilter.ref = tabla.ref
    return n


COLUMNAS_PUERTAS = ['EAN', 'Nombre', 'Marca', 'Precio compra (€)', 'Puerta', 'Motivo', 'Detalle', 'Caja',
                    'Precio caja (€)', 'EAN de la figura', 'Origen del EAN', 'Aviso del EAN']


def escribir_excel(foto, resultados, apartados, M, info):
    """El Excel del cruce de OSMA: DELANTE las hojas del viejo (`excel_como_el_viejo_osma`), DETRAS las del escaner 2
    con los mismos nombres y columnas que el de HEO (Resumen, Comparación, Varias fichas, Puertas y Puertas previas).
    «Comparación» va con su cabecera y sin filas: el escaner viejo no tiene escaneos de OSMA con que comparar
    (escaner_resultados, medido el 02-oct-2026). Solo lo calculado: aqui no se decide nada. → bytes.
    (Mapa de OSMA) `info['mapa']`: las fichas del mapa valoradas (`valorar_mapa`), que entran en «Análisis» (las que se
    venden) y en «Puertas» (todas) con la marca al final; `info['agotados']`: las filas de la hoja «Nuestros agotados en
    OSMA» (`agotados_nuestros`), la ULTIMA del libro."""
    from openpyxl.styles import Font
    por_foto = {f['id']: f for f in foto}
    mapa = info.get('mapa') or []
    asins_mapa = {r['asin_mapa'] for r in mapa}
    foto_mapa, res_mapa = filas_mapa_excel(mapa)
    foto_excel = foto_para_excel(foto, resultados, info['enlaces']) + foto_mapa
    wb = excel_como_el_viejo_osma(foto_excel, resultados, apartados, M, eleccion=info.get('eleccion'),
                                  paises=paises_de_la_hoja(info['params'], M), mapa=(res_mapa, asins_mapa))
    sustituidas = [r for r in resultados if r['puerta'] in e2.PUERTAS_ANALISIS and r.get('asin') in asins_mapa]

    def hoja(nombre, cabecera, filas, anchos=None):
        ws = wb.create_sheet(nombre)
        ws.append(cabecera)
        for c in ws[1]:
            c.font = Font(bold=True)
        for fila in filas:
            ws.append(fila)
        for letra, ancho in (anchos or {}).items():
            ws.column_dimensions[letra].width = ancho
        ws.freeze_panes = 'A2'
        return ws

    p = info['params']
    po = info.get('porte') or {}
    # (AM) Arriba, lo que se puede tocar: el pedido y el porte (celdas amarillas, con nombre), y lo que cambia con
    # ellos. Las filas 2 y 3 son las de CELDA_PEDIDO y CELDA_PORTE.
    an = wb['Análisis']
    cab_an = [c.value for c in an[1]]
    from openpyxl.utils import get_column_letter
    l_pais, l_dec = (get_column_letter(cab_an.index(n) + 1) for n in ('País', 'Decisión'))
    pais_cuenta = 'ES'
    # (Mapa de OSMA) Lo que pinta «Análisis»: las filas de siempre menos las que sustituye el mapa, y las del mapa.
    en_hoja = [r for r in resultados if r['puerta'] in e2.PUERTAS_ANALISIS and r.get('asin') not in asins_mapa] + res_mapa
    pasada_dec = {d: sum(1 for r in en_hoja if ((r.get('paises') or {}).get(pais_cuenta) or {}).get('decision') == d)
                  for d in ('COMPRAR', 'VALORAR')}
    vivas = {d: "=COUNTIFS('Análisis'!$%s:$%s,\"%s\",'Análisis'!$%s:$%s,\"%s\")"
                % (l_pais, l_pais, pais_cuenta, l_dec, l_dec, d) for d in ('COMPRAR', 'VALORAR')}
    pk = [r.get('pack_amazon') or {} for r in en_hoja]
    n_mult = sum(1 for r in en_hoja if r.get('pack') == 'amazon')
    n_dud_val = sum(1 for r, x in zip(en_hoja, pk) if x.get('estado') == PACK_DUDOSO and r['puerta'] == 'e')
    n_dud_no = sum(1 for r, x in zip(en_hoja, pk) if x.get('estado') == PACK_DUDOSO and r['puerta'] != 'e')
    filas_vivas = [
        ['Pedido previsto (€)', po.get('pedido_previsto')],
        ['Porte del envío (€)', po.get('gastos_envio')],
        ['Porte en el precio', '=IF(N(%s)>0,N(%s)/%s,0)' % (NOMBRE_PEDIDO, NOMBRE_PORTE, NOMBRE_PEDIDO)],
        ['Nota', 'Si el pedido no cabe en un palé, pon aquí el porte de los palés que sean. Al cambiar estas dos '
                 'celdas cambian «PA (€)», el beneficio y la «Decisión» de «Análisis», y estos recuentos.'],
        ['Recuento en %s' % pais_cuenta, 'con el pedido de arriba (fórmula)', 'en la pasada'],
        ['COMPRAR', vivas['COMPRAR'], pasada_dec['COMPRAR']],
        ['VALORAR', vivas['VALORAR'], pasada_dec['VALORAR']],
        ['Puertas y cuadre', 'con el pedido de la pasada'],
        ['Packs de Amazon detectados', '%d multiplicados · %d dudosos a VALORAR%s' % (
            n_mult, n_dud_val, (' · %d dudosos sin margen (NO COMPRAR)' % n_dud_no) if n_dud_no else '')]]
    filas_res = filas_vivas + [['Pasada', info['pasada']], ['Cruce', info['cruce']],
                 ['Proveedor', 'OSMA · artículos con marca y existencias'],
                 ['Descarga de OSMA', '%s (pasada %s)' % (info.get('descarga_en') or '—', info.get('disp_pasada') or '—')],
                 ['Porte en el precio (el de la puerta común, una sola vez)',
                  ('%s € de la factura %s (%s) ÷ pedido previsto %s € = %s'
                   % (_coma(po.get('gastos_envio')), po.get('factura') or '—', po.get('fecha') or '—',
                      _coma(po.get('pedido_previsto')), texto_pct(_dec(po['pct']) if po.get('pct') is not None else None)))
                  if po.get('pct') is not None else 'sin porte (sin pedido previsto o sin factura con porte)'],
                 ['Umbral de caídas (30 días)', '%s (> %d)' % (e2.texto_corte(p['umbral']), p['umbral'])],
                 ['Fichas compartidas (productos de la lista, por país)',
                  ' · '.join('%s %d' % (k, n) for k, n in (info.get('n_compartidas') or {}).items()) or '—'],
                 ['Países del filtro de ventas', ', '.join(p['paises_filtro'])],
                 ['Países que se calculan (si traen CSV)', ', '.join(p['paises_calculo'])],
                 ['Países con CSV', ', '.join(info['usados'])],
                 ['Códigos enlazados en la foto (el EAN de nuestra ficha va en la lista)', '%d códigos · %d EAN de '
                  'nuestras fichas (%d que no traía ya OSMA)'
                  % (len(info.get('enlaces_foto') or []), len(info.get('eans_fichas') or {}),
                     info.get('n_eans_fichas_nuevos') or 0)],
                 ['Fichas nuestras que quedan fuera: Keepa no conoce su EAN (no se rescatan)',
                  texto_fuera_de_keepa(info.get('fuera_keepa'))],
                 ['Packs nuestros valorados como unidades × precio', info.get('n_packs', 0)],
                 ['Fichas del mapa (código OSMA → ASIN → unidades), de códigos disponibles',
                  '%d valoradas (%d por el CSV de ASIN, %d por el de EAN, %d sin dato) · %d en «Análisis» · %d filas '
                  'normales sustituidas por la del mapa' % (
                      len(mapa), sum(1 for r in mapa if r.get('fuente') == 'asin'),
                      sum(1 for r in mapa if r.get('fuente') == 'ean'), sum(1 for r in mapa if not r.get('fuente')),
                      len(res_mapa), len(sustituidas))],
                 [HOJA_AGOTADOS, '%d (hoja al final del libro)' % len(info.get('agotados') or [])],
                 ['Catálogo de OSMA (descarga)', info['n_crudo']]]
    filas_res += [['Puerta previa · %s' % NOMBRE_PREVIA[x], info['previas'][x]] for x in e2.PUERTAS_PREVIAS]
    filas_res += [['Entradas (filas de la foto)', info['n_entradas']]]
    filas_res += [['Puerta %s · %s' % (x, e2.NOMBRE_PUERTA[x]), info['n_bd'][x]] for x in e2.PUERTAS]
    filas_res += [['Puertas previas + suma de puertas', sum(info['previas'].values()) + sum(info['n_bd'].values())],
                  ['Cuadra', 'SÍ' if info['cuadra'] else 'NO'],
                  ['Comparación con el escáner viejo', 'no se hace: el viejo no tiene escaneos de OSMA']]
    filas_res += [['CSV', '%s · %s · %s filas%s%s' % (
        f['nombre'], f.get('pais') or '¿?', f.get('filas') or 0,
        '' if f.get('usado') else ' · NO USADO', (' · ' + f['error']) if f.get('error') else '')]
        for f in info['ficheros']]
    filas_res += [['Aviso', a] for a in info['avisos']]
    wr = hoja('Resumen', ['Qué', 'Valor'], filas_res, {'A': 44, 'B': 90, 'C': 14})
    from openpyxl.styles import PatternFill
    amarillo = PatternFill(start_color='FFFF00', end_color='FFFF00', fill_type='solid')
    for celda in (CELDA_PEDIDO, CELDA_PORTE):
        wr[celda.replace('$', '')].fill = amarillo
        wr[celda.replace('$', '')].number_format = '#,##0.00'
        wr['A' + celda.split('$')[-1]].font = Font(bold=True)
    wr['B4'].number_format = '0.00%'
    poner_celdas_pedido(wb, wr)

    hoja('Comparación', ['EAN', 'Nombre', 'Categoría', 'Viejo', 'País viejo', 'Fecha viejo (UTC)', 'Excel viejo',
                         'Puerta nuevo', 'Nuevo', 'País nuevo', 'Fecha nuevo (UTC)', 'Diferencia de criterio',
                         'Explicación'], [], {'A': 15, 'B': 50, 'C': 16, 'G': 36, 'M': 110})

    varias = []
    for res in resultados:
        if res['puerta'] == 'b':
            f = por_foto[res['foto_id']]
            for fi in res['fichas']:
                varias.append([f['ean_original'], f['nombre'], fi['pais'], fi['asin'], fi['titulo'], fi['rank'],
                               fi['caidas_30d'], fi['precio_venta']])
    hoja('Varias fichas', ['EAN', 'Nombre OSMA', 'País', 'ASIN', 'Título Amazon', 'Puesto', 'Caídas 30 d',
                           'Precio venta (€)'], varias, {'A': 15, 'B': 50, 'E': 60})

    # (Mapa de OSMA) Detras de las de siempre, TODAS las fichas del mapa, con su marca en una columna nueva al final.
    hoja('Puertas', COLUMNAS_PUERTAS + [COLUMNA_MAPA],
         [[por_foto[r['foto_id']]['ean_original'], por_foto[r['foto_id']]['nombre'], por_foto[r['foto_id']]['marca'],
           r.get('pa', por_foto[r['foto_id']]['precio_unidad']), '%s · %s' % (r['puerta'], e2.NOMBRE_PUERTA[r['puerta']]),
           r['motivo'], r['detalle'], None, None, None, None, None, None] for r in resultados]
         + [[r['ean_original'] or r['asin_mapa'], r['nombre_osma'], None, r['pa'],
             '%s · %s' % (r['puerta'], e2.NOMBRE_PUERTA[r['puerta']]), r['motivo'], r['detalle'], None, None, None, None,
             None, '%s · %s' % (r['marca_mapa'], r['asin_mapa'])] for r in mapa],
         {'A': 15, 'B': 55, 'C': 16, 'E': 24, 'F': 16, 'G': 70, 'H': 26, 'M': 36})

    hoja('Puertas previas', ['EAN tal como vino', 'Nombre', 'Marca', 'Precio catálogo (€)', 'Puerta previa', 'Detalle'],
         [[a['ean_original'], a['nombre'], a['marca'], a['precio_catalogo'],
           NOMBRE_PREVIA.get(a['motivo'], a['motivo']), a['detalle']] for a in apartados],
         {'A': 18, 'B': 55, 'C': 16, 'E': 32, 'F': 80})

    # (Mapa de OSMA, Fernando, 04-oct-2026) La ULTIMA hoja: lo nuestro que OSMA tiene agotado hoy, que no se esconda.
    hoja(HOJA_AGOTADOS, COLUMNAS_AGOTADOS, info.get('agotados') or [], {'A': 12, 'B': 50, 'C': 14, 'D': 50, 'G': 18})
    salida = io.BytesIO()
    wb.save(salida)
    return salida.getvalue()


def ruta_excel(pasada, cruce, sello):
    """`osma/<pasada>/<cruce>/Escaner2_OSMA_<AAAAMMDD_HHMM>.xlsx`: la forma que admite escaner2_cruce.ruta_excel
    desde la migracion del encargo AG (v2)."""
    return '%s/%s/%s/Escaner2_OSMA_%s.xlsx' % (CARPETA, pasada, cruce, sello)

